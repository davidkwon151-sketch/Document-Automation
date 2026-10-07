"""Hand-built ZIP/XML expansions and intentional corruption, without fillers."""
from copy import deepcopy
from hashlib import sha256
from pathlib import Path
from zipfile import ZipFile, ZipInfo, ZIP_STORED

from docx import Document
from docx.oxml import OxmlElement
from lxml import etree
import pytest

from agent.repeat_check import verify_repeat_expansion, W, W14, HP, NS


def parts(path):
    with ZipFile(path) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def write(path, contents):
    with ZipFile(path, 'w') as archive:
        for name, data in contents.items():
            archive.writestr(name, data, compress_type=ZIP_STORED)
    return path


def docx_source(tmp_path):
    path = tmp_path / 'original.docx'
    doc = Document()
    doc.add_paragraph('원본 안내문')
    table = doc.add_table(rows=3, cols=2)
    table.style = 'Table Grid'
    table.cell(0, 0).text = '품목'; table.cell(0, 1).text = '수량'
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    table.cell(1, 0).text = '{{품목}}'; table.cell(1, 1).text = '{{수량}}'
    table.cell(1, 0).paragraphs[0].runs[0].bold = True
    table.cell(2, 0).text = '합계'; table.cell(2, 1).text = '{{합계}}'
    doc.add_table(rows=1, cols=1).cell(0, 0).text = '다른 표'
    doc.save(path)
    contents = parts(path)
    root = etree.fromstring(contents['word/document.xml'])
    prototype = root.find('.//w:tbl/w:tr[2]', NS)
    paragraph = prototype.find('.//w:p', NS)
    paragraph.set(f'{{{W14}}}paraId', '00000001')
    paragraph.set(f'{{{W14}}}textId', '00000002')
    # One native unselected choice preserves both stored and displayed options.
    cell = prototype.find('w:tc', NS)
    old = cell.find('w:p', NS)
    sdt = etree.Element(f'{{{W}}}sdt')
    properties = etree.SubElement(sdt, f'{{{W}}}sdtPr')
    etree.SubElement(properties, f'{{{W}}}id').set(f'{{{W}}}val', '17')
    etree.SubElement(properties, f'{{{W}}}showingPlcHdr')
    choice = etree.SubElement(properties, f'{{{W}}}dropDownList')
    for value, label in [('EXP', '수출'), ('DOM', '국내')]:
        item = etree.SubElement(choice, f'{{{W}}}listItem')
        item.set(f'{{{W}}}value', value); item.set(f'{{{W}}}displayText', label)
    content = etree.SubElement(sdt, f'{{{W}}}sdtContent')
    cell.remove(old); content.append(old); cell.append(sdt)
    contents['word/document.xml'] = etree.tostring(root)
    write(path, contents)
    table = root.find('.//w:tbl', NS)
    plan = {'table_id': 'docx:word/document.xml:' + root.getroottree().getpath(table), 'row': 2, 'count': 3}
    return path, contents, plan


def hwpx_source(tmp_path):
    path = tmp_path / 'original.hwpx'
    root = etree.Element('{http://www.hancom.co.kr/hwpml/2011/section}sec', nsmap={'hs': 'http://www.hancom.co.kr/hwpml/2011/section', 'hp': HP})
    def paragraph(parent, identifier, text):
        p = etree.SubElement(parent, f'{{{HP}}}p', id=str(identifier), paraPrIDRef='0', styleIDRef='0')
        run = etree.SubElement(p, f'{{{HP}}}run', charPrIDRef='0')
        etree.SubElement(run, f'{{{HP}}}t').text = text
        cache = etree.SubElement(p, f'{{{HP}}}linesegarray')
        etree.SubElement(cache, f'{{{HP}}}lineseg', vertpos='150', vertsize='200')
        return p, run
    paragraph(root, 9, '이전 안내 유지')
    anchor, run = paragraph(root, 10, '')
    table = etree.SubElement(run, f'{{{HP}}}tbl', id='80', rowCnt='3', colCnt='2', cellSpacing='5', lock='0')
    etree.SubElement(table, f'{{{HP}}}sz', width='2000', height='1000', heightRelTo='ABSOLUTE', protect='0')
    etree.SubElement(table, f'{{{HP}}}inMargin', top='10', bottom='20', left='0', right='0')
    for row_index, values in enumerate([['품목', '수량'], ['{{품목}}', '{{수량}}'], ['합계', '안내']]):
        row = etree.SubElement(table, f'{{{HP}}}tr')
        for column, text in enumerate(values):
            cell = etree.SubElement(row, f'{{{HP}}}tc', header='1' if row_index == 0 else '0', protect='0', name='', hasMargin='1' if column == 0 else '0', borderFillIDRef='0')
            sub = etree.SubElement(cell, f'{{{HP}}}subList', id='5' if column == 0 else '', hasTextRef='0', hasNumRef='0', linkListIDRef='0')
            p, _ = paragraph(sub, 20 + row_index * 2 + column, text)
            etree.SubElement(cell, f'{{{HP}}}cellAddr', rowAddr=str(row_index), colAddr=str(column))
            etree.SubElement(cell, f'{{{HP}}}cellSpan', rowSpan='1', colSpan='1')
            etree.SubElement(cell, f'{{{HP}}}cellSz', width='1000', height='100')
            etree.SubElement(cell, f'{{{HP}}}cellMargin', top='7', bottom='8', left='0', right='0')
    paragraph(root, 99, '뒤쪽 안내 유지')
    contents = {'mimetype': b'application/hwp+zip', 'Contents/section0.xml': etree.tostring(root),
                'Contents/header.xml': b'<styles id="50"/>', 'Preview/PrvText.txt': b'preserved preview'}
    write(path, contents)
    plan = {'table_id': 'hwpx:Contents/section0.xml:' + root.getroottree().getpath(table), 'row': 2, 'count': 3}
    return path, contents, plan


def hand_expansion(tmp_path, source_factory, count=3):
    original, contents, plan = source_factory(tmp_path)
    plan['count'] = count
    kind, part, xpath = plan['table_id'].split(':', 2)
    root = etree.fromstring(contents[part]); table = root.xpath(xpath, namespaces=root.nsmap)[0]
    namespace = W if kind == 'docx' else HP
    rows = table.findall(f'{{{namespace}}}tr'); prototype = rows[1]
    position = table.index(prototype)
    next_id = 2000
    for offset in range(1, count):
        clone = deepcopy(prototype)
        for node in clone.iter():
            if kind == 'docx':
                for attribute in [f'{{{W14}}}paraId', f'{{{W14}}}textId']:
                    if attribute in node.attrib:
                        node.set(attribute, f'{next_id:08X}'); next_id += 1
                if node.tag == f'{{{W}}}id' and node.getparent().tag == f'{{{W}}}sdtPr':
                    node.set(f'{{{W}}}val', str(next_id)); next_id += 1
            else:
                if node.tag in {f'{{{HP}}}p', f'{{{HP}}}subList'} and node.get('id', ''):
                    node.set('id', str(next_id)); next_id += 1
                if node.tag == f'{{{HP}}}cellAddr':
                    node.set('rowAddr', str(1 + offset))
        table.insert(position + offset, clone)
    if kind == 'hwpx' and count > 1:
        for row in rows[2:]:
            for address in row.findall(f'{{{HP}}}tc/{{{HP}}}cellAddr'):
                address.set('rowAddr', str(int(address.get('rowAddr')) + count - 1))
        table.set('rowCnt', str(3 + count - 1))
        # Source cell #2 uses table margin 30: 150+200+30=380 >365.
        table.find(f'{{{HP}}}sz').set('height', str(1000 + (count - 1) * 385))
        for child in list(root)[1:]:
            for cache in child.findall(f'{{{HP}}}linesegarray'):
                child.remove(cache)
    contents[part] = etree.tostring(root)
    prepared = write(tmp_path / ('prepared.' + kind), contents)
    return original, prepared, plan


def mutate(prepared, plan, change):
    contents = parts(prepared)
    kind, part, path = plan['table_id'].split(':', 2)
    root = etree.fromstring(contents[part]); table = root.xpath(path, namespaces=root.nsmap)[0]
    change(root, table)
    contents[part] = etree.tostring(root)
    write(prepared, contents)


@pytest.mark.parametrize('factory', [docx_source, hwpx_source])
@pytest.mark.parametrize('count', [1, 2, 3, 200])
def test_independent_hand_expansion_preserves_original_and_plan(tmp_path, factory, count):
    original, prepared, plan = hand_expansion(tmp_path, factory, count)
    before = original.read_bytes(); expected = deepcopy(plan)
    result = verify_repeat_expansion(original, prepared, plan)
    assert result['status'] == 'passed' and result['native_visual_qa'] == 'pending'
    assert result['rows'] == {'original': 3, 'prepared': 3 + count - 1, 'added': count - 1}
    assert result['original_sha256'] == sha256(before).hexdigest()
    assert result['prepared_sha256'] == sha256(prepared.read_bytes()).hexdigest()
    assert original.read_bytes() == before and plan == expected


@pytest.mark.parametrize('factory', [docx_source, hwpx_source])
@pytest.mark.parametrize('which', ['prototype', 'clone', 'footer', 'outside', 'delete', 'duplicate'])
def test_corrupted_rows_and_unrelated_text_are_blocked(tmp_path, factory, which):
    original, prepared, plan = hand_expansion(tmp_path, factory)
    namespace = W if factory is docx_source else HP
    def change(root, table):
        rows = table.findall(f'{{{namespace}}}tr')
        if which in {'prototype', 'clone', 'footer'}:
            row = rows[1 if which == 'prototype' else 2 if which == 'clone' else -1]
            row.find(f'.//{{{namespace}}}t').text = '변조된 텍스트'
        elif which == 'outside':
            root.find(f'.//{{{namespace}}}t').text = '원본 안내 삭제'
        elif which == 'delete':
            table.remove(rows[-1])
        else:
            table.append(deepcopy(rows[2]))
    mutate(prepared, plan, change)
    with pytest.raises(ValueError):
        verify_repeat_expansion(original, prepared, plan)


@pytest.mark.parametrize('factory', [docx_source, hwpx_source])
@pytest.mark.parametrize('operation', ['changed', 'added', 'removed'])
def test_other_zip_parts_must_be_identical(tmp_path, factory, operation):
    original, prepared, plan = hand_expansion(tmp_path, factory)
    contents = parts(prepared)
    if operation == 'changed':
        name = 'word/styles.xml' if factory is docx_source else 'Preview/PrvText.txt'
        contents[name] += b'changed'
    elif operation == 'added':
        contents['unexpected.xml'] = b'<extra/>'
    else:
        del contents['word/styles.xml' if factory is docx_source else 'Contents/header.xml']
    write(prepared, contents)
    with pytest.raises(ValueError, match='부품'):
        verify_repeat_expansion(original, prepared, plan)


@pytest.mark.parametrize('attribute', ['paraId', 'textId', 'sdt'])
@pytest.mark.parametrize('value', ['1', 'bad', '00000000', '80000000'])
def test_docx_clone_id_collision_format_range_blocked(tmp_path, attribute, value):
    original, prepared, plan = hand_expansion(tmp_path, docx_source)
    def change(root, table):
        clone = table.findall(f'{{{W}}}tr')[2]
        if attribute == 'sdt':
            node = clone.find('.//w:sdtPr/w:id', NS)
            node.set(f'{{{W}}}val', '17' if value == '1' else '2147483648' if value == '80000000' else value)
        else:
            clone.find('.//w:p', NS).set(f'{{{W14}}}{attribute}', value)
    mutate(prepared, plan, change)
    with pytest.raises(ValueError, match='ID'):
        verify_repeat_expansion(original, prepared, plan)


@pytest.mark.parametrize('which', ['font', 'merge', 'option', 'state', 'other_table'])
def test_docx_clone_properties_and_native_choices_cannot_change(tmp_path, which):
    original, prepared, plan = hand_expansion(tmp_path, docx_source)
    def change(root, table):
        clone = table.findall(f'{{{W}}}tr')[2]
        if which == 'font':
            clone.find('.//w:rPr/w:b', NS).set(f'{{{W}}}val', '0')
        elif which == 'merge':
            props = clone.find('w:tc/w:tcPr', NS)
            etree.SubElement(props, f'{{{W}}}gridSpan').set(f'{{{W}}}val', '2')
        elif which == 'option':
            clone.find('.//w:listItem', NS).set(f'{{{W}}}value', 'OTHER')
        elif which == 'state':
            clone.find('.//w:sdtPr', NS).remove(clone.find('.//w:showingPlcHdr', NS))
        else:
            root.findall('.//w:tbl', NS)[1].find('.//w:t', NS).text = '다른 표 변경'
    mutate(prepared, plan, change)
    with pytest.raises(ValueError):
        verify_repeat_expansion(original, prepared, plan)


@pytest.mark.parametrize('which', ['rowCnt', 'clone_addr', 'footer_addr', 'height', 'id_collision', 'font', 'merge', 'inner_cache', 'before_cache', 'after_cache'])
def test_hwpx_only_declared_addresses_ids_height_and_caches_may_change(tmp_path, which):
    original, prepared, plan = hand_expansion(tmp_path, hwpx_source)
    def change(root, table):
        clone = table.findall(f'{{{HP}}}tr')[2]
        if which == 'rowCnt': table.set('rowCnt', '99')
        elif which == 'clone_addr': clone.find('.//hp:cellAddr', NS).set('rowAddr', '1')
        elif which == 'footer_addr': table.findall(f'{{{HP}}}tr')[-1].find('.//hp:cellAddr', NS).set('rowAddr', '2')
        elif which == 'height': table.find('hp:sz', NS).set('height', '1000')
        elif which == 'id_collision': clone.find('.//hp:p', NS).set('id', '50')
        elif which == 'font': clone.find('.//hp:run', NS).set('charPrIDRef', '9')
        elif which == 'merge': clone.find('.//hp:cellSpan', NS).set('rowSpan', '2')
        elif which == 'inner_cache': clone.find('.//hp:p', NS).remove(clone.find('.//hp:linesegarray', NS))
        elif which == 'before_cache': root[0].remove(root[0].find('hp:linesegarray', NS))
        elif which == 'after_cache': etree.SubElement(root[-1], f'{{{HP}}}linesegarray')
    mutate(prepared, plan, change)
    with pytest.raises(ValueError):
        verify_repeat_expansion(original, prepared, plan)


@pytest.mark.parametrize('patch', [{'row': True}, {'count': True}, {'row': 0}, {'count': 0}, {'count': 201}, {'row': '2'}, {'count': 1.0}, {'extra': 1}, {'table_id': None}, {'table_id': 'docx:word/document.xml://w:tbl'}, {'table_id': 'docx:../word/document.xml:/w:document'}])
def test_invalid_plan_is_not_interpreted(tmp_path, patch):
    original, prepared, plan = hand_expansion(tmp_path, docx_source)
    plan.update(patch)
    with pytest.raises(ValueError): verify_repeat_expansion(original, prepared, plan)


@pytest.mark.parametrize('factory', [docx_source, hwpx_source])
def test_same_path_and_hard_link_are_rejected(tmp_path, factory):
    original, prepared, plan = hand_expansion(tmp_path, factory)
    with pytest.raises(ValueError, match='같음'): verify_repeat_expansion(original, original, plan)
    linked = tmp_path / ('linked' + original.suffix)
    linked.hardlink_to(original)
    with pytest.raises(ValueError, match='같음'): verify_repeat_expansion(original, linked, plan)


@pytest.mark.parametrize('bad', ['../outside.xml', '/absolute.xml', 'word\\bad.xml', 'C:bad.xml'])
def test_zip_path_guard(tmp_path, bad):
    original, prepared, plan = hand_expansion(tmp_path, docx_source)
    # ZipInfo normally normalizes Windows separators before writing; construct
    # the actual malformed archive name instead of testing a normalized name.
    with ZipFile(prepared, 'a') as archive:
        info = ZipInfo('placeholder.xml'); info.filename = bad
        archive.writestr(info, b'<x/>')
    with pytest.raises(ValueError, match='경로'): verify_repeat_expansion(original, prepared, plan)


def test_duplicate_zip_entries_and_dtd_are_rejected(tmp_path):
    original, prepared, plan = hand_expansion(tmp_path, docx_source)
    with ZipFile(prepared, 'a') as archive:
        with pytest.warns(UserWarning): archive.writestr('word/document.xml', b'<duplicate/>')
    with pytest.raises(ValueError, match='중복'): verify_repeat_expansion(original, prepared, plan)
    contents = parts(original); contents['word/document.xml'] = b'<!DOCTYPE x [<!ENTITY secret SYSTEM "file:///secret">]><x>&secret;</x>'
    write(prepared, contents)
    with pytest.raises(ValueError, match='DTD'): verify_repeat_expansion(original, prepared, plan)


def test_64_mib_guard_is_enforced_without_allocating_a_large_file(tmp_path, monkeypatch):
    original, prepared, plan = hand_expansion(tmp_path, docx_source)
    import agent.repeat_check as checker
    monkeypatch.setattr(checker, 'MAX_BYTES', 100)
    with pytest.raises(ValueError, match='64MiB'): verify_repeat_expansion(original, prepared, plan)


def test_checker_does_not_call_transformers(tmp_path, monkeypatch):
    from templates import repeat_docx, repeat_hwpx
    def forbidden(*args, **kwargs): raise AssertionError('Transformer cannot certify itself')
    monkeypatch.setattr(repeat_docx, 'transform', forbidden)
    monkeypatch.setattr(repeat_hwpx, 'transform', forbidden)
    for index, factory in enumerate([docx_source, hwpx_source]):
        directory = tmp_path / str(index); directory.mkdir()
        original, prepared, plan = hand_expansion(directory, factory)
        assert verify_repeat_expansion(original, prepared, plan)['status'] == 'passed'


def patch_both(original, prepared, plan, callback):
    for path in [original, prepared]:
        contents = parts(path)
        kind, part, xpath = plan['table_id'].split(':', 2)
        root = etree.fromstring(contents[part]); table = root.xpath(xpath, namespaces=root.nsmap)[0]
        callback(root, table)
        contents[part] = etree.tostring(root); write(path, contents)


@pytest.mark.parametrize('which', ['protected', 'header', 'total', 'vertical', 'next_continue', 'bookmark', 'binding', 'lock', 'unknown_id', 'signature', 'checkbox_missing', 'checkbox_selected', 'bad_options'])
def test_docx_ineligible_original_is_not_certified_even_if_protection_preserved(tmp_path, which):
    original, prepared, plan = hand_expansion(tmp_path, docx_source)
    if which in {'header', 'total'}:
        plan['row'] = 1 if which == 'header' else 3
    elif which in {'protected', 'signature'}:
        for path in [original, prepared]:
            contents = parts(path)
            if which == 'protected':
                settings = etree.fromstring(contents['word/settings.xml'])
                node = etree.SubElement(settings, f'{{{W}}}documentProtection')
                node.set(f'{{{W}}}enforcement', '1'); node.set(f'{{{W}}}edit', 'forms')
                contents['word/settings.xml'] = etree.tostring(settings)
            else: contents['_xmlsignatures/sig1.xml'] = b'<signature/>'
            write(path, contents)
    else:
        def change(root, table):
            prototype = table.findall(f'{{{W}}}tr')[1]
            if which in {'vertical', 'next_continue'}:
                row = prototype if which == 'vertical' else table.findall(f'{{{W}}}tr')[2]
                props = row.find('w:tc/w:tcPr', NS)
                etree.SubElement(props, f'{{{W}}}vMerge').set(f'{{{W}}}val', 'restart' if which == 'vertical' else 'continue')
            elif which == 'bookmark':
                etree.SubElement(prototype.find('.//w:p', NS), f'{{{W}}}bookmarkStart').set(f'{{{W}}}id', '500')
            elif which in {'binding', 'lock'}:
                etree.SubElement(prototype.find('.//w:sdtPr', NS), f'{{{W}}}' + ('dataBinding' if which == 'binding' else 'lock'))
            elif which == 'unknown_id': prototype.set('{urn:unknown}objectId', '42')
            elif which == 'bad_options':
                options = prototype.findall('.//w:listItem', NS); options[1].set(f'{{{W}}}value', options[0].get(f'{{{W}}}value'))
            else:
                properties = prototype.find('.//w:sdtPr', NS)
                box = etree.SubElement(properties, f'{{{W14}}}checkbox')
                if which == 'checkbox_selected': etree.SubElement(box, f'{{{W14}}}checked').set(f'{{{W14}}}val', '1')
        patch_both(original, prepared, plan, change)
    with pytest.raises(ValueError): verify_repeat_expansion(original, prepared, plan)


@pytest.mark.parametrize('which', ['protected', 'header', 'total', 'rowspan', 'cellname', 'linkref', 'pagebreak', 'tracking', 'caption', 'zone', 'relative_height', 'security'])
def test_hwpx_ineligible_original_is_not_certified(tmp_path, which):
    original, prepared, plan = hand_expansion(tmp_path, hwpx_source)
    if which in {'header', 'total'}:
        plan['row'] = 1 if which == 'header' else 3
    elif which == 'security':
        for path in [original, prepared]:
            contents = parts(path); contents['Contents/security.xml'] = b'<docSecurity/>'; write(path, contents)
    else:
        def change(root, table):
            prototype = table.findall(f'{{{HP}}}tr')[1]; cell = prototype.find('hp:tc', NS)
            if which == 'protected': cell.set('protect', '1')
            elif which == 'rowspan': cell.find('hp:cellSpan', NS).set('rowSpan', '2')
            elif which == 'cellname': cell.set('name', 'reference')
            elif which == 'linkref': cell.find('hp:subList', NS).set('linkListIDRef', '99')
            elif which == 'pagebreak': cell.find('.//hp:p', NS).set('pageBreak', '1')
            elif which == 'tracking': cell.find('.//hp:p', NS).set('paraTcId', '0')
            elif which == 'caption': etree.SubElement(table, f'{{{HP}}}caption')
            elif which == 'zone': etree.SubElement(table, f'{{{HP}}}cellzoneList')
            elif which == 'relative_height': table.find('hp:sz', NS).set('heightRelTo', 'PERCENT')
        patch_both(original, prepared, plan, change)
    with pytest.raises(ValueError): verify_repeat_expansion(original, prepared, plan)


def test_cross_part_id_collision_is_detected(tmp_path):
    original, prepared, plan = hand_expansion(tmp_path, docx_source)
    for path in [original, prepared]:
        contents = parts(path)
        contents['word/header9.xml'] = f'<w:hdr xmlns:w="{W}" xmlns:w14="{W14}"><w:p w14:paraId="000007D1"/></w:hdr>'.encode()
        write(path, contents)
    with pytest.raises(ValueError, match='충돌'): verify_repeat_expansion(original, prepared, plan)


@pytest.mark.parametrize('control', ['unchecked_checkbox', 'empty_combobox'])
def test_unselected_native_controls_are_preserved_in_copies(tmp_path, control):
    original, prepared, plan = hand_expansion(tmp_path, docx_source)
    for path in [original, prepared]:
        contents = parts(path); root = etree.fromstring(contents['word/document.xml'])
        for properties in root.findall('.//w:sdtPr', NS):
            choice = properties.find('w:dropDownList', NS)
            if control == 'empty_combobox':
                choice.tag = f'{{{W}}}comboBox'
                for option in list(choice): choice.remove(option)
            else:
                properties.remove(choice)
                box = etree.SubElement(properties, f'{{{W14}}}checkbox')
                etree.SubElement(box, f'{{{W14}}}checked').set(f'{{{W14}}}val', '0')
        contents['word/document.xml'] = etree.tostring(root); write(path, contents)
    assert verify_repeat_expansion(original, prepared, plan)['status'] == 'passed'
