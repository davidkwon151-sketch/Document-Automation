from copy import deepcopy
from hashlib import sha256
from io import BytesIO
from unittest.mock import Mock

from docx import Document
from openpyxl import Workbook
from PIL import Image
from pypdf import PdfReader, PdfWriter
from pypdf.constants import UserAccessPermissions
import pytest
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

from agent.multimodal_intake import (collect_multimodal, confirm_intake, generation_sources,
                                     validate_generation_source)
from llm.client import LLMError


def image_file(tmp_path, name='원본.png'):
    path = tmp_path / name
    Image.new('RGB', (120, 60), 'white').save(path)
    return path


def scan_file(tmp_path, pages=1):
    image = image_file(tmp_path)
    path = tmp_path / '스캔.pdf'
    pdf = canvas.Canvas(str(path), pagesize=(200, 200))
    for _ in range(pages):
        pdf.drawImage(ImageReader(str(image)), 10, 10, 150, 100)
        pdf.showPage()
    pdf.save()
    return path


def test_office_and_table_sources_share_full_provenance(tmp_path):
    txt = tmp_path / '지시.txt'
    txt.write_text('제품명: 시험정\n제형: 정제\n매출: 120만원', encoding='utf-8')
    csv = tmp_path / '거래.csv'
    csv.write_text('Order No.,품명,수량\nORD-01,시험정,10', encoding='utf-8')
    docx = tmp_path / '사양.docx'
    document = Document()
    document.add_paragraph('제품명: 시험정')
    document.save(docx)
    xlsx = tmp_path / '실적.xlsx'
    book = Workbook()
    book.active.title = '거래'
    book.active.append(['거래번호', '수량'])
    book.active.append(['ORD-01', 10])
    book.save(xlsx)
    result = collect_multimodal([txt, csv, docx, xlsx])
    assert all(file['status'] == 'ready' for file in result['files'])
    sources = generation_sources(result)
    assert sources and {source['filename'] for source in sources} == {path.name for path in [txt, csv, docx, xlsx]}
    for source in sources:
        assert source['id'] == source['source_id']
        assert source['quote'] == source['text']
        assert source['source_context'][source['context_start']:source['context_end']] == source['text']
        assert source['document_sha256'] == sha256((tmp_path / source['filename']).read_bytes()).hexdigest()
        assert source['source_location']['location'] == source['location']
        assert len(source['verification_fingerprint']) == 64
        assert source['page'] is None
    assert any(source['sheet'] == '거래' and 'ORD-01' in source['text'] for source in sources)
    assert result['confirmations'] == []


def test_image_model_is_opt_in_and_confirmation_binds_exact_text(tmp_path):
    path = image_file(tmp_path)
    client = Mock()
    client.read_image_json.return_value = {'본문': '제품명: 시험정\n보관조건: 2~8℃',
                                          '표 목록': [], '불확실한 항목': ['소문자 확인']}
    deferred = collect_multimodal([path], client=client)
    assert deferred['files'][0]['status'] == 'deferred'
    assert deferred['files'][0]['manual_transcription_available']
    assert not generation_sources(deferred)
    client.read_image_json.assert_not_called()
    result = collect_multimodal([path], client=client, allow_ocr=True)
    assert result['files'][0]['status'] == 'needs_confirmation'
    assert result['files'][0]['extraction_mode'] == 'image_reading'
    assert not generation_sources(result)
    receipts = [{'source_id': source['source_id'], 'fingerprint': source['verification_fingerprint']}
                for source in result['sources']]
    with pytest.raises(ValueError, match='명시적'):
        confirm_intake(result, receipts)
    confirmed = confirm_intake(result, receipts, confirmed=True)
    accepted = generation_sources(confirmed)
    assert len(accepted) == len(result['sources'])
    assert all(source['page'] == 1 and not source['requires_verification'] and not source['uncertain_items']
               for source in accepted)
    assert all(source['verification_receipt']['original_uncertain_items'] == ['소문자 확인']
               for source in accepted)
    assert all(source['requires_verification'] for source in result['sources'])
    assert result['confirmations'] == []
    assert client.read_image_json.call_count == 1


@pytest.mark.parametrize('field,value', [('text', '변조 999mg'), ('page', 2),
                                        ('document_sha256', '0' * 64), ('context_start', 2)])
def test_source_edit_invalidates_current_confirmation(tmp_path, field, value):
    path = image_file(tmp_path)
    digest = sha256(path.read_bytes()).hexdigest()
    result = collect_multimodal([path], transcriptions={digest: [{'page': 1, 'text': '제품명: 시험정'}]})
    source = result['sources'][0]
    result = confirm_intake(result, [{'source_id': source['source_id'],
                                     'fingerprint': source['verification_fingerprint']}], confirmed=True)
    result['sources'][0][field] = value
    with pytest.raises(ValueError, match='변경됨'):
        generation_sources(result)


def test_partial_confirmation_leaves_other_chunks_out(tmp_path):
    path = image_file(tmp_path)
    digest = sha256(path.read_bytes()).hexdigest()
    result = collect_multimodal([path], transcriptions={digest: [{'page': 1, 'text': '제품명: 시험정\n보관조건: 2~8℃'}]})
    source = result['sources'][0]
    checked = confirm_intake(result, [{'source_id': source['source_id'],
                                      'fingerprint': source['verification_fingerprint']}], confirmed=True)
    assert [row['source_id'] for row in generation_sources(checked)] == [source['source_id']]
    with pytest.raises(ValueError):
        confirm_intake(result, [{'source_id': source['source_id'], 'fingerprint': 'wrong'}], confirmed=True)
    corrupted = deepcopy(checked)
    corrupted['confirmations'][0]['document_sha256'] = '0' * 64
    with pytest.raises(ValueError, match='현재 원자료'):
        generation_sources(corrupted)
    corrupted = deepcopy(checked)
    corrupted['confirmations'].append(deepcopy(corrupted['confirmations'][0]))
    with pytest.raises(ValueError):
        generation_sources(corrupted)


def test_manual_image_transcription_has_original_hash_and_no_model_call(tmp_path):
    path = image_file(tmp_path)
    digest = sha256(path.read_bytes()).hexdigest()
    client = Mock()
    result = collect_multimodal([path], client=client, transcriptions={digest: [{'page': 1, 'text': '수량: 15개'}]})
    client.read_image_json.assert_not_called()
    assert result['files'][0]['extraction_mode'] == 'manual_transcription'
    assert result['sources'][0]['document_sha256'] == digest
    assert result['sources'][0]['source_context'] == '수량: 15개'
    assert not generation_sources(result)
    assert sha256(path.read_bytes()).hexdigest() == digest
    with pytest.raises(ValueError, match='연결되지 않은'):
        collect_multimodal([path], transcriptions={'0' * 64: [{'page': 1, 'text': '불일치'}]})


def test_all_scan_pages_need_explicit_transcription_and_confirmation(tmp_path):
    path = scan_file(tmp_path, 2)
    digest = sha256(path.read_bytes()).hexdigest()
    incomplete = collect_multimodal([path], transcriptions={digest: [{'page': 1, 'text': '제품명: 시험정'}]})
    assert incomplete['files'][0]['reason_code'] == 'invalid_transcription'
    assert not incomplete['sources']
    result = collect_multimodal([path], transcriptions={digest: [
        {'page': 2, 'text': '보관조건: 2~8℃'}, {'page': 1, 'text': '제품명: 시험정'}]})
    assert [source['page'] for source in result['sources']] == [1, 2]
    assert not generation_sources(result)


def test_extract_restricted_pdf_is_not_ocr_or_manual_transcription_bypassed(tmp_path):
    original = scan_file(tmp_path)
    protected = tmp_path / '보호.pdf'
    writer = PdfWriter(clone_from=original)
    writer.encrypt('', 'owner-secret', permissions_flag=UserAccessPermissions.PRINT)
    with protected.open('wb') as stream:
        writer.write(stream)
    digest = sha256(protected.read_bytes()).hexdigest()
    client = Mock()
    for transcriptions in (None, {digest: [{'page': 1, 'text': '제품명: 시험정'}]}):
        result = collect_multimodal([protected], client=client, allow_ocr=True,
                                    transcriptions=transcriptions)
        assert result['files'][0]['reason_code'] == 'protected_document'
        assert not result['files'][0]['manual_transcription_available']
        assert not generation_sources(result)
    client.read_image_json.assert_not_called()


def test_failed_image_model_is_deferred_without_provider_message_or_fallback(tmp_path):
    path = image_file(tmp_path)
    client = Mock()
    client.read_image_json.side_effect = LLMError('secret-provider-sk-example', kind='response')
    result = collect_multimodal([path], client=client, allow_ocr=True)
    assert result['files'][0]['reason_code'] == 'image_model_unavailable'
    assert 'secret-provider' not in str(result)
    assert not generation_sources(result)
    assert client.read_image_json.call_count == 1


def test_bad_image_and_unsupported_files_are_explicitly_deferred(tmp_path):
    image = tmp_path / '가짜.png'
    image.write_bytes(b'not an image')
    audio = tmp_path / '회의.wav'
    audio.write_bytes(b'fake audio')
    client = Mock()
    result = collect_multimodal([image, audio], client=client, allow_ocr=True)
    assert [file['reason_code'] for file in result['files']] == ['invalid_image', 'unsupported_format']
    assert not result['sources']
    client.read_image_json.assert_not_called()


def test_mixed_pdf_manual_transcription_preserves_original_text_page(tmp_path):
    image = image_file(tmp_path)
    path = tmp_path / '혼합.pdf'
    pdf = canvas.Canvas(str(path), pagesize=(200, 200))
    pdf.drawString(10, 150, 'Quantity: 10')
    pdf.showPage()
    pdf.drawImage(ImageReader(str(image)), 10, 10, 150, 100)
    pdf.save()
    with pytest.warns(UserWarning):
        deferred = collect_multimodal([path])
    assert deferred['files'][0]['transcription_pages'] == [2]
    assert not deferred['sources']
    digest = sha256(path.read_bytes()).hexdigest()
    result = collect_multimodal([path], transcriptions={digest: [{'page': 2, 'text': '제품명: 시험정'}]})
    assert [source['page'] for source in result['sources']] == [1, 2]
    assert [source['text'] for source in generation_sources(result)] == ['Quantity: 10']
    source = next(source for source in result['sources'] if source['page'] == 2)
    checked = confirm_intake(result, [{'source_id': source['source_id'],
                                      'fingerprint': source['verification_fingerprint']}], confirmed=True)
    assert [source['text'] for source in generation_sources(checked)] == ['Quantity: 10', '제품명: 시험정']
    assert sha256(path.read_bytes()).hexdigest() == digest


def test_single_generation_source_can_be_independently_revalidated(tmp_path):
    path = image_file(tmp_path)
    digest = sha256(path.read_bytes()).hexdigest()
    intake = collect_multimodal([path], transcriptions={digest: [{'page': 1, 'text': '수량: 10개'}]})
    source = intake['sources'][0]
    with pytest.raises(ValueError, match='확인이 완료'):
        validate_generation_source(source)
    intake = confirm_intake(intake, [{'source_id': source['source_id'],
                                     'fingerprint': source['verification_fingerprint']}], confirmed=True)
    accepted = generation_sources(intake)[0]
    assert validate_generation_source(accepted)
    for field, value in [('text', '수량: 11개'), ('page', 2), ('id', 'OTHER'),
                          ('context_text', '변조 원문'), ('document_sha256', '0' * 64)]:
        modified = {**accepted, field: value}
        with pytest.raises(ValueError):
            validate_generation_source(modified)
    no_receipt = deepcopy(accepted)
    no_receipt.pop('verification_receipt')
    with pytest.raises(ValueError):
        validate_generation_source(no_receipt)
    modified = deepcopy(accepted)
    modified['verification_receipt']['confirmed'] = False
    with pytest.raises(ValueError):
        validate_generation_source(modified)
    modified = deepcopy(accepted)
    modified['verification_receipt']['review_scope'] = '법정 제출 승인'
    with pytest.raises(ValueError):
        validate_generation_source(modified)


def test_confirmed_image_source_feeds_actual_ra_preparation_contract(tmp_path):
    from agent.ra_workflows import prepare_ra_workflow
    path = image_file(tmp_path)
    digest = sha256(path.read_bytes()).hexdigest()
    intake = collect_multimodal([path], transcriptions={digest: [{'page': 1, 'text': '제품명: SyntheticDrugA'}]})
    profile = {'domain': 'pharmaceutical_ra', 'ra_workflow': 'product_approval',
               'ra_product_name': 'SyntheticDrugA', 'document_kind': 'application',
               'fields': [{'id': 'drug', 'value_key': '제품명', 'label': '제품명',
                           'required': True, 'input_required': False, 'evidence_role': 'product_name'}]}
    assert not generation_sources(intake)
    original = intake['sources'][0]
    checked = confirm_intake(intake, [{'source_id': original['source_id'],
                                      'fingerprint': original['verification_fingerprint']}], confirmed=True)
    sources = generation_sources(checked)
    result = prepare_ra_workflow(profile, sources, source_bindings={
        '제품명': {'source_id': sources[0]['source_id'], 'quote': 'SyntheticDrugA'}}, direct_values={})
    assert result['ready_for_output_check'], result['review']
    assert result['actual_model_requests'] == 0
    assert not result['submission_ready']
    assert result['evidence']['제품명']['source']['verification_receipt']['document_sha256'] == digest


def test_invalid_image_json_offer_manual_transcription_without_guessing(tmp_path):
    path = image_file(tmp_path)
    client = Mock()
    client.read_image_json.return_value = {'본문': 123, '표 목록': []}
    result = collect_multimodal([path], client=client, allow_ocr=True)
    assert result['files'][0]['reason_code'] == 'invalid_ocr'
    assert result['files'][0]['manual_transcription_available']
    assert not result['sources']
