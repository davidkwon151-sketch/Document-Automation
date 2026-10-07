"""Actual synthetic files + mocked model; not live AI or human KPI evidence."""
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
import json

from docx import Document
import pytest

from agent.retrieve import chunk_documents, load_documents
from agent.review import CITATION_PATTERN
from app import ra_auto_service as service
from evals.run import MockEvaluationClient
from llm.client import LLMError
from templates.compatibility import analyze_template


@pytest.fixture
def bundle(tmp_path):
    template = tmp_path / 'Synthetic_RA_Form.docx'
    document = Document()
    document.add_heading('합성 RA 자동 기입 시험 양식', 0)
    table = document.add_table(rows=0, cols=2)
    for key in ['제품명', '성상', '저장방법', '포장단위', '신청인']:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = key, '{{' + key + '}}'
    document.save(template)
    profile = analyze_template(template)
    profile.update(configured=True, auto_mapping_confirmed=True, domain='pharmaceutical_ra',
                   ra_workflow='product_approval', document_kind='application')
    for field in profile['fields']:
        field['required'] = False
        field['narrative_style_required'] = False
        if field['value_key'] == '제품명':
            field['evidence_role'] = 'product_name'
    raw = tmp_path / 'Synthetic_RA_Evidence.txt'
    raw.write_text('제품명: SyntheticDrugA\n제형·함량: 정제 10 mg\n성상: 흰색 정제\n저장방법: 실온 보관\n포장단위: 30정\n신청인: 근거에서 추정하면 안 되는 이름', encoding='utf-8')
    sources = chunk_documents(load_documents([raw], allow_ocr=False))
    return template, profile, sources


def options(bundle, **changes):
    template, profile, sources = bundle
    options = dict(instruction='RA 담당자가 합성 시험 제품의 신청서 선정 항목을 작성해줘',
        template_path=template, profile=profile, sources=sources,
        selected_keys=['제품명', '성상', '저장방법', '포장단위', '신청인'],
        product_name='SyntheticDrugA', variant='정제 10 mg', field_values={'신청인': '사용자 직접 시험 입력'})
    options.update(changes)
    return options


def source_copy(bundle, **changes):
    args = options(bundle, **changes)
    proposal = service.propose_ra_auto(args['template_path'], args['profile'], args['sources'], args['selected_keys'],
                                      product_name=args['product_name'], variant=args['variant'])
    return service.generate_ra_auto(**args, mode='source_copy', confirmed_proposals=proposal['fingerprint'])


def test_actual_five_field_table_bulk_copy_and_independent_saved_docx(bundle):
    template, profile, sources = bundle
    original = template.read_bytes()
    result = source_copy(bundle)
    assert result['ready_for_output_check'], result['review']
    assert result['target_coverage']['filled_count'] == result['target_coverage']['target_count'] == 5
    assert result['actual_model_requests'] == result['model_request_attempts'] == 0
    assert result['locked_fields']['신청인']['value'] == '사용자 직접 시험 입력'
    output = service.export_ra_auto(result, confirmed=True)
    reopened = Document(BytesIO(output['document']))
    values = [row.cells[1].text for row in reopened.tables[0].rows]
    assert values == ['SyntheticDrugA', '흰색 정제', '실온 보관', '30정', '사용자 직접 시험 입력']
    assert template.read_bytes() == original
    assert output['output_verification']['status'] == 'passed'
    assert not output['submission_ready'] and not result['human_kpi_measured']
    evidence = json.loads(output['evidence'])
    assert evidence['auto']['source_records'] == sources
    assert evidence['target_coverage']['target_count'] == 5


def test_optional_selected_blank_blocks_empty_success_and_missing_question(bundle):
    result = source_copy(bundle, selected_keys=['성상', '저장방법'], sources=bundle[2][:1], field_values={})
    assert not result['ready_for_output_check']
    assert result['target_coverage']['target_count'] == 2
    assert result['target_coverage']['filled_count'] == 0
    assert set(result['target_coverage']['missing_keys']) == {'성상', '저장방법'}
    assert len(result['questions']) <= 2
    with pytest.raises(ValueError, match='오류'):
        service.export_ra_auto(result, confirmed=True)


def test_whole_candidate_confirmation_is_bound_to_current_source_and_profile(bundle):
    args = options(bundle)
    with pytest.raises(ValueError, match='확인'):
        service.generate_ra_auto(**args, mode='source_copy')
    proposal = service.propose_ra_auto(args['template_path'], args['profile'], args['sources'], args['selected_keys'],
                                      product_name=args['product_name'], variant=args['variant'])
    args['sources'] = deepcopy(args['sources'])
    args['sources'][-1]['text'] += ' 변경'
    args['sources'][-1].pop('context_text', None)
    args['sources'][-1].pop('context_start', None)
    args['sources'][-1].pop('context_end', None)
    with pytest.raises(ValueError, match='확인'):
        service.generate_ra_auto(**args, mode='source_copy', confirmed_proposals=proposal['fingerprint'])


@pytest.mark.parametrize('attack', ['model_origin', 'ocr', 'sha', 'source_id', 'context'])
def test_source_guards_before_any_model_request(bundle, attack):
    args = options(bundle)
    args['sources'] = deepcopy(args['sources'])
    if attack == 'model_origin': args['sources'][0]['source_kind'] = 'model'
    if attack == 'ocr': args['sources'][0]['requires_verification'] = True
    if attack == 'sha': args['sources'][0]['document_sha256'] = 'bad'
    if attack == 'source_id': args['sources'].append(deepcopy(args['sources'][0]))
    if attack == 'context': args['sources'][0]['context_text'] = '다른 문맥'
    class Never:
        def generate_json(self, *args): raise AssertionError('must validate before request')
    with pytest.raises(ValueError):
        service.generate_ra_auto(**args, mode='live', client=Never())


@pytest.mark.parametrize('attack', ['draft', 'source', 'profile', 'template', 'selected'])
def test_changed_data_or_targets_cannot_reuse_export_review(bundle, attack):
    result = source_copy(bundle)
    if attack == 'draft': result['draft']['포장단위'] = '999정 [S1]'
    if attack == 'source': result['auto']['source_records'][0]['text'] = '다른 제품'
    if attack == 'profile': result['template_profile']['fields'][-1]['input_required'] = False
    if attack == 'selected': result['auto']['selected_keys'] = ['성상']
    if attack == 'template': bundle[0].write_bytes(b'changed')
    with pytest.raises(ValueError):
        service.export_ra_auto(result, confirmed=True)


def test_new_mapping_and_nonempty_targets_require_explicit_confirmation(bundle):
    template, profile, sources = bundle
    profile = deepcopy(profile)
    profile.pop('auto_mapping_confirmed')
    with pytest.raises(ValueError, match='매핑'):
        service.target_profile(template, profile, ['성상'], product_name='SyntheticDrugA', variant='정제 10 mg')
    with pytest.raises(ValueError, match='선택'):
        service.target_profile(template, bundle[1], [], product_name='SyntheticDrugA', variant='정제 10 mg')


def test_personal_value_is_direct_and_missing_person_is_not_inferred(bundle):
    result = source_copy(bundle, field_values={})
    assert not result['draft']['신청인']
    assert '신청인' in result['target_coverage']['missing_keys']
    with pytest.raises(ValueError, match='직접 입력'):
        source_copy(bundle, field_values={'성상': '가짜 사실'})


class PipelineMock(MockEvaluationClient):
    def __init__(self, *, empty=False):
        super().__init__({'mock_brief': {'목적': 'SyntheticDrugA 제품명 성상 저장방법 포장단위',
           '보고 대상': 'RA 담당자', '보고서 유형': '결과보고서', '마감': '', '분량': '1쪽',
           '부족한 정보': [], '질문': []}})
        self.calls, self.empty = [], empty

    def generate_json(self, name, payload):
        self.calls.append((name, deepcopy(payload)))
        if name == 'draft':
            source = next(s for s in payload['sources'] if s['filename'] != '사용자 입력')
            quote = source['text'] + ' [' + source['source_id'] + ']'
            draft = {'제목': '합성 RA 시험', '요약': quote, '본문': quote}
            values = {'제품명': 'SyntheticDrugA', '성상': '흰색 정제', '저장방법': '실온 보관', '포장단위': '30정'}
            for field in payload.get('template_profile', {}).get('fields', []):
                label, key = field['label'], field['value_key']
                if label not in values:
                    continue
                source = next((s for s in payload['sources'] if label + ':' in s['text']), None)
                draft[key] = (values[label] + ' [' + source['source_id'] + ']') if source and not self.empty else ''
            return draft
        if name == 'grounding':
            by_id = {s['source_id']: s['text'] for s in payload['sources']}
            return {'claims': [{'field': item['field'], 'line': item['line'], 'status': 'supported',
                'evidence': [{'source_id': identifier, 'quote': by_id[identifier]}
                             for identifier in CITATION_PATTERN.findall(item['text']) if identifier in by_id]}
                for item in payload['claims']]}
        if name == 'completeness':
            return {'issues': [], 'checked_fields': list(payload['draft'])}
        return super().generate_json(name, payload)


def test_actual_pipeline_uses_central_prompts_flat_fields_and_retains_original_sources(bundle):
    client = PipelineMock()
    result = service.generate_ra_auto(**options(bundle), mode='live', client=client)
    assert result['ready_for_output_check'], result['review']
    assert result['target_coverage']['filled_count'] == 5
    assert {'brief', 'draft', 'grounding', 'completeness'} <= {name for name, _ in client.calls}
    used = {s['source_id']: s for s in result['sources'] if s['filename'] != '사용자 입력'}
    for source in bundle[2]:
        if source['source_id'] in used:
            assert {k: v for k, v in used[source['source_id']].items() if k != 'score'} == {k: v for k, v in source.items() if k != 'score'}
    assert result['model_request_attempts'] > 0 and result['actual_model_requests'] == 0
    output = service.export_ra_auto(result, confirmed=True)
    assert output['output_verification']['status'] == 'passed'
    assert Document(BytesIO(output['document'])).tables[0].rows[3].cells[1].text == '30정'


def test_provider_failure_propagates_without_copy_or_mock_fallback(bundle):
    class Failure:
        def generate_json(self, *args, **kwargs):
            raise LLMError('LLM 호출 실패 (quota)', kind='quota')
    with pytest.raises(LLMError):
        service.generate_ra_auto(**options(bundle), mode='live', client=Failure())


def test_live_edit_review_never_rewrites_and_preserves_first_ai_baseline(bundle):
    first = service.generate_ra_auto(**options(bundle), mode='live', client=PipelineMock())
    edited = deepcopy(first['draft'])
    edited['제목'] = '사용자가 바꾼 합성 RA 검토 제목'
    client = PipelineMock()
    reviewed = service.review_ra_auto(first, edited, client=client)
    assert reviewed['ready_for_output_check'], reviewed['review']
    assert {name for name, _ in client.calls} == {'grounding', 'completeness'}
    assert reviewed['metrics']['baseline_draft'] == first['metrics']['baseline_draft']
    assert reviewed['metrics']['draft_started_at'] == first['metrics']['draft_started_at']
    assert first['draft']['제목'] == '합성 RA 시험'
    assert service.export_ra_auto(reviewed, confirmed=True)['output_verification']['status'] == 'passed'
    with pytest.raises(ValueError, match='의미'):
        service.review_ra_auto(first, edited)


def test_locked_direct_value_cannot_be_changed_by_editing_ai_draft(bundle):
    first = service.generate_ra_auto(**options(bundle), mode='live', client=PipelineMock())
    edited = deepcopy(first['draft'])
    edited['신청인'] = 'AI가 대신 만든 이름 [' + first['locked_fields']['신청인']['source_id'] + ']'
    reviewed = service.review_ra_auto(first, edited, client=PipelineMock())
    assert not reviewed['ready_for_output_check']
    with pytest.raises(ValueError):
        service.export_ra_auto(reviewed, confirmed=True)


def test_original_numeric_label_survives_confirmed_unique_json_key_and_filling(bundle):
    p = deepcopy(bundle[1])
    next(f for f in p['fields'] if f['value_key'] == '포장단위')['value_key'] = '포장 1'
    args = options(bundle, profile=p, selected_keys=['제품명', '성상', '저장방법', '포장 1', '신청인'])
    result = service.generate_ra_auto(**args, mode='live', client=PipelineMock())
    assert result['ready_for_output_check'], result['review']
    exported = service.export_ra_auto(result, confirmed=True)
    assert Document(BytesIO(exported['document'])).tables[0].rows[3].cells[1].text == '30정'
    # Explicit original label is in the reviewed profile and its fingerprint.
    changed = deepcopy(result)
    next(f for f in changed['template_profile']['fields'] if f['value_key'] == '포장 1')['label'] = '다른 수치'
    with pytest.raises(ValueError):
        service.export_ra_auto(changed, confirmed=True)


def test_numeric_label_does_not_accept_same_amount_from_other_original_item(bundle):
    from agent.review import inspect_draft
    profile = {'fields': [{'id': 'p', 'label': '포장단위', 'value_key': '포장 1'}]}
    draft = {'제목': '합성 검토', '요약': '○ 추가 확인 필요함', '본문': '○ 추가 확인 필요함', '포장 1': '30정 [S1]'}
    sources = [{'source_id': 'S1', 'text': '투여수량: 30정'}]
    warnings = inspect_draft(draft, sources, template_profile=profile)
    assert any(item['code'] == 'number_mismatch' and item['field'] == '포장 1' for item in warnings)


def test_ai_optional_blank_selected_field_is_never_ready(bundle):
    with pytest.raises(ValueError, match='필수 양식'):
        service.generate_ra_auto(**options(bundle), mode='live', client=PipelineMock(empty=True))


def test_confirmed_chunk_pipeline_guards_apply_before_brief(bundle):
    from agent.pipeline import run_pipeline
    sources = deepcopy(bundle[2]); sources[0]['origin'] = 'ai_generated'
    client = PipelineMock()
    with pytest.raises(ValueError):
        run_pipeline('작성', source_records=sources, client=client)
    assert client.calls == []


def test_missing_required_evidence_after_top_forty_is_blocked_not_fabricated(bundle):
    """Current live retrieval cap is explicit; no false complete blank output."""
    from agent.pipeline import run_pipeline
    sources = [{'source_id': 'S' + str(i).zfill(3), 'filename': 'synthetic.txt', 'page': 1,
                'document_sha256': 'a' * 64, 'text': '동일 검색어 자료', 'location': str(i)} for i in range(41)]
    sources[-1]['text'] = '성상: 흰색 정제'
    class FlatVectors(PipelineMock):
        def embed(self, texts):
            return [[1.0] for _ in texts]
    client = FlatVectors()
    p = service.target_profile(bundle[0], bundle[1], ['성상'], product_name='SyntheticDrugA', variant='')
    # Query terms give all noise chunks a larger score than the actual late item.
    client.case['mock_brief']['목적'] = '동일 검색어 자료'
    with pytest.raises((ValueError, StopIteration)):
        run_pipeline('동일 검색어 자료', source_records=sources, template_profile=p,
                     ra_workflow='product_approval', client=client, semantic_review=True)
    draft_payload = next(payload for name, payload in client.calls if name == 'draft')
    assert len(draft_payload['sources']) == 6  # One field uses max(6, field_count*2), capped at 40.
    assert 'S040' not in {s['source_id'] for s in draft_payload['sources']}


def test_original_required_and_relation_dependencies_cannot_be_dropped(bundle):
    p = deepcopy(bundle[1])
    p['fields'][0]['required'] = True
    for field in p['fields'][1:3]:
        field['validation'] = {'type': 'integer'}
    p['constraints'] = {'relations': [{'kind': 'less_equal', 'left': '성상', 'right': '저장방법'}]}
    runtime = service.target_profile(bundle[0], p, ['신청인'], product_name='SyntheticDrugA', variant='')
    assert runtime['auto_target_keys'] == ['신청인', '제품명']
    assert {'성상', '저장방법'} <= {f['value_key'] for f in runtime['fields']}
    assert runtime['constraints'] == p['constraints']


def test_confirmed_original_image_receipt_survives_pipeline_and_tamper_blocks(bundle, tmp_path):
    from PIL import Image
    from agent.multimodal_intake import collect_multimodal, confirm_intake, generation_sources
    from agent.pipeline import run_pipeline
    path = tmp_path / 'Synthetic_scan.png'
    Image.new('RGB', (100, 100), 'white').save(path)
    digest = sha256(path.read_bytes()).hexdigest()
    intake = collect_multimodal([path], transcriptions={digest: [{'page': 1, 'text': '제품명: SyntheticDrugA\n성상: 흰색 정제'}]})
    receipts = [{'source_id': s['source_id'], 'fingerprint': s['verification_fingerprint']} for s in intake['sources']]
    accepted = generation_sources(confirm_intake(intake, receipts, confirmed=True))
    client = PipelineMock()
    p = service.target_profile(bundle[0], bundle[1], ['제품명', '성상'], product_name='SyntheticDrugA', variant='')
    # The stand-in only emits registered fields for this deliberately small case.
    class Small(PipelineMock):
        def generate_json(self, name, payload):
            if name == 'draft':
                s = next(s for s in payload['sources'] if '제품명:' in s['text'])
                a = next(s for s in payload['sources'] if '성상:' in s['text'])
                return {'제목': '합성 검토', '요약': s['text'] + ' [' + s['source_id'] + ']',
                        '본문': a['text'] + ' [' + a['source_id'] + ']',
                        '제품명': 'SyntheticDrugA [' + s['source_id'] + ']', '성상': '흰색 정제 [' + a['source_id'] + ']'}
            return super().generate_json(name, payload)
    result = run_pipeline('SyntheticDrugA 제품명 성상', source_records=accepted, template_profile=p,
                          ra_workflow='product_approval', client=Small(), semantic_review=True)
    assert all('verification_receipt' in s for s in result['sources'])
    assert {s['source_id'] for s in result['sources']} == {s['source_id'] for s in accepted}
    accepted[0]['verification_receipt']['document_sha256'] = '0' * 64
    with pytest.raises(ValueError):
        run_pipeline('작성', source_records=accepted, client=client)
    assert client.calls == []
