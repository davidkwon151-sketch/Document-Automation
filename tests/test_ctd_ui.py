"""CTD workspace keeps exports tied to current confirmed source material."""

from io import BytesIO
import json
from zipfile import ZipFile

from docx import Document
from PIL import Image
import pytest
import streamlit as st
from streamlit.testing.v1 import AppTest

from app import ctd_ui


class Upload:
    def __init__(self, name, data):
        self.name, self.data = name, data

    def getvalue(self):
        return self.data


def uploads(files, template=None):
    return lambda *args, **kwargs: (template if kwargs.get('key') == 'ctd_template_upload'
                                    else files)


def package():
    return {
        'sections': [
            {'section_id': 'M1_ADMIN', 'title': '행정', 'status': 'proposed',
             'draft': '원문 제품명: 시험정 5 mg', 'evidence': [{
                 'source_id': 'S1', 'filename': '원자료.txt', 'document_sha256': 'a' * 64,
                 'page': 1, 'sheet': None, 'location': '페이지 1',
                 'quote': '제품명: 시험정 5 mg', 'start': 0, 'end': 14}]},
            {'section_id': '3.2.P.8.3', 'title': '안정성 자료', 'status': 'missing',
             'draft': '', 'evidence': []},
            {'section_id': '3.2.S.4', 'title': '원료의약품 관리', 'status': 'deferred',
             'draft': '', 'evidence': []}],
        'missing_sections': ['3.2.P.8.3'], 'ambiguous_sections': [],
        'deferred_sections': ['3.2.S.4'],
        'manual_checks': ['담당자 원문 확인'], 'deferred_sources': [],
        'submission_ready': False, 'actual_model_requests': 0,
    }


def test_export_is_marked_working_draft_and_keeps_evidence_and_missing_sections():
    result = ctd_ui.export_ctd_working_draft(package())
    document = Document(BytesIO(result['document']))
    text = '\n'.join(paragraph.text for paragraph in document.paragraphs)
    assert '제출용 완성본 아님' in text
    assert 'M1_ADMIN 행정' in text and '3.2.P.8.3 안정성 자료' in text
    assert '원문 제품명: 시험정 5 mg' in text and '근거 자료 없음' in text
    assert '근거 후보 있음 · 자동 작성 보류' in text
    evidence = json.loads(result['evidence'])
    assert evidence['sections'][0]['evidence'][0]['document_sha256'] == 'a' * 64
    assert evidence['missing_sections'] == ['3.2.P.8.3']
    assert evidence['deferred_sections'] == ['3.2.S.4']
    with ZipFile(BytesIO(result['zip'])) as archive:
        assert set(archive.namelist()) == {'CTD_Module1_3_작업초안.docx', 'CTD_출처와_누락_검토.json'}
        assert archive.read('CTD_출처와_누락_검토.json') == result['evidence']


def test_export_rejects_submission_ready_claim():
    item = package()
    item['submission_ready'] = True
    with pytest.raises(ValueError, match='제출용'):
        ctd_ui.export_ctd_working_draft(item)


def app():
    return AppTest.from_string('from app.ctd_ui import render_ctd\nrender_ctd()', default_timeout=30).run()


def test_ctd_upload_generates_only_for_current_product_and_source(monkeypatch):
    files = [Upload('원자료.txt', '제품명: 시험정 5 mg\n안정성: 확인 자료 없음'.encode())]
    monkeypatch.setattr(st, 'file_uploader', uploads(files))
    monkeypatch.setattr(ctd_ui, 'prepare_ctd_package', lambda sources, **kwargs: package())
    screen = app()
    assert not screen.exception
    next(widget for widget in screen.text_input if '정확한 제품명' in widget.label).set_value('시험정')
    screen.run()
    next(widget for widget in screen.button if '작업 초안 만들기' in widget.label).click()
    screen.run()
    assert not screen.exception
    assert len(screen.get('download_button')) == 3
    assert screen.session_state['ctd_result']['submission_ready'] is False
    files[0].data = '제품명: 다른정 5 mg'.encode()
    screen.run()
    assert 'ctd_result' not in screen.session_state
    assert not screen.get('download_button')


def test_ctd_real_upload_carries_stability_value_into_working_draft(monkeypatch):
    files = [Upload('품질자료.txt',
                    '제품명: 시험정 5 mg\n\n3.2.P.8.3: 안정성 자료\n안정성 시험: 12개월'.encode())]
    monkeypatch.setattr(st, 'file_uploader', uploads(files))
    screen = app()
    next(widget for widget in screen.text_input if '정확한 제품명' in widget.label).set_value('시험정 5 mg')
    next(widget for widget in screen.multiselect if 'Module 1·3 절' in widget.label).set_value(['3.2.P.8.3'])
    screen.run()
    next(widget for widget in screen.button if '작업 초안 만들기' in widget.label).click()
    screen.run()
    assert not screen.exception
    section = screen.session_state['ctd_result']['sections'][0]
    assert section['status'] == 'proposed'
    assert '안정성 시험: 12개월' in section['draft']
    assert any('12개월' in evidence['quote'] for evidence in section['evidence'])
    assert len(screen.get('download_button')) == 3


def _set_mapping(screen, section_id, source_id):
    next(widget for widget in screen.selectbox if widget.label == '연결할 CTD 절').set_value(section_id)
    screen.run()
    next(widget for widget in screen.selectbox if widget.label == '원문 조각 선택').set_value(source_id)
    screen.run()
    next(widget for widget in screen.checkbox if '원본의 전체 문구' in widget.label).check()
    screen.run()


def test_manual_source_choice_rebuilds_and_source_or_product_change_expires_mapping(monkeypatch):
    files = [Upload('품질자료.txt', '제품명: 시험정 5 mg\n\n안정성 시험: 12개월'.encode())]
    monkeypatch.setattr(st, 'file_uploader', uploads(files))
    screen = app()
    next(widget for widget in screen.text_input if '정확한 제품명' in widget.label).set_value('시험정 5 mg')
    next(widget for widget in screen.multiselect if 'Module 1·3 절' in widget.label).set_value(['3.2.P.8.3'])
    screen.run()
    next(widget for widget in screen.button if '작업 초안 만들기' in widget.label).click()
    screen.run()
    assert screen.session_state['ctd_result']['sections'][0]['status'] == 'missing'
    source_id = next(source['source_id'] for source in screen.session_state['ctd_mm_intake']['sources']
                     if source['text'] == '안정성 시험: 12개월')
    _set_mapping(screen, '3.2.P.8.3', source_id)
    assert 'ctd_result' not in screen.session_state
    assert not screen.get('download_button')
    next(widget for widget in screen.button if '확인한 원문으로 절 다시 작성' in widget.label).click()
    screen.run()
    assert not screen.exception
    section = screen.session_state['ctd_result']['sections'][0]
    assert section['status'] == 'manual_check'
    assert '안정성 시험: 12개월' in section['draft']
    assert screen.session_state['ctd_mapping'] == {'3.2.P.8.3': [source_id]}
    assert len(screen.get('download_button')) == 3
    files[0].data = '제품명: 다른정 5 mg\n\n안정성 시험: 24개월'.encode()
    screen.run()
    assert screen.session_state['ctd_mapping'] == {}
    assert 'ctd_result' not in screen.session_state
    assert not screen.get('download_button')
    next(widget for widget in screen.text_input if '정확한 제품명' in widget.label).set_value('다른정 5 mg')
    screen.run()
    assert screen.session_state['ctd_mapping'] == {}


def test_same_source_cannot_be_selected_for_two_ctd_sections(monkeypatch):
    files = [Upload('품질자료.txt', '제품명: 시험정\n\n배치 분석: 10 mg'.encode())]
    monkeypatch.setattr(st, 'file_uploader', uploads(files))
    screen = app()
    next(widget for widget in screen.text_input if '정확한 제품명' in widget.label).set_value('시험정')
    next(widget for widget in screen.multiselect if 'Module 1·3 절' in widget.label).set_value(
        ['3.2.P.5.4', '3.2.P.8.3'])
    screen.run()
    source_id = next(source['source_id'] for source in screen.session_state['ctd_mm_intake']['sources']
                     if source['text'] == '배치 분석: 10 mg')
    _set_mapping(screen, '3.2.P.5.4', source_id)
    next(widget for widget in screen.button if '확인한 원문으로 절 다시 작성' in widget.label).click()
    screen.run()
    _set_mapping(screen, '3.2.P.8.3', source_id)
    assert any('이미 다른 CTD 절' in item.value for item in screen.error)
    assert next(widget for widget in screen.button if '확인한 원문으로 절 다시 작성' in widget.label).disabled
    assert screen.session_state['ctd_mapping'] == {'3.2.P.5.4': [source_id]}
    next(widget for widget in screen.multiselect if 'Module 1·3 절' in widget.label).set_value(['3.2.S.4'])
    screen.run()
    assert screen.session_state['ctd_mapping'] == {}
    assert 'ctd_result' not in screen.session_state


def test_product_and_variant_are_only_applied_after_user_click(monkeypatch):
    files = [Upload('품질자료.txt',
                    '제품명: 시험정 5 mg\n제형: 정제\n함량: 5 mg\n\n3.2.P.8.3: 안정성 자료\n시험: 12개월'.encode())]
    monkeypatch.setattr(st, 'file_uploader', uploads(files))
    screen = app()
    assert not screen.exception
    assert next(widget for widget in screen.text_input if '정확한 제품명' in widget.label).value == ''
    assert next(widget for widget in screen.text_input if '선택 제형' in widget.label).value == ''
    assert next(widget for widget in screen.selectbox if widget.label == '원문 제품명 후보').value == '시험정 5 mg'
    next(widget for widget in screen.button if '확인한 제품명 입력칸에 적용' in widget.label).click()
    screen.run()
    assert not screen.exception
    assert next(widget for widget in screen.text_input if '정확한 제품명' in widget.label).value == '시험정 5 mg'
    assert next(widget for widget in screen.text_input if '선택 제형' in widget.label).value == '정제 5 mg'


def test_ai_section_proposal_requires_separate_review_then_records_real_request_count(monkeypatch):
    files = [Upload('품질자료.txt', '제품명: 시험정\n\n안정성 시험: 12개월'.encode())]
    monkeypatch.setattr(st, 'file_uploader', uploads(files))
    calls = []
    monkeypatch.setattr(ctd_ui, 'LLMClient', lambda: calls.append('client') or object())

    def suggest(intake, **kwargs):
        calls.append(('suggest', kwargs['product_name'], len(intake['sources'])))
        identifier = next(source['source_id'] for source in intake['sources']
                          if source['text'] == '안정성 시험: 12개월')
        return {'section_map': {'3.2.P.8.3': [identifier]}, 'unmapped_source_ids': [],
                'requires_confirmation': True, 'mapping_origin': 'model_proposal',
                'actual_model_requests': 1, 'model_request_attempts': 1}

    monkeypatch.setattr(ctd_ui, 'suggest_ctd_section_map', suggest)
    screen = app()
    next(widget for widget in screen.text_input if '정확한 제품명' in widget.label).set_value('시험정')
    next(widget for widget in screen.multiselect if 'Module 1·3 절' in widget.label).set_value(['3.2.P.8.3'])
    screen.run()
    assert not calls
    next(widget for widget in screen.button if widget.label == 'AI 절 연결 제안').click()
    screen.run()
    assert not screen.exception
    assert calls == ['client', ('suggest', '시험정', 2)]
    assert screen.session_state['ctd_mapping'] == {}
    assert 'ctd_result' not in screen.session_state
    assert next(widget for widget in screen.button if '확인한 AI 연결' in widget.label).disabled
    next(widget for widget in screen.checkbox if '제안된 원문 전체' in widget.label).check()
    screen.run()
    next(widget for widget in screen.button if '확인한 AI 연결' in widget.label).click()
    screen.run()
    assert not screen.exception
    assert screen.session_state['ctd_result']['sections'][0]['status'] == 'manual_check'
    assert screen.session_state['ctd_result']['mapping_model_requests'] == 1
    assert screen.session_state['ctd_result']['actual_model_requests'] == 0
    assert screen.session_state['ctd_result']['mapping_origin'] == 'model_proposal_confirmed_by_user'
    assert len(screen.get('download_button')) == 3
    next(widget for widget in screen.text_input if '정확한 제품명' in widget.label).set_value('다른정')
    screen.run()
    assert screen.session_state['ctd_mapping'] == {}
    assert 'ctd_ai_proposal' not in screen.session_state
    assert not screen.get('download_button')


def test_ai_mapping_provider_failure_is_safe_and_preserves_manual_path(monkeypatch):
    from llm.client import LLMError

    files = [Upload('품질자료.txt', '제품명: 시험정\n\n안정성 시험: 12개월'.encode())]
    monkeypatch.setattr(st, 'file_uploader', uploads(files))
    def fail():
        raise LLMError('RAW_PROVIDER_SECRET', kind='quota')
    monkeypatch.setattr(ctd_ui, 'LLMClient', fail)
    screen = app()
    next(widget for widget in screen.text_input if '정확한 제품명' in widget.label).set_value('시험정')
    screen.run()
    next(widget for widget in screen.button if widget.label == 'AI 절 연결 제안').click()
    screen.run()
    assert not screen.exception
    assert any('API 잔액' in item.value for item in screen.error)
    assert all('RAW_PROVIDER_SECRET' not in item.value for item in screen.error)
    assert 'ctd_ai_proposal' not in screen.session_state
    assert any(widget.label == '연결할 CTD 절' for widget in screen.selectbox)


def test_substance_source_requires_named_user_link_and_only_opens_s_section(monkeypatch):
    files = [
        Upload('완제.txt', '제품명: 시험정\n\n3.2.P.8.3: 안정성 자료\n시험: 12개월'.encode()),
        Upload('원료.txt', '원료의약품명: 원료A\n\n3.2.S.4: 원료의약품 관리\n규격: 98%'.encode()),
    ]
    monkeypatch.setattr(st, 'file_uploader', uploads(files))
    screen = app()
    next(widget for widget in screen.text_input if '정확한 제품명' in widget.label).set_value('시험정')
    next(widget for widget in screen.multiselect if 'Module 1·3 절' in widget.label).set_value(
        ['3.2.S.4', '3.2.P.8.3'])
    screen.run()
    next(widget for widget in screen.button if '작업 초안 만들기' in widget.label).click()
    screen.run()
    before = {item['section_id']: item for item in screen.session_state['ctd_result']['sections']}
    assert before['3.2.S.4']['status'] == 'deferred'
    assert before['3.2.P.8.3']['status'] == 'proposed'
    digest = next(source['document_sha256'] for source in screen.session_state['ctd_mm_intake']['sources']
                  if source['filename'] == '원료.txt')
    next(widget for widget in screen.selectbox if widget.label == '원료 원본 파일').set_value(digest)
    screen.run()
    assert 'ctd_result' not in screen.session_state
    assert not screen.get('download_button')
    assert next(widget for widget in screen.selectbox if widget.label == '원문 원료명').value == '원료A'
    next(widget for widget in screen.checkbox if '이 원료가 선택 완제 제품' in widget.label).check()
    screen.run()
    next(widget for widget in screen.button if '원료–완제 관계로 S 절 다시 작성' in widget.label).click()
    screen.run()
    assert not screen.exception
    after = {item['section_id']: item for item in screen.session_state['ctd_result']['sections']}
    assert after['3.2.S.4']['status'] == 'manual_check'
    assert '규격: 98%' in after['3.2.S.4']['draft']
    assert any(item.get('source_scope') == 'confirmed_substance_link'
               for item in after['3.2.S.4']['evidence'])
    assert all(item['filename'] == '완제.txt' for item in after['3.2.P.8.3']['evidence'])
    assert screen.session_state['ctd_substance_links'][digest] == {
        'substance_name': '원료A', 'product_name': '시험정', 'confirmed': True}
    captured = {}
    monkeypatch.setattr(ctd_ui, 'LLMClient', lambda: object())
    def inspect_links(intake, **kwargs):
        captured.update(kwargs)
        return {'section_map': {}, 'unmapped_source_ids': [], 'requires_confirmation': True,
                'mapping_origin': 'model_proposal', 'actual_model_requests': 1,
                'model_request_attempts': 1}
    monkeypatch.setattr(ctd_ui, 'suggest_ctd_section_map', inspect_links)
    next(widget for widget in screen.button if widget.label == 'AI 절 연결 제안').click()
    screen.run()
    assert captured['confirmed_substance_links'] == screen.session_state['ctd_substance_links']
    next(widget for widget in screen.text_input if '정확한 제품명' in widget.label).set_value('다른정')
    screen.run()
    assert screen.session_state['ctd_substance_links'] == {}
    assert not screen.get('download_button')


def test_uploaded_ctd_docx_fills_exact_section_and_keeps_sidecar(monkeypatch):
    original = Document()
    original.add_paragraph('사내 CTD 작업 양식 · 공식 제출본 아님')
    original.add_paragraph('{{3.2.P.8.3}}')
    original.add_paragraph('')  # Common layout spacer is preserved, not treated as an input.
    stream = BytesIO()
    original.save(stream)
    template = Upload('ctd.docx', stream.getvalue())
    files = [Upload('품질자료.txt', '제품명: 시험정\n\n3.2.P.8.3: 안정성 자료\n시험: 12개월'.encode())]
    monkeypatch.setattr(st, 'file_uploader', uploads(files, template))
    screen = app()
    next(widget for widget in screen.button if '양식 자리표시자 절만 선택' in widget.label).click()
    screen.run()
    assert next(widget for widget in screen.multiselect if 'Module 1·3 절' in widget.label).value == ['3.2.P.8.3']
    next(widget for widget in screen.text_input if '정확한 제품명' in widget.label).set_value('시험정')
    next(widget for widget in screen.multiselect if 'Module 1·3 절' in widget.label).set_value(
        ['3.2.P.8.3', '3.2.P.5.4'])
    screen.run()
    assert not screen.exception
    next(widget for widget in screen.button if '작업 초안 만들기' in widget.label).click()
    screen.run()
    next(widget for widget in screen.button if '업로드한 CTD 양식에 초안 기입' in widget.label).click()
    screen.run()
    assert not screen.exception
    output = screen.session_state['ctd_template_export']
    text = '\n'.join(p.text for p in Document(BytesIO(output['document'])).paragraphs)
    assert '시험: 12개월' in text
    assert '{{3.2.P.8.3}}' not in text
    assert '사내 CTD 작업 양식 · 공식 제출본 아님' in text
    evidence = json.loads(output['evidence'])
    assert evidence['submission_ready'] is False
    assert evidence['evidence']['3.2.P.8.3'][0]['filename'] == '품질자료.txt'
    assert evidence['rendered_sections'] == ['3.2.P.8.3']
    assert '3.2.P.5.4' in evidence['missing_sections']
    assert len(screen.get('download_button')) == 5
    template.data = stream.getvalue() + b'changed'
    screen.run()
    assert 'ctd_template_export' not in screen.session_state
    assert not screen.get('download_button')


def test_unconfirmed_image_does_not_trigger_paid_model_or_generate(monkeypatch):
    image = BytesIO()
    Image.new('RGB', (50, 30), 'white').save(image, format='PNG')
    monkeypatch.setattr(st, 'file_uploader', uploads([Upload('스캔.png', image.getvalue())]))
    from app import multimodal_ui
    monkeypatch.setattr(multimodal_ui, 'LLMClient', lambda: pytest.fail('implicit model call'))
    screen = app()
    next(widget for widget in screen.text_input if '정확한 제품명' in widget.label).set_value('시험정')
    screen.run()
    assert not screen.exception
    assert any('미확인' in warning.value for warning in screen.warning)
    assert not screen.get('download_button')
    assert 'ctd_result' not in screen.session_state
