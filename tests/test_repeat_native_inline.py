"""Inline Word choices do not swallow adjacent, position-bound repeat slots."""
from copy import deepcopy
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
import pytest

from agent.output_check import verify_output
from templates import fill_compatible_template
from templates.repeat_fields import repeat_profile
from templates.repeat_rows import (expansion_binding, inspect_repeat_tables,
                                   prepare_repeat_template, verify_prepared_template)


NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}


def _choice(paragraph, initial, labels):
    sdt = OxmlElement('w:sdt'); properties = OxmlElement('w:sdtPr')
    for name, value in [('id', '81'), ('alias', '등급')]:
        node = OxmlElement('w:' + name); node.set(qn('w:val'), value); properties.append(node)
    properties.append(OxmlElement('w:showingPlcHdr'))
    options = OxmlElement('w:dropDownList')
    for value, label in zip(['A', 'B'], labels):
        node = OxmlElement('w:listItem'); node.set(qn('w:value'), value)
        node.set(qn('w:displayText'), label); options.append(node)
    properties.append(options); content = OxmlElement('w:sdtContent')
    run = OxmlElement('w:r'); text = OxmlElement('w:t'); text.text = initial
    run.append(text); content.append(run); sdt.extend([properties, content]); paragraph._p.append(sdt)


def _token(paragraph, split):
    if split:
        paragraph.add_run('{{품').bold = True
        paragraph.add_run('목}}').italic = True
    else:
        paragraph.add_run('{{품목}}').bold = True


def _fixture(tmp_path, position, split, duplicate, token_label, token_initial=False):
    original = tmp_path / 'original.docx'; prepared = tmp_path / 'prepared.docx'
    document = Document(); document.add_paragraph('원본 안내 유지; 선택·반복 입력 시험용')
    table = document.add_table(rows=3, cols=2); table.style = 'Table Grid'
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    table.cell(0, 0).text = '재료와 등급'; table.cell(0, 1).text = '수량'
    paragraph = table.cell(1, 0).paragraphs[0]
    labels = ['긴 선택 {{품목}} 표시 A', 'B'] if token_label else ['원래보다 훨씬 긴 등급 표시 A', 'B']
    initial = '{{품목}}' if token_initial else '선택'
    if position == 'before':
        paragraph.add_run('등급: '); _choice(paragraph, initial, labels); paragraph.add_run('; ')
    paragraph.add_run('재료: '); _token(paragraph, split)
    if duplicate:
        paragraph.add_run(' / 두번째: '); _token(paragraph, split)
    if position == 'after':
        paragraph.add_run('; 등급: '); _choice(paragraph, initial, labels)
    paragraph.add_run('; 단위 kg')
    table.cell(1, 1).text = '{{수량}} kg'
    table.cell(2, 0).text = '합계'; table.cell(2, 1).text = '고정 단위 kg'
    document.save(original)
    plan = {'table_id': inspect_repeat_tables(original)[0]['id'], 'row': 2, 'count': 2}
    expansion = prepare_repeat_template(original, prepared, plan)
    profile = repeat_profile(prepared, plan, expansion_binding(expansion))
    values = {}; products = {1: [], 2: []}
    for field in profile['fields']:
        row = field['repeat_info']['row']
        if field.get('control_type') == 'choice':
            value = ['A', 'B'][row - 1]
        elif field.get('placeholder_key') == '수량':
            value = str(row + 1)
        else:
            value = f'PRODUCT-{row}-{len(products[row]) + 1}'; products[row].append(value)
        values[field['value_key']] = value
    return original, prepared, expansion, profile, values, labels, products


def _root(path):
    with ZipFile(path) as archive:
        return etree.fromstring(archive.read('word/document.xml'))


@pytest.mark.parametrize('position', ['before', 'after'])
@pytest.mark.parametrize('split', [False, True])
@pytest.mark.parametrize('duplicate', [False, True])
@pytest.mark.parametrize('token_label', [False, True])
def test_inline_choices_and_adjacent_position_slots_are_independently_verified(
        tmp_path, position, split, duplicate, token_label):
    original, prepared, expansion, profile, values, labels, products = _fixture(
        tmp_path, position, split, duplicate, token_label)
    before = {path: path.read_bytes() for path in [original, prepared]}
    original_profile = deepcopy(profile); original_values = deepcopy(values)
    output = tmp_path / 'filled.docx'
    assert verify_prepared_template(prepared, profile, expansion)['status'] == 'passed'
    fill_compatible_template(prepared, values, output, profile=profile)
    assert verify_output(prepared, output, values, profile=profile)['status'] == 'passed'
    rows = _root(output).xpath('.//w:tbl/w:tr', namespaces=NS)
    assert len(rows) == 4
    for row_number in [1, 2]:
        cell = rows[row_number].find('w:tc', NS)
        text = ''.join(cell.xpath('.//w:t/text()', namespaces=NS))
        body = '재료: ' + products[row_number][0]
        if duplicate:
            body += ' / 두번째: ' + products[row_number][1]
        expected = (f'등급: {labels[row_number - 1]}; {body}' if position == 'before'
                    else f'{body}; 등급: {labels[row_number - 1]}') + '; 단위 kg'
        assert text == expected
        options = cell.xpath('.//w:dropDownList/w:listItem', namespaces=NS)
        assert [(node.get(qn('w:value')), node.get(qn('w:displayText'))) for node in options] == list(zip(['A', 'B'], labels))
        quantity = ''.join(rows[row_number].findall('w:tc', NS)[1].xpath('.//w:t/text()', namespaces=NS))
        assert quantity == f'{row_number + 1} kg'
    assert all(path.read_bytes() == raw for path, raw in before.items())
    assert profile == original_profile and values == original_values


@pytest.mark.parametrize('position', ['before', 'after'])
def test_unselected_native_placeholder_is_not_an_adjacent_repeat_slot(tmp_path, position):
    _, prepared, _, profile, values, _, _ = _fixture(
        tmp_path, position, True, True, True, token_initial=True)
    slots = [field for field in profile['fields'] if field['kind'] == 'docx_repeat_placeholder']
    assert len(slots) == 6  # two material tokens and one quantity token per row
    assert len({field['value_key'] for field in profile['fields']}) == len(profile['fields'])
    output = tmp_path / 'filled.docx'
    fill_compatible_template(prepared, values, output, profile=profile)
    assert verify_output(prepared, output, values, profile=profile)['status'] == 'passed'
    selected = _root(output).xpath('.//w:sdtContent//w:t/text()', namespaces=NS)
    assert selected == ['긴 선택 {{품목}} 표시 A', 'B']
