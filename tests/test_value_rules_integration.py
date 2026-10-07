"""Validate real serialization paths, mapping changes and user edits without APIs."""

from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile

from docx import Document
import pytest

from agent.brief import profile_fields
from agent.grounding import evidence_fingerprint
from agent.output_check import verify_output
from agent.pipeline import build_downloads, review_result, validate_template_values
from agent.retrieve import chunk_documents
from agent.template_learning import learn_template, load_learned_profile, save_learned_profile
from app.form_config import (configure_profile, configure_value_rules, configure_relations,
                             mapping_rows, validation_rows, relation_rows, value_rule_help)
from templates import analyze_template, fill_compatible_template

ROOT = Path(__file__).resolve().parents[1]


def typed_profile(path):
    profile = analyze_template(path)
    rules = {'제목': {'type': 'date'}, '요약': {'type': 'number', 'min': '0'},
             '본문': {'type': 'integer', 'min': 0}}
    for field in profile['fields']:
        if field['value_key'] in rules:
            field.update(validation=rules[field['value_key']], narrative_style_required=False)
    return profile


@pytest.mark.parametrize('suffix', ['docx', 'hwpx', 'xlsx', 'pptx', 'pdf'])
def test_formats_validate_before_write_and_reopen_without_changing_text(tmp_path, suffix):
    path = ROOT / f'samples/sample_company_form.{suffix}'
    profile = typed_profile(path)
    values = {'제목': '2024-02-29', '요약': '0.1', '본문': '3'}
    before = path.read_bytes()
    invalid_output = tmp_path / f'invalid.{suffix}'
    with pytest.raises(ValueError, match='제목|날짜'):
        fill_compatible_template(path, values | {'제목': '2025-02-29'}, invalid_output, profile=profile)
    assert not invalid_output.exists()
    output = fill_compatible_template(path, values, tmp_path / path.name, profile=profile)
    checked = verify_output(path, output, values, profile=profile)
    assert next(check for check in checked['checks'] if check['name'] == 'form_value_rules')['fields'] == 3
    assert path.read_bytes() == before
    assert values == {'제목': '2024-02-29', '요약': '0.1', '본문': '3'}
    with pytest.raises(ValueError, match='검증 실패'):
        verify_output(path, output, values | {'본문': 'NaN'}, profile=profile)


def amount_form(tmp_path):
    path = tmp_path / 'budget.docx'
    doc = Document()
    for key in ('금액A', '금액B', '합계', '시작일', '종료일'):
        doc.add_paragraph(key + ': {{' + key + '}}')
    doc.save(path)
    profile = analyze_template(path)
    for field in profile['fields']:
        field.update(narrative_style_required=False,
                     validation={'type': 'date'} if '일' in field['value_key']
                     else {'type': 'number', 'unit': '원', 'min': '0'})
    profile['constraints'] = {'relations': [
        {'kind': 'sum', 'total': '합계', 'parts': ['금액A', '금액B']},
        {'kind': 'date_order', 'start': '시작일', 'end': '종료일'}]}
    return path, profile


def amounts():
    return {'금액A': '0.1원', '금액B': '0.2원', '합계': '0.3원',
            '시작일': '2024-02-29', '종료일': '2024-03-01'}


def result_for(profile):
    values = amounts()
    draft = {'제목': '시험 문서', '요약': '□ 자료를 확인함 [Sbase]', '본문': '○ 자료를 확인함 [Sbase]'}
    sources = [{'source_id': 'Sbase', 'text': '자료를 확인함'}]
    for index, (key, value) in enumerate(values.items()):
        identifier = f'Sv{index}'
        sources.append({'source_id': identifier, 'text': value})
        draft[key] = value + f' [{identifier}]'
    return {'status': 'ready', 'draft': draft, 'sources': sources, 'template_profile': profile}


def test_manual_edit_is_blocked_even_when_wrong_total_exists_in_evidence(tmp_path):
    path, profile = amount_form(tmp_path)
    result = result_for(profile)
    assert not review_result(result)['blocking']
    exports = build_downloads(result, confirmed=True, template_paths={'docx': path})
    assert exports['docx'].startswith(b'PK')
    result['sources'].append({'source_id': 'Swrong', 'text': '9원'})
    result['draft']['합계'] = '9원 [Swrong]'
    checked = review_result(result)
    assert checked['blocking']
    assert any('sum' in issue['code'] and issue['field'] == '합계' for issue in checked['warnings'])
    assert result['draft']['합계'] == '9원 [Swrong]'  # Never silently recompute submitted figures.
    with pytest.raises(ValueError, match='오류'):
        build_downloads(result, confirmed=True, template_paths={'docx': path})
    with pytest.raises(ValueError, match='합계'):
        validate_template_values(result['draft'], profile)


def test_mapping_renames_constraints_and_cannot_drop_relation_member(tmp_path):
    path, profile = amount_form(tmp_path)
    rows = mapping_rows(profile)
    for row in rows:
        row['채울 값'] += '_새항목'
    configured, mapping = configure_profile(profile, rows)
    assert configured['constraints']['relations'][0]['total'] == '합계_새항목'
    assert profile['constraints']['relations'][0]['total'] == '합계'
    values = {key + '_새항목': value for key, value in amounts().items()}
    output = fill_compatible_template(path, values, tmp_path / 'mapped.docx', profile=profile, mapping=mapping)
    assert verify_output(path, output, values, profile=profile, mapping=mapping)['status'] == 'passed'
    with pytest.raises(ValueError, match='합계|규칙'):
        fill_compatible_template(path, values | {'합계_새항목': '9원'}, tmp_path / 'bad.docx', profile=profile, mapping=mapping)
    with pytest.raises(ValueError):
        fill_compatible_template(path, values, tmp_path / 'omitted.docx', profile=profile,
                                 mapping={key: value for key, value in mapping.items() if value != '합계_새항목'})


def test_value_editors_roundtrip_exact_units_and_validate_new_relations(tmp_path):
    _, profile = amount_form(tmp_path)
    rows = validation_rows(profile)
    configured = configure_value_rules(profile, rows)
    assert [field['validation'] for field in configured['fields']] == [
        dict(field['validation'], **({'unit_location': 'value'} if 'unit' in field['validation']
                                    else {'date_formats': ['YYYY-MM-DD']})) for field in profile['fields']]
    configured = configure_relations(configured, relation_rows(configured))
    assert configured['constraints']['relations'] == profile['constraints']['relations']
    amount = next(field for field in configured['fields'] if field['value_key'] == '합계')
    assert '단위 포함: 원' in value_rule_help(amount)
    rows[0]['단위 위치'] = '양식에 인쇄됨'
    configured = configure_value_rules(profile, rows)
    assert configured['fields'][0]['validation']['unit_location'] == 'label'
    assert '숫자만 입력' in value_rule_help(configured['fields'][0])
    with pytest.raises(ValueError):
        configure_value_rules(profile, rows[:-1])
    with pytest.raises(ValueError):
        configure_relations(profile, [{'검사': '이하', '기준 항목': '없는칸', '비교 항목': '합계'}])


def test_saved_file_tampering_with_a_numeric_cell_is_detected_independently(tmp_path):
    path, profile = amount_form(tmp_path)
    output = fill_compatible_template(path, amounts(), tmp_path / 'saved.docx', profile=profile)
    with ZipFile(output) as archive:
        parts = [(item, archive.read(item.filename)) for item in archive.infolist()]
    with ZipFile(output, 'w') as archive:
        for item, data in parts:
            archive.writestr(item, data.replace('0.3원'.encode(), '9원'.encode())
                            if item.filename == 'word/document.xml' else data)
    with pytest.raises(ValueError, match='불일치'):
        verify_output(path, output, amounts(), profile=profile)


class TypedMappingClient:
    def generate_json(self, prompt, payload):
        return {'fields': [dict(id=field['id'], kind=field['kind'], value_key=field['label'],
                                required=False, input_required=False, input_mode='source_grounded',
                                max_chars=100, confidence=.9, validation={'type': 'date'})
                           for field in payload['fields']]}


def test_unknown_form_date_proposal_is_confirmed_and_cached_without_values(tmp_path):
    path = tmp_path / 'new.docx'
    doc = Document()
    doc.add_paragraph('{{신청일}}')
    doc.save(path)
    cache = tmp_path / 'cache'
    profile = learn_template(path, TypedMappingClient(), cache_dir=cache)
    assert profile['fields'][0]['validation'] == {'type': 'date'}
    assert profile['fields'][0]['narrative_style_required'] is False
    with pytest.raises(ValueError, match='확인'):
        save_learned_profile(path, profile, cache_dir=cache)
    profile['values'] = {'신청일': 'PRIVATE'}
    target = save_learned_profile(path, profile, cache_dir=cache, user_confirmed=True)
    assert 'PRIVATE' not in target.read_text(encoding='utf-8')
    restored = load_learned_profile(path, cache)
    assert restored['fields'][0]['validation'] == {'type': 'date'}
    with pytest.raises(ValueError):
        fill_compatible_template(path, {'신청일': '2026-02-30'}, tmp_path / 'bad.docx', profile=restored)


def test_bad_rule_profile_fails_before_model_payload_can_use_it():
    profile = {'fields': [{'value_key': '인원', 'validation': {'type': 'integer', 'min': float('nan')}}]}
    with pytest.raises(ValueError):
        profile_fields(profile)
    with pytest.raises(ValueError, match='공백'):
        profile_fields({'fields': [{'value_key': ' 인원 ', 'validation': {'type': 'integer'}}]})


def test_chunk_context_has_exact_offsets_and_changes_source_proof():
    text = '4.2 Posology\nThe initial dose is 500 mg twice daily.\nHowever, a lower dose may be assessed.'
    document = {'파일명': '규정.pdf', '본문': text, '표 목록': [],
                '페이지/시트 정보': [{'페이지': 3, '본문': text, '표 목록': []}]}
    sources = chunk_documents([document])
    assert all(source['context_text'][source['context_start']:source['context_end']] == source['text']
               for source in sources)
    initial = next(source for source in sources if '500 mg' in source['text'])
    changed = deepcopy(document)
    changed['페이지/시트 정보'][0]['본문'] = text.replace('lower dose', 'higher dose')
    second = next(source for source in chunk_documents([changed]) if '500 mg' in source['text'])
    assert initial['source_id'] != second['source_id']
    assert evidence_fingerprint({'본문': 'dose'}, [initial]) != evidence_fingerprint({'본문': 'dose'}, [second])
    assert initial['page'] == second['page'] == 3


def test_calendar_editor_checks_separate_year_month_day_without_filling_missing_values(tmp_path):
    path = tmp_path / 'split-date.docx'
    doc = Document()
    for key in ('년', '월', '일'):
        doc.add_paragraph('{{' + key + '}}')
    doc.save(path)
    profile = analyze_template(path)
    for field in profile['fields']:
        field['validation'] = {'type': 'integer'}
    configured = configure_relations(profile, [{'검사': '년·월·일', '기준 항목': '년',
        '비교 항목': '월 | 일', '허용 오차': ''}])
    assert relation_rows(configured)[0]['비교 항목'] == '월 | 일'
    values = {'년': '2024', '월': '2', '일': '29'}
    output = fill_compatible_template(path, values, tmp_path / 'leap.docx', profile=configured)
    assert verify_output(path, output, values, profile=configured)['status'] == 'passed'
    with pytest.raises(ValueError, match='날짜|달력'):
        fill_compatible_template(path, values | {'년': '2025'}, tmp_path / 'bad.docx', profile=configured)
    with pytest.raises(ValueError):
        configure_relations(profile, [{'검사': '년·월·일', '기준 항목': '년', '비교 항목': '월'}])

