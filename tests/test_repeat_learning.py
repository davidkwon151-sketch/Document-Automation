from copy import deepcopy
import json

from docx import Document
import pytest

from agent.template_learning import learn_template, save_learned_profile, load_learned_profile
from templates import analyze_template
from templates.repeat_fields import repeat_profile
from templates.repeat_rows import inspect_repeat_tables, prepare_repeat_template, expansion_binding
from templates.repeat_rules import inherit_repeat_rules


def inherited_form(tmp_path, *, count=2):
    document = Document()
    document.add_paragraph('신청인: {{신청인}}')
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = '비용 내역'
    table.rows[0].cells[1].text = '금액(원)'
    source = tmp_path / 'source.docx'
    document.save(source)
    original = analyze_template(source)
    for field in original['fields']:
        field.update(required=True, max_chars=20, input_required=False,
                     input_mode='source_grounded')
        if field['kind'] == 'placeholder':
            field.update(input_required=True, input_mode='user_provided')
        elif '/w:tc[2]' in field['id']:
            field['validation'] = {'type': 'integer', 'min': 0}
    plan = {'table_id': inspect_repeat_tables(source)[0]['id'], 'row': 2, 'count': count}
    prepared = tmp_path / 'prepared.docx'
    expansion = prepare_repeat_template(source, prepared, plan)
    raw = repeat_profile(prepared, plan, expansion_binding(expansion))
    profile = inherit_repeat_rules(source, raw, original, prepared_path=prepared)
    return source, prepared, original, profile


class MappingClient:
    def __init__(self, profile, *, mutation=None):
        self.fields = {field['id']: field for field in profile['fields']}
        self.mutation = mutation
        self.calls = []

    def generate_json(self, prompt, payload):
        self.calls.append(deepcopy(payload))
        result = []
        for context in payload['fields']:
            field = self.fields[context['id']]
            proposed = {key: deepcopy(field[key]) for key in (
                'id', 'kind', 'required', 'input_required', 'max_chars')}
            proposed.update(value_key='확인_' + field['value_key'],
                            input_mode=field.get('input_mode', 'source_grounded'), confidence=0.98)
            if field.get('validation'):
                proposed['validation'] = deepcopy(field['validation'])
            if self.mutation == 'length':
                proposed['max_chars'] = field['max_chars'] + 1
            elif self.mutation == 'required':
                proposed['required'] = False
            result.append(proposed)
        return {'fields': result}


def test_inherited_mapping_keeps_constraints_and_invalidates_changed_rule_cache(tmp_path):
    source, prepared, original, profile = inherited_form(tmp_path)
    client = MappingClient(profile)
    learned = learn_template(prepared, client=client, force=True, base_profile=profile)
    assert learned['repeat_source_profile'] == profile['repeat_source_profile']
    assert all(field['max_chars'] == 20 for field in learned['fields'])
    assert all(field['value_key'].startswith('확인_') for field in learned['fields'])
    cache = tmp_path / 'cache'
    save_learned_profile(prepared, learned, cache_dir=cache, user_confirmed=True)
    fresh = MappingClient(profile)
    assert learn_template(prepared, client=fresh, cache_dir=cache, base_profile=profile)['learning']['origin'] == 'confirmed_cache'
    assert fresh.calls == []
    changed = deepcopy(original)
    for field in changed['fields']:
        field['max_chars'] = 18
    base = repeat_profile(prepared, profile['repeat_expansion']['plan'], profile['repeat_expansion'])
    updated = inherit_repeat_rules(source, base, changed, prepared_path=prepared)
    refresh = MappingClient(updated)
    changed_result = learn_template(prepared, client=refresh, cache_dir=cache, base_profile=updated)
    assert refresh.calls and changed_result['learning']['origin'] == 'ai_proposal'
    assert all(field['max_chars'] == 18 for field in changed_result['fields'])


@pytest.mark.parametrize('mutation', ['length', 'required'])
def test_ai_proposal_cannot_weaken_inherited_rules(tmp_path, mutation):
    _, prepared, _, profile = inherited_form(tmp_path)
    with pytest.raises(ValueError):
        learn_template(prepared, client=MappingClient(profile, mutation=mutation), force=True, base_profile=profile)


def test_recursive_source_snapshot_does_not_store_answers_or_local_paths(tmp_path):
    source, prepared, original, profile = inherited_form(tmp_path)
    original.update(values={'신청인': 'PRIVATE_PERSON'}, qa_values={'금액': 'PRIVATE_QA'},
                    source_path=str(source), output_path='PRIVATE_OUTPUT')
    raw = repeat_profile(prepared, profile['repeat_expansion']['plan'], profile['repeat_expansion'])
    profile = inherit_repeat_rules(source, raw, original, prepared_path=prepared)
    learned = learn_template(prepared, client=MappingClient(profile), force=True, base_profile=profile)
    path = save_learned_profile(prepared, learned, cache_dir=tmp_path / 'cache', user_confirmed=True)
    text = path.read_text(encoding='utf-8')
    for secret in ('PRIVATE_PERSON', 'PRIVATE_QA', 'PRIVATE_OUTPUT', str(tmp_path)):
        assert secret not in text
    stored = json.loads(text)['profile']
    assert stored['repeat_source_profile']['source_sha256'] == original['source_sha256']
    assert any(field.get('validation') for field in stored['fields'])


def test_count_one_mapping_never_overwrites_identical_sha_original_cache(tmp_path):
    source, prepared, original, profile = inherited_form(tmp_path, count=1)
    assert source.read_bytes() == prepared.read_bytes()
    cache = tmp_path / 'cache'
    original_path = save_learned_profile(source, original, cache_dir=cache, user_confirmed=True)
    original_cache_bytes = original_path.read_bytes()
    learned = learn_template(prepared, client=MappingClient(profile), force=True, base_profile=profile)
    repeated_path = save_learned_profile(prepared, learned, cache_dir=cache, user_confirmed=True)
    assert repeated_path != original_path
    assert original_path.read_bytes() == original_cache_bytes
    restored_original = load_learned_profile(source, cache)
    restored_repeat = load_learned_profile(prepared, cache, repeat_expansion=profile['repeat_expansion'])
    assert restored_original['fields'] == original['fields']
    assert restored_original.get('repeat_expansion') is None
    assert restored_repeat['repeat_expansion'] == profile['repeat_expansion']
    cached_client = MappingClient(profile)
    cached = learn_template(prepared, client=cached_client, cache_dir=cache, base_profile=profile)
    assert cached_client.calls == [] and cached['learning']['origin'] == 'confirmed_cache'
