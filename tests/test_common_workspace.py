"""Five general forms use the real source, review, template and tenant gates."""
from base64 import b64encode
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import time
from uuid import uuid4
from zipfile import ZipFile

from docx import Document
from fastapi.testclient import TestClient
import pytest

from app import web_api
from parsers import parse_file
from templates.common_presets import FORMS, common_profile, template_path


SECRET = 'test-only-common-secret-' + 'b' * 32


class Model:
    def embed(self, texts):
        return [[1.0, 0.0] for _ in texts]

    def generate_json(self, name, payload):
        if name == 'brief':
            return {'목적': '업무 확인', '보고 대상': '팀장', '보고서 유형': '결과보고서',
                    '마감': '', '분량': '1쪽', '부족한 정보': [], '질문': [], '양식 항목': {}}
        if name == 'draft':
            source = payload['sources'][0]
            line = f"□ {source['text']} [{source['source_id']}]"
            return {'제목': payload['template_profile']['common_context']['title'],
                    '요약': line, '본문': line}
        if name == 'review':
            return payload['draft']
        if name == 'boss_review':
            return {'questions': ['후속 작업은 무엇인가요?', '담당자는 누구인가요?', '추가 확인은 필요한가요?'],
                    'answers': []}
        if name == 'grounding':
            source = payload['sources'][0]
            return {'claims': [{'field': claim['field'], 'line': claim['line'], 'status': 'supported',
                                'evidence': [{'source_id': source['source_id'], 'quote': source['text']}]}
                               for claim in payload['claims']]}
        if name == 'completeness':
            return {'issues': [], 'checked_fields': sorted(payload['draft'])}
        raise AssertionError(name)


def call(api, method, path, payload=None, *, user='writer-a'):
    body = b'' if payload is None else json.dumps(payload, ensure_ascii=False).encode()
    stamp, nonce = str(int(time.time())), uuid4().hex
    headers = {'X-RA-User': user, 'X-RA-Time': stamp, 'X-RA-Nonce': nonce,
               'X-RA-Signature': web_api.signature(SECRET, method, path, stamp, nonce, user, body)}
    if payload is not None:
        headers['Content-Type'] = 'application/json'
    return api.request(method, path, content=body, headers=headers)


@pytest.mark.parametrize('form_id', FORMS)
def test_examples_have_exact_usable_fields(form_id):
    path = template_path(form_id)
    profile = common_profile(path, form_id)
    assert {item['value_key'] for item in profile['fields']} == {'제목', '요약', '본문'}
    assert profile['source_sha256'] == sha256(path.read_bytes()).hexdigest()
    assert FORMS[form_id][0] in '\n'.join(p.text for p in Document(path).paragraphs)


@pytest.mark.parametrize('form_id', FORMS)
def test_common_api_generates_checks_exports_and_isolates_user(tmp_path, form_id):
    api = TestClient(web_api.create_app(root=tmp_path / 'private', secret=SECRET, client_factory=Model),
                     raise_server_exceptions=False)
    created = call(api, 'POST', '/api/common/jobs', {'form_type': form_id,
        'instruction': '제공한 사실만 사용해 작성해줘', 'notes': '현장 점검을 완료함', 'sources': []})
    assert created.status_code == 200, created.text
    identifier = created.json()['job_id']
    assert created.json()['form_title'] == FORMS[form_id][0]
    assert call(api, 'GET', f'/api/common/jobs/{identifier}', user='writer-b').status_code == 404
    response = call(api, 'POST', f'/api/common/jobs/{identifier}/generate', {})
    assert response.status_code == 202, response.text
    for _ in range(200):
        state = call(api, 'GET', f'/api/common/jobs/{identifier}').json()
        if state['status'] != 'processing':
            break
        time.sleep(.02)
    assert state['status'] == 'ready', state
    assert state['result']['draft']['제목'] == FORMS[form_id][0]
    assert state['evidence'][0]['document_sha256'] == sha256('현장 점검을 완료함'.encode()).hexdigest()
    assert call(api, 'POST', f'/api/common/jobs/{identifier}/export', {'confirmed': True},
                user='writer-b').status_code == 404
    exported = call(api, 'POST', f'/api/common/jobs/{identifier}/export', {'confirmed': True})
    assert exported.status_code == 200, exported.text
    with ZipFile(BytesIO(exported.content)) as archive:
        doc = Document(BytesIO(archive.read(f'{form_id}_draft.docx')))
        text = '\n'.join(paragraph.text for paragraph in doc.paragraphs)
        assert FORMS[form_id][0] in text and '현장 점검을 완료함' in text
        evidence = json.loads(archive.read('출처와_검수기록.json'))
        assert evidence['sources'][0]['document_sha256'] == sha256('현장 점검을 완료함'.encode()).hexdigest()
    assert call(api, 'DELETE', f'/api/common/jobs/{identifier}').status_code == 200


def test_custom_form_requires_exact_placeholders_and_key_is_never_saved(tmp_path):
    api = TestClient(web_api.create_app(root=tmp_path / 'private', secret=SECRET, client_factory=Model),
                     raise_server_exceptions=False)
    bad = Document()
    bad.add_paragraph('{{제목}}')
    buffer = BytesIO(); bad.save(buffer)
    response = call(api, 'POST', '/api/common/jobs', {'form_type': 'weekly_report',
        'instruction': '작성', 'notes': '업무를 완료함', 'sources': [],
        'template': {'name': 'company.docx', 'base64': b64encode(buffer.getvalue()).decode()}})
    assert response.status_code == 422
    assert not list((tmp_path / 'private').rglob('state.json'))
    created = call(api, 'POST', '/api/common/jobs', {'form_type': 'weekly_report',
        'instruction': '작성', 'notes': '업무를 완료함', 'sources': []})
    identifier = created.json()['job_id']
    call(api, 'POST', f'/api/common/jobs/{identifier}/generate', {'gemini_api_key': 'TEST_KEY_ONLY'})
    saved = next((tmp_path / 'private').rglob('state.json')).read_text(encoding='utf-8')
    assert 'TEST_KEY_ONLY' not in saved


def test_changed_source_or_unreviewed_revision_cannot_be_downloaded(tmp_path):
    root = tmp_path / 'private'
    api = TestClient(web_api.create_app(root=root, secret=SECRET, client_factory=Model),
                     raise_server_exceptions=False)
    created = call(api, 'POST', '/api/common/jobs', {'form_type': 'weekly_report',
        'instruction': '이번 주 업무를 보고해줘', 'notes': '현장 점검을 완료함', 'sources': []})
    identifier = created.json()['job_id']
    call(api, 'POST', f'/api/common/jobs/{identifier}/generate', {})
    for _ in range(200):
        state = call(api, 'GET', f'/api/common/jobs/{identifier}').json()
        if state['status'] != 'processing':
            break
        time.sleep(.02)
    assert state['status'] == 'ready'
    draft = dict(state['result']['draft'])
    draft['제목'] = '주간보고서 수정본'
    response = call(api, 'POST', f'/api/common/jobs/{identifier}/review', {'draft': draft})
    assert response.status_code == 202
    for _ in range(200):
        state = call(api, 'GET', f'/api/common/jobs/{identifier}').json()
        if state['status'] != 'processing':
            break
        time.sleep(.02)
    assert state['status'] == 'ready', state
    assert state['result']['draft']['제목'] == '주간보고서 수정본'
    source = root / sha256(b'writer-a').hexdigest() / identifier / 'sources' / '사용자_입력_메모.txt'
    source.write_text('다른 사실을 입력함', encoding='utf-8')
    assert call(api, 'POST', f'/api/common/jobs/{identifier}/export', {'confirmed': True}).status_code == 409


@pytest.mark.parametrize('path', [template_path('weekly_report'),
                                      Path('templates/generic_document.hwpx')])
def test_valid_uploaded_company_form_keeps_its_format_and_fills_original_slots(tmp_path, path):
    api = TestClient(web_api.create_app(root=tmp_path / 'private', secret=SECRET, client_factory=Model),
                     raise_server_exceptions=False)
    original = path.read_bytes()
    created = call(api, 'POST', '/api/common/jobs', {
        'form_type': 'weekly_report', 'instruction': '확인된 업무를 보고해줘',
        'notes': '현장 점검을 완료함', 'sources': [],
        'template': {'name': 'company' + path.suffix, 'base64': b64encode(original).decode()}})
    assert created.status_code == 200, created.text
    identifier = created.json()['job_id']
    assert created.json()['template_origin'] == 'user_uploaded'
    assert call(api, 'POST', f'/api/common/jobs/{identifier}/generate', {}).status_code == 202
    for _ in range(200):
        state = call(api, 'GET', f'/api/common/jobs/{identifier}').json()
        if state['status'] != 'processing':
            break
        time.sleep(.02)
    assert state['status'] == 'ready', state
    exported = call(api, 'POST', f'/api/common/jobs/{identifier}/export', {'confirmed': True})
    assert exported.status_code == 200, exported.text
    with ZipFile(BytesIO(exported.content)) as archive:
        filled = tmp_path / ('filled' + path.suffix)
        filled.write_bytes(archive.read('weekly_report_draft' + path.suffix))
        text = parse_file(filled)['본문']
        assert '현장 점검을 완료함' in text
        assert '{{본문}}' not in text
        evidence = json.loads(archive.read('출처와_검수기록.json'))
        assert evidence['template_sha256'] == sha256(original).hexdigest()
