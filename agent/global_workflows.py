"""Source-confirmed trade documents; no inferred prices, FX, parties or approvals.

The bundled forms are project examples, not KITA or company official forms.
Confirmed company profiles can use the same preparation/export functions.
"""
from copy import deepcopy
from decimal import Decimal, localcontext
from hashlib import sha256
import json
from pathlib import Path
import re
from tempfile import TemporaryDirectory

from agent.brief import model_profile, profile_fields
from agent.business import inspect_business_draft
from agent.field_citations import is_selection_field, split_field_citations
from agent.ra_workflows import _bound_quote, _source_records
from templates.value_rules import inspect_form_values, parse_scalar

OFFICIAL_REFERENCES = [
    {'title': '한국무역협회 상업송장', 'url': 'https://www.kita.net/board/format/formatDetail.do?postIndex=1720212',
     'checked_on': '2026-10-04', 'supports': '견적송장과 상업송장 역할 및 대표 기재 항목'},
    {'title': '한국무역협회 포장명세서', 'url': 'https://www.kita.net/board/format/formatDetail.do?postIndex=1863940',
     'checked_on': '2026-10-04', 'supports': '포장 단위별 물품·수량·순중량·총중량 항목'},
    {'title': '관세청 포장명세서 용어', 'url': 'https://www.customs.go.kr/kcs/ad/tr/trTermView.do?mi=2902&termId=3303',
     'checked_on': '2026-10-04', 'supports': '송장과 포장명세서의 물품 명세·규격 단위 대조'},
]
GLOBAL_WORKFLOWS = {
    'proforma_invoice': {'title': 'Proforma Invoice', 'label': '견적송장', 'stage': '견적', 'business_workflow': 'trade_sales'},
    'commercial_invoice': {'title': 'Commercial Invoice', 'label': '상업송장', 'stage': '거래 문서', 'business_workflow': 'trade_sales'},
    'packing_list': {'title': 'Packing List', 'label': '포장명세서', 'stage': '포장 명세', 'business_workflow': 'trade_sales'},
    'overseas_activity': {'title': '해외영업 활동보고', 'label': '해외영업 활동보고', 'stage': '내부 활동보고', 'business_workflow': 'overseas_business'},
}
TRANSACTION = re.compile(r'(?im)(?:^|[;|])\s*(?:Order\s+No\.?|PO\s+No\.?|거래번호|주문번호)\s*[:：=]?\s*([^\r\n;|]+)')
SENSITIVE = re.compile(r'서명|동의|승인|계좌|수취인|signature|signed by|consent|account|beneficiary', re.I)


def _fingerprint(value):
    return sha256(json.dumps(value, ensure_ascii=False, sort_keys=True).encode()).hexdigest()


def _spec(key, *, aliases=(), required=True, direct=False, numeric=False):
    field = {'value_key': key, 'label': key, 'source_labels': [key, *aliases], 'required': required,
             'input_required': direct, 'narrative_style_required': False, 'max_chars': 30 if numeric else 400}
    if numeric:
        field['validation'] = {'type': 'number', 'min': 0, 'decimal_places': 8}
    return field


def workflow_spec(workflow, rows=1):
    """Return explicit field/row contracts; priority is representative, not a frequency rank."""
    if workflow not in GLOBAL_WORKFLOWS or type(rows) is not int or not 1 <= rows <= 20:
        raise ValueError('대표 업무와 품목 행 수 1~20을 선택해야 함')
    common = [_spec('Order No.', aliases=('거래번호', '주문번호')),
              _spec('Document Date', aliases=('작성일', '발행일'))]
    common[1]['validation'] = {'type': 'date', 'date_formats': ['YYYY-MM-DD']}
    relations, groups, item_rows = [], [], []
    if workflow == 'overseas_activity':
        fields = [common[0], _spec('보고 기간'), _spec('결론'), _spec('주요 활동'), _spec('다음 계획'),
                  _spec('보고 대상', direct=True), _spec('작성자', required=False, direct=True)]
        for field in fields:
            if field['value_key'] in {'결론', '주요 활동', '다음 계획'}:
                field.update(narrative_style_required=True, max_chars=1800)
    else:
        fields = [*common, _spec('Seller', aliases=('수출자',)), _spec('Buyer', aliases=('수입자',)),
                  _spec('Destination', aliases=('목적지',)), _spec('Port of Loading', aliases=('선적항',), required=False),
                  _spec('Signed By', aliases=('서명',), required=False, direct=True)]
        money = workflow != 'packing_list'
        if money:
            fields.extend([_spec('Currency', aliases=('통화',)), _spec('Incoterms', aliases=('인도조건',)),
                           _spec('Named Place', aliases=('지정장소',)), _spec('Payment Terms', aliases=('결제조건',))])
        columns = [('Description', '품명'), ('Quantity', '수량'), ('Quantity Unit', '수량단위')]
        columns += [('Unit Price', '단가'), ('Amount', '금액')] if money else [
            ('Packages', '포장개수'), ('Net Weight', '순중량'), ('Gross Weight', '총중량'), ('Weight Unit', '중량단위')]
        for row in range(1, rows + 1):
            item = {}
            for column, korean in columns:
                key = f'{column}[{row}]'
                item[column] = key
                fields.append(_spec(key, aliases=(f'{korean}[{row}]',), required=False,
                                    numeric=column in {'Quantity', 'Unit Price', 'Amount', 'Packages', 'Net Weight', 'Gross Weight'}))
            item_rows.append(item)
            groups.append({'kind': 'all_or_none', 'fields': list(item.values())})
            if not money:
                relations.append({'kind': 'less_equal', 'left': item['Net Weight'], 'right': item['Gross Weight']})
        if money:
            fields.append(_spec('Total Amount', aliases=('총금액',), numeric=True))
            # Only used rows enter this sum; empty optional rows never become 0 values.
        else:
            fields.extend([_spec('Shipping Marks', aliases=('화인',), required=False),
                           _spec('Package Type', aliases=('포장종류',), required=False)])
    return {'global_workflow': workflow, 'domain': 'business_support',
            'business_workflow': GLOBAL_WORKFLOWS[workflow]['business_workflow'],
            'document_kind': 'report' if workflow == 'overseas_activity' else 'other',
            'template_origin': 'project_example', 'citation_mode': 'sidecar',
            'fields': fields, 'global_item_rows': item_rows,
            'constraints': {'groups': groups, 'relations': relations},
            'official_references': deepcopy(OFFICIAL_REFERENCES), 'usage_frequency': 'not_measured'}


def create_global_template(workflow, path, *, rows=1):
    """Create a new example DOCX and bind field rules to real blank input locations."""
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from templates.compatibility import analyze_template
    target = Path(path)
    if target.suffix.lower() != '.docx' or target.exists():
        raise ValueError('새 DOCX 사본 경로가 필요함; 기존 양식을 덮어쓰지 않음')
    spec = workflow_spec(workflow, rows)
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.left_margin = section.right_margin = Inches(.75)
    normal = document.styles['Normal']
    normal.font.name, normal.font.size = '맑은 고딕', Pt(10)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
    document.styles['Title'].font.color.rgb = RGBColor(0, 0, 0)
    document.styles['Title'].font.underline = False
    for border in document.styles['Title'].element.xpath('./w:pPr/w:pBdr'):
        border.getparent().remove(border)
    document.add_paragraph(GLOBAL_WORKFLOWS[workflow]['title'], 'Title')
    document.add_paragraph('프로젝트 제작 예제 양식입니다. 확인한 거래 원자료의 지정 항목을 기입하며, 서명과 제출 승인은 담당자가 확인합니다.')
    row_keys = {key for row in spec['global_item_rows'] for key in row.values()}
    header_fields = [field for field in spec['fields'] if field['value_key'] not in row_keys]
    table = document.add_table(rows=0, cols=2)
    table.style = 'Table Grid'
    table.columns[0].width, table.columns[1].width = Inches(2), Inches(5)
    def light_borders(table):
        borders = OxmlElement('w:tblBorders')
        for name in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            edge = OxmlElement('w:' + name)
            for attribute, value in (('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')):
                edge.set(qn('w:' + attribute), value)
            borders.append(edge)
        table._tbl.tblPr.append(borders)
    light_borders(table)
    def input_cell(cell, key):
        # Native blank SDTs let optional fields/unused rows stay genuinely blank.
        # Ordinary {{tokens}} are mandatory in the existing placeholder engine.
        control, properties, content = (OxmlElement('w:' + name) for name in ('sdt', 'sdtPr', 'sdtContent'))
        alias = OxmlElement('w:alias'); alias.set(qn('w:val'), key)
        properties.append(alias); properties.append(OxmlElement('w:text'))
        run, text = OxmlElement('w:r'), OxmlElement('w:t')
        text.text = ' '; text.set(qn('xml:space'), 'preserve')
        run.append(text); content.append(run)
        control.append(properties); control.append(content)
        cell.paragraphs[0]._p.append(control)
    for field in header_fields:
        cells = table.add_row().cells
        cells[0].text = field['label']
        input_cell(cells[1], field['value_key'])
    if spec['global_item_rows']:
        document.add_paragraph('Items')
        columns = list(spec['global_item_rows'][0])
        table = document.add_table(rows=1, cols=len(columns))
        table.style = 'Table Grid'
        light_borders(table)
        for cell, column in zip(table.rows[0].cells, columns):
            cell.text = column
            shading = OxmlElement('w:shd'); shading.set(qn('w:fill'), 'E7EEF5')
            cell._tc.get_or_add_tcPr().append(shading)
        repeat = OxmlElement('w:tblHeader')
        table.rows[0]._tr.get_or_add_trPr().append(repeat)
        for item in spec['global_item_rows']:
            for cell, key in zip(table.add_row().cells, item.values()):
                input_cell(cell, key)
    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(target)
    analyzed = analyze_template(target)
    keyed = {field['value_key']: field for field in spec['fields']}
    fields = [dict(field, **keyed[field['value_key']]) for field in analyzed['fields'] if field['value_key'] in keyed]
    if len(fields) != len(keyed) or len({field['id'] for field in fields}) != len(fields):
        raise ValueError('예제 양식의 고유 입력 위치가 선언 항목과 다름')
    return {**analyzed, **spec, 'fields': fields}


def _structured_values(source, records):
    """Pair actual CSV headers and a physical data row, including numbered headers."""
    if (not source.get('filename', '').lower().endswith(('.csv', '.tsv'))
            or not re.match(r'^행 [2-9][0-9]*,|^행 1[0-9]+,', source.get('location', ''))):
        return []
    headers = [record for record in records.values() if record.get('document_sha256') == source.get('document_sha256')
               and record.get('filename') == source.get('filename') and record.get('sheet') == source.get('sheet')
               and record.get('location', '').startswith('행 1,')]
    if len(headers) != 1:
        return []
    labels = [value.strip() for value in headers[0]['text'].split('|')]
    parts = [match for match in re.finditer(r'[^|]+', source.get('text', ''))]
    if len(labels) != len(parts) or len(set(labels)) != len(labels):
        return []
    return [(label, match[0].strip(), match.start() + len(match[0]) - len(match[0].lstrip()))
            for label, match in zip(labels, parts)]


def _transaction_scope(source, records, selected):
    if any(record.get('filename') == source.get('filename')
           and record.get('document_sha256') != source.get('document_sha256') for record in records.values()):
        return False  # A reused name cannot silently merge two different originals.
    structured = _structured_values(source, records)
    if structured:
        explicit = {value for label, value, _ in structured if label in {'Order No.', 'Order No', 'PO No.', '거래번호', '주문번호'}}
        return explicit == {selected}
    # Explicit document-wide declarations are used only when the same original
    # SHA contains one transaction; a filename or metadata label is not evidence.
    declared = {match[1].strip() for record in records.values()
                if record.get('document_sha256') == source.get('document_sha256')
                for match in TRANSACTION.finditer(record.get('context_text', record.get('text', '')))}
    return declared == {selected}


def _global_profile(profile):
    cleaned = model_profile(profile)
    fields = profile_fields(cleaned)
    if (not cleaned or cleaned.get('global_workflow') not in GLOBAL_WORKFLOWS or not fields
            or len({field['value_key'] for field in fields}) != len(fields)
            or len({field.get('id') for field in fields}) != len(fields)):
        raise ValueError('고유 원위치와 확인된 매핑을 가진 글로벌 업무 양식이 필요함')
    selected = cleaned.get('global_transaction_id')
    if not isinstance(selected, str) or not selected.strip() or selected != selected.strip():
        raise ValueError('실제 원자료와 대조할 거래번호를 명시해야 함')
    return cleaned, fields


def _global_sources(sources):
    from agent.multimodal_intake import validate_generation_source, validate_source_origin
    records = _source_records(sources)
    for source in records.values():
        validate_source_origin(source)
        # Existing ordinary parser records remain valid. Accepted multimodal
        # sources carry a receipt/fingerprint and must preserve its exact text.
        if 'verification_fingerprint' in source or 'verification_receipt' in source:
            validate_generation_source(source)
    return records


def _check_bound_role(field, binding, source, records, fields):
    """A real source ID cannot authorize Seller→Buyer or a different item row."""
    start = binding['start']
    end = start + len(binding['quote'])
    roles, full_values = set(), set()
    aliases = {label: item['value_key'] for item in fields
               for label in item.get('source_labels', [item['label']])}
    for label, value, position in _structured_values(source, records):
        if label in aliases and start < position + len(value) and end > position:
            roles.add(aliases[label])
            if start == position and end == position + len(value):
                full_values.add(aliases[label])
    context = source.get('context_text', source['text'])
    offset = source.get('context_start', 0)
    begin, finish = start + offset, end + offset
    for label, role in aliases.items():
        delimiter = r'\s+' if '/행 ' in source.get('location', '') else r'\s*[:：=|]\s*'
        prefix = r'(?:^|;)\s*' if '/행 ' in source.get('location', '') else r'(?m)^\s*'
        value_pattern = r'(?P<value>[^\r\n;]+)' if '/행 ' in source.get('location', '') else r'(?P<value>[^\r\n]+)'
        for match in re.finditer(prefix + re.escape(label) + delimiter + value_pattern, context):
            if begin < match.end('value') and finish > match.start('value'):
                roles.add(role)
                raw = match['value']
                position = match.start('value') + len(raw) - len(raw.lstrip())
                if begin == position and finish == position + len(raw.strip()):
                    full_values.add(role)
    if roles != {field['value_key']}:
        raise ValueError('인용한 원문 항목·거래 당사자 역할·품목 행이 대상 입력칸과 다름')
    if full_values != {field['value_key']}:
        raise ValueError('회사·당사자명과 기입값은 해당 원문 항목의 전체 값을 보존해야 함; 앞부분·일부 수치만 기입하지 않음')


def propose_global_bindings(profile, sources):
    """Offer exact labelled values in one transaction scope, without applying them."""
    cleaned, fields = _global_profile(profile)
    records = _global_sources(sources)
    results, bindings = [], {}
    for field in fields:
        key = field['value_key']
        candidates = []
        if field.get('input_required') or SENSITIVE.search(field['label']) or is_selection_field(field):
            results.append({'value_key': key, 'status': 'direct_input', 'candidates': []})
            continue
        labels = field.get('source_labels', [field['label']])
        if not isinstance(labels, list) or any(not isinstance(label, str) or not label for label in labels):
            raise ValueError('확인한 원문 항목명 목록이 필요함')
        for source in records.values():
            text = source.get('text', '')
            if not isinstance(text, str) or not _transaction_scope(source, records, cleaned['global_transaction_id']):
                continue
            for label, quote, start in _structured_values(source, records):
                if label in labels and quote:
                    binding = {'source_id': source['source_id'], 'quote': quote, 'start': start}
                    try:
                        _bound_quote(binding, source)
                    except ValueError:
                        continue
                    candidates.append(binding)
            for label in labels:
                patterns = [re.compile(r'(?m)^\s*' + re.escape(label) + r'\s*[:：=]\s*(?P<value>[^\r\n]+)'),
                            re.compile(r'(?m)^\s*' + re.escape(label) + r'\s*\|\s*(?P<value>[^|\r\n]+)$')]
                if '/행 ' in source.get('location', ''):
                    # The existing table chunker binds actual header labels to
                    # values with " ; ". Plain prose semicolons are not split.
                    patterns = [re.compile(r'(?:^|;)\s*' + re.escape(label) + r'\s+(?P<value>[^;\r\n]+)')]
                for pattern in patterns:
                    for match in pattern.finditer(text):
                        raw, quote = match['value'], match['value'].strip()
                        binding = {'source_id': source['source_id'], 'quote': quote,
                                   'start': match.start('value') + len(raw) - len(raw.lstrip())}
                        try:
                            _bound_quote(binding, source)
                        except ValueError:
                            continue
                        candidates.append(binding)
        candidates = list({(item['source_id'], item['start'], item['quote']): item for item in candidates}.values())
        status = 'proposed' if len(candidates) == 1 else 'ambiguous' if candidates else 'missing'
        if status == 'proposed':
            bindings[key] = candidates[0]
        results.append({'value_key': key, 'status': status, 'candidates': candidates})
    return {'fields': results, 'source_bindings': bindings, 'requires_confirmation': True,
            'confirmed': False, 'actual_model_requests': 0, 'usage_frequency': 'not_measured'}


def _trade_relations(values, profile):
    issues = []
    def issue(code, key, message):
        issues.append({'code': code, 'field': key, 'severity': 'error', 'message': message})
    rows = profile.get('global_item_rows', [])
    known = {field['value_key'] for field in profile['fields']}
    if (not isinstance(rows, list) or any(not isinstance(row, dict) or any(key not in known for key in row.values()) for row in rows)
            or len([key for row in rows for key in row.values()]) != len({key for row in rows for key in row.values()})):
        raise ValueError('원래 품목 행 계약이 잘못됨')
    used = [row for row in rows if any(values.get(key, '').strip() for key in row.values())]
    if rows and not used:
        issue('global_no_items', '', '실제 품목 행을 최소 1개 제공해야 함')
    if profile['global_workflow'] in {'proforma_invoice', 'commercial_invoice'}:
        currency = values.get('Currency', '')
        if not re.fullmatch(r'[A-Z]{3}', currency):
            issue('global_currency', 'Currency', '통화를 명시적 세 글자 코드로 확인해야 함; $와 환율을 추정하지 않음')
        totals = []
        for row in used:
            if not all(values.get(row.get(name, '')) for name in ('Quantity', 'Unit Price', 'Amount')):
                continue  # Required group check reports missing values.
            try:
                numbers = [parse_scalar(values[row[name]], {'type': 'number', 'min': 0})
                           for name in ('Quantity', 'Unit Price', 'Amount')]
            except ValueError:
                continue
            with localcontext() as context:
                context.prec = max(28, sum(len(number.as_tuple().digits) for number in numbers) + 4)
                if numbers[0] * numbers[1] != numbers[2]:
                    issue('global_line_amount', row['Amount'], '제공된 수량×단가와 금액이 다름; 값은 임의 수정하지 않음')
            totals.append(numbers[2])
        if len(totals) == len(used) and used and values.get('Total Amount'):
            try:
                total = parse_scalar(values['Total Amount'], {'type': 'number', 'min': 0})
                with localcontext() as context:
                    context.prec = max(28, sum(len(number.as_tuple().digits) for number in [*totals, total]) + 4)
                    if sum(totals, Decimal(0)) != total:
                        issue('global_total_amount', 'Total Amount', '제공된 총금액과 품목 금액 합계가 다름; 공란·차액을 자동 생성하지 않음')
            except ValueError:
                pass
    if profile['global_workflow'] == 'overseas_activity':
        for key in ('결론', '주요 활동', '다음 계획'):
            for line in values.get(key, '').splitlines():
                if line.strip() and (not re.match(r'^[□○-]\s*', line.strip()) or not line.rstrip().endswith(('함', '임'))):
                    issue('global_report_style', key, '보고 문장은 원자료에 근거한 개조식·함/임 종결로 제공해야 함')
        if len([line for line in values.get('결론', '').splitlines() if line.strip()]) > 3:
            issue('global_summary_length', '결론', '결론 요약은 3줄 이내로 제공해야 함')
    return issues


def prepare_global_workflow(profile, sources, source_bindings=None, direct_values=None):
    """Prepare cited JSON with exact original spans, scalar/row/business checks."""
    cleaned, fields = _global_profile(profile)
    records = _global_sources(sources)
    keyed = {field['value_key']: field for field in fields}
    bindings, direct = source_bindings or {}, direct_values or {}
    if (not isinstance(bindings, dict) or not isinstance(direct, dict) or set(bindings) - keyed.keys()
            or set(direct) - keyed.keys() or set(bindings) & set(direct)):
        raise ValueError('존재하는 항목에 연결·직접 입력 중 한 방식만 적용해야 함')
    draft, evidence, locked = {key: '' for key in keyed}, {}, {}
    for key, binding in bindings.items():
        field = keyed[key]
        if field.get('input_required') or SENSITIVE.search(field['label']) or is_selection_field(field):
            raise ValueError('직접 확인·서명·선택 항목을 자동 기입할 수 없음')
        if not isinstance(binding, dict) or binding.get('source_id') not in records:
            raise ValueError('실제 출처 ID로 연결해야 함')
        source = records[binding['source_id']]
        quote, start = _bound_quote(binding, source)
        if not _transaction_scope(source, records, cleaned['global_transaction_id']):
            raise ValueError('원문의 실제 거래번호가 선택 거래와 다르거나 여러 거래가 섞임')
        _check_bound_role(field, {'quote': quote, 'start': start}, source, records, fields)
        draft[key] = '\n'.join(line + f" [{source['source_id']}]" if line.strip() else line for line in quote.split('\n'))
        evidence[key] = {'source_id': source['source_id'], 'quote': quote, 'start': start,
                         'end': start + len(quote), 'field_id': field['id'], 'label': field['label'], 'source': deepcopy(source)}
    for key, value in direct.items():
        if not keyed[key].get('input_required') or not isinstance(value, str):
            raise ValueError('직접 입력으로 선언한 칸의 실제 문자열만 허용함')
        if not value.strip():
            continue
        identifier = 'SU' + _fingerprint((key, value))[:20]
        if identifier in records:
            raise ValueError('직접 입력 출처 ID가 원자료와 충돌함')
        records[identifier] = {'source_id': identifier, 'filename': '사용자 입력', 'source_type': 'user_input',
                               'location': key, 'text': keyed[key]['label'] + ': ' + value}
        draft[key] = '\n'.join(line + f' [{identifier}]' if line.strip() else line for line in value.split('\n'))
        locked[key] = {'value': value, 'source_id': identifier}
    plain = {key: split_field_citations(value, keyed[key], literal=locked.get(key, {}).get('value'))[0]
             for key, value in draft.items()}
    literals = {key: item['value'] for key, item in locked.items()}
    issues = inspect_form_values(draft, cleaned, literal_values=literals) + _trade_relations(plain, cleaned)
    for key, field in keyed.items():
        if field.get('max_chars') and len(plain[key]) > field['max_chars']:
            issues.append({'code': 'global_field_length', 'field': key, 'severity': 'error', 'message': '원칸 분량을 넘음; 내용 삭제 없이 보완해야 함'})
    # Native code-looking selections are literal data and already checked above.
    business = inspect_business_draft({key: value for key, value in draft.items() if not is_selection_field(keyed[key])},
                                     list(records.values()), profile=cleaned)
    issues += business['issues']
    missing = [key for key, field in keyed.items() if field.get('required', True) and not plain[key].strip()]
    blocking = bool(missing or any(issue['severity'] == 'error' for issue in issues))
    return {'mode': 'verified_field_copy', 'draft': draft, 'plain_values': plain, 'evidence': evidence,
            'locked_fields': locked, 'sources': list(records.values()), 'template_profile': cleaned,
            'review': {'blocking': blocking, 'issues': issues}, 'missing_fields': missing,
            'questions': [key + '의 실제 원자료 또는 직접 입력값을 제공해 주시겠습니까?' for key in missing[:2]],
            'ready_for_output_check': not blocking, 'submission_ready': False, 'actual_model_requests': 0,
            'fingerprint': _fingerprint((cleaned, draft, evidence, list(records.values()), locked)),
            'scope': '확인한 거래 원문의 양식 기입임; AI 생성·통관 승인·빈도 순위·실사용 수정률은 별도 검증함'}


def export_global_workflow(template, profile, sources, source_bindings=None, direct_values=None):
    """Re-prepare and independently verify original positions immediately before download."""
    from agent.output_check import verify_output
    from templates.compatibility import fill_compatible_template
    source = Path(template)
    original = source.read_bytes()
    if profile.get('source_sha256') != sha256(original).hexdigest():
        raise ValueError('원본 양식 SHA가 확인 매핑과 다름')
    prepared = prepare_global_workflow(profile, sources, source_bindings, direct_values)
    if not prepared['ready_for_output_check']:
        raise ValueError('출처·필수 항목·거래·수치·관계 오류를 먼저 해결해야 함')
    with TemporaryDirectory(prefix='global-output-') as directory:
        target = Path(directory) / ('작성본' + source.suffix)
        fill_compatible_template(source, prepared['plain_values'], target, profile=prepared['template_profile'])
        proof = verify_output(source, target, prepared['plain_values'], profile=prepared['template_profile'])
        data = target.read_bytes()
        if (source.read_bytes() != original or proof['status'] != 'passed'
                or proof['sha']['template'] != profile['source_sha256']
                or proof['sha']['output'] != sha256(data).hexdigest()):
            raise ValueError('출력과 독립 원위치 검수 증거가 다름')
    sidecar = {**prepared, 'output_verification': proof, 'native_visual_review': 'not_run'}
    return {'document': data, 'filename': profile['global_workflow'] + source.suffix,
            'evidence': json.dumps(sidecar, ensure_ascii=False, indent=2).encode('utf-8'),
            'output_verification': proof}


def blank_global_input(workflow, rows=1):
    """A blank labelled input sheet; no invented seller, quantity, approval or fact."""
    spec = workflow_spec(workflow, rows)
    return ('\n'.join(field['source_labels'][0] + ': ' for field in spec['fields'] if not field['input_required']) + '\n').encode('utf-8')
