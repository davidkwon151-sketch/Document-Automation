"""저장된 결과물을 다시 열어 값과 원본 서식 보존을 독립 검사함."""

from collections import Counter
from copy import deepcopy
from decimal import Decimal
from hashlib import sha256
from math import ceil
from pathlib import Path
import json
import re
from zipfile import ZipFile

from lxml import etree
from openpyxl import load_workbook
from pypdf import PdfReader
from pypdf.generic import ArrayObject, ContentStream, DecodedStreamObject, IndirectObject, NameObject, TextStringObject

from parsers.extract import _check_zip, _xml
from templates.compatibility import _selected_fields, analyze_template
from templates.value_rules import inspect_form_values, mapped_rule_profile


TOKEN = re.compile(r"\{\{\s*([^{}]+?)\s*\}\}")


def _require(condition, message):
    if not condition:
        raise ValueError("출력 검증 실패: " + message)


def _normalized(text):
    return re.sub(r"\s+", "", str(text))


def _text(node):
    """Word/HWPX 줄바꿈도 문자열로 복원함. 채우기 함수는 사용하지 않음."""
    result = []
    for element in node.iter():
        local = etree.QName(element).localname
        if local == "t":
            result.append(element.text or "")
        elif local in {"br", "cr", "lineBreak"}:
            result.append("\n")
        elif local == "tab":
            result.append("\t")
        elif local in {'nbSpace', 'fwSpace', 'hyphen'}:
            result.append({'nbSpace': ' ', 'fwSpace': '　', 'hyphen': '-'}[local])
        if element.getparent() is not None and etree.QName(element.getparent()).localname == "t":
            result.append(element.tail or "")
    return "".join(result)


def _office_placeholder_text(node, replacements=None, *, retain_controls=True):
    """독립 검사에서 네이티브/잠금 SDT의 표시 문자열을 {{...}}로 재해석하지 않음."""
    special = "*[local-name()='sdt'][*[local-name()='sdtPr']/*[local-name()='dropDownList' or local-name()='comboBox' or local-name()='lock' or local-name()='date' or local-name()='picture' or local-name()='group' or local-name()='repeatingSection' or local-name()='repeatingSectionItem']]"
    if node.xpath("ancestor::" + special):
        return (_text(node) if retain_controls else ''), []
    original_text = _text(node)
    clone = deepcopy(node)
    protected = {}
    for index, content in enumerate(clone.xpath(".//" + special + "/*[local-name()='sdtContent']")):
        texts = content.xpath(".//*[local-name()='t']")
        if not texts:
            continue
        marker = f'\ue000DOCXCONTROL{index}\ue001'
        while marker in original_text:
            marker += '\ue000'
        protected[marker] = _text(content)
        for text in texts:
            text.text = ''
        texts[0].text = marker
    text = _text(clone)
    tokens = TOKEN.findall(text)
    if replacements is not None:
        text = TOKEN.sub(replacements if callable(replacements) else lambda match: replacements.get(match[1].strip(), match[0]), text)
    if retain_controls:
        for marker, original in protected.items():
            text = text.replace(marker, original)
    return text, tokens


def _office_source_matches(paragraph, kind):
    """Mask control text at its original length; source offsets remain exact."""
    if kind != 'docx':
        return list(TOKEN.finditer(_text(paragraph)))
    special = "*[local-name()='sdt'][*[local-name()='sdtPr']/*[local-name()='dropDownList' or local-name()='comboBox' or local-name()='lock' or local-name()='dataBinding' or local-name()='date' or local-name()='picture' or local-name()='checkbox' or local-name()='group' or local-name()='repeatingSection' or local-name()='repeatingSectionItem']]"
    if paragraph.xpath('ancestor::' + special):
        return []
    clone = deepcopy(paragraph)
    for content in clone.xpath('.//' + special + "/*[local-name()='sdtContent']"):
        for node in content.xpath(".//*[local-name()='t']"):
            node.text = ' ' * len(node.text or '')
    return list(TOKEN.finditer(_text(clone)))


def _contains(actual, expected, label):
    _require(_normalized(expected) in _normalized(actual), f"입력값 누락 또는 변경: {label}")
    # 개조식의 줄 구분을 공백 정규화로 숨기지 않음.
    if "\n" in expected:
        lines = [_normalized(line) for line in actual.splitlines()]
        _require(all(any(_normalized(line) in present for present in lines)
                     for line in expected.splitlines() if line.strip()), f"줄바꿈 누락: {label}")


def _target(root, field):
    _, _, xpath = field["id"].split(":", 2)
    if field['kind'] in {'docx_repeat_placeholder', 'hwpx_repeat_placeholder'}:
        family, part, _ = field['id'].split(':', 2)
        xpath = field.get('xml_path')
        expected = 'repeat_' + field['kind'].split('_', 1)[0]
        _require(family == expected and isinstance(xpath, str)
                 and field['id'] == f"{expected}:{part}:{xpath}#slot:{field.get('placeholder_start')}",
                 '반복 자리표시자 위치 기록 불일치')
    nodes = root.xpath(xpath, namespaces={k: v for k, v in root.nsmap.items() if k})
    _require(len(nodes) == 1, f"입력칸 구조 변경: {field['label']}")
    return nodes[0]


def _same_position(root, original):
    """원본의 태그별 형제 순번으로 찾음. 새 문단이 뒤에 추가되어도 기존 위치는 고정됨."""
    path = []
    node = original
    while node.getparent() is not None:
        parent = node.getparent()
        siblings = [child for child in parent if child.tag == node.tag]
        path.append((node.tag, siblings.index(node)))
        node = parent
    _require(root.tag == node.tag, "원본 XML 루트 변경")
    for tag, index in reversed(path):
        children = [child for child in root if child.tag == tag]
        _require(index < len(children), "원본 입력 문단 위치 손실")
        root = children[index]
    return root


def _estimated_lines(value, width, size):
    capacity = max(width / size, 1)
    return sum(max(1, ceil(sum(1 if ord(char) > 127 else 0.5 for char in line) / capacity))
               for line in value.splitlines() or [""])


def _docx_height(target, value, label):
    rows = target.xpath("ancestor-or-self::*[local-name()='tr']")
    if not rows:
        return
    heights = rows[-1].xpath("./*[local-name()='trPr']/*[local-name()='trHeight']")
    if not heights or not any(key.endswith("}hRule") and val == "exact" for key, val in heights[0].attrib.items()):
        return
    height = float(next(val for key, val in heights[0].attrib.items() if key.endswith("}val"))) / 20
    cells = target.xpath("ancestor-or-self::*[local-name()='tc']")
    widths = cells[-1].xpath("./*[local-name()='tcPr']/*[local-name()='tcW']") if cells else []
    sizes = target.xpath(".//*[local-name()='rPr']/*[local-name()='sz']/@*[local-name()='val']")
    size = min(map(float, sizes)) / 2 if sizes else 11
    lines = len(value.splitlines()) or 1
    if widths and sizes:
        width = float(next(val for key, val in widths[0].attrib.items() if key.endswith("}w"))) / 20
        lines = _estimated_lines(value, width, size)
    _require(lines * size <= height * 1.6, f"DOCX 고정 행 높이에 입력 내용이 명백히 넘침: {label}")


def _properties(root, kind):
    names = {"rPr", "pPr", "tblPr", "tblGrid", "trPr", "tcPr", "sectPr"} if kind == "docx" else {
        "cellAddr", "cellSpan", "cellSz", "cellMargin", "pagePr", "secPr"}
    serialized = Counter()
    for node in root.iter():
        local = etree.QName(node).localname
        if local in names:
            serialized[(local, etree.tostring(node, method="c14n", exclusive=True, with_comments=False))] += 1
        if local in {"tbl", "tr", "tc", "p", "r", "run"}:
            serialized[(local, tuple(sorted(node.attrib.items())))] += 1
    return serialized


def _run_position_properties(original, filled, kind):
    """Word/HWPX는 개행을 run 내부에 넣으므로 기존 run 위치는 바뀌지 않음."""
    for node in original.iter():
        local=etree.QName(node).localname
        if local not in {'r','run','p'}:
            continue
        actual=_same_position(filled,node)
        _require(dict(node.attrib)==dict(actual.attrib),f"{kind.upper()} 원본 문단·런 위치 속성 변경")
        property_name='pPr' if local=='p' else 'rPr'
        old=[etree.tostring(child,method='c14n',exclusive=True) for child in node if etree.QName(child).localname==property_name]
        new=[etree.tostring(child,method='c14n',exclusive=True) for child in actual if etree.QName(child).localname==property_name]
        _require(old==new,f"{kind.upper()} 원본 문단·런 위치의 글꼴·서식 변경")


def _pptx_paragraph_properties(original, filled, name):
    """원본 child 순서를 지키며 동일 run 서식의 새 줄+run 쌍만 허용함."""
    def skeleton(node):
        local=etree.QName(node).localname
        if local in {'r','fld'}:
            properties=tuple(etree.tostring(child,method='c14n',exclusive=True) for child in node if etree.QName(child).localname!='t')
            return node.tag,tuple(sorted(node.attrib.items())),properties
        return etree.tostring(node,method='c14n',exclusive=True)
    old,new=list(original),list(filled)
    if not any(etree.QName(node).localname in {'r','fld'} for node in old) and not _text(original).strip():
        # 처음부터 비어 있던 문단의 새 run은 끝 문단 속성의 글꼴을 상속함.
        run=etree.Element('{http://schemas.openxmlformats.org/drawingml/2006/main}r')
        end=next((node for node in old if etree.QName(node).localname=='endParaRPr'),None)
        if end is not None:
            properties=deepcopy(end)
            properties.tag='{http://schemas.openxmlformats.org/drawingml/2006/main}rPr'
            run.append(properties)
        old.insert(len(old)-int(end is not None),run)
    budget=sum(etree.QName(node).localname=='br' for node in new)-sum(etree.QName(node).localname=='br' for node in old)
    index=0
    for expected in old:
        while index<len(new) and skeleton(expected)!=skeleton(new[index]):
            _require(budget>0 and index>0 and index+1<len(new) and etree.QName(new[index]).localname=='br'
                     and etree.QName(new[index+1]).localname=='r' and etree.QName(new[index-1]).localname=='r',
                     f"PPTX 원본 런 위치의 글꼴·서식 변경: {name}")
            _require(skeleton(new[index+1])==skeleton(new[index-1]),f"PPTX 새 개행 런의 글꼴·서식 변경: {name}")
            previous=new[index-1].find('{http://schemas.openxmlformats.org/drawingml/2006/main}rPr')
            br_properties=list(new[index])
            _require(len(br_properties)==int(previous is not None) and (previous is None or etree.tostring(previous,method='c14n',exclusive=True)==etree.tostring(br_properties[0],method='c14n',exclusive=True)),
                     f"PPTX 새 줄 나눔의 글꼴·서식 변경: {name}")
            index+=2
            budget-=1
        _require(index<len(new),f"PPTX 원본 런 속성 손실: {name}")
        index+=1
    _require(index==len(new) and budget==0,f"PPTX 원본 런 순서 또는 서식 변경: {name}")


def _office_values(before, after, kind, selected):
    replacements = {field["label"]: value for field, value in selected if field["kind"] == "placeholder"}
    legacy_roots, legacy_paragraphs = {}, {}
    repeat_kinds = {'docx_repeat_placeholder', 'hwpx_repeat_placeholder'}
    repeat_slots = {}
    for field, value in selected:
        if field['kind'] not in repeat_kinds:
            continue
        _require(field['kind'] == kind + '_repeat_placeholder', '반복 입력 형식 불일치')
        part = field['id'].split(':', 2)[1]
        original = _target(_xml(before[part]), field)
        _require(etree.QName(original).localname == 'p', '반복 자리표시자 원본 문단 손실')
        text = _text(original)
        start, end = field.get('placeholder_start'), field.get('placeholder_end')
        _require(text == field.get('anchor_text') and type(start) is int and type(end) is int,
                 '반복 자리표시자 원문 불일치')
        _require(any(match.start() == start and match.end() == end and match[1].strip() == field.get('placeholder_key')
                     for match in _office_source_matches(original, kind)), '반복 자리표시자 원본 범위 불일치')
        slots = repeat_slots.setdefault((part, field['xml_path']), {})
        _require((start, end) not in slots, '반복 자리표시자 중복 매핑')
        slots[(start, end)] = value
    for field, value in selected:
        if field["kind"] == "placeholder":
            continue
        part = field["id"].split(":", 2)[1]
        target = _target(_xml(after[part]), field)
        if kind == 'docx':
            original_target = _target(_xml(before[part]), field)
            ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
            native = original_target.xpath("ancestor-or-self::w:sdt[w:sdtPr/w:dropDownList or w:sdtPr/w:comboBox] | .//w:sdt[w:sdtPr/w:dropDownList or w:sdtPr/w:comboBox]", namespaces=ns)
            if native and field['kind'] not in repeat_kinds:
                _require(len(native) == 1 and original_target.tag == '{'+ns['w']+'}sdtContent'
                         and original_target.getparent() is native[0], f"원본 선택 컨트롤 밖의 입력 위치 지정: {field['label']}")
                combo = bool(native[0].xpath('./w:sdtPr/w:comboBox', namespaces=ns))
                _require(field['kind'] == ('docx_combobox' if combo else 'docx_choice')
                         and field.get('control_type') == ('combobox' if combo else 'choice'),
                         f"원본 선택 컨트롤 종류와 프로파일 불일치: {field['label']}")
        if field['kind'] in repeat_kinds:
            _contains(_text(target), value, field['label'])
            if kind == 'docx':
                _docx_height(target, value, field['label'])
            continue  # Entire paragraph and every unchanged neighbour checked below.
        if field["kind"] == "docx_legacy_text":
            from templates.compatibility import _legacy_text_region
            nodes = _legacy_text_region(target)
            _require(nodes is not None and ''.join(node.text or '' for node in nodes) == value,
                     f"FORMTEXT 결과 값 또는 필드 구조 불일치: {field['label']}")
            expected_root = legacy_roots.setdefault(part,_xml(before[part]))
            expected_begin = _target(expected_root,field)
            expected_nodes = _legacy_text_region(expected_begin)
            _require(bool(expected_nodes),f"FORMTEXT 원본 필드 구조 손실: {field['label']}")
            expected_nodes[0].text=value
            for node in expected_nodes[1:]:
                node.text=''
            paragraph=expected_begin.xpath("ancestor::*[local-name()='p'][1]")[0]
            legacy_paragraphs[(part,paragraph.getroottree().getpath(paragraph))]=paragraph
            for tag in ('ffData','instrText','fldChar'):
                original_paragraph=target.xpath("ancestor::*[local-name()='p'][1]")[0]
                _require([etree.tostring(node,method='c14n',exclusive=True) for node in paragraph.xpath(f".//*[local-name()='{tag}']")]==
                         [etree.tostring(node,method='c14n',exclusive=True) for node in original_paragraph.xpath(f".//*[local-name()='{tag}']")],
                         f"FORMTEXT 지시문·보호·필드 속성 변경: {field['label']}")
        elif field["kind"] == "docx_checkbox":
            checked = target.getparent().xpath("./*[local-name()='sdtPr']/*[local-name()='checkbox']/*[local-name()='checked']")
            _require(len(checked) == 1 and next(iter(checked[0].attrib.values()), None) == ("1" if value == "true" else "0"),
                     f"체크박스 상태 불일치: {field['label']}")
            expected_root = legacy_roots.setdefault(part, _xml(before[part]))
            expected_target = _target(expected_root, field)
            state_name = 'checkedState' if value == 'true' else 'uncheckedState'
            state = expected_target.getparent().xpath(f"./*[local-name()='sdtPr']/*[local-name()='checkbox']/*[local-name()='{state_name}']")
            try:
                code = state[0].get('{http://schemas.microsoft.com/office/word/2010/wordml}val') if state else None
                glyph = chr(int(code, 16)) if code is not None else ('☒' if value == 'true' else '☐')
            except (ValueError, TypeError) as exc:
                raise ValueError('출력 검증 실패: 원본 체크 표시 문자 오류') from exc
            _require(_text(target) == glyph, f"체크박스 표시 문자 불일치: {field['label']}")
            nodes = expected_target.xpath(".//*[local-name()='t']")
            _require(bool(nodes), '원본 체크박스 표시 구조 손실')
            nodes[0].text = glyph
            for node in nodes[1:]:
                node.text = ''
            paragraphs = expected_target.xpath("ancestor::*[local-name()='p'][1]") or expected_target.xpath(".//*[local-name()='p']")
            if paragraphs:
                paragraph = paragraphs[0]
                legacy_paragraphs[(part, paragraph.getroottree().getpath(paragraph))] = paragraph
        elif field["kind"] in {"docx_choice", "docx_combobox"}:
            # 원본 XML을 직접 읽음. 채우기 helper의 선택/표시 계산을 재사용하지 않음.
            original_root = _xml(before[part])
            original_target = _target(original_root, field)
            original_properties = original_target.getparent().find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sdtPr")
            actual_properties = target.getparent().find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sdtPr")
            _require(original_properties is not None and actual_properties is not None, f"선택 컨트롤 속성 손실: {field['label']}")
            original_choices = original_properties.xpath("./*[local-name()='dropDownList' or local-name()='comboBox']")
            actual_choices = actual_properties.xpath("./*[local-name()='dropDownList' or local-name()='comboBox']")
            _require(len(original_choices) == len(actual_choices) == 1, f"선택 컨트롤 정의 중복 또는 손실: {field['label']}")
            original_choice, actual_choice = original_choices[0], actual_choices[0]
            word = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
            combo = etree.QName(original_choice).localname == 'comboBox'
            _require(field['kind'] == ('docx_combobox' if combo else 'docx_choice'), f"선택 컨트롤 유형 변경: {field['label']}")
            _require(not original_properties.xpath("./*[local-name()='lock' or local-name()='dataBinding']")
                     and not original_target.xpath("ancestor::*[local-name()='sdt']/*[local-name()='sdtPr']/*[local-name()='lock']"), f"잠김/XML 연결 선택 컨트롤 변경: {field['label']}")
            _require(not _text(original_target).strip() or bool(original_properties.xpath("./*[local-name()='showingPlcHdr']")), f"기존 선택값 변경: {field['label']}")
            if 'word/settings.xml' in before:
                protection = _xml(before['word/settings.xml']).xpath(".//*[local-name()='documentProtection']")
                _require(not any(node.get(word+'enforcement') in {'1','true','on'} for node in protection), f"DOCX 양식 보호 범위 밖의 선택 입력: {field['label']}")
            _require(bool(value.strip()) and not any(c in value for c in '\r\n\t'), f"선택 입력값은 한 줄이어야 함: {field['label']}")
            items = [{'value': item.get(word+'value'), 'label': item.get(word+'displayText', item.get(word+'value'))}
                     for item in original_choice]
            _require(all(item['value'] and item['label'] for item in items)
                     and len({item['value'] for item in items}) == len(items)
                     and len({item['label'] for item in items}) == len(items), f"원본 선택 항목이 유효하지 않음: {field['label']}")
            _require(field.get('options') == [item['value'] for item in items] and field.get('choice_items') == items
                     and field.get('allow_custom') is combo, f"원본 선택 정의와 프로파일 불일치: {field['label']}")
            matches = [item['label'] for item in items if item['value'] == value]
            _require(bool(matches) or combo, f"목록에 없는 선택값: {field['label']}")
            display = matches[0] if matches else value
            _require(_text(target) == display, f"선택 표시 문구 불일치: {field['label']}")
            _require(actual_choice.get(word+'lastValue') == value, f"선택 lastValue 불일치: {field['label']}")
            for properties in (original_properties, actual_properties):
                for choice in properties.xpath("./*[local-name()='dropDownList' or local-name()='comboBox']"):
                    choice.attrib.pop(word+'lastValue', None)
                for placeholder in properties.xpath("./*[local-name()='showingPlcHdr']"):
                    properties.remove(placeholder)
            _require(etree.tostring(original_properties, method='c14n', exclusive=True) == etree.tostring(actual_properties, method='c14n', exclusive=True),
                     f"원본 옵션·alias·스타일·선택 속성 변경: {field['label']}")
            expected_root = legacy_roots.setdefault(part, _xml(before[part]))
            expected_target = _target(expected_root, field)
            nodes = expected_target.xpath(".//*[local-name()='t']")
            _require(bool(nodes), f"원본 선택 표시 문자 구조 손실: {field['label']}")
            nodes[0].text = display
            for node in nodes[1:]:
                node.text = ''
            paragraphs = expected_target.xpath("ancestor::*[local-name()='p'][1]") or expected_target.xpath(".//*[local-name()='p']")
            if paragraphs:
                paragraph = paragraphs[0]
                legacy_paragraphs[(part, paragraph.getroottree().getpath(paragraph))] = paragraph
            _docx_height(target, display, field['label'])
        else:
            _contains(_text(target), value, field["label"])
            if field["kind"] == "docx_inline":
                expected = field["anchor_text"][:field["blank_start"]] + value + field["anchor_text"][field["blank_end"]:]
                _require(_text(target) == expected, f"입력칸의 주변 문구 또는 값 변경: {field['label']}")
            elif field["kind"] in {"docx_cell", "docx_sdt", "hwpx_cell"}:
                paragraphs = target.xpath("descendant-or-self::*[local-name()='p']")
                _require(_text(paragraphs[0] if paragraphs else target) == value, f"입력칸의 첫 문단 값 불일치: {field['label']}")
            if kind == "docx":
                _docx_height(target, value, field["label"])
    for (part,location),expected in legacy_paragraphs.items():
        if (part, location) in repeat_slots:
            continue  # Source-order slot replacement below also checks the control.
        actual = _same_position(_xml(after[part]),expected)
        expected_text = _office_placeholder_text(expected, replacements)[0] if kind == 'docx' else TOKEN.sub(lambda match:replacements.get(match[1].strip(),match[0]),_text(expected))
        _require(_text(actual)==expected_text,f"FORMTEXT 주변 안내·문단 값 불일치: {part}")
    for name, raw in before.items():
        if not (name.endswith(".xml") and (name.startswith("word/") if kind == "docx" else re.fullmatch(r"Contents/(section|masterpage)\d+\.xml", name))):
            continue
        original, filled = _xml(raw), _xml(after[name])
        original_paragraphs = original.xpath(".//*[local-name()='p']")
        filled_paragraphs = filled.xpath(".//*[local-name()='p']")
        changed_paragraphs = set()
        for field, _ in selected:
            if field["kind"] != "placeholder" and field["id"].split(":", 2)[1] == name:
                if field["kind"] == "docx_append":
                    continue  # 새 문단 추가는 기존 안내 문단을 수정할 권한이 아님.
                target = _target(original, field)
                changed_paragraphs.update(target.xpath("descendant-or-self::*[local-name()='p']"))
                changed_paragraphs.update(target.xpath("ancestor::*[local-name()='p']"))
        for paragraph in original_paragraphs:
            if paragraph not in changed_paragraphs and not TOKEN.search(_text(paragraph)):
                actual = _same_position(filled, paragraph)
                _require(_text(actual) == _text(paragraph),
                         f"원본 안내·항목명 텍스트 위치 손실 또는 변경: {name} {paragraph.getroottree().getpath(paragraph)}")
        for paragraph in original_paragraphs:
            text, tokens = _office_placeholder_text(paragraph) if kind == 'docx' else (_text(paragraph), TOKEN.findall(_text(paragraph)))
            if not tokens:
                continue
            location = paragraph.getroottree().getpath(paragraph)
            slots = repeat_slots.get((name, location))
            if slots:
                source_matches = iter(_office_source_matches(paragraph, kind))
                def replace_slot(match):
                    source_match = next(source_matches, None)
                    _require(source_match is not None and source_match[1].strip() == match[1].strip(),
                             '반복 문단의 원본 토큰 순서 변경')
                    span = (source_match.start(), source_match.end())
                    if span in slots:
                        return slots[span]
                    _require(match[1].strip() in replacements, f'반복 문단의 자리표시자 매핑 누락: {name} {location}')
                    return replacements[match[1].strip()]
                expected_node = _same_position(legacy_roots[name], paragraph) if kind == 'docx' and name in legacy_roots else paragraph
                expected = _office_placeholder_text(expected_node, replace_slot)[0] if kind == 'docx' else TOKEN.sub(replace_slot, text)
                _require(next(source_matches, None) is None, '반복 문단의 원본 토큰 누락')
                actual = _text(_same_position(filled, paragraph))
                _require(actual == expected, f'반복 문단 값 또는 주변 문구 불일치: {name} {location}')
                continue
            _require(all(token.strip() in replacements for token in tokens), f"원본 자리표시자 매핑 누락: {name}")
            expected_node = _same_position(legacy_roots[name], paragraph) if kind == 'docx' and name in legacy_roots else paragraph
            expected = _office_placeholder_text(expected_node, replacements)[0] if kind == 'docx' else TOKEN.sub(lambda match: replacements[match[1].strip()], text)
            location = paragraph.getroottree().getpath(paragraph)
            actual = _text(_same_position(filled, paragraph))
            _require(actual == expected, f"원본 문단 값 또는 주변 문구 불일치: {name} {location}")
            _contains(actual, expected, name)
        # 사용자 값의 문자 그대로 {{...}}는 허용하며 원본의 남은 토큰만 차단함.
        remaining = _office_placeholder_text(filled, retain_controls=False)[0] if kind == 'docx' else _text(filled)
        for _, value in selected:
            remaining = remaining.replace(value, "")
        original_keys = {key.strip() for paragraph in original_paragraphs for key in (_office_placeholder_text(paragraph)[1] if kind == 'docx' else TOKEN.findall(_text(paragraph)))}
        _require(not (original_keys & {key.strip() for key in TOKEN.findall(remaining)}), f"치환되지 않은 원본 자리표시자: {name}")


def _xlsx_checks(template, output, selected):
    from templates.xlsx_validation import XlsxLists, raw_text
    source, filled = load_workbook(template, data_only=False), load_workbook(output, data_only=False)
    pending = []
    try:
        _require(source.sheetnames == filled.sheetnames, "XLSX 시트 손실 또는 순서 변경")
        with ZipFile(template) as archive:
            lists = XlsxLists(archive)
            for field, value in selected:
                _, part, coordinate, *_ = field["id"].split(":", 3)
                try:
                    lists.check(part, coordinate, field, value)
                except ValueError as exc:
                    _require(False, str(exc))
            workbook = _xml(archive.read("xl/workbook.xml"))
            precision = workbook.xpath("./*[local-name()='calcPr']/@fullPrecision")
            displayed_precision = bool(precision and precision[0] in {'0', 'false', 'off'})
            sheets = workbook.xpath(".//*[local-name()='sheet']/@name")
            rels = _xml(archive.read("xl/_rels/workbook.xml.rels"))
            targets = {item.get("Id"): item.get("Target") for item in rels}
            part_sheets = {}
            for element, name in zip(workbook.xpath(".//*[local-name()='sheet']"), sheets):
                rid = next(value for key, value in element.attrib.items() if key.endswith("}id"))
                target = targets[rid].lstrip("/")
                part_sheets[target if target.startswith("xl/") else "xl/" + target] = name
            original_cells = {part: {cell.get('r'): cell for cell in lists.root(part).xpath(
                "./*[local-name()='sheetData']/*[local-name()='row']/*[local-name()='c']")}
                for part in part_sheets}
        groups = {}
        for field, value in selected:
            _, part, coordinate, *_ = field["id"].split(":", 3)
            groups.setdefault((part_sheets[part], coordinate), []).append((field, value))
        with ZipFile(output) as archive:
            output_cells = {part: {cell.get('r'): cell for cell in _xml(archive.read(part)).xpath(
                ".//*[local-name()='sheetData']/*[local-name()='row']/*[local-name()='c']")}
                for part in part_sheets}
        for name in source.sheetnames:
            old, new = source[name], filled[name]
            _require(set(map(str, old.merged_cells.ranges)) == set(map(str, new.merged_cells.ranges)), f"XLSX 병합 변경: {name}")
            for row in old:
                for cell in row:
                    actual = new[cell.coordinate]
                    part = next(part for part, sheet in part_sheets.items() if sheet == name)
                    created_choice = ((name, cell.coordinate) in groups and cell.coordinate not in original_cells[part]
                                      and groups[(name, cell.coordinate)][0][0].get('control_type') == 'choice')
                    if not created_choice:
                        _require(cell._style == actual._style, f"XLSX 셀 서식 변경: {name}!{cell.coordinate}")
                    if cell.data_type == "f":
                        # Array/data-table formulas are distinct Python objects on
                        # each read. Compare their original XML text and attributes.
                        original_xml = original_cells[part].get(cell.coordinate)
                        actual_xml = output_cells[part].get(cell.coordinate)
                        original_formula = original_xml.find('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}f') if original_xml is not None else None
                        actual_formula = actual_xml.find('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}f') if actual_xml is not None else None
                        _require(original_formula is not None and actual_formula is not None
                                 and etree.tostring(original_formula, method='c14n', exclusive=True)
                                 == etree.tostring(actual_formula, method='c14n', exclusive=True),
                                 f"XLSX 수식 변경: {name}!{cell.coordinate}")
                    elif (name, cell.coordinate) not in groups:
                        _require(cell.value == actual.value, f"XLSX 수정 대상 밖의 값 변경: {name}!{cell.coordinate}")
            for (sheet, coordinate), entries in groups.items():
                if sheet != name:
                    continue
                value = entries[0][1]
                field = entries[0][0]
                rule = field.get('validation', {})
                if field.get('control_type') == 'choice':
                    part = field['id'].split(':', 3)[1]
                    actual_xml = output_cells[part].get(coordinate)
                    _require(actual_xml is not None and actual_xml.get('t') == 'inlineStr'
                             and raw_text(actual_xml, []) == value,
                             f'XLSX 목록 선택값 또는 문자열 타입 불일치: {name}!{coordinate}')
                if entries[0][0]["kind"] == "xlsx_placeholder":
                    replacement = {field["label"]: value for field, value in entries}
                    value = TOKEN.sub(lambda match: replacement[match[1].strip()], str(old[coordinate].value))
                numeric = (field['kind'] == 'xlsx_cell' and rule.get('type') in {'integer', 'number'}
                           and (old[coordinate].value is None or not str(old[coordinate].value).strip())
                           and (not rule.get('unit') or rule.get('unit_location', 'value') == 'label')
                           and not re.search(r'\[S[A-Za-z0-9_-]+\]', value))
                if numeric:
                    # Re-read exact XML independently; openpyxl float conversion can hide
                    # a changed final digit. Do not call the filler's canonicalizer.
                    expected = Decimal(value.strip().replace(',', ''))
                    _require(not displayed_precision, f'Excel 표시 정밀도 설정의 숫자 입력은 검증할 수 없음: {name}!{coordinate}')
                    magnitude = expected.copy_abs()
                    significant = len(''.join(map(str, expected.as_tuple().digits)).rstrip('0')) or 1
                    _require(significant <= 15 and magnitude <= Decimal('9.99999999999999E307')
                             and (magnitude == 0 or magnitude >= Decimal('2.2251E-308')),
                             f'Excel 숫자 정밀도·저장 범위 위반: {name}!{coordinate}')
                    part = field['id'].split(':', 3)[1]
                    actual_xml = output_cells[part].get(coordinate)
                    _require(actual_xml is not None and actual_xml.get('t') == 'n',
                             f'XLSX 숫자 셀 타입 불일치: {name}!{coordinate}')
                    stored = actual_xml.xpath("./*[local-name()='v']/text()")
                    _require(len(stored) == 1 and re.fullmatch(r'[+-]?[0-9]+(?:\.[0-9]+)?', stored[0])
                             and not actual_xml.xpath("./*[local-name()='is']"),
                             f'XLSX 숫자 XML 구조 불일치: {name}!{coordinate}')
                    _require(Decimal(stored[0]) == expected, f'XLSX 숫자 값 불일치: {name}!{coordinate}')
                else:
                    _require(new[coordinate].value == value, f"XLSX 입력값 누락 또는 변경: {name}!{coordinate}")
                cell = old[coordinate]
                height = old.row_dimensions[cell.row].height
                merged = next((item for item in old.merged_cells.ranges if coordinate in item), None)
                size = cell.font.sz or 11
                lines = len(value.splitlines()) or 1
                if cell.alignment.wrap_text:
                    first = merged.min_col if merged else cell.column
                    last = merged.max_col if merged else cell.column
                    def column_width(column):
                        for dimension in old.column_dimensions.values():
                            if dimension.min is not None and dimension.max is not None and dimension.min <= column <= dimension.max:
                                return dimension.width
                        return old.sheet_format.defaultColWidth or 13
                    # Excel 열폭은 기본 글꼴 숫자 폭 단위임. 11pt 기본 글꼴의 7px를
                    # 보수적 근사로 환산하며 확정적인 세로 넘침만 차단함.
                    width = max(sum(column_width(column) for column in range(first,last+1))*5.25-6,1)
                    lines = _estimated_lines(value,width,size)
                if cell.alignment.shrink_to_fit or cell.alignment.textRotation:
                    pending.append(f"{name}!{coordinate}: 축소·회전 텍스트는 원본 프로그램에서 확인 필요함")
                    continue
                if height is not None:
                    if merged and merged.max_row > merged.min_row:
                        height = sum(old.row_dimensions[row].height or old.sheet_format.defaultRowHeight or 15
                                     for row in range(merged.min_row, merged.max_row + 1))
                    _require(lines * size <= height * 1.6,
                             f"XLSX 고정 행 높이에 입력 내용이 명백히 넘침: {name}!{coordinate}")
                    if lines*size > height or not cell.alignment.wrap_text:
                        pending.append(f"{name}!{coordinate}: 글꼴·가로 표시와 경계 분량은 원본 프로그램에서 확인 필요함")
                else:
                    pending.append(f"{name}!{coordinate}: 자동 행 높이의 표시 분량은 원본 프로그램에서 확인 필요함")
        return pending
    finally:
        source.close()
        filled.close()


def _pptx_checks(before, after, selected):
    replacements = {field["label"]: value for field, value in selected if field["kind"] == "placeholder"}
    for name, raw in before.items():
        if not re.fullmatch(r"ppt/slides/slide\d+\.xml", name):
            continue
        original, filled = _xml(raw), _xml(after[name])
        mutable = []
        for field, value in selected:
            if field["kind"] == "placeholder" or field["id"].split(":", 2)[1] != name:
                continue
            old_target, new_target = _target(original, field), _target(filled, field)
            paragraphs = old_target.xpath("descendant-or-self::*[local-name()='p']")
            new_paragraphs = new_target.xpath("descendant-or-self::*[local-name()='p']")
            _require(bool(paragraphs) and len(paragraphs) == len(new_paragraphs), f"PPTX 입력 문단 구조 변경: {field['label']}")
            actual = "\n".join(_text(p) for p in new_paragraphs)
            _require(actual == value, f"PPTX 입력칸 값 불일치: {name} {field['label']}")
            _contains(actual, value, field["label"])
            mutable.extend(paragraphs)
        for paragraph in original.xpath(".//*[local-name()='p']"):
            text = _text(paragraph)
            tokens = TOKEN.findall(text)
            if not tokens:
                continue
            _require(all(token.strip() in replacements for token in tokens), f"PPTX 원본 자리표시자 매핑 누락: {name}")
            expected = TOKEN.sub(lambda match: replacements[match[1].strip()], text)
            actual = _text(_same_position(filled, paragraph))
            _require(actual == expected, f"PPTX 원본 문단 값 또는 주변 문구 불일치: {name} {paragraph.getroottree().getpath(paragraph)}")
            _contains(actual, expected, name)
            mutable.append(paragraph)
        for paragraph in dict.fromkeys(mutable):
            new_paragraph = _same_position(filled, paragraph)
            _pptx_paragraph_properties(paragraph,new_paragraph,name)
            # 새 줄/런은 허용하되 기존 단락·런·폰트 속성은 삭제 또는 변경하지 않음.
            def formatting(node):
                return Counter(etree.tostring(item, method="c14n", exclusive=True) for item in node.xpath(
                    ".//*[local-name()='pPr' or local-name()='rPr' or local-name()='endParaRPr']"))
            _require(not (formatting(paragraph) - formatting(new_paragraph)), f"PPTX 문단·런·폰트 서식 변경: {name}")
            for node in (paragraph, new_paragraph):
                for child in list(node):
                    node.remove(child)
        _require(etree.tostring(original, method="c14n") == etree.tostring(filled, method="c14n"),
                 f"PPTX 수정 대상 밖의 슬라이드·도형·표·병합·위치 구조 변경: {name}")


def _zip_checks(template, output, kind, selected, checks):
    with ZipFile(template) as original, ZipFile(output) as result:
        _check_zip(original)
        _check_zip(result)
        _require(set(original.namelist()) == set(result.namelist()), "ZIP 부품 손실 또는 추가")
        before = {name: original.read(name) for name in original.namelist()}
        after = {name: result.read(name) for name in result.namelist()}
    changed = [name for name in before if before[name] != after[name]]
    for name in changed:
        editable = (kind == "docx" and re.fullmatch(r"word/(document|header\d+|footer\d+|footnotes|endnotes)\.xml", name)) or (
            kind == "hwpx" and (re.fullmatch(r"Contents/(section|masterpage)\d+\.xml", name) or name == "Preview/PrvText.txt")) or (
            kind == "xlsx" and re.fullmatch(r"xl/worksheets/sheet\d+\.xml", name)) or (
            kind == "pptx" and re.fullmatch(r"ppt/slides/slide\d+\.xml", name))
        _require(editable, f"글꼴·스타일·테마·그림·차트 또는 기타 부품 변경: {name}")
        if kind == "xlsx" and name.endswith(".xml"):
            from templates.xlsx_validation import XlsxLists, inherited_style
            from openpyxl.utils.cell import coordinate_to_tuple
            old_xml, new_xml = _xml(before[name]), _xml(after[name])
            selected_coordinates = {field["id"].split(":", 3)[2] for field, _ in selected
                                    if field["id"].split(":", 3)[1] == name}
            # Only a verified native list may materialize a previously absent cell.
            old_cells = {cell.get('r') for cell in old_xml.xpath("./*[local-name()='sheetData']/*[local-name()='row']/*[local-name()='c']")}
            created = selected_coordinates - old_cells
            if created:
                with ZipFile(template) as archive:
                    lists = XlsxLists(archive)
                    for coordinate in created:
                        entries = [(field, value) for field, value in selected
                                   if field['id'].split(':', 3)[1:3] == [name, coordinate]]
                        _require(len(entries) == 1, f'XLSX 신규 목록 셀의 매핑이 모호함: {coordinate}')
                        try:
                            native = lists.check(name, coordinate, *entries[0])
                            style = inherited_style(old_xml, coordinate)
                        except ValueError as exc:
                            _require(False, str(exc))
                        _require(native is not None, f'XLSX 승인되지 않은 신규 셀: {coordinate}')
                        cells = new_xml.xpath("./*[local-name()='sheetData']/*[local-name()='row']/*[local-name()='c'][@r=$coordinate]", coordinate=coordinate)
                        expected_attributes = {'r': coordinate, 't': 'inlineStr'}
                        if style is not None:
                            expected_attributes['s'] = style
                        _require(len(cells) == 1 and dict(cells[0].attrib) == expected_attributes,
                                 f'XLSX 신규 목록 셀 주소·타입·상속 서식 불일치: {coordinate}')
                        cell, row = cells[0], cells[0].getparent()
                        row_number = re.search(r'\d+$', coordinate)[0]
                        _require(row.get('r') == row_number, f'XLSX 신규 셀 행 주소 불일치: {coordinate}')
                        children = list(cell)
                        _require(len(children) == 1 and etree.QName(children[0]).localname == 'is',
                                 f'XLSX 신규 목록 셀 구조 불일치: {coordinate}')
                        rows = list(row.getparent())
                        _require([int(item.get('r')) for item in rows] == sorted(int(item.get('r')) for item in rows),
                                 f'XLSX 신규 목록 행 순서 불일치: {coordinate}')
                        cell_order = [coordinate_to_tuple(item.get('r')) for item in row]
                        _require(cell_order == sorted(cell_order), f'XLSX 신규 목록 셀 순서 불일치: {coordinate}')
                        row.remove(cell)
                        old_rows = old_xml.xpath("./*[local-name()='sheetData']/*[local-name()='row'][@r=$row]", row=row_number)
                        if not old_rows and not len(row):
                            _require(dict(row.attrib) == {'r': row_number}, f'XLSX 신규 목록 행 속성 변경: {coordinate}')
                            row.getparent().remove(row)
            for root in (old_xml, new_xml):
                for cell in root.xpath(".//*[local-name()='c']"):
                    if cell.get("r") in selected_coordinates:
                        cell.attrib.pop("t", None)
                        for child in list(cell):
                            if etree.QName(child).localname in {"v", "is"}:
                                cell.remove(child)
            _require(etree.tostring(old_xml, method="c14n") == etree.tostring(new_xml, method="c14n"),
                     f"XLSX 수식·병합·행 높이·열 너비·그림 참조 또는 기타 구조 변경: {name}")
        if name.endswith(".xml") and kind not in {"xlsx", "pptx"}:
            _run_position_properties(_xml(before[name]),_xml(after[name]),kind)
            old_properties, new_properties = _properties(_xml(before[name]), kind), _properties(_xml(after[name]), kind)
            _require(not (old_properties - new_properties), f"표·셀·병합·단락·글꼴·쪽 서식 변경: {name}")
    if kind == "xlsx":
        pending = _xlsx_checks(template, output, selected)
        checks.append({"name":"xlsx_layout_fit", "status":"pending" if pending else "passed", "warnings":pending,
                       "scope":"명시된 행 높이·병합 폭·줄바꿈의 보수적 추정이며 전체 렌더 검증은 아님"})
    elif kind == "pptx":
        _pptx_checks(before, after, selected)
    else:
        _office_values(before, after, kind, selected)
    checks.extend([{"name": "zip_parts_and_assets", "status": "passed", "changed_text_parts": changed},
                   {"name": "styles_tables_sections_formulas", "status": "passed"},
                   {"name": "saved_values", "status": "passed", "fields": len(selected)}])


def _pdf_field_name(widget):
    names = []
    while widget is not None:
        if widget.get("/T"):
            names.insert(0, str(widget["/T"]))
        parent = widget.get("/Parent")
        widget = parent.get_object() if parent else None
    return ".".join(names)


def _pdf_inherited_value(node, key, default=None):
    seen = set()
    while node is not None:
        _require(id(node) not in seen, 'PDF 선택 필드의 부모 연결이 순환함')
        seen.add(id(node))
        if key in node:
            return node[key]
        parent = node.get('/Parent')
        node = parent.get_object() if parent else None
    return default


def _pdf_choice_state(original, filled, value, label):
    """원본 PDF 옵션으로 저장값을 검증함. 작성기의 선택 해석을 재사용하지 않음."""
    flags = int(_pdf_inherited_value(original, '/Ff', 0))
    _require(not flags & 1, f'PDF 원본 읽기 전용 선택칸의 변경을 허용할 수 없음: {label}')
    combo, editable, multi = bool(flags & (1 << 17)), bool(flags & (1 << 18)), bool(flags & (1 << 21))
    _require(not (multi and combo) and not (editable and not combo), f'PDF 선택 플래그 조합을 확인할 수 없음: {label}')
    options = _pdf_inherited_value(original, '/Opt', [])
    _require(isinstance(options, list) and (bool(options) or combo and editable), f'PDF 선택 원본 옵션이 없음: {label}')
    pairs = []
    for item in options:
        pair = list(item) if isinstance(item, list) else [item, item]
        _require(len(pair) == 2 and all(isinstance(part, TextStringObject) for part in pair),
                 f'PDF 선택 원본 옵션 형식을 확인할 수 없음: {label}')
        pairs.append(tuple(str(part) for part in pair))
    codes = [pair[0] for pair in pairs]
    usable = [code for code, display in pairs if code.strip()]
    _require(all(code.strip() or not display.strip() for code, display in pairs)
             and all(code == code.strip() and display.strip() for code, display in pairs if code.strip())
             and (bool(usable) or combo and editable), f'PDF 원본 선택 코드·표시 문구가 모호함: {label}')
    _require(len(set(usable)) == len(usable), f'PDF 선택 export 코드가 중복됨: {label}')
    # The selected expected value is already canonical. Reinterpreting a
    # custom export literal such as [Scustom] here would erase actual data.
    if multi:
        try:
            requested = json.loads(value)
        except (TypeError, ValueError):
            requested = None
        _require(isinstance(requested, list) and bool(requested)
                 and all(type(code) is str and code in usable for code in requested)
                 and len(set(requested)) == len(requested), f'PDF 다중 선택 코드 배열이 잘못됨: {label}')
        indices = sorted(codes.index(code) for code in requested)
        expected = [codes[index] for index in indices]
        actual = _pdf_inherited_value(filled, '/V')
        _require(isinstance(actual, ArrayObject) and all(isinstance(code, TextStringObject) for code in actual)
                 and list(actual) == expected, f'PDF 다중 선택 canonical /V 불일치: {label}')
    else:
        _require(value in usable or combo and editable and value.strip(), f'PDF 원본 목록 밖의 선택값: {label}')
        indices = [codes.index(value)] if value in usable else []
        expected = [value]
        actual = _pdf_inherited_value(filled, '/V')
        _require(isinstance(actual, TextStringObject) and actual == value, f'PDF 선택 canonical /V 불일치: {label}')
    actual_indices = _pdf_inherited_value(filled, '/I')
    if indices:
        _require(isinstance(actual_indices, ArrayObject)
                 and all(type(index).__name__ == 'NumberObject' for index in actual_indices)
                 and list(actual_indices) == indices, f'PDF 선택 /I 인덱스 불일치: {label}')
    else:
        _require(actual_indices is None, f'PDF 자유 입력 선택값의 /I가 남음: {label}')
    return {'combo': combo, 'multi': multi, 'pairs': pairs, 'indices': indices,
            'display': [pairs[index][1] for index in indices] if indices else expected}


def _pdf_appearance_text(appearance, reader):
    """선택칸의 새 AP 텍스트만 독립 판독함. 보존된 배경 XObject는 제외함."""
    from pypdf import PageObject

    stream = ContentStream(appearance, reader)
    stream.operations = [(operands, operator) for operands, operator in stream.operations if operator != b'Do']
    content = DecodedStreamObject()
    content.set_data(stream.get_data())
    page = PageObject.create_blank_page(width=1000, height=1000)
    page[NameObject('/Contents')] = content
    page[NameObject('/Resources')] = appearance.get('/Resources', {}).get_object()
    return page.extract_text() or ''


def _pdf_choice_visible_bounds(appearance, reader, width, height, label):
    from io import BytesIO
    from pypdf import PageObject, PdfWriter
    import pypdfium2 as pdfium

    # 판독할 페이지에는 새 문자만 넣어 원래 배경 문자로 누락을 숨길 수 없게 함.
    content = ContentStream(appearance, reader)
    content.operations = [(args, op) for args, op in content.operations if op != b'Do']
    page = PageObject.create_blank_page(width=width, height=height)
    page[NameObject('/Contents')] = content
    page[NameObject('/Resources')] = appearance['/Resources']
    writer, buffer = PdfWriter(), BytesIO()
    writer.add_page(page)
    writer.write(buffer)
    document = pdfium.PdfDocument(buffer.getvalue())
    try:
        rendered = document[0]
        text = rendered.get_textpage()
        try:
            for index in range(text.count_chars()):
                if not text.get_text_range(index, 1).strip():
                    continue
                left, bottom, right, top = text.get_charbox(index)
                _require(-.01 <= left <= right <= width + .01 and -.01 <= bottom <= top <= height + .01,
                         f'PDF 선택 문구가 표시 영역에서 잘림: {label}')
        finally:
            text.close()
            rendered.close()
    finally:
        document.close()


def _pdf_choice_background(old_widget, new_widget, before, after, label):
    """선택 표시 밖 원래 appearance 배경과 그 자원은 그대로 보존해야 함."""
    def fingerprint(value, ancestors=frozenset()):
        value = value.get_object() if hasattr(value, 'get_object') else value
        if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
            return Decimal(str(value))
        if isinstance(value, (dict, list, tuple)):
            _require(id(value) not in ancestors, f'PDF 선택 배경 자원의 순환 연결: {label}')
            ancestors = ancestors | {id(value)}
            if isinstance(value, (list, tuple)):
                return tuple(fingerprint(item, ancestors) for item in value)
            properties = tuple((str(key), fingerprint(item, ancestors)) for key, item in sorted(value.items())
                               if key != '/Length' or not hasattr(value, 'get_data'))
            return (properties, value.get_data()) if hasattr(value, 'get_data') else properties
        return (type(value).__name__, str(value))

    old = old_widget.get('/AP', {}).get('/N')
    new = new_widget['/AP']['/N'].get_object()
    xobjects = new.get('/Resources', {}).get('/XObject', {})
    _require(set(xobjects) == ({'/OriginalAppearance'} if old is not None else set()),
             f'PDF 선택 원본 배경 그림 연결 변경: {label}')
    if old is None:
        return
    old, background = old.get_object(), xobjects['/OriginalAppearance'].get_object()
    original_stream = ContentStream(old, before)
    preserved, marked, in_text = [], 0, False
    for operands, operator in original_stream.operations:
        if operator in {b'BMC', b'BDC'} and (marked or operands and str(operands[0]) == '/Tx'):
            marked += 1
            continue
        if operator == b'EMC' and marked:
            marked -= 1
            continue
        if marked:
            continue
        if operator == b'BT':
            in_text = True
        elif operator == b'ET':
            in_text = False
        elif not in_text:
            preserved.append((operands, operator))
    _require(not marked and not in_text
             and fingerprint(preserved) == fingerprint(ContentStream(background, after).operations),
             f'PDF 선택 원본 배경 내용 변경: {label}')
    properties = lambda node: {key: value for key, value in node.items() if key not in {'/Length', '/Filter', '/DecodeParms'}}
    _require(fingerprint(properties(old)) == fingerprint(properties(background)),
             f'PDF 선택 원본 배경 자원·속성 변경: {label}')


def _pdf_choice_appearance(widget, reader, state, label):
    """새 선택 표시의 문구·가시 구간·강조 행을 저장 코드와 별도로 검사함."""
    appearance = widget['/AP']['/N'].get_object()
    width = float(widget['/Rect'][2] - widget['/Rect'][0])
    height = float(widget['/Rect'][3] - widget['/Rect'][1])
    bbox = [float(number) for number in appearance.get('/BBox', [])]
    _require(len(bbox) == 4 and all(abs(a - b) <= 1e-6 for a, b in zip(bbox, [0, 0, width, height]))
             and list(appearance.get('/Matrix', [1, 0, 0, 1, 0, 0])) == [1, 0, 0, 1, 0, 0],
             f'PDF 선택 appearance 위치·크기 불일치: {label}')
    operations = ContentStream(appearance, reader).operations
    marked, tx_count, tx_active = [], 0, False
    font_size, color, path, highlights, sizes, text_count = None, None, [], [], [], 0
    original_da = DecodedStreamObject()
    original_da.set_data(state['da'].encode('latin-1'))
    source_color, source_size = (b'g', (0.,)), None
    for operands, operator in ContentStream(original_da, reader).operations:
        if operator in {b'g', b'rg', b'k'}:
            source_color = (operator, tuple(float(number) for number in operands))
        elif operator == b'Tf':
            source_size = float(operands[1])
    for operands, operator in operations:
        if operator in {b'BMC', b'BDC'}:
            marked.append(tx_active)
            if operands and str(operands[0]) == '/Tx':
                _require(not tx_active, f'PDF 선택 표시 영역이 중첩됨: {label}')
                tx_count += 1
                tx_active = True
        elif operator == b'EMC':
            _require(bool(marked), f'PDF 선택 표시 영역 연결이 잘못됨: {label}')
            tx_active = marked.pop()
        elif operator == b'Do':
            _require(not tx_active and list(operands) == ['/OriginalAppearance'],
                     f'PDF 선택 표시를 다른 그림으로 대체함: {label}')
        elif operator == b'cm':
            _require(list(operands) == [1, 0, 0, 1, 0, 0], f'PDF 선택 문구의 표시 변환 변경: {label}')
        elif operator == b'Tf':
            font_size = float(operands[1])
        elif operator in {b'g', b'rg', b'k'}:
            color = (operator, tuple(float(number) for number in operands))
        elif operator in {b'Tj', b'TJ', b"'", b'"'}:
            _require(tx_active and font_size is not None and font_size >= 4 and color == source_color
                     and (not source_size or abs(font_size - source_size) < .001),
                     f'PDF 선택 문구의 실제 표시 속성 불일치: {label}')
            sizes.append(font_size)
            text_count += 1
        elif operator == b're' and tx_active:
            path.append(tuple(float(number) for number in operands))
        elif operator in {b'f', b'f*'} and tx_active:
            _require(color == (b'rg', (.6, .75, .95)) and len(path) == 1,
                     f'PDF 선택 목록의 강조 표시 속성 불일치: {label}')
            highlights.extend(path)
            path = []
        elif operator == b'n' and tx_active:
            _require(not path, f'PDF 선택 강조 경로가 지워짐: {label}')
        elif operator in {b'S', b's', b'B', b'B*', b'b', b'b*'} and tx_active:
            _require(False, f'PDF 선택 영역에 예상하지 않은 경로 표시가 있음: {label}')
        elif operator in {b'Tr', b'gs', b'W', b'W*'}:
            _require(False, f'PDF 선택 문구를 숨기거나 잘라내는 표시 속성: {label}')
        elif not tx_active and operator not in {b'q', b'Q', b'BT', b'ET', b'TL'}:
            _require(False, f'PDF 선택 표시 영역 밖에 예상하지 않은 내용이 있음: {label}')
    _require(tx_count == 1 and not marked and not tx_active and text_count and len(set(sizes)) == 1 and not path,
             f'PDF 선택 표시 문구 또는 영역 누락: {label}')
    size = sizes[0]
    if state['combo']:
        expected_text = state['display'][0]
        expected_highlights = []
    else:
        top = _pdf_inherited_value(widget, '/TI', 0)
        _require(type(top).__name__ == 'NumberObject' or type(top) is int,
                 f'PDF 선택 목록 /TI 형식 불일치: {label}')
        capacity = int((height - 4) // (size * 1.3))
        _require(capacity > 0 and 0 <= top < len(state['pairs'])
                 and all(top <= index < top + capacity for index in state['indices']),
                 f'PDF 선택한 문구가 표시 구간에서 누락됨: {label}')
        expected_text = ' '.join(display for _, display in state['pairs'][top:top + capacity])
        expected_highlights = [(2, height - 2 - (index - top + 1) * size * 1.3, width - 4, size * 1.3)
                               for index in state['indices']]
    actual_text = _pdf_appearance_text(appearance, reader)
    _require(' '.join(actual_text.split()) == ' '.join(expected_text.split()),
             f'PDF 선택 코드와 표시 문구 불일치: {label}')
    rounded = lambda rectangles: [tuple(round(number, 3) for number in rect) for rect in rectangles]
    _require(rounded(highlights) == rounded(expected_highlights),
             f'PDF 선택 목록의 강조 행 불일치: {label}')
    _pdf_choice_visible_bounds(appearance, reader, width, height, label)


def _pdf_form_structure(reader, expected, *, verify_values=False, display_values=None):
    """폼 트리와 페이지 위젯의 연결·원위치를 값/appearance 변경과 별도로 대조함."""
    def fingerprint(value, ancestors=frozenset()):
        value = value.get_object() if hasattr(value, 'get_object') else value
        if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
            return ('number', Decimal(str(value)))  # PDF 저장의 1.0→1 정규화는 같은 좌표/속성임.
        if isinstance(value, (dict, list)):
            identity = id(value)
            _require(identity not in ancestors, 'PDF 입력 메타데이터의 순환 참조를 확인할 수 없음')
            ancestors = ancestors | {identity}
            if isinstance(value, list):
                return tuple(fingerprint(item, ancestors) for item in value)
            data = tuple((str(key), fingerprint(item, ancestors)) for key, item in sorted(value.items())
                         if key != '/Length' or not hasattr(value, 'get_data'))
            return (data, value.get_data()) if hasattr(value, 'get_data') else data
        return (type(value).__name__, str(value))

    def properties(node, name):
        mutable = {'/V', '/AP', '/AS', '/I'} if name in expected else set()
        if name in expected and str(_pdf_inherited_value(node, '/FT', '')) == '/Ch' and not int(_pdf_inherited_value(node, '/Ff', 0)) & (1 << 17):
            mutable.add('/TI')  # 선택 목록의 보이는 첫 행은 아래 AP/선택 범위 검사에서 검증함.
        if name in expected and any(ord(character) > 127 for character in displayed(name)):
            mutable.add('/DA')  # 실제 문자용 내장 Type0 글꼴은 아래 /DR 보존 검사와 함께 허용함.
            mutable.add('/DR')  # 선택 필드의 local 자원은 아래에서 별도 대조함.
        data = {key: value for key, value in node.items()
                if key not in mutable | {'/Parent', '/Kids', '/P'}}
        if node.get('/Subtype') == '/Popup':
            parent = node.get('/Parent')
            parent = parent.get_object() if parent is not None else None
            _require(isinstance(parent, dict) and id(parent) in annotation_locations
                     and parent.get('/Subtype') in markup_types
                     and parent.get('/Popup') is not None and parent['/Popup'] is node,
                     'PDF Popup의 원래 주석 Parent 연결 불일치')
            data['/Parent'] = ('markup_annotation', *annotation_locations[id(parent)][:2])
        if node.get('/Popup') is not None:
            popup = node['/Popup']
            _require(node.get('/Subtype') in markup_types and id(node) in annotation_locations
                     and isinstance(popup, dict) and popup.get('/Subtype') == '/Popup'
                     and popup.get('/Parent') is not None and popup['/Parent'] is node,
                     'PDF 주석의 Popup 연결 불일치')
            owner = annotation_locations[id(node)]
            popup_position = annotation_locations.get(id(popup))
            _require(popup_position is None or popup_position[0] == owner[0], 'PDF Popup의 원래 페이지 위치 불일치')
            _require(popup.get('/P') is None or popup.get('/P') == owner[2], 'PDF Popup의 페이지 연결 불일치')
            # ISO 32000-1 12.5.6.14: Popup/Parent is a reciprocal annotation
            # link. Inspect its complete properties without recursing into a page.
            data['/Popup'] = (popup_position[:2] if popup_position else None, '/P' in popup,
                              {key: value for key, value in popup.items() if key not in {'/Parent', '/P'}})
        return fingerprint(data)

    def local_resources(node):
        if node.get('/DR') is None:
            return None
        resources = node['/DR'].get_object()
        fonts = resources.get('/Font', {}).get_object() if resources.get('/Font') else {}
        return ({str(key): fingerprint(value) for key, value in resources.items() if key != '/Font'},
                {str(key): fingerprint(value) for key, value in fonts.items()})

    root = reader.trailer['/Root'].get('/AcroForm')
    _require(root is not None, 'PDF canonical 폼 트리가 없음')
    form = root.get_object()
    tree, locations, names, local_fonts = [], {}, set(), {}
    markup_types = {'/Text', '/FreeText', '/Line', '/Square', '/Circle', '/Polygon', '/PolyLine',
                    '/Highlight', '/Underline', '/Squiggly', '/StrikeOut', '/Stamp', '/Caret',
                    '/Ink', '/FileAttachment', '/Sound', '/Redact'}
    annotation_locations = {}
    for page_number, page in enumerate(reader.pages):
        for ordinal, reference in enumerate(page.get('/Annots', [])):
            node = reference.get_object()
            _require(isinstance(node, dict) and id(node) not in annotation_locations,
                     'PDF 주석의 중복·모호한 페이지 위치')
            annotation_locations[id(node)] = (page_number, ordinal, page.indirect_reference)

    def displayed(name):
        return (display_values or {}).get(name, expected[name])

    def visit(reference, path, prefix='', parent=None):
        node = reference.get_object()
        _require(isinstance(node, dict) and id(node) not in locations, 'PDF 폼 트리의 중복·모호한 연결')
        if parent is None:
            _require('/Parent' not in node, 'PDF 최상위 필드의 Parent 연결이 모호함')
        else:
            link = node.get('/Parent')
            _require(link is not None and link.get_object() is parent, 'PDF canonical 필드/위젯 Parent 연결 불일치')
        name = '.'.join(part for part in (prefix, str(node.get('/T', ''))) if part)
        if '/T' in node:
            _require(name not in names, 'PDF 같은 이름의 canonical 필드가 중복됨')
            names.add(name)
        locations[id(node)] = (path, name, node)
        tree.append((path, name, properties(node, name), '/Parent' in node, '/P' in node))
        if name in expected and any(ord(character) > 127 for character in displayed(name)):
            local_fonts[path] = local_resources(node)
            if verify_values:
                da = str(node.get('/DA', ''))
                if not da and node.get('/Parent'):
                    da = str(node['/Parent'].get('/DA', ''))
                match = re.search(r'(/CompatCJK[0-9a-f]{12})\s+[\d.]+\s+Tf', da)
                _require(match is not None, f'PDF 한글 필드의 표시 글꼴 연결 누락: {name}')
                global_dr = form.get('/DR', {}).get_object()
                font = global_dr.get('/Font', {}).get_object().get(match[1])
                font = font.get_object() if font is not None else None
                _require(font is not None and font.get('/Subtype') == '/Type0' and font.get('/Encoding') == '/Identity-H'
                         and font.get('/ToUnicode') is not None and font.get('/DescendantFonts'),
                         f'PDF 한글 필드의 실제 Type0 글꼴 누락: {name}')
                descendant = font['/DescendantFonts'][0].get_object()
                _require(descendant.get('/CIDToGIDMap') == '/Identity',
                         f'PDF 한글 필드의 편집 글리프 연결 불일치: {name}')
                descriptor = descendant.get('/FontDescriptor')
                embedded = descriptor.get_object().get('/FontFile2') if descriptor else None
                _require(embedded is not None and bool(embedded.get_object().get_data()),
                         f'PDF 한글 필드의 내장 글꼴 누락: {name}')
                if local_fonts[path] is not None:
                    _require(local_fonts[path][1].get(match[1]) == fingerprint(font),
                             f'PDF 한글 필드의 local/global 글꼴 연결 불일치: {name}')
        for index, child in enumerate(node.get('/Kids', [])):
            visit(child, path + (index,), name, node)

    for index, field in enumerate(form.get('/Fields', [])):
        visit(field, (index,))
    _require(names == set(reader.get_fields() or {}), 'PDF canonical 필드 이름/트리가 일치하지 않음')
    annotations, seen_widgets, page_repairs = [], set(), []
    for page_number, page in enumerate(reader.pages):
        for ordinal, reference in enumerate(page.get('/Annots', [])):
            node = reference.get_object()
            page_link = node.get('/P')
            missing_source_link = (not verify_values and node.get('/Subtype') == '/Widget'
                                   and isinstance(page_link, IndirectObject) and page_link.get_object() is None)
            _require(page_link is None or page_link == page.indirect_reference or missing_source_link,
                     f'PDF 위젯/주석의 페이지 연결 불일치: {page_number + 1}쪽')
            if node.get('/Subtype') == '/Widget':
                _require(id(node) in locations and id(node) not in seen_widgets,
                         'PDF canonical 트리 밖 또는 중복 Widget 연결')
                seen_widgets.add(id(node))
                path, name, _ = locations[id(node)]
                if missing_source_link:
                    page_repairs.append({'page': page_number + 1, 'annotation_index': ordinal, 'field': name})
                annotations.append((page_number, ordinal, path, name, properties(node, name), '/P' in node))
                if verify_values and name in expected and '/V' in node and '/T' in node:
                    if str(_pdf_inherited_value(node, '/FT', '')) != '/Ch':
                        _require(str(node['/V']) == expected[name], f'PDF Widget /V 불일치: {name}')
            else:
                annotations.append((page_number, ordinal, None, '', properties(node, ''), '/P' in node))
    _require(seen_widgets == {identity for identity, (_, _, node) in locations.items()
                             if node.get('/Subtype') == '/Widget'}, 'PDF canonical Widget의 페이지 배치가 누락됨')
    form_properties = fingerprint({key: value for key, value in form.items()
                                   if key not in {'/Fields', '/DR', '/NeedAppearances'}})
    resources = form.get('/DR', {}).get_object() if form.get('/DR') else {}
    resource_properties = {str(key): fingerprint(value) for key, value in resources.items() if key != '/Font'}
    fonts = resources.get('/Font', {}).get_object() if resources.get('/Font') else {}
    font_properties = {str(key): fingerprint(value) for key, value in fonts.items()}
    return tree, annotations, form_properties, resource_properties, font_properties, local_fonts, page_repairs


def _pdf_compare_form_structure(before, after, expected, display_values=None):
    old = _pdf_form_structure(before, expected, display_values=display_values)
    new = _pdf_form_structure(after, expected, verify_values=True, display_values=display_values)
    _require(old[:4] == new[:4], 'PDF 원본 필드/위젯의 위치·종류·옵션·속성 또는 연결 변경')
    need_appearances = after.trailer['/Root']['/AcroForm'].get('/NeedAppearances')
    _require(need_appearances is None or str(need_appearances) == 'False',
             'PDF 저장된 appearance 대신 뷰어의 표시 재생성에 의존함')
    _require(all(new[4].get(name) == value for name, value in old[4].items()), 'PDF 원본 폼 글꼴 자원 변경')
    added = set(new[4]) - set(old[4])
    fallback = (('/BaseFont', ('NameObject', '/Helvetica')), ('/Encoding', ('NameObject', '/WinAnsiEncoding')),
                ('/Name', ('NameObject', '/Helvetica')), ('/Subtype', ('NameObject', '/Type1')),
                ('/Type', ('NameObject', '/Font')))
    rendered = {name: (display_values or {}).get(name, value) for name, value in expected.items()}
    _require(not added or (any(any(ord(character) > 127 for character in value) for value in rendered.values())
                           and all(re.fullmatch(r'/CompatCJK[0-9a-f]{12}', name)
                                   or name == '/Helvetica' and new[4][name] == fallback for name in added)),
             'PDF 허용하지 않은 폼 글꼴 자원 추가')
    _require(set(old[5]) == set(new[5]), 'PDF 선택 local 글꼴 필드 변경')
    for path, original in old[5].items():
        filled = new[5][path]
        _require((original is None) == (filled is None), 'PDF 원본 local 글꼴 자원의 추가 또는 삭제')
        if original is None:
            continue
        _require(original[0] == filled[0] and all(filled[1].get(name) == value for name, value in original[1].items()),
                 'PDF 원본 local 글꼴 또는 기타 자원 변경')
        _require(all(re.fullmatch(r'/CompatCJK[0-9a-f]{12}', name) and filled[1][name] == new[4].get(name)
                     for name in filled[1].keys() - original[1].keys()),
                 'PDF 선택 필드에 허용하지 않은 local 글꼴 자원 추가')
    return old[6]


def _pdf_render(path, page_number):
    import pypdfium2 as pdfium
    document = pdfium.PdfDocument(str(path))
    try:
        document.init_forms()
        page = document[page_number]
        try:
            bitmap = page.render(scale=1, draw_annots=True)
            try:
                return bitmap.to_pil().convert("RGB").copy()
            finally:
                bitmap.close()
        finally:
            page.close()
    finally:
        document.close()


def _pdf_checks(template, output, profile, selected, checks, annex_entries=None):
    from PIL import ImageChops, ImageDraw
    before, after = PdfReader(template), PdfReader(output)
    _require(not after.is_encrypted, "출력 PDF가 암호화됨")
    if annex_entries:
        _require(len(after.pages) > len(before.pages), 'PDF 별첨 페이지 누락')
        from templates.pdf_annex import verify_annex
        checks.append(verify_annex(after, len(before.pages), annex_entries))
        import pypdfium2 as pdfium
        document = pdfium.PdfDocument(str(output))
        try:
            for index in range(len(before.pages), len(after.pages)):
                page = document[index]
                text = page.get_textpage()
                try:
                    width, height = page.get_size()
                    for character in range(text.count_chars()):
                        if not text.get_text_range(character, 1).strip():
                            continue
                        left, bottom, right, top = text.get_charbox(character)
                        _require(24 <= left <= right <= width - 24 and 20 <= bottom <= top <= height - 20,
                                 f'PDF 별첨 글자가 페이지 작성 영역 밖에 있음: {index + 1}쪽')
                finally:
                    text.close()
                    page.close()
        finally:
            document.close()
        checks.append({'name': 'pdf_annex_visible_bounds', 'status': 'passed', 'pages': len(after.pages) - len(before.pages)})
    else:
        _require(len(before.pages) == len(after.pages), "PDF 페이지 수 변경")
    for index, (original, filled) in enumerate(zip(before.pages, after.pages)):
        old_size = tuple(round(float(value), 3) for value in (original.mediabox.width, original.mediabox.height))
        if profile.get("render_mode") == "overlay" and original.rotation % 180:
            old_size = old_size[::-1]
        new_size = tuple(round(float(value), 3) for value in (filled.mediabox.width, filled.mediabox.height))
        _require(old_size == new_size, f"PDF 페이지 크기 변경: {index + 1}쪽")
    selected_pages = set()
    regions = {}
    if profile.get("render_mode") == "overlay":
        for field, value in selected:
            page = int(field["page"]) - 1
            _require(0 <= page < len(after.pages), "PDF 입력 페이지 누락")
            _contains(after.pages[page].extract_text() or "", value, field["label"])
            regions.setdefault(page, []).append(((field["x"], field["y"], field["x"] + field["width"], field["y"] + field["height"]), True, field["label"]))
            _require(_normalized(before.pages[page].extract_text() or "") in _normalized(after.pages[page].extract_text() or ""), f"PDF 원본 텍스트 손실: {page + 1}쪽")
            selected_pages.add(page)
    else:
        old_fields, new_fields = before.get_fields() or {}, after.get_fields() or {}
        _require(set(old_fields) == set(new_fields), "PDF 폼 필드 손실 또는 추가")
        expected = {field["id"][4:]: value for field, value in selected}
        for name in old_fields.keys() - expected.keys():
            _require(str(old_fields[name].get("/V", "")) == str(new_fields[name].get("/V", "")), f"PDF 수정 대상 밖의 필드 변경: {name}")
        choice_states, display_values = {}, {}
        for field, value in selected:
            name = field["id"][4:]
            _require(name in new_fields, f"PDF canonical 필드 누락: {field['label']}")
            if str(old_fields[name].get('/FT')) == '/Ch':
                old_node = old_fields[name].indirect_reference.get_object()
                new_node = new_fields[name].indirect_reference.get_object()
                state = _pdf_choice_state(old_node, new_node, value, field['label'])
                state['da'] = str(_pdf_inherited_value(old_node, '/DA', before.trailer['/Root']['/AcroForm'].get('/DA', '')))
                choice_states[name] = state
                display_values[name] = ''.join(state['display'] if state['combo'] else [display for _, display in state['pairs']])
            else:
                _require(str(new_fields[name].get("/V", "")) == value, f"PDF canonical /V 불일치: {field['label']}")
        page_repairs = _pdf_compare_form_structure(before, after, expected, display_values)
        if page_repairs:
            checks.append({'name': 'pdf_dangling_widget_page_links_repaired', 'status': 'passed',
                           'widgets': page_repairs, 'count': len(page_repairs),
                           'scope': '원본의 존재하지 않는 페이지 참조만 고유한 원래 주석·canonical 위치로 복구함. 출력의 유효 페이지 연결을 독립 대조함'})
        for index, (original, filled) in enumerate(zip(before.pages, after.pages)):
            _require((original.get_contents().get_data() if original.get_contents() else b"") ==
                     (filled.get_contents().get_data() if filled.get_contents() else b""), f"PDF 원본 배경 내용 변경: {index + 1}쪽")
            for ordinal, reference in enumerate(filled.get("/Annots", [])):
                widget = reference.get_object()
                name = _pdf_field_name(widget)
                if name not in expected:
                    continue
                selected_pages.add(index)
                rect = [float(number) for number in widget.get("/Rect", [])]
                _require(len(rect) == 4, f"PDF 표시 영역 누락: {name}")
                height = float(filled.mediabox.height)
                changed = str(old_fields.get(name, {}).get("/V", "")) != expected[name]
                if name in choice_states:
                    old_node = old_fields[name].indirect_reference.get_object()
                    new_node = new_fields[name].indirect_reference.get_object()
                    changed = any(str(_pdf_inherited_value(old_node, key)) != str(_pdf_inherited_value(new_node, key))
                                  for key in ('/V', '/I'))
                    _pdf_choice_state(old_node, widget, expected[name], name)
                    if choice_states[name]['combo']:
                        old_value = str(_pdf_inherited_value(old_node, '/V', ''))
                        old_display = next((display for code, display in choice_states[name]['pairs'] if code == old_value), old_value)
                        changed = old_display != choice_states[name]['display'][0]
                if str(new_fields[name].get('/FT')) == '/Btn':
                    original_widget = original['/Annots'][ordinal].get_object()
                    changed = str(original_widget.get('/AS', '')) != str(widget.get('/AS', ''))
                regions.setdefault(index, []).append(((rect[0], height - rect[3], rect[2], height - rect[1]),
                    changed, name))
                appearances = widget.get("/AP", {}).get("/N")
                _require(appearances is not None, f"PDF 표시 appearance 누락: {name}")
                if str(new_fields[name].get("/FT")) == "/Btn":
                    value = expected[name]
                    # 라디오의 비선택 widget은 /Off를 유지해야 함.
                    allowed = set(appearances.get_object())
                    state = value if value in allowed else "/Off"
                    _require(str(widget.get("/AS")) == state, f"PDF 체크박스/라디오 표시 상태 불일치: {name}")
                else:
                    _require(bool(appearances.get_object().get_data()), f"PDF 빈 표시 appearance: {name}")
                    if name in choice_states:
                        _pdf_choice_background(original['/Annots'][ordinal].get_object(), widget, before, after, name)
                        _pdf_choice_appearance(widget, after, choice_states[name], name)
        checks.extend([{"name": "pdf_source_fields_widgets_unchanged", "status": "passed"},
                       {"name": "pdf_canonical_values_widget_states", "status": "passed"}])
        if choice_states:
            checks.append({'name': 'pdf_choice_export_display_indices', 'status': 'passed', 'fields': len(choice_states),
                           'scope': '원본 코드·표시 문구·선택 인덱스·가시 목록 강조를 독립 대조함. 모든 PDF 뷰어의 편집 동작은 별도임'})
    for index in selected_pages:
        original, filled = _pdf_render(template, index), _pdf_render(output, index)
        _require(original.size == filled.size, f"PDF 렌더 크기 변경: {index + 1}쪽")
        difference = ImageChops.difference(original, filled).getbbox()
        # 동일한 기존값을 다시 쓴 폼은 픽셀 변화가 필요하지 않음.
        if profile.get("render_mode") == "overlay" or any(changed for _, changed, _ in regions.get(index, [])):
            _require(difference is not None, f"PDF 입력 내용의 시각적 표시 없음: {index + 1}쪽")
        for box, changed, label in regions.get(index, []):
            if changed:
                _require(ImageChops.difference(original.crop(box), filled.crop(box)).getbbox() is not None,
                         f"PDF 입력칸의 시각적 표시 없음: {label}")
    for index in range(len(before.pages)):
        original, filled = _pdf_render(template, index), _pdf_render(output, index)
        difference = ImageChops.difference(original, filled)
        mask = ImageDraw.Draw(difference)
        for box, _, _ in regions.get(index, []):
            left, top, right, bottom = box
            mask.rectangle((max(0, left - 2), max(0, top - 2), right + 2, bottom + 2), fill=(0, 0, 0))
        _require(difference.getbbox() is None, f'PDF 원본의 작성 영역 밖 배치·그림·문구 변경: {index + 1}쪽')
    checks.append({'name': 'pdf_overlay_outside_regions_unchanged' if profile.get('render_mode') == 'overlay'
                   else 'pdf_form_outside_regions_unchanged', 'status': 'passed', 'pages': len(before.pages)})
    checks.extend([{"name": "pdf_pages_dimensions", "status": "passed", "pages": len(after.pages)},
                   {"name": "pdfium_visual_presence", "status": "passed", "pages": len(selected_pages),
                    "scope": "입력 페이지 렌더 변화 검사임. 원본 프로그램의 전체 배치·글자 식별 검증은 별도임"}])


def verify_output(template, output, values, *, profile=None, mapping=None) -> dict:
    """성공은 JSON 보고서, 실패는 다운로드를 차단할 ValueError를 반환함."""
    source, result = Path(template), Path(output)
    _require(source.resolve() != result.resolve(), "원본을 덮어쓸 수 없음")
    _require(source.is_file() and result.is_file(), "원본 또는 출력 파일이 없음")
    _require(0 < source.stat().st_size <= 64 * 1024 * 1024 and 0 < result.stat().st_size <= 64 * 1024 * 1024, "파일 크기 제한(64MiB) 위반")
    kind = source.suffix.lower().lstrip(".")
    _require(kind in {"docx", "hwpx", "xlsx", "pdf", "pptx"} and result.suffix.lower() == source.suffix.lower(), "검증할 출력 형식이 일치하지 않음")
    _require(isinstance(values, dict) and all(isinstance(k, str) and isinstance(v, str) for k, v in values.items()), "입력값은 문자열 사전이어야 함")
    original_sha = sha256(source.read_bytes()).hexdigest()
    profile = profile if profile is not None else analyze_template(source)
    _require(profile.get("source_sha256") == original_sha, "프로파일과 원본 SHA-256 불일치")
    _require(profile.get("format") == kind, "프로파일 형식 불일치")
    rules_profile = mapped_rule_profile(profile, mapping)
    if profile.get('repeat_expansion'):
        from templates.repeat_rows import validate_repeat_fields
        validate_repeat_fields(source, rules_profile)
    literal_values = ({field['value_key']: values.get(field['value_key'], '')
                       for field in rules_profile.get('fields', [])
                       if field.get('control_type') in {'checkbox', 'radio', 'choice', 'combobox'}}
                      if profile.get('citation_mode') == 'sidecar' else None)
    value_issues = inspect_form_values(values, rules_profile, literal_values=literal_values)
    _require(not value_issues, '양식 입력값 검증 실패: ' + ' / '.join(f"{issue['field']}: {issue['message']}" for issue in value_issues))
    try:
        from templates.pdf_annex import plan_annex
        printed, annex_entries = plan_annex(profile, values, mapping)
        selected = _selected_fields(profile, printed, mapping)
        checks = [{"name": "profile_source_sha256", "status": "passed"}]
        checks.append({'name': 'form_value_rules', 'status': 'passed',
                       'fields': len({field['value_key'] for field in rules_profile.get('fields', [])
                                      if field.get('validation') or field.get('control_type') in {'choice', 'radio'} and field.get('options')}),
                       'relations': len(rules_profile.get('constraints', {}).get('relations', [])),
                       'groups': len(rules_profile.get('constraints', {}).get('groups', [])),
                       'scope': '예상 값의 입력 규칙과 재열기한 동일 입력 위치의 정확한 값을 함께 검사함'})
        if kind == "pdf":
            _pdf_checks(source, result, profile, selected, checks, annex_entries)
        else:
            _zip_checks(source, result, kind, selected, checks)
        _require(sha256(source.read_bytes()).hexdigest() == original_sha, "검사 중 원본이 변경됨")
        output_sha = sha256(result.read_bytes()).hexdigest()
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(f"출력 검증 실패: 저장 파일 재열기 오류 ({exc})") from exc
    checks.append({"name": "original_unchanged_during_check", "status": "passed"})
    layout_pending = any(check["name"] == "xlsx_layout_fit" and check["status"] == "pending" for check in checks)
    checks.append({"name": "obvious_fixed_height_overflow", "status": "pending" if layout_pending else ("passed" if kind in {"docx", "xlsx"} else "not_applicable"),
                   "scope": "명시된 고정 높이만 보수적으로 추정함. 글꼴·폭에 따른 전체 배치는 원본 프로그램 확인이 필요함"})
    return {"status": "passed", "format": kind, "sha": {"template": original_sha, "output": output_sha},
            "checks": checks, "resource_kind": profile.get("resource_kind", "unknown"),
            "native_visual_qa": "pending", "legal_compliance_certified": False}
