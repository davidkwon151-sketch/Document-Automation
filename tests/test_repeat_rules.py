"""반복 준비본에도 확인된 위치·입력권한·관계 규칙을 유지해야 함."""

from copy import deepcopy
from hashlib import sha256
import json
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
import pytest

from agent.output_check import verify_output
from templates import analyze_template, fill_compatible_template, load_form_profile
from templates.fill import TemplateError
from templates.repeat_docx import inventory
from templates.repeat_fields import repeat_profile
from templates.repeat_rows import prepare_repeat_template, expansion_binding
from templates.repeat_rules import inherit_repeat_rules
from templates.value_rules import inspect_form_values


ROOT=Path(__file__).resolve().parents[1]
NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}


def read_parts(path):
    with ZipFile(path) as archive:
        return {name:archive.read(name) for name in archive.namelist()}


def setup(tmp_path, *, placeholders=False):
    doc=Document()
    doc.add_paragraph('고정 안내 {{발행일}} / {{종료일}}')
    table=doc.add_table(rows=3,cols=4)
    for cell,label in zip(table.rows[0].cells,['첫 금액(원)','둘째 금액(원)','계 금액(원)','결재 상태']):
        cell.text=label
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    if placeholders:
        for cell,label in zip(table.rows[1].cells[:3],['첫 금액','둘째 금액','계 금액']):
            cell.text='{{'+label+'}}'
    sdt,props,content=OxmlElement('w:sdt'),OxmlElement('w:sdtPr'),OxmlElement('w:sdtContent')
    alias=OxmlElement('w:alias'); alias.set(qn('w:val'),'결재 상태')
    identifier=OxmlElement('w:id'); identifier.set(qn('w:val'),'12')
    props.extend([alias,identifier,OxmlElement('w:showingPlcHdr')])
    choices=OxmlElement('w:dropDownList')
    for value,label in [('A','검토 전'),('B','검토 완료')]:
        item=OxmlElement('w:listItem'); item.set(qn('w:value'),value); item.set(qn('w:displayText'),label)
        choices.append(item)
    props.append(choices)
    content.append(table.rows[1].cells[3].add_paragraph('항목 선택')._p)
    sdt.extend([props,content]); table.rows[1].cells[3]._tc.append(sdt)
    source=tmp_path/'source.docx'; doc.save(source)
    original=analyze_template(source)
    approved=[]
    for field in original['fields']:
        field=deepcopy(field)
        if field['kind']=='placeholder' and field['value_key'] in {'발행일','종료일'}:
            field.update(validation={'type':'date'},max_chars=10,required=True,input_required=True,input_mode='user_provided',narrative_style_required=False)
            approved.append(field)
        elif field['kind']=='docx_choice' or '/w:tr[2]/' in field['id']:
            if field['kind']!='docx_choice':
                column=int(field['id'].split('/w:tc[')[1].split(']')[0])
                field.update(value_key=['첫 금액','둘째 금액','계 금액'][column-1])
            approved.append(field)
        elif placeholders and field['kind']=='placeholder':
            approved.append(field)
        elif '/w:tr[3]/w:tc[1]' in field['id']:
            field.update(value_key='바깥 확인란',label='바깥 확인란',max_chars=20,input_required=True,input_mode='user_provided',narrative_style_required=False)
            approved.append(field)
    for field in approved:
        if field['value_key'] in {'첫 금액','둘째 금액','계 금액'}:
            field.update(validation={'type':'integer','unit':'원','unit_location':'label','min':0},required=True,
                         input_required=False,input_mode='source_grounded',max_chars=16,narrative_style_required=False)
    if not any(field['value_key']=='바깥 확인란' for field in approved):
        approved.append({'id':'docx:word/document.xml:/w:document/w:body/w:tbl/w:tr[3]/w:tc[1]',
                         'kind':'docx_cell','label':'바깥 확인란','value_key':'바깥 확인란','required':False,
                         'max_chars':20,'input_required':True,'input_mode':'user_provided','narrative_style_required':False})
    original['fields']=approved
    original.update(domain='business_support',business_workflow='government_grant',citation_mode='sidecar',document_kind='신청서')
    original['constraints']={'relations':[{'kind':'sum','total':'계 금액','parts':['첫 금액','둘째 금액']},
                                        {'kind':'date_order','start':'발행일','end':'종료일'}],
                             'groups':[{'kind':'all_or_none','fields':['첫 금액','둘째 금액','계 금액'],'required_fields':['첫 금액','둘째 금액']}]}
    plan={'table_id':inventory(read_parts(source))[0]['id'],'row':2,'count':3}
    prepared=tmp_path/'prepared.docx'
    binding=expansion_binding(prepare_repeat_template(source,prepared,plan))
    raw=repeat_profile(prepared,plan,binding)
    return source,prepared,original,raw


def inherited(tmp_path, **kwargs):
    source,prepared,original,raw=setup(tmp_path,**kwargs)
    return source,prepared,original,inherit_repeat_rules(source,raw,original,prepared_path=prepared)


@pytest.mark.parametrize('placeholders',[False,True])
def test_each_row_inherits_exact_rules_and_native_choices_and_outside_scope(tmp_path,placeholders):
    source,prepared,original,raw=setup(tmp_path,placeholders=placeholders)
    before=(source.read_bytes(),prepared.read_bytes(),deepcopy(original),deepcopy(raw))
    profile=inherit_repeat_rules(source,raw,original,prepared_path=prepared)
    assert (source.read_bytes(),prepared.read_bytes(),original,raw)==before
    assert profile['domain']=='business_support' and profile['citation_mode']=='sidecar'
    assert len(profile['constraints']['relations'])==4
    sums=[relation for relation in profile['constraints']['relations'] if relation['kind']=='sum']
    for row in range(1,4):
        fields=[field for field in profile['fields'] if field.get('repeat_info',{}).get('row')==row]
        assert len(fields)==4
        for field in fields:
            if field['kind']=='docx_choice':
                assert field['options']==['A','B'] and field['input_required'] is True
                assert field['choice_items']==[{'value':'A','label':'검토 전'},{'value':'B','label':'검토 완료'}]
            else:
                assert field['validation']=={'type':'integer','unit':'원','unit_location':'label','min':0}
                assert field['required'] is True and field['max_chars']==16 and field['narrative_style_required'] is False
        assert {sums[row-1]['total'],*sums[row-1]['parts']}=={field['value_key'] for field in fields if field['kind']!='docx_choice'}
    suffix=next(field for field in profile['fields'] if field['value_key']=='바깥 확인란')
    assert '/w:tr[5]/w:tc[1]' in suffix['id'] and suffix['input_required'] is True and suffix['max_chars']==20
    assert not any('/w:tr[5]/w:tc[2]' in field['id'] for field in profile['fields'])


def test_inherited_scalar_sum_group_and_choice_errors_are_detected(tmp_path):
    _,_,_,profile=inherited(tmp_path)
    values={'발행일':'2026-01-01','종료일':'2026-02-01'}
    for field in profile['fields']:
        if field.get('repeat_info'):
            column=field['repeat_info']['column']
            values[field['value_key']]={1:'2',2:'3',3:'5',4:'A'}[column]
    assert not inspect_form_values(values,profile)
    first=[field for field in profile['fields'] if field.get('repeat_info',{}).get('row')==1]
    total=next(field['value_key'] for field in first if field['repeat_info']['column']==3)
    changed=dict(values); changed[total]='6'
    assert any('sum' in issue['code'] for issue in inspect_form_values(changed,profile))
    changed=dict(values); changed[total]='-1'
    assert inspect_form_values(changed,profile)
    changed=dict(values); changed[next(field['value_key'] for field in first if field['kind']=='docx_choice')]='검토 전'
    assert inspect_form_values(changed,profile)
    changed=dict(values); changed[next(field['value_key'] for field in first if field['repeat_info']['column']==2)]=''
    assert any(issue['code']=='form_group_required' for issue in inspect_form_values(changed,profile))


def test_original_absent_backprojects_positions_and_retains_rules(tmp_path):
    source,prepared,original,raw=setup(tmp_path)
    expected=inherit_repeat_rules(source,raw,original,prepared_path=prepared)
    actual=inherit_repeat_rules(None,raw,original,prepared_path=prepared)
    assert actual['fields']==expected['fields'] and actual['constraints']==expected['constraints']
    assert actual['repeat_rule_verification']=={'original_file_checked':False,'prepared_file_checked':True}


def test_original_only_builds_in_memory_without_creating_a_file(tmp_path):
    source,prepared,original,raw=setup(tmp_path)
    expected=inherit_repeat_rules(source,raw,original,prepared_path=prepared)
    before=set(tmp_path.iterdir())
    actual=inherit_repeat_rules(source,raw,original)
    assert actual['fields']==expected['fields'] and set(tmp_path.iterdir())==before


def test_source_policy_cache_has_no_values_or_local_paths_but_keeps_xml_and_choice_codes(tmp_path):
    source,prepared,original,raw=setup(tmp_path)
    original.update(source_path=str(source),documents=[{'text':'PRIVATE DATA'}],field_values={'결재 상태':'PRIVATE DATA'})
    native=next(field for field in original['fields'] if field['kind']=='docx_choice')
    native.update(xml_path=native['id'].split(':',2)[2],value='PRIVATE DATA',default_value='PRIVATE DATA')
    profile=inherit_repeat_rules(source,raw,original,prepared_path=prepared)
    cached=profile['repeat_source_profile']
    text=json.dumps(cached,ensure_ascii=False)
    assert 'PRIVATE DATA' not in text and str(source) not in text
    saved=next(field for field in cached['fields'] if field['kind']=='docx_choice')
    assert saved['xml_path']==native['xml_path'] and saved['choice_items']==native['choice_items']
    assert 'source_path' not in cached and 'field_values' not in cached


@pytest.mark.parametrize('change',['original_sha','prepared_sha','file','duplicate_id','repeated_source','kind','options','direct_mode','readonly'])
def test_binding_scope_and_native_policy_forgery_are_rejected(tmp_path,change):
    source,prepared,original,raw=setup(tmp_path)
    native=next(field for field in original['fields'] if field['kind']=='docx_choice')
    if change=='original_sha': original['source_sha256']='a'*64
    elif change=='prepared_sha': raw['source_sha256']='b'*64
    elif change=='file': prepared.write_bytes(prepared.read_bytes()+b'changed')
    elif change=='duplicate_id': original['fields'].append(deepcopy(original['fields'][0]))
    elif change=='repeated_source': original['repeat_expansion']=deepcopy(raw['repeat_expansion'])
    elif change=='kind': native['kind']='docx_sdt'; native.pop('options'); native.pop('control_type')
    elif change=='options': native['options']=['C']
    elif change=='direct_mode': native['input_mode']='source_grounded'
    elif change=='readonly': original['fields'][0]['readonly']=True
    with pytest.raises((TemplateError,ValueError)):
        inherit_repeat_rules(source,raw,original,prepared_path=prepared)


@pytest.mark.parametrize('kind',['sum','group'])
def test_crossing_repeated_and_outside_dependencies_is_an_explicit_error(tmp_path,kind):
    source,prepared,original,raw=setup(tmp_path)
    if kind=='group':
        original['constraints']['groups']=[{'kind':'all_or_none','fields':['첫 금액','발행일'],'required_fields':['첫 금액']}]
    else:
        outside=next(field for field in original['fields'] if field['value_key']=='바깥 확인란')
        outside['validation']={'type':'integer','unit':'원','unit_location':'label'}
        original['constraints']['relations']=[{'kind':'sum','total':'계 금액','parts':['첫 금액','바깥 확인란']}]
    with pytest.raises(TemplateError,match='바깥|원본 행'):
        inherit_repeat_rules(source,raw,original,prepared_path=prepared)


def test_newly_selected_unregistered_row_fields_remain_confirmation_candidates(tmp_path):
    source,prepared,original,raw=setup(tmp_path)
    original['fields']=[field for field in original['fields'] if not '/w:tr[2]/' in field['id']]
    original['constraints']={'relations':[],'groups':[]}
    profile=inherit_repeat_rules(source,raw,original,prepared_path=prepared)
    assert len(profile['repeat_unconfirmed_fields'])==12
    assert len([field for field in profile['fields'] if field.get('repeat_info')])==12
    assert any('매핑 확인' in warning for warning in profile['warnings'])


@pytest.mark.parametrize('name,table,row',[
    ('business_kotra_digital_2026','/w:document/w:body/w:tbl',13),
    ('business_startup_2026','/w:document/w:body/w:tbl[5]',15),
])
def test_real_public_docx_registered_exact_locations_survive_raw_heuristic_gaps(tmp_path,name,table,row):
    source=ROOT/'data/public_templates/business'/f'{name}.docx'
    if not source.exists(): pytest.skip('공식 공개 원본이 이 체크아웃에 없음')
    before=source.read_bytes()
    original=load_form_profile(source)
    assert original is not None
    plan={'table_id':'docx:word/document.xml:'+table,'row':row,'count':3}
    prepared=tmp_path/'prepared.docx'
    binding=expansion_binding(prepare_repeat_template(source,prepared,plan))
    profile=inherit_repeat_rules(source,repeat_profile(prepared,plan,binding),original,prepared_path=prepared)
    repeated=[field for field in profile['fields'] if field.get('repeat_info')]
    assert repeated and all(field['value_key'] for field in repeated)
    for field in original['fields']:
        if '/w:tr['+str(row)+']/' not in field['id']:
            saved=next(item for item in profile['fields'] if item['value_key']==field['value_key'])
            assert saved['kind']==field['kind'] and saved['id']==field['id']
            for key in ('required','input_required','input_mode','max_chars','narrative_style_required','validation'):
                if key in field: assert saved[key]==field[key]
    values={field['value_key']:f"QA{field['repeat_info']['row']}" for field in repeated}
    for field in profile['fields']:
        values.setdefault(field['value_key'],'QA' if field.get('required') else '')
    output=fill_compatible_template(prepared,values,tmp_path/'filled.docx',profile=profile)
    assert verify_output(prepared,output,values,profile=profile)['status']=='passed'
    assert source.read_bytes()==before
    assert sha256(source.read_bytes()).hexdigest()==binding['original_sha256']


@pytest.mark.parametrize('original_available',[True,False])
def test_hwpx_exact_cell_positions_inherit_per_row_sum_rules(tmp_path,original_available):
    from templates.repeat_hwpx import inventory as hwpx_inventory
    source=tmp_path/'source.hwpx'
    source.write_bytes((ROOT/'samples/sample_company_form.hwpx').read_bytes())
    original=analyze_template(source)
    original['fields']=[field for field in original['fields'] if field['kind']=='hwpx_cell']
    assert len(original['fields'])==3
    for field,key in zip(original['fields'],['수량1','수량2','합계수량']):
        field.update(value_key=key,required=True,validation={'type':'integer','min':0},max_chars=8,
                     input_mode='source_grounded',input_required=False,narrative_style_required=False)
    original['constraints']={'relations':[{'kind':'sum','total':'합계수량','parts':['수량1','수량2']}]}
    plan={'table_id':hwpx_inventory(read_parts(source))[0]['id'],'row':2,'count':3}
    prepared=tmp_path/'prepared.hwpx'
    binding=expansion_binding(prepare_repeat_template(source,prepared,plan))
    raw=repeat_profile(prepared,plan,binding)
    before=(source.read_bytes(),prepared.read_bytes())
    profile=inherit_repeat_rules(source if original_available else None,raw,original,prepared_path=prepared)
    assert len(profile['fields'])==9 and len(profile['constraints']['relations'])==3
    values={field['value_key']:{1:'1',2:'2',3:'3'}[field['repeat_info']['column']] for field in profile['fields']}
    assert not inspect_form_values(values,profile)
    values[profile['constraints']['relations'][1]['total']]='4'
    assert any('sum' in issue['code'] for issue in inspect_form_values(values,profile))
    assert (source.read_bytes(),prepared.read_bytes())==before


def test_native_source_cannot_be_disguised_as_its_surrounding_blank_cell(tmp_path):
    source,prepared,original,raw=setup(tmp_path)
    native=next(field for field in original['fields'] if field['kind']=='docx_choice')
    native.update(id=native['id'].split('/w:sdt')[0],kind='docx_cell')
    for key in ('control_type','options','choice_items','allow_custom'):
        native.pop(key,None)
    with pytest.raises(TemplateError,match='위장'):
        inherit_repeat_rules(source,raw,original,prepared_path=prepared)


def test_nonexistent_approved_xml_location_is_never_replaced_by_a_nearby_raw_field(tmp_path):
    source,prepared,original,raw=setup(tmp_path)
    native=next(field for field in original['fields'] if field['kind']=='docx_choice')
    native['id']=native['id']+'/w:p[99]'
    with pytest.raises(TemplateError,match='실제 XML에 없음'):
        inherit_repeat_rules(source,raw,original,prepared_path=prepared)


def test_approved_blank_paragraph_fallback_does_not_copy_saved_personal_values(tmp_path):
    source,prepared,original,raw=setup(tmp_path)
    field=next(field for field in original['fields'] if field['value_key']=='바깥 확인란')
    field.update(id=field['id']+'/w:p',kind='docx_paragraph',value='PRIVATE USER',default='PRIVATE USER',source_path=str(source))
    profile=inherit_repeat_rules(source,raw,original,prepared_path=prepared)
    assert 'PRIVATE USER' not in json.dumps(profile,ensure_ascii=False)
    saved=next(item for item in profile['fields'] if item['value_key']=='바깥 확인란')
    assert saved['kind']=='docx_paragraph' and saved['id'].endswith('/w:tr[5]/w:tc[1]/w:p')
    assert 'value' not in saved and 'source_path' not in saved
