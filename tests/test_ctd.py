"""Synthetic CTD excerpts test provenance and abstention, not filing suitability."""

from hashlib import sha256

from docx import Document
from openpyxl import Workbook
import pytest
from reportlab.pdfgen import canvas

from agent.ctd import (CTD_SECTIONS, prepare_ctd_package, propose_ctd_product_options,
                       propose_ctd_substance_options)
from agent.multimodal_intake import collect_multimodal, generation_sources


def _sources(paths):
    return generation_sources(collect_multimodal(paths))


def _txt(path, text):
    path.write_text(text, encoding='utf-8')
    return path


def _section(result, identifier):
    return next(item for item in result['sections'] if item['section_id'] == identifier)


def test_ctd_multiline_evidence_is_exact_and_not_a_submission(tmp_path):
    path = _txt(tmp_path / 'stability.txt',
                '제품명: 시험정\n제형: 정제\n함량: 5 mg\n\n'
                '3.2.P.8.3: 안정성 자료\n배치번호: B-001\n시험 시점: 12개월\n함량: 98.7%\n')
    result = prepare_ctd_package(_sources([path]), product_name='시험정',
                                 product_variant='정제 5 mg', selected_sections=['3.2.P.8.3'])
    section = _section(result, '3.2.P.8.3')
    assert section['status'] == 'proposed'
    assert '12개월' in section['draft'] and '98.7%' in section['draft']
    assert all(item['document_sha256'] == sha256(path.read_bytes()).hexdigest()
               for item in section['evidence'])
    assert all(item['quote'] in section['draft'] and item['source_id'].startswith('S')
               for item in section['evidence'])
    assert not result['submission_ready'] and result['actual_model_requests'] == 0
    assert result['coverage']['selected_sections'] == 1


def test_ctd_other_product_and_other_section_never_fill_selected_field(tmp_path):
    alpha = _txt(tmp_path / 'alpha.txt', '제품명: ALPHA정\n3.2.S.4: 원료 규격 99.0%\n')
    beta = _txt(tmp_path / 'beta.txt', '제품명: BETA정\n3.2.P.5: 완제 규격 95.0%\n')
    result = prepare_ctd_package(_sources([alpha, beta]), product_name='ALPHA정',
                                 selected_sections=['3.2.S.4', '3.2.P.5'])
    assert '99.0%' in _section(result, '3.2.S.4')['draft']
    assert _section(result, '3.2.P.5')['status'] == 'deferred'
    assert result['deferred_sections'] == ['3.2.P.5']
    assert not _section(result, '3.2.P.5')['draft']
    assert any(item.get('section_id') == '3.2.P.5' for item in result['deferred_sources'])


def test_ctd_same_field_conflict_stays_ambiguous_but_different_time_is_retained(tmp_path):
    conflict = _txt(tmp_path / 'conflict.txt', '제품명: ALPHA정\n'
                    '3.2.P.1: 주성분 10 mg/정\n3.2.P.1: 주성분 20 mg/정\n')
    result = prepare_ctd_package(_sources([conflict]), product_name='ALPHA정',
                                 selected_sections=['3.2.P.1'])
    section = _section(result, '3.2.P.1')
    assert section['status'] == 'ambiguous' and not section['draft']
    assert len(section['evidence']) == 2
    selected_id = section['evidence'][0]['source_id']
    narrowed = prepare_ctd_package(_sources([conflict]), product_name='ALPHA정',
                                   selected_sections=['3.2.P.1'],
                                   section_map={'3.2.P.1': [selected_id]})
    assert _section(narrowed, '3.2.P.1')['status'] == 'proposed'
    assert '20 mg' not in _section(narrowed, '3.2.P.1')['draft']

    timepoints = _txt(tmp_path / 'timepoints.txt', '제품명: ALPHA정\n'
                      '3.2.P.8.3: 안정성 자료\n시점: 0개월\n시험결과: 99.1%\n'
                      '시점: 6개월\n시험결과: 98.7%\n')
    checked = prepare_ctd_package(_sources([timepoints]), product_name='ALPHA정',
                                  selected_sections=['3.2.P.8.3'])
    assert _section(checked, '3.2.P.8.3')['status'] == 'proposed'
    assert '99.1%' in _section(checked, '3.2.P.8.3')['draft']
    assert '98.7%' in _section(checked, '3.2.P.8.3')['draft']

    batches = _txt(tmp_path / 'batches.txt', '제품명: ALPHA정\n'
                   '3.2.P.5.4: 배치 분석\n배치번호: B-001\n시험결과: 98.7%\n'
                   '배치번호: B-002\n시험결과: 99.1%\n')
    two_batches = prepare_ctd_package(_sources([batches]), product_name='ALPHA정',
                                      selected_sections=['3.2.P.5.4'])
    assert _section(two_batches, '3.2.P.5.4')['status'] == 'proposed'
    same_batch = _txt(tmp_path / 'same_batch.txt', '제품명: ALPHA정\n'
                      '3.2.P.5.4: 배치 분석\n배치번호: B-001\n시험결과: 98.7%\n'
                      '시험결과: 99.1%\n')
    conflicted = prepare_ctd_package(_sources([same_batch]), product_name='ALPHA정',
                                     selected_sections=['3.2.P.5.4'])
    assert _section(conflicted, '3.2.P.5.4')['status'] == 'ambiguous'


def test_ctd_m1_keeps_non_personal_cells_without_signature(tmp_path):
    path = _txt(tmp_path / 'admin.txt', '제품명: ALPHA정\n'
                '신청서: 신청 구분: 신규 | 신청인: 홍길동 | 제조소: TEST INC.\n'
                '서명: 홍길동\n신청 유형: 제조판매\n')
    result = prepare_ctd_package(_sources([path]), product_name='ALPHA정',
                                 selected_sections=['M1_ADMIN'])
    section = _section(result, 'M1_ADMIN')
    assert section['status'] == 'proposed'
    assert '신청 구분: 신규' in section['draft'] and '제조소: TEST INC.' in section['draft']
    assert '홍길동' not in section['draft'] and '서명' not in section['draft']
    assert any('직접 확인' in item['reason'] for item in result['deferred_sources'])

    split = _txt(tmp_path / 'admin_split.txt', '제품명: ALPHA정\n'
                 '신청서: 신청 구분: 신규 | 신청인 | 홍길동 | 제조소: TEST INC.\n'
                 '서명\n홍길동\n신청 유형: 제조판매\n')
    checked = prepare_ctd_package(_sources([split]), product_name='ALPHA정',
                                  selected_sections=['M1_ADMIN'])
    assert _section(checked, 'M1_ADMIN')['status'] == 'proposed'
    assert '홍길동' not in _section(checked, 'M1_ADMIN')['draft']
    assert '제조소: TEST INC.' in _section(checked, 'M1_ADMIN')['draft']


def test_ctd_office_parsers_and_high_value_sections_do_not_invent_crosswalk(tmp_path):
    docx = tmp_path / 'formula.docx'
    document = Document()
    document.add_paragraph('제품명: ALPHA정')
    document.add_paragraph('3.2.P.3.2: 배치 처방')
    document.add_paragraph('배치번호: B-001')
    document.add_paragraph('투입량: 10 g / 1000정')
    document.save(docx)

    xlsx = tmp_path / 'stability.xlsx'
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(['제품명', 'ALPHA정'])
    sheet.append(['3.2.P.8.3', '안정성 자료'])
    sheet.append(['배치번호', 'B-002'])
    sheet.append(['시험 시점', '12개월'])
    sheet.append(['시험결과', '99.1%'])
    workbook.save(xlsx)

    result = prepare_ctd_package(_sources([docx, xlsx]), product_name='ALPHA정',
                                 selected_sections=['3.2.P.3.2', '3.2.P.8.3'])
    formula = _section(result, '3.2.P.3.2')
    stability = _section(result, '3.2.P.8.3')
    assert formula['status'] == 'proposed' and '10 g / 1000정' in formula['draft']
    assert stability['status'] == 'proposed' and '99.1%' in stability['draft']
    assert 'B-001' not in stability['draft'] and 'B-002' not in formula['draft']
    assert any(item['sheet'] for item in stability['evidence'])
    assert not result['submission_ready']


def test_ctd_pdf_coa_and_xlsx_stability_keep_distinct_batches(tmp_path):
    pdf = tmp_path / 'coa.pdf'
    drawing = canvas.Canvas(str(pdf))
    for offset, text in enumerate(('Product name: ALPHA', '3.2.P.5.4: Batch analyses',
                                   'Batch ID: B-001', 'Assay: 98.7%')):
        drawing.drawString(48, 780 - offset * 20, text)
    drawing.save()
    xlsx = tmp_path / 'stability.xlsx'
    workbook = Workbook()
    sheet = workbook.active
    for row in (('Product name', 'ALPHA'), ('3.2.P.8.3', 'Stability data'),
                ('Batch ID', 'B-002'), ('Time point', '12 months'), ('Assay', '99.1%')):
        sheet.append(row)
    workbook.save(xlsx)
    result = prepare_ctd_package(_sources([pdf, xlsx]), product_name='ALPHA',
                                 selected_sections=['3.2.P.5.4', '3.2.P.8.3'])
    assay = _section(result, '3.2.P.5.4')
    stability = _section(result, '3.2.P.8.3')
    assert assay['status'] == 'proposed' and '98.7%' in assay['draft']
    assert stability['status'] == 'proposed' and '99.1%' in stability['draft']
    assert 'B-001' not in stability['draft'] and 'B-002' not in assay['draft']
    assert assay['evidence'][0]['page'] == 1


def test_ctd_rejects_unconfirmed_or_tampered_sources_and_unknown_sections(tmp_path):
    sources = _sources([_txt(tmp_path / 's.txt', '제품명: ALPHA정\n3.2.P.1: 10 mg/정\n')])
    sources[1]['requires_verification'] = True
    result = prepare_ctd_package(sources, product_name='ALPHA정',
                                 selected_sections=['3.2.P.1'])
    assert _section(result, '3.2.P.1')['status'] == 'deferred'
    assert result['deferred_sources']
    with pytest.raises(ValueError):
        prepare_ctd_package([], product_name='ALPHA정', selected_sections=['2.3'])
    assert '2.3' not in {item['section_id'] for item in CTD_SECTIONS}


def test_ctd_product_options_require_literal_single_product(tmp_path):
    one = _txt(tmp_path / 'one.txt', '제품명: ALPHA정\n제형: 정제\n함량: 5 mg\n'
               '3.2.P.5.4: 배치 분석\n시험결과: 98.7%\n')
    options = propose_ctd_product_options(_sources([one]))
    assert options['proposed_product_name'] == 'ALPHA정'
    assert options['proposed_product_variant'] == '정제 5 mg'
    assert options['requires_confirmation']
    assert options['product_names'][0]['evidence'][0]['quote'] == 'ALPHA정'
    other = _txt(tmp_path / 'other.txt', '제품명: BETA정\n3.2.P.1: 20 mg/정\n')
    mixed = propose_ctd_product_options(_sources([one, other]))
    assert mixed['proposed_product_name'] is None
    assert mixed['proposed_product_variant'] is None
    assert {item['value'] for item in mixed['product_names']} == {'ALPHA정', 'BETA정'}


def test_ctd_confirmed_api_link_is_s_only_and_original_bound(tmp_path):
    coa = tmp_path / 'api_coa.pdf'
    drawing = canvas.Canvas(str(coa))
    for offset, text in enumerate(('Material name: API-X', '3.2.S.4: Control of drug substance',
                                   'Batch ID: API-B-01', 'Assay: 99.5%')):
        drawing.drawString(48, 780 - offset * 20, text)
    drawing.save()
    dp = _txt(tmp_path / 'finished.txt', '제품명: ALPHA정\n제형: 정제\n함량: 5 mg\n'
              '3.2.P.5.4: 배치 분석\n배치번호: DP-B-01\n시험결과: 98.7%\n')
    sources = _sources([coa, dp])
    options = propose_ctd_substance_options(sources)
    assert len(options) == 1 and options[0]['proposed_substance_name'] == 'API-X'
    assert options[0]['document_sha256'] == sha256(coa.read_bytes()).hexdigest()
    without = prepare_ctd_package(sources, product_name='ALPHA정', product_variant='정제 5 mg',
                                  selected_sections=['3.2.S.4', '3.2.P.5.4'])
    assert _section(without, '3.2.S.4')['status'] == 'deferred'
    link = {sha256(coa.read_bytes()).hexdigest():
            {'substance_name': 'API-X', 'product_name': 'ALPHA정', 'confirmed': True}}
    result = prepare_ctd_package(sources, product_name='ALPHA정', product_variant='정제 5 mg',
                                 selected_sections=['3.2.S.4', '3.2.P.5.4'],
                                 confirmed_substance_links=link)
    api = _section(result, '3.2.S.4')
    finished = _section(result, '3.2.P.5.4')
    assert api['status'] == 'manual_check' and '99.5%' in api['draft']
    assert api['evidence'][0]['source_scope'] == 'confirmed_substance_link'
    assert finished['status'] == 'proposed' and '98.7%' in finished['draft']
    assert '99.5%' not in finished['draft'] and 'DP-B-01' not in api['draft']
    assert any('원료–완제 연결' in notice for notice in result['manual_checks'])
    with pytest.raises(ValueError, match='원료명'):
        prepare_ctd_package(sources, product_name='ALPHA정',
                            selected_sections=['3.2.S.4'],
                            confirmed_substance_links={sha256(coa.read_bytes()).hexdigest():
                                {'substance_name': 'API-Y', 'product_name': 'ALPHA정', 'confirmed': True}})


def test_ctd_api_link_rejects_conflicting_names_in_same_original(tmp_path):
    mixed = _txt(tmp_path / 'mixed.txt', '원료명: API-X\n제품명: API-Y\n'
                 '3.2.S.4: 원료의약품 관리\n시험결과: 99.5%\n')
    sources = _sources([mixed])
    assert propose_ctd_substance_options(sources)[0]['proposed_substance_name'] is None
    with pytest.raises(ValueError, match='원료명'):
        prepare_ctd_package(sources, product_name='ALPHA정',
                            selected_sections=['3.2.S.4'],
                            confirmed_substance_links={sha256(mixed.read_bytes()).hexdigest():
                                {'substance_name': 'API-X', 'product_name': 'ALPHA정', 'confirmed': True}})
