"""Native list round trips use raw OOXML fixtures, without network or Excel."""
from copy import deepcopy
from hashlib import sha256
from pathlib import Path
from zipfile import ZipFile

from lxml import etree
from openpyxl import load_workbook
import pytest

from agent.output_check import verify_output
from templates import TemplateError, analyze_template, fill_compatible_template
from templates.xlsx_validation import NS, XlsxLists


ROOT = Path(__file__).resolve().parents[1]
PART = 'xl/worksheets/sheet1.xml'


def rewrite(path, part, transform):
    with ZipFile(path) as archive:
        entries = [(info, archive.read(info.filename)) for info in archive.infolist()]
    with ZipFile(path, 'w') as archive:
        for info, data in entries:
            archive.writestr(info, transform(data) if info.filename == part else data)


def cell(row, coordinate, value=None, *, kind='inlineStr', style=None):
    attributes = {'r': coordinate}
    if style is not None:
        attributes['s'] = str(style)
    node = etree.SubElement(row, f'{{{NS}}}c', **attributes)
    if value is not None:
        node.set('t', kind)
        if kind == 'inlineStr':
            etree.SubElement(etree.SubElement(node, f'{{{NS}}}is'), f'{{{NS}}}t').text = value
        else:
            etree.SubElement(node, f'{{{NS}}}v').text = value
    return node


def book(tmp_path, formula='"승인,반려,승인"', *, target='B2', missing=None,
         names=(), source_values=('승인', '반려'), source_kinds=(), row_style=False, col_style=False):
    # Modify only disposable XML fixtures; production filling keeps all ZIP assets.
    with ZipFile(ROOT / 'samples/sample_company_form.xlsx') as archive:
        parts = {info.filename: archive.read(info.filename) for info in archive.infolist()}
    workbook = etree.fromstring(parts['xl/workbook.xml'])
    sheet = workbook.find(f'./{{{NS}}}sheets/{{{NS}}}sheet')
    sheet.set('name', '입력')
    etree.SubElement(sheet.getparent(), f'{{{NS}}}sheet', name="선택'목록", sheetId='2',
                     **{'{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id': 'rIdChoices'})
    defined = workbook.find(f'{{{NS}}}definedNames')
    if defined is None:
        defined = etree.Element(f'{{{NS}}}definedNames')
        workbook.insert(list(workbook).index(sheet.getparent()) + 1, defined)
    for name, expression, scope in names:
        node = etree.SubElement(defined, f'{{{NS}}}definedName', name=name)
        if scope is not None:
            node.set('localSheetId', str(scope))
        node.text = expression
    parts['xl/workbook.xml'] = etree.tostring(workbook)
    rels = etree.fromstring(parts['xl/_rels/workbook.xml.rels'])
    etree.SubElement(rels, '{http://schemas.openxmlformats.org/package/2006/relationships}Relationship',
                     Id='rIdChoices', Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet',
                     Target='worksheets/sheet2.xml')
    parts['xl/_rels/workbook.xml.rels'] = etree.tostring(rels)
    types = etree.fromstring(parts['[Content_Types].xml'])
    etree.SubElement(types, '{http://schemas.openxmlformats.org/package/2006/content-types}Override',
                     PartName='/xl/worksheets/sheet2.xml',
                     ContentType='application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml')
    parts['[Content_Types].xml'] = etree.tostring(types)
    root = etree.Element(f'{{{NS}}}worksheet', nsmap={None: NS})
    if col_style:
        etree.SubElement(etree.SubElement(root, f'{{{NS}}}cols'), f'{{{NS}}}col', min='2', max='2', style='1', width='20')
    data = etree.SubElement(root, f'{{{NS}}}sheetData')
    for number in (1, 2, 3, 5):
        if missing == 'row' and number == 2:
            continue
        row = etree.SubElement(data, f'{{{NS}}}row', r=str(number))
        if row_style and number == 2:
            row.set('s', '1')
            row.set('customFormat', '1')
        cell(row, f'A{number}', '구분' if number == 2 else f'항목{number}')
        if not (missing in {'cell', 'row'} and number == 2):
            cell(row, f'B{number}', style=1)
        if number == 3:
            calculated = cell(row, 'C3')
            etree.SubElement(calculated, f'{{{NS}}}f').text = '1+2'
    validation = etree.SubElement(etree.SubElement(root, f'{{{NS}}}dataValidations', count='1'),
                                  f'{{{NS}}}dataValidation', type='list', sqref=target, allowBlank='1')
    etree.SubElement(validation, f'{{{NS}}}formula1').text = formula
    parts[PART] = etree.tostring(root)
    options = etree.Element(f'{{{NS}}}worksheet', nsmap={None: NS})
    data = etree.SubElement(options, f'{{{NS}}}sheetData')
    for index, value in enumerate(source_values, 1):
        row = etree.SubElement(data, f'{{{NS}}}row', r=str(index))
        kind = source_kinds[index - 1] if source_kinds else 'inlineStr'
        cell(row, f'A{index}', value, kind=kind)
    parts['xl/worksheets/sheet2.xml'] = etree.tostring(options)
    path = tmp_path / 'original.xlsx'
    with ZipFile(path, 'w') as archive:
        for name, contents in parts.items():
            archive.writestr(name, contents)
    return path


def field(profile, coordinate='B2'):
    return next(item for item in profile['fields'] if item['id'] == f'xlsx:{PART}:{coordinate}')


def fill(source, tmp_path, value='승인', *, profile=None, coordinate='B2'):
    profile = profile or analyze_template(source)
    entry = field(profile, coordinate)
    values = {'선택': value}
    mapping = {entry['id']: '선택'}
    output = fill_compatible_template(source, values, tmp_path / 'filled.xlsx', mapping, profile)
    return output, values, mapping, profile


def test_inline_options_exact_deduplicated_and_assets_unchanged(tmp_path):
    source = book(tmp_path)
    original = source.read_bytes()
    profile = analyze_template(source)
    choice = field(profile)
    assert choice['control_type'] == 'choice'
    assert choice['input_required'] is True and choice['narrative_style_required'] is False
    assert choice['options'] == ['승인', '반려']
    assert choice['choice_items'] == [{'value': text, 'label': text} for text in choice['options']]
    output, values, mapping, _ = fill(source, tmp_path, profile=profile)
    assert verify_output(source, output, values, mapping=mapping, profile=profile)['status'] == 'passed'
    assert source.read_bytes() == original
    with ZipFile(source) as before, ZipFile(output) as after:
        assert all(before.read(name) == after.read(name) for name in before.namelist() if name != PART)


def test_inline_doubled_quote_is_a_literal_option_character(tmp_path):
    source = book(tmp_path, '"Alpha ""Q"",Beta"')
    profile = analyze_template(source)
    assert field(profile)['options'] == ['Alpha "Q"', 'Beta']
    output, values, mapping, _ = fill(source, tmp_path, 'Alpha "Q"', profile=profile)
    assert verify_output(source, output, values, mapping=mapping, profile=profile)['status'] == 'passed'


@pytest.mark.parametrize('value', ['미정', '승인함', '승인 [UNKNOWN]', '승인\n반려'])
def test_non_option_strings_are_rejected(tmp_path, value):
    source = book(tmp_path)
    with pytest.raises(TemplateError, match='선택|목록'):
        fill(source, tmp_path, value)
    assert not (tmp_path / 'filled.xlsx').exists()


@pytest.mark.parametrize('formula', ["'선택''목록'!$A$1:$A$2", "='선택''목록'!$A$1:$A$2", '선택목록'])
def test_quoted_unicode_other_sheet_and_workbook_named_range(tmp_path, formula):
    source = book(tmp_path, formula, names=[('선택목록', "'선택''목록'!$A$1:$A$2", None)])
    assert field(analyze_template(source))['options'] == ['승인', '반려']
    output, values, mapping, profile = fill(source, tmp_path, '반려')
    assert verify_output(source, output, values, mapping=mapping, profile=profile)['status'] == 'passed'


def test_local_name_shadows_workbook_and_other_sheet_scope(tmp_path):
    source = book(tmp_path, '목록', names=[('목록', "'선택''목록'!$A$1", None),
                                           ('목록', "'선택''목록'!$A$2", 0),
                                           ('목록', "'선택''목록'!$A$1", 1)])
    assert field(analyze_template(source))['options'] == ['반려']
    output, values, mapping, profile = fill(source, tmp_path, '반려')
    assert verify_output(source, output, values, mapping=mapping, profile=profile)['status'] == 'passed'


def test_same_sheet_literal_range(tmp_path):
    source = book(tmp_path, '$A$1:$A$2')
    assert field(analyze_template(source))['options'] == ['항목1', '구분']


def test_single_target_relative_literal_is_read_without_invented_recalculation(tmp_path):
    source = book(tmp_path, 'A1:A2')
    assert field(analyze_template(source))['options'] == ['항목1', '구분']


@pytest.mark.parametrize('formula,names', [('A1:A2', []), ('상대목록', [('상대목록', '입력!A1:A2', None)])])
def test_relative_multi_target_or_defined_name_is_unresolved_instead_of_fixed_wrong_options(tmp_path, formula, names):
    source = book(tmp_path, formula, target='B2:B3', names=names)
    profile = analyze_template(source)
    assert field(profile)['control_type'] == field(profile, 'B3')['control_type'] == 'choice_unresolved'
    assert '상대참조' in field(profile)['unsupported_reason']


def test_multiline_label_keeps_original_location_but_suggestion_is_normalized(tmp_path):
    source = book(tmp_path)
    def label(data):
        root = etree.fromstring(data)
        root.xpath(".//s:c[@r='A2']/s:is/s:t", namespaces={'s': NS})[0].text = '선택\n구분'
        return etree.tostring(root)
    rewrite(source, PART, label)
    profile = analyze_template(source)
    assert field(profile)['label'] == '선택\n구분'
    assert field(profile)['value_key'] == '선택 구분'
    assert field(profile)['id'] == f'xlsx:{PART}:B2'
    output, values, mapping, _ = fill(source, tmp_path, profile=profile)
    assert verify_output(source, output, values, mapping=mapping, profile=profile)['status'] == 'passed'
    original = load_workbook(source)
    rendered = load_workbook(output)
    assert original.active['A2'].value == rendered.active['A2'].value == '선택\n구분'
    original.close()
    rendered.close()


def test_normalized_suggestion_collision_keeps_distinct_ids_and_rejects_conflicting_lists(tmp_path):
    source = book(tmp_path)
    def labels(data):
        root = etree.fromstring(data)
        root.xpath(".//s:c[@r='A2']/s:is/s:t", namespaces={'s': NS})[0].text = '구\n분'
        root.xpath(".//s:c[@r='A3']/s:is/s:t", namespaces={'s': NS})[0].text = '구 분'
        parent = root.find(f'{{{NS}}}dataValidations')
        parent.set('count', '2')
        dv = etree.SubElement(parent, f'{{{NS}}}dataValidation', type='list', sqref='B3')
        etree.SubElement(dv, f'{{{NS}}}formula1').text = '"예,아니오"'
        return etree.tostring(root)
    rewrite(source, PART, labels)
    profile = analyze_template(source)
    assert field(profile)['value_key'] == field(profile, 'B3')['value_key'] == '구 분'
    assert field(profile)['id'] != field(profile, 'B3')['id']
    assert field(profile)['options'] != field(profile, 'B3')['options']
    with pytest.raises(TemplateError, match='상충'):
        fill_compatible_template(source, {'구 분': '승인'}, tmp_path / 'filled.xlsx', profile=profile)


def test_x14_literal_list_extension_round_trip_keeps_extension_bytes(tmp_path):
    source = book(tmp_path, "'선택''목록'!$A$1:$A$2")
    def extension(data):
        root = etree.fromstring(data)
        native = root.find(f'{{{NS}}}dataValidations')
        root.remove(native)
        x14 = 'http://schemas.microsoft.com/office/spreadsheetml/2009/9/main'
        xm = 'http://schemas.microsoft.com/office/excel/2006/main'
        ext = etree.SubElement(etree.SubElement(root, f'{{{NS}}}extLst'), f'{{{NS}}}ext', uri='{test}', nsmap={'x14': x14, 'xm': xm})
        dv = etree.SubElement(etree.SubElement(ext, f'{{{x14}}}dataValidations', count='1'), f'{{{x14}}}dataValidation', type='list')
        etree.SubElement(etree.SubElement(dv, f'{{{x14}}}formula1'), f'{{{xm}}}f').text = "'선택''목록'!$A$1:$A$2"
        etree.SubElement(dv, f'{{{xm}}}sqref').text = 'B2'
        return etree.tostring(root)
    rewrite(source, PART, extension)
    output, values, mapping, profile = fill(source, tmp_path)
    assert field(profile)['control_type'] == 'choice'
    assert verify_output(source, output, values, mapping=mapping, profile=profile)['status'] == 'passed'
    with ZipFile(source) as before, ZipFile(output) as after:
        xpath = './*[local-name()="extLst"]'
        assert etree.tostring(etree.fromstring(before.read(PART)).xpath(xpath)[0]) == etree.tostring(etree.fromstring(after.read(PART)).xpath(xpath)[0])


def test_boundary_space_options_are_unresolved_and_cannot_be_silently_trimmed(tmp_path):
    source = book(tmp_path, '"승인, 반려"')
    profile = analyze_template(source)
    assert field(profile)['control_type'] == 'choice_unresolved'
    assert '공백' in field(profile)['unsupported_reason']
    with pytest.raises(TemplateError, match='목록|공백'):
        fill(source, tmp_path, '반려', profile=profile)


def test_raw_zero_long_numeric_identifier_boolean_and_duplicate(tmp_path):
    source = book(tmp_path, "'선택''목록'!$A$1:$A$5",
                  source_values=('0', '12345678901234567890', '1', '0', '0'),
                  source_kinds=('n', 'n', 'b', 'b', 'n'))
    profile = analyze_template(source)
    assert field(profile)['options'] == ['0', '12345678901234567890', 'TRUE', 'FALSE']
    output, values, mapping, _ = fill(source, tmp_path, '12345678901234567890', profile=profile)
    assert verify_output(source, output, values, mapping=mapping, profile=profile)['status'] == 'passed'


@pytest.mark.parametrize('formula', ['INDIRECT("A1:A2")', '[Book.xlsx]Sheet!$A$1:$A$2',
    "'선택''목록'!$A:$A", "'선택''목록'!$A$1:$A$2,'선택''목록'!$A$4", 'OFFSET(A1,0,0,2)',
    '알수없는목록', '""'])
def test_unresolved_native_lists_never_become_free_text(tmp_path, formula):
    source = book(tmp_path, formula)
    profile = analyze_template(source)
    choice = field(profile)
    assert choice['control_type'] == 'choice_unresolved' and choice['options'] == []
    assert choice['unsupported_reason'] and any('목록 확인 필요' in item for item in profile['warnings'])
    forged = deepcopy(profile)
    forged_field = field(forged)
    for key in ('control_type', 'options', 'choice_items', 'xlsx_list'):
        forged_field.pop(key, None)
    with pytest.raises(TemplateError, match='목록 확인 필요'):
        fill(source, tmp_path, '임의 문구', profile=forged)


@pytest.mark.parametrize('key,value', [('options', ['승인', '반려', '미정']),
    ('xlsx_list', {'sqref': 'B3', 'formula1': '"승인,반려"', 'sha256': '0' * 64}),
    ('choice_items', [{'label': '승인', 'value': '반려'}]), ('control_type', 'text')])
def test_profile_options_address_and_type_cannot_override_native_rule(tmp_path, key, value):
    source = book(tmp_path)
    profile = analyze_template(source)
    field(profile)[key] = value
    with pytest.raises(TemplateError, match='프로파일|선택'):
        fill(source, tmp_path, profile=profile)


@pytest.mark.parametrize('missing,row_style,col_style', [('cell', True, False), ('cell', False, True), ('row', False, True)])
def test_missing_physical_cell_or_row_is_discovered_and_inherits_style(tmp_path, missing, row_style, col_style):
    source = book(tmp_path, missing=missing, row_style=row_style, col_style=col_style)
    digest = sha256(source.read_bytes()).hexdigest()
    output, values, mapping, profile = fill(source, tmp_path)
    assert field(profile)['control_type'] == 'choice'
    assert verify_output(source, output, values, mapping=mapping, profile=profile)['status'] == 'passed'
    with ZipFile(output) as archive:
        root = etree.fromstring(archive.read(PART))
        created = root.xpath(".//s:c[@r='B2']", namespaces={'s': NS})[0]
        assert created.attrib == {'r': 'B2', 's': '1', 't': 'inlineStr'}
    assert sha256(source.read_bytes()).hexdigest() == digest


@pytest.mark.parametrize('mutation', ['style', 'extra_cell', 'row_property', 'existing_loss', 'order', 'dv_formula'])
def test_created_cell_exception_still_rejects_structure_and_validation_tampering(tmp_path, mutation):
    source = book(tmp_path, missing='cell', row_style=True)
    output, values, mapping, profile = fill(source, tmp_path)
    def damage(data):
        root = etree.fromstring(data)
        selected = root.xpath(".//s:c[@r='B2']", namespaces={'s': NS})[0]
        row = selected.getparent()
        if mutation == 'style':
            selected.set('s', '0')
        elif mutation == 'extra_cell':
            cell(row, 'C2', '무단')
        elif mutation == 'row_property':
            row.set('ht', '3')
        elif mutation == 'existing_loss':
            row.remove(row[0])
        elif mutation == 'order':
            row.remove(selected)
            row.insert(0, selected)
        else:
            root.find(f'.//{{{NS}}}formula1').text = '"승인,미정"'
        return etree.tostring(root)
    rewrite(output, PART, damage)
    with pytest.raises(ValueError, match='XLSX'):
        verify_output(source, output, values, mapping=mapping, profile=profile)


def test_formula_merged_continuation_and_protected_sheet_are_not_choices(tmp_path):
    source = book(tmp_path, target='B2:C3')
    def mutate(data):
        root = etree.fromstring(data)
        merge = etree.Element(f'{{{NS}}}mergeCells', count='1')
        etree.SubElement(merge, f'{{{NS}}}mergeCell', ref='B2:C2')
        root.insert(list(root).index(root.find(f'{{{NS}}}dataValidations')), merge)
        return etree.tostring(root)
    rewrite(source, PART, mutate)
    profile = analyze_template(source)
    choices = {item['id'].split(':')[2] for item in profile['fields'] if item.get('control_type') == 'choice'}
    assert choices == {'B2', 'B3'}
    rewrite(source, PART, lambda data: data.replace(b'<sheetData>', b'<sheetProtection sheet="1"/><sheetData>'))
    assert not any(item['id'].startswith(f'xlsx:{PART}:') for item in analyze_template(source)['fields'])


def test_source_formula_options_are_not_evaluated_from_cached_values(tmp_path):
    source = book(tmp_path, "'선택''목록'!$A$1:$A$2")
    def mutate(data):
        root = etree.fromstring(data)
        first = root.xpath(".//s:c[@r='A1']", namespaces={'s': NS})[0]
        etree.SubElement(first, f'{{{NS}}}f').text = '"승인"'
        return etree.tostring(root)
    rewrite(source, 'xl/worksheets/sheet2.xml', mutate)
    assert field(analyze_template(source))['control_type'] == 'choice_unresolved'


def test_phonetic_annotations_are_not_appended_to_actual_list_values(tmp_path):
    source = book(tmp_path, "'선택''목록'!$A$1:$A$2")
    def annotate(data):
        root = etree.fromstring(data)
        inline = root.xpath(".//s:c[@r='A1']/s:is", namespaces={'s': NS})[0]
        etree.SubElement(etree.SubElement(inline, f'{{{NS}}}rPh', sb='0', eb='2'), f'{{{NS}}}t').text = 'phonetic'
        return etree.tostring(root)
    rewrite(source, 'xl/worksheets/sheet2.xml', annotate)
    assert field(analyze_template(source))['options'] == ['승인', '반려']


def test_large_target_does_not_expand_unbounded_cells_but_existing_is_closed(tmp_path):
    source = book(tmp_path, target='B:B')
    profile = analyze_template(source)
    assert field(profile)['control_type'] == 'choice_unresolved'
    assert len(profile['fields']) < 50


def test_output_value_is_rechecked_at_exact_original_choice_cell(tmp_path):
    source = book(tmp_path)
    output, values, mapping, profile = fill(source, tmp_path)
    def mutate(data):
        root = etree.fromstring(data)
        root.xpath(".//s:c[@r='B2']/s:is/s:t", namespaces={'s': NS})[0].text = '반려'
        return etree.tostring(root)
    rewrite(output, PART, mutate)
    with pytest.raises(ValueError, match='선택값|입력값'):
        verify_output(source, output, values, mapping=mapping, profile=profile)
