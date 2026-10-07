"""External CTD working drafts use tenant-bound, checked real sample files."""

from base64 import b64encode
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
from zipfile import ZipFile

from docx import Document

from app import web_api
from test_web_api import SECRET, request
from fastapi.testclient import TestClient


SAMPLES = Path(__file__).resolve().parents[1] / 'samples'
SECTIONS = ['3.2.P.3.2', '3.2.P.5.4', '3.2.P.8.3']
INPUTS = {'product_name': '예시정', 'product_variant': '정제 5 mg',
          'selected_sections': SECTIONS}


def payload():
    names = ['ctd_demo_batch_formula.txt', 'ctd_demo_batch_analysis.txt',
             'ctd_demo_stability.txt']
    def file(name):
        return {'name': name, 'base64': b64encode((SAMPLES / name).read_bytes()).decode()}
    return {'template': file('ctd_demo_template.docx'), 'sources': [file(n) for n in names]}


def test_web_ctd_prepare_export_and_tenant_isolation(tmp_path):
    root = tmp_path / 'private'
    api = TestClient(web_api.create_app(root=root, secret=SECRET),
                     raise_server_exceptions=False)
    section_ids = {item['section_id'] for item in request(api, 'GET', '/api/ctd/sections').json()['sections']}
    assert set(SECTIONS) <= section_ids
    uploaded = request(api, 'POST', '/api/jobs', payload())
    assert uploaded.status_code == 200, uploaded.text
    identifier = uploaded.json()['job_id']
    path = f'/api/jobs/{identifier}/ctd'
    assert request(api, 'POST', path + '/prepare', INPUTS, user='user-B').status_code == 404
    prepared = request(api, 'POST', path + '/prepare', INPUTS)
    assert prepared.status_code == 200, prepared.text
    package = prepared.json()['ctd']['package']
    assert package['coverage']['proposed_sections'] == 3
    assert all(item['evidence'] for item in package['sections'])
    assert request(api, 'POST', path + '/export', {'confirmed': True,
        'fingerprint': package['fingerprint']}, user='user-B').status_code == 404
    assert request(api, 'POST', path + '/export', {'confirmed': True,
        'fingerprint': '0' * 64}).status_code == 409
    exported = request(api, 'POST', path + '/export', {'confirmed': True,
        'fingerprint': package['fingerprint']})
    assert exported.status_code == 200, exported.text
    with ZipFile(BytesIO(exported.content)) as archive:
        document = Document(BytesIO(archive.read('CTD_작업초안.docx')))
        content = '\n'.join([p.text for p in document.paragraphs] +
                            [cell.text for table in document.tables
                             for row in table.rows for cell in row.cells])
        assert 'DEMO-B01' in content and '98.7%' in content
        evidence = json.loads(archive.read('CTD_출처와_누락.json'))
        assert evidence['output_check']['status'] == 'passed'
        assert evidence['submission_ready'] is False
    template = root / sha256(b'user-A').hexdigest() / identifier / 'template' / 'ctd_demo_template.docx'
    assert template.read_bytes() == (SAMPLES / template.name).read_bytes()
    # Current source bytes are checked again at export, not just at upload.
    source = root / sha256(b'user-A').hexdigest() / identifier / 'sources' / 'ctd_demo_stability.txt'
    source.write_text('tampered', encoding='utf-8')
    assert request(api, 'POST', path + '/export', {'confirmed': True,
        'fingerprint': package['fingerprint']}).status_code == 409


def test_claude_ctd_preview_reads_only_own_confirmed_job(tmp_path):
    api = TestClient(web_api.create_app(root=tmp_path / 'private', secret=SECRET),
                     raise_server_exceptions=False)
    identifier = request(api, 'POST', '/api/jobs', payload()).json()['job_id']
    issued = request(api, 'POST', '/api/mcp/credentials', {}).json()['token']
    def call(token, job_id):
        return request(api, 'POST', '/mcp', {'jsonrpc': '2.0', 'id': 1,
            'method': 'tools/call', 'params': {'name': 'ra_ctd_preview',
            'arguments': {'job_id': job_id, **INPUTS}}}, user='mcp:' + token).json()['result']
    result = call(issued, identifier)
    assert result['isError'] is False
    assert result['structuredContent']['package']['coverage']['proposed_sections'] == 3
    assert 'ctd' not in request(api, 'GET', f'/api/jobs/{identifier}').json()
    other = request(api, 'POST', '/api/mcp/credentials', {}, user='user-B').json()['token']
    assert call(other, identifier)['isError'] is True
