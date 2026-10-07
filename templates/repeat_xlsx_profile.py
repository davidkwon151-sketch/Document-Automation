"""XLSX 반복 입력 위치와 확인된 원본 규칙을 연결함. 파일을 쓰지 않음."""

from copy import deepcopy
from hashlib import sha256
from io import BytesIO
from pathlib import Path
import re
from zipfile import ZipFile

from openpyxl.utils import get_column_letter, range_boundaries

from parsers.extract import _check_zip, _xml
from .compatibility import analyze_template, _xlsx_cell_text, _xlsx_parts
from .fill import PLACEHOLDER, TemplateError
from .repeat_fields import _binding, _repeat_key
from .repeat_rules import META, RULES, _without_values_and_paths
from .value_rules import validate_rule_profile
from .xlsx_validation import NS, XlsxLists, editable_blank


SHEET = re.compile(r'xl/worksheets/sheet[1-9][0-9]*\.xml\Z')


def _plan(plan):
    if (not isinstance(plan, dict) or set(plan) != {'table_id', 'row', 'count'}
            or not isinstance(plan['table_id'], str) or not plan['table_id'].startswith('xlsx:')
            or not SHEET.fullmatch(plan['table_id'][5:]) or type(plan['row']) is not int
            or type(plan['count']) is not int or not 1 <= plan['count'] <= 200
            or not 1 <= plan['row'] <= plan['row'] + plan['count'] - 1 <= 1048576):
        raise TemplateError('XLSX 반복 입력 계획의 시트·행·개수가 유효하지 않음')
    return plan['table_id'][5:]


def _location(field):
    try:
        prefix, part, coordinate, *tokens = field['id'].split(':', 3)
        column, row, last_column, last_row = range_boundaries(coordinate)
    except (KeyError, TypeError, ValueError, AttributeError) as exc:
        raise TemplateError('XLSX 입력 ID의 실제 셀 주소가 유효하지 않음') from exc
    if (prefix != 'xlsx' or not SHEET.fullmatch(part) or len(tokens) > 1
            or not re.fullmatch(r'[A-Z]{1,3}[1-9][0-9]{0,6}', coordinate)
            or not 1 <= column <= 16384 or not 1 <= row <= 1048576
            or (column, row) != (last_column, last_row)
            or field.get('kind') != ('xlsx_placeholder' if tokens else 'xlsx_cell')
            or tokens and not tokens[0].strip()):
        raise TemplateError('XLSX 입력 종류와 실제 셀 주소가 다름')
    return part, coordinate, column, row, tokens[0] if tokens else None


def _roots(archive):
    if any(name.startswith('_xmlsignatures/') for name in archive.namelist()):
        raise TemplateError('전자서명된 XLSX 원본 입력은 승계할 수 없음')
    shared, names = _xlsx_parts(archive)
    return shared, {name: _xml(archive.read(name)) for name in names}


def _cell(root, coordinate):
    nodes = root.xpath('./s:sheetData/s:row/s:c[@r=$coordinate]', namespaces={'s': NS}, coordinate=coordinate)
    if len(nodes) > 1:
        raise TemplateError('XLSX 실제 입력 셀 주소가 중복됨')
    if nodes and int(nodes[0].getparent().get('r')) != range_boundaries(coordinate)[1]:
        raise TemplateError('XLSX 실제 셀과 부모 행의 주소가 다름')
    return nodes[0] if nodes else None


def _header_rows(rows, shared, before):
    """Header hints only: ignore explicitly marked examples and numbered data rows."""
    result = []
    for row in rows:
        if int(row.get('r')) >= before:
            continue
        cells = [(cell, _xlsx_cell_text(cell, shared).strip())
                 for cell in row.findall(f'{{{NS}}}c')]
        printed = [(cell, text) for cell, text in cells if text]
        if not printed or any(PLACEHOLDER.search(text) for _, text in printed):
            continue
        first, text = printed[0]
        if re.fullmatch(r'\(?예\s*시\)?|example', text, re.IGNORECASE):
            continue
        if (len(printed) > 1 and first.get('t', 'n') == 'n'
                and first.find(f'{{{NS}}}f') is None
                and re.fullmatch(r'[1-9][0-9]{0,2}', text) and int(text) <= 200):
            continue
        result.append(row)
    return sorted(result, key=lambda row: int(row.get('r')), reverse=True)


def _check_field(roots, shared, lists, field, *, native_coordinate=None, exact_native=True):
    """실제 원문 위치/값/보호와 native 목록을 확인하며 프로파일 선언을 믿지 않음."""
    part, coordinate, _, _, token = _location(field)
    root = roots.get(part)
    if root is None or root.find(f'{{{NS}}}sheetProtection') is not None:
        raise TemplateError('XLSX 입력 시트가 없거나 편집 보호됨')
    if field.get('readonly') or field.get('read_only') or field.get('editable') is False:
        raise TemplateError('읽기 전용 XLSX 원본 입력은 승계할 수 없음')
    cell = _cell(root, coordinate)
    if cell is not None and cell.find(f'{{{NS}}}f') is not None:
        raise TemplateError('XLSX 수식 셀은 원본 입력으로 승계할 수 없음')
    column, row, _, _ = range_boundaries(coordinate)
    for merged in root.findall(f'./{{{NS}}}mergeCells/{{{NS}}}mergeCell'):
        x1, y1, x2, y2 = range_boundaries(merged.get('ref'))
        if x1 <= column <= x2 and y1 <= row <= y2 and (column, row) != (x1, y1):
            raise TemplateError('XLSX 병합 이어진 셀은 원본 입력으로 승계할 수 없음')
    rule = lists.at(part, native_coordinate or coordinate)
    if token is not None:
        text = _xlsx_cell_text(cell, shared)
        if rule is not None or cell is None or token not in {match[1].strip() for match in PLACEHOLDER.finditer(text)}:
            raise TemplateError('XLSX 원본 자리표시자의 동일 위치를 확인할 수 없음')
        if field.get('label') != token or field.get('required') is False:
            raise TemplateError('XLSX 자리표시자 이름·필수 조건을 바꿀 수 없음')
    elif not editable_blank(root, coordinate, shared) or cell is None and rule is None:
        raise TemplateError('XLSX 기존 고정 값 또는 없는 일반 입력 셀을 승계할 수 없음')
    if rule is not None:
        if rule['error']:
            raise TemplateError('XLSX 원본 선택 목록 확인 필요: ' + rule['error'])
        expected = {'control_type': 'choice', 'options': rule['options'],
                    'choice_items': [{'value': value, 'label': value} for value in rule['options']]}
        if any(field.get(key) != value for key, value in expected.items()):
            raise TemplateError('XLSX native 원본 선택 옵션·입력 종류를 바꿀 수 없음')
        if exact_native and field.get('xlsx_list') != lists.metadata(rule):
            raise TemplateError('XLSX native 원본 목록 주소·정의가 프로파일과 다름')
        if field.get('input_required') is not True or field.get('input_mode', 'user_provided') != 'user_provided':
            raise TemplateError('XLSX 원본 선택은 사용자 직접 입력이어야 함')
    elif field.get('control_type') in {'choice', 'choice_unresolved'} or field.get('xlsx_list'):
        raise TemplateError('XLSX 실제 원본에 없는 native 목록 선언임')


def _summaries(fields, count, constraints, warnings):
    summaries = []
    for row in range(1, count + 1):
        members = [field['value_key'] for field in fields if field.get('repeat_info', {}).get('row') == row]
        required = [field['value_key'] for field in fields if field.get('repeat_info', {}).get('row') == row and field.get('required') is True]
        if required and not any(set(group['fields']) == set(members) for group in constraints.get('groups', [])):
            constraints.setdefault('groups', []).append({'kind': 'all_or_none', 'fields': members, 'required_fields': required})
        elif members and not required:
            warnings.append(f'반복 {row}행: 원본 필수 항목 정의 없음; 행 필수 조건은 사용자 확인 필요')
        summaries.append({'row': row, 'fields': members, 'required_fields': required,
                          'required_definition_pending': bool(members and not required)})
    return summaries


def xlsx_repeat_profile(prepared_path, plan, binding):
    """준비본 분석 결과를 Excel 실제 행/열에 연결한 고유 입력 키로 반환함."""
    path = Path(prepared_path)
    part = _plan(plan)
    if path.suffix.lower() != '.xlsx':
        raise TemplateError('XLSX 반복 프로파일에는 XLSX 준비 파일이 필요함')
    raw = _binding(path, plan, binding)
    profile = deepcopy(analyze_template(path))
    if not profile.get('supported') or profile.get('source_sha256') != binding['prepared_sha256'] or path.read_bytes() != raw:
        raise TemplateError('XLSX 준비본 입력 위치 분석 또는 SHA 확인 실패')
    with ZipFile(BytesIO(raw)) as archive:
        _check_zip(archive)
        shared, roots = _roots(archive)
        if part not in roots:
            raise TemplateError('XLSX 준비본에 반복 대상 시트가 없음')
        if roots[part].find(f'{{{NS}}}sheetProtection') is not None:
            raise TemplateError('XLSX 반복 대상 시트가 편집 보호됨')
        rows = roots[part].findall(f'./{{{NS}}}sheetData/{{{NS}}}row')
        numbers = [int(row.get('r')) for row in rows]
        if len(set(numbers)) != len(numbers) or any(row not in numbers for row in range(plan['row'], plan['row'] + plan['count'])):
            raise TemplateError('XLSX 준비본의 실제 반복 행이 없거나 중복됨')
        fields, used = profile['fields'], set()
        occupied = {_location(field)[:2] for field in fields}
        for row in rows:
            if not plan['row'] <= int(row.get('r')) < plan['row'] + plan['count']:
                continue
            for cell in row.findall(f'{{{NS}}}c'):
                coordinate = cell.get('r')
                if (part, coordinate) in occupied or not editable_blank(roots[part], coordinate, shared):
                    continue
                label = f'빈 셀 {coordinate}'
                fields.append({'id': f'xlsx:{part}:{coordinate}', 'label': label, 'value_key': label,
                               'kind': 'xlsx_cell', 'required': False, 'input_required': False, 'confidence': 0.1})
        for field in fields:
            field_part, _, _, row, _ = _location(field)
            if field_part != part or not plan['row'] <= row < plan['row'] + plan['count']:
                used.add(field['value_key'])
        headers = _header_rows(rows, shared, plan['row'])
        for field in fields:
            field_part, _, column, row, token = _location(field)
            if field_part != part or not plan['row'] <= row < plan['row'] + plan['count']:
                continue
            # Native controls and tokens already have explicit names. Blank rows
            # use the prototype's header rather than a new clone's blank neighbour.
            label = field['label']
            if token is None and not field.get('control_type'):
                for header in headers:
                    text = _xlsx_cell_text(_cell(roots[part], f'{get_column_letter(column)}{header.get("r")}'), shared).strip()
                    if text and not PLACEHOLDER.search(text):
                        label = text
                        break
                field['label'] = label
            relative = row - plan['row'] + 1
            field.update(value_key=_repeat_key(field['value_key'] if token or field.get('control_type') else label, relative, column, used),
                         repeat_info={'row': relative, 'column': column, 'source_label': label})
            field['suggested_mapping'] = field['value_key']
        if not any(field.get('repeat_info') for field in fields):
            raise TemplateError('XLSX 반복 행에서 실제 입력칸을 찾지 못함')
    constraints = profile.setdefault('constraints', {})
    profile['repeat_rows'] = _summaries(fields, plan['count'], constraints, profile.setdefault('warnings', []))
    profile['repeat_expansion'] = deepcopy(binding)
    validate_rule_profile(profile)
    return profile


def _backproject(roots, part, start, count):
    """원본 없는 재열기에서 clone 제거/원위치 역투영. 실제 원본 SHA 검수와 구분함."""
    roots = deepcopy(roots)
    root = roots[part]
    delta, end = count - 1, start + count - 1
    for row in list(root.findall(f'./{{{NS}}}sheetData/{{{NS}}}row')):
        number = int(row.get('r'))
        if start < number <= end:
            row.getparent().remove(row)
        elif number > end:
            row.set('r', str(number - delta))
            for cell in row.findall(f'{{{NS}}}c'):
                column, actual, _, _ = range_boundaries(cell.get('r'))
                cell.set('r', f'{get_column_letter(column)}{actual - delta}')
    merges = root.find(f'{{{NS}}}mergeCells')
    if merges is not None:
        for node in list(merges):
            x1, y1, x2, y2 = range_boundaries(node.get('ref'))
            if start < y1 <= end:
                merges.remove(node)
            else:
                y1 -= delta if y1 > end else 0
                y2 -= delta if y2 > end else 0
                node.set('ref', f'{get_column_letter(x1)}{y1}:{get_column_letter(x2)}{y2}')
    return roots


def inherit_xlsx_repeat_rules(original_path, prepared_profile, original_profile, *, prepared_path=None):
    """확인 원본의 위치/권한/규칙을 복제하고 모호한 행간 관계를 거부함."""
    if original_profile.get('repeat_expansion'):
        raise TemplateError('이미 확장한 XLSX 원본의 중첩 반복 승계는 미지원')
    result = _without_values_and_paths(deepcopy(prepared_profile))
    binding = result.get('repeat_expansion', {})
    prepared = prepared_path or prepared_profile.get('source_path')
    if not prepared:
        raise TemplateError('XLSX 규칙 승계에는 준비 파일 경로가 필요함')
    path = Path(prepared)
    part = _plan(binding.get('plan'))
    plan = binding['plan']
    if original_profile.get('format') != 'xlsx' or result.get('format') != 'xlsx' or original_profile.get('source_sha256') != binding.get('original_sha256'):
        raise TemplateError('XLSX 반복 승계의 원본/준비 형식 또는 SHA가 다름')
    raw = _binding(path, plan, binding)
    if result.get('source_sha256') != binding['prepared_sha256']:
        raise TemplateError('XLSX 준비 프로파일 SHA가 반복 기록과 다름')
    source_policy = _without_values_and_paths(original_profile)
    validate_rule_profile(source_policy)
    with ZipFile(BytesIO(raw)) as prepared_archive:
        _check_zip(prepared_archive)
        prepared_shared, prepared_roots = _roots(prepared_archive)
        prepared_lists = XlsxLists(prepared_archive)
        source_archive = None
        try:
            if original_path is not None:
                original_raw = Path(original_path).read_bytes()
                if sha256(original_raw).hexdigest() != binding['original_sha256']:
                    raise TemplateError('XLSX 실제 원본 SHA가 반복 기록과 다름')
                source_archive = ZipFile(BytesIO(original_raw))
                _check_zip(source_archive)
                source_shared, source_roots = _roots(source_archive)
                source_lists = XlsxLists(source_archive)
            else:
                source_shared = prepared_shared
                source_roots = _backproject(prepared_roots, part, plan['row'], plan['count'])
                source_lists = prepared_lists
            available = {field['id']: field for field in result.get('fields', [])}
            if len(available) != len(result.get('fields', [])) or len({field['id'] for field in source_policy['fields']}) != len(source_policy['fields']):
                raise TemplateError('XLSX 원본/준비 입력 ID가 중복됨')
            for field in available.values():
                _check_field(prepared_roots, prepared_shared, prepared_lists, field)
                field_part, _, column, row, _ = _location(field)
                relative = row - plan['row'] + 1 if field_part == part and plan['row'] <= row < plan['row'] + plan['count'] else 0
                info = field.get('repeat_info')
                if bool(info) != bool(relative) or info and (info.get('row') != relative or info.get('column') != column):
                    raise TemplateError('XLSX 반복 입력의 실제 행·열과 반복 위치가 다름')
            assignments = {}
            delta = plan['count'] - 1
            for original in source_policy['fields']:
                source_part, coordinate, column, row, token = _location(original)
                shifted = row + delta if source_part == part and row > plan['row'] else row
                native_coordinate = f'{get_column_letter(column)}{shifted}' if original_path is None else None
                _check_field(source_roots, source_shared, source_lists, original, native_coordinate=native_coordinate, exact_native=original_path is not None)
                destinations = [(relative, plan['row'] + relative - 1) for relative in range(1, plan['count'] + 1)] if source_part == part and row == plan['row'] else [(0, shifted)]
                for relative, target_row in destinations:
                    identifier = f'xlsx:{source_part}:{get_column_letter(column)}{target_row}' + (f':{token}' if token is not None else '')
                    field = available.get(identifier)
                    if field is None:
                        # A confirmed blank cell can outrank a raw label heuristic;
                        # this still checks the exact source and destination cell.
                        field = deepcopy(original)
                        field['id'] = identifier
                        if relative:
                            field['repeat_info'] = {'row': relative, 'column': column, 'source_label': original['label']}
                        rule = prepared_lists.at(source_part, f'{get_column_letter(column)}{target_row}')
                        if rule is not None:
                            field['xlsx_list'] = prepared_lists.metadata(rule)
                        available[identifier] = field
                        result['fields'].append(field)
                    _check_field(prepared_roots, prepared_shared, prepared_lists, field)
                    for key in ('control_type', 'options', 'choice_items', 'allow_custom'):
                        if key in original and original[key] != field.get(key):
                            raise TemplateError('XLSX 원본 native 선택 범위를 준비본에서 바꿀 수 없음')
                    if identifier in assignments:
                        raise TemplateError('서로 다른 XLSX 원본 규칙의 입력 위치가 겹침')
                    assignments[identifier] = (original, relative)
        finally:
            if source_archive is not None:
                source_archive.close()
    fields = [field for field in result['fields'] if field.get('repeat_info') or field['id'] in assignments
              or field['kind'] == 'xlsx_placeholder' and field.get('required')]
    used = {original['value_key'] for original, relative in assignments.values() if not relative}
    used.update(field['value_key'] for field in fields if field['id'] not in assignments)
    destinations = {}
    for field in fields:
        if field['id'] not in assignments:
            continue
        original, relative = assignments[field['id']]
        for key in RULES:
            if key in original:
                field[key] = deepcopy(original[key])
        if field['kind'] == 'xlsx_placeholder' and field.get('required') is False:
            raise TemplateError('XLSX 원본 자리표시자의 필수 조건을 완화할 수 없음')
        field['label'] = original['label']
        if relative:
            field['repeat_info']['source_label'] = original['label']
            field['value_key'] = _repeat_key(original['value_key'], relative, field['repeat_info']['column'], used)
        else:
            field['value_key'] = original['value_key']
        field['suggested_mapping'] = field['value_key']
        destinations.setdefault(original['value_key'], {}).setdefault(relative, set()).add(field['value_key'])
    def rewrite(item, keys):
        mapping, scopes = {}, set()
        for key in keys:
            positions = destinations.get(key, {})
            if not positions or any(len(values) != 1 for values in positions.values()):
                raise TemplateError('XLSX 원본 관계의 동일 값을 모호하게 나눌 수 없음')
            scopes.add('outside' if set(positions) == {0} else 'repeat' if set(positions) == set(range(1, plan['count'] + 1)) else 'mixed')
            mapping[key] = positions
        if scopes == {'outside'}:
            rows = [0]
        elif scopes == {'repeat'}:
            rows = range(1, plan['count'] + 1)
        else:
            raise TemplateError('XLSX 반복행·바깥 또는 복수 원본 행을 연결하는 관계는 승계 미지원')
        output = []
        for row in rows:
            names = {key: next(iter(mapping[key][row])) for key in keys}
            cloned = deepcopy(item)
            for name in ('total', 'left', 'right', 'start', 'end', 'year', 'month', 'day'):
                if name in cloned:
                    cloned[name] = names[cloned[name]]
            for name in ('parts', 'fields', 'required_fields'):
                if name in cloned:
                    cloned[name] = [names[key] for key in cloned[name]]
            output.append(cloned)
        return output
    constraints = deepcopy(source_policy.get('constraints', {}))
    constraints.update(relations=[], groups=[])
    for relation in source_policy.get('constraints', {}).get('relations', []):
        keys = [relation[name] for name in ('total', 'left', 'right', 'start', 'end', 'year', 'month', 'day') if name in relation] + relation.get('parts', [])
        constraints['relations'].extend(rewrite(relation, keys))
    for group in source_policy.get('constraints', {}).get('groups', []):
        constraints['groups'].extend(rewrite(group, group['fields']))
    result.update(fields=fields, constraints=constraints)
    result['repeat_rows'] = _summaries(fields, plan['count'], constraints, result.setdefault('warnings', []))
    for key in META:
        if key in source_policy:
            result[key] = deepcopy(source_policy[key])
    result['warnings'] = list(dict.fromkeys([*result.get('warnings', []), *source_policy.get('warnings', [])]))
    result['repeat_source_profile'] = source_policy
    result['repeat_rule_verification'] = {'original_file_checked': original_path is not None, 'prepared_file_checked': True}
    unresolved = [field['id'] for field in fields if field.get('repeat_info') and field['id'] not in assignments]
    if unresolved:
        result['repeat_unconfirmed_fields'] = unresolved
        result.setdefault('warnings', []).append(f'반복행 신규 입력칸 {len(unresolved)}개는 원본 확인 매핑에 없음; 사용자 매핑 확인 필요')
    outside_tokens = [field['id'] for field in fields if field['kind'] == 'xlsx_placeholder'
                      and not field.get('repeat_info') and field['id'] not in assignments]
    if outside_tokens:
        result.setdefault('warnings', []).append(f'반복행 바깥 필수 자리표시자 {len(outside_tokens)}개는 사용자 매핑 확인 필요')
    validate_rule_profile(result)
    return result
