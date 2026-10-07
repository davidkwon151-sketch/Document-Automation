"""Registered source metadata and signed-cache authority contracts; synthetic data only."""

from copy import deepcopy
from hashlib import sha256

from docx import Document
import pytest

from agent import template_learning
from templates.value_rules import mapped_rule_profile


def profile():
    fields = []
    for identifier, label, scope, kinds in [
        ('input-purpose', '임상시험 목적', '같은 원문 블록의 임상시험 목적 항목임', ['임상시험계획서', '임상시험 변경계획서']),
        ('input-introduction', '서론', '같은 원문 블록의 DSUR 서론 항목임', ['DSUR 원자료', '임상시험 안전성 보고 원자료']),
    ]:
        fields.append({'id': identifier, 'value_key': label, 'label': label,
                       'kind': 'docx_blank', 'required': False, 'input_required': False,
                       'input_mode': 'source_grounded', 'max_chars': 100, 'confidence': 1.0,
                       'evidence_scope': scope, 'evidence_document_labels': kinds})
    return {'format': 'docx', 'source_sha256': 'a' * 64, 'configured': True,
            'fields': fields, 'domain': 'pharmaceutical_ra', 'document_kind': 'application',
            'ra_workflow': 'clinical_trial'}


@pytest.mark.parametrize('key', ['evidence_scope', 'evidence_document_labels'])
@pytest.mark.parametrize('change', ['removed', 'none', 'empty', 'other-field', 'weakened'])
def test_declared_source_scope_cannot_disappear_relax_or_move_to_another_position(key, change):
    authority = profile()
    proposed = deepcopy(authority)
    field = proposed['fields'][0]
    if change == 'removed':
        field.pop(key)
    elif change == 'none':
        field[key] = None
    elif change == 'empty':
        field[key] = '' if key == 'evidence_scope' else []
    elif change == 'other-field':
        field[key] = deepcopy(proposed['fields'][1][key])
    else:
        field[key] = '제품 자료이면 허용함' if key == 'evidence_scope' else field[key] + ['시판 제품 설명서']
    before = deepcopy((proposed, authority))
    assert template_learning.reusable_profile(proposed, authority) is False
    assert (proposed, authority) == before


def test_scope_field_original_label_is_not_replaced_by_another_item():
    authority = profile()
    proposed = deepcopy(authority)
    proposed['fields'][0]['label'] = authority['fields'][1]['label']
    assert template_learning.reusable_profile(proposed, authority) is False


def test_declared_label_scope_array_order_is_bound_as_registered_metadata():
    authority = profile()
    proposed = deepcopy(authority)
    proposed['fields'][0]['evidence_document_labels'].reverse()
    assert not template_learning.reusable_profile(proposed, authority)


def test_original_position_id_binding_survives_reordered_fields():
    authority = profile()
    proposed = deepcopy(authority)
    proposed['fields'].reverse()
    assert template_learning.reusable_profile(proposed, authority)
    proposed['fields'][0]['evidence_scope'] = proposed['fields'][1]['evidence_scope']
    assert not template_learning.reusable_profile(proposed, authority)


def test_confirmed_value_key_rename_and_stronger_other_rules_remain_usable():
    authority = profile()
    proposed = mapped_rule_profile(authority, {'input-purpose': '임상 목적 새값', 'input-introduction': '서론 새값'})
    proposed['fields'][0].update(max_chars=50, required=True)
    proposed['citation_mode'] = 'sidecar'
    proposed['long_text_policy'] = 'annex'
    before = deepcopy((proposed, authority))
    assert template_learning.reusable_profile(proposed, authority)
    assert (proposed, authority) == before


def test_no_scope_authority_retains_existing_confirmation_and_new_constraints():
    authority = profile()
    for field in authority['fields']:
        field.pop('evidence_scope')
        field.pop('evidence_document_labels')
    proposed = deepcopy(authority)
    proposed['fields'][0]['value_key'] = '사용자 확정 항목'
    proposed['fields'][0]['evidence_scope'] = '사용자가 추가 확인한 좁은 범위'
    assert template_learning.reusable_profile(proposed, authority)


@pytest.fixture
def actual_cache(tmp_path):
    source = tmp_path / 'synthetic-template.docx'
    document = Document()
    document.add_paragraph('합성 입력1')
    document.add_paragraph('합성 입력2')
    document.save(source)
    authority = profile()
    authority['source_sha256'] = sha256(source.read_bytes()).hexdigest()
    return source, authority, tmp_path / 'private-cache'


def test_current_contract_roundtrip_renames_keys_and_strips_value_path_and_qa_payloads(actual_cache):
    source, authority, directory = actual_cache
    proposed = mapped_rule_profile(authority, {'input-purpose': '목적 수정', 'input-introduction': '서론 수정'})
    proposed.update(values={'성명': 'PRIVATE_SYNTHETIC_NAME'}, answers={'질문': 'PRIVATE_SYNTHETIC_ANSWER'},
                    sources=[{'filename': 'C:/PRIVATE_SYNTHETIC_PATH/evidence.pdf'}],
                    documents=[{'본문': 'PRIVATE_SYNTHETIC_DOCUMENT'}], qa={'secret': 'PRIVATE_SYNTHETIC_QA'})
    proposed['fields'][0].update(value='PRIVATE_SYNTHETIC_VALUE', default_value='PRIVATE_SYNTHETIC_DEFAULT',
                                qa_values={'x': 'PRIVATE_SYNTHETIC_FIELD_QA'})
    source_before = sha256(source.read_bytes()).hexdigest()
    saved = template_learning.save_learned_profile(source, proposed, cache_dir=directory, user_confirmed=True)
    cache_text = saved.read_text(encoding='utf-8')
    assert 'PRIVATE_SYNTHETIC_' not in cache_text
    loaded = template_learning.load_learned_profile(source, directory)
    assert template_learning.reusable_profile(loaded, authority)
    assert loaded['fields'][0]['value_key'] == '목적 수정'
    assert loaded['fields'][0]['evidence_scope'] == authority['fields'][0]['evidence_scope']
    assert loaded['fields'][0]['evidence_document_labels'] == authority['fields'][0]['evidence_document_labels']
    assert sha256(source.read_bytes()).hexdigest() == source_before


@pytest.mark.parametrize('change', ['scope-added', 'scope-changed', 'labels-changed', 'original-label-changed'])
def test_valid_hmac_legacy_cache_does_not_override_updated_registered_contract(actual_cache, change):
    source, authority, directory = actual_cache
    legacy = deepcopy(authority)
    if change == 'scope-added':
        legacy['fields'][0].pop('evidence_scope')
        legacy['fields'][0].pop('evidence_document_labels')
    template_learning.save_learned_profile(source, legacy, cache_dir=directory, user_confirmed=True)
    loaded = template_learning.load_learned_profile(source, directory)
    updated = deepcopy(authority)
    if change == 'scope-changed':
        updated['fields'][0]['evidence_scope'] += ' 및 특정 원자료의 확인 범위임'
    elif change == 'labels-changed':
        updated['fields'][0]['evidence_document_labels'] = ['개정된 임상시험 원자료']
    elif change == 'original-label-changed':
        updated['fields'][0]['label'] = '임상시험의 새 목적'
    assert loaded['learning']['origin'] == 'confirmed_cache'
    assert template_learning.reusable_profile(loaded, updated) is False


def test_optional_unused_position_can_stay_unmapped():
    authority = profile()
    proposed = deepcopy(authority)
    proposed['fields'] = proposed['fields'][:1]
    assert template_learning.reusable_profile(proposed, authority)
