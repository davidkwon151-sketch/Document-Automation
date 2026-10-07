"""기존 사내 양식의 입력 위치를 분석하고 명시적으로 매핑해 채움."""

from io import BytesIO
from hashlib import sha256
from copy import deepcopy
from decimal import Decimal
import os
from pathlib import Path
import re
import tempfile
from zipfile import ZipFile

from lxml import etree
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries
from pypdf import PdfReader, PdfWriter
from pypdf.generic import (ArrayObject, DecodedStreamObject, DictionaryObject, FloatObject, IndirectObject,
                          NameObject, NumberObject, TextStringObject, ContentStream)
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
import pdfplumber

from parsers.extract import _check_zip, _path, _xml
from .fill import PLACEHOLDER, WORD_NS, XML_SPACE, TemplateError, fill_template, _segments, _set_segment
from .value_rules import inspect_form_values, mapped_rule_profile
from .pdf_choices import choice_metadata, parse_choice_value


SHEET_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
CHECKBOX_NS = "http://schemas.microsoft.com/office/word/2010/wordml"
SUGGESTIONS = {"제목": "제목", "건명": "제목", "보고제목": "제목", "문서제목": "제목",
               "품의제목": "제목", "title": "제목", "subject": "제목", "요약": "요약",
               "핵심요약": "요약", "개요": "요약", "summary": "요약", "본문": "본문",
               "내용": "본문", "보고내용": "본문", "주요내용": "본문", "추진내용": "본문",
               "body": "본문", "content": "본문"}


def _suggest(label):
    normalized = re.sub(r"[\s:：*()（）]", "", label).lower()
    return SUGGESTIONS.get(normalized)


def _field(identifier, label, kind, *, required=False, confidence=None):
    suggestion = _suggest(label)
    return {"id": identifier, "label": label, "kind": kind,
            "value_key": suggestion or label, "suggested_mapping": suggestion,
            "required": required, "input_required": label in {"이름", "성명", "소속", "제출처", "연락처", "작성자", "부서"},
            "confidence": confidence if confidence is not None else (0.9 if suggestion else 0.35)}


def _text(element):
    return "".join(element.xpath(".//*[local-name()='t']/text()"))


def _parts(archive, kind):
    if kind == "docx":
        if "word/document.xml" not in archive.namelist():
            raise TemplateError("DOCX 본문이 없음")
        return [name for name in archive.namelist() if name.startswith("word/") and name.endswith(".xml")]
    if archive.read("mimetype").strip() != b"application/hwp+zip":
        raise TemplateError("HWPX mimetype이 올바르지 않음")
    return [name for name in archive.namelist()
            if re.fullmatch(r"Contents/(section|masterpage)\d+\.xml", name)]


def _blank_marker(text):
    """Explicit blank text only; ordinary punctuation/content is never replaced."""
    return bool(text) and (not text.strip() or bool(re.fullmatch(r"[\s_\.\u2026\u00af\u2014\u2013]{4,}", text)))


def _legacy_text_region(begin):
    """Locate an enabled, non-nested legacy FORMTEXT result within its paragraph."""
    namespace = {"w": WORD_NS}
    if begin.tag != f"{{{WORD_NS}}}fldChar" or begin.get(f"{{{WORD_NS}}}fldCharType") != "begin" or begin.xpath("ancestor::w:del | ancestor::w:sdt[w:sdtPr/w:lock]", namespaces=namespace):
        return None
    data = begin.find(f"{{{WORD_NS}}}ffData")
    if data is None or data.find(f"{{{WORD_NS}}}textInput") is None:
        return None
    enabled = data.find(f"{{{WORD_NS}}}enabled")
    if (enabled is not None and enabled.get(f"{{{WORD_NS}}}val") in {"0", "false", "off"}) or begin.get(f"{{{WORD_NS}}}fldLock") in {"1", "true", "on"}:
        return None
    paragraphs = begin.xpath("ancestor::w:p[1]", namespaces=namespace)
    if not paragraphs:
        return None
    items = list(paragraphs[0].iter())
    instruction, nodes, separated = [], [], False
    for item in items[items.index(begin) + 1:]:
        if item.tag == f"{{{WORD_NS}}}fldChar":
            state = item.get(f"{{{WORD_NS}}}fldCharType")
            if state == "begin":
                return None
            if state == "separate":
                separated = True
            elif state == "end":
                return nodes if nodes and "".join(instruction).strip() == "FORMTEXT" else None
        elif item.tag == f"{{{WORD_NS}}}instrText" and not separated:
            instruction.append(item.text or "")
        elif item.tag == f"{{{WORD_NS}}}t" and separated:
            nodes.append(item)
        elif separated and item.tag in {f"{{{WORD_NS}}}tab", f"{{{WORD_NS}}}br", f"{{{WORD_NS}}}drawing"}:
            return None
    return None


_LEGACY_PRIVATE = re.compile(
    r"성명|이름|작성자|담당자|대표자|소속|부서|주소|전화|연락|메일|생년|주민|서명|직인|도장|인감|계좌|비밀번호|투표|동의|위임|승인|결재|"
    r"(?<![A-Za-z0-9])(?:name|(?:first|last|full|given|family)[ _-]*name|author|prepared[ _-]*by|responsible|department|applicant|signature|signed[ _-]*by|"
    r"address|e-?mail|phone|telephone|fax|consent|vote|account|password|ssn|tax[ _-]*(?:id|identification)|approval|approver|approved[ _-]*by|authorized[ _-]*by|"
    r"contact[ _-]*person|bearbeiter|abteilung|telefon|telefax|unterschrift)(?![A-Za-z0-9])", re.I)
_LEGACY_CONTENT = re.compile(
    r"^(?:(?:specific|product|technical|business|project|report|work|item|material)[ /-]+){0,3}"
    r"(?:description|specifications?|characteristics|content|body|details)(?:\b|$)|"
    r"^(?:(?:supplier|siemens|product|item)[ /-]+)?material(?:[ /-]+(?:no\.?|number|code|type|composition))?[ :：]*$|"
    r"^(?:제품|상품|재료|재질|원재료|사업|프로젝트|업무|보고)?\s*(?:설명|내용|본문|사양|규격|특성|세부내용|기술사양)(?:\s*[:：]|$)|"
    r"^(?:(?:Spezifische\s+)?Produktmerkmale|Produktbeschreibung|Spezifikation|Beschreibung)(?:\b|$)|"
    r"^(?:Lieferanten-|Siemens-)?Material-(?:Nr\.?|Nummer)[ :：]*$", re.I)


def _legacy_neighbor_label(begin, paragraph):
    """Read observed labels in the same cell or immediately preceding cell only."""
    namespace = {"w": WORD_NS}
    cells = begin.xpath("ancestor::w:tc[1]", namespaces=namespace)
    if not cells:
        return ""
    cell = cells[0]
    paragraphs = cell.xpath("./w:p", namespaces=namespace)
    if paragraph in paragraphs:
        labels = []
        for previous in reversed(paragraphs[:paragraphs.index(paragraph)]):
            if previous.xpath(".//w:fldChar | .//w:sdt", namespaces=namespace):
                break
            text = _text(previous).strip()
            if text and not _blank_marker(text):
                labels.insert(0, text)
            if len(labels) == 2:
                break
        label = " ".join(labels).strip()
        if label:
            return label
    previous = cell.getprevious()
    if previous is not None and previous.tag == f"{{{WORD_NS}}}tc" and not previous.xpath(".//w:fldChar | .//w:sdt", namespaces=namespace):
        return _text(previous).strip()
    return ""


def _legacy_text_fields(root, part):
    fields, tree = [], root.getroottree()
    for begin in root.xpath(".//w:fldChar[@w:fldCharType='begin']", namespaces={"w": WORD_NS}):
        nodes = _legacy_text_region(begin)
        if not nodes or not _blank_marker("".join(node.text or "" for node in nodes)):
            continue
        paragraph = begin.xpath("ancestor::w:p[1]", namespaces={"w": WORD_NS})[0]
        prefix = []
        for item in paragraph.iter():
            if item is begin:
                break
            if item.tag == f"{{{WORD_NS}}}t":
                prefix.append(item.text or "")
        name = begin.find(f"{{{WORD_NS}}}ffData/{{{WORD_NS}}}name")
        name = name.get(f"{{{WORD_NS}}}val", "") if name is not None else ""
        observed_label = "".join(prefix).strip() or _legacy_neighbor_label(begin, paragraph)
        # A recognizable business-content label allows grounding; an unknown TextN,
        # personal/approval label or instruction-like sentence stays direct input.
        grounded = bool(observed_label and len(observed_label) <= 220 and
                        not _LEGACY_PRIVATE.search(observed_label) and _LEGACY_CONTENT.search(observed_label))
        label = observed_label.rstrip(":：") or f"{name or 'FORMTEXT'} (입력칸 {len(fields) + 1})"
        field = _field(f"docx:{part}:{tree.getpath(begin)}", label, "docx_legacy_text", confidence=1.0)
        field.update({"input_required": not grounded, "input_mode": "source_grounded" if grounded else "user_provided",
                      "narrative_style_required": False, "form_field_name": name,
                      "anchor_text": "".join(node.text or "" for node in nodes)})
        maximum = begin.find(f"{{{WORD_NS}}}ffData/{{{WORD_NS}}}textInput/{{{WORD_NS}}}maxLength")
        if maximum is not None and maximum.get(f"{{{WORD_NS}}}val", "").isdigit() and int(maximum.get(f"{{{WORD_NS}}}val")) > 0:
            field["max_chars"] = int(maximum.get(f"{{{WORD_NS}}}val"))
        fields.append(field)
    return fields


def _office_fields(archive, kind, warnings):
    fields, placeholders = [], set()
    forms_only = False
    if kind == "docx" and "word/settings.xml" in archive.namelist():
        settings = _xml(archive.read("word/settings.xml"))
        protected = settings.xpath(".//w:documentProtection[@w:enforcement='1' or @w:enforcement='true' or @w:enforcement='on']",
                                   namespaces={"w": WORD_NS})
        if protected:
            forms_only = all(item.get(f"{{{WORD_NS}}}edit") == "forms" for item in protected)
            if not forms_only:
                warnings.append("편집 보호가 설정된 DOCX임. 보호 해제한 사본을 사용해야 함")
                return []
            warnings.append("양식 보호를 유지하며 활성화된 빈 FORMTEXT 결과만 입력함. 일반 본문·선택상자는 변경하지 않음")
    for part in _parts(archive, kind):
        root = _xml(archive.read(part))
        tree = root.getroottree()
        if kind == "docx":
            fields.extend(_legacy_text_fields(root, part))
            if forms_only:
                continue
        for paragraph in root.xpath(".//*[local-name()='p']"):
            # Nested 표/글상자 문단은 따로 처리해 placeholder를 중복 추출하지 않음.
            from .fill import _segments
            candidate = paragraph
            if kind == "docx":
                excluded = "w:sdt[w:sdtPr/w:lock or w:sdtPr/w:dropDownList or w:sdtPr/w:comboBox or w:sdtPr/w:date or w:sdtPr/w:picture or w:sdtPr/w:group or w:sdtPr/*[local-name()='repeatingSection' or local-name()='repeatingSectionItem']]"
                if paragraph.xpath("ancestor::" + excluded, namespaces={"w": WORD_NS}):
                    continue
                candidate = deepcopy(paragraph)
                for control in candidate.xpath(".//" + excluded, namespaces={"w": WORD_NS}):
                    control.getparent().remove(control)
            paragraph_text = "".join(item.value for item in _segments(candidate, kind))
            for match in PLACEHOLDER.finditer(paragraph_text):
                key = match[1].strip()
                if key not in placeholders:
                    fields.append(_field("placeholder:" + key, key, "placeholder", required=True, confidence=1.0))
                    placeholders.add(key)
        for table in root.xpath(".//*[local-name()='tbl']"):
            rows = table.xpath("./*[local-name()='tr']")
            previous = []
            for row in rows:
                cells = row.xpath("./*[local-name()='tc']")
                for index, cell in enumerate(cells):
                    if kind == "docx" and cell.xpath(".//w:sdt", namespaces={"w": WORD_NS}):
                        continue
                    if _text(cell).strip() and not (kind == "docx" and _blank_marker(_text(cell))):
                        continue
                    if kind == "docx" and cell.xpath(".//w:fldChar", namespaces={"w": WORD_NS}):
                        continue
                    if kind == "docx" and cell.xpath(".//w:vMerge[not(@w:val) or @w:val='continue']", namespaces={"w": WORD_NS}):
                        continue
                    if kind == "hwpx" and cell.get("protect") == "1":
                        continue
                    label = _text(cells[index - 1]).strip() if index else ""
                    if not label and index < len(previous):
                        label = _text(previous[index]).strip()
                    if not PLACEHOLDER.search(label):
                        identifier = f"{kind}:{part}:{tree.getpath(cell)}"
                        fields.append(_field(identifier, label or f"빈 표 셀 {len(fields) + 1}", f"{kind}_cell"))
                previous = cells
        if kind == "docx":
            for paragraph in root.xpath(".//w:p", namespaces={"w": WORD_NS}):
                if paragraph.xpath(".//w:fldChar | .//w:sdt", namespaces={"w": WORD_NS}):
                    continue
                segments = _segments(paragraph, kind)
                text = "".join(segment.value for segment in segments)
                match = re.fullmatch(r"([^:\n：]{1,80}[:：])([ \u00a0\u2002\u2003_\.\u2026\u00af]{3,})", text)
                if match and _blank_marker(match[2]) and all(segment.editable for segment in segments):
                    field = _field(f"docx:{part}:{tree.getpath(paragraph)}", match[1].rstrip(":： "), "docx_inline")
                    field.update({"anchor_text": text, "blank_start": match.start(2), "blank_end": match.end(2)})
                    fields.append(field)
            for control in root.xpath(".//w:sdt", namespaces={"w": WORD_NS}):
                properties = control.find(f"{{{WORD_NS}}}sdtPr")
                content = control.find(f"{{{WORD_NS}}}sdtContent")
                if properties is None or content is None:
                    continue
                alias = properties.find(f"{{{WORD_NS}}}alias")
                tag = properties.find(f"{{{WORD_NS}}}tag")
                label = (alias if alias is not None else tag)
                label = label.get(f"{{{WORD_NS}}}val", "") if label is not None else ""
                if not label:
                    continue
                if properties.find(f"{{{WORD_NS}}}lock") is not None or control.xpath("ancestor::w:sdt/w:sdtPr/w:lock", namespaces={"w": WORD_NS}):
                    warnings.append(f"잠긴 콘텐츠 컨트롤은 자동 입력에서 제외함: {label}")
                    continue
                checkbox = properties.find(f"{{{CHECKBOX_NS}}}checkbox")
                if checkbox is not None:
                    field = _field(f"{kind}:{part}:{tree.getpath(content)}", label, "docx_checkbox")
                    field.update({"control_type": "checkbox", "options": ["true", "false"], "input_required": True})
                    fields.append(field)
                    continue
                if properties.xpath("./w:date | ./w:picture | ./w:group | ./*[local-name()='repeatingSection' or local-name()='repeatingSectionItem']", namespaces={"w": WORD_NS}):
                    warnings.append(f"날짜·그림·그룹·반복 콘텐츠 컨트롤은 자동 입력에서 제외함: {label}")
                    continue
                choices = properties.xpath("./w:dropDownList | ./w:comboBox", namespaces={"w": WORD_NS})
                if choices:
                    from .docx_choices import native_choice
                    try:
                        metadata = native_choice(properties, content)
                    except TemplateError as exc:
                        warnings.append(f"선택 컨트롤을 자동 입력에서 제외함: {label}: {exc}")
                        continue
                    if _text(content).strip() and properties.find(f"{{{WORD_NS}}}showingPlcHdr") is None:
                        warnings.append(f"기존 내용이 있는 선택 컨트롤은 자동 입력에서 제외함: {label}")
                        continue
                    field = _field(f"{kind}:{part}:{tree.getpath(content)}", label, metadata.pop("kind"))
                    field.update(metadata)
                    fields.append(field)
                    continue
                if PLACEHOLDER.search(_text(content)):
                    continue
                if _text(content).strip() and properties.find(f"{{{WORD_NS}}}showingPlcHdr") is None:
                    warnings.append(f"기존 내용이 있는 콘텐츠 컨트롤은 자동 입력에서 제외함: {label}")
                    continue
                fields.append(_field(f"{kind}:{part}:{tree.getpath(content)}", label, "docx_sdt"))
        for paragraph in root.xpath(".//*[local-name()='p']"):
            ancestors = paragraph.xpath("ancestor::*[local-name()='tc' or local-name()='sdtContent' or local-name()='subList']")
            if ancestors or _text(paragraph).strip():
                continue
            if kind == "docx" and paragraph.xpath(".//w:fldChar", namespaces={"w": WORD_NS}):
                continue
            # 그림/표가 있는 빈 anchor 문단을 실제 입력칸으로 오인하지 않음.
            if paragraph.xpath(".//*[local-name()='tbl' or local-name()='drawing' or local-name()='pic' or local-name()='sdt']"):
                continue
            fields.append(_field(f"{kind}:{part}:{tree.getpath(paragraph)}",
                                 f"빈 본문 문단 {len(fields) + 1}", f"{kind}_paragraph", confidence=0.1))
    return fields


def _xlsx_cell_text(cell, shared):
    if cell is None:
        return ""
    if cell.get("t") == "s":
        value = cell.find(f"{{{SHEET_NS}}}v")
        return shared[int(value.text)] if value is not None and value.text else ""
    if cell.get("t") == "inlineStr":
        return "".join(cell.xpath(".//*[local-name()='t']/text()"))
    value = cell.find(f"{{{SHEET_NS}}}v")
    return value.text or "" if value is not None else ""


def _xlsx_parts(archive):
    shared = []
    if "xl/sharedStrings.xml" in archive.namelist():
        shared = ["".join(item.itertext()) for item in _xml(archive.read("xl/sharedStrings.xml"))]
    parts = [name for name in archive.namelist() if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", name)]
    if not parts:
        raise TemplateError("XLSX 시트가 없음")
    return shared, parts


def _xlsx_fields(archive, warnings):
    from .xlsx_validation import XlsxLists, editable_blank
    fields = []
    shared, parts = _xlsx_parts(archive)
    lists = XlsxLists(archive)
    for part in parts:
        root = _xml(archive.read(part))
        if root.find(f"{{{SHEET_NS}}}sheetProtection") is not None:
            warnings.append(f"편집 보호된 시트는 입력에서 제외함: {part}")
            continue
        cells = {cell.get("r"): cell for cell in root.findall(f".//{{{SHEET_NS}}}sheetData/{{{SHEET_NS}}}row/{{{SHEET_NS}}}c")}
        merged = [range_boundaries(item.get("ref")) for item in root.findall(f".//{{{SHEET_NS}}}mergeCell")]
        coordinates = set(cells) | lists.candidates(part, warnings)
        for coordinate in sorted(coordinates, key=lambda item: (range_boundaries(item)[1], range_boundaries(item)[0])):
            cell = cells.get(coordinate)
            if cell is not None and cell.find(f"{{{SHEET_NS}}}f") is not None:
                continue  # 수식은 사용자 문구로 덮어쓰지 않음.
            text = _xlsx_cell_text(cell, shared)
            native_list = lists.at(part, coordinate)
            if native_list is not None:
                if not editable_blank(root, coordinate, lists.shared):
                    continue
                letters, row = re.fullmatch(r"([A-Z]+)(\d+)", coordinate).groups()
                column, row = column_index_from_string(letters), int(row)
                left = f"{get_column_letter(column - 1)}{row}" if column > 1 else ""
                above = f"{letters}{row - 1}" if row > 1 else ""
                label = _xlsx_cell_text(cells.get(left), shared).strip() or _xlsx_cell_text(cells.get(above), shared).strip()
                item = _field(f"xlsx:{part}:{coordinate}", label or f"목록 선택 {coordinate}", "xlsx_cell")
                options = native_list["options"] if not native_list["error"] else []
                item.update({"control_type": "choice_unresolved" if native_list["error"] else "choice",
                             "options": options, "choice_items": [{"value": value, "label": value} for value in options],
                             "input_required": True, "narrative_style_required": False, "xlsx_list": lists.metadata(native_list)})
                if native_list["error"]:
                    item["unsupported_reason"] = native_list["error"]
                fields.append(item)
                continue
            matches = list(PLACEHOLDER.finditer(text))
            for match in matches:
                key = match[1].strip()
                fields.append(_field(f"xlsx:{part}:{coordinate}:{key}", key, "xlsx_placeholder", required=True, confidence=1.0))
            if text.strip() or matches:
                continue
            letters, row = re.fullmatch(r"([A-Z]+)(\d+)", coordinate).groups()
            column, row = column_index_from_string(letters), int(row)
            if any(x1 <= column <= x2 and y1 <= row <= y2 and (column, row) != (x1, y1)
                   for x1, y1, x2, y2 in merged):
                continue
            left = f"{get_column_letter(column - 1)}{row}" if column > 1 else ""
            above = f"{letters}{row - 1}" if row > 1 else ""
            label = _xlsx_cell_text(cells.get(left), shared).strip() or _xlsx_cell_text(cells.get(above), shared).strip()
            if not PLACEHOLDER.search(label):
                fields.append(_field(f"xlsx:{part}:{coordinate}", label or f"빈 셀 {coordinate}", "xlsx_cell"))
    for item in fields:
        # Keep the original multi-line label; only normalize this new suggestion.
        item["value_key"] = " ".join(item["value_key"].split())
    return fields


def _pdf_profile(source, warnings):
    reader = PdfReader(source)
    if reader.is_encrypted:
        warnings.append("암호화된 PDF는 지원하지 않음. 잠금 해제한 사본을 사용해야 함")
        return [], "encrypted", []
    if _pdf_signed(reader):
        warnings.append("전자서명 또는 인증된 PDF임. 서명을 무효화하지 않도록 입력을 차단함")
        return [], "signed", _pdf_pages(reader)
    acroform = reader.trailer["/Root"].get("/AcroForm")
    if acroform and "/XFA" in acroform.get_object():
        warnings.append("XFA 양식은 원본을 유지하고 페이지 위에 입력함. XFA 우선 렌더링을 하는 뷰어의 결과는 별도 확인해야 함")
        return _pdf_overlay_candidates(source), "overlay", _pdf_pages(reader)
    fields = []
    for key, field in (reader.get_fields() or {}).items():
        flags = int(_pdf_inherited(field, '/Ff', 0))
        field_type = _pdf_inherited(field, '/FT', '')
        if field_type not in {"/Tx", "/Btn", "/Ch"} or flags & 1 or (field_type == "/Btn" and flags & 65536):
            continue
        label = str(field.get("/TU") or field.get("/TM") or key)
        item = _field("pdf:" + key, label, "pdf_form", required=bool(flags & 2))
        if field_type == "/Btn":
            options = {str(option) for option in field.get("_States_", []) if str(option) != "/Off"}
            widgets = list(field.get("/Kids", [])) + [reference for page in reader.pages
                       for reference in page.get("/Annots", []) if _pdf_field_name(reference.get_object()) == key]
            for widget in widgets:
                normal = widget.get_object().get("/AP", {}).get("/N", {})
                if isinstance(normal, DictionaryObject):
                    options.update(str(option) for option in normal if str(option) != "/Off")
            item.update({"control_type": "radio" if flags & 32768 else "checkbox",
                         "options": sorted(options), "input_required": True})
        elif field_type == "/Ch":
            try:
                item.update(choice_metadata(_pdf_inherited(field, '/Opt', []), flags))
            except TemplateError as exc:
                item.update(control_type='choice_unresolved', unsupported_reason=str(exc))
                warnings.append(f"PDF 선택 목록을 확인할 수 없어 입력을 보류함: {label}: {exc}")
            item.update(input_required=True, input_mode='user_provided', narrative_style_required=False)
        fields.append(item)
    if not fields:
        warnings.append("정적·스캔 PDF의 원본 페이지 위에 입력함. 자동 제안 영역을 확인하거나 입력 영역을 지정해야 함")
        return _pdf_overlay_candidates(source), "overlay", _pdf_pages(reader)
    else:
        warnings.append("PDF 입력칸의 배경·테두리·크기를 유지함. 입력 문자를 표시할 수 있는 한글 TTF 글꼴을 내장하며 부족한 문자는 동봉 글꼴을 사용해 글꼴이 달라질 수 있음. 장문 넘침은 자동 검사함")
    return fields, "acroform", _pdf_pages(reader)


def _pdf_signed(reader):
    return bool(reader.trailer["/Root"].get("/Perms")) or any(
        field.get("/FT") == "/Sig" and field.get("/V") for field in (reader.get_fields() or {}).values())


def _pdf_pages(reader):
    pages = []
    for number, page in enumerate(reader.pages, 1):
        width, height = float(page.mediabox.width), float(page.mediabox.height)
        if page.rotation % 180:
            width, height = height, width
        pages.append({"page": number, "width": width, "height": height})
    return pages


def _pdf_overlay_candidates(source):
    fields = []
    with pdfplumber.open(source) as document:
        for page_number, page in enumerate(document.pages, 1):
            for word in page.extract_words():
                label = word["text"].rstrip(":：")
                if not _suggest(label):
                    continue
                width = float(page.width) - float(word["x1"]) - 35
                if width < 40:
                    continue
                field = _field(f"pdf-overlay:{page_number}:{len(fields)}", label, "pdf_overlay", confidence=0.55)
                field.update({"page": page_number, "x": float(word["x1"]) + 8,
                              "y": float(word["top"]), "width": width,
                              "height": min(60.0 if _suggest(label) == "본문" else 25.0, float(page.height) - float(word["top"]) - 20),
                              "font_size": 10.0})
                fields.append(field)
    return fields


def _manual_pdf_fields(manual_fields, pages):
    fields = []
    sizes = {page["page"]: page for page in pages}
    for index, item in enumerate(manual_fields):
        label = str(item.get("label", item.get("value_key", f"입력 영역 {index + 1}")))
        field = _field(str(item.get("id", f"pdf-manual:{index}")), label, "pdf_overlay", confidence=1.0)
        field.update(item)
        field["kind"] = "pdf_overlay"
        field["value_key"] = str(item.get("value_key", field["value_key"]))
        page = int(item["page"])
        x, y, width, height = (float(item[key]) for key in ("x", "y", "width", "height"))
        size = sizes.get(page)
        if size is None or x < 0 or y < 0 or width <= 0 or height <= 0 or x + width > size["width"] or y + height > size["height"]:
            raise TemplateError("PDF 입력 영역은 해당 페이지 안에 있어야 함")
        field.update({"page": page, "x": x, "y": y, "width": width, "height": height,
                      "font_size": float(item.get("font_size", 10))})
        fields.append(field)
    return fields


def analyze_template(path: str | Path, manual_fields: list[dict] | None = None) -> dict:
    """자동 인식 가능한 칸과 수동 매핑이 필요한 칸을 동일 프로파일로 반환함."""
    source = Path(path)
    kind = source.suffix.lower().lstrip(".")
    profile = {"format": kind, "fields": [], "warnings": [], "supported": False}
    if kind not in {"docx", "hwpx", "xlsx", "pdf", "pptx"}:
        profile["warnings"].append("지원 형식은 DOCX/HWPX/XLSX/PPTX/입력칸 PDF임. HWP→HWPX, XLS→XLSX, PPT→PPTX 변환 후 사용할 수 있음")
        return profile
    try:
        source = _path(source)
        profile["source_sha256"] = sha256(source.read_bytes()).hexdigest()
        if kind == "pdf":
            fields, render_mode, pages = _pdf_profile(source, profile["warnings"])
            profile.update({"render_mode": render_mode, "pages": pages,
                            "coordinate_system": "1쪽부터, 페이지 좌상단 원점, 포인트(72pt/in)"})
            if manual_fields:
                fields = _manual_pdf_fields(manual_fields, pages)
                profile["render_mode"] = "overlay"
        else:
            with ZipFile(source) as archive:
                _check_zip(archive)
                if kind == "pptx":
                    from .pptx import pptx_fields
                    fields = pptx_fields(archive, profile["warnings"])
                else:
                    fields = _xlsx_fields(archive, profile["warnings"]) if kind == "xlsx" else _office_fields(archive, kind, profile["warnings"])
        # 동일 칸이 content control·빈 셀로 중복 잡히면 더 명시적인 control을 우선함.
        profile["fields"] = list({field["id"]: field for field in fields}.values())
        profile["supported"] = bool(profile["fields"]) or (kind == "pdf" and profile.get("render_mode") == "overlay")
        if kind in {"docx", "hwpx", "xlsx"}:
            profile["warnings"].append("인접 항목명으로 제안한 매핑은 확인이 필요함. 입력 길이에 따른 페이지·행 높이·셀 넘침은 원본 프로그램에서 확인해야 함")
        if not fields and kind != "pdf":
            profile["warnings"].append("자리표시자 또는 항목명이 있는 빈 입력칸을 찾지 못함. 입력칸을 추가하거나 {{제목}}·{{요약}}·{{본문}}을 넣어야 함")
    except Exception as exc:
        profile["warnings"].append(f"양식을 분석할 수 없음: {exc}")
    return profile


def _selected_fields(profile, values, mapping):
    available = {field["id"]: field for field in profile["fields"]}
    if mapping is not None and any(key not in available for key in mapping):
        raise TemplateError("양식에 없는 입력칸 ID가 매핑에 포함됨")
    selected = []
    for field in profile["fields"]:
        key = mapping.get(field["id"]) if mapping is not None else field["value_key"]
        if key is None or key == "":
            if field["required"]:
                raise TemplateError(f"필수 입력칸 매핑이 누락됨: {field['label']}")
            continue
        if key not in values or not values[key].strip():
            if field["required"] or (mapping is not None and key not in values):
                raise TemplateError(f"매핑된 값이 누락됨: {field['label']} → {key}")
            continue
        value = values[key]
        control = field.get("control_type")
        if control == 'choice_unresolved':
            raise TemplateError(f"원본 선택 목록을 확인할 수 없어 기입을 보류함: {field['label']}")
        pdf_choice = field.get('kind') == 'pdf_form' and control in {'choice', 'combobox'}
        if pdf_choice:
            from .value_rules import plain_form_value
            literal = value if profile.get('citation_mode') == 'sidecar' else None
            value = plain_form_value(value, field, literal=literal)
            parse_choice_value(field, value, literal=literal)
        elif control in {'checkbox', 'radio', 'choice', 'combobox'}:
            from .value_rules import plain_form_value
            value = plain_form_value(value, field)
        if field.get("max_chars") and len(value) > field["max_chars"]:
            raise TemplateError(f"입력 길이가 제한을 넘음: {field['label']} ({field['max_chars']}자)")
        if control == "checkbox":
            if field["kind"] == "docx_checkbox":
                if value.lower() not in {"true", "false"}:
                    raise TemplateError("체크박스 값은 true 또는 false여야 함")
                value = value.lower()
            elif value.lower() in {"true", "false"}:
                if value.lower() == "true" and len(field.get("options", [])) != 1:
                    raise TemplateError("체크박스의 선택값을 명시해야 함")
                value = field["options"][0] if value.lower() == "true" else "/Off"
            elif value not in {"/Off", *field.get("options", [])}:
                raise TemplateError("지원하지 않는 체크박스 선택값임")
        elif not pdf_choice and control in {"radio", "choice"} and value not in field.get("options", []):
            raise TemplateError(f"목록에 없는 선택값임: {field['label']}")
        selected.append((field, value))
    if not selected:
        raise TemplateError("채울 수 있는 값과 입력칸의 매핑이 없음")
    return selected


def _rewrite_zip(source, target, changes):
    target.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary = tempfile.mkstemp(dir=target.parent, suffix=target.suffix)
    os.close(descriptor)
    try:
        with ZipFile(source) as archive, ZipFile(temporary, "w") as result:
            result.comment = archive.comment
            for item in archive.infolist():
                replacement = changes.get(item.filename, archive.read(item.filename))
                if replacement is not None:
                    result.writestr(item, replacement)
        os.replace(temporary, target)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


def _office_fill(source, target, selected):
    roots, substitutions, shielded = {}, {}, {}
    repeat_kinds = {'docx_repeat_placeholder', 'hwpx_repeat_placeholder'}
    # Higher offsets first keep every earlier source offset stable, even when
    # several placeholders share a run or have different replacement lengths.
    selected = sorted(selected, key=lambda item: (
        0 if item[0]['kind'] in repeat_kinds else 1,
        item[0]['id'].split(':', 2)[1] if item[0]['kind'] in repeat_kinds else '',
        item[0].get('xml_path', ''), -item[0].get('placeholder_start', 0)))
    with ZipFile(source) as archive:
        reserved_tokens = {field['label'] for field, _ in selected if field['kind'] == 'placeholder'}
        for name in _parts(archive, source.suffix.lower().lstrip('.')):
            original = _xml(archive.read(name))
            for paragraph in original.xpath(".//*[local-name()='p']"):
                text = ''.join(segment.value for segment in _segments(paragraph, source.suffix.lower().lstrip('.')))
                reserved_tokens.update(match[1].strip() for match in PLACEHOLDER.finditer(text))
        def mapped_token(index):
            token = f'__mapped_{index}'
            while token in reserved_tokens:
                token += '_'
            reserved_tokens.add(token)
            return token
        if source.suffix.lower() == ".docx" and "word/settings.xml" in archive.namelist():
            settings = _xml(archive.read("word/settings.xml"))
            protection = settings.xpath(".//w:documentProtection[@w:enforcement='1' or @w:enforcement='true' or @w:enforcement='on']", namespaces={"w": WORD_NS})
            if protection and (any(item.get(f"{{{WORD_NS}}}edit") != "forms" for item in protection) or any(field["kind"] != "docx_legacy_text" for field, _ in selected)):
                raise TemplateError("DOCX 보호 범위 밖의 입력은 허용되지 않음")
        for index, (field, value) in enumerate(selected):
            if field["kind"] == "placeholder":
                substitutions[field["label"]] = value
                continue
            kind, part, xpath = field["id"].split(":", 2)
            if field['kind'] in repeat_kinds:
                kind = field['kind'].split('_', 1)[0]
                xpath = field.get('xml_path')
                if (kind != source.suffix.lower().lstrip('.') or not isinstance(xpath, str)
                        or field['id'] != f"repeat_{kind}:{part}:{xpath}#slot:{field.get('placeholder_start')}"):
                    raise TemplateError('반복 자리표시자 위치 기록이 올바르지 않음')
            root = roots.setdefault(part, _xml(archive.read(part)))
            namespace_map = {key: val for key, val in root.nsmap.items() if key}
            targets = root.xpath(xpath, namespaces=namespace_map)
            if len(targets) != 1:
                raise TemplateError("입력칸 위치가 변경되어 양식을 다시 분석해야 함")
            cell = targets[0]
            if cell.tag == f"{{{WORD_NS}}}sdtContent" and cell.getparent().find(f"{{{WORD_NS}}}sdtPr") is not None:
                properties = cell.getparent().find(f"{{{WORD_NS}}}sdtPr")
                native = properties.xpath("./w:dropDownList | ./w:comboBox", namespaces={"w": WORD_NS})
                if native or field["kind"] in {"docx_choice", "docx_combobox"}:
                    from .docx_choices import fill_native_choice
                    fill_native_choice(properties, cell, field, value)
                    continue
                if properties.xpath("./w:lock | ./w:dataBinding | ./w:date | ./w:picture | ./w:group | ./*[local-name()='repeatingSection' or local-name()='repeatingSectionItem']", namespaces={"w": WORD_NS}):
                    raise TemplateError("잠김·연결·미지원 콘텐츠 컨트롤은 변경할 수 없음")
            if field["kind"] == "docx_legacy_text":
                nodes = _legacy_text_region(cell)
                if not nodes or "".join(node.text or "" for node in nodes) != field.get("anchor_text") or not _blank_marker(field.get("anchor_text", "")):
                    raise TemplateError("FORMTEXT 입력 위치 또는 원문이 변경됨")
                if "\n" in value or "\t" in value:
                    raise TemplateError("FORMTEXT 입력값은 한 줄이어야 함")
                nodes[0].text = value
                nodes[0].set(XML_SPACE, "preserve")
                for node in nodes[1:]:
                    node.text = ""
                continue
            if field["kind"] == "docx_checkbox":
                checkbox = cell.getparent().find(f"{{{WORD_NS}}}sdtPr/{{{CHECKBOX_NS}}}checkbox")
                if checkbox is None:
                    raise TemplateError("DOCX 체크박스 구조가 변경됨")
                checked = checkbox.find(f"{{{CHECKBOX_NS}}}checked")
                if checked is None:
                    checked = etree.SubElement(checkbox, f"{{{CHECKBOX_NS}}}checked")
                checked.set(f"{{{CHECKBOX_NS}}}val", "1" if value == "true" else "0")
                state = checkbox.find(f"{{{CHECKBOX_NS}}}{'checkedState' if value == 'true' else 'uncheckedState'}")
                glyph = chr(int(state.get(f"{{{CHECKBOX_NS}}}val"), 16)) if state is not None else ("☒" if value == "true" else "☐")
                nodes = cell.xpath(".//w:t", namespaces={"w": WORD_NS})
                if not nodes:
                    raise TemplateError("DOCX 체크박스의 표시 문자 구조가 없음")
                nodes[0].text = glyph
                for node in nodes[1:]:
                    node.text = ""
                continue
            if field["kind"] == "docx_inline" or field['kind'] in repeat_kinds:
                segments = _segments(cell, kind)
                text = "".join(segment.value for segment in segments)
                if field['kind'] in repeat_kinds:
                    original_root = _xml(archive.read(part))
                    original = original_root.xpath(xpath, namespaces={k: v for k, v in original_root.nsmap.items() if k})
                    if len(original) != 1 or etree.QName(cell).localname != 'p':
                        raise TemplateError('반복 자리표시자의 원본 문단이 없음')
                    original_text = ''.join(segment.value for segment in _segments(original[0], kind))
                    start, end = field.get('placeholder_start'), field.get('placeholder_end')
                    matches = list(PLACEHOLDER.finditer(original_text))
                    if (original_text != field.get('anchor_text') or type(start) is not int or type(end) is not int
                            or not any(match.start() == start and match.end() == end and match[1].strip() == field.get('placeholder_key') for match in matches)
                            or text[start:end] != original_text[start:end]):
                        raise TemplateError('반복 자리표시자의 원문 또는 입력 범위가 변경됨')
                else:
                    start, end = field["blank_start"], field["blank_end"]
                    if text != field["anchor_text"] or not 0 <= start < end <= len(text) or not _blank_marker(text[start:end]):
                        raise TemplateError("DOCX 원문 또는 빈 입력 범위가 변경됨")
                position, affected = 0, []
                for segment in segments:
                    next_position = position + len(segment.value)
                    if position < end and next_position > start:
                        if not segment.editable:
                            raise TemplateError("빈 입력 범위가 제어 문자를 가로지름")
                        affected.append((segment, position))
                    position = next_position
                if not affected:
                    raise TemplateError("빈 입력 범위의 텍스트 구조가 없음")
                first, first_start = affected[0]
                last, last_start = affected[-1]
                prefix, suffix = first.value[:start-first_start], last.value[end-last_start:]
                token = mapped_token(index)
                substitutions[token] = value
                _set_segment(first, prefix + "{{" + token + "}}" + (suffix if first is last else ""), kind)
                for segment, _ in affected[1:-1]:
                    _set_segment(segment, "", kind)
                if first is not last:
                    _set_segment(last, suffix, kind)
                continue
            token = mapped_token(index)
            substitutions[token] = value
            paragraphs = [cell] if etree.QName(cell).localname == "p" else cell.xpath(".//*[local-name()='p']")
            if field["kind"] == "docx_append":
                if _text(cell) != field.get("anchor_text") or etree.QName(cell).localname != "tc":
                    raise TemplateError("본문 입력 셀의 원문 안내가 변경됨")
                paragraph = etree.SubElement(cell, f"{{{WORD_NS}}}p")
                if paragraphs:
                    properties = paragraphs[0].find(f"{{{WORD_NS}}}pPr")
                    if properties is not None:
                        paragraph.append(deepcopy(properties))
                    run = etree.SubElement(paragraph, f"{{{WORD_NS}}}r")
                    properties = paragraphs[0].find(f"{{{WORD_NS}}}r/{{{WORD_NS}}}rPr")
                    if properties is not None:
                        run.append(deepcopy(properties))
            elif paragraphs:
                paragraph = paragraphs[0]
            elif kind == "docx" and field["kind"] == "docx_sdt" and cell.xpath("ancestor::w:p", namespaces={"w": WORD_NS}):
                paragraph = cell  # 문단 안의 inline 콘텐츠 컨트롤은 run을 직접 유지함.
            elif kind == "docx":
                paragraph = etree.SubElement(cell, f"{{{WORD_NS}}}p")
            else:
                raise TemplateError("HWPX 빈 입력칸에 문단 구조가 없음")
            text_nodes = paragraph.xpath(".//*[local-name()='t']")
            if not text_nodes:
                namespace = WORD_NS if kind == "docx" else etree.QName(paragraph).namespace
                runs = paragraph.xpath("./*[local-name()='r' or local-name()='run']")
                if runs:
                    run = runs[0]
                else:
                    run = etree.SubElement(paragraph, f"{{{namespace}}}{'r' if kind == 'docx' else 'run'}")
                    if kind == "hwpx":
                        run.set("charPrIDRef", "0")
                text_nodes = [etree.SubElement(run, f"{{{namespace}}}t")]
            for text in text_nodes:
                text.text = ""
                for child in list(text):
                    text.remove(child)
            text_nodes[0].text = "{{" + token + "}}"
            # showingPlcHdr는 실제값 입력 후 제거하되 sdtPr의 스타일·alias는 보존함.
            if field["kind"] == "docx_sdt":
                control = cell.getparent()
                for element in control.xpath("./w:sdtPr/w:showingPlcHdr", namespaces={"w": WORD_NS}):
                    element.getparent().remove(element)
        # 일반 {{...}} 치환이 선택 표시/잠금·미지원 SDT의 원문을 다시 해석하지 않도록 보존함.
        if substitutions and source.suffix.lower() == ".docx":
            special = ".//w:sdt[w:sdtPr/w:dropDownList or w:sdtPr/w:comboBox or w:sdtPr/w:lock or w:sdtPr/w:date or w:sdtPr/w:picture or w:sdtPr/w:group or w:sdtPr/*[local-name()='repeatingSection' or local-name()='repeatingSectionItem']]/w:sdtContent"
            for part in _parts(archive, "docx"):
                root = roots.get(part)
                if root is None:
                    root = _xml(archive.read(part))
                contents = root.xpath(special, namespaces={"w": WORD_NS})
                if contents:
                    roots[part] = root
                for content in contents:
                    shielded.setdefault(part, []).append((root.getroottree().getpath(content), deepcopy(content)))
                    for node in content.xpath(".//w:t", namespaces={"w": WORD_NS}):
                        node.text = ""  # 구조/런/속성은 유지하고 치환 후 전체 content를 되돌림.
        changes = {name: etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True) for name, root in roots.items()}
    if not substitutions:
        _rewrite_zip(source, target, changes)
        return
    descriptor, temporary = tempfile.mkstemp(suffix=source.suffix)
    os.close(descriptor)
    try:
        _rewrite_zip(source, Path(temporary), changes)
        if shielded:
            with tempfile.TemporaryDirectory() as directory:
                filled = Path(directory) / ("filled" + source.suffix)
                fill_template(temporary, substitutions, filled)
                restored = {}
                with ZipFile(filled) as archive:
                    for part, contents in shielded.items():
                        root = _xml(archive.read(part))
                        for xpath, content in contents:
                            nodes = root.xpath(xpath, namespaces={k: v for k, v in root.nsmap.items() if k})
                            if len(nodes) != 1:
                                raise TemplateError("보존한 DOCX 콘텐츠 컨트롤 위치가 변경됨")
                            nodes[0].getparent().replace(nodes[0], content)
                        restored[part] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                _rewrite_zip(filled, target, restored)
        else:
            fill_template(temporary, substitutions, target)
    finally:
        os.unlink(temporary)


def _xlsx_numeric_text(field, value):
    """Canonical numeric XML for explicitly typed blank cells; never convert IDs.

    Decimal values lose grouping/leading plus/trailing fraction zeros only.
    Excel's 15 significant digit and conservative published entry limits apply;
    IEEE binary arithmetic/recalculation is separate from this exact XML value.
    """
    rule = field.get('validation', {})
    if (field['kind'] != 'xlsx_cell' or rule.get('type') not in {'integer', 'number'}
            or rule.get('unit') and rule.get('unit_location', 'value') != 'label'
            or re.search(r'\[S[A-Za-z0-9_-]+\]', value)):
        return None
    number = Decimal(value.strip().replace(',', ''))
    significant = len(''.join(map(str, number.as_tuple().digits)).rstrip('0')) or 1
    magnitude = number.copy_abs()
    if (significant > 15 or magnitude > Decimal('9.99999999999999E307')
            or magnitude != 0 and magnitude < Decimal('2.2251E-308')):
        raise TemplateError(f"Excel 숫자 정밀도·저장 범위를 초과함: {field['label']} (15 유효자리)")
    text = format(number, 'f')
    return ('0' if not number else text.rstrip('0').rstrip('.') if '.' in text else text)


def _xlsx_fill(source, target, selected):
    from .xlsx_validation import XlsxLists, create_choice_cell
    with ZipFile(source) as archive:
        shared, _ = _xlsx_parts(archive)
        lists = XlsxLists(archive)
        precision = _xml(archive.read('xl/workbook.xml')).find(f'{{{SHEET_NS}}}calcPr')
        displayed_precision = precision is not None and precision.get('fullPrecision') in {'0', 'false', 'off'}
        roots, selections = {}, {}
        for field, value in selected:
            parts = field["id"].split(":", 3)
            _, part, coordinate = parts[:3]
            selections.setdefault((part, coordinate), []).append((field, value))
        for (part, coordinate), entries in selections.items():
            root = roots.setdefault(part, _xml(archive.read(part)))
            try:
                native_list = lists.check(part, coordinate, entries[0][0], entries[0][1])
            except ValueError as exc:
                raise TemplateError(str(exc)) from exc
            cells = root.xpath(f".//s:sheetData/s:row/s:c[@r='{coordinate}']", namespaces={"s": SHEET_NS})
            if not cells and native_list is not None:
                try:
                    cells = [create_choice_cell(root, coordinate)]
                except ValueError as exc:
                    raise TemplateError(str(exc)) from exc
            if len(cells) != 1:
                raise TemplateError("XLSX 입력칸 위치가 변경됨")
            cell = cells[0]
            text = _xlsx_cell_text(cell, shared)
            if cell.find(f"{{{SHEET_NS}}}f") is not None:
                raise TemplateError('XLSX 원본 수식 셀은 입력값으로 덮어쓸 수 없음')
            # A native list export value stays exact text, including IDs and zero.
            numeric = _xlsx_numeric_text(entries[0][0], entries[0][1]) if not text.strip() and native_list is None else None
            if numeric is not None and displayed_precision:
                raise TemplateError('Excel 표시 정밀도 설정으로 숫자 값이 반올림될 수 있어 입력을 차단함')
            rich_source = None
            if cell.get("t") == "inlineStr":
                rich_source = cell.find(f"{{{SHEET_NS}}}is")
            elif cell.get("t") == "s" and "xl/sharedStrings.xml" in archive.namelist():
                value_node = cell.find(f"{{{SHEET_NS}}}v")
                rich_source = _xml(archive.read("xl/sharedStrings.xml"))[int(value_node.text)]
            rich = deepcopy(rich_source) if rich_source is not None else None
            if entries[0][0]["kind"] == "xlsx_placeholder":
                values = {field["label"]: value for field, value in entries}
                if rich is not None:
                    rich.tag = f"{{{SHEET_NS}}}is"
                    _xlsx_replace_rich_text(rich, values)
                else:
                    text = PLACEHOLDER.sub(lambda match: values[match[1].strip()], text)
            else:
                text = entries[0][1]
                rich = None
            for child in list(cell):
                if etree.QName(child).localname in {"v", "is"}:
                    cell.remove(child)
            cell.set("t", "n" if numeric is not None else "inlineStr")
            if numeric is not None:
                etree.SubElement(cell, f"{{{SHEET_NS}}}v").text = numeric
            elif rich is not None:
                cell.append(rich)
            else:
                inline = etree.SubElement(cell, f"{{{SHEET_NS}}}is")
                node = etree.SubElement(inline, f"{{{SHEET_NS}}}t")
                node.set(XML_SPACE, "preserve")
                node.text = text.replace("\r\n", "\n").replace("\r", "\n")
        changes = {name: etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True) for name, root in roots.items()}
    _rewrite_zip(source, target, changes)


def _xlsx_replace_rich_text(inline, values):
    nodes = inline.xpath("./s:t | ./s:r/s:t", namespaces={"s": SHEET_NS})
    original = "".join(node.text or "" for node in nodes)
    starts, offset = [], 0
    for node in nodes:
        starts.append(offset)
        offset += len(node.text or "")
    lengths = [len(node.text or "") for node in nodes]
    for match in reversed(list(PLACEHOLDER.finditer(original))):
        affected = [index for index in range(len(nodes)) if starts[index] < match.end() and starts[index] + lengths[index] > match.start()]
        first, last = affected[0], affected[-1]
        prefix = (nodes[first].text or "")[:match.start() - starts[first]]
        suffix = (nodes[last].text or "")[match.end() - starts[last]:]
        replacement = values[match[1].strip()].replace("\r\n", "\n").replace("\r", "\n")
        nodes[first].text = prefix + replacement + (suffix if first == last else "")
        for index in affected[1:-1]:
            nodes[index].text = ""
        if first != last:
            nodes[last].text = suffix
        for index in affected:
            nodes[index].set(XML_SPACE, "preserve")


def _pdf_page_link_repairs(reader):
    """Recover only dangling Widget /P links whose real placement is unique.

    The canonical tree and every page's /Annots remain authoritative. A valid
    different page, orphan or duplicate is never interpreted as a repair.
    """
    if reader.is_encrypted or _pdf_signed(reader):
        raise TemplateError('보호·서명된 PDF의 페이지 연결을 복구할 수 없음')
    form = reader.trailer['/Root'].get('/AcroForm')
    if form is None:
        return []
    canonical, names = {}, set()

    def visit(reference, parent=None, prefix=''):
        node = reference.get_object()
        if not isinstance(node, dict) or id(node) in canonical:
            raise TemplateError('PDF canonical 필드의 연결이 중복되거나 모호함')
        link = node.get('/Parent')
        if (parent is None and link is not None or parent is not None and
                (link is None or link.get_object() is not parent)):
            raise TemplateError('PDF canonical 필드의 Parent 연결이 모호함')
        name = '.'.join(part for part in (prefix, str(node.get('/T', ''))) if part)
        if '/T' in node:
            if name in names:
                raise TemplateError('PDF canonical 필드 이름이 중복됨')
            names.add(name)
        canonical[id(node)] = node
        for child in node.get('/Kids', []):
            visit(child, node, name)

    for reference in form.get_object().get('/Fields', []):
        visit(reference)
    if names != set(reader.get_fields() or {}):
        raise TemplateError('PDF canonical 필드 이름과 실제 트리가 일치하지 않음')
    seen, repairs = set(), []
    for number, page in enumerate(reader.pages):
        for ordinal, reference in enumerate(page.get('/Annots', [])):
            node = reference.get_object()
            if not isinstance(node, dict):
                raise TemplateError('PDF 페이지 주석의 연결을 확인할 수 없음')
            widget = node.get('/Subtype') == '/Widget'
            if widget:
                if id(node) not in canonical or id(node) in seen:
                    raise TemplateError('PDF Widget과 canonical 필드의 페이지 연결이 모호함: 중복 또는 원본 트리 밖 위젯')
                seen.add(id(node))
            link = node.get('/P')
            if link is None or link == page.indirect_reference:
                continue
            if widget and isinstance(link, IndirectObject) and link.get_object() is None:
                repairs.append((number, ordinal))
            else:
                raise TemplateError('PDF 위젯/주석이 실제 다른 페이지 또는 잘못된 페이지를 가리킴')
    if seen != {identity for identity, node in canonical.items() if node.get('/Subtype') == '/Widget'}:
        raise TemplateError('PDF canonical Widget의 실제 페이지 배치가 누락됨')
    return repairs


def _pdf_fill(source, target, selected):
    reader = PdfReader(source)
    repairs = _pdf_page_link_repairs(reader)
    writer = PdfWriter()
    writer.clone_document_from_reader(reader)
    for page, ordinal in repairs:
        writer.pages[page]['/Annots'][ordinal].get_object()[NameObject('/P')] = writer.pages[page].indirect_reference
    values = {field["id"][4:]: value for field, value in selected}
    source_fields = reader.get_fields() or {}
    choices = {}
    immutable = ('control_type', 'options', 'choice_items', 'allow_custom', 'multiselect',
                 'selection_encoding', 'pdf_choice_flags', 'pdf_blank_options')
    for field, value in selected:
        name = field['id'][4:]
        native_field = source_fields.get(name)
        if native_field is None:
            raise TemplateError('PDF 원본 입력칸이 없음')
        if _pdf_inherited(native_field, '/FT', '') == '/Ch':
            flags = int(_pdf_inherited(native_field, '/Ff', 0))
            native = choice_metadata(_pdf_inherited(native_field, '/Opt', []), flags)
            if any(field.get(key) != native.get(key) for key in immutable) or flags & 1:
                raise TemplateError('원본 PDF 선택 조건을 프로파일로 바꿀 수 없음')
            # _selected_fields already separated sidecar provenance. Its value
            # is the actual one-line export/custom literal, including [S...].
            choices[name] = (native, parse_choice_value(native, value, literal=value))
            linked = 0
            for page in reader.pages:
                for reference in page.get('/Annots', []):
                    widget = reference.get_object()
                    if _pdf_field_name(widget) != name:
                        continue
                    node, seen = widget, set()
                    while node is not None and id(node) not in seen:
                        seen.add(id(node))
                        if node.indirect_reference == native_field.indirect_reference:
                            linked += 1
                            break
                        node = node.get('/Parent')
                        node = node.get_object() if node else None
                    else:
                        raise TemplateError('PDF 선택 위젯과 canonical 원본 필드의 연결이 모호함')
            if not linked:
                raise TemplateError('PDF 선택 필드에 실제 페이지 위젯이 없음')
        elif field.get('control_type') in {'choice', 'combobox'}:
            raise TemplateError('PDF 원본 입력 종류와 선택 프로파일이 일치하지 않음')
    for page in reader.pages:
        for reference in page.get('/Annots', []):
            widget = reference.get_object()
            name = _pdf_field_name(widget)
            if name not in values or _pdf_inherited(widget, '/FT', '') != '/Tx' or any(ord(c) > 127 for c in values[name]):
                continue
            da = str(_pdf_inherited(widget, '/DA', reader.trailer['/Root']['/AcroForm'].get('/DA', '')))
            match = re.search(r'(/[^\s]+)\s+([\d.]+)\s+Tf', da)
            resources = _pdf_inherited(widget, '/DR', reader.trailer['/Root']['/AcroForm'].get('/DR', {}))
            if match is None or resources.get('/Font', {}).get(match[1]) is None:
                raise TemplateError(f'PDF 입력칸의 원본 글꼴 자원을 확인할 수 없어 기입을 보류함: {name}')
    original_appearances = {}
    for page in writer.pages:
        for reference in page.get("/Annots", []):
            widget = reference.get_object()
            name = _pdf_field_name(widget)
            if name in choices or name in values and any(ord(character) > 127 for character in values[name]):
                appearance = widget.get("/AP", {}).get("/N")
                if appearance is not None:
                    # pypdf는 기존 appearance stream 자체를 수정하므로 빈 배경을 먼저 복제함.
                    original = appearance.get_object()
                    background = DecodedStreamObject()
                    stream = ContentStream(original, writer)
                    operations, inside_text, choice_content = [], False, 0
                    for operands, operator in stream.operations:
                        if name in choices:
                            if operator in {b'BMC', b'BDC'}:
                                if choice_content or operands and str(operands[0]) == '/Tx':
                                    choice_content += 1
                                    continue
                            elif operator == b'EMC' and choice_content:
                                choice_content -= 1
                                continue
                            if choice_content:
                                continue
                        if operator == b"BT":
                            inside_text = True
                        elif operator == b"ET":
                            inside_text = False
                        elif not inside_text:
                            operations.append((operands, operator))
                    stream.operations = operations
                    background.set_data(stream.get_data())
                    for key, value in original.items():
                        if key not in {"/Length", "/Filter", "/DecodeParms"}:
                            background[key] = value
                    original_appearances[id(widget)] = writer._add_object(background)
    # The custom embedded appearance below handles Unicode. Giving those values
    # to pypdf first creates an unnecessary, sometimes unencodable font fallback.
    unicode_values = {name: value for name, value in values.items()
                      if name not in choices and any(ord(character) > 127 for character in value)}
    ordinary_values = {name: '' if name in unicode_values else value for name, value in values.items() if name not in choices}
    if ordinary_values:
        writer.update_page_form_field_values(None, ordinary_values, auto_regenerate=False)
    elif choices:
        writer.set_need_appearances_writer(False)
    pending = list(writer._root_object["/AcroForm"]["/Fields"])
    while pending:
        reference = pending.pop()
        node = reference.get_object()
        name = _pdf_field_name(node)
        pending.extend(node.get('/Kids', []))
        if name in choices and ('/T' in node or '/V' in node):
            native, codes = choices[name]
            indices = _pdf_choice_indices(native, codes)
            node[NameObject('/V')] = (ArrayObject([TextStringObject(code) for code in codes])
                                     if native['multiselect'] else TextStringObject(codes[0]))
            if indices:
                node[NameObject('/I')] = ArrayObject([NumberObject(index) for index in indices])
            else:
                node.pop(NameObject('/I'), None)
        elif name in unicode_values and ('/T' in node or '/V' in node):
            node[NameObject('/V')] = TextStringObject(unicode_values[name])
        if name in values and node.get("/FT") == "/Btn":
            node[NameObject("/V")] = NameObject(values[name])
    painted_choices, choice_tops = set(), {}
    for page in writer.pages:
        for reference in page.get("/Annots", []):
            widget = reference.get_object()
            original = original_appearances.get(id(widget))
            name = _pdf_field_name(widget)
            if name in choices:
                native, codes = choices[name]
                widget_native = choice_metadata(_pdf_inherited(widget, '/Opt', []), int(_pdf_inherited(widget, '/Ff', 0)))
                if widget_native != native or _pdf_inherited(widget, '/FT', '') != '/Ch':
                    raise TemplateError('PDF 선택 위젯과 원본 필드의 종류·목록이 일치하지 않음')
                top = _pdf_choice_appearance(writer, widget, original, native, codes)
                if name in choice_tops and choice_tops[name] != top:
                    raise TemplateError('같은 PDF 선택 필드의 위젯 표시 범위가 서로 달라 기입을 보류함')
                choice_tops[name] = top
                painted_choices.add(name)
            elif original is not None:
                name = _pdf_field_name(widget)
                _pdf_unicode_appearance(writer, widget, original, values[name])
            elif _pdf_field_name(widget) in unicode_values:
                # A blank native appearance generated above is also a safe base
                # for originally appearance-less text fields, with native border.
                appearance = widget.get('/AP', {}).get('/N')
                if appearance is None or not hasattr(appearance.get_object(), 'get_data'):
                    raise TemplateError('PDF 한글 입력칸의 원래 표시/테두리를 확인할 수 없음')
                _pdf_unicode_appearance(writer, widget, appearance, unicode_values[_pdf_field_name(widget)])
    if painted_choices != set(choices):
        raise TemplateError('PDF 선택 필드의 실제 페이지 입력 위치가 없음')
    for page in writer.pages:
        for reference in page.get('/Annots', []):
            widget = reference.get_object()
            name = _pdf_field_name(widget)
            if name in values and name not in choices and _pdf_inherited(widget, '/FT', '') == '/Tx':
                da = str(_pdf_inherited(widget, '/DA', writer.root_object['/AcroForm'].get('/DA', '')))
                native = re.search(r'(/[^\s]+)\s+([\d.]+)\s+Tf', da)
                _pdf_text_appearance_fit(widget, values[name], name,
                    adjust_auto=name not in unicode_values and native is not None and float(native[2]) == 0)
    target.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary = tempfile.mkstemp(dir=target.parent, suffix=".pdf")
    os.close(descriptor)
    try:
        with open(temporary, "wb") as stream:
            writer.write(stream)
        final_fields = PdfReader(temporary).get_fields() or {}
        if any((list(final_fields.get(name, {}).get('/V', [])) != choices[name][1]
                if name in choices and choices[name][0]['multiselect'] else
                str(final_fields.get(name, {}).get('/V', '')) != (choices[name][1][0] if name in choices else value))
               for name, value in values.items()):
            raise TemplateError("PDF 폼 값 저장 검증에 실패함")
        os.replace(temporary, target)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


def _pdf_text_appearance_fit(widget, value, label, *, adjust_auto=False):
    """Check the actual saved glyph positions, including pypdf's ASCII AP.

    Native fixed font sizes and alignment can overflow while /V remains whole.
    Read the generated appearance before committing the output; never truncate.
    """
    import pypdfium2 as pdfium

    appearance = widget.get('/AP', {}).get('/N')
    if appearance is None or not hasattr(appearance.get_object(), 'get_data'):
        raise TemplateError(f'PDF 입력칸의 저장된 표시를 확인할 수 없음: {label}')
    appearance = appearance.get_object()
    rectangle, box = widget.get('/Rect', []), appearance.get('/BBox', [])
    if len(rectangle) != 4 or len(box) != 4:
        raise TemplateError(f'PDF 입력칸의 표시 경계가 모호함: {label}')
    width, height = float(rectangle[2]-rectangle[0]), float(rectangle[3]-rectangle[1])
    if list(map(float, box)) != [0., 0., width, height]:
        if any(abs(float(a)-b) > 1e-4 for a, b in zip(box, [0., 0., width, height])):
            raise TemplateError(f'PDF 입력칸의 원래 표시 좌표를 해석할 수 없음: {label}')
    if widget.get('/MK', {}).get('/R', 0) or list(map(float, appearance.get('/Matrix', [1, 0, 0, 1, 0, 0]))) != [1., 0., 0., 1., 0., 0.]:
        raise TemplateError(f'회전한 PDF 입력칸의 표시 경계를 확인할 수 없음: {label}')
    writer, buffer = PdfWriter(), BytesIO()
    page = writer.add_blank_page(width=width, height=height)
    page[NameObject('/Resources')] = appearance.get('/Resources', DictionaryObject()).clone(writer)
    content = DecodedStreamObject()
    content.set_data(bytes(appearance.get_data()))
    page[NameObject('/Contents')] = writer._add_object(content)
    writer.write(buffer)
    document = pdfium.PdfDocument(buffer.getvalue())
    glyphs = []
    try:
        rendered = document[0]
        text = rendered.get_textpage()
        try:
            normalized = lambda s: re.sub(r'\s+', '', s)
            if normalized(value) != normalized(text.get_text_range()):
                raise TemplateError(f'PDF 입력칸의 원문 전체 표시를 확인할 수 없음: {label}')
            for index in range(text.count_chars()):
                if not text.get_text_range(index, 1).strip():
                    continue
                left, bottom, right, top = text.get_charbox(index)
                glyphs.append((left, bottom, right, top))
        finally:
            text.close()
            rendered.close()
    finally:
        document.close()
    if any(not (-.01 <= left <= right <= width+.01 and -.01 <= bottom <= top <= height+.01)
           for left, bottom, right, top in glyphs):
        left = min(b[0] for b in glyphs); bottom = min(b[1] for b in glyphs)
        right = max(b[2] for b in glyphs); top = max(b[3] for b in glyphs)
        # Native auto sizing can fit all lines but place a descender just below
        # the boundary. Translate its first text origin only when the complete
        # glyph rectangle already fits. No font, wrapping, text or DA changes.
        if adjust_auto and right-left <= width-.02 and top-bottom <= height-.02:
            dx = max(.01-left, 0.) if left < 0 else min(width-.01-right, 0.)
            dy = max(.01-bottom, 0.) if bottom < 0 else min(height-.01-top, 0.)
            stream = ContentStream(appearance, appearance.indirect_reference.pdf)
            text_open = False
            for args, operator in stream.operations:
                if operator == b'BT': text_open = True
                elif operator == b'ET': text_open = False
                elif text_open and operator == b'Td' and len(args) == 2:
                    args[0] = FloatObject(float(args[0])+dx); args[1] = FloatObject(float(args[1])+dy)
                    appearance.set_data(stream.get_data())
                    return _pdf_text_appearance_fit(widget, value, label)
                elif text_open and operator == b'Tm' and len(args) == 6:
                    args[4] = FloatObject(float(args[4])+dx); args[5] = FloatObject(float(args[5])+dy)
                    appearance.set_data(stream.get_data())
                    return _pdf_text_appearance_fit(widget, value, label)
        raise TemplateError(f'PDF 입력칸 글자가 표시 경계를 넘침: {label}')


def _pdf_choice_indices(native, codes):
    blanks = {item['index'] for item in native['pdf_blank_options']}
    indices = [index for index in range(len(native['options']) + len(blanks)) if index not in blanks]
    positions = dict(zip(native['options'], indices))
    return [positions[code] for code in codes if code in positions]


def _pdf_choice_font(writer, widget, labels):
    da = str(_pdf_inherited(widget, '/DA', writer._root_object['/AcroForm'].get('/DA', '')))
    found = re.search(r'(/[^\s]+)\s+([\d.]+)\s+Tf', da)
    if not found:
        raise TemplateError('PDF 선택칸의 원본 글꼴·크기를 확인할 수 없음')
    resource, size = found[1], float(found[2])
    if size < 0 or size > 100:
        raise TemplateError('PDF 선택칸의 글꼴 크기가 올바르지 않음')
    if any(ord(character) > 127 for character in ''.join(labels)):
        bundled = Path(__file__).with_name('fonts') / 'NotoSansKR-Regular.ttf'
        return _pdf_font({'font_path': str(bundled)} if size == 0 and bundled.is_file() else {}, ''.join(labels)), size, True
    resources = _pdf_inherited(widget, '/DR', writer._root_object['/AcroForm'].get('/DR', {}))
    font = resources.get('/Font', {}).get(resource)
    base = str(font.get_object().get('/BaseFont', '')).lstrip('/') if font is not None else ''
    if base not in {'Helvetica', 'Helvetica-Bold', 'Helvetica-Oblique', 'Helvetica-BoldOblique',
                    'Times-Roman', 'Times-Bold', 'Times-Italic', 'Times-BoldItalic',
                    'Courier', 'Courier-Bold', 'Courier-Oblique', 'Courier-BoldOblique'}:
        # Resolve source embedded TTFs rather than silently substituting an ASCII font.
        descriptor = font.get_object().get('/FontDescriptor', {}) if font is not None else {}
        embedded = descriptor.get('/FontFile2')
        if embedded is None:
            raise TemplateError('PDF 선택칸의 원본 글꼴을 표시용 글꼴로 해석할 수 없음')
        data = embedded.get_object().get_data()
        base = 'ChoiceFont' + sha256(data).hexdigest()[:12]
        if base not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(base, BytesIO(data)))
    return base, size, False


def _pdf_choice_appearance(writer, widget, original, native, codes):
    rectangle = widget.get('/Rect', [])
    if len(rectangle) != 4 or widget.get('/MK', {}).get('/R', 0):
        raise TemplateError('회전하거나 위치가 모호한 PDF 선택칸은 지원하지 않음')
    width, height = float(rectangle[2] - rectangle[0]), float(rectangle[3] - rectangle[1])
    if not 4 < width < 100000 or not 4 < height < 100000:
        raise TemplateError('PDF 선택칸의 크기가 올바르지 않음')
    if original is None:
        raise TemplateError('PDF 선택칸의 원래 표시·테두리를 확인할 수 없음')
    items = native['choice_items']
    indices = _pdf_choice_indices(native, native['options'])
    by_index = {index: item['label'] for index, item in zip(indices, items)}
    by_index.update({item['index']: item['label'] for item in native['pdf_blank_options']})
    selected = _pdf_choice_indices(native, codes)
    combo = bool(native['pdf_choice_flags'] & (1 << 17))
    labels = [next((item['label'] for item in items if item['value'] == codes[0]), codes[0])] if combo else list(by_index.values())
    font, requested, unicode = _pdf_choice_font(writer, widget, labels)
    size = requested or 12.0
    if not requested:
        size = min(size, (height - 4) / 1.3, (width - 4) / max(pdfmetrics.stringWidth(label, font, 1) or 1 for label in labels))
    if size < 4:
        raise TemplateError('PDF 선택 표시가 4pt보다 작아 원문 전체 표시를 보류함')
    ascent, descent = pdfmetrics.getAscentDescent(font, size)
    leading = size * 1.3
    if ascent - descent > height - 4:
        raise TemplateError('PDF 선택 문구가 입력칸 높이를 넘음')
    if combo:
        shown = [(None, labels[0])]
        top = None
    else:
        capacity = int((height - 4) // leading)
        if not capacity:
            raise TemplateError('PDF 목록의 한 줄을 모두 표시할 수 없음')
        original_top = _pdf_inherited(widget, '/TI', 0)
        if not isinstance(original_top, int) or not 0 <= original_top < len(by_index):
            raise TemplateError('PDF 목록의 원본 표시 시작 위치가 올바르지 않음')
        top = int(original_top)
        if not all(top <= index < top + capacity for index in selected):
            top = min(min(selected), max(0, len(by_index) - capacity))
        if not all(top <= index < top + capacity for index in selected):
            raise TemplateError('다중 선택한 문구를 같은 목록 영역에 모두 표시할 수 없음')
        shown = [(index, by_index[index]) for index in range(top, min(top + capacity, len(by_index)))]
    if any(pdfmetrics.stringWidth(label, font, size) > width - 4 for _, label in shown):
        raise TemplateError('PDF 선택 문구가 입력칸 폭을 넘음')
    buffer = BytesIO()
    drawing = canvas.Canvas(buffer, pagesize=(width, height), invariant=1)
    drawing.addLiteral('/Tx BMC')
    drawing.setFont(font, size)
    da = str(_pdf_inherited(widget, '/DA', writer._root_object['/AcroForm'].get('/DA', '')))
    color = re.findall(r'((?:[-+]?\d*\.?\d+\s+){1,4})(rg|g|k)\b', da)
    for relative, (index, label) in enumerate(shown):
        if index in selected:
            drawing.setFillColorRGB(.6, .75, .95)
            drawing.rect(2, height - 2 - (relative + 1) * leading, width - 4, leading, stroke=0, fill=1)
        components, operator = ([float(v) for v in color[-1][0].split()], color[-1][1]) if color else ([0], 'g')
        if operator == 'rg' and len(components) == 3:
            drawing.setFillColorRGB(*components)
        elif operator == 'k' and len(components) == 4:
            drawing.setFillColorCMYK(*components)
        elif operator == 'g' and len(components) == 1:
            drawing.setFillGray(*components)
        else:
            raise TemplateError('PDF 선택칸의 원본 글자 색상을 확인할 수 없음')
        y = ((height - ascent - descent) / 2 if combo else height - 2 - ascent - relative * leading)
        alignment = int(_pdf_inherited(widget, '/Q', 0))
        if alignment == 1:
            drawing.drawCentredString(width / 2, y, label)
        elif alignment == 2:
            drawing.drawRightString(width - 2, y, label)
        else:
            drawing.drawString(2, y, label)
    drawing.addLiteral('EMC')
    drawing.showPage(); drawing.save(); buffer.seek(0)
    page = PdfReader(buffer).pages[0]
    resources = page['/Resources'].clone(writer)
    data = page.get_contents().get_data()
    if original is not None:
        resources.setdefault(NameObject('/XObject'), DictionaryObject())[NameObject('/OriginalAppearance')] = original
        data = b'q /OriginalAppearance Do Q\n' + data
    appearance = DecodedStreamObject()
    appearance.set_data(data)
    appearance.update({NameObject('/Type'): NameObject('/XObject'), NameObject('/Subtype'): NameObject('/Form'),
                      NameObject('/BBox'): ArrayObject([FloatObject(0), FloatObject(0), FloatObject(width), FloatObject(height)]),
                      NameObject('/Resources'): resources})
    old_ap = widget.get('/AP', DictionaryObject()).get_object()
    widget[NameObject('/AP')] = DictionaryObject(dict(old_ap))
    widget['/AP'][NameObject('/N')] = writer._add_object(appearance)
    if top is not None:
        owner = widget
        while '/T' not in owner and owner.get('/Parent'):
            owner = owner['/Parent']
        owner[NameObject('/TI')] = NumberObject(top)
        if '/TI' in widget:
            widget[NameObject('/TI')] = NumberObject(top)
    if unicode:
        resource = _pdf_cjk_form_font(writer, font, ''.join(label for _, label in shown))
        widget[NameObject('/DA')] = TextStringObject(f'{resource} {0 if not requested else size:g} Tf 0 g')
        _pdf_cjk_local_resources(writer, widget, resource)
        if widget.get('/Parent'):
            widget['/Parent'][NameObject('/DA')] = widget['/DA']
            _pdf_cjk_local_resources(writer, widget['/Parent'], resource)
    return top


def _pdf_field_name(widget):
    names, visited = [], set()
    current = widget
    while current is not None and id(current) not in visited:
        visited.add(id(current))
        if current.get("/T"):
            names.append(str(current["/T"]))
        parent = current.get("/Parent")
        current = parent.get_object() if parent else None
    return ".".join(reversed(names))


def _pdf_inherited(widget, key, default):
    seen = set()
    while widget is not None:
        identity = id(widget)
        if identity in seen:
            raise TemplateError('PDF 입력칸의 부모 연결이 순환함')
        seen.add(identity)
        if key in widget:
            return widget[key]
        parent = widget.get('/Parent')
        widget = parent.get_object() if parent else None
    return default


def _pdf_unicode_appearance(writer, widget, original, value):
    """빈 입력칸의 원래 배경/테두리를 유지하고 한글 TTF appearance를 내장함."""
    rectangle = widget["/Rect"]
    width, height = float(rectangle[2] - rectangle[0]), float(rectangle[3] - rectangle[1])
    da = str(_pdf_inherited(widget, "/DA", writer._root_object['/AcroForm'].get('/DA', '/Helv 10 Tf')))
    found = re.search(r"\s([\d.]+)\s+Tf", da)
    requested = float(found[1]) if found else 10.0
    automatic = requested == 0
    size = 12.0 if automatic else requested
    # The bundled font's actual full bounding box also leaves room for
    # descenders when native viewers choose an automatic size in small fields.
    bundled = Path(__file__).with_name('fonts') / 'NotoSansKR-Regular.ttf'
    font = _pdf_font({'font_path': str(bundled)} if automatic and bundled.is_file() else {}, value)
    multiline = bool(int(_pdf_inherited(widget, '/Ff', 0)) & 4096)
    if automatic and not multiline:
        if '\n' in value or '\r' in value:
            raise TemplateError('한 줄 PDF 입력칸에 여러 줄을 기입할 수 없음')
        size = min(size, (height - 4) / 1.3,
                   (width - 8) / (pdfmetrics.stringWidth(value, font, 1) or 1))
        lines = [value]
    else:
        lines = _wrap_pdf_text(value, font, size, width - 8)
        while automatic and size > 4 and len(lines) * size * 1.3 + 4 > height:
            size -= 0.25
            lines = _wrap_pdf_text(value, font, size, width - 8)
    if automatic and size < 4:
        raise TemplateError('PDF 자동 글꼴 크기가 4pt보다 작아 원문 전체 표시를 보류함')
    if len(lines) * size * 1.3 + 4 > height:
        raise TemplateError("PDF 입력칸에 내용이 넘침. 더 넓은 입력칸이나 수동 PDF 영역을 사용해야 함")
    buffer = BytesIO()
    drawing = canvas.Canvas(buffer, pagesize=(width, height), invariant=1)
    drawing.setFont(font, size)
    drawing.setFillColorRGB(0, 0, 0)
    alignment = int(_pdf_inherited(widget, '/Q', 0))
    for index, line in enumerate(lines):
        y = height - 4 - size - index * size * 1.3
        if alignment == 1:
            drawing.drawCentredString(width / 2, y, line)
        elif alignment == 2:
            drawing.drawRightString(width - 4, y, line)
        else:
            drawing.drawString(4, y, line)
    drawing.showPage()
    drawing.save()
    buffer.seek(0)
    page = PdfReader(buffer).pages[0]
    resources = page["/Resources"].clone(writer)
    resources.setdefault(NameObject("/XObject"), DictionaryObject())[NameObject("/OriginalAppearance")] = original
    appearance = DecodedStreamObject()
    appearance.set_data(b"q /OriginalAppearance Do Q\n" + page.get_contents().get_data())
    appearance[NameObject("/Type")] = NameObject("/XObject")
    appearance[NameObject("/Subtype")] = NameObject("/Form")
    appearance[NameObject("/BBox")] = ArrayObject([FloatObject(0), FloatObject(0), FloatObject(width), FloatObject(height)])
    appearance[NameObject("/Resources")] = resources
    widget["/AP"][NameObject("/N")] = writer._add_object(appearance)
    # 폼 뷰어가 /V로 재렌더링해도 '?'가 겹쳐 보이지 않도록 실제 입력 글꼴도 내장함.
    resource_name = _pdf_cjk_form_font(writer, font, value)
    widget[NameObject("/DA")] = TextStringObject(f"{resource_name} {0 if automatic else size:g} Tf 0 g")
    _pdf_cjk_local_resources(writer, widget, resource_name)
    if widget.get("/Parent"):
        widget["/Parent"][NameObject("/DA")] = widget["/DA"]
        _pdf_cjk_local_resources(writer, widget["/Parent"], resource_name)


def _pdf_cjk_local_resources(writer, node, resource_name):
    # Local /DR takes precedence over AcroForm /DR. Clone the dictionaries:
    # source widgets can share them, including widgets outside the selected field.
    if node.get('/DR') is None:
        return
    original = node['/DR'].get_object()
    resources = DictionaryObject(dict(original))
    fonts = DictionaryObject(dict(original.get('/Font', DictionaryObject()).get_object()))
    fonts[resource_name] = writer._root_object['/AcroForm']['/DR']['/Font'].get(resource_name)
    resources[NameObject('/Font')] = writer._add_object(fonts)
    node[NameObject('/DR')] = writer._add_object(resources)


def _pdf_cjk_form_font(writer, font, value=''):
    face = pdfmetrics.getFont(font).face
    # Appearance regeneration encodes the embedded TTF's glyph IDs. Keep CID
    # identity and the reverse Unicode map consistent with those IDs.
    characters = {}
    for character, glyph in sorted(face.charToGlyph.items()):
        if 0 < glyph < 65536:
            characters.setdefault(glyph, character)
    if any(characters.get(face.charToGlyph.get(ord(character))) != ord(character)
           for character in value if character not in '\r\n\t'):
        raise TemplateError('PDF 글리프의 유니코드 역매핑이 원문과 일치하지 않아 편집 가능한 입력을 보류함')
    form = writer._root_object["/AcroForm"].get_object()
    resources = form.setdefault(NameObject("/DR"), DictionaryObject()).get_object()
    fonts = resources.setdefault(NameObject("/Font"), DictionaryObject()).get_object()
    resource_name = NameObject("/CompatCJK" + sha256(('gid-identity-1|' + font).encode()).hexdigest()[:12])
    if resource_name in fonts:
        return resource_name
    # A widget may share the global dictionaries. A global font addition must
    # not silently modify unselected widgets' local resources.
    resources = DictionaryObject(dict(resources))
    fonts = DictionaryObject(dict(fonts))
    resources[NameObject('/Font')] = writer._add_object(fonts)
    form[NameObject('/DR')] = writer._add_object(resources)
    font_file = DecodedStreamObject()
    font_file.set_data(Path(face.filename).read_bytes())
    font_file[NameObject("/Length1")] = NumberObject(len(font_file.get_data()))
    descriptor = DictionaryObject({NameObject("/Type"): NameObject("/FontDescriptor"),
        NameObject("/FontName"): resource_name, NameObject("/Flags"): NumberObject(32),
        NameObject("/FontBBox"): ArrayObject([FloatObject(value) for value in face.bbox]),
        NameObject("/ItalicAngle"): FloatObject(face.italicAngle), NameObject("/Ascent"): FloatObject(face.ascent),
        NameObject("/Descent"): FloatObject(face.descent), NameObject("/CapHeight"): FloatObject(face.capHeight),
        NameObject("/StemV"): NumberObject(80), NameObject("/FontFile2"): writer._add_object(font_file.flate_encode())})
    widths = ArrayObject()
    for glyph, character in sorted(characters.items()):
        widths.extend([NumberObject(glyph), ArrayObject([FloatObject(face.charWidths.get(character, 1000))])])
    descendant = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/CIDFontType2"),
        NameObject("/BaseFont"): resource_name, NameObject("/FontDescriptor"): writer._add_object(descriptor),
        NameObject("/CIDSystemInfo"): DictionaryObject({NameObject("/Registry"): TextStringObject("Adobe"),
            NameObject("/Ordering"): TextStringObject("Identity"), NameObject("/Supplement"): NumberObject(0)}),
        NameObject("/DW"): NumberObject(1000), NameObject("/W"): widths,
        NameObject("/CIDToGIDMap"): NameObject('/Identity')})
    unicode_map = DecodedStreamObject()
    entries = [f'<{glyph:04X}> <{chr(character).encode("utf-16-be").hex().upper()}>\n'
               for glyph, character in sorted(characters.items())]
    cmap = ('/CIDInit /ProcSet findresource begin 12 dict begin begincmap\n'
            '/CIDSystemInfo << /Registry (Adobe) /Ordering (UCS) /Supplement 0 >> def\n'
            '/CMapName /CompatIdentity def /CMapType 2 def\n'
            '1 begincodespacerange <0000> <FFFF> endcodespacerange\n')
    for start in range(0, len(entries), 100):
        block = entries[start:start + 100]
        cmap += f'{len(block)} beginbfchar\n' + ''.join(block) + 'endbfchar\n'
    unicode_map.set_data((cmap + 'endcmap CMapName currentdict /CMap defineresource pop end end').encode('ascii'))
    type0 = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type0"),
        NameObject("/BaseFont"): resource_name, NameObject("/Encoding"): NameObject("/Identity-H"),
        NameObject("/DescendantFonts"): ArrayObject([writer._add_object(descendant)]),
        NameObject("/ToUnicode"): writer._add_object(unicode_map.flate_encode())})
    fonts[resource_name] = writer._add_object(type0)
    return resource_name


def _pdf_font(field, value=""):
    configured = field.get("font_path")
    candidates = [Path(configured)] if configured else [Path("C:/Windows/Fonts/malgun.ttf"),
        Path("/usr/share/fonts/truetype/nanum/NanumGothic.ttf"),
        Path("/System/Library/Fonts/Supplemental/Arial Unicode.ttf"),
        Path(__file__).with_name("fonts") / "NotoSansKR-Regular.ttf"]
    characters = {ord(character) for character in value if character not in "\n\r\t"}
    missing = None
    for path in candidates:
        if not path.is_file():
            continue
        name = "FormFont" + sha256(str(path.resolve()).encode()).hexdigest()[:12]
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name, str(path)))
        missing = characters - pdfmetrics.getFont(name).face.charToGlyph.keys()
        if not missing:
            return name
    if missing:
        raise TemplateError("선택한 PDF 글꼴에 표시할 수 없는 문자가 있음: " + chr(min(missing)))
    raise TemplateError("PDF 입력에 사용할 한글 TTF 글꼴이 없음. 입력 영역에 font_path를 지정해야 함")


def _wrap_pdf_text(value, font, size, width):
    lines = []
    face = pdfmetrics.getFont(font).face
    for character in set(value) - {"\n", "\r", "\t"}:
        if ord(character) not in face.charToGlyph:
            raise TemplateError(f"선택한 PDF 글꼴에 표시할 수 없는 문자가 있음: {character}")
    for paragraph in value.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        line = ""
        for character in paragraph.replace("\t", "    "):
            if pdfmetrics.stringWidth(character, font, size) > width:
                raise TemplateError("PDF 입력 영역의 폭이 한 글자보다 좁음")
            if line and pdfmetrics.stringWidth(line + character, font, size) > width:
                lines.append(line)
                line = character
            else:
                line += character
        lines.append(line)
    return lines


def _pdf_overlay_fill(source, target, selected, profile, annex_entries=None):
    reader = PdfReader(source)
    writer = PdfWriter()
    writer.clone_document_from_reader(reader)
    pages = _pdf_pages(reader)
    # 프로파일을 전달한 호출도 좌표·범위를 다시 검증함.
    checked = _manual_pdf_fields([field for field, _ in selected], pages)
    by_page = {}
    for field, (_, value) in zip(checked, selected):
        size = float(field["font_size"])
        if not 1 <= size <= 100:
            raise TemplateError("PDF 글꼴 크기는 1~100pt이어야 함")
        font = _pdf_font(field, value)
        lines = _wrap_pdf_text(value, font, size, field["width"])
        line_height = size * 1.3
        if len(lines) * line_height > field["height"]:
            raise TemplateError(f"PDF 입력 영역에 내용이 넘침: {field['label']}. 영역을 늘리거나 분량을 줄여야 함")
        by_page.setdefault(field["page"], []).append((field, font, lines, line_height))
    for page_number, entries in by_page.items():
        page = writer.pages[page_number - 1]
        if page.rotation:
            page.transfer_rotation_to_content()
        width, height = float(page.mediabox.width), float(page.mediabox.height)
        buffer = BytesIO()
        overlay = canvas.Canvas(buffer, pagesize=(width, height), invariant=1)
        for field, font, lines, line_height in entries:
            overlay.setFont(font, field["font_size"])
            for index, line in enumerate(lines):
                overlay.drawString(field["x"], height - field["y"] - field["font_size"] - index * line_height, line)
        overlay.showPage()
        overlay.save()
        buffer.seek(0)
        overlay_page = PdfReader(buffer).pages[0]
        page.merge_translated_page(overlay_page, float(page.mediabox.left), float(page.mediabox.bottom))
    from .pdf_annex import append_annex
    append_annex(writer, annex_entries or [])
    target.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary = tempfile.mkstemp(dir=target.parent, suffix=".pdf")
    os.close(descriptor)
    try:
        with open(temporary, "wb") as stream:
            writer.write(stream)
        os.replace(temporary, target)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


def fill_compatible_template(template_path: str | Path, values: dict[str, str], output_path: str | Path,
                             mapping: dict[str, str] | None = None, profile: dict | None = None) -> Path:
    """mapping은 프로파일 field.id → values의 키임. 원본을 덮어쓰지 않음."""
    source, target = Path(template_path), Path(output_path)
    if source.resolve() == target.resolve() or source.suffix.lower() != target.suffix.lower():
        raise TemplateError("원본과 다른 경로에 같은 형식으로 저장해야 함")
    if not isinstance(values, dict) or any(not isinstance(k, str) or not isinstance(v, str) for k, v in values.items()):
        raise TemplateError("값은 dict[str, str]이어야 함")
    if any(re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", value) for value in values.values()):
        raise TemplateError("허용되지 않는 제어 문자가 포함됨")
    if mapping is not None and (not isinstance(mapping, dict) or any(not isinstance(key, str) or not isinstance(value, str) for key, value in mapping.items())):
        raise TemplateError("매핑은 입력칸 ID → 값의 키인 dict[str, str]이어야 함")
    profile = profile if profile is not None else analyze_template(source)
    if profile.get("source_sha256") and sha256(source.read_bytes()).hexdigest() != profile["source_sha256"]:
        raise TemplateError("양식 파일이 분석 이후 변경됨. 현재 양식을 다시 분석해야 함")
    if profile.get("format") != source.suffix.lower().lstrip("."):
        raise TemplateError("양식 프로파일 형식이 원본과 일치하지 않음")
    if not profile["supported"]:
        raise TemplateError(" / ".join(profile["warnings"]))
    try:
        rules_profile = mapped_rule_profile(profile, mapping)
        if profile.get('repeat_expansion'):
            from .repeat_rows import validate_repeat_fields
            validate_repeat_fields(source, rules_profile)
        literal_values = ({field['value_key']: values.get(field['value_key'], '')
                           for field in rules_profile.get('fields', [])
                           if field.get('control_type') in {'checkbox', 'radio', 'choice', 'combobox'}}
                          if profile.get('citation_mode') == 'sidecar' else None)
        value_issues = inspect_form_values(values, rules_profile, literal_values=literal_values)
    except ValueError as exc:
        raise TemplateError(f'양식 입력 규칙이 잘못됨: {exc}') from exc
    if value_issues:
        raise TemplateError('양식 입력값 검증 실패: ' + ' / '.join(f"{issue['field']}: {issue['message']}" for issue in value_issues))
    if profile["format"] == "pdf":
        reader = PdfReader(source)
        if reader.is_encrypted:
            raise TemplateError("암호화된 PDF는 변경할 수 없음")
        if _pdf_signed(reader):
            raise TemplateError("전자서명 또는 인증된 PDF는 변경할 수 없음")
    from .pdf_annex import plan_annex
    printed, annex_entries = plan_annex(profile, values, mapping)
    selected = _selected_fields(profile, printed, mapping)
    try:
        if profile["format"] in {"docx", "hwpx"}:
            _office_fill(source, target, selected)
        elif profile["format"] == "xlsx":
            _xlsx_fill(source, target, selected)
        elif profile["format"] == "pptx":
            from .pptx import fill_pptx_selected
            fill_pptx_selected(source, target, selected)
        elif profile.get("render_mode") == "overlay":
            _pdf_overlay_fill(source, target, selected, profile, annex_entries)
        else:
            _pdf_fill(source, target, selected)
    except TemplateError:
        raise
    except Exception as exc:
        raise TemplateError(f"양식 입력에 실패함: {exc}") from exc
    return target
