"""Offline UI gates plus real parser, filler and output SHA checks."""
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path

from docx import Document
import pytest
from streamlit.testing.v1 import AppTest
from app import ra_workflow_ui as ui
from templates.compatibility import analyze_template


@pytest.fixture
def registered(tmp_path,monkeypatch):
    monkeypatch.setattr(ui,'ROOT',tmp_path)
    template=tmp_path/'공식.docx'
    doc=Document();doc.add_paragraph('제품명: {{제품명}}');doc.add_paragraph('신청인: {{신청인}}');doc.save(template)
    profile=analyze_template(template)
    for field in profile['fields']:
        field['required']=False
        field['input_required']=field['value_key']=='신청인'
        field['narrative_style_required']=False
    profile.update(domain='pharmaceutical_ra',ra_workflow='product_registration',citation_mode='sidecar')
    p=tmp_path/'profile.json';p.write_text(json.dumps(profile,ensure_ascii=False),encoding='utf-8')
    entry={'id':'trial','title':'공식 시험 양식','source_path':template.name,'source_sha256':sha256(template.read_bytes()).hexdigest(),
           'profile_path':p.name,'profile_sha256':sha256(p.read_bytes()).hexdigest(),'coverage_note':'지정 입력칸 시험임',
           'attachments':[{'title':'첨부 증빙','condition':'해당하는 경우 확인', 'source_quote':'첨부 증빙 안내 원문',
                           'source_sha256':sha256(template.read_bytes()).hexdigest(), 'page':1}]}
    catalog=tmp_path/'catalog.json';catalog.write_text(json.dumps({'workflows':[entry]}),encoding='utf-8')
    monkeypatch.setattr(ui,'CATALOG',catalog)
    return entry,profile,template,p


def receipt(profile,*,blocking=False,missing=False):
    return {'ready_for_output_check':not blocking and not missing,'review':{'blocking':blocking,'issues':[]},
            'draft':{'제품명':'제품 A [S1]','신청인':'확인 사용자 [SU1]'},'template_profile':deepcopy(profile),
            'locked_fields':{'신청인':{'value':'확인 사용자','source_id':'SU1'}},'questions':['제품명 원자료를 제공해 주세요']*3 if missing else [],
            'sources':[],'mode':'verified_field_copy','actual_model_requests':0,'submission_ready':False}


def app():
    return AppTest.from_string('from app.ra_workflow_ui import render\nrender()',default_timeout=20).run()


def clean(at):
    assert not at.exception,[e.message for e in at.exception]


def test_ctd_workspace_is_available_from_ra_main_screen():
    screen = app()
    clean(screen)
    assert any('CTD Module 1·3' in tab.label for tab in screen.tabs)
    assert any('원자료의 정확한 제품명' in field.label for field in screen.text_input)
    assert not any('CTD 작업 초안 DOCX' in button.label for button in screen.get('download_button'))


def test_resolver_checks_both_original_and_profile_sha(registered):
    entry,profile,template,p=registered
    assert ui.resolve_workflow(entry)[1]==profile
    p.write_text('{}',encoding='utf-8')
    with pytest.raises(ValueError,match='매핑'):ui.resolve_workflow(entry)
    p.write_text(json.dumps(profile,ensure_ascii=False),encoding='utf-8')
    template.write_bytes(b'changed')
    with pytest.raises(ValueError,match='원본'):ui.resolve_workflow(entry)


def test_resolver_rejects_path_escape(registered):
    entry,*_=registered
    entry={**entry,'source_path':'../outside.docx'}
    with pytest.raises(ValueError,match='경로'):ui.resolve_workflow(entry)


def test_uploaded_docx_keeps_actual_sha_location_and_never_uses_ocr(monkeypatch):
    doc=Document();doc.add_paragraph('제품 A 함량 75 mg임');stream=BytesIO();doc.save(stream)
    class Upload:
        name='원자료.docx'
        def getvalue(self):return stream.getvalue()
    from llm.client import LLMClient
    def forbidden(*a,**k):raise AssertionError('No external AI in confirmed entry')
    monkeypatch.setattr(LLMClient,'generate_json',forbidden)
    monkeypatch.setattr(LLMClient,'read_image_json',forbidden)
    chunks=ui.parse_uploads([Upload()])
    assert chunks and chunks[0]['document_sha256']==sha256(stream.getvalue()).hexdigest()
    assert chunks[0]['filename']=='원자료.docx' and chunks[0]['location'] and chunks[0]['page'] is None


@pytest.mark.parametrize('names',[['../outside.txt'],['duplicate.txt','duplicate.txt']])
def test_uploaded_filename_path_and_duplicates_rejected(names):
    class Upload:
        def __init__(self,name):self.name=name
        def getvalue(self):return b'fixture'
    with pytest.raises(ValueError,match='파일명'):ui.parse_uploads([Upload(name) for name in names])


def test_real_filler_and_independent_output_sha_bound_to_current_source(registered,monkeypatch):
    entry,profile,template,_=registered
    original=template.read_bytes()
    calls=[]
    def prepare(p,s,**kwargs):calls.append(kwargs);return receipt(p)
    monkeypatch.setattr(ui,'prepare_ra_workflow',prepare)
    result=ui.export_workflow(entry,profile,[],{}, {})
    actual=Document(BytesIO(result['document']))
    assert actual.paragraphs[0].text=='제품명: 제품 A'
    evidence=json.loads(result['evidence'])
    assert evidence['draft']['제품명']=='제품 A [S1]' and not evidence['submission_ready']
    assert evidence['output_verification']['sha']['output']==sha256(result['document']).hexdigest()
    assert template.read_bytes()==original and calls==[{'source_bindings':{},'direct_values':{}}]


@pytest.mark.parametrize('blocking,missing',[(True,False),(False,True)])
def test_export_rechecks_current_values_and_blocks_errors(registered,monkeypatch,blocking,missing):
    entry,profile,*_=registered
    monkeypatch.setattr(ui,'prepare_ra_workflow',lambda p,*a,**k:receipt(p,blocking=blocking,missing=missing))
    def forbidden(*a,**k):raise AssertionError('Filler must not run')
    monkeypatch.setattr(ui,'fill_compatible_template',forbidden)
    with pytest.raises(ValueError,match='오류'):ui.export_workflow(entry,profile,[],{}, {})


def test_output_verification_failure_never_returns_download(registered,monkeypatch):
    entry,profile,*_=registered
    monkeypatch.setattr(ui,'prepare_ra_workflow',lambda p,*a,**k:receipt(p))
    monkeypatch.setattr(ui,'verify_output',lambda *a,**k:{'status':'failed'})
    with pytest.raises(ValueError,match='독립 검수'):ui.export_workflow(entry,profile,[],{}, {})


def test_changed_mapping_cannot_be_exported(registered):
    entry,profile,*_=registered
    changed=deepcopy(profile);changed['fields'][0]['value_key']='다른 값'
    with pytest.raises(ValueError,match='매핑'):ui.export_workflow(entry,changed,[],{}, {})


def test_ui_never_downloads_before_values_verification_and_confirmation(registered,monkeypatch):
    _,profile,*_=registered
    monkeypatch.setattr(ui,'prepare_ra_workflow',lambda p,*a,**k:receipt(p))
    at=app();clean(at)
    assert at.button(key='rw_export').disabled
    assert not any(item.proto.label=='작성 문서 다운로드' for item in at.get('download_button'))
    next(item for item in at.text_input if item.label == '신청인').set_value('직접 입력 신청인').run();clean(at)
    at.button(key='rw_prepare').click().run();clean(at)
    assert at.button(key='rw_export').disabled
    confirm=next(item for item in at.checkbox if '기입 항목과 직접 입력값' in item.label)
    confirm.check().run();clean(at)
    assert not at.button(key='rw_export').disabled


def test_ui_input_change_invalidates_old_result_and_download(registered,monkeypatch):
    _,profile,*_=registered
    monkeypatch.setattr(ui,'prepare_ra_workflow',lambda p,*a,**k:receipt(p))
    monkeypatch.setattr(ui,'export_workflow',lambda *a,**k:{'document':b'checked fixture','filename':'test.docx','evidence':b'{}'})
    at=app();next(item for item in at.text_input if item.label == '신청인').set_value('신청인').run()
    at.button(key='rw_prepare').click().run()
    next(item for item in at.checkbox if '기입 항목과 직접 입력값' in item.label).check().run()
    at.button(key='rw_export').click().run();clean(at)
    assert any(item.proto.label=='작성 문서 다운로드' for item in at.get('download_button'))
    next(item for item in at.text_input if item.label == '신청인').set_value('다른 신청인').run();clean(at)
    assert at.button(key='rw_export').disabled
    assert not any(item.proto.label=='작성 문서 다운로드' for item in at.get('download_button'))


def test_ui_missing_questions_max_two_and_blocks_download(registered,monkeypatch):
    _,profile,*_=registered
    monkeypatch.setattr(ui,'prepare_ra_workflow',lambda p,*a,**k:receipt(p,missing=True))
    at=app();next(item for item in at.text_input if item.label == '신청인').set_value('신청인').run();at.button(key='rw_prepare').click().run();clean(at)
    assert len([item for item in at.warning if '원자료를 제공' in item.value])==2
    assert at.button(key='rw_export').disabled


def test_ui_empty_values_cannot_prepare_download(registered):
    at=app();at.button(key='rw_prepare').click().run();clean(at)
    assert any('한 개 이상' in item.value for item in at.error)
    assert at.button(key='rw_export').disabled


def proposal_app(monkeypatch):
    from agent import ra_autofill
    sources = [{'source_id':'S1','filename':'원자료.txt','page':1,'location':'1쪽','text':'제품명: 제품 A'},
               {'source_id':'S2','filename':'다른자료.txt','page':2,'location':'2쪽','text':'제품명: 제품 B'}]
    proposal = {'fields':[{'value_key':'제품명','status':'proposed','binding':{'source_id':'S1','quote':'제품 A'},
                           'reason':'확인할 후보 1개임'}],
                'source_bindings':{'제품명':{'source_id':'S1','quote':'제품 A'}}, 'requires_confirmation':True}
    monkeypatch.setattr(ra_autofill,'propose_ra_bindings',lambda *a,**k:proposal)
    monkeypatch.setattr(ui, 'render_multimodal_upload', lambda *a, **k: {
        'sources': ui.st.session_state.get('rw_sources', []), 'changed': False, 'ready': True})
    at=app();at.session_state['rw_sources']=sources;at.run();clean(at)
    at.button(key='rw_suggest').click().run();clean(at)
    return at


def apply_proposal(at):
    next(item for item in at.multiselect if item.label=='확인해서 적용할 후보 항목').set_value(['제품명']).run()
    next(item for item in at.checkbox if item.label=='선택한 후보의 제품·항목·원문 위치를 확인함').check().run()
    at.button(key='rw_apply_suggestions').click().run();clean(at)


def test_batch_proposals_cannot_apply_until_explicit_selection_and_confirmation(registered,monkeypatch):
    _,profile,*_=registered
    calls=[]
    def prepare(*a,**k):calls.append(k);return receipt(profile)
    monkeypatch.setattr(ui,'prepare_ra_workflow',prepare)
    at=proposal_app(monkeypatch)
    assert at.button(key='rw_apply_suggestions').disabled and not calls
    next(item for item in at.multiselect if item.label=='확인해서 적용할 후보 항목').set_value(['제품명']).run()
    assert at.button(key='rw_apply_suggestions').disabled
    at.button(key='rw_prepare').click().run()
    assert not calls and at.button(key='rw_export').disabled
    apply_proposal(at);at.button(key='rw_prepare').click().run();clean(at)
    assert calls[-1]['source_bindings']=={'제품명':{'source_id':'S1','quote':'제품 A'}}
    assert at.button(key='rw_export').disabled


@pytest.mark.parametrize('changed_key',['rw_product','rw_variant'])
def test_changed_product_or_variant_discards_confirmed_batch_and_result(registered,monkeypatch,changed_key):
    _,profile,*_=registered
    calls=[]
    def prepare(*a,**k):calls.append(k);return receipt(profile)
    monkeypatch.setattr(ui,'prepare_ra_workflow',prepare)
    at=proposal_app(monkeypatch);apply_proposal(at);at.button(key='rw_prepare').click().run()
    assert len(calls)==1
    at.text_input(key=changed_key).set_value('선택이 변경됨').run();clean(at)
    at.button(key='rw_prepare').click().run();clean(at)
    assert len(calls)==1 and at.button(key='rw_export').disabled
    assert not at.session_state.get('rw_confirmed_suggestions',{})


def test_unconfirmed_manual_edit_replaces_old_confirmed_batch_binding(registered,monkeypatch):
    _,profile,*_=registered
    calls=[]
    def prepare(*a,**k):calls.append(k);return receipt(profile)
    monkeypatch.setattr(ui,'prepare_ra_workflow',prepare)
    at=proposal_app(monkeypatch);apply_proposal(at)
    next(item for item in at.selectbox if item.label=='사용할 원문 조각').set_value('S2').run()
    at.button(key='rw_prepare').click().run();clean(at)
    assert not calls and at.button(key='rw_export').disabled
    next(item for item in at.checkbox if item.label=='이 문구의 항목·제품·범위와 기입 위치를 확인함').check().run()
    at.button(key='rw_prepare').click().run();clean(at)
    assert calls[-1]['source_bindings']['제품명']=={'source_id':'S2','quote':'제품명: 제품 B'}


def test_manual_source_ranking_preserves_access_to_all_pages():
    sources=[{'source_id':f'S{i}','filename':'원자료.txt','text':f'문구 {i}', 'location':f'문단 {i}'} for i in range(75)]
    ranked=ui.rank_field_sources({'label':'제품명'},sources)
    assert len(ranked)==75 and {source['source_id'] for source in ranked}=={source['source_id'] for source in sources}


@pytest.mark.parametrize('changed_key',['rw_product','rw_variant'])
def test_attachment_confirmation_does_not_carry_into_a_different_product(registered,changed_key):
    at=app();clean(at)
    next(item for item in at.selectbox if item.label=='담당자 확인 상태').set_value('checked').run()
    assert next(item for item in at.selectbox if item.label=='담당자 확인 상태').value=='checked'
    at.text_input(key=changed_key).set_value('다른 제품 범위').run();clean(at)
    assert next(item for item in at.selectbox if item.label=='담당자 확인 상태').value=='pending'


def test_attachment_confirmation_does_not_carry_into_changed_sources(registered, monkeypatch):
    monkeypatch.setattr(ui, 'render_multimodal_upload', lambda *a, **k: {
        'sources': ui.st.session_state.get('rw_sources', []), 'changed': False, 'ready': True})
    at=app();next(item for item in at.selectbox if item.label=='담당자 확인 상태').set_value('not_applicable').run()
    at.session_state['rw_sources']=[{'source_id':'S9','filename':'새첨부.txt','text':'새 원자료'}]
    at.run();clean(at)
    assert next(item for item in at.selectbox if item.label=='담당자 확인 상태').value=='pending'


def test_multimodal_pending_source_blocks_form_preparation_and_old_download(registered, monkeypatch):
    monkeypatch.setattr(ui, 'render_multimodal_upload', lambda *a, **k: {
        'sources': [], 'changed': True, 'ready': False})
    at = app(); clean(at)
    assert any('미확인 이미지' in item.value for item in at.warning)
    assert not any(button.key == 'rw_prepare' for button in at.button)
    assert not any(item.proto.label == '작성 문서 다운로드' for item in at.get('download_button'))


def test_export_revalidates_confirmed_image_receipt_before_preparation(registered, tmp_path, monkeypatch):
    from PIL import Image
    from agent.multimodal_intake import collect_multimodal, confirm_intake, generation_sources
    entry, profile, *_ = registered
    path = tmp_path / '실제원본.png'
    Image.new('RGB', (20, 20), 'white').save(path)
    digest = sha256(path.read_bytes()).hexdigest()
    intake = collect_multimodal([path], transcriptions={digest: [{'page': 1, 'text': '제품명: 시험 제품'}]})
    source = intake['sources'][0]
    checked = confirm_intake(intake, [{'source_id': source['source_id'],
                                      'fingerprint': source['verification_fingerprint']}], confirmed=True)
    tampered = generation_sources(checked)[0]
    tampered['text'] = '제품명: 다른 제품'
    def forbidden(*args, **kwargs):
        raise AssertionError('Changed source must fail before prepare/fill')
    monkeypatch.setattr(ui, 'prepare_ra_workflow', forbidden)
    with pytest.raises(ValueError):
        ui.export_workflow(entry, profile, [tampered], {}, {})


def test_unchecking_final_confirmation_removes_existing_downloads(registered, monkeypatch):
    _, profile, *_ = registered
    monkeypatch.setattr(ui, 'prepare_ra_workflow', lambda p, *a, **k: receipt(p))
    monkeypatch.setattr(ui, 'export_workflow', lambda *a, **k: {
        'document': b'checked fixture', 'filename': 'test.docx', 'evidence': b'{}'})
    at = app()
    at.text_input(key='rw_direct_' + sha256(('trial' + profile['fields'][1]['id']).encode()).hexdigest()[:16]).set_value('신청인').run()
    at.button(key='rw_prepare').click().run()
    confirmation = next(item for item in at.checkbox if '기입 항목과 직접 입력값' in item.label)
    confirmation.check().run()
    at.button(key='rw_export').click().run(); clean(at)
    assert any(item.proto.label == '작성 문서 다운로드' for item in at.get('download_button'))
    next(item for item in at.checkbox if '기입 항목과 직접 입력값' in item.label).uncheck().run(); clean(at)
    assert at.button(key='rw_export').disabled
    assert 'rw_exports' not in at.session_state
    assert not any(item.proto.label == '작성 문서 다운로드' for item in at.get('download_button'))
