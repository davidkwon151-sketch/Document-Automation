from pathlib import Path
from zipfile import ZipFile

from docx import Document
from lxml import etree
from openpyxl import load_workbook
from pypdf import PdfReader, PdfWriter
from pypdf.generic import NameObject
import pytest

from agent.output_check import verify_output
from samples.generate import VALUES
from templates import analyze_template, fill_compatible_template


ROOT = Path(__file__).resolve().parents[1]


def rewrite_part(path, name, transform):
    with ZipFile(path) as archive:
        parts = {item.filename: (item, archive.read(item.filename)) for item in archive.infolist()}
    with ZipFile(path, "w") as archive:
        for key, (item, data) in parts.items():
            archive.writestr(item, transform(data) if key == name else data)


@pytest.mark.parametrize("relative", ["templates/result_report.docx", "templates/result_report.hwpx", *[
    f"samples/sample_company_form.{suffix}" for suffix in ("docx", "hwpx", "xlsx", "pdf", "pptx")]])
def test_saved_output_reopens_and_preserves_original(tmp_path, relative):
    source = ROOT / relative
    original = source.read_bytes()
    output = fill_compatible_template(source, VALUES, tmp_path / source.name)
    result = verify_output(source, output, VALUES)
    assert result["status"] == "passed"
    assert len(result["sha"]["output"]) == 64
    assert result["native_visual_qa"] == "pending"
    assert result["legal_compliance_certified"] is False
    assert source.read_bytes() == original


def test_profile_hash_mismatch_and_source_overwrite_are_blocked(tmp_path):
    source = ROOT / "templates/result_report.docx"
    output = fill_compatible_template(source, VALUES, tmp_path / source.name)
    profile = analyze_template(source) | {"source_sha256": "0" * 64}
    with pytest.raises(ValueError, match="SHA-256"):
        verify_output(source, output, VALUES, profile=profile)
    with pytest.raises(ValueError, match="덮어쓸"):
        verify_output(source, source, VALUES)


@pytest.mark.parametrize("part", ["word/styles.xml", "word/fontTable.xml", "word/theme/theme1.xml"])
def test_protected_zip_assets_are_checked_independently(tmp_path, part):
    source = ROOT / "templates/result_report.docx"
    output = fill_compatible_template(source, VALUES, tmp_path / source.name)
    rewrite_part(output, part, lambda data: data.replace(b"Calibri", b"Changed") + b" ")
    with pytest.raises(ValueError, match="부품 변경"):
        verify_output(source, output, VALUES)


def test_zip_member_loss_is_blocked(tmp_path):
    source = ROOT / "templates/result_report.docx"
    output = fill_compatible_template(source, VALUES, tmp_path / source.name)
    with ZipFile(output) as archive:
        parts = {name: archive.read(name) for name in archive.namelist() if name != "word/fontTable.xml"}
    with ZipFile(output, "w") as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    with pytest.raises(ValueError, match="부품 손실"):
        verify_output(source, output, VALUES)


@pytest.mark.parametrize("replacement,pattern", [("망가진 값", "입력값|문단 값"), ("{{제목}}", "자리표시자|문단 값")])
def test_missing_or_unreplaced_saved_value_is_blocked(tmp_path, replacement, pattern):
    source = ROOT / "templates/result_report.docx"
    output = fill_compatible_template(source, VALUES, tmp_path / source.name)
    rewrite_part(output, "word/document.xml", lambda data: data.replace(VALUES["제목"].encode(), replacement.encode()))
    with pytest.raises(ValueError, match=pattern):
        verify_output(source, output, VALUES)


def test_user_literal_placeholder_is_not_an_unresolved_template_token(tmp_path):
    source = ROOT / "templates/result_report.docx"
    values = VALUES | {"제목": "{{요약}} 문자를 그대로 쓰는 제목"}
    output = fill_compatible_template(source, values, tmp_path / source.name)
    assert verify_output(source, output, values)["status"] == "passed"


@pytest.mark.parametrize("suffix", ["docx", "hwpx"])
def test_placeholder_values_in_wrong_paragraph_are_rejected(tmp_path, suffix):
    source = ROOT / f"templates/result_report.{suffix}"
    values = VALUES | {"제목": "서로 다른 제목", "요약": "서로 다른 요약"}
    output = fill_compatible_template(source, values, tmp_path / source.name)
    part = "word/document.xml" if suffix == "docx" else "Contents/section0.xml"
    def swap(data):
        return data.replace(values["제목"].encode(), b"SWAP_SENTINEL").replace(values["요약"].encode(), values["제목"].encode()).replace(b"SWAP_SENTINEL", values["요약"].encode())
    rewrite_part(output, part, swap)
    with pytest.raises(ValueError, match="문단 값.*불일치"):
        verify_output(source, output, values)


@pytest.mark.parametrize("changed", ["prefix", "suffix", "swap"])
def test_multiple_placeholders_and_surrounding_literals_stay_in_same_paragraph(tmp_path, changed):
    document = Document()
    document.add_paragraph("앞 문구 {{제목}} / {{요약}} 뒤 문구")
    document.add_paragraph("{{제목}}")
    source = tmp_path / "multiple.docx"
    document.save(source)
    values = {"제목": "제목값", "요약": "요약값"}
    output = fill_compatible_template(source, values, tmp_path / "filled.docx")
    assert verify_output(source, output, values)["status"] == "passed"
    def mutate(data):
        if changed == "swap":
            return data.replace("제목값 / 요약값".encode(), "요약값 / 제목값".encode())
        original = "앞 문구" if changed == "prefix" else "뒤 문구"
        return data.replace(original.encode(), "바뀐 문구".encode())
    rewrite_part(output, "word/document.xml", mutate)
    with pytest.raises(ValueError, match="문단 값.*불일치"):
        verify_output(source, output, values)


def test_mapped_cell_value_in_other_cell_is_rejected(tmp_path):
    source = ROOT / "samples/sample_company_form.docx"
    values = VALUES | {"제목": "제목값", "요약": "요약값"}
    output = fill_compatible_template(source, values, tmp_path / source.name)
    rewrite_part(output, "word/document.xml", lambda data: data.replace("제목값".encode(), b"TEMP").replace("요약값".encode(), "제목값".encode()).replace(b"TEMP", "요약값".encode()))
    with pytest.raises(ValueError, match="입력값|첫 문단"):
        verify_output(source, output, values)


def test_docx_original_cell_and_font_properties_cannot_change(tmp_path):
    source = ROOT / "templates/result_report.docx"
    output = fill_compatible_template(source, VALUES, tmp_path / source.name)
    def mutate(data):
        root = etree.fromstring(data)
        node = root.xpath(".//*[local-name()='tcW']")[0]
        node.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}w", "1")
        return etree.tostring(root)
    rewrite_part(output, "word/document.xml", mutate)
    with pytest.raises(ValueError, match="서식 변경"):
        verify_output(source, output, VALUES)


def test_unrelated_docx_form_label_change_is_detected(tmp_path):
    source = ROOT / "templates/result_report.docx"
    output = fill_compatible_template(source, VALUES, tmp_path / source.name)
    rewrite_part(output, "word/document.xml", lambda data: data.replace("핵심 요약".encode(), "다른 항목".encode()))
    with pytest.raises(ValueError, match="항목명"):
        verify_output(source, output, VALUES)


def test_xlsx_formula_and_merge_mutations_are_detected(tmp_path):
    source = ROOT / "samples/sample_company_form.xlsx"
    output = fill_compatible_template(source, VALUES, tmp_path / source.name)
    rewrite_part(output, "xl/worksheets/sheet1.xml", lambda data: data.replace(b"<f>1+2</f>", b"<f>9+9</f>"))
    with pytest.raises(ValueError, match="수식"):
        verify_output(source, output, VALUES)


def test_xlsx_column_width_changes_are_detected(tmp_path):
    source = ROOT / "samples/sample_company_form.xlsx"
    output = fill_compatible_template(source, VALUES, tmp_path / source.name)
    def mutate(data):
        root = etree.fromstring(data)
        root.xpath(".//*[local-name()='col']")[0].set("width", "1")
        return etree.tostring(root)
    rewrite_part(output, "xl/worksheets/sheet1.xml", mutate)
    with pytest.raises(ValueError, match="열 너비"):
        verify_output(source, output, VALUES)


def test_pdf_canonical_value_is_reopened_and_checked(tmp_path):
    source = ROOT / "samples/sample_company_form.pdf"
    output = fill_compatible_template(source, VALUES, tmp_path / source.name)
    writer = PdfWriter()
    writer.clone_document_from_reader(PdfReader(output))
    writer.update_page_form_field_values(None, {"title": "wrong title"})
    with output.open("wb") as stream:
        writer.write(stream)
    with pytest.raises(ValueError, match="canonical /V"):
        verify_output(source, output, VALUES)


def test_pdf_korean_overlay_values_and_visual_presence(tmp_path):
    source = ROOT / "samples/sample_static_form.pdf"
    fields = [{"label": key, "value_key": key, "page": 1, "x": 55, "y": y,
               "width": 470, "height": height, "font_size": 10} for key, y, height in (
                   ("제목", 95, 25), ("요약", 165, 70), ("본문", 305, 150))]
    profile = analyze_template(source, fields) | {"resource_kind": "layout_reference"}
    output = fill_compatible_template(source, VALUES, tmp_path / source.name, profile=profile)
    report = verify_output(source, output, VALUES, profile=profile)
    assert report["resource_kind"] == "layout_reference"
    assert any(item["name"] == "pdfium_visual_presence" for item in report["checks"])


def test_pdf_checkbox_widget_state_is_checked(tmp_path):
    source = ROOT / "samples/sample_selection_form.pdf"
    profile = analyze_template(source)
    checkbox = next(field for field in profile["fields"] if field.get("control_type") == "checkbox")
    values = {"Consent": "true", "Decision": "/reject", "Department": "Operations"}
    mapping = None
    output = fill_compatible_template(source, values, tmp_path / source.name, profile=profile, mapping=mapping)
    writer = PdfWriter()
    writer.clone_document_from_reader(PdfReader(output))
    for reference in writer.pages[0]["/Annots"]:
        widget = reference.get_object()
        if widget.get("/T") == checkbox["id"][4:]:
            widget[NameObject("/AS")] = NameObject("/Off")
    with output.open("wb") as stream:
        writer.write(stream)
    with pytest.raises(ValueError, match="표시 상태"):
        verify_output(source, output, values, profile=profile, mapping=mapping)


def test_obvious_xlsx_fixed_row_height_overflow_is_blocked(tmp_path):
    source = tmp_path / "fixed.xlsx"
    workbook = load_workbook(ROOT / "samples/sample_company_form.xlsx")
    workbook.active.row_dimensions[3].height = 15
    workbook.save(source)
    workbook.close()
    output = fill_compatible_template(source, VALUES, tmp_path / "filled.xlsx")
    with pytest.raises(ValueError, match="고정 행 높이"):
        verify_output(source, output, VALUES)


def test_obvious_docx_fixed_row_height_overflow_is_blocked(tmp_path):
    from docx.enum.table import WD_ROW_HEIGHT_RULE
    from docx.shared import Pt
    source = tmp_path / "fixed.docx"
    document = Document(ROOT / "samples/sample_company_form.docx")
    document.tables[0].rows[2].height = Pt(10)
    document.tables[0].rows[2].height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    document.save(source)
    output = fill_compatible_template(source, VALUES, tmp_path / "filled.docx")
    with pytest.raises(ValueError, match="고정 행 높이"):
        verify_output(source, output, VALUES)


def test_pptx_cell_and_shape_values_preserve_split_runs_and_newlines(tmp_path):
    source = ROOT / "samples/sample_company_form.pptx"
    values = VALUES | {"작성자": "시험홍길동", "검토 메모": "□ 검토함\n○ 확인함"}
    output = fill_compatible_template(source, values, tmp_path / source.name)
    assert verify_output(source, output, values)["status"] == "passed"
    with ZipFile(source) as old, ZipFile(output) as new:
        assert old.read("ppt/slides/slide2.xml") == new.read("ppt/slides/slide2.xml")
        assert b"<a:br" in new.read("ppt/slides/slide1.xml")


@pytest.mark.parametrize("part", ["ppt/theme/theme1.xml", "ppt/presentation.xml", "ppt/slides/_rels/slide1.xml.rels"])
def test_pptx_theme_slide_order_and_relationships_cannot_change(tmp_path, part):
    source = ROOT / "samples/sample_company_form.pptx"
    output = fill_compatible_template(source, VALUES, tmp_path / source.name)
    rewrite_part(output, part, lambda data: data + b" ")
    with pytest.raises(ValueError, match="부품 변경"):
        verify_output(source, output, VALUES)


@pytest.mark.parametrize("mutate_kind", ["other_slide", "shape_position", "table_merge", "font"])
def test_pptx_non_target_content_geometry_merges_and_fonts_are_preserved(tmp_path, mutate_kind):
    source = ROOT / "samples/sample_company_form.pptx"
    output = fill_compatible_template(source, VALUES, tmp_path / source.name)
    part = "ppt/slides/slide2.xml" if mutate_kind == "other_slide" else "ppt/slides/slide1.xml"
    def mutate(data):
        root = etree.fromstring(data)
        if mutate_kind == "other_slide":
            node = root.xpath(".//*[local-name()='t']")[0]
            node.text = "임의 변경"
        elif mutate_kind == "shape_position":
            root.xpath(".//*[local-name()='off']")[0].set("x", "1")
        elif mutate_kind == "table_merge":
            root.xpath(".//*[local-name()='tc']")[0].set("gridSpan", "2")
        else:
            root.xpath(".//*[local-name()='rPr']")[0].set("sz", "100")
        return etree.tostring(root)
    rewrite_part(output, part, mutate)
    with pytest.raises(ValueError, match="PPTX.*변경"):
        verify_output(source, output, VALUES)


def test_pptx_values_in_another_placeholder_location_are_rejected(tmp_path):
    source = ROOT / "samples/sample_company_form.pptx"
    values = VALUES | {"제목": "제목값", "요약": "요약값"}
    output = fill_compatible_template(source, values, tmp_path / source.name)
    rewrite_part(output, "ppt/slides/slide1.xml", lambda data: data.replace("제목값".encode(), b"TEMP").replace("요약값".encode(), "제목값".encode()).replace(b"TEMP", "요약값".encode()))
    with pytest.raises(ValueError, match="문단 값.*불일치"):
        verify_output(source, output, values)


@pytest.mark.parametrize('suffix',['docx','hwpx'])
def test_static_form_labels_cannot_swap_positions_even_with_same_text_counts(tmp_path,suffix):
    source=ROOT/f'templates/result_report.{suffix}'
    output=fill_compatible_template(source,VALUES,tmp_path/source.name)
    part='word/document.xml' if suffix=='docx' else 'Contents/section0.xml'
    def swap(data):
        root=etree.fromstring(data)
        nodes=[]
        for paragraph in root.xpath(".//*[local-name()='p']"):
            texts=paragraph.xpath(".//*[local-name()='t']")
            if len(texts)==1 and texts[0].text and texts[0].text.strip() and '{{' not in texts[0].text:
                nodes.append(texts[0])
        first,second=nodes[:2]
        first.text,second.text=second.text,first.text
        return etree.tostring(root)
    rewrite_part(output,part,swap)
    with pytest.raises(ValueError,match='항목명|문단 값'):
        verify_output(source,output,VALUES)


def _wrap_workbook(path,*,merged=False,height=15):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment,Font
    workbook=Workbook()
    sheet=workbook.active
    sheet['A1']='본문'
    if merged:
        sheet.merge_cells('B1:E1')
    sheet['B1'].alignment=Alignment(wrap_text=True)
    sheet['B1'].font=Font(name='맑은 고딕',size=11)
    for column in 'BCDE':
        sheet.column_dimensions[column].width=12
    if height is not None:
        sheet.row_dimensions[1].height=height
    workbook.save(path)
    workbook.close()


def test_xlsx_wrapped_single_line_long_value_in_fixed_short_row_is_blocked(tmp_path):
    source=tmp_path/'narrow.xlsx'
    _wrap_workbook(source)
    values={'본문':'가'*200}
    output=fill_compatible_template(source,values,tmp_path/'filled.xlsx')
    with pytest.raises(ValueError,match='고정 행 높이'):
        verify_output(source,output,values)


def test_xlsx_wrapped_text_uses_merged_column_width_and_allows_short_values(tmp_path):
    source=tmp_path/'merged.xlsx'
    _wrap_workbook(source,merged=True,height=30)
    values={'본문':'가'*38}
    output=fill_compatible_template(source,values,tmp_path/'filled.xlsx')
    assert verify_output(source,output,values)['status']=='passed'


def test_xlsx_auto_height_is_pending_rather_than_claimed_layout_pass(tmp_path):
    source=tmp_path/'auto.xlsx'
    _wrap_workbook(source,height=None)
    values={'본문':'짧은 시험값'}
    output=fill_compatible_template(source,values,tmp_path/'filled.xlsx')
    report=verify_output(source,output,values)
    assert report['status']=='passed'
    assert next(check for check in report['checks'] if check['name']=='obvious_fixed_height_overflow')['status']=='pending'


@pytest.mark.parametrize('mutation',[None,'label','instruction','enabled'])
def test_legacy_formtext_values_and_unchanged_field_instructions_are_checked(tmp_path,mutation):
    namespace='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
    q=lambda name:f'{{{namespace}}}{name}'
    document=Document()
    paragraph=document.add_paragraph('Name: ')
    begin=etree.SubElement(paragraph.add_run()._r,q('fldChar'))
    begin.set(q('fldCharType'),'begin')
    data=etree.SubElement(begin,q('ffData'))
    etree.SubElement(data,q('name')).set(q('val'),'Applicant')
    etree.SubElement(data,q('enabled')).set(q('val'),'1')
    etree.SubElement(data,q('textInput'))
    etree.SubElement(paragraph.add_run()._r,q('instrText')).text=' FORMTEXT '
    etree.SubElement(paragraph.add_run()._r,q('fldChar')).set(q('fldCharType'),'separate')
    paragraph.add_run('________')
    etree.SubElement(paragraph.add_run()._r,q('fldChar')).set(q('fldCharType'),'end')
    paragraph.add_run(' / user supplied')
    source=tmp_path/'legacy.docx'
    document.save(source)
    profile=analyze_template(source)
    field=next(item for item in profile['fields'] if item['kind']=='docx_legacy_text')
    values={'name':'TEST USER'}
    mapping={field['id']:'name'}
    output=fill_compatible_template(source,values,tmp_path/'filled.docx',profile=profile,mapping=mapping)
    if mutation:
        def mutate(raw):
            root=etree.fromstring(raw)
            if mutation=='label':
                root.xpath(".//*[local-name()='t']")[0].text='Wrong label: '
            elif mutation=='instruction':
                root.xpath(".//*[local-name()='instrText']")[0].text=' FORMULA '
            else:
                root.xpath(".//*[local-name()='enabled']")[0].set(q('val'),'0')
            return etree.tostring(root)
        rewrite_part(output,'word/document.xml',mutate)
        with pytest.raises(ValueError,match='FORMTEXT'):
            verify_output(source,output,values,profile=profile,mapping=mapping)
    else:
        assert verify_output(source,output,values,profile=profile,mapping=mapping)['status']=='passed'


def test_docx_font_properties_cannot_swap_between_existing_runs(tmp_path):
    document=Document()
    paragraph=document.add_paragraph()
    paragraph.add_run('Title: ').bold=False
    paragraph.add_run('{{제목}}').bold=True
    source=tmp_path/'fonts.docx'
    document.save(source)
    values={'제목':'검증 제목'}
    output=fill_compatible_template(source,values,tmp_path/'filled.docx')
    assert verify_output(source,output,values)['status']=='passed'
    def swap(raw):
        root=etree.fromstring(raw)
        runs=root.xpath(".//*[local-name()='p']/*[local-name()='r']")
        first,second=runs[:2]
        one=first.xpath("./*[local-name()='rPr']")[0]
        two=second.xpath("./*[local-name()='rPr']")[0]
        first.remove(one);second.remove(two)
        first.insert(0,two);second.insert(0,one)
        return etree.tostring(root)
    rewrite_part(output,'word/document.xml',swap)
    with pytest.raises(ValueError,match='런 위치'):
        verify_output(source,output,values)


def test_pptx_bold_properties_cannot_swap_between_label_and_value_runs(tmp_path):
    source=ROOT/'samples/sample_company_form.pptx'
    output=fill_compatible_template(source,VALUES,tmp_path/source.name)
    def swap(raw):
        root=etree.fromstring(raw)
        namespace={'p':'http://schemas.openxmlformats.org/presentationml/2006/main','a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
        runs=root.xpath('//p:sp[2]/p:txBody/a:p/a:r',namespaces=namespace)
        first=runs[0]
        second=next(run for run in runs if run.find('{'+namespace['a']+'}rPr').get('b')=='1')
        one=first.find('{'+namespace['a']+'}rPr');two=second.find('{'+namespace['a']+'}rPr')
        first.remove(one);second.remove(two);first.insert(0,two);second.insert(0,one)
        return etree.tostring(root)
    rewrite_part(output,'ppt/slides/slide1.xml',swap)
    with pytest.raises(ValueError,match='런 위치'):
        verify_output(source,output,VALUES)


def test_hwpx_existing_character_style_reference_is_bound_to_run_position(tmp_path):
    source=ROOT/'templates/result_report.hwpx'
    output=fill_compatible_template(source,VALUES,tmp_path/source.name)
    def mutate(raw):
        root=etree.fromstring(raw)
        run=root.xpath(".//*[local-name()='run']")[0]
        run.set('charPrIDRef','999')
        return etree.tostring(root)
    rewrite_part(output,'Contents/section0.xml',mutate)
    with pytest.raises(ValueError,match='런 위치|위치 속성'):
        verify_output(source,output,VALUES)


def test_pptx_new_multiline_run_in_previously_blank_shape_inherits_original_font(tmp_path):
    source=ROOT/'samples/sample_company_form.pptx'
    profile=analyze_template(source)
    field=next(item for item in profile['fields'] if item['kind']=='pptx_text')
    values=VALUES|{field['value_key']:'□ 첫 줄을 확인함\n○ 둘째 줄을 검토함'}
    output=fill_compatible_template(source,values,tmp_path/source.name,profile=profile)
    assert verify_output(source,output,values,profile=profile)['status']=='passed'
