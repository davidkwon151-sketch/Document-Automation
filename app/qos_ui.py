"""DMF originals to CTD 2.3.S source-bound working document."""

from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import tempfile

from docx import Document
import streamlit as st

from agent.ctd_qos import QOS_S_SECTIONS, prepare_qos_package, propose_dmf_options
from agent.ctd_template import ctd_template_section_ids, fill_ctd_working_template
from app.multimodal_ui import render_multimodal_upload
from templates.compatibility import analyze_template


def build_qos_template(section_ids):
    """A clearly labelled project example, with only currently evidenced slots."""
    catalogue = {item['section_id']: item for item in QOS_S_SECTIONS}
    if (not isinstance(section_ids, list) or not section_ids
            or len(section_ids) != len(set(section_ids))
            or any(value not in catalogue for value in section_ids)):
        raise ValueError('기입할 2.3.S 절이 필요함')
    document = Document()
    document.styles['Normal'].font.name = '맑은 고딕'
    document.add_heading('CTD 2.3.S 원료의약품 품질요약 · 검토용', 0)
    document.add_paragraph('프로젝트 제작 예제 양식 · 원문 발췌 기반 작업 초안 · 제출용 완성본 아님')
    table = document.add_table(rows=0, cols=2)
    for identifier in section_ids:
        cells = table.add_row().cells
        cells[0].text = identifier + ' ' + catalogue[identifier]['title']
        cells[1].text = '{{' + identifier + '}}'
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def render_qos():
    st.subheader('제조원 DMF → CTD 2.3.S 원료의약품 품질요약')
    st.caption('PDF·DOCX·HWPX·XLSX·TXT·사진/스캔의 확인된 3.2.S 원문을 해당 2.3.S 절에 자동 옮깁니다. '
               '사진·스캔은 원본 대조 확인이 필요합니다. 출력은 발췌 기반 검토용 초안입니다.')
    product = st.text_input('선택 완제 품목명', key='qos_product').strip()
    variant = st.text_input('완제 제형·함량 (해당 시)', key='qos_variant').strip()
    section_ids = [item['section_id'] for item in QOS_S_SECTIONS]
    selected = st.multiselect('작성할 2.3.S 절', section_ids, default=section_ids,
                              key='qos_sections')
    uploaded = st.file_uploader('기입할 사내/업무용 2.3.S 양식 (선택)',
                                type=['docx', 'hwpx'], key='qos_form')
    form_bytes = uploaded.getvalue() if uploaded else None
    form_sha = sha256(form_bytes).hexdigest() if form_bytes else ''
    intake = render_multimodal_upload(key_prefix='qos', context={
        'product': product, 'variant': variant, 'selected': selected, 'form_sha': form_sha})
    if intake['changed']:
        st.session_state.pop('qos_package', None)
        st.session_state.pop('qos_output', None)
    try:
        options = propose_dmf_options(intake['sources']) if intake['sources'] else []
    except ValueError as exc:
        st.error(str(exc))
        return
    by_sha = {item['document_sha256']: item for item in options
              if item['substance_names'] and item['manufacturer_names']}
    chosen = st.multiselect('3.2.S 원문과 원료명·제조원 선언이 있는 DMF 원본',
                            list(by_sha), format_func=lambda digest:
                            by_sha[digest]['filename'] + ' · SHA ' + digest[:12],
                            key='qos_dmf_sha')
    substances = {item['value'] for digest in chosen
                  for item in by_sha[digest]['substance_names']}
    manufacturers = {item['value'] for digest in chosen
                     for item in by_sha[digest]['manufacturer_names']}
    substance = st.selectbox('DMF 원문에 선언된 원료명', [''] + sorted(substances),
                             key='qos_substance') if chosen else ''
    manufacturer = st.selectbox('DMF 원문에 선언된 원료 제조원',
                                [''] + sorted(manufacturers), key='qos_manufacturer') if chosen else ''
    for digest in chosen:
        with st.expander(by_sha[digest]['filename'] + ' · 원료/제조원 선언 근거'):
            for kind in ('substance_names', 'manufacturer_names'):
                for option in by_sha[digest][kind]:
                    for evidence in option['evidence']:
                        st.caption(f"{evidence['source_id']} · {evidence.get('page') or evidence.get('sheet') or '본문'}"
                                   f" · SHA {digest}")
                        st.code(evidence['quote'], language=None)
    confirm_key = sha256(json.dumps({'source': intake['signature'], 'chosen': chosen,
                                     'product': product, 'substance': substance,
                                     'manufacturer': manufacturer}, ensure_ascii=False,
                                    sort_keys=True).encode()).hexdigest()[:20]
    confirmed = st.checkbox('원료명·제조원 원문과 선택 완제의 실제 원료 관계를 대조했음',
                            key='qos_confirm_' + confirm_key)
    can_prepare = bool(product and selected and chosen and substance and manufacturer
                       and confirmed and intake['ready'])
    if st.button('DMF 근거로 2.3.S 작업 초안 준비', disabled=not can_prepare,
                 key='qos_prepare'):
        links = {digest: {'substance_name': substance, 'manufacturer_name': manufacturer,
                          'product_name': product, 'confirmed': True} for digest in chosen}
        try:
            package = prepare_qos_package(intake['sources'], product_name=product,
                                          product_variant=variant, dmf_links=links,
                                          selected_sections=selected)
            st.session_state['qos_package'] = package
            st.session_state['qos_package_signature'] = sha256(json.dumps({
                'input': intake['signature'], 'form': form_sha, 'product': product,
                'variant': variant, 'selected': selected, 'links': links,
            }, ensure_ascii=False, sort_keys=True).encode()).hexdigest()
            st.session_state.pop('qos_output', None)
        except ValueError as exc:
            st.session_state.pop('qos_package', None)
            st.error(str(exc))
    links = {digest: {'substance_name': substance, 'manufacturer_name': manufacturer,
                      'product_name': product, 'confirmed': True} for digest in chosen}
    signature = sha256(json.dumps({'input': intake['signature'], 'form': form_sha,
                                   'product': product, 'variant': variant,
                                   'selected': selected, 'links': links},
                                  ensure_ascii=False, sort_keys=True).encode()).hexdigest()
    package = st.session_state.get('qos_package')
    if not package or st.session_state.get('qos_package_signature') != signature or not confirmed:
        return
    covered = package['coverage']['proposed_sections'] + package['coverage']['needs_manual_check']
    st.metric('원문이 연결된 절', f"{covered}/{len(package['sections'])}")
    st.caption('이 수치는 출처 연결 범위이며 QOS 작성 완료율이나 제출 적합성 판정이 아닙니다. '
               '본문 생성 모델 호출 0회 · 원문 그대로 기입')
    for section in package['sections']:
        with st.expander(section['section_id'] + ' · ' + section['title'] +
                         ' · ' + section['status']):
            st.text(section['draft'] or '연결된 원문 없음 · 담당자 보완 필요')
            for evidence in section['evidence']:
                st.caption(f"{evidence['filename']} · {evidence.get('page') or evidence.get('sheet') or '본문'}"
                           f" · {evidence['source_id']} · SHA {evidence['document_sha256']}")
    for key, label in (('missing_sections', '근거 없음'),
                       ('ambiguous_sections', '상충'), ('deferred_sections', '보류')):
        if package[key]:
            st.warning(label + ': ' + ', '.join(package[key]))
    for check in package['manual_checks']:
        st.caption('담당자 확인: ' + check)
    fillable = [item['section_id'] for item in package['sections']
                if item['status'] in {'proposed', 'manual_check'} and item['draft']]
    if uploaded:
        with tempfile.TemporaryDirectory(prefix='qos-form-analyze-') as directory:
            path = Path(directory) / ('원본' + Path(uploaded.name).suffix.lower())
            path.write_bytes(form_bytes)
            try:
                form_ids = ctd_template_section_ids(analyze_template(path))
            except ValueError as exc:
                st.error(str(exc))
                return
        if not form_ids or set(form_ids) - set(fillable):
            st.warning('업로드 양식의 절마다 확인된 3.2.S 원문이 필요합니다. 누락 절은 기입하지 않습니다.')
            return
        template_bytes = form_bytes
        suffix = Path(uploaded.name).suffix.lower()
    elif fillable:
        template_bytes = build_qos_template(fillable)
        suffix = '.docx'
        st.caption('기본 양식은 확인 근거가 있는 절만 담는 프로젝트 예제입니다.')
    else:
        st.warning('기입 가능한 3.2.S 원문이 없습니다.')
        return
    if st.button('2.3.S 양식에 자동 기입·저장 후 검수', key='qos_fill'):
        try:
            with tempfile.TemporaryDirectory(prefix='qos-fill-') as directory:
                form = Path(directory) / ('원본' + suffix)
                target = Path(directory) / ('2.3.S_검토초안' + suffix)
                form.write_bytes(template_bytes)
                result = fill_ctd_working_template(
                    form, target, package, sources=intake['sources'],
                    product_name=product, product_variant=variant,
                    confirmed_dmf_links=links)
                st.session_state['qos_output'] = {
                    'signature': signature, 'document': target.read_bytes(),
                    'evidence': result['evidence_path'].read_bytes(), 'suffix': suffix,
                    'template_sha': result['template_sha256'],
                    'output_sha': result['output_sha256']}
        except ValueError as exc:
            st.error(str(exc))
    output = st.session_state.get('qos_output')
    if (output and output['signature'] == signature
            and output['template_sha'] == sha256(template_bytes).hexdigest()
            and output['output_sha'] == sha256(output['document']).hexdigest()):
        st.download_button('기입한 2.3.S 검토 초안', output['document'],
                           file_name='CTD_2.3.S_DMF_검토초안' + output['suffix'],
                           key='qos_download')
        st.download_button('원문 SHA·페이지·누락·제조원 확인 기록', output['evidence'],
                           file_name='CTD_2.3.S_DMF_검토초안.sources.json',
                           key='qos_evidence_download')
