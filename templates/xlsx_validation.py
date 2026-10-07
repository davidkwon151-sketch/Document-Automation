"""Resolve bounded native XLSX list controls without calculating Excel formulas."""

from hashlib import sha256
import posixpath
import re

from lxml import etree
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries

from parsers.extract import _xml


NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
X14_NS = "http://schemas.microsoft.com/office/spreadsheetml/2009/9/main"
MAX_CELLS = 2048
CELL = r"\$?[A-Za-z]{1,3}\$?[1-9][0-9]{0,6}"
RANGE = re.compile(rf"({CELL})(?::({CELL}))?\Z")


def _relative(reference):
    return any(not re.fullmatch(r"\$[A-Za-z]{1,3}\$[1-9][0-9]{0,6}", item) for item in reference.split(":"))


def raw_text(cell, shared):
    """Keep numeric XML spelling, including zero and long identifiers."""
    if cell is None:
        return ""
    if cell.find(f"{{{NS}}}f") is not None:
        raise ValueError("수식 목록 원문은 계산하지 않음")
    value = cell.find(f"{{{NS}}}v")
    text = value.text or "" if value is not None else ""
    if cell.get("t") == "s":
        return shared[int(text)]
    if cell.get("t") == "inlineStr":
        return "".join(cell.xpath("./s:is/s:t/text() | ./s:is/s:r/s:t/text()", namespaces={"s": NS}))
    if cell.get("t") == "b":
        if text not in {"0", "1"}:
            raise ValueError("목록의 논리값이 잘못됨")
        return "TRUE" if text == "1" else "FALSE"
    if cell.get("t") in {"e", "d"}:
        raise ValueError("오류·날짜 타입 목록은 확인이 필요함")
    return text


def _bounds(reference, *, bounded=True):
    if not RANGE.fullmatch(reference):
        if bounded:
            raise ValueError("한정된 A1 셀 범위만 지원함")
        bounds = range_boundaries(reference.replace("$", ""))
        return (bounds[0] or 1, bounds[1] or 1, bounds[2] or 16384, bounds[3] or 1048576)
    bounds = range_boundaries(reference.replace("$", "").upper())
    x1, y1, x2, y2 = bounds
    if not (1 <= x1 <= x2 <= 16384 and 1 <= y1 <= y2 <= 1048576):
        raise ValueError("셀 범위가 Excel 주소 범위를 벗어남")
    return bounds


def inherited_style(root, coordinate):
    column, row, _, _ = _bounds(coordinate)
    rows = root.xpath("./s:sheetData/s:row[@r=$row]", namespaces={"s": NS}, row=str(row))
    if rows and rows[0].get("customFormat") in {"1", "true"} and rows[0].get("s") is not None:
        return rows[0].get("s")
    styles = {item.get("style") for item in root.findall(f"./{{{NS}}}cols/{{{NS}}}col")
              if int(item.get("min")) <= column <= int(item.get("max")) and item.get("style") is not None}
    if len(styles) > 1:
        raise ValueError("겹친 열 서식을 확인해야 함")
    return next(iter(styles), None)


def editable_blank(root, coordinate, shared):
    if root.find(f"{{{NS}}}sheetProtection") is not None:
        return False
    x, y, _, _ = _bounds(coordinate)
    for merged in root.findall(f"./{{{NS}}}mergeCells/{{{NS}}}mergeCell"):
        x1, y1, x2, y2 = _bounds(merged.get("ref"))
        if x1 <= x <= x2 and y1 <= y <= y2 and (x, y) != (x1, y1):
            return False
    cells = root.xpath("./s:sheetData/s:row/s:c[@r=$coordinate]", namespaces={"s": NS}, coordinate=coordinate)
    if len(cells) > 1:
        return False
    try:
        return not raw_text(cells[0] if cells else None, shared).strip()
    except (ValueError, IndexError):
        return False


class XlsxLists:
    def __init__(self, archive):
        self.archive = archive
        workbook = _xml(archive.read("xl/workbook.xml"))
        rels = _xml(archive.read("xl/_rels/workbook.xml.rels"))
        targets = {node.get("Id"): node.get("Target") for node in rels if node.get("TargetMode") != "External"}
        self.sheets = {}
        for index, node in enumerate(workbook.findall(f"./{{{NS}}}sheets/{{{NS}}}sheet")):
            rid = node.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
            target = targets.get(rid, "")
            part = target.lstrip("/") if target.startswith("/") else posixpath.normpath("xl/" + target)
            self.sheets[node.get("name")] = (part, index)
        self.names = list(workbook.findall(f"./{{{NS}}}definedNames/{{{NS}}}definedName"))
        self.shared = []
        if "xl/sharedStrings.xml" in archive.namelist():
            self.shared = ["".join(node.xpath("./s:t/text() | ./s:r/s:t/text()", namespaces={"s": NS}))
                           for node in _xml(archive.read("xl/sharedStrings.xml"))]
        self.roots, self.rules = {}, {}

    def root(self, part):
        if part not in self.roots:
            self.roots[part] = _xml(self.archive.read(part))
        return self.roots[part]

    def _options(self, formula, part, *, named=False):
        expression = formula.strip().removeprefix("=")
        if re.fullmatch(r'"(?:[^\"]|\"\")*"', expression):
            return expression[1:-1].replace('""', '"').split(",")
        # Only literal ranges and a single name are evaluated; no formula engine.
        match = re.fullmatch(rf"(?:(?:'((?:[^']|'')+)'|([^'!\[\],()]+))!)?({CELL}(?::{CELL})?)", expression)
        if match:
            if named and _relative(match[3]):
                raise ValueError("상대참조 이름 정의의 기준 셀은 확인이 필요함")
            sheet_name = match[1].replace("''", "'") if match[1] is not None else match[2]
            target = self.sheets.get(sheet_name, (part, None))[0] if sheet_name is not None else part
            if sheet_name is not None and sheet_name not in self.sheets:
                raise ValueError("목록 시트가 없음")
            x1, y1, x2, y2 = _bounds(match[3])
            if (x2 - x1 + 1) * (y2 - y1 + 1) > MAX_CELLS or x1 != x2 and y1 != y2:
                raise ValueError("목록은 2048셀 이내 한 행 또는 한 열이어야 함")
            root = self.root(target)
            cells = {cell.get("r"): cell for cell in root.findall(f"./{{{NS}}}sheetData/{{{NS}}}row/{{{NS}}}c")}
            return [raw_text(cells.get(f"{get_column_letter(x)}{y}"), self.shared)
                    for y in range(y1, y2 + 1) for x in range(x1, x2 + 1)]
        if not named and re.fullmatch(r"[^\W\d][\w.\\]*", expression, re.UNICODE):
            index = next((index for name, (sheet_part, index) in self.sheets.items() if sheet_part == part), None)
            matching = [node for node in self.names if (node.get("name") or "").casefold() == expression.casefold()]
            local = [node for node in matching if node.get("localSheetId") == str(index)]
            choices = local or [node for node in matching if node.get("localSheetId") is None]
            if len(choices) != 1:
                raise ValueError("이름 정의가 없거나 범위가 모호함")
            return self._options(choices[0].text or "", part, named=True)
        raise ValueError("외부·동적 함수·다중 영역·무한 목록은 확인이 필요함")

    def sheet_rules(self, part):
        if part in self.rules:
            return self.rules[part]
        rules = []
        for node in self.root(part).xpath(".//*[local-name()='dataValidation'][@type='list']"):
            sqref = node.get("sqref") or "".join(node.xpath("./*[local-name()='sqref']/text()"))
            formula = "".join(node.xpath("./*[local-name()='formula1']//text()"))
            item = {"sqref": sqref, "formula1": formula,
                    "sha256": sha256(etree.tostring(node, method="c14n")).hexdigest(), "options": [], "error": ""}
            try:
                if node.tag not in {f"{{{NS}}}dataValidation", f"{{{X14_NS}}}dataValidation"}:
                    raise ValueError("확장형 목록은 확인이 필요함")
                for area in sqref.split():
                    _bounds(area)
                if not sqref:
                    raise ValueError("목록 적용 주소가 없음")
                expression = formula.strip().removeprefix("=")
                reference = expression.rsplit("!", 1)[-1]
                if RANGE.fullmatch(reference) and _relative(reference):
                    areas = sqref.split()
                    bounds = _bounds(areas[0])
                    if len(areas) != 1 or bounds[0:2] != bounds[2:4]:
                        raise ValueError("여러 입력칸의 상대참조 목록은 기준 확인이 필요함")
                values = self._options(formula, part)
                item["options"] = list(dict.fromkeys(value for value in values if value != ""))
                if not item["options"]:
                    raise ValueError("목록에 확인된 선택값이 없음")
                if any(value != value.strip() for value in item["options"]):
                    raise ValueError("선택값 앞뒤 공백은 현재 입력 규칙과 충돌하여 확인이 필요함")
            except (ValueError, KeyError, IndexError) as exc:
                item["error"] = str(exc)
            rules.append(item)
        self.rules[part] = rules
        return rules

    def at(self, part, coordinate):
        x, y, _, _ = _bounds(coordinate)
        matches = []
        for rule in self.sheet_rules(part):
            for area in rule["sqref"].split():
                try:
                    x1, y1, x2, y2 = _bounds(area, bounded=False)
                except ValueError:
                    continue
                if x1 <= x <= x2 and y1 <= y <= y2:
                    matches.append(rule)
                    break
        if len(matches) > 1:
            return {"sqref": " ".join(rule["sqref"] for rule in matches), "formula1": "",
                    "sha256": "", "options": [], "error": "중복 적용된 목록을 확인해야 함"}
        return matches[0] if matches else None

    def candidates(self, part, warnings):
        coordinates = set()
        for rule in self.sheet_rules(part):
            if rule["error"]:
                warnings.append(f"XLSX 목록 확인 필요: {part} {rule['sqref']} / {rule['error']}")
            for area in rule["sqref"].split():
                try:
                    x1, y1, x2, y2 = _bounds(area)
                    count = (x2 - x1 + 1) * (y2 - y1 + 1)
                    if count > MAX_CELLS or count + len(coordinates) > MAX_CELLS:
                        raise ValueError("목록 적용 범위는 2048셀까지만 자동 탐색함")
                    coordinates.update(f"{get_column_letter(x)}{y}" for y in range(y1, y2 + 1) for x in range(x1, x2 + 1))
                except ValueError as exc:
                    warnings.append(f"XLSX 목록 입력칸 탐색 제한: {part} {area} / {exc}")
        return coordinates

    def metadata(self, rule):
        return {key: rule[key] for key in ("sqref", "formula1", "sha256")}

    def check(self, part, coordinate, field, value):
        rule = self.at(part, coordinate)
        if rule is None:
            if field.get("control_type") in {"choice", "choice_unresolved"}:
                raise ValueError("원본에 목록 선택 규칙이 없음")
            return None
        if rule["error"]:
            raise ValueError("XLSX 목록 확인 필요: " + rule["error"])
        if not editable_blank(self.root(part), coordinate, self.shared):
            raise ValueError("XLSX 목록 원본의 보호·수식·병합·기존 값은 변경할 수 없음")
        if (field.get("control_type") != "choice" or field.get("options") != rule["options"]
                or field.get("xlsx_list") != self.metadata(rule)
                or field.get("choice_items") != [{"value": option, "label": option} for option in rule["options"]]):
            raise ValueError("XLSX 목록 프로파일과 원본 규칙·선택값이 일치하지 않음")
        if value not in rule["options"]:
            raise ValueError("XLSX 목록에 없는 선택값임")
        return rule


def create_choice_cell(root, coordinate):
    """Create only a previously verified DV target, inheriting row/column style."""
    column, row_number, _, _ = _bounds(coordinate)
    data = root.find(f"{{{NS}}}sheetData")
    if data is None:
        raise ValueError("XLSX sheetData가 없음")
    rows = data.xpath("./s:row[@r=$row]", namespaces={"s": NS}, row=str(row_number))
    if len(rows) > 1:
        raise ValueError("XLSX 행 주소가 중복됨")
    if rows:
        row = rows[0]
    else:
        row = etree.Element(f"{{{NS}}}row", r=str(row_number))
        position = next((i for i, node in enumerate(data) if int(node.get("r")) > row_number), len(data))
        data.insert(position, row)
    attributes = {"r": coordinate}
    style = inherited_style(root, coordinate)
    if style is not None:
        attributes["s"] = style
    cell = etree.Element(f"{{{NS}}}c", **attributes)
    position = next((i for i, node in enumerate(row) if _bounds(node.get("r"))[0] > column), len(row))
    row.insert(position, cell)
    return cell
