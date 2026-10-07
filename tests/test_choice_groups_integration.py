"""Related row completeness, exact selection rendering and confirmed reuse."""
from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path

from docx import Document
import pytest

from agent.pipeline import build_downloads, review_result, validate_template_values
from agent.output_check import verify_output
from agent.template_learning import save_learned_profile, load_learned_profile
from app.form_config import (configure_groups, group_rows, configure_profile, mapping_rows,
                             validation_rows, configure_value_rules, input_group_issues,
                             selection_options, selection_label)
from templates import analyze_template, fill_compatible_template


def row_form(tmp_path, *, direct=False):
    path = tmp_path / 'items.docx'
    doc = Document()
    doc.add_heading('품목 명세', 1)
    table = doc.add_table(rows=3, cols=3)
    table.style = 'Table Grid'
    for cell, label in zip(table.rows[0].cells, ('품목', '수량', '금액')):
        cell.text = label
    doc.save(path)
    profile = analyze_template(path)
    fields = [field for field in profile['fields'] if field['kind'] == 'docx_cell']
    assert len(fields) == 6
    for index, field in enumerate(fields):
        key = ('품목', '수량', '금액')[index % 3] + str(index // 3 + 1)
        field.update(value_key=key, label=key, required=False, input_required=direct,
                     input_mode='user_provided' if direct else 'source_grounded',
                     max_chars=200, confidence=1.0, narrative_style_required=False)
        if index % 3:
            field['validation'] = {'type': 'integer' if index % 3 == 1 else 'number', 'min': 0}
    profile['fields'] = fields
    profile['constraints'] = {'groups': [
        {'kind': 'all_or_none', 'fields': ['품목1', '수량1', '금액1']},
        {'kind': 'all_or_none', 'fields': ['품목2', '수량2', '금액2']}]}
    return path, profile


def make_result(profile, values):
    draft = {'제목': '품목 보고서', '요약': '□ 자료를 확인함 [Sbase]', '본문': '○ 자료를 확인함 [Sbase]'}
    sources = [{'source_id': 'Sbase', 'text': '자료를 확인함'}]
    for index, (key, value) in enumerate(values.items()):
        draft[key] = value + f' [Sv{index}]' if value else ''
        if value:
            sources.append({'source_id': f'Sv{index}', 'text': value})
    return {'status': 'ready', 'draft': draft, 'sources': sources, 'template_profile': profile}


def test_complete_first_row_and_unused_second_row_export_without_inventing_zeroes(tmp_path):
    path, profile = row_form(tmp_path)
    values = {'품목1': '시험기구', '수량1': '2', '금액1': '12.50', '품목2': '', '수량2': '', '금액2': ''}
    result = make_result(profile, values)
    assert not review_result(result)['blocking']
    output = build_downloads(result, confirmed=True, template_paths={'docx': path})['docx']
    document = Document(BytesIO(output))
    assert [cell.text for cell in document.tables[0].rows[2].cells] == ['', '', '']
    assert '시험기구' in document.tables[0].cell(1, 0).text
    assert result['draft']['수량2'] == '' and result['draft']['금액2'] == ''
    saved = tmp_path / 'complete.docx'
    saved.write_bytes(output)
    report = verify_output(path, saved, result['draft'], profile=profile)
    rules = next(check for check in report['checks'] if check['name'] == 'form_value_rules')
    assert rules['fields'] == 4 and rules['groups'] == 2


@pytest.mark.parametrize('missing', ['품목1', '수량1', '금액1'])
def test_partial_row_blocks_generation_export_and_independent_verification(tmp_path, missing):
    path, profile = row_form(tmp_path)
    complete = {'품목1': '시험기구', '수량1': '2', '금액1': '12.50'}
    filled = fill_compatible_template(path, complete, tmp_path / 'filled.docx', profile=profile)
    partial = complete | {missing: ''}
    result = make_result(profile, partial)
    assert review_result(result)['blocking']
    with pytest.raises(ValueError, match=missing):
        validate_template_values(result['draft'], profile)
    with pytest.raises(ValueError):
        build_downloads(result, confirmed=True, template_paths={'docx': path})
    target = tmp_path / 'partial.docx'
    target.write_bytes(b'prior output')
    before = path.read_bytes()
    with pytest.raises(ValueError, match=missing):
        fill_compatible_template(path, partial, target, profile=profile)
    assert target.read_bytes() == b'prior output' and path.read_bytes() == before
    with pytest.raises(ValueError, match=missing):
        verify_output(path, filled, partial, profile=profile)


def test_group_editor_uses_exact_json_names_and_keeps_other_relations(tmp_path):
    _, profile = row_form(tmp_path)
    profile['constraints']['relations'] = [{'kind': 'less_equal', 'left': '수량1', 'right': '수량2'}]
    saved = deepcopy(profile)
    configured = configure_groups(profile, group_rows(profile))
    assert configured == profile and profile == saved
    optional = configure_groups(profile, [{'함께 입력할 항목 JSON': '["품목1", "수량1", "금액1"]',
                                          '사용 시 필수 항목 JSON': '["품목1", "수량1"]'}])
    assert optional['constraints']['relations'] == profile['constraints']['relations']
    assert optional['constraints']['groups'][0]['required_fields'] == ['품목1', '수량1']
    for raw in ('품목1 | 수량1', '[]', '["없는 칸"]', '["품목1", "품목1"]'):
        with pytest.raises(ValueError):
            configure_groups(profile, [{'함께 입력할 항목 JSON': raw}])


def test_confirmed_group_mapping_rebinds_and_cache_stores_no_actual_values(tmp_path):
    path, profile = row_form(tmp_path)
    profile['configured'] = True
    rows = mapping_rows(profile)
    for row in rows:
        row['채울 값'] = '새 ' + row['채울 값']
    configured, mapping = configure_profile(profile, rows)
    assert configured['constraints']['groups'][0]['fields'] == ['새 품목1', '새 수량1', '새 금액1']
    configured['demo_values'] = {'새 품목1': 'PRIVATE-DO-NOT-CACHE'}
    saved = save_learned_profile(path, configured, mapping, cache_dir=tmp_path / 'cache', user_confirmed=True)
    loaded = load_learned_profile(path, tmp_path / 'cache')
    assert loaded['constraints']['groups'] == configured['constraints']['groups']
    assert 'PRIVATE-DO-NOT-CACHE' not in saved.read_text(encoding='utf-8')
    assert loaded['learning']['engine_version'] == 'template-mapping-6'
    rows[0]['채울 값'] = ''
    with pytest.raises(ValueError):
        configure_profile(profile, rows)


def test_direct_group_checks_skip_fields_that_need_later_grounded_generation(tmp_path):
    _, profile = row_form(tmp_path, direct=True)
    issues = input_group_issues({'품목1': '시험기구'}, profile)
    assert {issue['field'] for issue in issues} == {'수량1', '금액1'}
    profile['fields'][1]['input_required'] = False
    assert not input_group_issues({'품목1': '시험기구'}, profile)


def test_selection_editor_and_labels_preserve_export_values_and_never_default_to_yes():
    field = {'id': 'choice', 'label': '구분', 'value_key': '구분', 'kind': 'docx_choice',
             'control_type': 'choice', 'options': ['EXP', 'DOM'],
             'choice_items': [{'value': 'EXP', 'label': '수출'}, {'value': 'DOM', 'label': '국내'}],
             'required': False, 'input_required': True}
    profile = {'fields': [field]}
    assert selection_options(field) == ['', 'EXP', 'DOM']
    assert selection_label(field, 'EXP') == '수출'
    configured = configure_value_rules(profile, validation_rows(profile))
    assert configured['fields'][0]['validation'] == {'type': 'choice', 'options': ['EXP', 'DOM']}
    assert configured['fields'][0]['input_required']
    bad = validation_rows(profile)
    bad[0]['선택값 JSON'] = '["EXP", "DOM", "승인"]'
    with pytest.raises(ValueError):
        configure_value_rules(profile, bad)
    assert selection_options({'kind': 'docx_checkbox', 'control_type': 'checkbox', 'options': ['true', 'false']}) == ['', 'true', 'false']


def test_native_choice_rendering_removes_only_trailing_citations_keeps_json_and_exact_checks(tmp_path):
    from openpyxl import Workbook, load_workbook
    from openpyxl.worksheet.datavalidation import DataValidation
    path = tmp_path / 'classification.xlsx'
    workbook = Workbook()
    workbook.active['A1'] = '구분'
    workbook.active['B1'] = ''
    rule = DataValidation(type='list', formula1='"EXP,DOM"')
    workbook.active.add_data_validation(rule)
    rule.add('B1')
    workbook.save(path)
    profile = analyze_template(path)
    field = next(field for field in profile['fields'] if field['id'].endswith(':B1'))
    field.update(required=False, max_chars=3, narrative_style_required=False)
    profile['fields'] = [field]
    result = make_result(profile, {'구분': 'EXP'})
    assert not review_result(result)['blocking']
    before = path.read_bytes()
    output = build_downloads(result, confirmed=True, template_paths={'xlsx': path})['xlsx']
    file = tmp_path / 'result.xlsx'
    file.write_bytes(output)
    checked = load_workbook(file)
    assert checked.active['B1'].value == 'EXP'
    checked.close()
    assert result['draft']['구분'].startswith('EXP [S') and result['sources'][1]['text'] == 'EXP'
    assert path.read_bytes() == before
    report = verify_output(path, file, result['draft'], profile=profile)
    rules = next(check for check in report['checks'] if check['name'] == 'form_value_rules')
    assert rules['fields'] == 1 and rules['groups'] == 0
    edited = deepcopy(result)
    edited['sources'].append({'source_id': 'Sbad', 'text': 'OTHER'})
    edited['draft']['구분'] = 'OTHER [Sbad]'
    assert review_result(edited)['blocking']
    with pytest.raises(ValueError):
        build_downloads(edited, confirmed=True, template_paths={'xlsx': path})


def test_registered_profile_keeps_approved_locations_and_cannot_drop_native_choices(tmp_path, monkeypatch):
    from openpyxl import Workbook
    from openpyxl.worksheet.datavalidation import DataValidation
    import templates.profiles as registered
    path = tmp_path / 'registered.xlsx'
    workbook = Workbook()
    workbook.active['A1'] = '구분'
    workbook.active['B1'] = ''
    workbook.active['A2'] = '별도 구분'
    workbook.active['B2'] = ''
    rule = DataValidation(type='list', formula1='"EXP,DOM"')
    workbook.active.add_data_validation(rule)
    rule.add('B1:B2')
    workbook.save(path)
    profile = analyze_template(path)
    field = deepcopy(next(field for field in profile['fields'] if field['id'].endswith(':B1')))
    field.update(value_key='확인한 구분', input_required=False)
    for key in ('control_type', 'options', 'choice_items', 'xlsx_list'):
        field.pop(key, None)
    profile.update(fields=[field], resource_kind='blank_form', schema_version=1)
    directory = tmp_path / 'profiles'
    directory.mkdir()
    record = directory / 'registered.json'
    record.write_text(json.dumps(profile, ensure_ascii=False), encoding='utf-8')
    monkeypatch.setattr(registered, 'PROFILE_DIRECTORY', directory)
    loaded = registered.load_form_profile(path)
    assert len(loaded['fields']) == 1
    found = loaded['fields'][0]
    assert found['value_key'] == '확인한 구분' and found['options'] == ['EXP', 'DOM']
    assert found['input_required'] and found['control_type'] == 'choice' and found['xlsx_list']
    assert profile['fields'][0]['input_required'] is False
    with pytest.raises(ValueError):
        fill_compatible_template(path, {'확인한 구분': 'OTHER'}, tmp_path / 'bad.xlsx', profile=loaded)
    profile['fields'][0]['validation'] = {'type': 'choice', 'options': ['EXP', 'DOM', 'OTHER']}
    record.write_text(json.dumps(profile, ensure_ascii=False), encoding='utf-8')
    with pytest.raises(ValueError):
        registered.load_form_profile(path)


def test_invalid_direct_values_and_partial_user_rows_fail_before_model_or_parsing(tmp_path):
    from agent.pipeline import run_pipeline
    class NoCalls:
        def generate_json(self, *args, **kwargs):
            raise AssertionError('Invalid direct input must not call a model')
    _, profile = row_form(tmp_path, direct=True)
    with pytest.raises(ValueError, match='수량1'):
        run_pipeline('품목을 작성해줘', client=NoCalls(), template_profile=profile,
                     field_values={'품목1': '시험기구', '수량1': '둘', '금액1': '12.50'})
    with pytest.raises(ValueError, match='금액1'):
        run_pipeline('품목을 작성해줘', client=NoCalls(), template_profile=profile,
                     answers={'품목1': '시험기구', '수량1': '2'})
    choice = {'fields': [{'value_key': '구분', 'label': '구분', 'required': False, 'input_required': True,
                         'control_type': 'choice', 'options': ['EXP', 'DOM']}]}
    with pytest.raises(ValueError, match='선택값'):
        run_pipeline('구분 작성', client=NoCalls(), template_profile=choice, field_values={'구분': 'OTHER'})


def test_required_plain_field_cannot_disappear_from_mapping_rows_or_direct_mapping():
    from templates.value_rules import mapped_rule_profile
    profile = {'fields': [
        {'id': 'a', 'kind': 'docx_cell', 'label': '성명', 'value_key': '성명', 'required': True},
        {'id': 'b', 'kind': 'docx_cell', 'label': '제목', 'value_key': '제목', 'required': False}]}
    before = deepcopy(profile)
    with pytest.raises(ValueError, match='필수'):
        configure_profile(profile, mapping_rows(profile)[1:])
    with pytest.raises(ValueError, match='필수'):
        mapped_rule_profile(profile, {'b': '제목'})
    renamed = mapped_rule_profile(profile, {'a': '신청인 성명', 'b': '문서 제목'})
    assert renamed['fields'][0]['required'] is True and profile == before
