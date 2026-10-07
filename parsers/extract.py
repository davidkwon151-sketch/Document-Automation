"""출처 위치를 유지하는 원자료 파서. DOCX/HWPX 쪽 번호는 추정하지 않음."""

from collections import Counter
from datetime import date, datetime
from pathlib import Path
import re
import warnings
from zipfile import BadZipFile, ZipFile

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from lxml import etree
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pypdf import PdfReader
from pypdf.constants import UserAccessPermissions
from pypdf.errors import FileNotDecryptedError
import pdfplumber


MAX_FILE_BYTES = 30 * 1024 * 1024
MAX_PACKAGE_BYTES = 100 * 1024 * 1024


class ParseError(ValueError):
    """사용자에게 표시할 파싱 오류와 기계 판별용 코드."""

    def __init__(self, message: str, code: str = "invalid_file"):
        super().__init__(message)
        self.code = code


def _path(path: str | Path) -> Path:
    path = Path(path)
    if not path.is_file():
        raise ParseError(f"파일을 찾을 수 없음: {path.name}", "missing_file")
    if path.stat().st_size == 0:
        raise ParseError(f"빈 파일임: {path.name}", "empty_file")
    if path.stat().st_size > MAX_FILE_BYTES:
        raise ParseError("파일 크기는 30 MiB 이하이어야 함", "too_large")
    return path


def _check_zip(archive: ZipFile) -> None:
    if sum(item.file_size for item in archive.infolist()) > MAX_PACKAGE_BYTES:
        raise ParseError("압축 해제 크기는 100 MiB 이하이어야 함", "too_large")
    if len(set(archive.namelist())) != len(archive.namelist()):
        raise ParseError("중복된 ZIP 항목이 포함됨")


def _xml(data: bytes):
    if b"<!DOCTYPE" in data.upper() or b"<!ENTITY" in data.upper():
        raise ParseError("DTD 또는 외부 엔티티가 포함된 XML은 지원하지 않음")
    return etree.fromstring(data, etree.XMLParser(resolve_entities=False, no_network=True))


def _block(text: str, location: str, *, page=None, sheet=None, tables=None) -> dict:
    return {"페이지": page, "시트": sheet, "위치": location,
            "본문": text, "표 목록": tables or []}


def _table(rows: list[list[str]], location: str, *, page=None, sheet=None) -> dict:
    return {"행": rows, "위치": location, "페이지": page, "시트": sheet}


def _result(path: Path, blocks: list[dict], tables: list[dict]) -> dict:
    text = "\n".join(block["본문"] for block in blocks if block["본문"].strip())
    if not text.strip():
        raise ParseError(f"추출 가능한 본문이 없음: {path.name}", "empty_document")
    return {"파일명": path.name, "본문": text, "표 목록": tables, "페이지/시트 정보": blocks}


def parse_file(path: str | Path) -> dict:
    """확장자를 확인하고 동일한 네 필드 구조로 추출함."""
    path = _path(path)
    if path.suffix.lower() == '.hwp':
        from .hwp import parse_hwp
        return parse_hwp(path)
    if path.suffix.lower() == '.pptx':
        from .pptx import parse_pptx
        return parse_pptx(path)
    parser = {".pdf": parse_pdf, ".xlsx": parse_xlsx,
              ".docx": parse_docx, ".hwpx": parse_hwpx}.get(path.suffix.lower())
    if parser is None:
        raise ParseError(f"지원하지 않는 파일 형식임: {path.suffix}", "unsupported_format")
    return parser(path)


def parse_pdf(path: str | Path) -> dict:
    """PDF 콘텐츠 스트림 읽기 순서를 보존하고 다른 물리적 순서를 기록함.

    동일 y 좌표의 좌우 열을 한 문장으로 합치지 않도록 텍스트 흐름을
    사용함. 이는 문서 작성자의 콘텐츠 순서이며 의미상 정답이라는 인증은
    아님. 두 추출의 문자/빈도가 다르면 손실을 승인하지 않고 원래 물리적
    순서로 돌아가 경고함. 표의 행/열과 출처 페이지는 별도로 유지함.
    """
    path = _path(path)
    blocks, tables = [], []
    try:
        # Normal blank-password reading is permitted only when the original
        # grants ordinary text extraction too. Never supply a password, remove
        # protection, or use OCR to work around an extraction restriction.
        reader = PdfReader(path)
        if reader.is_encrypted:
            permissions = reader.user_access_permissions
            if permissions is None or not permissions & UserAccessPermissions.EXTRACT:
                raise ParseError("PDF의 텍스트 추출 허용을 확인할 수 없음. 보호를 유지하고 "
                                 "자동 읽기를 보류함", "protected_document")
            try:
                len(reader.pages)
            except FileNotDecryptedError as exc:
                raise ParseError("암호가 필요한 PDF는 보호를 유지하고 자동 읽기를 보류함",
                                 "protected_document") from exc
        with pdfplumber.open(path) as document:
            for number, page in enumerate(document.pages, 1):
                physical_text = page.extract_text() or ""
                flow_text = page.extract_text(use_text_flow=True) or ""
                physical_chars = re.sub(r"\s+", "", physical_text)
                flow_chars = re.sub(r"\s+", "", flow_text)
                same_characters = Counter(physical_chars) == Counter(flow_chars)
                has_control_characters = bool(re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]",
                                                        physical_text + flow_text))
                use_flow = same_characters and not has_control_characters
                text = flow_text if use_flow else physical_text
                page_tables = []
                for index, rows in enumerate(page.extract_tables(), 1):
                    table = _table([[cell or "" for cell in row] for row in rows],
                                   f"페이지 {number}/표 {index}", page=number)
                    page_tables.append(table)
                    tables.append(table)
                if not text.strip():
                    warnings.warn(f"{path.name} 페이지 {number}: 텍스트가 없어 OCR 확인 필요함",
                                  UserWarning, stacklevel=2)
                block = _block(text, f"페이지 {number}", page=number, tables=page_tables)
                if physical_chars != flow_chars or has_control_characters:
                    block["PDF 읽기순서"] = {
                        "선택 방식": "PDF 콘텐츠 스트림" if use_flow else "물리적 줄순서",
                        "물리적 줄순서 본문": physical_text,
                        "콘텐츠 스트림 본문": flow_text,
                        "문자와 빈도 보존": same_characters,
                        "제어문자 없음": not has_control_characters,
                        "확인 필요": True,
                        "주의": "두 읽기순서가 다르므로 열·표·문단의 의미 연결을 원문과 확인해야 함",
                    }
                if not same_characters:
                    warnings.warn(f"{path.name} 페이지 {number}: 읽기순서 추출의 문자와 빈도가 달라 "
                                  "물리적 줄순서를 유지함. 원문 확인 필요함", UserWarning, stacklevel=2)
                elif has_control_characters:
                    warnings.warn(f"{path.name} 페이지 {number}: 제어문자가 있어 "
                                  "물리적 줄순서를 유지함. 원문 확인 필요함", UserWarning, stacklevel=2)
                if not use_flow:
                    reason = ("PDF 읽기순서 추출의 문자 손실·중복을 원문과 대조해야 함"
                              if not same_characters else "PDF 제어문자가 포함된 본문을 원문과 대조해야 함")
                    block.update({"검증 필요": True, "불확실한 항목": [reason],
                                  "추출 방식": "PDF 물리적 줄순서 / 원문 확인 대기"})
                    for table in page_tables:
                        table["검증 필요"] = True
                if re.search(r"\(cid:\d+\)", text) or any(
                        re.search(r"\(cid:\d+\)", str(table)) for table in page_tables):
                    reason = "PDF 글꼴의 문자 매핑을 복원하지 못함. CID 문자와 주변 문구를 원문에서 확인해야 함"
                    block["검증 필요"] = True
                    block.setdefault("불확실한 항목", []).append(reason)
                    block["추출 방식"] = "PDF 문자 매핑 / 원문 확인 대기"
                    for table in page_tables:
                        if re.search(r"\(cid:\d+\)", str(table)):
                            table["검증 필요"] = True
                    warnings.warn(f"{path.name} 페이지 {number}: {reason}", UserWarning, stacklevel=2)
                blocks.append(block)
    except ParseError:
        raise
    except Exception as exc:
        raise ParseError(f"PDF를 읽을 수 없음: {path.name}") from exc
    if not any(block["본문"].strip() for block in blocks):
        raise ParseError("PDF에 추출할 텍스트가 없음. 스캔본은 OCR이 필요함", "ocr_required")
    return _result(path, blocks, tables)


def _cell_text(cell) -> str:
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    # 단위가 없는 숫자로 변환하지 않고 원래 값과 표시 형식을 유지함.
    if isinstance(value, (int, float)) and "%" in cell.number_format:
        decimal = re.search(r"0\.([0#]+).*%", cell.number_format)
        precision = len(decimal[1]) if decimal else 0
        return f"{value * 100:.{precision}f}%"
    return str(value)


def parse_xlsx(path: str | Path) -> dict:
    path = _path(path)
    blocks, tables = [], []
    workbook = None
    try:
        with ZipFile(path) as archive:
            _check_zip(archive)
        # 수식 자체를 근거 수치로 오인하지 않도록 캐시가 없는 수식을 거부함.
        workbook = load_workbook(path, read_only=True, data_only=False)
        cached = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in workbook.worksheets:
                rows = []
                for row_index, cells in enumerate(sheet.iter_rows(), 1):
                    values = []
                    for cell in cells:
                        if cell.data_type == "f":
                            cached_cell = cached[sheet.title][cell.coordinate]
                            if cached_cell.value is None:
                                raise ParseError(f"{sheet.title}!{cell.coordinate}: 수식 계산값이 없음. "
                                                 "Excel에서 계산 후 저장해야 함", "uncalculated_formula")
                            values.append(_cell_text(cached_cell))
                        else:
                            values.append(_cell_text(cell))
                    while values and not values[-1]:
                        values.pop()
                    if not values:
                        continue
                    rows.append(values)
                    location = f"시트 {sheet.title}!A{row_index}:{get_column_letter(len(values))}{row_index}"
                    blocks.append(_block(" | ".join(values), location, sheet=sheet.title))
                if rows:
                    tables.append(_table(rows, f"시트 {sheet.title}", sheet=sheet.title))
        finally:
            cached.close()
    except ParseError:
        raise
    except Exception as exc:
        raise ParseError(f"XLSX를 읽을 수 없음: {path.name}") from exc
    finally:
        if workbook is not None:
            workbook.close()
    return _result(path, blocks, tables)


def parse_docx(path: str | Path) -> dict:
    path = _path(path)
    blocks, tables = [], []
    try:
        with ZipFile(path) as archive:
            _check_zip(archive)
        document = Document(path)

        def read_container(container, prefix):
            element = container.element.body if hasattr(container, "element") else container._tc
            paragraph_number = table_number = 0
            for child in element:
                local_name = etree.QName(child).localname
                if local_name == "p":
                    paragraph_number += 1
                    text = Paragraph(child, container).text
                    if text.strip():
                        blocks.append(_block(text, f"{prefix}/문단 {paragraph_number}"))
                elif local_name == "tbl":
                    table_number += 1
                    table = Table(child, container)
                    location = f"{prefix}/표 {table_number}"
                    rows = [[cell.text for cell in row.cells] for row in table.rows]
                    tables.append(_table(rows, location))
                    seen = set()
                    for row_index, row in enumerate(table.rows, 1):
                        for col_index, cell in enumerate(row.cells, 1):
                            if cell._tc in seen:
                                continue
                            seen.add(cell._tc)
                            read_container(cell, f"{location}/행 {row_index}/열 {col_index}")

        read_container(document, "본문")
    except ParseError:
        raise
    except Exception as exc:
        raise ParseError(f"DOCX를 읽을 수 없음: {path.name}") from exc
    return _result(path, blocks, tables)


def _hwpx_text(paragraph) -> str:
    """해당 문단 자체의 run/t만 읽어 표·머리말의 중복 추출을 방지함."""
    parts = []
    for run in paragraph:
        if etree.QName(run).localname != "run":
            continue
        for text in run:
            if etree.QName(text).localname != "t":
                continue
            parts.append(text.text or "")
            for child in text:
                name = etree.QName(child).localname
                parts.append({"lineBreak": "\n", "tab": "\t", "nbSpace": " ",
                              "fwSpace": "　", "hyphen": "-"}.get(name, child.text or ""))
                parts.append(child.tail or "")
    return "".join(parts)


def parse_hwpx(path: str | Path) -> dict:
    path = _path(path)
    blocks, tables = [], []
    try:
        with ZipFile(path) as archive:
            _check_zip(archive)
            names = archive.namelist()
            required = {"mimetype", "version.xml", "Contents/header.xml", "Contents/content.hpf",
                        "META-INF/container.xml", "Contents/section0.xml"}
            if not required.issubset(names) or archive.read("mimetype").strip() != b"application/hwp+zip":
                raise ParseError("정상적인 HWPX 패키지가 아님")
            header = _xml(archive.read("Contents/header.xml"))
            sections = sorted((name for name in names if re.fullmatch(r"Contents/section\d+\.xml", name)),
                              key=lambda name: int(re.search(r"section(\d+)", name)[1]))
            if int(header.get("secCnt", len(sections))) != len(sections):
                raise ParseError("HWPX 구역 개수가 header와 일치하지 않음")
            for section in sections:
                root = _xml(archive.read(section))
                tree = root.getroottree()
                for paragraph in root.iter():
                    if etree.QName(paragraph).localname != "p":
                        continue
                    text = _hwpx_text(paragraph)
                    if text.strip():
                        blocks.append(_block(text, f"{section}:{tree.getpath(paragraph)}"))
                for index, table in enumerate(root.xpath(".//*[local-name()='tbl']"), 1):
                    rows = []
                    for row in table.xpath("./*[local-name()='tr']"):
                        cells = []
                        for cell in row.xpath("./*[local-name()='tc']"):
                            paragraphs = cell.xpath("./*[local-name()='subList']/*[local-name()='p']")
                            cells.append("\n".join(_hwpx_text(p) for p in paragraphs))
                        rows.append(cells)
                    tables.append(_table(rows, f"{section}/표 {index}"))
    except ParseError:
        raise
    except (BadZipFile, KeyError, ValueError, etree.XMLSyntaxError, OSError) as exc:
        raise ParseError(f"HWPX를 읽을 수 없음: {path.name}") from exc
    return _result(path, blocks, tables)
