"""개인정보 없는 파서·양식 fixture 재생성: python -m samples.generate."""

from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree
from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.platypus import Table, TableStyle


ROOT = Path(__file__).resolve().parents[1]
HP = "http://www.hancom.co.kr/hwpml/2011/paragraph"
SEED = ROOT / "samples/hancom_table_seed.hwpx"
VALUES = {"제목": "교육 실시 결과보고", "요약": "□ 교육 참석률 95%를 달성함 [S1]",
          "본문": "□ 교육 대상 100명 중 95명이 참석함 [S1]\n○ 만족도 4.5점을 기록함 [S2]\n- 후속 교육을 검토함"}


def _font(run, *, bold=False, size=11):
    run.font.name = "맑은 고딕"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def _docx_template(path: Path, label: str):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.5)
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "맑은 고딕", Pt(11)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    normal.paragraph_format.space_after = Pt(8)
    for style_name in ("Title", "Heading 1", "Heading 2"):
        document.styles[style_name].font.name = "맑은 고딕"
        document.styles[style_name].font.color.rgb = RGBColor(0, 0, 0)
        for border in document.styles[style_name]._element.findall(".//" + qn("w:pBdr")):
            border.getparent().remove(border)
    title = document.add_paragraph(style="Title")
    for text in ("{{제", "목}}"):
        _font(title.add_run(text), bold=True, size=18)
    document.add_paragraph(label)
    document.add_heading("핵심 요약", level=1)
    _font(document.add_paragraph().add_run("{{요약}}"))
    document.add_heading("주요 내용", level=1)
    _font(document.add_paragraph().add_run("{{본문}}"))
    table = document.add_table(rows=2, cols=2)
    table.autofit = False
    table.columns[0].width, table.columns[1].width = Cm(3), Cm(12.8)
    table.style = "Table Grid"
    for row in table.rows:
        for cell, width in zip(row.cells, (Cm(3), Cm(12.8))):
            cell.width = width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            properties = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right"):
                border = OxmlElement("w:" + edge)
                for key, value in {"val": "single", "sz": "4", "color": "D9D9D9"}.items():
                    border.set(qn("w:" + key), value)
                borders.append(border)
            properties.append(borders)
    generic = label.startswith('사내외 문서')
    for row, label_text, value in zip(table.rows, ("문서 제목", "문서 요약") if generic else ("보고 제목", "보고 요약"), ("{{제목}}", "{{요약}}")):
        _font(row.cells[0].paragraphs[0].add_run(label_text), bold=True)
        _font(row.cells[1].paragraphs[0].add_run(value))
    _font(section.header.paragraphs[0].add_run(("사내외 문서  " if generic else "사내 보고서  ") + "{{제목}}"), size=9)
    _font(section.footer.paragraphs[0].add_run("검토용 문서"), size=9)
    document.save(path)


def _hp_p(text: str):
    paragraph = etree.Element(f"{{{HP}}}p", id="0", paraPrIDRef="0", styleIDRef="0",
                              pageBreak="0", columnBreak="0", merged="0")
    run = etree.SubElement(paragraph, f"{{{HP}}}run", charPrIDRef="0")
    etree.SubElement(run, f"{{{HP}}}t").text = text
    return paragraph


def _hwpx_sample(path: Path, label: str, *, source=False):
    with ZipFile(SEED) as archive:
        root = etree.fromstring(archive.read("Contents/section0.xml"))
        first = root[0]
        table = root.find(f".//{{{HP}}}tbl")
        table.getparent().remove(table)
        for child in list(first):
            if etree.QName(child).localname == "linesegarray":
                first.remove(child)
        for child in list(root)[1:]:
            root.remove(child)
        first.append(deepcopy(_hp_p("교육 실시 결과" if source else "{{제")[0]))
        if not source:
            first.append(deepcopy(_hp_p("목}}")[0]))  # run 분할 자리표시자 fixture
        root.append(_hp_p(label))
        root.append(_hp_p("교육 대상 100명 중 95명이 참석함" if source else "{{요약}}"))
        root.append(_hp_p("만족도 4.5점을 기록함" if source else "{{본문}}"))
        table_paragraph = _hp_p("")
        table_paragraph[0].insert(0, table)
        root.append(table_paragraph)
        cell_texts = ("항목", "대상", "참석", "교육", "100명", "95명") if source else (
            "보고 제목", "보고 요약", "검토", "{{제목}}", "{{요약}}", "검토 필요")
        if not source and label.startswith('사내외 문서'):
            cell_texts = ('문서 제목', '문서 요약', '검토', '{{제목}}', '{{요약}}', '검토 필요')
        for cell, text in zip(table.findall(f".//{{{HP}}}tc"), cell_texts):
            paragraph = cell.find(f"{{{HP}}}subList/{{{HP}}}p")
            for child in list(paragraph):
                paragraph.remove(child)
            paragraph.append(_hp_p(text)[0])
            size = cell.find(f"{{{HP}}}cellSz")
            size.set("height", "1800")
        table.find(f"{{{HP}}}sz").set("height", "4200")
        changes = {"Contents/section0.xml": etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True),
                   "Preview/PrvText.txt": ("교육 실시 결과" if source else "{{제목}}\n{{요약}}\n{{본문}}").encode("utf-8")}
        with ZipFile(path, "w") as result:
            for item in archive.infolist():
                result.writestr(item, changes.get(item.filename, archive.read(item.filename)))


def generate(output_root: Path = ROOT):
    """전체 fixture를 지정 경로에 생성하며 seed 원본은 변경하지 않음."""
    samples, templates = output_root / "samples", output_root / "templates"
    samples.mkdir(parents=True, exist_ok=True)
    templates.mkdir(parents=True, exist_ok=True)
    for stem, label in {"result_report": "결과보고서", "weekly_report": "주간업무보고", "approval_request": "품의서",
                        "generic_document": "사내외 문서 (일반 검토용)"}.items():
        _docx_template(templates / f"{stem}.docx", label)
        _hwpx_sample(templates / f"{stem}.hwpx", label)

    document = Document()
    _font(document.add_paragraph().add_run("교육 대상 100명 중 95명이 참석함"))
    _font(document.add_paragraph().add_run("만족도 4.5점을 기록함"))
    table = document.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, (("항목", "실적"), ("참석", "95명"))):
        for cell, value in zip(row.cells, values):
            _font(cell.paragraphs[0].add_run(value))
    document.save(samples / "source.docx")
    _hwpx_sample(samples / "source.hwpx", "교육 실적 원자료", source=True)

    workbook = Workbook()
    workbook.active.title = "교육 실적"
    for values in (("항목", "실적"), ("교육 대상", "100명"), ("교육 참석", "95명"), ("교육 참석률", 0.95)):
        workbook.active.append(values)
    workbook.active["B4"].number_format = "0%"
    satisfaction = workbook.create_sheet("만족도")
    satisfaction.append(["교육 만족도", "4.5점"])
    workbook.save(samples / "source.xlsx")

    font_path = Path("C:/Windows/Fonts/malgun.ttf")
    if not font_path.exists():
        raise FileNotFoundError("PDF fixture 생성에 필요한 한글 글꼴 경로를 samples/generate.py에서 지정해야 함")
    pdfmetrics.registerFont(TTFont("FixtureKorean", str(font_path)))
    pdf = canvas.Canvas(str(samples / "source.pdf"), pagesize=(595, 842), invariant=1)
    pdf.setFont("FixtureKorean", 14)
    pdf.drawString(60, 780, "교육 실시 실적 원자료")
    pdf.setFont("FixtureKorean", 11)
    pdf.drawString(60, 745, "교육 대상 100명 중 95명이 참석함")
    table = Table([["항목", "실적"], ["교육 참석", "95명"]], colWidths=[200, 200], rowHeights=[28, 28])
    table.setStyle(TableStyle([("FONTNAME", (0, 0), (-1, -1), "FixtureKorean"),
                               ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                               ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
    table.wrapOn(pdf, 400, 100)
    table.drawOn(pdf, 60, 640)
    pdf.showPage()
    pdf.setFont("FixtureKorean", 11)
    pdf.drawString(60, 780, "교육 만족도 4.5점을 기록함")
    pdf.save()
    from templates import fill_template
    for suffix in ("docx", "hwpx"):
        fill_template(templates / f"result_report.{suffix}", VALUES, samples / f"filled_result_report.{suffix}")
    generate_company_forms(output_root)


def generate_company_forms(output_root: Path = ROOT):
    """자리표시자가 없는 회사 양식 및 입력칸/정적 PDF 호환 fixture."""
    samples = output_root / "samples"
    samples.mkdir(parents=True, exist_ok=True)
    document = Document()
    title = document.add_paragraph("교육 결과보고서", style="Title")
    _font(title.runs[0], bold=True, size=18)
    title_style = document.styles["Title"]
    for border in title_style._element.findall(".//" + qn("w:pBdr")):
        border.getparent().remove(border)
    table = document.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    for row, label in zip(table.rows, ("제목", "요약", "내용", "작성자")):
        _font(row.cells[0].paragraphs[0].add_run(label), bold=True)
        _font(row.cells[1].paragraphs[0].add_run(""))
        row.cells[0].width, row.cells[1].width = Cm(3), Cm(12)
    document.save(samples / "sample_company_form.docx")
    with ZipFile(SEED) as archive:
        root = etree.fromstring(archive.read("Contents/section0.xml"))
        for cell, label in zip(root.findall(f".//{{{HP}}}tc"), ("제목", "요약", "내용", "", "", "")):
            paragraph = cell.find(f"{{{HP}}}subList/{{{HP}}}p")
            for child in list(paragraph):
                paragraph.remove(child)
            paragraph.append(_hp_p(label)[0])
        with ZipFile(samples / "sample_company_form.hwpx", "w") as result:
            for item in archive.infolist():
                result.writestr(item, etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                                if item.filename == "Contents/section0.xml" else archive.read(item.filename))
    workbook = Workbook()
    workbook.active.title = "결과보고서"
    from openpyxl.styles import Alignment, Border, Side, Font
    for row, label in enumerate(("제목", "요약", "내용", "작성자"), 1):
        workbook.active.cell(row, 1, label)
        cell = workbook.active.cell(row, 2)
        cell.font = Font(name="맑은 고딕", size=11)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.border = Border(bottom=Side(style="thin"))
        workbook.active.row_dimensions[row].height = 35 if row != 3 else 75
    workbook.active["C5"] = "=1+2"
    workbook.active.column_dimensions["A"].width = 16
    workbook.active.column_dimensions["B"].width = 75
    workbook.save(samples / "sample_company_form.xlsx")
    pdf = canvas.Canvas(str(samples / "sample_company_form.pdf"), pagesize=(595, 842), invariant=1)
    pdf.setFont("Helvetica", 18)
    pdf.drawString(50, 790, "Company Report")
    for name, label, y, height in (("title", "Title", 730, 30), ("summary", "Summary", 650, 50), ("body", "Body", 450, 150)):
        pdf.setFont("Helvetica", 11)
        pdf.drawString(50, y + height + 10, label)
        pdf.acroForm.textfield(name=name, tooltip=label, x=50, y=y, width=480, height=height,
                              fontSize=10, fieldFlags="multiline" if name != "title" else "")
    pdf.showPage()
    pdf.save()
    static = canvas.Canvas(str(samples / "sample_static_form.pdf"), pagesize=(595, 842), invariant=1)
    static.setFont("FixtureKorean", 14)
    static.drawString(50, 790, "교육 결과보고서")
    static.setFont("FixtureKorean", 11)
    for label, y, height in (("제목", 720, 30), ("요약", 600, 80), ("내용", 380, 160)):
        static.drawString(50, y + height + 8, label)
        static.rect(50, y, 480, height)
    static.showPage()
    static.save()
    selection = canvas.Canvas(str(samples / "sample_selection_form.pdf"), pagesize=(595, 842), invariant=1)
    selection.setFont("Helvetica", 13)
    selection.drawString(50, 790, "Selection compatibility fixture - NOT FOR SUBMISSION")
    selection.acroForm.checkbox(name="consent", tooltip="Consent", x=50, y=720, size=18, checked=False)
    for index, option in enumerate(("approve", "reject")):
        selection.drawString(80, 665 - index * 35, option)
        selection.acroForm.radio(name="decision", tooltip="Decision", value=option,
                                 selected=False, x=50, y=660 - index * 35, buttonStyle="circle")
    selection.acroForm.choice(name="department", tooltip="Department", x=50, y=570, width=200,
                             height=25, options=["Planning", "Operations"], value="Planning")
    selection.showPage()
    selection.save()


if __name__ == "__main__":
    generate()
