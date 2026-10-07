"""Exact typed XLSX XML values; formula preservation is not native recalculation."""
from copy import deepcopy
from decimal import Decimal, getcontext
from zipfile import ZipFile

from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.workbook.properties import CalcProperties
from openpyxl.styles import Alignment, Font
import pytest

from agent.output_check import verify_output
from templates import TemplateError, analyze_template, fill_compatible_template


NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}


def blank_form(tmp_path, *, rule=None):
    source = tmp_path / 'numeric_source.xlsx'
    book = Workbook()
    sheet = book.active
    sheet.title = '신청서'
    for row, label in enumerate(['금액', '예산', '등록번호', '날짜'], 1):
        sheet.cell(row, 1, label)
        cell = sheet.cell(row, 2)
        cell.font = Font(name='맑은 고딕', size=11, bold=True)
        cell.alignment = Alignment(wrap_text=True)
        cell.number_format = '#,##0.00' if row <= 2 else '@'
    sheet['C1'] = '합계 {{금액}}원'
    sheet['B5'] = '=SUM(B1:B2)'
    sheet.merge_cells('D1:E1')
    sheet['D1'] = '원본 안내문'
    book.save(source)
    book.close()
    profile = analyze_template(source)
    for field in profile['fields']:
        if field['kind'] == 'xlsx_cell' and field['value_key'] in {'금액', '예산'}:
            field['validation'] = deepcopy(rule or {'type': 'number', 'unit': '원', 'unit_location': 'label'})
    return source, profile


def cells(path):
    with ZipFile(path) as archive:
        root = etree.fromstring(archive.read('xl/worksheets/sheet1.xml'))
    return {cell.get('r'): cell for cell in root.xpath('.//s:sheetData/s:row/s:c', namespaces=NS)}


def replace_cell(path, coordinate, mutate):
    with ZipFile(path) as archive:
        entries = [(item, archive.read(item.filename)) for item in archive.infolist()]
    with ZipFile(path, 'w') as archive:
        for item, raw in entries:
            if item.filename == 'xl/worksheets/sheet1.xml':
                root = etree.fromstring(raw)
                mutate(root.xpath(f'.//s:c[@r="{coordinate}"]', namespaces=NS)[0])
                raw = etree.tostring(root)
            archive.writestr(item, raw)


@pytest.mark.parametrize('raw,canonical', [('1,234', '1234'), ('-1,234.50', '-1234.5'), ('0.1', '0.1'),
                                        ('+001.2300', '1.23'), ('-0.00', '0'),
                                        ('999999999999999', '999999999999999'),
                                        ('0.123456789012345', '0.123456789012345'),
                                        ('1000000000000000', '1000000000000000')])
def test_typed_blank_cells_write_exact_numeric_xml_and_preserve_template(tmp_path, raw, canonical):
    source, profile = blank_form(tmp_path)
    values = {'금액': raw, '예산': '0.2', '등록번호': '001234', '날짜': '2024-02-29'}
    original = source.read_bytes()
    original_values = values.copy()
    output = fill_compatible_template(source, values, tmp_path / 'filled.xlsx', profile=profile)
    xml = cells(output)
    assert xml['B1'].get('t') == 'n' and xml['B1'].find('s:v', NS).text == canonical
    assert Decimal(xml['B1'].find('s:v', NS).text) == Decimal(raw.replace(',', ''))
    assert xml['B2'].get('t') == 'n' and xml['B5'].find('s:f', NS).text == 'SUM(B1:B2)'
    assert xml['B4'].get('t') == 'inlineStr' and xml['B3'].get('t') == 'inlineStr'
    assert verify_output(source, output, values, profile=profile)['status'] == 'passed'
    assert source.read_bytes() == original and values == original_values
    with ZipFile(source) as old, ZipFile(output) as new:
        assert old.namelist() == new.namelist()
        assert all(old.read(part) == new.read(part) for part in old.namelist()
                   if part != 'xl/worksheets/sheet1.xml')
    before, after = load_workbook(source), load_workbook(output)
    try:
        assert before.active['B1']._style == after.active['B1']._style
        assert before.active['B1'].number_format == after.active['B1'].number_format == '#,##0.00'
        assert after.active['B5'].value == '=SUM(B1:B2)'
        assert str(before.active.merged_cells) == str(after.active.merged_cells)
        assert after.active['B3'].value == '001234'
    finally:
        before.close()
        after.close()


def test_explicit_integer_without_unit_and_renamed_mapping_are_numeric(tmp_path):
    source, profile = blank_form(tmp_path, rule={'type': 'integer'})
    mapping = {field['id']: '인원' for field in profile['fields'] if field['value_key'] == '금액'}
    output = fill_compatible_template(source, {'인원': '-1,234'}, tmp_path / 'filled.xlsx', mapping=mapping, profile=profile)
    assert cells(output)['B1'].get('t') == 'n'
    assert verify_output(source, output, {'인원': '-1,234'}, profile=profile, mapping=mapping)['status'] == 'passed'


@pytest.mark.parametrize('rule,raw', [({'type': 'number'}, '123 [S1]'),
                                  ({'type': 'number', 'unit': '원'}, '123원'),
                                  ({'type': 'date'}, '2024-02-29')])
def test_citations_value_units_and_declared_dates_remain_original_text(tmp_path, rule, raw):
    source, profile = blank_form(tmp_path, rule=rule)
    values = {'금액': raw}
    output = fill_compatible_template(source, values, tmp_path / 'filled.xlsx', profile=profile)
    assert cells(output)['B1'].get('t') == 'inlineStr'
    assert ''.join(cells(output)['B1'].xpath('.//s:t/text()', namespaces=NS)) == raw
    assert verify_output(source, output, values, profile=profile)['status'] == 'passed'


def test_numeric_validation_on_mixed_placeholder_does_not_convert_the_cell(tmp_path):
    source, profile = blank_form(tmp_path)
    field = next(field for field in profile['fields'] if field['kind'] == 'xlsx_placeholder')
    field['validation'] = {'type': 'number'}
    # Its shared key must have the same rule as the ordinary numeric input.
    for other in profile['fields']:
        if other['value_key'] == '금액':
            other['validation'] = {'type': 'number'}
    values = {'금액': '1,234'}
    output = fill_compatible_template(source, values, tmp_path / 'filled.xlsx', profile=profile)
    assert cells(output)['B1'].get('t') == 'n'
    assert cells(output)['C1'].get('t') == 'inlineStr'
    assert verify_output(source, output, values, profile=profile)['status'] == 'passed'


@pytest.mark.parametrize('raw', ['1000000000000001', '0.1234567890123456', '9' * 309,
                               '0.' + '0' * 308 + '1'])
def test_excel_precision_and_range_are_explicitly_blocked_without_string_fallback(tmp_path, raw):
    source, profile = blank_form(tmp_path)
    target = tmp_path / 'filled.xlsx'
    with pytest.raises(TemplateError, match='Excel 숫자'):
        fill_compatible_template(source, {'금액': raw}, target, profile=profile)
    assert not target.exists()


@pytest.mark.parametrize('tamper', ['inlineStr', 's', 'wrong_value', 'rounded_float_collision', 'exponent'])
def test_saved_numeric_type_and_exact_xml_tampering_are_detected_independently(tmp_path, tamper):
    source, profile = blank_form(tmp_path)
    values = {'금액': '0.123456789012345'}
    output = fill_compatible_template(source, values, tmp_path / 'filled.xlsx', profile=profile)
    def change(cell):
        if tamper in {'inlineStr', 's'}:
            cell.set('t', tamper)
            cell.find('s:v', NS).text = '0' if tamper == 's' else '0.123456789012345'
        elif tamper == 'wrong_value':
            cell.find('s:v', NS).text = '0.223456789012345'
        elif tamper == 'rounded_float_collision':
            # Same Python float, different exact decimal: XML must catch this.
            cell.find('s:v', NS).text = '0.1234567890123450000000001'
        else:
            cell.find('s:v', NS).text = '1.23456789012345E-1'
    replace_cell(output, 'B1', change)
    with pytest.raises(ValueError):
        verify_output(source, output, values, profile=profile)


def test_independent_checker_does_not_call_the_filler_canonicalizer(tmp_path, monkeypatch):
    source, profile = blank_form(tmp_path)
    values = {'금액': '0.1'}
    output = fill_compatible_template(source, values, tmp_path / 'filled.xlsx', profile=profile)
    def fail(*args):
        raise AssertionError('filler called during independent check')
    monkeypatch.setattr('templates.compatibility._xlsx_numeric_text', fail)
    assert verify_output(source, output, values, profile=profile)['status'] == 'passed'


def test_decimal_canonicalization_does_not_round_under_low_ambient_precision(tmp_path):
    source, profile = blank_form(tmp_path)
    precision = getcontext().prec
    try:
        getcontext().prec = 2
        values = {'금액': '1234567890123.45'}
        output = fill_compatible_template(source, values, tmp_path / 'filled.xlsx', profile=profile)
        assert cells(output)['B1'].find('s:v', NS).text == '1234567890123.45'
        assert verify_output(source, output, values, profile=profile)['status'] == 'passed'
        assert getcontext().prec == 2
    finally:
        getcontext().prec = precision


def test_forged_numeric_mapping_cannot_write_a_formula_cell(tmp_path):
    source, profile = blank_form(tmp_path)
    fake = deepcopy(next(field for field in profile['fields'] if field['id'].endswith(':B1')))
    fake.update(id=fake['id'].replace(':B1', ':B5'), value_key='수식')
    profile['fields'] = [fake]
    with pytest.raises(TemplateError, match='수식'):
        fill_compatible_template(source, {'수식': '123'}, tmp_path / 'filled.xlsx', profile=profile)


def test_precision_as_displayed_setting_is_preserved_and_blocks_numeric_input(tmp_path):
    source, _ = blank_form(tmp_path)
    book = load_workbook(source)
    book.calculation = CalcProperties(fullPrecision=False)
    book.save(source)
    book.close()
    profile = analyze_template(source)
    for field in profile['fields']:
        if field['kind'] == 'xlsx_cell' and field['value_key'] == '금액':
            field['validation'] = {'type': 'number'}
    with pytest.raises(TemplateError, match='표시 정밀도'):
        fill_compatible_template(source, {'금액': '1.234'}, tmp_path / 'filled.xlsx', profile=profile)
    # Original string inputs remain supported without changing calculation settings.
    values = {'등록번호': '001', '금액': '1.234 [S1]'}
    output = fill_compatible_template(source, values, tmp_path / 'text.xlsx', profile=profile)
    assert verify_output(source, output, values, profile=profile)['status'] == 'passed'
    with pytest.raises(ValueError, match='표시 정밀도'):
        verify_output(source, output, dict(values, 금액='1.234'), profile=profile)
    book = load_workbook(output)
    assert book.calculation.fullPrecision is False
    book.close()


def test_actual_registered_history_form_has_numeric_grade_and_preserves_original(tmp_path):
    from pathlib import Path
    import json
    root = Path(__file__).resolve().parents[1]
    source = root / 'data/public_templates/government/history-competition-2026.xlsx'
    if not source.exists():
        pytest.skip('공개 원본의 재배포 없이 로컬 확보된 파일만 실제 회귀함')
    profile = json.loads((root / 'templates/profiles/history-competition-2026.json').read_text(encoding='utf-8'))
    original = source.read_bytes()
    values = profile['demo_values']
    output = fill_compatible_template(source, values, tmp_path / 'history.xlsx', profile=profile)
    xml = cells(output)
    assert xml['G9'].get('t') == 'n' and xml['G9'].find('s:v', NS).text == values['학년']
    assert verify_output(source, output, values, profile=profile)['status'] == 'passed'
    assert source.read_bytes() == original
