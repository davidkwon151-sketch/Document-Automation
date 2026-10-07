"""Data-first RA writing with explicit targets and guarded, current downloads."""
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import tempfile

import streamlit as st

from app.form_config import configure_profile, mapping_rows, selection_label, selection_options, value_rule_help
from app.multimodal_ui import render_multimodal_upload
from app.ra_workflow_ui import load_workflows, resolve_workflow
from agent.documents import DOCUMENT_KINDS
from llm.client import LLMClient, LLMError
from templates.compatibility import analyze_template


def fingerprint(value):
    return sha256(json.dumps(value, ensure_ascii=False, sort_keys=True, allow_nan=False).encode()).hexdigest()


def clear_outputs(*, include_proposal=True):
    names = ['ra_auto_result', 'ra_auto_exports', 'ra_auto_result_signature']
    if include_proposal:
        names.append('ra_auto_proposal')
    for name in names:
        st.session_state.pop(name, None)


def example_uploads():
    """Clearly synthetic inputs help try the actual parser/filler without a model."""
    from docx import Document
    document = Document()
    document.add_paragraph('프로젝트 제작 합성 시험 양식 · 실제 회사/기관 양식 아님')
    table = document.add_table(rows=4, cols=2)
    for row, label in zip(table.rows, ['제품명', '성상', '포장단위', '성명']):
        row.cells[0].text = label
        row.cells[1].text = '{{' + label + '}}'
    stream = BytesIO()
    document.save(stream)
    data = ('제품명,항목,값\nSyntheticDrugA,제품명,SyntheticDrugA\n'
            'SyntheticDrugA,성상,흰색 분말\nSyntheticDrugA,포장단위,합성 시험 포장\n').encode('utf-8-sig')
    return stream.getvalue(), data


def _error(exc):
    # Provider messages may contain credentials or private request text.
    if isinstance(exc, LLMError):
        labels = {'quota': 'API 잔액·할당량 부족', 'billing_limit': 'API 사용 한도 도달',
                  'rate_limit': 'API 일시 요청 속도 제한', 'configuration': '모델 연결 설정 필요',
                  'connection': '모델 서버 연결 실패', 'timeout': '모델 요청 시간 초과'}
        st.error('실제 AI 작성을 완료하지 못했습니다: ' + labels.get(exc.kind, '모델 응답 오류') +
                 '. 모델·결제 설정을 확인해 주세요. 작성 모드를 자동으로 바꾸지 않습니다.')
    elif isinstance(exc, ValueError):
        st.error(str(exc))
    else:
        st.error('작업을 완료하지 못했습니다. 양식과 원자료를 확인해 주세요.')


def uploaded_template(upload):
    """Keep only a session-owned temporary original; no values enter mapping caches."""
    name = upload.name
    if Path(name).name != name or '/' in name or '\\' in name:
        raise ValueError('양식 파일명에 경로를 포함할 수 없음')
    data = upload.getvalue()
    suffix = Path(name).suffix.lower()
    if suffix not in {'.docx', '.hwpx', '.xlsx', '.pdf'} or not isinstance(data, bytes) or not 0 < len(data) <= 30 * 1024 * 1024:
        raise ValueError('양식은 비어 있지 않은 30 MiB 이하 DOCX/HWPX/XLSX/PDF이어야 함')
    digest = sha256(data).hexdigest()
    if st.session_state.get('ra_auto_uploaded_sha') != digest:
        previous = st.session_state.pop('ra_auto_template_directory', None)
        if previous is not None:
            previous.cleanup()
        directory = tempfile.TemporaryDirectory(prefix='ra-auto-template-')
        path = Path(directory.name) / ('원본양식' + suffix)
        path.write_bytes(data)
        st.session_state.update(ra_auto_template_directory=directory, ra_auto_uploaded_path=str(path),
                                ra_auto_uploaded_sha=digest, ra_auto_raw_profile=analyze_template(path))
        clear_outputs()
    path = Path(st.session_state['ra_auto_uploaded_path'])
    if sha256(path.read_bytes()).hexdigest() != digest:
        raise ValueError('업로드한 양식 원본이 변경됨')
    return path, deepcopy(st.session_state['ra_auto_raw_profile'])


def _template_choice():
    kind = st.radio('양식 가져오기', ['등록된 RA 양식', '내 양식 업로드'], horizontal=True, key='ra_auto_template_kind')
    if kind == '등록된 RA 양식':
        entries = load_workflows()
        identifier = st.selectbox('사용할 양식', [entry['id'] for entry in entries],
                                  format_func=lambda value: next(entry['title'] for entry in entries if entry['id'] == value),
                                  key='ra_auto_registered_form')
        entry = next(row for row in entries if row['id'] == identifier)
        path, profile = resolve_workflow(entry)
        profile = {**profile, 'document_kind': profile.get('document_kind', entry.get('document_kind', 'application'))}
        st.caption('문서 종류: ' + DOCUMENT_KINDS[profile['document_kind']]['title'] + ' · 등록된 양식의 종류를 적용합니다.')
        st.caption(entry.get('coverage_note', '등록된 입력 위치만 작성하며, 원본의 나머지 항목은 별도 확인해야 합니다.'))
        st.download_button('빈 원본 양식', path.read_bytes(), file_name=path.name, key='ra_auto_original')
        return path, profile
    upload = st.file_uploader('내가 사용하는 사내외 양식', type=['docx', 'hwpx', 'xlsx', 'pdf'], key='ra_auto_template_upload')
    if upload is None:
        clear_outputs()
        st.info('원본 양식을 첨부하면 작성할 입력 위치를 분석합니다. HWP 원본은 HWP→HWPX 탭에서 먼저 변환해 주세요.')
        return None, None
    path, raw = uploaded_template(upload)
    for warning in raw.get('warnings', []):
        st.caption(warning)
    if not raw.get('fields'):
        clear_outputs()
        st.warning('자동으로 찾을 수 있는 입력칸이 없습니다. PDF 작성 영역·복잡한 배치는 기존 문서 작성 화면에서 먼저 지정해 주세요.')
        st.link_button('새 양식 작성 영역·매핑 화면 열기', 'http://127.0.0.1:8501/')
        return None, None
    digest = raw['source_sha256']
    kind = st.selectbox('내 양식의 문서 종류', list(DOCUMENT_KINDS), index=1,
                        format_func=lambda value: DOCUMENT_KINDS[value]['title'],
                        key='ra_auto_document_kind_' + digest[:16])
    raw['document_kind'] = kind
    st.caption(DOCUMENT_KINDS[kind]['style'])
    with st.expander('처음 쓰는 양식 · 입력칸 매핑 확인', expanded=True):
        st.caption('원본 위치와 항목을 확인해 채울 값 이름을 지정합니다. 개인정보·서명·승인은 직접 입력을 선택합니다.')
        rows = mapping_rows(raw)
        # Preserve detected item names as suggestions, never approve them automatically.
        for row, field in zip(rows, raw['fields']):
            if not row['채울 값']:
                row['채울 값'] = field['value_key']
        rows = st.data_editor(rows, disabled=['입력칸 ID', '항목'], hide_index=True,
                              key='ra_auto_mapping_' + digest[:16])
        profile, _ = configure_profile(raw, rows)
        profile['citation_mode'] = 'sidecar'
        token = fingerprint(profile)
        confirmed = st.checkbox('원본 입력 위치·항목명·필수 조건·직접 입력 구분을 확인함',
                                key='ra_auto_mapping_confirm_' + token[:16])
        if not confirmed:
            clear_outputs()
            st.info('새 양식의 매핑을 확인한 뒤 작성할 수 있습니다.')
            return None, None
        profile.update(user_confirmed=True, mapping_confirmed=True, auto_mapping_confirmed=True,
                       domain='pharmaceutical_ra', ra_workflow='product_approval')
        return path, profile


def _direct_inputs(profile, selected):
    values = {}
    fields = [field for field in profile['fields'] if field['value_key'] in selected and field.get('input_required')]
    if not fields:
        st.caption('선택한 항목에 직접 입력이 필요한 칸이 없습니다.')
    for field in fields:
        key = 'ra_auto_direct_' + fingerprint((profile['source_sha256'], field['id']))[:16]
        options = selection_options(field)
        help_text = value_rule_help(field)
        if options is not None:
            if field.get('multiselect'):
                value = json.dumps(st.multiselect(field['label'], options, key=key, help=help_text), ensure_ascii=False)
            else:
                value = st.selectbox(field['label'], options, key=key, help=help_text,
                                     format_func=lambda value, item=field: selection_label(item, value))
        else:
            value = st.text_input(field['label'], key=key, help=help_text)
        if value.strip() and value != '[]':
            values[field['value_key']] = value
    return values


def _show_result(result, profile):
    selected = [field for field in profile['fields'] if field.get('required')]
    draft = result.get('draft', {})
    filled = sum(bool(draft.get(field['value_key'], '').strip()) for field in selected)
    st.metric('이번 작성 대상 기입률', f'{filled}/{len(selected)}칸',
              help='선택 항목과 원본 필수 항목의 값 존재 비율입니다. 법정 제출 적합성이나 원본 전체 완료율이 아닙니다.')
    st.caption(f"실제 모델 요청 {result.get('actual_model_requests', 0)}회 · 제출 준비 완료: 판정하지 않음")
    for question in result.get('questions', [])[:2]:
        st.warning(question)
    review = result.get('review', {})
    for issue in [*review.get('issues', []), *review.get('warnings', [])]:
        if isinstance(issue, dict):
            st.warning(str(issue.get('message', issue.get('code', '검수 보류'))))
        else:
            st.warning(str(issue))
    if not draft:
        st.warning('아직 초안이 없습니다. 표시된 질문에 맞춰 원자료·직접 입력 또는 작성 요청의 목적·보고 대상을 보완하고 다시 작성해 주세요.')
        return {}
    edited = deepcopy(draft)
    token = fingerprint(draft)[:16]
    fields = {field['value_key'] for field in profile['fields']}
    for key, value in draft.items():
        if key not in fields:
            continue
        edited[key] = st.text_area(key + ' · 초안', value=value,
                                  key='ra_auto_draft_' + token + '_' + fingerprint(key)[:12],
                                  disabled=result.get('mode') == 'source_copy')
    return edited


def render_ra_auto():
    from app.ra_auto_service import target_profile, propose_ra_auto, generate_ra_auto, review_ra_auto, export_ra_auto

    st.subheader('데이터로 자동 작성')
    st.caption('1 양식 선택 → 2 데이터 첨부 → 3 직접 입력 → 4 일괄 작성 → 5 검수·다운로드')
    st.info('원문 일괄 기입은 확인한 원문을 옮기는 기능입니다. 실제 AI 작성은 모델을 호출합니다. 모든 결과는 담당자 검토용이며 제출 완료를 보장하지 않습니다.')
    with st.expander('빠른 시험 · 합성 양식과 데이터 내려받기'):
        template_example, data_example = example_uploads()
        st.caption('아래 파일은 프로젝트 제작 시험 자료이며 실제 의약품 정보나 공식 기관 양식이 아닙니다. 내 양식 업로드에 DOCX, 원자료에 CSV를 올리고 제품명 SyntheticDrugA를 입력해 시험할 수 있습니다.')
        st.download_button('합성 4칸 시험 양식 DOCX', template_example, file_name='합성시험양식.docx', key='ra_auto_example_template')
        st.download_button('합성 원자료 CSV', data_example, file_name='합성원자료.csv', key='ra_auto_example_data')
    try:
        st.markdown('**1. 양식과 이번 작성 대상**')
        path, original = _template_choice()
        if original is None:
            return
        form_identity = fingerprint((str(path), original))
        if st.session_state.get('ra_auto_form_identity') != form_identity:
            for key in list(st.session_state):
                if key.startswith(('ra_auto_direct_', 'ra_auto_draft_', 'ra_auto_final_', 'ra_auto_selection_')):
                    st.session_state.pop(key, None)
            clear_outputs()
            st.session_state.pop('ra_auto_active_metrics', None)
            st.session_state['ra_auto_form_identity'] = form_identity
        if st.button('새 문서 시작 · 최초 초안 기준 새로 기록', key='ra_auto_new_document'):
            clear_outputs()
            st.session_state.pop('ra_auto_active_metrics', None)
            for key in list(st.session_state):
                if key.startswith(('ra_auto_draft_', 'ra_auto_final_')):
                    st.session_state.pop(key, None)
            st.info('새 문서의 기준을 시작합니다. 현재 양식과 첨부 자료는 유지합니다.')
        names = list(dict.fromkeys(field['value_key'] for field in original['fields']))
        labels = {field['value_key']: field['label'] for field in original['fields']}
        default = [key for key in names if any(field['value_key'] == key and not field.get('input_required') for field in original['fields'])]
        chosen = st.multiselect('이번에 자동 작성할 항목', names, default=default,
                                format_func=lambda value: labels[value], key='ra_auto_selection_' + form_identity[:16])
        required = [field['value_key'] for field in original['fields'] if field.get('required')]
        selected = list(dict.fromkeys([*chosen, *required]))
        if not selected:
            clear_outputs()
            st.warning('이번에 작성할 항목을 하나 이상 선택해 주세요.')
            return
        if set(required) - set(chosen):
            st.caption('원본 필수 항목은 작성 대상에 자동 포함합니다: ' + ', '.join(labels[key] for key in required if key not in chosen))
        product = st.text_input('제품명 · 원자료에 명시된 이름', key='ra_auto_product')
        variant = st.text_input('제형·함량 · 원자료에 명시된 경우만 입력', key='ra_auto_variant')
        mode_label = st.radio('작성 방식', ['원문 일괄 기입', '실제 AI 작성'], horizontal=True, key='ra_auto_mode')
        mode = 'source_copy' if mode_label == '원문 일괄 기입' else 'live'
        if not product.strip():
            clear_outputs()
            st.info('다른 품목의 자료를 섞지 않도록 원자료의 제품명을 입력해 주세요.')
            return
        profile = target_profile(path, original, selected, product_name=product.strip(), variant=variant.strip())
        st.markdown('**2. 작성에 사용할 데이터**')
        intake = render_multimodal_upload(key_prefix='ra_auto_sources', context={
            'form': form_identity, 'profile': profile, 'product': product, 'variant': variant, 'mode': mode})
        if intake['changed']:
            clear_outputs()
        if not intake['ready']:
            clear_outputs()
            st.warning('사진·스캔 원문 확인 또는 전사문 저장을 마친 뒤 작성할 수 있습니다. 이전 출력은 사용할 수 없습니다.')
            return
        st.markdown('**3. 직접 입력이 필요한 항목만 입력**')
        direct = _direct_inputs(profile, selected)
        instruction = st.text_area('작성 요청', value='첨부 원자료를 근거로 선택한 양식의 작성 대상 항목을 작성해 주세요.', key='ra_auto_instruction')
        current = fingerprint({'template': form_identity, 'profile': profile, 'sources': intake['signature'],
                               'selected': selected, 'direct': direct, 'mode': mode, 'instruction': instruction})
        if st.session_state.get('ra_auto_input_signature') != current:
            clear_outputs()
            st.session_state['ra_auto_input_signature'] = current
        st.markdown('**4. 선택한 항목을 한 번에 작성**')
        confirmed_proposals = None
        can_generate = bool(intake['sources'] or direct) and bool(instruction.strip())
        if mode == 'source_copy':
            st.caption('항목명·제품·함량이 명확한 정확 인용 후보만 기입합니다. 여러 후보나 부족 자료는 자동으로 결정하지 않습니다.')
            if st.button('원자료에서 일괄 기입 후보 찾기', disabled=not can_generate, key='ra_auto_propose'):
                clear_outputs()
                st.session_state['ra_auto_proposal'] = propose_ra_auto(path, original, intake['sources'], selected,
                                                                  product_name=product.strip(), variant=variant.strip())
            proposal = st.session_state.get('ra_auto_proposal')
            if proposal:
                statuses = {'proposed': '원문 확인 필요', 'missing': '자료 부족', 'ambiguous': '복수 후보·보류',
                            'blocked': '검수 보류', 'direct_input': '직접 입력'}
                st.dataframe([{'항목': labels.get(row['value_key'], row['value_key']),
                               '상태': statuses.get(row['status'], row['status']),
                               '기입할 정확 원문': (row.get('binding') or {}).get('quote', ''),
                               '확인 사항': row.get('reason', '')} for row in proposal.get('fields', [])], hide_index=True)
                by_id = {source['source_id']: source for source in intake['sources']}
                with st.expander('일괄 기입할 원문·출처 위치 확인', expanded=True):
                    for key, binding in proposal.get('source_bindings', {}).items():
                        source = by_id[binding['source_id']]
                        st.write(labels.get(key, key) + ' → ' + binding['quote'])
                        st.caption(source['filename'] + ' · ' + str(source.get('location') or source.get('page') or source.get('sheet')))
                        st.code(source['text'], language=None)
                confirmed = st.checkbox('일괄 후보의 제품·항목·원문 위치와 직접 입력을 확인함',
                                        key='ra_auto_proposal_confirm_' + proposal['fingerprint'][:16])
                if confirmed:
                    confirmed_proposals = proposal['fingerprint']
                else:
                    clear_outputs(include_proposal=False)
            can_generate = can_generate and confirmed_proposals is not None
        else:
            st.caption('설정된 실제 모델로 작성·의미 검수를 수행합니다. 호출 실패 시 원문 기입·mock 결과로 자동 대체하지 않습니다.')
        if st.button('선택 항목 일괄 작성·검수', disabled=not can_generate, key='ra_auto_generate'):
            clear_outputs(include_proposal=False)
            with st.spinner('선택한 입력칸의 근거·제품·수치·필수 조건을 검수하고 있습니다.'):
                client = LLMClient() if mode == 'live' else None
                result = generate_ra_auto(instruction, path, original, intake['sources'], selected,
                                          product_name=product.strip(), variant=variant.strip(), mode=mode,
                                          client=client, field_values=direct, confirmed_proposals=confirmed_proposals,
                                          previous_metrics=st.session_state.get('ra_auto_active_metrics') if mode == 'live' else None)
            st.session_state.update(ra_auto_result=result, ra_auto_result_signature=current)
            if mode == 'live' and result.get('metrics'):
                st.session_state['ra_auto_active_metrics'] = deepcopy(result['metrics'])
        result = st.session_state.get('ra_auto_result')
        if not result or st.session_state.get('ra_auto_result_signature') != current:
            return
        st.markdown('**5. 초안 검수와 담당자 확인 후 다운로드**')
        edited = _show_result(result, profile)
        unsaved = edited != result.get('draft', {})
        if unsaved:
            st.session_state.pop('ra_auto_exports', None)
            st.warning('초안이 수정되었습니다. 현재 내용으로 다시 검수하기 전에는 다운로드할 수 없습니다.')
        if edited and mode == 'live' and st.button('현재 초안 다시 검수', key='ra_auto_review'):
            st.session_state.pop('ra_auto_exports', None)
            client = LLMClient() if mode == 'live' else None
            reviewed = review_ra_auto(result, edited, client=client)
            st.session_state['ra_auto_result'] = reviewed
            if reviewed.get('metrics'):
                st.session_state['ra_auto_active_metrics'] = deepcopy(reviewed['metrics'])
            st.rerun()
        if edited and mode == 'source_copy' and unsaved:
            st.caption('원문 기입 문장을 임의 편집한 결과는 승인하지 않습니다. 원자료 또는 선택 항목을 변경한 뒤 후보 확인부터 다시 진행해 주세요.')
        elif edited and mode == 'source_copy':
            st.caption('확인한 원문 기입값을 표시합니다. 내용을 바꾸려면 원자료 또는 선택 항목을 수정하고 후보를 다시 확인해 주세요.')
        review = result.get('review', {})
        ready = bool(edited) and not unsaved and not review.get('blocking', True) and result.get('status') == 'ready'
        confirmation_key = 'ra_auto_final_' + fingerprint((current, result))[:16]
        confirmed = st.checkbox('현재 기입값·출처·선택 대상과 남은 미기입 항목을 검토함', key=confirmation_key)
        if not confirmed or not ready:
            st.session_state.pop('ra_auto_exports', None)
        if st.button('검수한 문서 파일 만들기', disabled=not (ready and confirmed), key='ra_auto_export'):
            st.session_state['ra_auto_exports'] = export_ra_auto(result, confirmed=True)
        exported = st.session_state.get('ra_auto_exports')
        if exported and ready and confirmed:
            st.download_button('작성 문서 다운로드', exported['document'], file_name=exported['filename'], key='ra_auto_download')
            st.download_button('원자료·검수 기록 다운로드', exported['evidence'], file_name='작성_근거검수.json', key='ra_auto_evidence')
        if not ready:
            st.caption('부족 자료·빈 작성 대상·검수 오류가 있으면 출력을 보류합니다. 원본 전체 항목과 서명·첨부 요건은 담당자가 별도 확인해야 합니다.')
        st.caption('실사용 수정률·상사 반려 횟수는 이 화면에서 측정 완료로 표시하지 않습니다. 반려·재작업 기록은 기존 문서 작성 화면에서 지원합니다.')
    except Exception as exc:
        clear_outputs()
        _error(exc)


render = render_ra_auto
