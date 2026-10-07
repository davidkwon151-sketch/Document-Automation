"""삽입 이동과 복사 참조를 구분하고 원본 ZIP 부품을 보존해야 함."""

from copy import deepcopy
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.table import Table
from openpyxl.workbook.defined_name import DefinedName
import pytest

from templates.fill import TemplateError
from templates.repeat_xlsx import inventory, transform, S, P, X14, XM


NS={'s':S}
ROOT=Path(__file__).resolve().parents[1]


def fixture(*,table=False,chart=False,protected_other=False):
    workbook=Workbook(); sheet=workbook.active; sheet.title='입력 양식'
    sheet.append(['고정 제목'])
    sheet.append(['품목','수량','단가','금액','선택'])
    sheet['A3']='{{품목}}'
    for address in ['B3','C3','E3','F3','G3']: sheet[address].fill=PatternFill('solid',fgColor='DDDDDD')
    sheet['D3']='=B3*C3+$G$5+Other!A3'
    sheet['A4']='고정 기존 행'; sheet['B4']=2; sheet['C4']=3; sheet['D4']='=B4*C4'
    sheet['A5']='합계'; sheet['D5']='=SUM(D3:D4)'; sheet['G5']=50
    sheet['H5']='=SUM(B3:B3)'
    sheet.merge_cells('F3:G3'); sheet.merge_cells('F4:G4')
    sheet.row_dimensions[3].height=27
    sheet.freeze_panes='B5'; sheet.print_area='A1:H5'
    sheet.row_breaks.append(Break(id=4))
    whole=DataValidation(type='whole',operator='between',formula1='0',formula2='100')
    whole.add('B3:C3'); sheet.add_data_validation(whole)
    choice=DataValidation(type='list',formula1='"A,B"'); choice.add('E3'); sheet.add_data_validation(choice)
    sheet.conditional_formatting.add('B3:B4',FormulaRule(formula=['B3<0'],fill=PatternFill('solid',fgColor='FF0000')))
    other=workbook.create_sheet('Other')
    other['A1']="='입력 양식'!$D$5+SUM('입력 양식'!B3:B3)"
    other['A3']=4; other['A4']=5; other['A5']=6
    if protected_other:
        locked=workbook.create_sheet('Locked'); locked['A1']='=1+2'; locked.protection.sheet=True
    workbook.defined_names.add(DefinedName('InputRows',attr_text="'입력 양식'!$B$3:$C$4"))
    workbook.defined_names.add(DefinedName('OneRow',attr_text="'입력 양식'!$B$3:$B$3"))
    if table:
        item=Table(displayName='InputTable',ref='A2:E5'); item.totalsRowCount=1; sheet.add_table(item)
    if chart:
        item=BarChart(); item.add_data(Reference(sheet,min_col=2,min_row=2,max_row=4),titles_from_data=True)
        sheet.add_chart(item,'J7')
    stream=BytesIO(); workbook.save(stream)
    with ZipFile(BytesIO(stream.getvalue())) as archive:
        return {name:archive.read(name) for name in archive.namelist()}


def plan(row=3,count=3):
    return {'table_id':'xlsx:xl/worksheets/sheet1.xml','row':row,'count':count}


def result(parts,command=None):
    changes=transform(parts,command or plan())
    return {name:changes.get(name,value) for name,value in parts.items() if changes.get(name,value) is not None},changes


def root(parts,name='xl/worksheets/sheet1.xml'):
    return etree.fromstring(parts[name])


def formula(parts,address,sheet='xl/worksheets/sheet1.xml'):
    return root(parts,sheet).xpath(f'.//s:c[@r="{address}"]/s:f/text()',namespaces=NS)[0]


def replace_xml(parts,name,mutate):
    node=root(parts,name); mutate(node); parts[name]=etree.tostring(node)


def save(parts,path):
    with ZipFile(path,'w') as archive:
        for name,data in parts.items(): archive.writestr(name,data)


def test_inventory_uses_real_sparse_excel_row_numbers_and_never_modifies_parts():
    parts=fixture(); before=deepcopy(parts)
    replace_xml(parts,'xl/worksheets/sheet1.xml',lambda xml:xml.find(f'{{{S}}}sheetData').remove(xml.find(f'{{{S}}}sheetData/{{{S}}}row')))
    before=deepcopy(parts)
    record=inventory(parts)[0]
    assert record['label']=='입력 양식' and record['rows'][0]['index']==2
    assert record['row_count']==4 and record['max_row']==5
    assert next(row for row in record['rows'] if row['index']==3)['editable']
    assert not next(row for row in record['rows'] if row['index']==5)['editable']
    assert parts==before


def test_formula_move_then_copy_respects_absolute_and_other_sheet_references():
    parts=fixture(); before=deepcopy(parts)
    output,changes=result(parts)
    assert formula(output,'D3')=='B3*C3+$G$7+Other!A3'
    assert formula(output,'D4')=='B4*C4+$G$7+Other!A4'
    assert formula(output,'D5')=='B5*C5+$G$7+Other!A5'
    assert formula(output,'D6')=='B6*C6'
    assert formula(output,'D7')=='SUM(D3:D6)'
    assert formula(output,'H7')=='SUM(B3:B5)'
    assert formula(output,'A1','xl/worksheets/sheet2.xml')=="'입력 양식'!$D$7+SUM('입력 양식'!B3:B5)"
    assert parts==before
    for name in parts:
        if name not in changes: assert output[name]==parts[name]
    assert output['xl/styles.xml']==parts['xl/styles.xml']


def test_row_style_merges_dv_cf_print_ranges_names_pane_and_breaks_are_preserved():
    parts=fixture(); output,_=result(parts)
    xml=root(output); source=root(parts)
    rows=xml.findall(f'{{{S}}}sheetData/{{{S}}}row')
    assert [int(row.get('r')) for row in rows]==list(range(1,8))
    prototype=source.xpath('.//s:row[@r="3"]',namespaces=NS)[0]
    for number in [3,4,5]:
        copied=xml.xpath(f'.//s:row[@r="{number}"]',namespaces=NS)[0]
        assert {key:value for key,value in copied.attrib.items() if key!='r'}=={key:value for key,value in prototype.attrib.items() if key!='r'}
        assert copied.xpath('./s:c[@r="B'+str(number)+'"]/@s',namespaces=NS)==prototype.xpath('./s:c[@r="B3"]/@s',namespaces=NS)
    assert set(xml.xpath('./s:mergeCells/s:mergeCell/@ref',namespaces=NS))=={'F3:G3','F4:G4','F5:G5','F6:G6'}
    assert xml.xpath('./s:mergeCells/@count',namespaces=NS)==['4']
    assert xml.xpath('.//s:dataValidation/@sqref',namespaces=NS)==['B3:C5','E3:E5']
    assert xml.xpath('./s:conditionalFormatting/@sqref',namespaces=NS)==['B3:B6']
    assert xml.xpath('.//s:cfRule/s:formula/text()',namespaces=NS)==['B3<0']
    assert xml.xpath('.//s:pane/@topLeftCell',namespaces=NS)==['B7']
    assert xml.xpath('.//s:pane/@ySplit',namespaces=NS)==['6']
    assert xml.xpath('./s:rowBreaks/s:brk/@id',namespaces=NS)==['6']
    names=root(output,'xl/workbook.xml').xpath('.//s:definedName/text()',namespaces=NS)
    assert "'입력 양식'!$B$3:$C$6" in names and "'입력 양식'!$B$3:$B$5" in names
    assert "'입력 양식'!$A$1:$H$7" in names


def test_table_and_chart_ranges_and_anchors_update_without_losing_assets(tmp_path):
    parts=fixture(table=True,chart=True); output,changes=result(parts)
    table=root(output,'xl/tables/table1.xml')
    assert table.get('ref')=='A2:E7' and table.find(f'{{{S}}}autoFilter').get('ref')=='A2:E7'
    assert table.xpath('./s:tableStyleInfo/@*',namespaces=NS)==root(parts,'xl/tables/table1.xml').xpath('./s:tableStyleInfo/@*',namespaces=NS)
    chart=root(output,'xl/charts/chart1.xml')
    assert "'입력 양식'!$B$3:$B$6" in chart.xpath('.//*[local-name()="f"]/text()')
    assert not chart.xpath('.//*[local-name()="numCache" or local-name()="strCache"]')
    drawing=root(output,'xl/drawings/drawing1.xml')
    assert drawing.xpath('.//*[local-name()="from"]/*[local-name()="row"]/text()')==['8']
    for name in parts:
        if name.endswith('.rels') or name=='xl/styles.xml': assert output[name]==parts[name]
    target=tmp_path/'prepared.xlsx'; save(output,target)
    actual=load_workbook(target,data_only=False)
    assert actual['입력 양식']['D5'].value=='=B5*C5+$G$7+Other!A5'
    assert actual['입력 양식'].tables['InputTable'].ref=='A2:E7'


def test_formula_caches_are_removed_and_recalculation_request_does_not_change_calc_mode():
    parts=fixture()
    for name in ['xl/worksheets/sheet1.xml','xl/worksheets/sheet2.xml']:
        def cache(xml):
            for cell in xml.xpath('.//s:c[s:f]',namespaces=NS): cell.find(f'{{{S}}}v').text='9999'
        replace_xml(parts,name,cache)
    replace_xml(parts,'xl/workbook.xml',lambda xml:xml.find(f'{{{S}}}calcPr').set('calcMode','manual'))
    output,_=result(parts)
    assert not root(output).xpath('.//s:c[s:f]/s:v',namespaces=NS)
    assert not root(output,'xl/worksheets/sheet2.xml').xpath('.//s:c[s:f]/s:v',namespaces=NS)
    calc=root(output,'xl/workbook.xml').find(f'{{{S}}}calcPr')
    assert calc.get('calcMode')=='manual' and calc.get('forceFullCalc')=='1' and calc.get('fullCalcOnLoad')=='1'


def test_calc_chain_only_deletes_chain_and_its_exact_relationship_and_content_type():
    parts=fixture()
    parts['xl/calcChain.xml']=f'<calcChain xmlns="{S}"><c r="D5" i="1"/></calcChain>'.encode()
    rel=root(parts,'xl/_rels/workbook.xml.rels')
    etree.SubElement(rel,f'{{{P}}}Relationship',Id='chain',Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/calcChain',Target='calcChain.xml')
    parts['xl/_rels/workbook.xml.rels']=etree.tostring(rel)
    content=root(parts,'[Content_Types].xml')
    etree.SubElement(content,'{'+content.nsmap[None]+'}Override',PartName='/xl/calcChain.xml',ContentType='application/vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml')
    parts['[Content_Types].xml']=etree.tostring(content)
    output,changes=result(parts)
    assert changes['xl/calcChain.xml'] is None and 'xl/calcChain.xml' not in output
    assert not root(output,'xl/_rels/workbook.xml.rels').xpath('//*[@Id="chain"]')
    assert not root(output,'[Content_Types].xml').xpath('//*[@PartName="/xl/calcChain.xml"]')
    assert len(root(output,'xl/_rels/workbook.xml.rels'))==len(rel)-1
    assert len(root(output,'[Content_Types].xml'))==len(content)-1


def test_unrelated_protected_sheet_is_byte_identical_but_changed_protected_formula_blocks():
    parts=fixture(protected_other=True)
    output,_=result(parts)
    assert output['xl/worksheets/sheet3.xml']==parts['xl/worksheets/sheet3.xml']
    replace_xml(parts,'xl/worksheets/sheet3.xml',lambda xml:xml.find(f'.//{{{S}}}f').__setattr__('text',"'입력 양식'!D5"))
    with pytest.raises(TemplateError,match='보호'): transform(parts,plan())


def test_x14_dv_sqref_full_height_and_qualified_formula_preserve_options_and_uid():
    parts=fixture()
    def extend(xml):
        container=etree.SubElement(xml,f'{{{S}}}extLst')
        ext=etree.SubElement(container,f'{{{S}}}ext',uri='{CCE6A557-97BC-4B89-ADB6-D9C93CAAB3DF}')
        values=etree.SubElement(ext,f'{{{X14}}}dataValidations',count='1')
        native=etree.SubElement(values,f'{{{X14}}}dataValidation',type='list',allowBlank='1',uid='same-id')
        formula=etree.SubElement(native,f'{{{X14}}}formula1'); etree.SubElement(formula,f'{{{XM}}}f').text="'입력 양식'!$B$4:$B$5"
        etree.SubElement(native,f'{{{XM}}}sqref').text='E3 E4:E1048576'
    replace_xml(parts,'xl/worksheets/sheet1.xml',extend)
    assert next(row for row in inventory(parts)[0]['rows'] if row['index']==3)['editable']
    output,_=result(parts)
    xml=root(output)
    assert xml.xpath('.//*[local-name()="sqref"]/text()')==['E3:E5 E6:E1048576']
    assert xml.xpath('.//*[local-name()="formula1"]/*[local-name()="f"]/text()')==["'입력 양식'!$B$6:$B$7"]
    assert xml.xpath('.//*[local-name()="dataValidation" and @uid]/@uid')==['same-id']


@pytest.mark.parametrize('count',[0,201,True,'3'])
def test_invalid_count_is_not_coerced(count):
    with pytest.raises(TemplateError): transform(fixture(),plan(count=count))


@pytest.mark.parametrize('change',['signature','protection','workbook_protection','array','shared','external','indirect','offset','metadata','unknown_ext','vertical_merge','row_missing','duplicate_row','cell_mismatch','last_row'])
def test_unsafe_sources_or_invalid_addresses_never_publish_changes(change):
    parts=fixture()
    command=plan()
    def mutate(xml):
        if change=='protection': etree.SubElement(xml,f'{{{S}}}sheetProtection',sheet='1')
        elif change in {'array','shared'}: xml.find(f'.//{{{S}}}f').set('t',change)
        elif change in {'indirect','offset'}: xml.find(f'.//{{{S}}}f').text=change.upper()+'("A3")'
        elif change=='metadata': xml.find(f'.//{{{S}}}c').set('cm','1')
        elif change=='unknown_ext': etree.SubElement(etree.SubElement(xml,f'{{{S}}}extLst'),f'{{{S}}}ext',uri='unknown')
        elif change=='vertical_merge': etree.SubElement(xml.find(f'{{{S}}}mergeCells'),f'{{{S}}}mergeCell',ref='H2:H3')
        elif change=='duplicate_row': xml.find(f'{{{S}}}sheetData').append(deepcopy(xml.find(f'{{{S}}}sheetData/{{{S}}}row')))
        elif change=='cell_mismatch': xml.find(f'.//{{{S}}}c').set('r','A999')
        elif change=='last_row': etree.SubElement(xml.find(f'{{{S}}}sheetData'),f'{{{S}}}row',r='1048576')
    if change=='signature': parts['_xmlsignatures/sig1.xml']=b'signed'
    elif change=='external': parts['xl/externalLinks/externalLink1.xml']=b'external'
    elif change=='workbook_protection': replace_xml(parts,'xl/workbook.xml',lambda xml:xml.find(f'{{{S}}}workbookProtection').set('lockStructure','1'))
    elif change=='row_missing': command['row']=99
    else: replace_xml(parts,'xl/worksheets/sheet1.xml',mutate)
    before=deepcopy(parts)
    with pytest.raises(TemplateError): transform(parts,command)
    assert parts==before


def test_count_one_is_a_byte_identical_noop():
    parts=fixture(); assert transform(parts,plan(count=1))=={}


@pytest.mark.parametrize('filename,row',[
    ('business_mss_india_application_2026.xlsx',6),
    ('business_kotra_salesforce_2026.xlsx',6),
])
def test_official_blank_business_forms_expand_known_native_input_rows(filename,row,tmp_path):
    source=ROOT/'data/public_templates/business'/filename
    if not source.exists(): pytest.skip('공개 원본 파일이 이 체크아웃에 없음')
    before=source.read_bytes()
    import json
    metadata=json.loads(source.with_suffix(source.suffix+'.source.json').read_text(encoding='utf-8'))
    assert metadata['sha256']==sha256(before).hexdigest() and metadata['resource_kind']=='blank_form'
    with ZipFile(BytesIO(before)) as archive: parts={name:archive.read(name) for name in archive.namelist()}
    command={'table_id':'xlsx:xl/worksheets/sheet1.xml','row':row,'count':3}
    output,changes=result(parts,command)
    assert changes and len(root(output).findall(f'{{{S}}}sheetData/{{{S}}}row'))==len(root(parts).findall(f'{{{S}}}sheetData/{{{S}}}row'))+2
    assert source.read_bytes()==before and sha256(source.read_bytes()).hexdigest()==sha256(before).hexdigest()
    assert output['xl/styles.xml']==parts['xl/styles.xml']
    target=tmp_path/'prepared.xlsx'; save(output,target)
    assert load_workbook(target,data_only=False).sheetnames


def test_formula_whitespace_and_string_literals_are_preserved_in_move_and_copy():
    parts=fixture()
    replace_xml(parts,'xl/worksheets/sheet1.xml',lambda xml:xml.xpath('.//s:c[@r="D3"]/s:f',namespaces=NS)[0].__setattr__('text','B3  *  C3 + $G$5 + LEN("A3   B5")'))
    output,_=result(parts)
    assert formula(output,'D3')=='B3  *  C3 + $G$7 + LEN("A3   B5")'
    assert formula(output,'D4')=='B4  *  C4 + $G$7 + LEN("A3   B5")'


def test_table_column_formulas_move_and_structured_row_refs_remain_native():
    parts=fixture(table=True)
    def columns(xml):
        column=xml.findall(f'{{{S}}}tableColumns/{{{S}}}tableColumn')[3]
        etree.SubElement(column,f'{{{S}}}calculatedColumnFormula').text='[@수량]*[@단가]+$G$5'
        etree.SubElement(column,f'{{{S}}}totalsRowFormula').text='SUM(InputTable[금액])+$G$5'
    replace_xml(parts,'xl/tables/table1.xml',columns)
    output,_=result(parts)
    assert root(output,'xl/tables/table1.xml').xpath('.//s:calculatedColumnFormula/text()',namespaces=NS)==['[@수량]*[@단가]+$G$7']
    assert root(output,'xl/tables/table1.xml').xpath('.//s:totalsRowFormula/text()',namespaces=NS)==['SUM(InputTable[금액])+$G$7']


def test_full_height_validation_caps_only_range_end_and_preserves_full_column_countif():
    parts=fixture()
    def validation(xml):
        node=etree.SubElement(xml.find(f'{{{S}}}dataValidations'),f'{{{S}}}dataValidation',type='custom',sqref='P5:Q1048576 N5:N1048576')
        etree.SubElement(node,f'{{{S}}}formula1').text='COUNTIF($P:$P,O4)<2'
    replace_xml(parts,'xl/worksheets/sheet1.xml',validation)
    output,_=result(parts)
    native=root(output).xpath('.//s:dataValidation[@type="custom"]',namespaces=NS)[0]
    assert native.get('sqref')=='P7:Q1048576 N7:N1048576'
    assert native.find(f'{{{S}}}formula1').text=='COUNTIF($P:$P,O6)<2'


@pytest.mark.parametrize('reference',['A1048576','SUM(A1:A1048576)'])
def test_out_of_grid_single_or_copied_formula_references_are_not_silently_truncated(reference):
    parts=fixture()
    replace_xml(parts,'xl/worksheets/sheet1.xml',lambda xml:xml.find(f'.//{{{S}}}f').__setattr__('text',reference))
    with pytest.raises(TemplateError,match='행|주소'): transform(parts,plan())


def test_prototype_hyperlinks_copy_relationship_without_changing_the_relationship_part():
    parts=fixture()
    def hyperlink(xml):
        nodes=etree.SubElement(xml,f'{{{S}}}hyperlinks')
        etree.SubElement(nodes,f'{{{S}}}hyperlink',ref='A3',location="'입력 양식'!H5",display='기존 안내')
    replace_xml(parts,'xl/worksheets/sheet1.xml',hyperlink)
    output,_=result(parts)
    links=root(output).xpath('./s:hyperlinks/s:hyperlink',namespaces=NS)
    assert [(node.get('ref'),node.get('location'),node.get('display')) for node in links]==[
        ('A3',"'입력 양식'!H7",'기존 안내'),('A4',"'입력 양식'!H7",'기존 안내'),('A5',"'입력 양식'!H7",'기존 안내')]


def test_xlsx_fixed_numbered_public_rows_stay_blocked_even_with_supported_x14_choices():
    source=ROOT/'data/public_templates/business/business_kotra_usedcars_2026.xlsx'
    if not source.exists(): pytest.skip('공개 원본 파일이 이 체크아웃에 없음')
    with ZipFile(source) as archive: parts={name:archive.read(name) for name in archive.namelist()}
    row=next(item for item in inventory(parts)[0]['rows'] if item['index']==2)
    assert not row['editable'] and '순번' in row['reason']
    with pytest.raises(TemplateError,match='순번'): transform(parts,plan(row=2))


def test_selected_drawing_boundary_is_rejected_and_distant_drawings_move():
    parts=fixture(chart=True)
    replace_xml(parts,'xl/drawings/drawing1.xml',lambda xml:xml.xpath('.//*[local-name()="from"]/*[local-name()="row"]')[0].__setattr__('text','2'))
    with pytest.raises(TemplateError,match='그림|차트'): transform(parts,plan())


def test_last_allowed_count_and_plan_schema_do_not_change_original():
    parts=fixture(); before=deepcopy(parts)
    output,_=result(parts,plan(count=200))
    assert len(root(output).findall(f'{{{S}}}sheetData/{{{S}}}row'))==204
    assert parts==before
    command=plan(); command['extra']=True
    with pytest.raises(TemplateError): transform(parts,command)


def test_other_sheet_internal_hyperlinks_and_case_insensitive_sheet_references_move():
    parts=fixture()
    def hyperlink(xml):
        nodes=etree.SubElement(xml,f'{{{S}}}hyperlinks')
        etree.SubElement(nodes,f'{{{S}}}hyperlink',ref='A3',location="'입력 양식'!H5")
        xml.find(f'.//{{{S}}}f').text="SUM('입력 양식'!B3:B3)+OTHER!A1"
    replace_xml(parts,'xl/worksheets/sheet2.xml',hyperlink)
    output,_=result(parts)
    other=root(output,'xl/worksheets/sheet2.xml')
    assert other.xpath('.//s:hyperlink/@ref',namespaces=NS)==['A3']
    assert other.xpath('.//s:hyperlink/@location',namespaces=NS)==["'입력 양식'!H7"]
    # 영어 대소문자 시트명은 같은 원본 시트를 가리킴.
    def rename(xml):
        xml.find(f'{{{S}}}sheets/{{{S}}}sheet').set('name','Main')
        for name in xml.findall(f'{{{S}}}definedNames/{{{S}}}definedName'):
            name.text=name.text.replace("'입력 양식'!","Main!")
    replace_xml(parts,'xl/workbook.xml',rename)
    replace_xml(parts,'xl/worksheets/sheet2.xml',lambda xml:xml.find(f'.//{{{S}}}hyperlink').set('location','Main!H5'))
    replace_xml(parts,'xl/worksheets/sheet2.xml',lambda xml:xml.find(f'.//{{{S}}}f').__setattr__('text',"main!D5+MAIN!B3"))
    output,_=result(parts)
    assert formula(output,'A1','xl/worksheets/sheet2.xml')=='main!D7+MAIN!B3'


def test_modified_internal_hyperlink_on_protected_other_sheet_is_blocked():
    parts=fixture(protected_other=True)
    def hyperlink(xml):
        nodes=etree.SubElement(xml,f'{{{S}}}hyperlinks')
        etree.SubElement(nodes,f'{{{S}}}hyperlink',ref='A3',location="'입력 양식'!D5")
    replace_xml(parts,'xl/worksheets/sheet3.xml',hyperlink)
    with pytest.raises(TemplateError,match='보호'): transform(parts,plan())


@pytest.mark.parametrize('cy,allowed',[(381000,True),(381001,False),(2000000,False)])
def test_above_one_cell_drawing_uses_height_and_row_offset_not_only_anchor_row(cy,allowed):
    parts=fixture(chart=True)
    def drawing(xml):
        anchor=xml[0]
        anchor.find('{http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing}from/{http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing}row').text='0'
        anchor.find('{http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing}ext').set('cy',str(cy))
    replace_xml(parts,'xl/drawings/drawing1.xml',drawing)
    if allowed:
        output,_=result(parts)
        assert output['xl/drawings/drawing1.xml']==parts['xl/drawings/drawing1.xml']
    else:
        with pytest.raises(TemplateError,match='높이'): transform(parts,plan())


@pytest.mark.parametrize('y,cy,allowed',[(0,381000,True),(381000,100,False),(500000,100,False)])
def test_absolute_drawing_requires_explicit_bounds_fully_above_prototype(y,cy,allowed):
    parts=fixture(chart=True)
    def drawing(xml):
        anchor=xml[0]; anchor.tag='{http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing}absoluteAnchor'
        anchor.remove(anchor[0])
        position=etree.Element('{http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing}pos',x='0',y=str(y)); anchor.insert(0,position)
        anchor.find('{http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing}ext').set('cy',str(cy))
    replace_xml(parts,'xl/drawings/drawing1.xml',drawing)
    if allowed:
        output,_=result(parts)
        assert output['xl/drawings/drawing1.xml']==parts['xl/drawings/drawing1.xml']
    else:
        with pytest.raises(TemplateError,match='절대'): transform(parts,plan())


def test_unknown_row_height_does_not_guess_above_drawing_fits():
    parts=fixture(chart=True)
    replace_xml(parts,'xl/drawings/drawing1.xml',lambda xml:xml.xpath('.//*[local-name()="from"]/*[local-name()="row"]')[0].__setattr__('text','0'))
    replace_xml(parts,'xl/worksheets/sheet1.xml',lambda xml:xml.remove(xml.find(f'{{{S}}}sheetFormatPr')))
    with pytest.raises(TemplateError,match='높이'): transform(parts,plan())
