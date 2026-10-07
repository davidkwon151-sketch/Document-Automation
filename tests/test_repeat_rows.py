from copy import deepcopy
from hashlib import sha256
import json
from pathlib import Path
from zipfile import ZipFile

from docx import Document
import pytest
from streamlit.testing.v1 import AppTest

from agent.output_check import verify_output
from agent.template_learning import learn_template
from templates import fill_compatible_template
from templates.repeat_rows import (inspect_repeat_tables, prepare_repeat_template,
                                   expansion_binding, verify_prepared_template)


def form(tmp_path):
    document = Document()
    document.add_paragraph('표 밖 {{이름}} · 원본 안내')
    table = document.add_table(rows=3, cols=2)
    table.rows[0].cells[0].text = '품목'
    table.rows[0].cells[1].text = '검토 내역'
    paragraph = table.rows[1].cells[0].paragraphs[0]
    for text, bold in [('제품: {{이', True), ('름}} / {{코드}}', False)]:
        paragraph.add_run(text).bold = bold
    table.rows[1].cells[1].text = '□ {{내용}}\n○ {{상태}}'
    table.rows[2].cells[0].text = '합계'
    table.rows[2].cells[1].text = '고정 문구 유지'
    source = tmp_path / 'source.docx'
    document.save(source)
    plan = {'table_id': inspect_repeat_tables(source)[0]['id'], 'row': 2, 'count': 3}
    return source, plan


def prepared_form(tmp_path):
    from templates.repeat_fields import repeat_profile
    source, plan = form(tmp_path)
    prepared = tmp_path / 'prepared.docx'
    metadata = prepare_repeat_template(source, prepared, plan)
    profile = repeat_profile(prepared, plan, expansion_binding(metadata))
    return source, prepared, metadata, profile


def test_full_row_fill_has_distinct_values_and_preserves_prefixes_fonts_source(tmp_path):
    source, prepared, metadata, profile = prepared_form(tmp_path)
    original_bytes = source.read_bytes()
    repeated = [field for field in profile['fields'] if field.get('repeat_info')]
    assert len(repeated) == 12
    assert len({field['value_key'] for field in repeated}) == 12
    values = {field['value_key']: f"값 {field['repeat_info']['row']}-{field['placeholder_key']}"
              if field.get('repeat_info') else '공통 이름' for field in profile['fields']}
    # A literal token inside a user value must survive without being interpreted.
    values[repeated[0]['value_key']] = '제품 A {{사용자 문자}}'
    output = tmp_path / 'filled.docx'
    fill_compatible_template(prepared, values, output, profile=profile)
    checked = verify_output(prepared, output, values, profile=profile)
    assert checked['status'] == 'passed'
    result = Document(output)
    assert result.paragraphs[0].text == '표 밖 공통 이름 · 원본 안내'
    assert len(result.tables[0].rows) == 5
    for index, row in enumerate(result.tables[0].rows[1:4], 1):
        assert row.cells[0].text.startswith('제품: ')
        assert ' / ' in row.cells[0].text
        assert row.cells[1].text.startswith('□ ')
        assert '\n○ ' in row.cells[1].text
        assert any(run.bold for run in row.cells[0].paragraphs[0].runs)
    assert result.tables[0].rows[-1].cells[1].text == '고정 문구 유지'
    assert source.read_bytes() == original_bytes
    assert verify_prepared_template(prepared, profile, metadata)['status'] == 'passed'


@pytest.mark.parametrize('mutation', ['source', 'prepared', 'plan', 'profile_sha', 'binding', 'path'])
def test_prepared_binding_rejects_stale_or_forged_expansion(tmp_path, mutation):
    source, prepared, metadata, profile = prepared_form(tmp_path)
    if mutation in {'source', 'prepared'}:
        path = source if mutation == 'source' else prepared
        with ZipFile(path) as archive:
            parts = {name: archive.read(name) for name in archive.namelist()}
        parts['word/document.xml'] = parts['word/document.xml'].replace('고정 문구'.encode(), '바뀐 문구'.encode())
        with ZipFile(path, 'w') as archive:
            for name, data in parts.items():
                archive.writestr(name, data)
    elif mutation == 'plan':
        metadata['plan']['count'] = 4
    elif mutation == 'profile_sha':
        profile['source_sha256'] = '0' * 64
    elif mutation == 'binding':
        profile['repeat_expansion']['original_sha256'] = '0' * 64
    else:
        metadata['prepared_path'] = str(source)
    with pytest.raises(ValueError):
        verify_prepared_template(prepared, profile, metadata)


def test_failed_independent_check_keeps_existing_target_and_source(tmp_path, monkeypatch):
    import agent.repeat_check as checks
    source, plan = form(tmp_path)
    original = source.read_bytes()
    target = tmp_path / 'kept.docx'
    target.write_bytes(b'previous output')
    def fail(*args):
        raise ValueError('independent check rejected')
    monkeypatch.setattr(checks, 'verify_repeat_expansion', fail)
    with pytest.raises(ValueError, match='rejected'):
        prepare_repeat_template(source, target, plan)
    assert target.read_bytes() == b'previous output'
    assert source.read_bytes() == original
    assert set(tmp_path.iterdir()) == {source, target}


@pytest.mark.parametrize('mutation', ['offset', 'anchor', 'key', 'id', 'duplicate'])
def test_repeat_placeholders_reject_scope_forgery_or_duplicate_slots(tmp_path, mutation):
    _, prepared, _, profile = prepared_form(tmp_path)
    field = next(item for item in profile['fields'] if item.get('repeat_info'))
    if mutation == 'offset':
        field['placeholder_start'] += 1
    elif mutation == 'anchor':
        field['anchor_text'] += '가짜'
    elif mutation == 'key':
        field['placeholder_key'] = '가짜'
    elif mutation == 'id':
        field['id'] += '가짜'
    else:
        copied = deepcopy(field)
        copied['id'] += '_duplicate'
        profile['fields'].append(copied)
    values = {item['value_key']: '검증값' for item in profile['fields']}
    with pytest.raises(ValueError):
        fill_compatible_template(prepared, values, tmp_path / 'wrong.docx', profile=profile)


def test_ai_mapping_keeps_row_identity_and_rebinds_completion_groups(tmp_path):
    _, prepared, _, profile = prepared_form(tmp_path)
    originals = {field['id']: field for field in profile['fields']}
    class Client:
        calls = []
        def generate_json(self, prompt, payload):
            self.calls.append(deepcopy(payload))
            return {'fields': [{
                'id': item['id'], 'kind': item['kind'],
                'value_key': '새_' + originals[item['id']]['value_key'],
                'required': originals[item['id']]['required'],
                'input_required': originals[item['id']]['input_required'],
                'input_mode': 'user_provided' if originals[item['id']]['input_required'] else 'source_grounded',
                'max_chars': 1000, 'confidence': 0.95} for item in payload['fields']]}
    client = Client()
    learned = learn_template(prepared, client=client, force=True, base_profile=profile)
    assert learned['repeat_expansion'] == profile['repeat_expansion']
    assert all(key.startswith('새_') for group in learned['constraints']['groups'] for key in group['fields'])
    contexts = [item for call in client.calls for item in call['fields'] if item['kind'].endswith('repeat_placeholder')]
    assert len(contexts) == 12
    assert {item['location']['repeat_info']['row'] for item in contexts} == {1, 2, 3}
    assert {field['id'] for field in learned['fields']} == set(originals)


def test_repeat_ui_is_opt_in_and_updates_prepared_count(tmp_path):
    source, plan = form(tmp_path)
    script = "from pathlib import Path\nimport streamlit as st\nfrom app.repeat_ui import repeat_options\n"
    script += f"path, profile, expansion = repeat_options(Path({str(source)!r}), Path({str(tmp_path / 'data')!r}))\n"
    script += "st.json({'path': str(path), 'count': expansion['plan']['count'] if expansion else None})\n"
    ui = AppTest.from_string(script, default_timeout=20).run()
    assert not ui.exception
    assert json.loads(ui.json[0].value)['count'] is None
    digest = sha256(source.read_bytes()).hexdigest()
    ui.checkbox(key=f'repeat_enabled_{digest}').check().run()
    assert not ui.exception
    ui.number_input(key=f'repeat_count_{digest}').set_value(3).run()
    assert not ui.exception
    shown = json.loads(ui.json[0].value)
    assert shown['count'] == 3
    assert len(Document(shown['path']).tables[0].rows) == 5


def test_hwpx_blank_rows_fill_individually_and_keep_package_styles(tmp_path):
    from templates.repeat_fields import repeat_profile
    source = Path(__file__).resolve().parents[1] / 'samples/sample_company_form.hwpx'
    original = source.read_bytes()
    record = inspect_repeat_tables(source)[0]
    plan = {'table_id': record['id'], 'row': 2, 'count': 3}
    prepared = tmp_path / 'prepared.hwpx'
    expansion = prepare_repeat_template(source, prepared, plan)
    profile = repeat_profile(prepared, plan, expansion_binding(expansion))
    fields = [field for field in profile['fields'] if field.get('repeat_info')]
    assert len(fields) == 9
    values = {field['value_key']: f"시험 {field['repeat_info']['row']}-{field['repeat_info']['column']}"
              for field in fields}
    output = tmp_path / 'filled.hwpx'
    fill_compatible_template(prepared, values, output, profile=profile)
    assert verify_output(prepared, output, values, profile=profile)['status'] == 'passed'
    with ZipFile(source) as before, ZipFile(output) as after:
        for name in ('Contents/header.xml', 'version.xml', 'Contents/content.hpf'):
            assert before.read(name) == after.read(name)
    assert source.read_bytes() == original


def test_export_rechecks_repeat_proof_before_writing(tmp_path, monkeypatch):
    import agent.pipeline as pipeline
    from evals.run import MockEvaluationClient
    _, prepared, expansion, profile = prepared_form(tmp_path)
    case = {'mock_brief': {'목적': '매출', '보고 대상': '팀장', '보고서 유형': '결과보고서',
                          '마감': '', '분량': '1쪽', '부족한 정보': [], '질문': []}}
    text = '매출 120만원을 달성함'
    documents = [{'파일명': '자료.txt', '본문': text, '표 목록': [],
                  '페이지/시트 정보': [{'본문': text, '표 목록': [], '페이지': 1}]}]
    result = pipeline.run_pipeline('매출 결과보고서', documents=documents, client=MockEvaluationClient(case))
    draft = {field['value_key']: '시험값' for field in profile['fields']}
    draft.update(result['draft'])
    result.update(draft=draft, template_profile=profile, template_expansion=expansion)
    monkeypatch.setattr(pipeline, 'review_result', lambda value: {'draft': draft, 'blocking': False})
    output = pipeline.build_downloads(result, confirmed=True, template_paths={'docx': prepared})
    assert output['docx']
    assert result['output_verification']['docx']['repeat_expansion']['status'] == 'passed'
    result.pop('template_expansion')
    with pytest.raises(ValueError, match='준비 기록'):
        pipeline.build_downloads(result, confirmed=True, template_paths={'docx': prepared})
