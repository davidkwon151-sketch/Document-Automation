"""Offline rendered-PDF evidence, deliberately independent of Office/API access."""
from copy import deepcopy
from hashlib import sha256
from pathlib import Path

from docx import Document
from pypdf import PdfReader
from reportlab.pdfgen import canvas
import pytest

from agent.native_review import review_native_output
from templates.compatibility import analyze_template


def files(tmp_path, fields=('Title',)):
    source, output = tmp_path / 'original.docx', tmp_path / 'filled.docx'
    document = Document()
    for field in fields:
        document.add_paragraph('{{' + field + '}}')
    document.save(source)
    document.paragraphs[0].text = 'WRITTEN EXAMPLE'
    document.save(output)
    profile = analyze_template(source)
    for field in profile['fields']:
        field['value_key'] = field['label']
    return source, output, profile


class Renderer:
    def __init__(self, source_lines, output_lines, *, cell_evidence=None):
        self.source_lines, self.output_lines = source_lines, output_lines
        self.cell_evidence = cell_evidence or {}
        self.calls = []

    def __call__(self, path, target, **kwargs):
        self.calls.append((path, deepcopy(kwargs)))
        assert kwargs['timeout'] == 60
        lines = self.source_lines if path.name.startswith('original') else self.output_lines
        pdf = canvas.Canvas(str(target), pagesize=(500, 700), invariant=1)
        for number, line in enumerate(lines):
            if isinstance(line, tuple):
                x, y, text = line
            else:
                x, y, text = 35, 650 - number * 20, line
            pdf.drawString(x, y, text)
        if not lines:
            pdf.showPage()
        pdf.save()
        return {'status': 'rendered', 'engine': 'fake_renderer', 'source_sha256': sha256(path.read_bytes()).hexdigest(),
                'source_unchanged': True, 'output_path': str(target), 'output_sha256': sha256(target.read_bytes()).hexdigest(),
                'page_count': len(PdfReader(target).pages), 'formula_recalculated': False, 'cell_evidence': deepcopy(self.cell_evidence)}


def test_native_private_preview_value_increment_and_inputs_immutable(tmp_path):
    source, output, profile = files(tmp_path)
    original, filled = source.read_bytes(), output.read_bytes()
    values = {'Title': 'Distinct report value'}
    renderer = Renderer(['Source heading'], ['Source heading', 'Distinct report value'])
    report = review_native_output(source, output, values, profile=profile, renderer=renderer)
    assert report['status'] == 'passed' and not report['blocking']
    assert report['checked_field_count'] == 1
    assert report['page_count'] == 1 and report['pdf_bytes'].startswith(b'%PDF-')
    assert sha256(report['pdf_bytes']).hexdigest() == report['pdf_sha256']
    assert report['actualoutput_sha256'] == sha256(filled).hexdigest()
    assert source.read_bytes() == original and output.read_bytes() == filled
    metadata = {key: value for key, value in report.items() if key != 'pdf_bytes'}
    assert str(tmp_path) not in str(metadata) and 'Distinct report value' not in str(metadata)
    assert not report['full_layout_certified'] and not report['legal_compliance_certified']


@pytest.mark.parametrize('before,after', [
    (['Static heading'], ['Static heading']),
    (['Distinct report value'], ['Distinct report value']),
    (['Distinct report value expanded'], ['Distinct report value expanded']),
])
def test_missing_value_cannot_hide_in_original_text_or_heading_prefix(tmp_path, before, after):
    source, output, profile = files(tmp_path)
    report = review_native_output(source, output, {'Title': 'Distinct report value'}, profile=profile, renderer=Renderer(before, after))
    assert report['status'] == 'failed' and report['blocking']
    assert any(issue['code'] == 'native_value_missing' for issue in report['issues'])
    assert 'pdf_bytes' not in report


def test_existing_same_value_requires_actual_count_increment(tmp_path):
    source, output, profile = files(tmp_path)
    renderer = Renderer(['Distinct report value'], ['Distinct report value', 'Distinct report value'])
    report = review_native_output(source, output, {'Title': 'Distinct report value'}, profile=profile, renderer=renderer)
    assert report['status'] == 'passed' and report['checked_field_count'] == 1


def test_duplicate_selected_fields_require_every_occurrence(tmp_path):
    source, output, profile = files(tmp_path, fields=('First', 'Second'))
    values = {'First': 'Distinct repeated value', 'Second': 'Distinct repeated value'}
    report = review_native_output(source, output, values, profile=profile, renderer=Renderer([], ['Distinct repeated value']))
    assert report['blocking']
    good = review_native_output(source, output, values, profile=profile, renderer=Renderer([], ['Distinct repeated value', 'Distinct repeated value']))
    assert good['status'] == 'passed' and good['checked_field_count'] == 2


@pytest.mark.parametrize('value', ['123', '1,234.00', 'Yes', 'OK'])
def test_short_numeric_or_common_values_are_unevaluated_warning(tmp_path, value):
    source, output, profile = files(tmp_path)
    report = review_native_output(source, output, {'Title': value}, profile=profile, renderer=Renderer([], [value]))
    assert report['status'] == 'warning' and not report['blocking']
    assert report['checked_field_count'] == 0 and report['coverage']['ambiguous_fields'] == 1
    assert 'pdf_bytes' in report


def test_closed_choice_display_label_not_export_value_is_not_false_missing(tmp_path):
    source, output, profile = files(tmp_path)
    profile['fields'][0].update(control_type='choice', options=['A'], choice_items=[{'value': 'A', 'label': 'Choice display'}])
    report = review_native_output(source, output, {'Title': 'A'}, profile=profile, renderer=Renderer([], ['Choice display']))
    assert report['status'] == 'warning' and not report['blocking']
    assert report['coverage']['ambiguous_fields'] == 1


def test_new_page_clipped_text_blocks_but_preexisting_clipping_only_warns(tmp_path):
    source, output, profile = files(tmp_path)
    value = 'Distinct report value'
    report = review_native_output(source, output, {'Title': value}, profile=profile, renderer=Renderer([], [(480, 650, value)]))
    assert report['blocking'] and any(issue['code'] == 'native_new_page_clipping' for issue in report['issues'])
    original_clip = (480, 650, 'Original boundary text')
    report = review_native_output(source, output, {'Title': value}, profile=profile, renderer=Renderer([original_clip], [original_clip, value]))
    assert report['status'] == 'warning' and not report['blocking']
    assert any(issue['code'] == 'native_source_clipping' for issue in report['issues'])


@pytest.mark.parametrize('position', [(480, 600), (470, 650)])
def test_existing_clipped_text_reflow_warns_without_falsely_blocking_new_safe_value(tmp_path, position):
    source, output, profile = files(tmp_path)
    value = 'Distinct report value'
    original = (480, 650, 'STATIC CLIPPING')
    moved = (*position, 'STATIC CLIPPING')
    report = review_native_output(source, output, {'Title': value}, profile=profile,
                                  renderer=Renderer([original], [moved, value]))
    assert report['status'] == 'warning' and not report['blocking']
    assert report['checked_field_count'] == 1 and 'pdf_bytes' in report
    assert any(issue['code'] == 'native_clipping_reflow_unknown' for issue in report['issues'])
    assert not any(issue['code'] == 'native_new_page_clipping' for issue in report['issues'])
    assert next(check for check in report['checks'] if check['name'] == 'all_pdf_pages_character_bounds')['status'] == 'pending'


def test_old_clipped_glyphs_cannot_mask_new_selected_clipping(tmp_path):
    from reportlab.pdfbase.pdfmetrics import stringWidth
    source, output, profile = files(tmp_path)
    value = 'CLIPPING'
    # Both PDFs have exactly the same clipped glyphs, but the new value is clipped.
    aligned = 480 - stringWidth('STATIC ', 'Helvetica', 12) + stringWidth('C', 'Helvetica', 12) - stringWidth('F', 'Helvetica', 12)
    original = (aligned, 650, 'STATIC FLIPPING')
    report = review_native_output(source, output, {'Title': value}, profile=profile,
                                  renderer=Renderer([original], ['STATIC FLIPPING', (480, 600, value)]))
    assert report['status'] == 'failed' and report['blocking']
    assert 'pdf_bytes' not in report
    assert any(issue['code'] == 'native_new_value_clipping' for issue in report['issues'])
    assert not any(issue['code'] == 'native_new_page_clipping' for issue in report['issues'])


def test_moved_original_clipping_still_cannot_mask_missing_input(tmp_path):
    source, output, profile = files(tmp_path)
    report = review_native_output(source, output, {'Title': 'Distinct report value'}, profile=profile,
                                  renderer=Renderer([(480, 650, 'STATIC CLIPPING')], [(480, 600, 'STATIC CLIPPING')]))
    assert report['status'] == 'failed' and 'pdf_bytes' not in report
    assert any(issue['code'] == 'native_value_missing' for issue in report['issues'])


def test_profile_sha_stale_blocks_before_renderer(tmp_path):
    source, output, profile = files(tmp_path)
    profile['source_sha256'] = '0' * 64
    renderer = Renderer([], ['Distinct report value'])
    report = review_native_output(source, output, {'Title': 'Distinct report value'}, profile=profile, renderer=renderer)
    assert report['blocking'] and not renderer.calls


@pytest.mark.parametrize('stage', ['source_sha', 'pdf_sha', 'page_count'])
def test_renderer_stale_proof_is_not_success(tmp_path, stage):
    source, output, profile = files(tmp_path)
    base = Renderer([], ['Distinct report value'])
    def renderer(path, target, **kwargs):
        result = base(path, target, **kwargs)
        result[{'source_sha': 'source_sha256', 'pdf_sha': 'output_sha256', 'page_count': 'page_count'}[stage]] = 999 if stage == 'page_count' else '0' * 64
        return result
    report = review_native_output(source, output, {'Title': 'Distinct report value'}, profile=profile, renderer=renderer)
    assert report['status'] == 'failed' and 'pdf_bytes' not in report


@pytest.mark.parametrize('which', ['source', 'output'])
def test_input_mutation_is_blocking_even_if_renderer_reports_unavailable(tmp_path, which):
    source, output, profile = files(tmp_path)
    def renderer(path, target, **kwargs):
        selected = source if which == 'source' else output
        selected.write_bytes(selected.read_bytes() + b'changed')
        return {'status': 'unavailable', 'engine': None}
    report = review_native_output(source, output, {'Title': 'Distinct report value'}, profile=profile, renderer=renderer)
    assert report['status'] == 'failed' and report['blocking']
    assert any(issue['code'] == 'native_inputs_changed' for issue in report['issues'])


@pytest.mark.parametrize('unavailable', ['missing', 'timeout', 'second'])
def test_unavailable_native_is_never_passed(tmp_path, unavailable):
    source, output, profile = files(tmp_path)
    base = Renderer([], ['Distinct report value'])
    def renderer(path, target, **kwargs):
        if unavailable == 'timeout':
            raise TimeoutError('private-name')
        if unavailable == 'second' and path == source:
            return base(path, target, **kwargs)
        return {'status': 'unavailable', 'engine': None}
    report = review_native_output(source, output, {'Title': 'Distinct report value'}, profile=profile, renderer=renderer)
    assert report['status'] == 'unavailable' and not report['blocking'] and 'pdf_bytes' not in report
    assert report['source_sha256'] and report['output_sha256']
    assert 'private-name' not in str(report)


def xlsx_files(tmp_path, *, numeric=False):
    from openpyxl import Workbook
    source, output = tmp_path / 'original.xlsx', tmp_path / 'filled.xlsx'
    book = Workbook(); sheet = book.active; sheet['A1'] = 'Quantity'; sheet['B1'].number_format = '#,##0.00'
    sheet.print_area = 'A1:B1'; book.save(source); sheet['B1'] = 1234 if numeric else 'Distinct cell value'; book.save(output); book.close()
    profile = analyze_template(source)
    profile['fields'] = [{'id': 'xlsx:xl/worksheets/sheet1.xml:B1', 'label': 'Quantity', 'value_key': 'Quantity', 'kind': 'xlsx_cell', 'required': True, 'input_required': True}]
    if numeric:
        profile['fields'][0]['validation'] = {'type': 'number'}
    evidence = {'xlsx:xl/worksheets/sheet1.xml:B1': {'text': '1,234.00' if numeric else 'Distinct cell value', 'row_hidden': False, 'column_hidden': False, 'sheet_hidden': False, 'within_print_area': True, 'has_formula': False, 'print_area': '$A$1:$B$1', 'number_format': '#,##0.00', 'value2': 1234 if numeric else 'Distinct cell value'}}
    return source, output, profile, evidence


def test_xlsx_requests_are_physical_and_numeric_display_remains_ambiguous(tmp_path):
    source, output, profile, evidence = xlsx_files(tmp_path, numeric=True)
    renderer = Renderer(['Quantity'], ['Quantity', '1,234.00'], cell_evidence=evidence)
    report = review_native_output(source, output, {'Quantity': '1234'}, profile=profile, renderer=renderer)
    assert report['status'] == 'warning' and not report['blocking']
    assert report['checked_field_count'] == 0
    assert renderer.calls[0][1]['cell_requests'] == [{'id': 'xlsx:xl/worksheets/sheet1.xml:B1', 'part': 'xl/worksheets/sheet1.xml', 'coordinate': 'B1'}]
    assert '1,234.00' not in str({key: value for key, value in report.items() if key != 'pdf_bytes'})


@pytest.mark.parametrize('change,value', [('row_hidden', True), ('column_hidden', True), ('sheet_hidden', True), ('within_print_area', False), ('text', '###'), ('text', '')])
def test_native_xlsx_hidden_outside_print_or_display_overflow_blocks(tmp_path, change, value):
    source, output, profile, evidence = xlsx_files(tmp_path)
    evidence[next(iter(evidence))][change] = value
    report = review_native_output(source, output, {'Quantity': 'Distinct cell value'}, profile=profile,
                                  renderer=Renderer(['Quantity'], ['Quantity', 'Distinct cell value'], cell_evidence=evidence))
    assert report['status'] == 'failed' and report['blocking']
    assert report['checked_field_count'] == 0 and report['coverage']['failed_fields'] == 1
    assert next(check for check in report['checks'] if check['name'] == 'selected_printed_text_increment')['status'] == 'failed'


def test_unknown_xlsx_print_area_warns_and_never_claims_exact_print_scope(tmp_path):
    source, output, profile, evidence = xlsx_files(tmp_path)
    evidence[next(iter(evidence))]['within_print_area'] = None
    report = review_native_output(source, output, {'Quantity': 'Distinct cell value'}, profile=profile,
                                  renderer=Renderer(['Quantity'], ['Quantity', 'Distinct cell value'], cell_evidence=evidence))
    assert report['status'] == 'warning'
    assert any(issue['code'] == 'native_cell_visibility_unknown' for issue in report['issues'])
    assert report['checked_field_count'] == 0 and report['coverage']['ambiguous_fields'] == 1


def test_repeated_cell_phrase_cannot_hide_one_unknown_print_scope(tmp_path):
    source, output, profile, evidence = xlsx_files(tmp_path)
    second = deepcopy(profile['fields'][0])
    second.update(id='xlsx:xl/worksheets/sheet1.xml:C1', value_key='Repeated')
    profile['fields'].append(second)
    evidence[second['id']] = deepcopy(evidence[profile['fields'][0]['id']])
    evidence[profile['fields'][0]['id']]['within_print_area'] = None
    report = review_native_output(source, output, {'Quantity': 'Distinct cell value', 'Repeated': 'Distinct cell value'},
                                  profile=profile, renderer=Renderer(['Quantity'], ['Quantity', 'Distinct cell value', 'Distinct cell value'], cell_evidence=evidence))
    assert report['status'] == 'warning' and report['checked_field_count'] == 0
    assert report['coverage']['ambiguous_fields'] == 2


def test_no_selected_fields_does_not_claim_value_inspection_passed(tmp_path):
    source, output, profile = files(tmp_path)
    profile['fields'] = []
    report = review_native_output(source, output, {}, profile=profile, renderer=Renderer(['Original heading'], ['Original heading']))
    assert report['status'] == 'failed' and report['blocking']
    assert report['checked_field_count'] == 0
    assert report['pdf_sha256'] is None


def test_mapping_profile_and_values_are_not_mutated(tmp_path):
    source, output, profile = files(tmp_path)
    mapping = {profile['fields'][0]['id']: 'Other'}
    values = {'Other': 'Distinct report value'}
    before = deepcopy((profile, mapping, values))
    report = review_native_output(source, output, values, profile=profile, mapping=mapping,
                                  renderer=Renderer([], ['Distinct report value']))
    assert report['status'] == 'passed'
    assert (profile, mapping, values) == before


def test_invalid_values_cannot_be_coerced_and_same_input_rejected(tmp_path):
    source, output, profile = files(tmp_path)
    for result, values in [(output, {'Title': True}), (source, {'Title': 'Distinct value'})]:
        report = review_native_output(source, result, values, profile=profile, renderer=Renderer([], []))
        assert report['status'] == 'failed'


def test_native_exception_cannot_leak_private_paths_or_document_text(tmp_path):
    source, output, profile = files(tmp_path)
    def renderer(*args, **kwargs):
        raise ValueError(str(tmp_path) + ' PRIVATE CONTENT')
    report = review_native_output(source, output, {'Title': 'Distinct report value'}, profile=profile, renderer=renderer)
    assert report['blocking'] and 'PRIVATE CONTENT' not in str(report) and str(tmp_path) not in str(report)


def test_guarded_native_unavailable_subclass_keeps_existing_core_compatibility(tmp_path, monkeypatch):
    import parsers.native as native
    class Guarded(ValueError):
        pass
    monkeypatch.setattr(native, 'NativeRenderingUnavailable', Guarded, raising=False)
    source, output, profile = files(tmp_path)
    def renderer(*args, **kwargs):
        raise Guarded('PRIVATE PROTECTED FIELD')
    report = review_native_output(source, output, {'Title': 'Distinct report value'}, profile=profile, renderer=renderer)
    assert report['status'] == 'unavailable' and not report['blocking']
    assert report['source_sha256'] and report['actualoutput_sha256']
    assert any(issue['code'] == 'native_rendering_guarded' for issue in report['issues'])
    assert 'PRIVATE PROTECTED FIELD' not in str(report)


def test_ordinary_value_error_is_failed_despite_native_guard_subclass(tmp_path):
    source, output, profile = files(tmp_path)
    def renderer(*args, **kwargs):
        raise ValueError('malformed ZIP or stale SHA')
    report = review_native_output(source, output, {'Title': 'Distinct report value'}, profile=profile, renderer=renderer)
    assert report['status'] == 'failed' and report['blocking']


def test_whitespace_wrap_is_normalized_without_header_prefix_hits(tmp_path):
    source, output, profile = files(tmp_path)
    report = review_native_output(source, output, {'Title': 'Distinct report value'}, profile=profile,
                                  renderer=Renderer(['Distinct report values'], ['Distinct report values', 'Distinct report', 'value']))
    assert report['status'] == 'passed'


def test_pdf_extraction_permission_is_not_bypassed_by_pdfium(tmp_path):
    source, output, profile = files(tmp_path)
    base = Renderer([], ['Distinct report value'])
    def renderer(path, target, **kwargs):
        from pypdf import PdfWriter
        from pypdf.constants import UserAccessPermissions
        metadata = base(path, target, **kwargs)
        writer = PdfWriter(clone_from=target)
        writer.encrypt('', owner_password='test-owner', permissions_flag=UserAccessPermissions.PRINT)
        buffer = target.with_suffix('.encrypted.pdf')
        writer.write(buffer); buffer.replace(target)
        metadata['output_sha256'] = sha256(target.read_bytes()).hexdigest()
        return metadata
    report = review_native_output(source, output, {'Title': 'Distinct report value'}, profile=profile, renderer=renderer)
    assert report['status'] == 'failed' and report['blocking'] and 'pdf_bytes' not in report


def test_multiline_long_value_scope_is_warning_instead_of_unproven_success(tmp_path):
    source, output, profile = files(tmp_path)
    value = 'Complete long paragraph ' * 40
    report = review_native_output(source, output, {'Title': value}, profile=profile, renderer=Renderer([], ['Partial paragraph only']))
    assert report['status'] == 'warning' and report['checked_field_count'] == 0
    assert report['coverage']['ambiguous_fields'] == 1


def test_empty_text_pdf_is_unavailable_for_text_checks_not_a_success(tmp_path):
    source, output, profile = files(tmp_path)
    report = review_native_output(source, output, {'Title': 'Distinct report value'}, profile=profile, renderer=Renderer([], []))
    assert report['status'] == 'warning' and report['coverage']['checked_fields'] == 0


def test_unknown_engine_string_does_not_leak_private_metadata(tmp_path):
    source, output, profile = files(tmp_path)
    base = Renderer([], ['Distinct report value'])
    def renderer(path, target, **kwargs):
        metadata = base(path, target, **kwargs); metadata['engine'] = 'PRIVATE ENGINE PATH'
        return metadata
    report = review_native_output(source, output, {'Title': 'Distinct report value'}, profile=profile, renderer=renderer)
    assert report['blocking'] and 'PRIVATE ENGINE PATH' not in str(report)


def test_numeric_native_format_rounding_or_currency_has_explicit_warning(tmp_path):
    source, output, profile, evidence = xlsx_files(tmp_path, numeric=True)
    evidence[next(iter(evidence))]['text'] = '$1,234.00'
    report = review_native_output(source, output, {'Quantity': '1234'}, profile=profile,
                                  renderer=Renderer(['Quantity'], ['Quantity', '$1,234.00'], cell_evidence=evidence))
    assert report['status'] == 'warning' and report['checked_field_count'] == 0
    assert any(issue['code'] == 'native_number_format_ambiguous' for issue in report['issues'])


def test_global_placeholder_requires_all_physical_printed_occurrences(tmp_path):
    source, output, profile = files(tmp_path, fields=('Title', 'Title'))
    assert len(profile['fields']) == 1
    bad = review_native_output(source, output, {'Title': 'Distinct report value'}, profile=profile,
                               renderer=Renderer([], ['Distinct report value']))
    assert bad['status'] == 'failed' and bad['coverage']['selected_fields'] == 2
    good = review_native_output(source, output, {'Title': 'Distinct report value'}, profile=profile,
                                renderer=Renderer([], ['Distinct report value', 'Distinct report value']))
    assert good['status'] == 'passed' and good['checked_field_count'] == 2


def test_split_run_global_tokens_still_have_exact_physical_denominator(tmp_path):
    source, output, _ = files(tmp_path)
    document = Document(source); paragraph = document.add_paragraph(); paragraph.add_run('{{Ti'); paragraph.add_run('tle}}'); document.save(source)
    profile = analyze_template(source); profile['fields'][0]['value_key'] = 'Title'
    report = review_native_output(source, output, {'Title': 'Distinct report value'}, profile=profile,
                                  renderer=Renderer([], ['Distinct report value']))
    assert report['blocking'] and report['coverage']['selected_fields'] == 2


def test_unknown_global_profile_field_is_failed_before_renderer(tmp_path):
    source, output, profile = files(tmp_path)
    profile['fields'][0].update(id='placeholder:Unknown', value_key='Unknown')
    renderer = Renderer([], ['Distinct report value'])
    report = review_native_output(source, output, {'Unknown': 'Distinct report value'}, profile=profile, renderer=renderer)
    assert report['blocking'] and not renderer.calls


@pytest.mark.parametrize('state,expected', [(0, 'passed'), (2, 'warning')])
def test_excel_calculation_pending_is_explicit_without_discarding_visible_value(tmp_path, state, expected):
    source, output, profile, evidence = xlsx_files(tmp_path)
    rendered = Renderer(['Quantity'], ['Quantity', 'Distinct cell value'], cell_evidence=evidence)
    def renderer(*args, **kwargs):
        return {**rendered(*args, **kwargs), 'calculation_state': state,
                'formula_recalculated': state == 0}
    report = review_native_output(source, output, {'Quantity': 'Distinct cell value'},
                                  profile=profile, renderer=renderer)
    assert report['status'] == expected and not report['blocking']
    assert report['formula_recalculated'] is (state == 0)
    assert report['calculation_state'] == state and 'pdf_bytes' in report
    assert ('native_calculation_pending' in [item['code'] for item in report['issues']]) is (state == 2)
