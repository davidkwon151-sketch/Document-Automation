"""Rebuild the 15 synthetic, non-personal evaluation attachments."""

import json
from pathlib import Path

from docx import Document
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parent


def build_cases():
    data = ROOT / "data"
    data.mkdir(exist_ok=True)
    subjects = [
        ("주간업무보고", "영업 매출", "매출 120만원을 달성함"),
        ("결과보고서", "교육 참석자", "참석자 32명이 교육을 완료함"),
        ("품의서", "장비 예산", "장비 예산 450만원임"),
        ("주간업무보고", "고객 문의", "고객 문의 18건을 처리함"),
        ("결과보고서", "행사 비용", "행사 비용 210만원임"),
        ("품의서", "출장 예산", "출장 예산 75만원임"),
        ("주간업무보고", "개발 작업", "개발 작업 7건을 완료함"),
        ("결과보고서", "설문 응답자", "설문 응답자 84명임"),
        ("품의서", "홍보 예산", "홍보 예산 330만원임"),
        ("주간업무보고", "운영 비용", "운영 비용 95만원임"),
        ("결과보고서", "영업 매출", "매출 1250만원을 달성함"),
        ("품의서", "교육 예산", "교육 예산 160만원임"),
        ("결과보고서", "고객 상담", "고객 상담 24건을 완료함"),
        ("결과보고서", "교육 참석자", "참석자 12명이 교육을 완료함"),
        ("주간업무보고", "영업 매출", "매출 140만원을 달성함"),
    ]
    cases = []
    for index, (report_type, purpose, fact) in enumerate(subjects, 1):
        suffix = "xlsx" if index % 2 else "docx"
        relative = f"data/case_{index:02}.{suffix}"
        path = ROOT / relative
        if suffix == "xlsx":
            workbook = Workbook()
            workbook.active.title = "근거자료"
            workbook.active.append([fact])
            workbook.save(path)
        else:
            document = Document()
            document.add_paragraph(fact)
            document.save(path)
        audience = "임원" if report_type == "품의서" else "팀장"
        instruction = f"{audience}에게 {purpose} {report_type}를 2026-10-06까지 1쪽으로 작성해줘"
        brief = {"목적": purpose, "보고 대상": audience, "보고서 유형": report_type, "마감": "2026-10-06", "분량": "1쪽", "부족한 정보": [], "질문": []}
        expected = {"status": "ready", "report_type": report_type, "source_text": fact, "fields": ["제목", "요약", "본문"], "max_summary_lines": 3}
        if index == 14:
            instruction = "첨부 자료로 보고서를 만들어줘"
            brief.update({"목적": "", "보고 대상": "", "보고서 유형": "", "마감": "", "분량": "", "부족한 정보": ["목적", "보고 대상", "보고서 유형"], "질문": ["어떤 목적으로 누구에게 보고하나요?", "주간업무보고, 결과보고서, 품의서 중 어떤 유형인가요?"]})
            expected["status"] = "needs_information"
        cases.append({"id": f"case_{index:02}", "instruction": instruction, "attachments": [relative], "expected": expected, "mock_brief": brief, "inject_wrong_number_in_mock": index == 15})
    (ROOT / "cases.json").write_text(json.dumps(cases, ensure_ascii=False, indent=2), encoding="utf-8")
    return cases


if __name__ == "__main__":
    build_cases()
    print("평가 케이스 15개와 DOCX/XLSX 첨부자료를 생성함")
