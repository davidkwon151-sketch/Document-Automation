from pathlib import Path
import subprocess

from docx import Document
from openpyxl import Workbook
import pytest

from parsers import legacy


def docx_engine(source, output):
    document = Document()
    document.add_paragraph("변환 결과: 매출 100만원임")
    document.save(output)


def rtf(tmp_path, name="원자료.rtf"):
    source = tmp_path / name
    source.write_bytes(b"{\\rtf1\\ansi original content}")
    return source


def test_injected_converter_uses_copy_and_preserves_original(tmp_path):
    source = rtf(tmp_path)
    original = source.read_bytes()
    called = []
    def engine(copied, output):
        called.append((copied, output))
        assert copied != source and copied.read_bytes() == original
        copied.write_text("converter changed its private copy", encoding="utf-8")
        docx_engine(copied, output)
    output = legacy.convert_legacy(source, tmp_path / "output", engine)
    assert output.name == "원자료.docx" and output.is_file()
    assert source.read_bytes() == original
    assert "100만원" in Document(output).paragraphs[0].text
    assert not called[0][0].parent.exists()


def test_xls_converts_to_valid_xlsx(tmp_path):
    source = tmp_path / "실적.xls"
    source.write_bytes(legacy.OLE_SIGNATURE + b"test fixture")
    def engine(copied, output):
        workbook = Workbook()
        workbook.active.append(["매출", 100])
        workbook.save(output)
    assert legacy.convert_legacy(source, engine=engine).suffix == ".xlsx"


def test_converter_failure_does_not_modify_original_or_save_partial_output(tmp_path):
    source = rtf(tmp_path)
    original = source.read_bytes()
    def engine(copied, output):
        output.write_bytes(b"partial")
        raise RuntimeError("test failure")
    with pytest.raises(RuntimeError):
        legacy.convert_legacy(source, engine=engine)
    assert source.read_bytes() == original
    assert not source.with_suffix(".docx").exists()


@pytest.mark.parametrize("data", [b"not a zip", b""])
def test_invalid_output_is_rejected(tmp_path, data):
    source = rtf(tmp_path)
    with pytest.raises(ValueError):
        legacy.convert_legacy(source, engine=lambda copied, output: output.write_bytes(data))
    assert not source.with_suffix(".docx").exists()


def test_existing_destination_is_not_overwritten(tmp_path):
    source = rtf(tmp_path)
    destination = source.with_suffix(".docx")
    destination.write_bytes(b"original destination")
    with pytest.raises(FileExistsError):
        legacy.convert_legacy(source, engine=docx_engine)
    assert destination.read_bytes() == b"original destination"


@pytest.mark.parametrize(("name", "data"), [("file.doc", b"not Office"), ("file.xls", b"not Office"), ("file.rtf", b"not RTF"), ("file.odt", b"not ODT"), ("file.bin", b"unsupported")])
def test_input_format_and_extension_are_validated(tmp_path, name, data):
    source = tmp_path / name
    source.write_bytes(data)
    with pytest.raises(ValueError):
        legacy.convert_legacy(source, engine=docx_engine)


def test_hwp_security_registration_is_not_bypassed(tmp_path):
    source = tmp_path / "보고서.hwp"
    source.write_bytes(legacy.OLE_SIGNATURE)
    called = []
    with pytest.raises(ValueError, match="HWPX"):
        legacy.convert_legacy(source, engine=lambda *args: called.append(args))
    assert not called


def test_file_size_cap_is_checked_before_conversion(tmp_path, monkeypatch):
    monkeypatch.setattr(legacy, "MAX_FILE_BYTES", 5)
    with pytest.raises(ValueError, match="50 MiB"):
        legacy.convert_legacy(rtf(tmp_path), engine=docx_engine)


def test_missing_converter_has_actionable_message(tmp_path, monkeypatch):
    monkeypatch.setattr(legacy, "available_converters", lambda: {name: {"available": False} for name in ("libreoffice", "word_com", "excel_com")})
    with pytest.raises(RuntimeError, match="DOCX/XLSX"):
        legacy.convert_legacy(rtf(tmp_path))


def test_subprocess_timeout_is_reported_and_owned_process_is_stopped(monkeypatch):
    calls = []
    class Process:
        pid = 123
        def communicate(self, timeout=None):
            if timeout is not None:
                raise subprocess.TimeoutExpired("converter", timeout)
            return ("", "")
        def kill(self):
            calls.append("killed")
    monkeypatch.setattr(legacy.subprocess, "Popen", lambda *args, **kwargs: Process())
    monkeypatch.setattr(legacy.subprocess, "run", lambda *args, **kwargs: calls.append(args[0]))
    monkeypatch.setattr(legacy.os, "killpg", lambda *args: calls.append(args), raising=False)
    with pytest.raises(TimeoutError, match="제한시간"):
        legacy._run_bounded(["converter"], 1)
    assert "killed" in calls


def test_office_worker_disables_macros_and_opens_read_only():
    assert "$app.AutomationSecurity = 3" in legacy.OFFICE_WORKER
    assert "$app.Visible = $false" in legacy.OFFICE_WORKER
    assert "$readOnly = $true" in legacy.OFFICE_WORKER
    assert ".Open([ref]$InputPath, [ref]$confirm, [ref]$readOnly, [ref]$recent)" in legacy.OFFICE_WORKER
    assert ".Open($InputPath, 0, $true)" in legacy.OFFICE_WORKER
