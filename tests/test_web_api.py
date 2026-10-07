"""Gateway isolation/security and actual synthetic file flow; no paid model calls."""
from base64 import b64encode
from hashlib import sha256
from io import BytesIO
import json
import time
import threading
from uuid import uuid4
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from fastapi.testclient import TestClient
import httpx
import pytest

from app import web_api
from llm.client import LLMError

SECRET = 'test-only-gateway-secret-' + 'a' * 32


@pytest.fixture
def api(tmp_path):
    return TestClient(web_api.create_app(root=tmp_path / 'private', secret=SECRET),
                      raise_server_exceptions=False)


def request(api, method, path, payload=None, *, user='user-A', nonce=None, stamp=None,
            signature=None, sign_body=None):
    body = b'' if payload is None else json.dumps(payload, ensure_ascii=False).encode()
    timestamp, nonce = str(stamp or int(time.time())), nonce or uuid4().hex
    headers = {'X-RA-User': user, 'X-RA-Time': timestamp, 'X-RA-Nonce': nonce,
               'X-RA-Signature': signature or web_api.signature(
                   SECRET, method, path, timestamp, nonce, user, body if sign_body is None else sign_body)}
    if payload is not None:
        headers['Content-Type'] = 'application/json'
    return api.request(method, path, headers=headers, content=body)


def upload_data():
    document = Document()
    document.add_paragraph('합성 RA 시험 양식 · 실제 기관 양식 아님')
    table = document.add_table(rows=0, cols=2)
    for label in ('제품명', '성상', '저장방법', '포장단위', '신청인'):
        cells = table.add_row().cells
        cells[0].text, cells[1].text = label, '{{' + label + '}}'
    stream = BytesIO()
    document.save(stream)
    evidence = ('제품명: SyntheticDrugA\n제형·함량: 정제 10 mg\n'
                '성상: 흰색 정제\n저장방법: 실온 보관\n포장단위: 30정').encode()
    return {'template': {'name': 'Synthetic_RA_Form.docx', 'base64': b64encode(stream.getvalue()).decode()},
            'sources': [{'name': 'Synthetic_RA_Evidence.txt', 'base64': b64encode(evidence).decode()}]}


def configured(api, *, field_values=None, user='user-A'):
    response = request(api, 'POST', '/api/jobs', upload_data(), user=user)
    assert response.status_code == 200, response.text
    view = response.json()
    rows = view['mapping_rows']
    for row, field in zip(rows, view['profile']['fields']):
        row['채울 값'] = field['value_key']
        if row['항목'] == '신청인':
            row['직접 입력'] = True
    payload = {'confirmed': True, 'rows': rows,
               'selected_keys': [row['채울 값'] for row in rows],
               'instruction': '합성 제품 신청서 항목을 작성해줘',
               'product_name': 'SyntheticDrugA', 'variant': '정제 10 mg',
               'field_values': {'신청인': '담당자 직접 시험 입력'} if field_values is None else field_values}
    identifier = view['job_id']
    response = request(api, 'POST', f'/api/jobs/{identifier}/configure', payload, user=user)
    assert response.status_code == 200, response.text
    return identifier, payload


def copied(api, identifier, *, user='user-A'):
    proposal = request(api, 'POST', f'/api/jobs/{identifier}/proposals', user=user).json()
    response = request(api, 'POST', f'/api/jobs/{identifier}/generate', {
        'mode': 'source_copy', 'confirmed_proposals': proposal['fingerprint'],
        'confirmed_keys': list(proposal['source_bindings'])}, user=user)
    assert response.status_code == 200, response.text
    return response.json(), proposal


def test_default_gemini_uses_visitor_key_only_in_memory(tmp_path, monkeypatch):
    api = TestClient(web_api.create_app(root=tmp_path / 'private', secret=SECRET),
                     raise_server_exceptions=False)
    identifier, _ = configured(api)
    path = f'/api/jobs/{identifier}/generate'
    assert request(api, 'POST', path, {'mode': 'live'}).status_code == 422
    seen = []
    class FakeGemini:
        provider = 'gemini'
        def __init__(self, *, provider, api_key):
            seen.append((provider, api_key))
    monkeypatch.setattr(web_api, 'LLMClient', FakeGemini)
    def fake_generate(*args, client, **kwargs):
        assert client.provider == 'gemini'
        return {'status': 'needs_revision', 'mode': 'live', 'draft': {'제목': '시험'},
                'review': {'issues': [], 'blocking': True}, 'ready_for_output_check': False}
    monkeypatch.setattr(web_api.writer, 'generate_ra_auto', fake_generate)
    key = 'test-visitor-gemini-key-0123456789'
    assert request(api, 'POST', path, {'mode': 'live', 'gemini_api_key': key}).status_code == 202
    for _ in range(100):
        result = request(api, 'GET', f'/api/jobs/{identifier}')
        if result.json()['status'] != 'processing':
            break
        time.sleep(.01)
    assert seen == [('gemini', key)]
    assert key not in result.text
    saved = (tmp_path / 'private' / sha256(b'user-A').hexdigest() / identifier / 'state.json').read_text(encoding='utf-8')
    assert key not in saved and 'gemini_api_key' not in saved
    assert request(api, 'GET', '/api/health').json()['default_ai'] == 'gemini_byok'


def test_entire_actual_file_path_upload_confirm_copy_verify_download(api):
    identifier, _ = configured(api)
    result, _ = copied(api, identifier)
    assert result['result']['target_coverage']['filled_count'] == 5
    assert result['result']['actual_model_requests'] == 0
    assert result['result']['metrics'] is None
    response = request(api, 'POST', f'/api/jobs/{identifier}/export', {'confirmed': True})
    assert response.status_code == 200, response.text
    assert response.headers['Cache-Control'] == 'no-store'
    with ZipFile(BytesIO(response.content)) as archive:
        document = Document(BytesIO(archive.read('RA_작성본.docx')))
        assert [row.cells[1].text for row in document.tables[0].rows] == [
            'SyntheticDrugA', '흰색 정제', '실온 보관', '30정', '담당자 직접 시험 입력']
        evidence = json.loads(archive.read('출처와_검수기록.json'))
        assert evidence['output_verification']['docx']['status'] == 'passed'
        assert evidence['submission_ready'] is False
        assert 'template_path' not in evidence['auto']
        assert all(s['document_sha256'] for s in evidence['auto']['source_records'])


@pytest.mark.parametrize('method,suffix,payload', [
    ('GET', '', None), ('POST', '/proposals', None),
    ('POST', '/generate', {'mode': 'source_copy'}),
    ('POST', '/export', {'confirmed': True}), ('DELETE', '', None)])
def test_cross_tenant_ids_do_not_read_write_or_download(api, method, suffix, payload):
    identifier, _ = configured(api)
    response = request(api, method, f'/api/jobs/{identifier}' + suffix, payload, user='user-B')
    assert response.status_code == 404
    assert request(api, 'GET', f'/api/jobs/{identifier}').status_code == 200


def test_no_browser_header_or_known_local_routes_authorize(api):
    assert api.get('/api/health', headers={'oai-authenticated-user-id': 'user-A'}).status_code == 401
    assert api.get('/health').status_code == 401
    assert api.get('/sources').status_code == 401
    assert request(api, 'GET', '/api/health').json()['status'] == 'ok'


def test_replay_stale_timestamp_body_and_identity_tampering(api):
    nonce = uuid4().hex
    assert request(api, 'GET', '/api/health', nonce=nonce).status_code == 200
    assert request(api, 'GET', '/api/health', nonce=nonce).status_code == 401
    assert request(api, 'GET', '/api/health', stamp=int(time.time()) - 300).status_code == 401
    assert request(api, 'POST', '/api/jobs', upload_data(), sign_body=b'{}').status_code == 401
    assert request(api, 'GET', '/api/health', signature='0' * 64).status_code == 401


@pytest.mark.parametrize('name', ['../x.docx', 'C:\\private.docx', 'x.docx:ads', 'NUL.docx', 'x.docx.', '/x.docx'])
def test_upload_path_and_windows_stream_injection_rejected(api, name):
    payload = upload_data()
    payload['template']['name'] = name
    assert request(api, 'POST', '/api/jobs', payload).status_code == 422


def test_archive_bomb_and_member_path_rejected_before_parsing(api):
    payload = upload_data()
    for name, data in [('word/document.xml', b'0' * 200_000), ('../outside', b'data')]:
        stream = BytesIO()
        with ZipFile(stream, 'w', ZIP_DEFLATED) as archive:
            archive.writestr(name, data)
        payload['template']['base64'] = b64encode(stream.getvalue()).decode()
        assert request(api, 'POST', '/api/jobs', payload).status_code == 422


def test_oversized_request_rejected_before_authentication(api):
    response = api.post('/api/jobs', content=b'0' * (web_api.MAX_BODY + 1))
    assert response.status_code == 413


def test_current_confirmation_and_configuration_invalidation(api):
    identifier, config = configured(api)
    view, proposal = copied(api, identifier)
    assert view['result']['ready_for_output_check']
    changed = {**config, 'product_name': 'OtherProduct'}
    assert request(api, 'POST', f'/api/jobs/{identifier}/configure', changed).status_code == 200
    assert request(api, 'POST', f'/api/jobs/{identifier}/export', {'confirmed': True}).status_code == 409
    assert request(api, 'POST', f'/api/jobs/{identifier}/generate', {
        'mode': 'source_copy', 'confirmed_proposals': proposal['fingerprint'],
        'confirmed_keys': list(proposal['source_bindings'])}).status_code == 422


def test_candidate_confirmation_cannot_omit_a_candidate(api):
    identifier, _ = configured(api)
    proposal = request(api, 'POST', f'/api/jobs/{identifier}/proposals').json()
    assert request(api, 'POST', f'/api/jobs/{identifier}/generate', {
        'mode': 'source_copy', 'confirmed_proposals': proposal['fingerprint'],
        'confirmed_keys': []}).status_code == 422


def test_missing_direct_field_blocks_export(api):
    identifier, _ = configured(api, field_values={})
    view, _ = copied(api, identifier)
    assert view['result']['target_coverage']['missing_keys'] == ['신청인']
    assert not view['result']['ready_for_output_check']
    assert request(api, 'POST', f'/api/jobs/{identifier}/export', {'confirmed': True}).status_code == 422


def test_quota_error_is_safe_and_invalidates_previous_download(tmp_path):
    def factory():
        raise LLMError('secret sk-test-do-not-expose C:\\private\\patient.txt', kind='quota')
    api = TestClient(web_api.create_app(root=tmp_path, secret=SECRET, client_factory=factory),
                     raise_server_exceptions=False)
    identifier, _ = configured(api)
    copied(api, identifier)
    response = request(api, 'POST', f'/api/jobs/{identifier}/generate', {'mode': 'live'})
    assert response.status_code == 202
    for _ in range(50):
        response = request(api, 'GET', f'/api/jobs/{identifier}')
        if response.json()['status'] != 'processing':
            break
        time.sleep(.01)
    assert response.json()['error']['code'] == 'llm_quota'
    assert 'sk-test' not in response.text and 'patient.txt' not in response.text
    assert request(api, 'POST', f'/api/jobs/{identifier}/export', {'confirmed': True}).status_code == 409


def test_stored_template_mutation_does_not_pass_final_verification(api, tmp_path):
    identifier, _ = configured(api)
    copied(api, identifier)
    owner = tmp_path / 'private' / sha256(b'user-A').hexdigest()
    (owner / identifier / 'template' / 'Synthetic_RA_Form.docx').write_bytes(b'changed')
    assert request(api, 'POST', f'/api/jobs/{identifier}/export', {'confirmed': True}).status_code == 422


def test_private_server_paths_are_not_in_api_views_or_download(api, tmp_path):
    identifier, _ = configured(api)
    response = request(api, 'GET', f'/api/jobs/{identifier}')
    assert str(tmp_path) not in response.text
    assert 'template_path' not in response.text


def test_hwp_uses_separate_authenticated_worker_with_checked_sha(tmp_path, monkeypatch):
    source = tmp_path / 'public.hwp'
    source.write_bytes(b'original-HWP')
    stream = BytesIO()
    with ZipFile(stream, 'w') as archive:
        archive.writestr('mimetype', 'application/hwp+zip')
        archive.writestr('Contents/section0.xml', '<section/>')
    data = stream.getvalue()
    class Client:
        def __init__(self, **kwargs):
            assert kwargs['trust_env'] is False and kwargs['follow_redirects'] is False
        def __enter__(self): return self
        def __exit__(self, *args): pass
        def post(self, url, headers, json):
            assert url == 'http://127.0.0.1:8601/convert'
            assert headers['Authorization'] == 'Bearer worker-secret'
            assert json['name'] == source.name
            return httpx.Response(200, json={'base64': b64encode(data).decode(),
                'source_sha256': sha256(source.read_bytes()).hexdigest(),
                'output_sha256': sha256(data).hexdigest()})
    monkeypatch.setattr(web_api.httpx, 'Client', Client)
    converted = web_api._convert_hwp(source, 'http://127.0.0.1:8601', 'worker-secret')
    assert converted.read_bytes() == data and source.read_bytes() == b'original-HWP'


def test_short_or_missing_secret_fails_closed(tmp_path):
    with pytest.raises(RuntimeError):
        web_api.create_app(root=tmp_path, secret='short')


@pytest.mark.parametrize('names', [('x.hwp', 'x.hwpx'), ('x.hwpx', 'x.hwp'), ('X.HWP', 'x.hwpx')])
def test_converted_source_name_collision_rejected_before_worker(api, names, monkeypatch):
    def never(*args): raise AssertionError('must reject before native request')
    monkeypatch.setattr(web_api, '_convert_hwp', never)
    payload = upload_data()
    payload['sources'] = [{'name': names[0], 'base64': b64encode(b'HWP').decode()},
                          {'name': names[1], 'base64': payload['template']['base64']}]
    assert request(api, 'POST', '/api/jobs', payload).status_code == 422


def test_signed_user_request_rate_is_bounded(api):
    for _ in range(60):
        assert request(api, 'GET', '/api/health').status_code == 200
    assert request(api, 'GET', '/api/health').status_code == 429
    assert request(api, 'GET', '/api/health', user='user-B').status_code == 200


def test_global_storage_cap_prevents_many_tenants(api, monkeypatch):
    monkeypatch.setattr(web_api, 'MAX_STORAGE_BYTES', 1)
    response = request(api, 'POST', '/api/jobs', upload_data())
    assert response.status_code == 429
    assert response.json()['detail']['error'] == 'server_storage_limit'


def test_async_owner_mutations_block_and_baseline_survives_failed_regeneration(tmp_path, monkeypatch):
    entered, release = threading.Event(), threading.Event()
    def factory():
        entered.set()
        release.wait(timeout=5)
        raise LLMError('private raw error', kind='quota')
    api = TestClient(web_api.create_app(root=tmp_path, secret=SECRET, client_factory=factory),
                     raise_server_exceptions=False)
    identifier, config = configured(api)
    copied(api, identifier)
    state_path = tmp_path / sha256(b'user-A').hexdigest() / identifier / 'state.json'
    state = json.loads(state_path.read_text(encoding='utf-8'))
    # Only the service owns this server-side record; the browser cannot supply a baseline.
    state['previous_metrics'] = {'first_draft_at': 'fixed-test-baseline', 'draft': {'제목': 'original'}}
    web_api._save(state_path, state)
    assert request(api, 'POST', f'/api/jobs/{identifier}/generate', {'mode': 'live'}).status_code == 202
    assert entered.wait(timeout=2)
    assert request(api, 'GET', f'/api/jobs/{identifier}').json()['status'] == 'processing'
    assert request(api, 'POST', f'/api/jobs/{identifier}/configure', config).status_code == 409
    assert request(api, 'DELETE', f'/api/jobs/{identifier}').status_code == 409
    assert request(api, 'POST', f'/api/jobs/{identifier}/export', {'confirmed': True}).status_code == 409
    release.set()
    for _ in range(50):
        if request(api, 'GET', f'/api/jobs/{identifier}').json()['status'] != 'processing': break
        time.sleep(.01)
    saved = json.loads(state_path.read_text(encoding='utf-8'))
    assert saved['previous_metrics']['first_draft_at'] == 'fixed-test-baseline'
    assert 'result' not in saved


def test_source_copy_review_does_not_construct_model(tmp_path):
    def never(): raise AssertionError('copy must not invoke model')
    api = TestClient(web_api.create_app(root=tmp_path, secret=SECRET, client_factory=never),
                     raise_server_exceptions=False)
    identifier, _ = configured(api)
    copied(api, identifier)
    assert request(api, 'POST', f'/api/jobs/{identifier}/review', {'draft': {}}).status_code == 422


def test_nonce_replay_is_rejected_after_server_instance_restart(tmp_path):
    nonce = uuid4().hex
    first = TestClient(web_api.create_app(root=tmp_path, secret=SECRET))
    assert request(first, 'GET', '/api/health', nonce=nonce).status_code == 200
    restarted = TestClient(web_api.create_app(root=tmp_path, secret=SECRET))
    assert request(restarted, 'GET', '/api/health', nonce=nonce).status_code == 401


def test_interrupted_background_state_is_recovered_without_old_download(tmp_path):
    first = TestClient(web_api.create_app(root=tmp_path, secret=SECRET))
    identifier, _ = configured(first)
    copied(first, identifier)
    state_path = tmp_path / sha256(b'user-A').hexdigest() / identifier / 'state.json'
    state = json.loads(state_path.read_text(encoding='utf-8'))
    state.update(processing=True, operation_owner='old-process', operation='generate')
    web_api._save(state_path, state)
    restarted = TestClient(web_api.create_app(root=tmp_path, secret=SECRET))
    view = request(restarted, 'GET', f'/api/jobs/{identifier}').json()
    assert view['status'] == 'failed' and view['error']['code'] == 'interrupted'
    assert 'result' not in view
    assert request(restarted, 'POST', f'/api/jobs/{identifier}/export', {'confirmed': True}).status_code == 409
    assert request(restarted, 'DELETE', f'/api/jobs/{identifier}').status_code == 200


def test_sha_registered_form_authority_is_used_before_mapping(tmp_path, monkeypatch):
    from templates.compatibility import analyze_template
    raw = upload_data()
    from base64 import b64decode
    path = tmp_path / 'registered.docx'
    path.write_bytes(b64decode(raw['template']['base64']))
    authority = analyze_template(path)
    authority.update(configured=True, domain='pharmaceutical_ra', ra_workflow='product_approval')
    for field in authority['fields']:
        field.update(max_chars=1000, input_mode='source_grounded')
    authority['fields'][0].update(evidence_role='product_name', required=True, max_chars=70)
    authority['fields'][-1].update(input_required=True, input_mode='user_provided')
    monkeypatch.setattr(web_api, 'load_form_profile', lambda p: authority)
    monkeypatch.setattr(web_api.writer, 'load_form_profile', lambda p, entry=None: authority)
    monkeypatch.setattr(__import__(__name__, fromlist=['upload_data']), 'upload_data', lambda: raw)
    api = TestClient(web_api.create_app(root=tmp_path / 'store', secret=SECRET), raise_server_exceptions=False)
    identifier, _ = configured(api)
    view = request(api, 'GET', f'/api/jobs/{identifier}').json()
    assert view['profile']['fields'][0]['evidence_role'] == 'product_name'
    assert view['profile']['fields'][0]['required'] is True
    assert view['profile']['fields'][0]['max_chars'] == 70
    assert view['profile']['fields'][-1]['input_required'] is True
    result, _ = copied(api, identifier)
    assert result['result']['ready_for_output_check']


def test_detected_applicant_requires_direct_ui_input_even_if_mapping_did_not_declare(api):
    response = request(api, 'POST', '/api/jobs', upload_data())
    view = response.json()
    rows = view['mapping_rows']
    for row, field in zip(rows, view['profile']['fields']):
        row['채울 값'] = field['value_key']
        row['직접 입력'] = False
    payload = {'confirmed': True, 'rows': rows, 'selected_keys': ['신청인'],
               'instruction': '직접 입력 칸 시험', 'product_name': 'SyntheticDrugA', 'variant': '',
               'field_values': {}}
    response = request(api, 'POST', f'/api/jobs/{view["job_id"]}/configure', payload)
    assert response.status_code == 200, response.text
    applicant = next(f for f in response.json()['profile']['fields'] if f['value_key'] == '신청인')
    assert applicant['input_required'] and applicant['input_mode'] == 'user_provided'


def test_private_metadata_cleaning_preserves_exact_facts_and_path_named_fields(tmp_path):
    literal = 'C:\\원자료\\실제 보관 위치.txt'
    source = {'text': literal, 'quote': literal, 'context_text': literal,
              'path': '원문 선언 값', 'arbitrary_provenance_key': {'label': literal}}
    value = {'template_path': str(tmp_path / 'private.docx'),
             'metadata': {'path': str(tmp_path / 'output.docx'), 'sha': {'output': 'a' * 64}},
             'draft': {'path': literal, 'template_path': literal},
             'field_values': {'path': literal}, 'sources': [source],
             'auto': {'source_records': [source], 'template_path': str(tmp_path / 'private.docx')},
             'profile': {'fields': [{'value_key': 'path', 'label': literal}]}}
    clean = web_api._clean(value, tmp_path)
    assert 'template_path' not in clean and 'path' not in clean['metadata']
    assert clean['metadata']['sha']['output'] == 'a' * 64
    assert clean['draft'] == value['draft'] and clean['field_values'] == value['field_values']
    assert clean['sources'] == [source] and clean['auto']['source_records'] == [source]
    assert clean['profile']['fields'][0] == {'value_key': 'path', 'label': literal}


def test_storage_quota_scan_tolerates_atomic_temporary_file_disappearance(tmp_path):
    class Gone:
        def is_file(self): return True
        def stat(self): raise FileNotFoundError('journal already removed')
    class Stable:
        def is_file(self): return True
        def stat(self):
            return type('Stat', (), {'st_size': 31})()
    class Directory:
        def rglob(self, pattern): return [Gone(), Stable()]
    assert web_api._storage_size(Directory()) == 31


def test_single_huge_plaintext_paragraph_rejected_before_chunk_creation(api, monkeypatch):
    from agent import multimodal_intake
    def never(*args, **kwargs): raise AssertionError('must bound context before chunking')
    monkeypatch.setattr(multimodal_intake, 'chunk_documents', never)
    payload = upload_data()
    payload['sources'][0]['base64'] = b64encode(b'A' * (web_api.MAX_BLOCK_CHARS + 1)).decode()
    response = request(api, 'POST', '/api/jobs', payload)
    assert response.status_code == 413
    assert response.json()['error'] == 'resource_limit'


def test_highly_expanded_source_package_rejected_before_chunk_creation(api, monkeypatch):
    from agent import multimodal_intake
    def never(*args, **kwargs): raise AssertionError('must bound parsed data before chunking')
    monkeypatch.setattr(multimodal_intake, 'chunk_documents', never)
    # Distinct short paragraphs stay below per-block bounds and pass ZIP ratio guard.
    document = Document()
    for i in range(4200):
        document.add_paragraph(f'문단 {i:05d}: ' + 'ABCD' * 128)
    stream = BytesIO(); document.save(stream)
    payload = upload_data()
    payload['sources'] = [{'name': 'Expanded_source.docx', 'base64': b64encode(stream.getvalue()).decode()}]
    response = request(api, 'POST', '/api/jobs', payload)
    assert response.status_code == 413, response.text


def test_oversized_state_write_keeps_previous_file_and_removes_temporary(tmp_path):
    path = tmp_path / 'state.json'
    web_api._save(path, {'original': 'safe'})
    original = path.read_bytes()
    with pytest.raises(ValueError):
        web_api._save(path, {'too_long': '가' * 1000}, max_bytes=100)
    assert path.read_bytes() == original
    assert not list(tmp_path.glob('*.tmp'))


def test_state_metadata_and_parse_data_count_toward_tenant_storage_cap(api, monkeypatch):
    # Inputs are ~40KiB; structured document+provenance makes state exceed this cap.
    raw = upload_data()
    incoming = len(__import__('base64').b64decode(raw['template']['base64']))
    incoming += len(__import__('base64').b64decode(raw['sources'][0]['base64']))
    monkeypatch.setattr(web_api, 'MAX_TENANT_BYTES', incoming + 10)
    response = request(api, 'POST', '/api/jobs', raw)
    assert response.status_code == 413, response.text


def test_intake_resource_bound_defaults_preserve_existing_source_records(tmp_path):
    from agent.multimodal_intake import collect_multimodal
    path = tmp_path / 'source.txt'; path.write_text('제품명: SyntheticDrugA\n성상: 흰색 정제', encoding='utf-8')
    original = collect_multimodal([path])
    guarded = collect_multimodal([path], max_source_chars=1000, max_block_chars=1000)
    assert original == guarded


def test_actual_answers_accumulate_by_exact_question_across_config_and_quota_failure(tmp_path, monkeypatch):
    captured = []
    def mock_writer(*args, **kwargs):
        captured.append(dict(kwargs['answers']))
        if len(captured) == 3:
            raise LLMError('private quota payload', kind='quota')
        return {'status': 'needs_revision', 'mode': 'live', 'metrics': None,
                'questions': ['보고 대상은 누구인가요?'] if len(captured) == 1 else ['마감은 언제인가요?'],
                'answers': kwargs['answers']}
    monkeypatch.setattr(web_api.writer, 'generate_ra_auto', mock_writer)
    api = TestClient(web_api.create_app(root=tmp_path, secret=SECRET, client_factory=lambda: object()))
    identifier, config = configured(api)
    def generate(answers):
        response = request(api, 'POST', f'/api/jobs/{identifier}/generate', {'mode': 'live', 'answers': answers})
        assert response.status_code == 202
        for _ in range(50):
            view = request(api, 'GET', f'/api/jobs/{identifier}').json()
            if view['status'] != 'processing': return view
            time.sleep(.01)
        pytest.fail('background mock did not finish')
    first = {'보고 목적은 무엇인가요?': '변경 자료 검토'}
    view = generate(first)
    assert view['user_answers'] == first
    # New question text is not an answer and is never automatically populated.
    assert '보고 대상은 누구인가요?' not in view['user_answers']
    changed = request(api, 'POST', f'/api/jobs/{identifier}/configure', config).json()
    assert changed['user_answers'] == first
    second = {'보고 대상은 누구인가요?': '부서장'}
    view = generate(second)
    assert captured[1] == {**first, **second}
    assert '마감은 언제인가요?' not in view['user_answers']
    view = generate({'마감은 언제인가요?': '담당자 지정 일정'})
    assert view['error']['code'] == 'llm_quota'
    assert view['user_answers'] == captured[2] == {**first, **second, '마감은 언제인가요?': '담당자 지정 일정'}
    # An omitted answer dictionary on retry also restores the cumulative record.
    response = request(api, 'POST', f'/api/jobs/{identifier}/generate', {'mode': 'live'})
    assert response.status_code == 202
    for _ in range(50):
        if request(api, 'GET', f'/api/jobs/{identifier}').json()['status'] != 'processing': break
        time.sleep(.01)
    assert captured[3] == captured[2]
    other = request(api, 'POST', '/api/jobs', upload_data()).json()
    assert other['user_answers'] == {} and other['job_id'] != identifier


@pytest.mark.parametrize('answers', [None, {'Q': 1}, {'': 'reply'}, {'Q': {'fake': 'fact'}},
                                    {'Q': 'a' * 10_001}, {'Q' * 2001: 'reply'}])
def test_answer_contract_rejects_nonflat_or_excessive_values(api, answers):
    identifier, _ = configured(api)
    response = request(api, 'POST', f'/api/jobs/{identifier}/generate', {'mode': 'live', 'answers': answers})
    assert response.status_code == 422
    assert request(api, 'GET', f'/api/jobs/{identifier}').json()['user_answers'] == {}


def test_cumulative_user_answers_cannot_become_ra_source_facts(api):
    identifier, _ = configured(api)
    proposal = request(api, 'POST', f'/api/jobs/{identifier}/proposals').json()
    response = request(api, 'POST', f'/api/jobs/{identifier}/generate', {
        'mode': 'source_copy', 'confirmed_proposals': proposal['fingerprint'],
        'confirmed_keys': list(proposal['source_bindings']), 'answers': {'성상': '가짜 사실'}})
    assert response.status_code == 422
    view = request(api, 'GET', f'/api/jobs/{identifier}').json()
    assert view['user_answers'] == {'성상': '가짜 사실'}
    assert 'result' not in view
    assert request(api, 'POST', f'/api/jobs/{identifier}/export', {'confirmed': True}).status_code == 409
