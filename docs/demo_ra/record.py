"""Capture the real RA CTD workspace with synthetic files and an isolated API."""

from __future__ import annotations

import asyncio
import json
import tempfile
import time
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from uuid import uuid4
from zipfile import ZipFile

from docx import Document
from fastapi.testclient import TestClient
from playwright.async_api import async_playwright

from app import web_api


ROOT = Path(__file__).resolve().parents[2]
DIST = ROOT / "share_site" / "dist"
SAMPLES = ROOT / "samples"
RAW = ROOT / ".runtime" / "demo_ra_capture"
SECRET = "ra-video-only-gateway-secret-" + "x" * 40
SOURCES = ("ctd_demo_batch_formula.txt", "ctd_demo_batch_analysis.txt",
           "ctd_demo_stability.txt")
SECTIONS = ("3.2.P.3.2", "3.2.P.5.4", "3.2.P.8.3")


def mime(path: Path) -> str:
    return {".html": "text/html; charset=utf-8", ".js": "application/javascript; charset=utf-8",
            ".css": "text/css; charset=utf-8", ".png": "image/png", ".svg": "image/svg+xml",
            ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"}.get(
                path.suffix, "application/octet-stream")


def checked_output(data: bytes) -> dict:
    with ZipFile(BytesIO(data)) as archive:
        report = json.loads(archive.read("CTD_출처와_누락.json"))
        doc = Document(BytesIO(archive.read("CTD_작업초안.docx")))
        text = "\n".join([p.text for p in doc.paragraphs] +
                         [cell.text for table in doc.tables for row in table.rows for cell in row.cells])
    assert report["output_check"]["status"] == "passed"
    assert report["submission_ready"] is False
    assert all(value in text for value in ("DEMO-B01", "98.7%", "98.2%"))
    assert all(f"{{{{{section}}}}}" not in text for section in SECTIONS)
    return {"output_sha256": sha256(data).hexdigest(), "output_check": "passed",
            "submission_ready": False, "contains": ["DEMO-B01", "98.7%", "98.2%"]}


async def record(raw: Path = RAW) -> None:
    raw.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="ra-video-api-", ignore_cleanup_errors=True) as private, \
            TestClient(web_api.create_app(root=private, secret=SECRET),
                       raise_server_exceptions=False) as api:
        async with async_playwright() as playwright:
            browser = await playwright.chromium.launch(
                executable_path=r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
                headless=True)
            page = await browser.new_page(viewport={"width": 1280, "height": 720},
                                          device_scale_factor=1, accept_downloads=True)

            async def route(request_route):
                request = request_route.request
                path = request.url.split("demo.invalid", 1)[-1].split("?", 1)[0]
                if path.startswith("/api/"):
                    body = request.post_data_buffer or b""
                    stamp, nonce, user = str(int(time.time())), uuid4().hex, "ra-demo-viewer"
                    headers = {"X-RA-User": user, "X-RA-Time": stamp, "X-RA-Nonce": nonce,
                               "X-RA-Signature": web_api.signature(
                                   SECRET, request.method, path, stamp, nonce, user, body)}
                    if body:
                        headers["Content-Type"] = "application/json"
                    reply = api.request(request.method, path, headers=headers, content=body)
                    await request_route.fulfill(status=reply.status_code, body=reply.content,
                                                content_type=reply.headers.get("content-type", "application/json"))
                    return
                target = DIST / ("workspace.html" if path in {"", "/"} else path.lstrip("/"))
                if not target.is_file() or not target.resolve().is_relative_to(DIST.resolve()):
                    await request_route.fulfill(status=404, body="Not found")
                    return
                await request_route.fulfill(body=target.read_bytes(), content_type=mime(target))

            await page.route("https://demo.invalid/**", route)
            await page.goto("https://demo.invalid/workspace.html", wait_until="networkidle")
            await page.wait_for_function("document.querySelector('#status').textContent.includes('서버 연결 완료')")
            captures = []

            async def snap(stage: str):
                filename = f"{len(captures):02d}.png"
                await page.screenshot(path=str(raw / filename), animations="disabled")
                captures.append({"stage": stage, "file": filename})

            await snap("workspace")
            await page.locator("#template").set_input_files(str(SAMPLES / "ctd_demo_template.docx"))
            await page.locator("#sources").set_input_files([str(SAMPLES / name) for name in SOURCES])
            await page.locator("#upload-section").scroll_into_view_if_needed()
            await snap("files_selected")
            await page.locator("#upload").click()
            await page.locator("#configuration:not([hidden])").wait_for()
            await page.locator("#configuration").scroll_into_view_if_needed()
            await snap("mapping")
            await page.locator("#ctd-workspace").evaluate("element => element.open = true")
            await page.locator("#ctd-product").fill("예시정")
            await page.locator("#ctd-variant").fill("정제 5 mg")
            await page.locator('[data-ctd-section="M1_ADMIN"]').uncheck()
            await page.locator("#ctd-product").scroll_into_view_if_needed()
            await snap("ctd_setup")
            await page.locator("#ctd-sections").scroll_into_view_if_needed()
            await snap("sections_selected")
            await page.locator("#ctd-prepare").click()
            await page.wait_for_function("document.querySelector('#ctd-result').textContent.includes('선택 3절')")
            assert "원문 후보 3절" in await page.locator("#ctd-result").inner_text()
            await page.locator("#ctd-result").scroll_into_view_if_needed()
            await snap("source_preview")
            await page.locator("#ctd-result article").last.scroll_into_view_if_needed()
            await snap("source_evidence")
            await page.locator("#ctd-confirm").check()
            await page.locator("#ctd-export").scroll_into_view_if_needed()
            await snap("export_ready")
            async with page.expect_download() as download_info:
                await page.locator("#ctd-export").click()
            download = await download_info.value
            await download.save_as(str(raw / "output.zip"))
            output = checked_output((raw / "output.zip").read_bytes())
            await page.locator("#status").scroll_into_view_if_needed()
            await snap("downloaded")
            provenance = {
                "capture": "real workspace UI and isolated FastAPI backend",
                "mode": "deterministic exact-source excerpt filling; no LLM call",
                "template": "synthetic project example, not an official submission form",
                "template_sha256": sha256((SAMPLES / "ctd_demo_template.docx").read_bytes()).hexdigest(),
                "source_sha256": {name: sha256((SAMPLES / name).read_bytes()).hexdigest() for name in SOURCES},
                "sections": list(SECTIONS), "coverage": "3/3 selected sections",
                **output,
            }
            (raw / "manifest.json").write_text(json.dumps(captures, ensure_ascii=False, indent=2), encoding="utf-8")
            (raw / "provenance.json").write_text(json.dumps(provenance, ensure_ascii=False, indent=2), encoding="utf-8")
            await browser.close()


if __name__ == "__main__":
    asyncio.run(record())
