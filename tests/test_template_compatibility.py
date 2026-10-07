from hashlib import sha256
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, Alignment
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DictionaryObject, NameObject, TextStringObject
import pytest
from reportlab.pdfgen import canvas

from parsers import parse_file
from samples.generate import VALUES
from templates import TemplateError, analyze_template, fill_compatible_template


ROOT = Path(__file__).resolve().parents[1]


def _parts(path):
    with ZipFile(path) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


@pytest.mark.parametrize("suffix", ["docx", "hwpx", "xlsx", "pdf"])
def test_company_form_without_placeholders_exposes_stable_mapping_profile(suffix):
    source = ROOT / f"samples/sample_company_form.{suffix}"
    profile = analyze_template(source)
    assert profile["supported"]
    assert profile["format"] == suffix
    assert {field["value_key"] for field in profile["fields"]} >= {"제목", "요약", "본문"}
    assert analyze_template(source) == profile
    assert profile["warnings"]
    for field in profile["fields"]:
        assert {"id", "label", "kind", "value_key", "required", "confidence"} <= field.keys()


@pytest.mark.parametrize("suffix", ["docx", "hwpx", "xlsx"])
def test_blank_company_form_fill_preserves_original_zip_style_parts_and_text(tmp_path, suffix):
    source = ROOT / f"samples/sample_company_form.{suffix}"
    digest = sha256(source.read_bytes()).hexdigest()
    before = _parts(source)
    output = fill_compatible_template(source, VALUES, tmp_path / f"filled.{suffix}")
    after = _parts(output)
    assert sha256(source.read_bytes()).hexdigest() == digest
    assert set(before) == set(after)
    expected_changes = {"word/document.xml"} if suffix == "docx" else (
        {"Contents/section0.xml", "Preview/PrvText.txt"} if suffix == "hwpx" else {"xl/worksheets/sheet1.xml"})
    assert {name for name in before if before[name] != after[name]} <= expected_changes
    if suffix != "xlsx":
        assert VALUES["본문"] in parse_file(output)["본문"]
    else:
        workbook = load_workbook(output)
        assert workbook.active["B1"].value == VALUES["제목"]
        assert workbook.active["B3"].value == VALUES["본문"]
        assert workbook.active["C5"].value == "=1+2"
        assert workbook.active["B3"].alignment.wrap_text
        assert workbook.active["B3"].font.name == "맑은 고딕"
        workbook.close()


def test_unknown_field_is_left_blank_until_explicit_mapping(tmp_path):
    source = ROOT / "samples/sample_company_form.docx"
    profile = analyze_template(source)
    author = next(field for field in profile["fields"] if field["label"] == "작성자")
    assert author["input_required"] is True
    mapping = {field["id"]: field["value_key"] for field in profile["fields"]}
    output = fill_compatible_template(source, VALUES | {"작성자": "홍길동"}, tmp_path / "filled.docx", mapping)
    assert Document(output).tables[0].cell(3, 1).text == "홍길동"
    assert author["required"] is False


def test_content_control_alias_maps_and_preserves_run_properties(tmp_path):
    document = Document()
    control = OxmlElement("w:sdt")
    properties = OxmlElement("w:sdtPr")
    alias = OxmlElement("w:alias")
    alias.set(qn("w:val"), "요약")
    properties.append(alias)
    properties.append(OxmlElement("w:showingPlcHdr"))
    content = OxmlElement("w:sdtContent")
    paragraph = document.add_paragraph("이곳에 입력")
    paragraph.runs[0].bold = True
    content.append(paragraph._p)
    control.append(properties)
    control.append(content)
    document._element.body.insert(0, control)
    source = tmp_path / "control.docx"
    document.save(source)
    profile = analyze_template(source)
    assert profile["fields"][0]["kind"] == "docx_sdt"
    output = fill_compatible_template(source, VALUES, tmp_path / "filled.docx")
    xml = etree.fromstring(_parts(output)["word/document.xml"])
    assert xml.xpath(".//w:sdt//w:t/text()", namespaces={"w": qn("w:p").split("}")[0][1:]}) == [VALUES["요약"]]
    assert xml.xpath(".//w:sdt//w:b", namespaces={"w": qn("w:p").split("}")[0][1:]})
    assert not xml.xpath(".//w:showingPlcHdr", namespaces={"w": qn("w:p").split("}")[0][1:]})


def test_blank_body_paragraph_can_be_manually_mapped_without_nested_paragraph(tmp_path):
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("").italic = True
    source = tmp_path / "blank.docx"
    document.save(source)
    profile = analyze_template(source)
    field = profile["fields"][0]
    output = fill_compatible_template(source, VALUES, tmp_path / "filled.docx", {field["id"]: "본문"})
    result = Document(output)
    assert result.paragraphs[0].text == VALUES["본문"]
    assert result.paragraphs[0].runs[0].italic


def test_xlsx_placeholders_formula_chart_and_merged_cells_are_preserved(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "제목"
    sheet["B1"] = "{{제목}} / {{요약}}"
    sheet["B1"].font = Font(name="맑은 고딕", bold=True)
    sheet["A2"] = "내용"
    sheet["B2"].alignment = Alignment(wrap_text=True)
    sheet.merge_cells("B2:C2")
    sheet["D1"] = 1
    sheet["D2"] = 2
    sheet["D3"] = "=SUM(D1:D2)"
    chart = BarChart()
    chart.add_data(Reference(sheet, min_col=4, min_row=1, max_row=2))
    sheet.add_chart(chart, "E2")
    source = tmp_path / "chart.xlsx"
    workbook.save(source)
    output = fill_compatible_template(source, VALUES, tmp_path / "filled.xlsx")
    before, after = _parts(source), _parts(output)
    assert {name for name in before if before[name] != after[name]} == {"xl/worksheets/sheet1.xml"}
    result = load_workbook(output)
    assert result.active["B1"].value == VALUES["제목"] + " / " + VALUES["요약"]
    assert result.active["B2"].value == VALUES["본문"]
    assert result.active["C2"].value is None
    assert result.active["D3"].value == "=SUM(D1:D2)"
    assert len(result.active._charts) == 1
    assert result.active["B1"].font.bold
    result.close()


def test_pdf_acroform_canonical_and_widget_values_appearances_and_background(tmp_path):
    source = ROOT / "samples/sample_company_form.pdf"
    digest = sha256(source.read_bytes()).hexdigest()
    output = fill_compatible_template(source, VALUES, tmp_path / "filled.pdf")
    reader = PdfReader(output)
    fields = reader.get_fields()
    for key, value_key in {"title": "제목", "summary": "요약", "body": "본문"}.items():
        assert fields[key]["/V"] == VALUES[value_key]
    for reference in reader.pages[0]["/Annots"]:
        widget = reference.get_object()
        assert widget["/V"] == fields[widget["/T"]]["/V"]
        assert widget["/AP"]["/N"].get_data()
        assert "/CompatCJK" in widget["/DA"]
        appearance = widget["/AP"]["/N"]
        assert "/OriginalAppearance" in appearance["/Resources"]["/XObject"]
    assert reader.pages[0].get_contents().get_data() == PdfReader(source).pages[0].get_contents().get_data()
    assert sha256(source.read_bytes()).hexdigest() == digest


def test_static_pdf_supports_manual_coordinates_korean_and_original_page_text(tmp_path):
    source = ROOT / "samples/sample_static_form.pdf"
    fields = [{"label": key, "value_key": key, "page": 1, "x": 55, "y": y,
               "width": 470, "height": height, "font_size": 10} for key, y, height in (
                   ("제목", 95, 25), ("요약", 165, 70), ("본문", 305, 150))]
    profile = analyze_template(source, manual_fields=fields)
    assert profile["render_mode"] == "overlay"
    output = fill_compatible_template(source, VALUES, tmp_path / "filled.pdf", profile=profile)
    text = PdfReader(output).pages[0].extract_text()
    assert "교육 결과보고서" in text
    assert "교육 실시 결과보고" in text
    assert "95명이 참석함" in text
    assert "후속 교육을 검토함" in text


def test_scanned_or_unlabelled_pdf_has_manual_overlay_path(tmp_path):
    source = tmp_path / "scan.pdf"
    document = canvas.Canvas(str(source), pagesize=(400, 600))
    document.rect(10, 10, 380, 580)
    document.showPage()
    document.save()
    profile = analyze_template(source)
    assert profile["supported"] is True and profile["fields"] == []
    profile = analyze_template(source, [{"page": 1, "x": 20, "y": 20, "width": 350,
                                        "height": 200, "value_key": "본문"}])
    output = fill_compatible_template(source, VALUES, tmp_path / "filled.pdf", profile=profile)
    assert "95명이 참석함" in PdfReader(output).pages[0].extract_text()


def test_xfa_pdf_keeps_xfa_data_and_offers_overlay_fallback(tmp_path):
    source = ROOT / "samples/sample_static_form.pdf"
    writer = PdfWriter()
    writer.clone_document_from_reader(PdfReader(source))
    writer._root_object[NameObject("/AcroForm")] = DictionaryObject({NameObject("/XFA"): TextStringObject("original XML data")})
    xfa_source = tmp_path / "xfa.pdf"
    with xfa_source.open("wb") as stream:
        writer.write(stream)
    profile = analyze_template(xfa_source, [{"page": 1, "x": 55, "y": 95, "width": 470,
                                            "height": 25, "value_key": "제목"}])
    assert profile["supported"] and profile["render_mode"] == "overlay"
    output = fill_compatible_template(xfa_source, VALUES, tmp_path / "filled.pdf", profile=profile)
    assert PdfReader(output).trailer["/Root"]["/AcroForm"]["/XFA"] == "original XML data"


def test_pdf_overflow_rejects_output_without_clipping(tmp_path):
    source = ROOT / "samples/sample_static_form.pdf"
    profile = analyze_template(source, [{"page": 1, "x": 55, "y": 95, "width": 50,
                                         "height": 12, "value_key": "본문", "font_size": 10}])
    output = tmp_path / "filled.pdf"
    output.write_bytes(b"previous file")
    with pytest.raises(TemplateError, match="넘침"):
        fill_compatible_template(source, VALUES, output, profile=profile)
    assert output.read_bytes() == b"previous file"


def test_pdf_coordinates_outside_page_are_not_silently_used():
    profile = analyze_template(ROOT / "samples/sample_static_form.pdf", [{"page": 1, "x": 580, "y": 20,
                              "width": 100, "height": 20, "value_key": "본문"}])
    assert profile["supported"] is False
    assert any("페이지 안" in warning for warning in profile["warnings"])


@pytest.mark.parametrize("suffix", ["hwp", "xls", "unknown"])
def test_legacy_formats_ask_for_conversion_instead_of_fake_support(tmp_path, suffix):
    path = tmp_path / f"old.{suffix}"
    path.write_bytes(b"legacy")
    profile = analyze_template(path)
    assert not profile["supported"]
    assert "변환" in profile["warnings"][0]


def test_invalid_mapping_and_original_overwrite_are_rejected(tmp_path):
    source = ROOT / "samples/sample_company_form.docx"
    with pytest.raises(TemplateError, match="원본과 다른"):
        fill_compatible_template(source, VALUES, source)
    with pytest.raises(TemplateError, match="없는 입력칸"):
        fill_compatible_template(source, VALUES, tmp_path / "filled.docx", {"fake": "본문"})


def test_placeholder_mapping_can_rename_business_field(tmp_path):
    document = Document()
    document.add_paragraph("{{영업 현황}}")
    source = tmp_path / "custom.docx"
    document.save(source)
    output = fill_compatible_template(source, VALUES, tmp_path / "filled.docx", {"placeholder:영업 현황": "본문"})
    assert Document(output).paragraphs[0].text == VALUES["본문"]


def test_profile_cannot_be_reused_with_changed_original(tmp_path):
    source = ROOT / "samples/sample_company_form.docx"
    profile = analyze_template(source)
    assert len(profile["source_sha256"]) == 64
    changed = tmp_path / "changed.docx"
    document = Document(source)
    document.add_paragraph("다른 양식임")
    document.save(changed)
    with pytest.raises(TemplateError, match="분석 이후 변경"):
        fill_compatible_template(changed, VALUES, tmp_path / "output.docx", profile=profile)


def test_inline_content_control_keeps_inline_run_layout(tmp_path):
    document = Document()
    paragraph = document.add_paragraph("요약 ")
    control = OxmlElement("w:sdt")
    properties = OxmlElement("w:sdtPr")
    alias = OxmlElement("w:alias")
    alias.set(qn("w:val"), "요약")
    properties.append(alias)
    properties.append(OxmlElement("w:showingPlcHdr"))
    content = OxmlElement("w:sdtContent")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "입력 안내"
    run.append(text)
    content.append(run)
    control.append(properties)
    control.append(content)
    paragraph._p.append(control)
    source = tmp_path / "inline.docx"
    document.save(source)
    output = fill_compatible_template(source, VALUES, tmp_path / "filled.docx")
    root = etree.fromstring(_parts(output)["word/document.xml"])
    ns = {"w": qn("w:p").split("}")[0][1:]}
    assert not root.xpath(".//w:sdtContent/w:p", namespaces=ns)
    assert root.xpath(".//w:sdtContent/w:r/w:t/text()", namespaces=ns) == [VALUES["요약"]]


def test_xlsx_split_rich_text_placeholder_keeps_run_fonts(tmp_path):
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    workbook = Workbook()
    workbook.active["A1"] = CellRichText("앞 ", TextBlock(InlineFont(b=True, rFont="맑은 고딕"), "{{제"),
                                        TextBlock(InlineFont(i=True), "목}} 뒤"))
    source = tmp_path / "rich.xlsx"
    workbook.save(source)
    original = etree.fromstring(_parts(source)["xl/worksheets/sheet1.xml"])
    output = fill_compatible_template(source, VALUES, tmp_path / "filled.xlsx")
    filled = etree.fromstring(_parts(output)["xl/worksheets/sheet1.xml"])
    old_properties = original.xpath(".//*[local-name()='rPr']")
    new_properties = filled.xpath(".//*[local-name()='rPr']")
    assert [etree.tostring(node) for node in old_properties] == [etree.tostring(node) for node in new_properties]
    result = load_workbook(output)
    assert result.active["A1"].value == "앞 " + VALUES["제목"] + " 뒤"
    result.close()
