from hashlib import sha256
import json
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from lxml import etree
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DictionaryObject, NameObject, TextStringObject
import pytest

from templates import TemplateError, analyze_template, fill_compatible_template, load_form_profile
from templates import profiles
from templates.compatibility import CHECKBOX_NS, WORD_NS

ROOT = Path(__file__).resolve().parents[1]
REGISTERED = sorted((ROOT / "templates/profiles").glob("*.json"))
SECOND_PAGE_TEST_VALUES = {
    'ra_law_form_23_pdf_v2_second_page': {
        '임상시험 목적': '시험 목적(제출금지)', '투여 방법': '시험 투여 방법(제출금지)'},
    'ra_law_form_32_pdf_v2_second_page': {
        '서론': '시험 서론(제출금지)',
        '일련 목록(Line listing)과 요약표의 데이터': '시험 표 데이터(제출금지)'},
}


def _test_values(profile):
    # QA values belong to tests, never the learned profile or model payload.
    if profile.get('entry_id') in SECOND_PAGE_TEST_VALUES:
        parent = json.loads((ROOT / 'templates/profiles' / (profile['parent_profile_id'] + '.json')).read_text(encoding='utf-8'))
        return _test_values(parent) | SECOND_PAGE_TEST_VALUES[profile['entry_id']]
    if 'demo_values' in profile:
        return profile['demo_values']
    manifest = json.loads((ROOT/'evals/ra_extended_public_templates.json').read_text(encoding='utf-8'))
    entry = next(item for item in manifest['documents'] if item['sha256'] == profile['source_sha256'])
    assert entry['full_form_filled'] is entry['submission_ready'] is False
    return entry['synthetic_values']


def test_registered_profiles_have_exact_source_binding_modes_and_non_submission_test_values():
    assert len(REGISTERED) >= 8
    for path in REGISTERED:
        profile = json.loads(path.read_text(encoding="utf-8"))
        assert len(profile["source_sha256"]) == 64
        assert profile["resource_kind"] == "blank_form"
        assert profile["submission_ready"] is False
        if 'demo_values' in profile:
            assert "실제 제출용" in profile["demo_notice"]
        else:
            assert _test_values(profile) and 'demo_notice' not in profile
        assert len({field["id"] for field in profile["fields"]}) == len(profile["fields"])
        for field in profile["fields"]:
            assert field["input_mode"] in {"user_provided", "source_grounded"}
            assert field["input_required"] == (field["input_mode"] == "user_provided")
            assert field["max_chars"] > 0
            assert field["value_key"] in _test_values(profile)


def test_profile_lookup_matches_content_hash_not_original_filename(tmp_path, monkeypatch):
    source = ROOT / "samples/sample_company_form.docx"
    directory = tmp_path / "profiles"
    directory.mkdir()
    profile = analyze_template(source) | {"schema_version": 1, "resource_kind": "blank_form"}
    (directory / "approved.json").write_text(json.dumps(profile), encoding="utf-8")
    monkeypatch.setattr(profiles, "PROFILE_DIRECTORY", directory)
    copied = tmp_path / "renamed.docx"
    copied.write_bytes(source.read_bytes())
    assert load_form_profile(copied, "approved") == profile
    assert load_form_profile(copied) == profile
    assert load_form_profile(copied, "unknown") is None
    copied.write_bytes(b"another document")
    assert load_form_profile(copied, "approved") is None
    with pytest.raises(ValueError):
        load_form_profile(source, "../approved")


@pytest.mark.parametrize("profile_path", REGISTERED, ids=lambda path: path.stem)
def test_registered_public_form_fills_with_fake_values_and_preserves_source(profile_path, tmp_path):
    registered = json.loads(profile_path.read_text(encoding="utf-8"))
    candidates = ([ROOT / registered["source_path"]] if registered.get("source_path") else [])
    candidates += [ROOT / "data/public_templates" / registered["source_filename"],
                  ROOT / "data/public_templates/government" / registered["source_filename"],
                  ROOT / "data/public_templates/international" / registered["source_filename"],
                  ROOT / "data/public_templates/ra" / registered["source_filename"],
                  ROOT / "data/public_templates/business" / registered["source_filename"],
                  ROOT / "data/public_templates/office" / registered["source_filename"],
                  ROOT / "data/public_templates/ra_corporate" / registered["source_filename"]]
    source = next((path for path in candidates if path.is_file()), None)
    if source is None:
        pytest.skip("공개 원본 corpus는 별도로 내려받음. 오프라인 테스트에서는 생략함")
    assert sha256(source.read_bytes()).hexdigest() == registered["source_sha256"]
    profile = load_form_profile(source, profile_path.stem)
    assert profile is not None
    before = source.read_bytes()
    mapping = {field["id"]: field["value_key"] for field in profile["fields"]}
    values = _test_values(profile)
    output = fill_compatible_template(source, values, tmp_path / source.name, mapping, profile)
    assert source.read_bytes() == before
    assert output.is_file()
    from evals.compatibility import _verify_output
    checked = _verify_output(output, profile["format"], set(values.values()))
    assert checked["confirmed_values"] >= 1
    if source.suffix in {".docx", ".hwpx", ".xlsx"}:
        with ZipFile(source) as original, ZipFile(output) as filled:
            for name in original.namelist():
                if name in {"word/document.xml", "Preview/PrvText.txt"} or name.startswith(("Contents/section", "xl/worksheets/")):
                    continue
                assert original.read(name) == filled.read(name)


def test_docx_inline_blank_replacement_preserves_labels_and_run_style(tmp_path):
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("담당자 : ").bold = True
    p.add_run("       ").italic = True
    p.add_run(" (인)")
    source = tmp_path / "source.docx"
    doc.save(source)
    text = p.text
    profile = analyze_template(source) | {"supported": True, "fields": [{
        "id": "docx:word/document.xml:/w:document/w:body/w:p", "label": "담당자", "kind": "docx_inline",
        "value_key": "담당자", "required": True, "anchor_text": text,
        "blank_start": 5, "blank_end": 12, "max_chars": 20}]}
    output = fill_compatible_template(source, {"담당자": "시험홍길동 {{literal}}"}, tmp_path / "filled.docx", profile=profile)
    paragraph = Document(output).paragraphs[0]
    assert paragraph.text == "담당자 :시험홍길동 {{literal}}  (인)"
    assert paragraph.runs[0].bold is True
    assert paragraph.runs[1].italic is True
    bad = profile | {"fields": [profile["fields"][0] | {"blank_start": 0}]}
    with pytest.raises(TemplateError, match="빈 입력"):
        fill_compatible_template(source, {"담당자": "시험"}, tmp_path / "bad.docx", profile=bad)


def test_docx_checkbox_updates_control_state_and_display_without_losing_properties(tmp_path):
    doc = Document()
    sdt = etree.Element(f"{{{WORD_NS}}}sdt", nsmap={"w14": CHECKBOX_NS})
    pr = etree.SubElement(sdt, f"{{{WORD_NS}}}sdtPr")
    alias = etree.SubElement(pr, f"{{{WORD_NS}}}alias")
    alias.set(f"{{{WORD_NS}}}val", "사용자 동의")
    checkbox = etree.SubElement(pr, f"{{{CHECKBOX_NS}}}checkbox")
    checked = etree.SubElement(checkbox, f"{{{CHECKBOX_NS}}}checked")
    checked.set(f"{{{CHECKBOX_NS}}}val", "0")
    for name, value in [("checkedState", "2612"), ("uncheckedState", "2610")]:
        state = etree.SubElement(checkbox, f"{{{CHECKBOX_NS}}}{name}")
        state.set(f"{{{CHECKBOX_NS}}}val", value)
    content = etree.SubElement(sdt, f"{{{WORD_NS}}}sdtContent")
    p = doc.add_paragraph("☐")
    p.runs[0].bold = True
    content.append(p._p)
    doc._element.body.insert(0, sdt)
    source = tmp_path / "source.docx"
    doc.save(source)
    profile = analyze_template(source)
    field = next(item for item in profile["fields"] if item["kind"] == "docx_checkbox")
    output = fill_compatible_template(source, {field["value_key"]: "true"}, tmp_path / "filled.docx", profile=profile)
    with ZipFile(output) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    assert root.find(f".//{{{CHECKBOX_NS}}}checked").get(f"{{{CHECKBOX_NS}}}val") == "1"
    assert root.find(f".//{{{WORD_NS}}}t").text == "☒"
    assert root.find(f".//{{{WORD_NS}}}b") is not None
    with pytest.raises(TemplateError, match="true"):
        fill_compatible_template(source, {field["value_key"]: "추정 동의"}, tmp_path / "bad.docx", profile=profile)


def test_pdf_checkbox_radio_and_choice_update_canonical_value_and_widget_appearance(tmp_path):
    source = ROOT / "samples/sample_selection_form.pdf"
    profile = analyze_template(source)
    fields = {item["id"]: item for item in profile["fields"]}
    assert fields["pdf:consent"]["control_type"] == "checkbox"
    assert fields["pdf:decision"]["control_type"] == "radio"
    assert fields["pdf:department"]["options"] == ["Planning", "Operations"]
    values = {"Consent": "true", "Decision": "/reject", "Department": "Operations"}
    output = fill_compatible_template(source, values, tmp_path / "selected.pdf", profile=profile)
    reader = PdfReader(output)
    actual = reader.get_fields()
    assert actual["consent"]["/V"] == "/Yes"
    assert actual["decision"]["/V"] == "/reject"
    assert actual["department"]["/V"] == "Operations"
    appearances = [str(ref.get_object().get("/AS")) for ref in reader.pages[0]["/Annots"]]
    assert "/Yes" in appearances and "/reject" in appearances and "/Off" in appearances
    with pytest.raises(TemplateError, match="선택값"):
        fill_compatible_template(source, values | {"Decision": "AI 임의 판단"}, tmp_path / "bad.pdf", profile=profile)
    output = fill_compatible_template(source, values | {"Consent": "false"}, tmp_path / "off.pdf", profile=profile)
    assert PdfReader(output).get_fields()["consent"]["/V"] == "/Off"


@pytest.mark.parametrize("signature", [False, True])
def test_signed_and_certified_pdf_cannot_be_modified_even_with_overridden_profile(tmp_path, signature):
    source = ROOT / "samples/sample_company_form.pdf"
    writer = PdfWriter()
    writer.clone_document_from_reader(PdfReader(source))
    if signature:
        field = writer._root_object["/AcroForm"]["/Fields"][0].get_object()
        field[NameObject("/FT")] = NameObject("/Sig")
        field[NameObject("/V")] = DictionaryObject({NameObject("/Type"): NameObject("/Sig")})
    else:
        writer._root_object[NameObject("/Perms")] = DictionaryObject({NameObject("/DocMDP"): DictionaryObject({NameObject("/Type"): NameObject("/Sig")})})
    protected = tmp_path / "certified.pdf"
    writer.write(protected)
    assert analyze_template(protected)["supported"] is False
    forged = analyze_template(source) | {"source_sha256": sha256(protected.read_bytes()).hexdigest(), "supported": True}
    with pytest.raises(TemplateError, match="전자서명"):
        fill_compatible_template(protected, {"제목": "시험"}, tmp_path / "bad.pdf", profile=forged)
