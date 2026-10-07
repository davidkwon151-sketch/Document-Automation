"""Four actual original-PDF regions; synthetic source/mock semantics, API zero."""
from copy import deepcopy
from hashlib import sha256
from pathlib import Path
import importlib.util
import json

from docx import Document
import pdfplumber
import pypdfium2 as pdfium
import pytest
from streamlit.testing.v1 import AppTest

from agent.field_citations import profile_field, split_field_citations
from agent.output_check import verify_output
from agent.pipeline import build_downloads, review_result
from agent.review import CITATION_PATTERN
from agent.retrieve import chunk_documents, load_documents
from parsers import parse_file
from templates import fill_compatible_template
from templates.fill import TemplateError
from app import ra_mvp_service as service

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'outputs/ra-v2-tests' / __import__('uuid').uuid4().hex[:8]
OUT.mkdir(parents=True, exist_ok=True)
spec = importlib.util.spec_from_file_location('original_clinical_test', ROOT / 'tests/test_ra_clinical_safety_pipeline.py')
base = importlib.util.module_from_spec(spec)
spec.loader.exec_module(base)
V2 = ('ra_law_form_23_pdf_v2_second_page', 'ra_law_form_32_pdf_v2_second_page')
BODY = {
    '임상시험 목적': '합성약A 제안 계획의 시험 목적은 자료 수집임',
    '투여 방법': '합성약A 제안 계획의 투여 방법은 1일 1회 5 mg 경구 투여임',
    '서론': '합성약A DSUR 자료는 검토 중임',
    '일련 목록(Line listing)과 요약표의 데이터': '참고 정보는 합성약A DSUR 검토 자료임\n합성약A 중대한 약물이상반응의 일련 목록은 집계 미확정임\n합성약A 중대한 이상사례의 누적 요약표는 12건임',
}

class SecondPageClient(base.ClinicalSafetyClient):
    def __init__(self, form_id):
        super().__init__(form_id.removesuffix('_v2_second_page'))
        self.v2_id = form_id

    def generate_json(self, name, payload):
        reply = super().generate_json(name, payload)
        if name == 'brief':
            reply['목적'] += ' 임상시험 목적 투여 방법 서론 일련 목록 중대한 약물이상반응 중대한 이상사례 누적 요약표 참고 정보 결론 추가 확인 필요함 결론 추가 확인 필요함'
        if name == 'draft':
            for field in payload['template_profile']['fields']:
                key = field['value_key']
                if key not in BODY:
                    continue
                rows = []
                for line in BODY[key].splitlines():
                    source = next(item for item in payload['sources'] if line in item['text'])
                    rows.append(f"{line} [{source['source_id']}]")
                reply[key] = '\n'.join(rows)
        return reply

@pytest.fixture
def evidence(tmp_path, request):
    path = tmp_path / 'synthetic-clinical-dsur-evidence.docx'
    doc = Document()
    doc.add_paragraph('합성 mock 시험자료임. 실제 임상시험·허가·안전성 결과가 아님.')
    for line in base.LINES.values():
        doc.add_paragraph(line)
    for key, value in BODY.items():
        form_id = request.node.callspec.params.get('form_id')
        if (key in ('임상시험 목적', '투여 방법')) != (form_id == V2[0]):
            continue
        kind = '임상시험계획서' if key in ('임상시험 목적', '투여 방법') else 'DSUR 원자료'
        header = ('Clinical Study Protocol SYN-01 Original' if key in ('임상시험 목적', '투여 방법') else 'Development Safety Update Report\nDSUR number: 1\nReporting period: 2026-01-01 to 2026-06-30')
        doc.add_paragraph(f'{header}\n자료 유형: {kind}\n제품명: 합성약A\n항목: {key}\n{value}')
        if key in ('임상시험 목적', '서론'):
            doc.add_paragraph(f'{header}\n자료 유형: {kind}\n제품명: 합성약B\n항목: {key}\n{value.replace("합성약A", "합성약B")}')
    doc.save(path)
    return path

def generate(form_id, evidence):
    client = SecondPageClient(form_id)
    result = service.generate_mvp('합성약A 제안 계획 및 DSUR 검토 중 상태를 원자료 그대로 작성함',
                                  form_id, paths=[evidence], client=client)
    assert result['status'] == 'ready', result.get('review')
    return client, result

def printed(result):
    return {key: split_field_citations(value, profile_field(result['template_profile'], key))[0]
            for key, value in result['draft'].items()}

def export(result):
    return build_downloads(result, confirmed=True, template_paths=result['template_paths'],
                           template_profiles={'pdf': result['template_profile']}, native_review='off')['pdf']

def test_registry_keeps_six_records_and_v1_profiles_identical():
    baseline = json.loads((ROOT / 'tests/fixtures/ra_second_page_v1_baseline.json').read_text(encoding='utf-8'))
    catalog = json.loads((ROOT / 'templates/ra_mvp_catalog.json').read_text(encoding='utf-8'))
    assert catalog['forms'][:6] == baseline['records']
    assert len(service.load_mvp_forms()) == 8
    assert len({r['source_sha256'] for r in catalog['forms']}) == 6
    assert catalog['summary']['unique_registered_physical_field_count'] == 84
    assert catalog['summary']['unique_public_demo_fields'] == 8
    for path, digest in baseline['profile_sha256'].items():
        assert sha256((ROOT / path).read_bytes()).hexdigest() == digest
    for form_id in V2:
        template, profile, record = service.resolve_mvp_form(form_id)
        prior = json.loads((ROOT / record['profile_path'].replace(form_id, form_id.removesuffix('_v2_second_page'))).read_text(encoding='utf-8'))
        assert profile['fields'][:-2] == prior['fields']
        assert len(profile['fields']) == len(prior['fields']) + 2
        assert profile['source_sha256'] == record['source_sha256'] == sha256(template.read_bytes()).hexdigest()
        assert template.name == profile['source_filename']
        assert profile['document_kind'] == record['document_kind']
        assert profile['ra_workflow'] == record['ra_workflow']
        assert all(f['input_mode'] == 'source_grounded' and not f['input_required']
                   and not f['required'] and f['max_chars'] < 500 and f['evidence_document_labels']
                   for f in profile['fields'][-2:])

@pytest.mark.parametrize('form_id', V2)
def test_actual_four_regions_with_sources_original_location_and_png(form_id, evidence):
    client, result = generate(form_id, evidence)
    assert not result['grounding']['blocking'] and not result['completeness']['blocking']
    for key in result['form_record']['user_input_keys']:
        assert result['draft'][key] == ''
    assert all(s['filename'] == evidence.name for s in result['sources'])
    assert all(s['location'] and s['context_text'] for s in result['sources'])
    values = printed(result)
    original = Path(result['template_paths']['pdf'])
    original_hash = sha256(original.read_bytes()).hexdigest()
    directory = OUT / 'artifacts' / form_id
    directory.mkdir(parents=True, exist_ok=True)
    payload = export(result)
    output = directory / 'filled.pdf'
    output.write_bytes(payload)
    evidence_copy = directory / evidence.name
    evidence_copy.write_bytes(evidence.read_bytes())
    check = verify_output(original, output, values, profile=result['template_profile'])
    assert check['status'] == 'passed', check
    with pdfplumber.open(original) as before, pdfplumber.open(output) as after:
        assert len(before.pages) == len(after.pages) == 2
        for field in result['template_profile']['fields'][-2:]:
            rect = (field['x'], field['y'], field['x'] + field['width'], field['y'] + field['height'])
            assert not before.pages[1].crop(rect).extract_text()
            actual = after.pages[1].crop(rect).extract_text()
            assert ''.join(actual.split()) == ''.join(values[field['value_key']].split())
            assert CITATION_PATTERN.findall(result['draft'][field['value_key']])
    # PDFium rendering stays in the pytest main thread, not a worker pool.
    document = pdfium.PdfDocument(output)
    for page_index in range(len(document)):
        page = document[page_index]
        image = page.render(scale=1.5).to_pil()
        image.save(directory / f'filled-{page_index + 1}.png')
        page.close()
    document.close()
    assert sha256(original.read_bytes()).hexdigest() == original_hash
    report = {'scope': 'synthetic/mock production expansion; actual clinical validity not evaluated', 'actual_api_requests': 0,
        'actual_model_responses': 0, 'human_kpi_observations': 0, 'new_original_acquisitions': 0,
        'original_sha256': original_hash, 'profile_sha256': result['form_record']['profile_sha256'],
        'evidence_sha256': sha256(evidence.read_bytes()).hexdigest(),
        'output_sha256': sha256(payload).hexdigest(), 'draft': result['draft'],
        'sources': result['sources'], 'fields': result['template_profile']['fields'][-2:],
        'source_bindings': {s['source_id']: {'document_sha256': sha256(evidence.read_bytes()).hexdigest(),
            'filename': s['filename'], 'page': s['page'], 'sheet': s['sheet'], 'location': s['location'],
            'quote': s['text'], 'context_start': s['context_start'], 'context_end': s['context_end']}
            for s in result['sources']},
        'independent_verification': check, 'actual_model_quality_evaluated': False,
        'native_hancom_acrobat_review': 'not_performed', 'submission_ready': False}
    (directory / 'check.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')

@pytest.mark.parametrize('form_id,key,before,after', [
    (V2[0], '투여 방법', '5 mg', '50 mg'),
    (V2[1], '일련 목록(Line listing)과 요약표의 데이터', '12건', '120건')])
def test_wrong_number_blocks_actual_export(form_id, key, before, after, evidence):
    client, result = generate(form_id, evidence)
    changed = {**result['draft'], key: result['draft'][key].replace(before, after)}
    checked = review_result(result, changed, client=client)
    assert checked['blocking']
    assert any(i['field'] == key and i['code'] in {'ra_quantity_mismatch', 'ra_number_mismatch', 'number_mismatch', 'ra_document_scope_unverified'} for i in checked['warnings'])
    result['draft'] = changed
    with pytest.raises(ValueError):
        export(result)

@pytest.mark.parametrize('form_id', V2)
def test_wrong_product_scope_known_source_blocks_export(form_id, evidence):
    client, result = generate(form_id, evidence)
    key = result['template_profile']['fields'][-2]['value_key']
    line = BODY[key]
    other = next(s for s in chunk_documents(load_documents([evidence])) if line.replace('합성약A', '합성약B') in s['text'])
    if other['source_id'] not in {s['source_id'] for s in result['sources']}:
        result['sources'].append(other)
    old = CITATION_PATTERN.findall(result['draft'][key])[0]
    changed = {**result['draft'], key: result['draft'][key].replace(f'[{old}]', f"[{other['source_id']}]")}
    checked = review_result(result, changed, client=client)
    assert checked['blocking'] and any(i['code'] == 'ra_document_scope_unverified' and i['field'] == key for i in checked['warnings'])
    result['draft'] = changed
    with pytest.raises(ValueError):
        export(result)

@pytest.mark.parametrize('form_id', V2)
def test_source_document_kind_cannot_be_product_leaflet(form_id, evidence):
    client, result = generate(form_id, evidence)
    for source in result['sources']:
        source['text'] = source['text'].replace('자료 유형: 임상시험계획서', '자료 유형: 제품 허가 설명서').replace('자료 유형: DSUR 원자료', '자료 유형: 제품 허가 설명서')
        source['context_text'] = source['context_text'].replace('자료 유형: 임상시험계획서', '자료 유형: 제품 허가 설명서').replace('자료 유형: DSUR 원자료', '자료 유형: 제품 허가 설명서')
        source['context_start'] = source['context_text'].index(source['text'])
        source['context_end'] = source['context_start'] + len(source['text'])
        source['context_sha256'] = sha256(source['context_text'].encode()).hexdigest()
    checked = review_result(result, client=client)
    assert checked['blocking']
    assert any(i['code'] == 'ra_document_scope_unverified' for i in checked['warnings'])
    with pytest.raises(ValueError):
        export(result)

@pytest.mark.parametrize('form_id', V2)
def test_exact_real_public_product_name_citation_is_not_trial_or_dsur_source(form_id):
    result = service.generate_mvp('공개 제품 설명서 데모', form_id, demo=True)
    key = result['template_profile']['fields'][-2]['value_key']
    product_key = result['form_record']['demo_field_keys'][0]
    result['draft'][key] = result['draft'][product_key]
    checked = review_result(result)
    assert checked['blocking']
    assert any(i['code'] == 'ra_document_scope_unverified' and i['field'] == key for i in checked['warnings'])
    with pytest.raises(ValueError):
        export(result)

@pytest.mark.parametrize('form_id', V2)
def test_only_document_kind_and_product_name_do_not_establish_section(form_id):
    from agent.ra import inspect_ra_draft
    _, profile, _ = service.resolve_mvp_form(form_id)
    field = profile['fields'][-2]
    kind = field['evidence_document_labels'][0]
    source = {'source_id': 'STestScope', 'filename': 'synthetic.docx', 'location': '문단 1',
              'text': f'자료 유형: {kind}\n제품명: 합성약A'}
    draft = {field['value_key']: '제품명: 합성약A [STestScope]'}
    checked = inspect_ra_draft(draft, [source], profile=profile)
    assert checked['blocking'] and any(i['code'] == 'ra_document_scope_unverified' for i in checked['issues'])

@pytest.mark.parametrize('form_id', V2)
@pytest.mark.parametrize('origin', ['filename_only', 'user_input'])
def test_filename_or_direct_answer_does_not_classify_actual_evidence(form_id, origin):
    from agent.ra import inspect_ra_draft
    _, profile, _ = service.resolve_mvp_form(form_id)
    field = profile['fields'][-2]
    line = BODY[field['value_key']]
    kind = field['evidence_document_labels'][0]
    text = f'항목: {field["label"]}\n{line}'
    filename = kind + '.docx'
    if origin == 'user_input':
        filename = '사용자 입력'
        text = f'자료 유형: {kind}\n제품명: 합성약A\n' + text
    source = {'source_id': 'STestOrigin', 'filename': filename, 'location': field['label'], 'text': text}
    checked = inspect_ra_draft({field['value_key']: line + ' [STestOrigin]'}, [source], profile=profile)
    assert checked['blocking'] and any(i['code'] == 'ra_document_scope_unverified' for i in checked['issues'])

@pytest.mark.parametrize('form_id', V2)
def test_guidance_title_is_not_the_explicit_source_document_kind(form_id):
    from agent.ra import inspect_ra_draft
    _, profile, _ = service.resolve_mvp_form(form_id)
    field = profile['fields'][-2]
    line = BODY[field['value_key']]
    kind = field['evidence_document_labels'][0]
    source = {'source_id': 'STestGuide', 'filename': 'synthetic-guide.docx', 'location': '문단 1',
              'text': f'자료 유형: {kind} 보완사례집\n항목: {field["label"]}\n{line}'}
    checked = inspect_ra_draft({field['value_key']: line + ' [STestGuide]'}, [source], profile=profile)
    assert checked['blocking'] and any(i['code'] == 'ra_document_scope_unverified' for i in checked['issues'])

@pytest.mark.parametrize('form_id', V2)
@pytest.mark.parametrize('misplacement', ['previous_section', 'following_section', 'changed_document_kind'])
def test_actual_quote_must_belong_to_nearest_preceding_declared_section(form_id, misplacement):
    from agent.ra import inspect_ra_draft
    _, profile, _ = service.resolve_mvp_form(form_id)
    field = profile['fields'][-2]
    line = BODY[field['value_key']]
    kind = field['evidence_document_labels'][0]
    text = f'자료 유형: {kind}\n항목: {field["label"]}\n'
    if misplacement == 'previous_section':
        text += f'앞 절의 다른 자료임\n항목: 다른 항목\n{line}'
    elif misplacement == 'following_section':
        text = f'자료 유형: {kind}\n항목: 다른 항목\n{line}\n항목: {field["label"]}\n다음 절임'
    else:
        text += f'앞 원자료임\n문서 종류: 제품 허가 설명서\n{line}'
    source = {'source_id': 'STestSection', 'filename': 'synthetic-sections.docx', 'location': '문단 1', 'text': text}
    checked = inspect_ra_draft({field['value_key']: line + ' [STestSection]'}, [source], profile=profile)
    assert checked['blocking'] and any(i['code'] == 'ra_document_scope_unverified' for i in checked['issues'])

@pytest.mark.parametrize('form_id', V2)
def test_missing_real_sources_requests_attachment_without_fact_promotion(form_id):
    class MissingSourceClient(SecondPageClient):
        def generate_json(self, name, payload):
            if name == 'brief':
                reply = super().generate_json(name, payload)
                reply['부족한 정보'] = ['후면 임상/안전성 원자료']
                reply['질문'] = ['해당 항목과 자료 유형이 명시된 임상시험계획서 또는 DSUR 원자료를 첨부해 주세요.']
                return reply
            raise AssertionError('No draft/model call when actual RA source is missing')
    result = service.generate_mvp('제품 설명서밖에 없으며 후면 자료는 없음', form_id,
                                  client=MissingSourceClient(form_id))
    assert result['status'] == 'needs_information'
    assert 1 <= len(result['questions']) <= 2 and not result.get('draft')

@pytest.mark.parametrize('form_id', V2)
def test_missing_second_page_source_stays_blank_in_public_demo(form_id):
    result = service.generate_mvp('공개 제품명 데모만 확인', form_id, demo=True)
    assert all(result['draft'][f['value_key']] == '' for f in result['template_profile']['fields'][-2:])
    assert set(f['value_key'] for f in result['template_profile']['fields'][-2:]) <= set(result['missing_form_fields'])
    assert result['submission_ready'] is False
    for f in result['template_profile']['fields'][-2:]:
        with pytest.raises(ValueError, match='직접 입력'):
            service.generate_mvp('임상 목적 임의 입력', form_id, demo=True, field_values={f['value_key']: '임의 값'})

@pytest.mark.parametrize('form_id', V2)
def test_no_clipping_in_real_filler_or_catalog_char_limit(form_id, tmp_path):
    path, profile, _ = service.resolve_mvp_form(form_id)
    field = profile['fields'][-1]
    target = tmp_path / 'overflow.pdf'
    value = '가' * (field['max_chars'] + 1)
    with pytest.raises(TemplateError, match='길이|분량|넘침|검증'):
        fill_compatible_template(path, {field['value_key']: value}, target, profile=profile)
    assert not target.exists()
    # Existing actual glyph/line-height check also blocks content below char cap.
    value = '\n'.join(['가'] * 10)
    assert len(value) < field['max_chars']
    with pytest.raises(TemplateError, match='넘침'):
        fill_compatible_template(path, {field['value_key']: value}, target, profile=profile)
    assert not target.exists()

@pytest.mark.parametrize('form_id', V2)
def test_v1_metrics_and_answer_do_not_carry_into_v2(form_id):
    old_id = form_id.removesuffix('_v2_second_page')
    result = service.generate_mvp('공개자료', old_id, demo=True)
    with pytest.raises(ValueError, match='다른 양식의 KPI'):
        service.generate_mvp('새 v2 양식', form_id, demo=True, previous_metrics=result['metrics'])
    new = service.generate_mvp('새 v2 양식', form_id, demo=True)
    assert new['run_id'] != result['run_id']
    assert not new['answers'] and new['metrics']['mvp_form_id'] == form_id

@pytest.mark.parametrize('form_id', V2)
def test_actual_ui_v1_v2_reset_and_direct_input_permissions(form_id, monkeypatch, tmp_path):
    monkeypatch.setenv('REPORT_AGENT_DATA_DIR', str(tmp_path))
    ui = AppTest.from_file(str(ROOT / 'app/ra_mvp_ui.py'), default_timeout=60).run()
    assert not ui.exception
    ui.selectbox(key='mvp_form').select(form_id.removesuffix('_v2_second_page')).run()
    ui.button(key='mvp_generate').click().run()
    assert not ui.exception
    assert 'mvp_result' in ui.session_state and 'mvp_metrics' in ui.session_state
    ui.session_state['mvp_answers'] = {'이전 질문': '이전 답변'}
    ui.session_state['mvp_exports'] = {'pdf': b'old'}
    ui.session_state['mvp_confirmed'] = True
    ui.selectbox(key='mvp_form').select(form_id).run()
    assert not ui.exception
    for key in ('mvp_result', 'mvp_metrics', 'mvp_answers', 'mvp_exports'):
        assert key not in ui.session_state
    assert 'mvp_confirmed' not in ui.session_state or not ui.session_state['mvp_confirmed']
    _, profile, record = service.resolve_mvp_form(form_id)
    keys = {control.key for control in ui.text_input}
    assert all(f'mvp_input_{form_id}_{key}' not in keys for key in BODY)
    assert all(f'mvp_input_{form_id}_{key}' in keys for key in record['user_input_keys'])
    assert record['source_based_keys'][-2:] == [f['value_key'] for f in profile['fields'][-2:]]

@pytest.fixture(autouse=True)
def actual_api_is_forbidden(monkeypatch):
    from llm.client import LLMClient
    def forbidden(*args, **kwargs):
        raise AssertionError('RA regression uses explicit mock and must not invoke actual API')
    for name in ('generate_json', 'read_image_json', 'embed'):
        monkeypatch.setattr(LLMClient, name, forbidden)
