"""합성 DOCX 반복행: 원본·고정 구조 보존과 위험한 복제 차단."""

from copy import deepcopy
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
import pytest

from templates.fill import TemplateError
from templates.repeat_docx import NS, W14, inventory, transform


def fixture(tmp_path):
    doc = Document()
    doc.add_paragraph('원본 안내; 제출시험 아님')
    table = doc.add_table(rows=4, cols=6)
    for cell, text in zip(table.rows[0].cells, ['항목','내용','수량','단위','선택','확인']):
        cell.text = text
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    row = table.rows[1]
    row.cells[0].merge(row.cells[1]).text = '제품'
    run = row.cells[0].add_paragraph().add_run('{{내역}}')
    run.bold = True
    run.font.name = '맑은 고딕'
    row.cells[3].text = '개'
    height = OxmlElement('w:trHeight')
    height.set(qn('w:val'), '360')
    height.set(qn('w:hRule'), 'atLeast')
    row._tr.get_or_add_trPr().append(height)
    paragraph = row.cells[0].paragraphs[0]._p
    paragraph.set(qn('w14:paraId'), '00000001')
    paragraph.set(qn('w14:textId'), '77777777')
    for cell, checked in [(row.cells[4], False),(row.cells[5], True)]:
        sdt = OxmlElement('w:sdt')
        props = OxmlElement('w:sdtPr')
        identity = OxmlElement('w:id')
        identity.set(qn('w:val'), '10' if not checked else '11')
        props.append(identity)
        if checked:
            control = OxmlElement('w14:checkbox')
            state = OxmlElement('w14:checked')
            state.set(qn('w14:val'), '0')
            control.append(state)
            text = '☐'
        else:
            props.append(OxmlElement('w:showingPlcHdr'))
            control = OxmlElement('w:dropDownList')
            for value, label in [('A','검토 전'),('B','확인 완료')]:
                item = OxmlElement('w:listItem')
                item.set(qn('w:value'), value)
                item.set(qn('w:displayText'), label)
                control.append(item)
            text = '선택 필요'
        props.append(control)
        content = OxmlElement('w:sdtContent')
        content.append(cell.add_paragraph(text)._p)
        sdt.extend([props, content])
        cell._tc.append(sdt)
    for cell, text in zip(table.rows[2].cells, ['합계','','','개','','']):
        cell.text = text
    for cell, text in zip(table.rows[3].cells, ['끝','원본 후속 행','-','-','-','-']):
        cell.text = text
    doc.add_paragraph('원본 종료 안내')
    source = tmp_path / 'repeat.docx'
    doc.save(source)
    with ZipFile(source) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    return source, parts


def plan(parts, count=3, row=2):
    return {'table_id': inventory(parts)[0]['id'], 'row': row, 'count': count}


def mutate(parts, callback):
    copy = dict(parts)
    root = etree.fromstring(copy['word/document.xml'])
    callback(root)
    copy['word/document.xml'] = etree.tostring(root)
    return copy


def normalized(row):
    row = deepcopy(row)
    for node in row.xpath('.//w:sdtPr/w:id', namespaces=NS):
        node.set(qn('w:val'), 'ID')
    for node in row.iter():
        for attr in [qn('w14:paraId'), qn('w14:textId')]:
            if attr in node.attrib:
                node.set(attr, 'ID')
    return etree.tostring(row, method='c14n')


def test_inventory_uses_observed_table_rows_and_original_columns(tmp_path):
    _, parts = fixture(tmp_path)
    record = inventory(parts)[0]
    assert record['id'] == 'docx:word/document.xml:' + record['xpath']
    assert record['part'] == 'word/document.xml' and record['row_count'] == 4
    assert [row['editable'] for row in record['rows']] == [False, True, False, False]
    assert record['rows'][1]['columns'] == ['제품{{내역}}', '', '개', '선택 필요', '☐']
    assert '사용자 확인' in record['rows'][1]['reason']


@pytest.mark.parametrize('count', [1,2,5,200])
def test_expansion_preserves_original_all_styles_and_suffix_only_fresh_ids(tmp_path, count):
    source, parts = fixture(tmp_path)
    original = dict(parts)
    source_bytes = source.read_bytes()
    edits = transform(parts, plan(parts, count))
    assert parts == original and source.read_bytes() == source_bytes
    if count == 1:
        assert edits == {}
        return
    assert set(edits) == {'word/document.xml'}
    old = etree.fromstring(parts['word/document.xml'])
    new = etree.fromstring(edits['word/document.xml'])
    old_rows = old.xpath('.//w:tbl/w:tr', namespaces=NS)
    new_rows = new.xpath('.//w:tbl/w:tr', namespaces=NS)
    assert len(new_rows) == len(old_rows) + count - 1
    for before, after in zip(old_rows[:2], new_rows[:2]):
        assert etree.tostring(before, method='c14n') == etree.tostring(after, method='c14n')
    for before, after in zip(old_rows[2:], new_rows[count + 1:]):
        assert etree.tostring(before, method='c14n') == etree.tostring(after, method='c14n')
    assert all(normalized(row) == normalized(old_rows[1]) for row in new_rows[2:count + 1])
    identifiers = new.xpath('.//w:sdtPr/w:id/@w:val', namespaces=NS)
    assert len(set(identifiers)) == len(identifiers)
    para_ids = new.xpath('//@w14:paraId | //@w14:textId', namespaces=NS)
    assert len(set(para_ids)) == len(para_ids)
    assert all(0 < int(value, 16) < 0x80000000 for value in para_ids)
    assert new.xpath('/w:document/w:body/w:p//w:t/text()', namespaces=NS) == old.xpath('/w:document/w:body/w:p//w:t/text()', namespaces=NS)
    exported = tmp_path / 'expanded.docx'
    with ZipFile(exported, 'w') as archive:
        for name, data in parts.items():
            archive.writestr(name, edits.get(name, data))
    reopened = Document(exported)
    assert len(reopened.tables[0].rows) == count + 3


@pytest.mark.parametrize('changes', [{'count':0},{'count':201},{'count':True},{'count':'2'},{'row':0},{'row':9},{'row':False},{'table_id':'docx:word/document.xml:/fake'},{'extra':1}])
def test_invalid_plan_is_rejected_without_mutation(tmp_path, changes):
    _, parts = fixture(tmp_path)
    before = dict(parts)
    command = plan(parts)
    command.update(changes)
    with pytest.raises(TemplateError):
        transform(parts, command)
    assert parts == before


@pytest.mark.parametrize('row', [1,3,4])
def test_header_total_and_noninput_rows_are_rejected(tmp_path, row):
    _, parts = fixture(tmp_path)
    with pytest.raises(TemplateError):
        transform(parts, plan(parts, row=row))


@pytest.mark.parametrize('tag', ['fldSimple','bookmarkStart','commentRangeStart','footnoteReference','drawing','hyperlink','permStart','ins'])
def test_reference_and_relationship_structures_block_expansion(tmp_path, tag):
    _, parts = fixture(tmp_path)
    altered = mutate(parts, lambda root: root.xpath('.//w:tbl/w:tr[2]/w:tc/w:p', namespaces=NS)[0].append(OxmlElement('w:' + tag)))
    assert not inventory(altered)[0]['rows'][1]['editable']
    with pytest.raises(TemplateError):
        transform(altered, plan(altered))


@pytest.mark.parametrize('index,state', [(2,'restart'),(2,'continue'),(3,'continue'),(3,None)])
def test_vertical_merge_prototype_and_continuation_boundary_block(tmp_path, index, state):
    _, parts = fixture(tmp_path)
    def add(root):
        props = root.xpath(f'.//w:tbl/w:tr[{index}]/w:tc/w:tcPr', namespaces=NS)[0]
        merge = OxmlElement('w:vMerge')
        if state:
            merge.set(qn('w:val'), state)
        props.append(merge)
    altered = mutate(parts, add)
    with pytest.raises(TemplateError, match='세로 병합'):
        transform(altered, plan(altered))


@pytest.mark.parametrize('mode', ['locked','bound','selected','filled','missing_checkbox_state','malformed_id'])
def test_unsafe_native_controls_are_source_classified(tmp_path, mode):
    _, parts = fixture(tmp_path)
    def add(root):
        props = root.xpath('.//w:tbl/w:tr[2]//w:sdtPr', namespaces=NS)[0]
        if mode in {'locked','bound'}:
            props.append(OxmlElement('w:lock' if mode == 'locked' else 'w:dataBinding'))
        elif mode == 'selected':
            root.xpath('.//w14:checked', namespaces=NS)[0].set(qn('w14:val'), '1')
        elif mode == 'missing_checkbox_state':
            state = root.xpath('.//w14:checked', namespaces=NS)[0]
            state.getparent().remove(state)
        elif mode == 'malformed_id':
            props.find(qn('w:id')).set(qn('w:val'), 'not-an-id')
        else:
            props.remove(props.find(qn('w:showingPlcHdr')))
    altered = mutate(parts, add)
    with pytest.raises(TemplateError):
        transform(altered, plan(altered))


@pytest.mark.parametrize('enforcement', ['1','true','on'])
def test_document_protection_remains_intact_and_blocks(tmp_path, enforcement):
    _, parts = fixture(tmp_path)
    root = etree.fromstring(parts['word/settings.xml'])
    protection = OxmlElement('w:documentProtection')
    protection.set(qn('w:enforcement'), enforcement)
    root.append(protection)
    parts['word/settings.xml'] = etree.tostring(root)
    with pytest.raises(TemplateError, match='보호'):
        transform(parts, plan(parts))


def test_signed_package_is_not_modified(tmp_path):
    _, parts = fixture(tmp_path)
    parts['_xmlsignatures/sig1.xml'] = b'<signature/>'
    with pytest.raises(TemplateError, match='서명'):
        transform(parts, plan(parts))


def test_xml_default_namespace_and_blank_plain_row_work_without_control_ids(tmp_path):
    _, parts = fixture(tmp_path)
    parts['word/document.xml'] = ('<document xmlns="' + NS['w'] + '"><body><tbl><tr><tc><p><r><t>항목</t></r></p></tc><tc><p/></tc></tr></tbl></body></document>').encode()
    edit = transform(parts, plan(parts, 2, row=1))
    assert len(etree.fromstring(edit['word/document.xml']).xpath('.//w:tr', namespaces=NS)) == 2


@pytest.mark.parametrize('state', ['0','false','off'])
def test_unchecked_native_states_remain_exactly_as_source(tmp_path, state):
    _, parts = fixture(tmp_path)
    altered = mutate(parts, lambda root: root.xpath('.//w14:checked', namespaces=NS)[0].set(qn('w14:val'), state))
    edits = transform(altered, plan(altered))
    assert etree.fromstring(edits['word/document.xml']).xpath('.//w14:checked/@w14:val', namespaces=NS) == [state] * 3


@pytest.mark.parametrize('existing', ['1','+1','00001',' 1 '])
def test_existing_ids_in_other_parts_are_reserved(tmp_path, existing):
    _, parts = fixture(tmp_path)
    parts['word/header9.xml'] = ('<w:hdr xmlns:w="' + NS['w'] + '" xmlns:w14="' + W14 + '"><w:sdt><w:sdtPr><w:id w:val="' + existing + '"/></w:sdtPr><w:sdtContent><w:p w14:paraId="00000002" w14:textId="00000003"/></w:sdtContent></w:sdt></w:hdr>').encode()
    edits = transform(parts, plan(parts))
    root = etree.fromstring(edits['word/document.xml'])
    assert '1' not in root.xpath('.//w:sdtPr/w:id/@w:val', namespaces=NS)
    assert not {'00000002','00000003'} & set(root.xpath('//@w14:paraId | //@w14:textId', namespaces=NS))
    assert set(edits) == {'word/document.xml'}


def test_opaque_relationship_attribute_is_never_copied(tmp_path):
    _, parts = fixture(tmp_path)
    altered = mutate(parts, lambda root: root.xpath('.//w:tbl/w:tr[2]', namespaces=NS)[0].set(qn('r:id'), 'rId999'))
    with pytest.raises(TemplateError, match='참조'):
        transform(altered, plan(altered))


def test_unknown_id_attribute_cannot_be_silently_duplicated(tmp_path):
    _, parts = fixture(tmp_path)
    altered = mutate(parts, lambda root: root.xpath('.//w:tbl/w:tr[2]/w:tc/w:p', namespaces=NS)[0].set('{urn:unknown}recordId', 'record-one'))
    with pytest.raises(TemplateError, match='식별자'):
        transform(altered, plan(altered))
