from copy import deepcopy
import json
from pathlib import Path

from docx import Document
import pytest

from agent import template_learning as learning


ROOT = Path(__file__).resolve().parents[1]


class FakeClient:
    def __init__(self, mutate=None, image_fields=None):
        self.calls = []
        self.mutate = mutate
        self.image_fields = image_fields

    def generate_json(self, prompt, payload):
        self.calls.append((prompt, payload))
        fields = []
        for original in payload['fields']:
            direct = bool(learning.DIRECT_INPUT.search(original['label'])) or bool(original.get('control_type'))
            fields.append(dict(id=original['id'], kind=original['kind'], value_key=original['label'], required=False,
                               input_required=direct, input_mode='user_provided' if direct else 'source_grounded',
                               max_chars=120, confidence=.9))
        response = {'fields': fields}
        if self.mutate:
            self.mutate(response)
        return response

    def read_image_json(self, image, *, prompt_name, payload):
        self.calls.append((prompt_name, payload))
        return {'fields': deepcopy(self.image_fields or [])}


@pytest.fixture
def unknown_form(tmp_path):
    path = tmp_path / 'unknown.docx'
    doc = Document()
    table = doc.add_table(rows=3, cols=2)
    for row, label in zip(table.rows, ['작성자', '예산', '사업 개요']):
        row.cells[0].text = label
    doc.save(path)
    return path


def test_registered_profile_skips_llm():
    path = ROOT / 'data/public_templates/celltrion_attendance_2026.docx'
    if not path.exists():
        pytest.skip('공개 원본은 오프라인 환경에 없을 수 있음')
    assert learning.learn_template(path, object())['learning']['origin'] == 'registered'


def test_unknown_mapping_keeps_original_and_numeric_style(unknown_form, tmp_path):
    before = unknown_form.read_bytes()
    client = FakeClient()
    profile = learning.learn_template(unknown_form, client, cache_dir=tmp_path/'cache')
    assert profile['learning']['needs_confirmation']
    assert client.calls[0][0] == 'template'
    assert all(field['id'] for field in client.calls[0][1]['fields'])
    assert next(f for f in profile['fields'] if f['label'] == '예산')['narrative_style_required'] is False
    assert next(f for f in profile['fields'] if f['label'] == '작성자')['input_mode'] == 'user_provided'
    assert unknown_form.read_bytes() == before


@pytest.mark.parametrize('mutation', [
    lambda r: r['fields'][0].update(id='invented'),
    lambda r: r['fields'][0].update(kind='pdf_overlay'),
    lambda r: r['fields'][0].update(max_chars=0),
    lambda r: r['fields'][0].update(max_chars=True),
    lambda r: r['fields'][0].update(max_chars=10001),
    lambda r: r['fields'][0].update(value='invented personal value'),
    lambda r: r['fields'][0].update(input_required=False, input_mode='source_grounded'),
    lambda r: r['fields'].append(deepcopy(r['fields'][0])),
    lambda r: r['fields'].pop(),
])
def test_invalid_ai_mapping_is_rejected(unknown_form, tmp_path, mutation):
    with pytest.raises(ValueError):
        learning.learn_template(unknown_form, FakeClient(mutation), cache_dir=tmp_path/'cache')


def test_confirmed_cache_reuse_and_tamper(unknown_form, tmp_path):
    cache = tmp_path/'cache'
    profile = learning.learn_template(unknown_form, FakeClient(), cache_dir=cache)
    with pytest.raises(ValueError, match='확인'):
        learning.save_learned_profile(unknown_form, profile, cache_dir=cache)
    profile['values'] = {'작성자': 'PRIVATE'}
    profile['fields'][0]['value'] = 'PRIVATE'
    target = learning.save_learned_profile(unknown_form, profile, cache_dir=cache, user_confirmed=True)
    assert 'PRIVATE' not in target.read_text(encoding='utf-8')
    assert learning.learn_template(unknown_form, object(), cache_dir=cache)['learning']['origin'] == 'confirmed_cache'
    payload = json.loads(target.read_text(encoding='utf-8'))
    payload['profile']['fields'][0]['max_chars'] = 999
    target.write_text(json.dumps(payload), encoding='utf-8')
    with pytest.raises(ValueError, match='변조'):
        learning.load_learned_profile(unknown_form, cache)


@pytest.mark.parametrize('private_key', ['field_values', 'draft', 'locked_fields', 'sources', 'documents'])
def test_confirmed_cache_and_model_metadata_exclude_actual_inputs(unknown_form, tmp_path, private_key):
    from agent.brief import model_profile
    profile = learning.learn_template(unknown_form, FakeClient(), cache_dir=tmp_path / 'cache')
    profile[private_key] = {'담당자': 'AUDIT-PRIVATE-USER-VALUE'}
    before = deepcopy(profile)
    target = learning.save_learned_profile(unknown_form, profile, cache_dir=tmp_path / 'cache', user_confirmed=True)
    assert 'AUDIT-PRIVATE-USER-VALUE' not in target.read_text(encoding='utf-8')
    assert private_key not in learning.load_learned_profile(unknown_form, tmp_path / 'cache')
    assert private_key not in model_profile(profile)
    assert profile == before


def test_cache_cannot_omit_an_independent_required_field(unknown_form, tmp_path):
    profile = learning.learn_template(unknown_form, FakeClient(), cache_dir=tmp_path / 'cache')
    profile['fields'][0]['required'] = True
    mapping = {field['id']: field['value_key'] for field in profile['fields'][1:]}
    with pytest.raises(ValueError, match='필수'):
        learning.save_learned_profile(unknown_form, profile, mapping, cache_dir=tmp_path / 'cache', user_confirmed=True)


def test_cache_version_and_hash_are_bound(unknown_form, tmp_path, monkeypatch):
    cache = tmp_path/'cache'
    profile = learning.learn_template(unknown_form, FakeClient(), cache_dir=cache)
    learning.save_learned_profile(unknown_form, profile, cache_dir=cache, user_confirmed=True)
    monkeypatch.setattr(learning, 'ENGINE_VERSION', 'next')
    assert learning.load_learned_profile(unknown_form, cache) is None
    other = tmp_path/'other.docx'
    Document().save(other)
    assert learning.load_learned_profile(other, cache) is None


def test_duplicate_value_can_fill_repeated_field(unknown_form, tmp_path):
    def repeat(response):
        for field in response['fields']:
            field.update(value_key='작성자', input_required=True, input_mode='user_provided', max_chars=40)
    profile = learning.learn_template(unknown_form, FakeClient(repeat), cache_dir=tmp_path/'cache')
    assert len({f['value_key'] for f in profile['fields']}) == 1


def test_same_value_conflicting_rules_rejected(unknown_form, tmp_path):
    def bad(response):
        for field in response['fields']:
            field.update(value_key='작성자', input_required=True, input_mode='user_provided')
        response['fields'][1]['max_chars'] = 10
    with pytest.raises(ValueError, match='같아야'):
        learning.learn_template(unknown_form, FakeClient(bad), cache_dir=tmp_path/'cache')


def test_pdf_image_coordinates_and_overlap(tmp_path, monkeypatch):
    monkeypatch.setattr('parsers.extended.render_pdf_page', lambda *args: b'image')
    source = ROOT/'samples/sample_static_form.pdf'
    client = FakeClient(image_fields=[
        {'label':'사업 개요', 'x':80, 'y':100, 'width':100, 'height':20},
        {'label':'일정', 'x':100, 'y':100, 'width':100, 'height':20}])
    profile = learning.learn_template(source, client, cache_dir=tmp_path/'cache')
    assert client.calls[0][0] == 'template_image'
    assert client.calls[0][1]['coordinate_system'] == 'top-left points 72pt/in'
    assert all(f['id'].startswith('learned-pdf:') for f in profile['fields'])
    assert any('겹침' in warning for warning in profile['warnings'])


@pytest.mark.parametrize('coordinates', [dict(x=-1,y=20,width=40,height=20),dict(x=5000,y=0,width=40,height=20),dict(x=1,y=2,width=float('nan'),height=20)])
def test_image_out_of_page_is_rejected(tmp_path, monkeypatch, coordinates):
    monkeypatch.setattr('parsers.extended.render_pdf_page', lambda *args: b'image')
    with pytest.raises(ValueError):
        learning.learn_template(ROOT/'samples/sample_static_form.pdf', FakeClient(image_fields=[dict(label='본문', **coordinates)]),cache_dir=tmp_path/'cache')


def test_unsupported_does_not_call_llm(tmp_path):
    path = tmp_path/'legacy.hwp'
    path.write_bytes(b'legacy')
    assert learning.learn_template(path, object(),cache_dir=tmp_path/'cache')['learning']['origin'] == 'unsupported'


@pytest.mark.parametrize('suffix', ['hwpx','xlsx'])
def test_other_office_formats_supply_context_and_stable_ids(suffix, tmp_path):
    client = FakeClient()
    profile = learning.learn_template(ROOT/f'samples/sample_company_form.{suffix}',client,cache_dir=tmp_path/'cache')
    assert profile['format'] == suffix and profile['fields']
    assert all(field['context'] for field in client.calls[0][1]['fields'])


def test_low_confidence_warns_without_claiming_confirmation(unknown_form,tmp_path):
    client = FakeClient(lambda response: response['fields'][0].update(confidence=.2))
    profile = learning.learn_template(unknown_form,client,cache_dir=tmp_path/'cache')
    assert any('확신도' in warning for warning in profile['warnings'])
    assert profile['learning']['needs_confirmation'] is True


@pytest.fixture
def mixed_language_form(tmp_path):
    path=tmp_path/'foreign_company.docx'
    document=Document()
    labels=['Prepared by / 작성자','Department','Applicant','Tax ID','Budget (예산)','Business purpose / 사업 개요']
    table=document.add_table(rows=len(labels),cols=2)
    for row,label in zip(table.rows,labels):row.cells[0].text=label
    document.save(path)
    return path


def test_mixed_korean_english_form_retains_keys_and_metadata_rules(mixed_language_form,tmp_path):
    before=mixed_language_form.read_bytes()
    profile=learning.learn_template(mixed_language_form,FakeClient(),cache_dir=tmp_path/'cache')
    fields={field['label']:field for field in profile['fields']}
    assert set(fields)=={'Prepared by / 작성자','Department','Applicant','Tax ID','Budget (예산)','Business purpose / 사업 개요'}
    assert all(field['value_key']==label for label,field in fields.items())
    for label in ('Prepared by / 작성자','Department','Applicant','Tax ID'):
        assert fields[label]['input_required'] and fields[label]['input_mode']=='user_provided'
        assert fields[label]['narrative_style_required'] is False
    assert not fields['Budget (예산)']['input_required']
    assert fields['Budget (예산)']['narrative_style_required'] is False
    assert fields['Business purpose / 사업 개요']['narrative_style_required'] is True
    assert mixed_language_form.read_bytes()==before


@pytest.mark.parametrize('label',['Name','Author','Prepared by','Prepared_by','PreparedBy','Department','Applicant',
    'Signature','E-mail','Phone','Consent','Vote','Account','SSN','Tax ID','Tax_ID','Tax Identification Number',
    'Employee ID','Date of birth','Social Security Number','Approved by','Approver','Name / 성명',
    'Requested by','FirstName','FullName','BankAccountNumber'])
def test_english_sensitive_label_cannot_be_remapped_into_generated_body(label,tmp_path):
    path=tmp_path/'sensitive.docx';document=Document();table=document.add_table(rows=1,cols=2)
    table.cell(0,0).text=label;document.save(path)
    def bypass(response):
        response['fields'][0].update(value_key='본문',input_required=False,input_mode='source_grounded')
    with pytest.raises(ValueError,match='직접 입력'):
        learning.learn_template(path,FakeClient(bypass),cache_dir=tmp_path/'cache')


def test_english_mapping_key_is_protected_even_with_neutral_korean_label(unknown_form,tmp_path):
    def bypass(response):
        response['fields'][1].update(value_key='Author',input_required=False,input_mode='source_grounded')
    with pytest.raises(ValueError,match='직접 입력'):
        learning.learn_template(unknown_form,FakeClient(bypass),cache_dir=tmp_path/'cache')


def test_old_confirmed_cache_is_invalidated_for_new_direct_input_rules(unknown_form,tmp_path,monkeypatch):
    cache=tmp_path/'cache'
    with monkeypatch.context() as old:
        old.setattr(learning,'ENGINE_VERSION','template-mapping-1')
        profile=learning.learn_template(unknown_form,FakeClient(),cache_dir=cache)
        learning.save_learned_profile(unknown_form,profile,cache_dir=cache,user_confirmed=True)
    assert learning.load_learned_profile(unknown_form,cache) is None
    client=FakeClient()
    assert learning.learn_template(unknown_form,client,cache_dir=cache)['learning']['needs_confirmation']
    assert client.calls


def test_pptx_mapping_supplies_table_row_and_shape_context_without_mutating_source(tmp_path):
    from templates import analyze_template
    source=ROOT/'samples/sample_company_form.pptx';before=source.read_bytes()
    client=FakeClient()
    profile=learning.learn_template(source,client,cache_dir=tmp_path/'cache')
    originals={field['id']:field for field in analyze_template(source)['fields']}
    supplied=client.calls[0][1]['fields']
    assert profile['format']=='pptx'
    assert any(field['kind']=='pptx_cell' for field in supplied)
    for field in supplied:
        assert field['id'] in originals and 0<len(field['context'])<=800
        if field['kind']=='pptx_cell':assert field['label'] in field['context']
    assert source.read_bytes()==before


def test_pptx_author_shape_retains_original_identity_in_context(tmp_path):
    from zipfile import ZipFile
    from lxml import etree
    from templates import analyze_template
    source=ROOT/'samples/sample_company_form.pptx'
    target=tmp_path/'author_slide.pptx'
    with ZipFile(source) as archive:
        parts={name:archive.read(name) for name in archive.namelist()}
    root=etree.fromstring(parts['ppt/slides/slide1.xml'])
    shapes=root.xpath(".//*[local-name()='sp' and .//*[local-name()='txBody']]")
    shape=shapes[0]
    identity=shape.xpath(".//*[local-name()='cNvPr']")[0];identity.set('name','Prepared by')
    for text in shape.xpath(".//*[local-name()='t']"):text.text=''
    parts['ppt/slides/slide1.xml']=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
    with ZipFile(target,'w') as archive:
        for name,data in parts.items():archive.writestr(name,data)
    fields=analyze_template(target)['fields'];contexts=learning._contexts(target,fields)
    item=next(field for field in contexts if field['kind']=='pptx_text' and field['label']=='Prepared by')
    assert 'Prepared by' in item['context']
    def bypass(response):
        field=next(field for field in response['fields'] if field['id']==item['id'])
        field.update(value_key='본문',input_required=False,input_mode='source_grounded')
    with pytest.raises(ValueError,match='직접 입력'):
        learning.learn_template(target,FakeClient(bypass),cache_dir=tmp_path/'cache')
