"""공개 원본 및 자체 fixture 양식 호환 검증. python -m evals.compatibility."""

from __future__ import annotations

import argparse
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import time

from pypdf import PdfReader, PdfWriter

from parsers import ParseError, parse_file
from templates import analyze_template, fill_compatible_template, load_form_profile


ROOT = Path(__file__).resolve().parents[1]
SUPPORTED = {"docx", "hwpx", "xlsx", "pdf", "pptx"}
PROVENANCES = {"official_catalog", "official_government", "official_tech", "official_international",
               "official_ra", "official_business", "official_office", "unregistered_local", "synthetic_fixture"}
CHECK_VALUES = {"제목": "양식 호환성 검증", "요약": "□ 양식 구조 검증용 문구임",
                "본문": "□ 양식 구조 검증용 문구임\n○ 사용자 검토를 요청함\n- 업무 자료로 교체함"}


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        while block := stream.read(1024 * 1024):
            digest.update(block)
    return digest.hexdigest()


def collect_jobs(directory: Path, *, group: str | None = None, provenance: str | None = None,
                 include_samples=True, include_government=True, include_tech=True, include_international=True,
                 include_ra=True, include_business=True, include_office=True, download=False) -> list[dict]:
    """검증된 카탈로그 항목을 로컬 파일과 연결하며 미등록 원본은 별도로 표시함."""
    from templates.catalog import download_public_template, list_public_templates
    if provenance is not None and provenance not in PROVENANCES:
        raise ValueError("등록된 출처 분류를 선택해야 함")
    if group is not None and provenance not in {None, "official_catalog"}:
        raise ValueError("기업집단 필터는 기업 공개 카탈로그에만 적용할 수 있음")
    directory = Path(directory)
    entries = list_public_templates(group) if provenance in {None, "official_catalog"} else []
    jobs, registered = [], set()
    for entry in entries:
        path = directory / entry["filename"]
        registered.add(path.name)
        download_error = None
        if download and not path.is_file() and entry.get("download_status") == "verified":
            try:
                path = download_public_template(entry, directory)
            except Exception as exc:
                download_error = str(exc)
        job = {"id": entry["id"], "path": str(path.resolve()), "format": path.suffix.lower().lstrip("."),
               "group": entry.get("group"), "organization": entry.get("affiliate") or entry.get("owner") or entry.get("group"),
               "source_url": entry.get("source_url"), "source_kind": entry.get("source_kind"),
               "provenance": "official_catalog", "usage": "blank_form" if entry.get("source_kind") == "public_form" else "layout_reference",
               "expected_sha256": entry.get("sha256"), "catalog_download_status": entry.get("download_status")}
        if download_error:
            job["download_error"] = download_error
        jobs.append(job)
    if group is None and include_government and provenance in {None, "official_government"}:
        from templates.government import list_government_templates, download_government_template
        for entry in list_government_templates():
            path = directory / "government" / entry["filename"]
            error = None
            if download and not path.is_file() and entry.get("download_status") == "verified":
                try:
                    path = download_government_template(entry, directory / "government")
                except Exception as exc:
                    error = str(exc)
            usage = "blank_form" if entry.get("resource_kind") == "blank_form" else "layout_reference"
            if "변경금지" in entry.get("notes", ""):
                usage = "read_only_form"
            job = {"id": entry["id"], "path": str(path.resolve()), "format": path.suffix.lstrip("."),
                   "group": entry["agency"], "organization": entry["agency"], "source_url": entry["source_url"],
                   "source_kind": entry["source_kind"], "provenance": "official_government", "usage": usage,
                   "expected_sha256": entry.get("sha256"), "catalog_download_status": entry.get("download_status")}
            if error:
                job["download_error"] = error
            jobs.append(job)
    if group is None and include_tech and provenance in {None, "official_tech"}:
        from templates.tech import list_tech_templates, download_tech_template
        for entry in list_tech_templates():
            path = directory / "tech" / entry["filename"]
            error = None
            if download and not path.is_file() and entry.get("download_status") == "verified":
                try:
                    path = download_tech_template(entry, directory / "tech")
                except Exception as exc:
                    error = str(exc)
            job = {"id": entry["id"], "path": str(path.resolve()), "format": path.suffix.lower().lstrip("."),
                   "group": entry.get("company"), "organization": entry.get("affiliate") or entry.get("owner"),
                   "source_url": entry.get("source_url"), "source_kind": entry.get("source_kind"),
                   "provenance": "official_tech", "usage": entry.get("resource_kind", "layout_reference"),
                   "use_classification": entry.get("use_classification"), "company_internal": entry.get("company_internal", False),
                   "expected_sha256": entry.get("sha256"), "catalog_download_status": entry.get("download_status")}
            if error:
                job["download_error"] = error
            jobs.append(job)
    if group is None and include_international and provenance in {None, "official_international"}:
        from templates.international import list_international_templates, download_international_template
        for entry in list_international_templates():
            path = directory / 'international' / entry['filename']
            error = None
            if download and not path.is_file():
                try:
                    path = download_international_template(entry, directory / 'international')
                except Exception as exc:
                    error = str(exc)
            job = {'id': entry['id'], 'path': str(path.resolve()), 'format': path.suffix.lower().lstrip('.'),
                   'group': entry['company'], 'organization': entry['company'], 'source_url': entry['source_url'],
                   'source_kind': entry['source_kind'], 'provenance': 'official_international',
                   'usage': entry['resource_kind'], 'use_classification': entry['use_classification'],
                   'company_internal': False, 'expected_sha256': entry['sha256'], 'catalog_download_status': entry['download_status']}
            if error:
                job['download_error'] = error
            jobs.append(job)
    if group is None and include_ra and provenance in {None, "official_ra"}:
        from templates.ra import list_ra_templates, download_ra_template
        for entry in list_ra_templates():
            path = directory / 'ra' / entry['filename']
            error = None
            if download and not path.is_file() and entry['download_status'] == 'verified':
                try:
                    path = download_ra_template(entry, directory / 'ra')
                except Exception as exc:
                    error = str(exc)
            job = {'id': entry['id'], 'path': str(path.resolve()), 'format': path.suffix.lower().lstrip('.'),
                   'group': '식품의약품안전처', 'organization': entry['publisher'], 'source_url': entry['source_url'],
                   'source_kind': entry['source_kind'], 'provenance': 'official_ra', 'usage': entry['resource_kind'],
                   'use_classification': entry['use_classification'], 'company_internal': False,
                   'workflows': entry['workflows'], 'expected_sha256': entry['sha256'],
                   'catalog_download_status': entry['download_status']}
            if error:
                job['download_error'] = error
            jobs.append(job)
    for prefix, enabled in (('business', include_business), ('office', include_office)):
        if group is not None or not enabled or provenance not in {None, "official_" + prefix}:
            continue
        from importlib import import_module
        try:
            module = import_module('templates.' + prefix)
        except ImportError:
            continue
        for entry in getattr(module, 'list_' + prefix + '_templates')():
            path = directory / prefix / entry['filename']
            error = None
            if download and not path.is_file() and entry['download_status'] == 'verified':
                try:
                    path = getattr(module, 'download_' + prefix + '_template')(entry, directory / prefix)
                except Exception as exc:
                    error = str(exc)
            job = {'id': entry['id'], 'path': str(path.resolve()), 'format': path.suffix.lower().lstrip('.'),
                   'group': entry['publisher'], 'organization': entry.get('organization', entry['publisher']),
                   'source_url': entry['source_url'], 'source_kind': entry['source_kind'],
                   'provenance': 'official_' + prefix, 'usage': entry['resource_kind'],
                   'use_classification': entry['use_classification'], 'company_internal': False,
                   'workflows': entry['workflows'], 'expected_sha256': entry['sha256'],
                   'catalog_download_status': entry['download_status']}
            if error:
                job['download_error'] = error
            jobs.append(job)
    if group is None and directory.exists() and provenance in {None, "unregistered_local"}:
        for path in sorted(directory.iterdir()):
            if path.is_file() and path.name not in registered and path.suffix.lower().lstrip(".") in SUPPORTED:
                jobs.append({"id": "local:" + path.name, "path": str(path.resolve()), "format": path.suffix.lower().lstrip("."),
                             "group": None, "organization": None, "source_url": None, "source_kind": None,
                             "provenance": "unregistered_local", "usage": "layout_reference"})
    if include_samples and provenance in {None, "synthetic_fixture"}:
        for name in ("sample_company_form.docx", "sample_company_form.hwpx", "sample_company_form.xlsx",
                     "sample_company_form.pdf", "sample_static_form.pdf", "sample_company_form.pptx"):
            path = ROOT / "samples" / name
            jobs.append({"id": "synthetic:" + name, "path": str(path.resolve()), "format": path.suffix.lstrip("."),
                         "group": "자체 테스트", "organization": None, "source_url": None, "source_kind": "synthetic_form",
                         "provenance": "synthetic_fixture", "usage": "blank_form"})
    return jobs


def _empty_result(job) -> dict:
    return {key: job.get(key) for key in ("id", "path", "format", "group", "organization", "source_url", "source_kind", "provenance", "usage", "use_classification", "company_internal")} | {
        "status": "pending", "checks": {}, "issues": [], "native_layout_review": "pending"}


def _prefix_pdf(source, temporary_directory, max_pages):
    reader = PdfReader(source)
    count = len(reader.pages)
    if count <= max_pages:
        return source, {"kind": "full", "checked_pages": count, "total_pages": count}
    writer = PdfWriter()
    for page in reader.pages[:max_pages]:
        writer.add_page(page)
    target = Path(temporary_directory) / "sampled.pdf"
    with target.open("wb") as stream:
        writer.write(stream)
    return target, {"kind": "first_pages", "checked_pages": max_pages, "total_pages": count}


def _manual_static_fields(profile):
    # samples/generate.py에서 만든 본문 빈 사각형의 안쪽 좌표임.
    return [{"id": "compatibility-demo-body", "label": "본문", "value_key": "본문",
             "kind": "pdf_overlay", "required": False, "page": 1, "x": 55.0, "y": 307.0,
             "width": 470.0, "height": 150.0, "font_size": 10.0}]


def check_file(job: dict, *, max_pdf_pages=3) -> dict:
    """원본을 변경하지 않는 파일별 검사. 완성 보고서는 채우지 않음."""
    started = time.perf_counter()
    result = _empty_result(job)
    source = Path(job["path"])
    if job.get("catalog_download_status") == "blocked_drm":
        result["checks"]["access"] = {"status": "pending", "reason": "drm_protected"}
        result["issues"].append("DRM 보호 원본임. 권한 있는 비보호 양식 사본이 필요하며 DRM 해제는 수행하지 않음")
        return result
    if not source.is_file():
        result["checks"]["available"] = {"status": "pending"}
        result["issues"].append(job.get("download_error") or "로컬 원본이 없어 검증 대기임")
        return result
    original_hash = file_sha256(source)
    result.update({"source_sha256": original_hash, "size_bytes": source.stat().st_size})
    result["checks"]["available"] = {"status": "passed"}
    expected_hash = job.get("expected_sha256")
    if expected_hash and expected_hash != original_hash:
        result["checks"]["provenance"] = {"status": "failed"}
        result["issues"].append("카탈로그 검증 원본 해시와 로컬 파일 해시가 다름")
        result["status"] = "failed"
        return result
    result["checks"]["provenance"] = {"status": "passed" if expected_hash else "pending"}
    try:
        with tempfile.TemporaryDirectory(prefix="report-compat-") as temporary:
            parse_source = source
            scope = {"kind": "full"}
            if job["format"] == "pdf" and job["usage"] == "layout_reference" and max_pdf_pages:
                parse_source, scope = _prefix_pdf(source, temporary, max_pdf_pages)
            try:
                parsed = parse_file(parse_source)
                result["checks"]["parse"] = {"status": "passed", "scope": scope,
                    "text_characters": len(parsed["본문"]), "tables": len(parsed["표 목록"]),
                    "location_blocks": len(parsed["페이지/시트 정보"])}
            except Exception as exc:
                code = exc.code if isinstance(exc, ParseError) else None
                pending = code in {"uncalculated_formula", "unsupported_format", "too_large", "empty_document", "ocr_required"}
                result["checks"]["parse"] = {"status": "pending" if pending else "failed", "scope": scope,
                                             "error": str(exc), "reason": code}
                result["issues"].append(f"본문 추출 {'대기' if pending else '실패'}: {exc}")
            profile_source = parse_source if job["usage"] == "layout_reference" else source
            registered = load_form_profile(source) if job["provenance"] in {"official_catalog", "official_government", "official_tech", "official_international", "official_ra", "official_business", "official_office"} else None
            profile = registered or analyze_template(profile_source)
            if registered:
                result["registered_profile"] = registered["entry_id"]
                result["native_layout_review"] = registered["verification"]["native_output"]
            result["checks"]["profile"] = {"status": "passed" if profile.get("supported") else "pending",
                "scope": scope, "fields": len(profile.get("fields", [])), "supported": profile.get("supported", False),
                "render_mode": profile.get("render_mode"), "warnings": profile.get("warnings", [])}
            if job["usage"] in {"layout_reference", "read_only_form"}:
                result["checks"]["fill"] = {"status": "not_applicable", "reason": "완성 문서 또는 변경 제한 원본은 읽기·구조 확인만 수행하며 본문을 덮어쓰지 않음"}
            elif (not registered and profile.get("render_mode") == "overlay"
                  and job["provenance"] in {"official_catalog", "official_government", "official_tech", "official_international", "official_ra", "official_business", "official_office"}):
                result["checks"]["fill"] = {"status": "pending", "reason": "자동 제안한 PDF 영역은 원본 배치를 확인한 후 등록하거나 사용자가 승인해야 함"}
            elif profile.get("supported"):
                values = registered["demo_values"] if registered else CHECK_VALUES
                mapping = {field["id"]: field["value_key"] for field in profile["fields"] if field["value_key"] in values}
                if job["id"] == "synthetic:sample_static_form.pdf":
                    profile = analyze_template(source, _manual_static_fields(profile))
                    mapping = {field["id"]: field["value_key"] for field in profile["fields"]}
                required_unmapped = [field["label"] for field in profile["fields"] if field.get("required") and field["id"] not in mapping]
                if required_unmapped or not mapping:
                    result["checks"]["fill"] = {"status": "pending", "reason": "업무별 수동 항목 매핑이 필요함", "unmapped_required": required_unmapped}
                else:
                    output = Path(temporary) / ("filled." + job["format"])
                    try:
                        fill_compatible_template(source, values, output, mapping, profile=profile)
                        expected_values = [values[key] for key in set(mapping.values())]
                        verification = _verify_output(output, job["format"], expected_values)
                        from agent.output_check import verify_output
                        independent = verify_output(source, output, values, profile=profile, mapping=mapping)
                        result["checks"]["fill"] = {"status": "passed", "mapped_fields": len(mapping),
                            "output_bytes": output.stat().st_size, "output_validation": verification, "independent_verification": independent,
                            "test_data_only": True, "mapping_method": "registered_verified_positions" if registered else "automatic_or_fixture_mapping"}
                    except Exception as exc:
                        result["checks"]["fill"] = {"status": "failed", "error": str(exc)}
                        result["issues"].append(f"양식 채우기 실패: {exc}")
            else:
                result["checks"]["fill"] = {"status": "pending", "reason": "HWP를 HWPX로 변환한 후 기입 위치 확인이 필요함" if job["format"] == "hwp" else "양식 입력 위치를 수동으로 지정해야 함"}
    except Exception as exc:
        result["issues"].append(str(exc))
        result["checks"]["execution"] = {"status": "failed", "error": str(exc)}
    unchanged = file_sha256(source) == original_hash
    result["checks"]["source_unmodified"] = {"status": "passed" if unchanged else "failed"}
    statuses = [check["status"] for name, check in result["checks"].items() if name != "provenance"]
    result["status"] = "failed" if "failed" in statuses else ("pending" if "pending" in statuses else "passed")
    result["elapsed_seconds"] = round(time.perf_counter() - started, 3)
    return result


def _verify_output(path, kind, expected_values=()):
    if kind == "pdf":
        reader = PdfReader(path)
        if not reader.pages:
            raise ValueError("PDF 출력 페이지가 없음")
        fields = reader.get_fields() or {}
        text = "\n".join(str(field.get("/V", "")) for field in fields.values())
        if not fields:
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        verification = {"pages": len(reader.pages), "form_fields": len(fields), "method": "PDF 구조·입력값 재열기"}
    elif kind == "xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            text = "\n".join(str(value) for sheet in workbook for row in sheet.iter_rows(values_only=True)
                             for value in row if value is not None)
            verification = {"sheets": len(workbook.sheetnames), "method": "XLSX 구조·입력값 재열기"}
        finally:
            workbook.close()
    else:
        text = parse_file(path)["본문"]
        if not text.strip():
            raise ValueError("출력 문서의 본문이 없음")
        verification = {"text_characters": len(text), "method": "DOCX/HWPX/PPTX 구조·입력값 재열기"}
    normalized = " ".join(text.split())
    if any(" ".join(value.split()) not in normalized for value in expected_values):
        raise ValueError("채운 입력값이 출력 문서에서 확인되지 않음")
    return verification | {"confirmed_values": len(expected_values)}


def check_isolated(job: dict, *, timeout=45, max_pdf_pages=3) -> dict:
    """파일별 전용 Python 프로세스만 종료할 수 있는 시간 제한 검사."""
    source = Path(job["path"])
    before_hash = file_sha256(source) if source.is_file() else None
    with tempfile.TemporaryDirectory(prefix="compat-worker-") as temporary:
        input_path, output_path = Path(temporary) / "job.json", Path(temporary) / "result.json"
        input_path.write_text(json.dumps(job, ensure_ascii=False), encoding="utf-8")
        command = [sys.executable, "-m", "evals.compatibility", "--worker", str(input_path),
                   "--worker-output", str(output_path), "--pdf-pages", str(max_pdf_pages)]
        try:
            completed = subprocess.run(command, cwd=ROOT, capture_output=True, timeout=timeout, check=False)
            if completed.returncode != 0 or not output_path.is_file():
                result = _empty_result(job)
                result["status"] = "failed"
                result["issues"].append("독립 파일 검사 프로세스가 실패함")
                result["checks"]["execution"] = {"status": "failed", "returncode": completed.returncode}
                _parent_hash_check(result, source, before_hash)
                return result
            return json.loads(output_path.read_text(encoding="utf-8"))
        except subprocess.TimeoutExpired:
            result = _empty_result(job)
            result["issues"].append(f"파일별 {timeout}초 검사 제한에 도달함. 별도 재검증이 필요함")
            result["checks"]["execution"] = {"status": "pending", "reason": "timeout"}
            _parent_hash_check(result, source, before_hash)
            return result


def _parent_hash_check(result, source, before_hash):
    if before_hash is not None:
        unchanged = source.is_file() and file_sha256(source) == before_hash
        result["source_sha256"] = before_hash
        result["checks"]["available"] = {"status": "passed"}
        result["checks"]["source_unmodified"] = {"status": "passed" if unchanged else "failed"}
        if not unchanged:
            result["status"] = "failed"
            result["issues"].append("검사 전후 원본 해시가 달라짐")


def summarize(results: list[dict]) -> dict:
    statuses = Counter(row["status"] for row in results)
    official = [row for row in results if row.get("provenance") == "official_catalog"]
    checked = [row for row in results if row.get("checks", {}).get("available", {}).get("status") == "passed"]
    parsed = [row for row in checked if row.get("checks", {}).get("parse", {}).get("status") == "passed"]
    filled = [row for row in checked if row.get("checks", {}).get("fill", {}).get("status") == "passed"]
    unique_hashes = {row['source_sha256'] for row in checked if row.get('source_sha256')}
    return {"checked_count": len(checked), "unique_source_sha256_count": len(unique_hashes),
            "duplicate_source_job_count": sum(bool(row.get('source_sha256')) for row in checked) - len(unique_hashes),
            "registered_count": len(results), "passed_count": statuses["passed"],
            "pending_count": statuses["pending"], "failed_count": statuses["failed"], "parse_passed_count": len(parsed),
            "full_parse_count": sum(row["checks"]["parse"].get("scope", {}).get("kind") == "full" for row in parsed),
            "sampled_parse_count": sum(row["checks"]["parse"].get("scope", {}).get("kind") == "first_pages" for row in parsed),
            "fill_passed_count": len(filled), "official_checked_count": sum(row.get("provenance") == "official_catalog" for row in checked),
            "official_group_count": len({row["group"] for row in checked if row.get("provenance") == "official_catalog" and row.get("group")}),
            "official_registered_groups": len({row["group"] for row in official if row.get("group")}),
            "official_checked_organizations": len({row["organization"] for row in checked if row.get("provenance") == "official_catalog" and row.get("organization")}),
            "synthetic_checked_count": sum(row.get("provenance") == "synthetic_fixture" for row in checked),
            "government_checked_count": sum(row.get("provenance") == "official_government" for row in checked),
            "government_agency_count": len({row["organization"] for row in checked if row.get("provenance") == "official_government"}),
            "tech_checked_count": sum(row.get("provenance") == "official_tech" for row in checked),
            "tech_company_count": len({row["organization"] for row in checked if row.get("provenance") == "official_tech" and row.get("organization")}),
            "international_checked_count": sum(row.get("provenance") == "official_international" for row in checked),
            "international_company_count": len({row['organization'] for row in checked if row.get('provenance') == 'official_international' and row.get('organization')}),
            "ra_checked_count": sum(row.get('provenance') == 'official_ra' for row in checked),
            "ra_profile_fill_count": sum(row.get('provenance') == 'official_ra' and bool(row.get('registered_profile')) for row in filled),
            "business_checked_count": sum(row.get('provenance') == 'official_business' for row in checked),
            "business_profile_fill_count": sum(row.get('provenance') == 'official_business' and bool(row.get('registered_profile')) for row in filled),
            "office_checked_count": sum(row.get('provenance') == 'official_office' for row in checked),
            "office_profile_fill_count": sum(row.get('provenance') == 'official_office' and bool(row.get('registered_profile')) for row in filled),
            "hwp_parse_passed_count": sum(row["format"] == "hwp" for row in parsed),
            "hwp_fill_pending_count": sum(row["format"] == "hwp" and row.get("checks", {}).get("fill", {}).get("status") == "pending" for row in checked),
            "registered_profile_fill_count": sum(bool(row.get("registered_profile")) for row in filled),
            "native_layout_verified_count": sum(str(row.get("native_layout_review", "")).startswith("verified_") for row in checked),
            "unique_native_layout_verified_count": len({row['source_sha256'] for row in checked if row.get('source_sha256') and str(row.get('native_layout_review', '')).startswith('verified_')}),
            "by_format": dict(Counter(row["format"] for row in checked)), "by_provenance": dict(Counter(row["provenance"] for row in checked)),
            "native_layout_review": "별도 검증 대기. 자동 구조 검사 통과는 실제 문서 프로그램의 모든 배치 통과를 뜻하지 않음"}


def run_corpus(jobs: list[dict], *, workers=4, timeout=45, max_pdf_pages=3) -> dict:
    if not 1 <= workers <= 4 or timeout <= 0 or max_pdf_pages < 0:
        raise ValueError("workers는 1~4, timeout은 양수, pdf-pages는 0 이상이어야 함")
    with ThreadPoolExecutor(max_workers=workers) as executor:
        results = list(executor.map(lambda job: check_isolated(job, timeout=timeout, max_pdf_pages=max_pdf_pages), jobs))
    return {"generated_at": datetime.now(timezone.utc).isoformat(),
            "method": "독립 프로세스별 자동 파싱·프로파일·빈 양식 치환·원본 해시 검증",
            "pdf_policy": {"reference_max_pages": max_pdf_pages, "zero_means_full": True},
            "summary": summarize(results), "results": results,
            "limitations": ["자체 샘플은 실제 회사 양식 검증 건수에 포함하지 않음",
                            "완성 공개 보고서는 빈 양식으로 간주하지 않으며 본문을 덮어쓰지 않음",
                            "앞쪽 페이지만 추출한 PDF는 전체 문서 파싱 검증과 별도로 집계함",
                            "회사 목록 등록 수와 실제 양식 파일 검증 수는 별개임"]}


def discovery_statistics(directory: Path) -> dict:
    """탐색 후보 다운로드를 채움 검증 성과와 구분하고 SHA 합집합으로 집계함."""
    runs, hashes, agencies, discovered = [], set(), set(), set()
    for manifest in sorted((Path(directory) / "government").glob("discovery*/discovery.json")):
        data = json.loads(manifest.read_text(encoding="utf-8"))
        downloaded = 0
        for entry in data.get("documents", []):
            discovered.add(entry.get("url"))
            if entry.get("agency"):
                agencies.add(entry["agency"])
            if entry.get("download_status") == "verified" and entry.get("sha256") and entry.get("filename") and (manifest.parent / entry["filename"]).is_file():
                hashes.add(entry["sha256"])
                downloaded += 1
        runs.append({"manifest": str(manifest.resolve()), "scope": data.get("scope"),
                     "discovered": len(data.get("documents", [])), "local_downloaded": downloaded,
                     "pages_read": len(data.get("pages", []))})
    return {"runs": runs, "discovered_unique_urls": len(discovered - {None}), "downloaded_unique_sha256": len(hashes),
            "agency_count": len(agencies), "fill_verified": 0, "classification": "unclassified",
            "note": "탐색·다운로드 후보이며 빈 양식 분류·매핑·실제 채움 검증 건수로 계산하지 않음"}


def main(argv=None):
    parser = argparse.ArgumentParser(description="공개 원본·자체 양식 호환 자동 검사")
    parser.add_argument("--directory", type=Path, default=ROOT / "data/public_templates")
    parser.add_argument("--output", type=Path, default=ROOT / "outputs/compatibility_report.json")
    parser.add_argument("--group")
    parser.add_argument("--provenance", choices=sorted(PROVENANCES), help="선택한 출처의 양식만 수집·검증함")
    parser.add_argument("--workers", type=int, default=4)
    parser.add_argument("--timeout", type=float, default=45)
    parser.add_argument("--pdf-pages", type=int, default=3)
    parser.add_argument("--download", action="store_true")
    parser.add_argument("--no-samples", action="store_true")
    parser.add_argument("--no-government", action="store_true")
    parser.add_argument("--no-tech", action="store_true")
    parser.add_argument("--no-international", action="store_true")
    parser.add_argument("--no-ra", action="store_true")
    parser.add_argument("--no-business", action="store_true")
    parser.add_argument("--no-office", action="store_true")
    parser.add_argument("--worker", type=Path, help=argparse.SUPPRESS)
    parser.add_argument("--worker-output", type=Path, help=argparse.SUPPRESS)
    args = parser.parse_args(argv)
    if args.worker:
        job = json.loads(args.worker.read_text(encoding="utf-8"))
        result = check_file(job, max_pdf_pages=args.pdf_pages)
        args.worker_output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
        return 0
    jobs = collect_jobs(args.directory, group=args.group, provenance=args.provenance, include_samples=not args.no_samples,
                        include_government=not args.no_government, include_tech=not args.no_tech,
                        include_international=not args.no_international, include_ra=not args.no_ra,
                        include_business=not args.no_business, include_office=not args.no_office, download=args.download)
    report = run_corpus(jobs, workers=args.workers, timeout=args.timeout, max_pdf_pages=args.pdf_pages)
    report["scope"] = {"provenance": args.provenance or "all", "group": args.group,
                       "downloads_requested": args.download}
    report["government_discovery"] = discovery_statistics(args.directory) if not args.group and not args.no_government and args.provenance in {None, "official_government"} else None
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report["summary"], ensure_ascii=False, indent=2))
    print(f"검증 기록: {args.output.resolve()}")
    return 1 if report["summary"]["failed_count"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
