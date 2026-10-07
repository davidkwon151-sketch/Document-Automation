"""네이티브 앱·네트워크 없이 렌더러의 신뢰 경계와 원자적 출력 계약을 검증함."""

import hashlib
import json
import subprocess
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest
from docx import Document
from openpyxl import Workbook
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DictionaryObject, NameObject, TextStringObject

from parsers import native


def source_file(tmp_path, suffix=".docx"):
    path = tmp_path / ("비공개 원자료" + suffix)
    if suffix == ".docx":
        document = Document()
        document.add_paragraph("한글 원문 150 mg")
        document.save(path)
    elif suffix == ".xlsx":
        workbook = Workbook()
        workbook.active.title = "업무 자료"
        workbook.active["A1"] = 1234
        workbook.active["A2"] = "=A1*2"
        workbook.save(path)
    else:
        package = {"[Content_Types].xml": b'<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Override PartName="/ppt/presentation.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"/></Types>', "_rels/.rels": b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="ppt/presentation.xml"/></Relationships>', "ppt/presentation.xml": b'<p:presentation xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"/>'}
        with zipfile.ZipFile(path, "w") as archive:
            for name, value in package.items():
                archive.writestr(name, value)
    return path


def pdf(path, pages=2):
    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=300, height=400)
    writer.write(path)


def fake_engine(copied, output, *, cell_requests):
    assert copied.name.startswith("input") and copied != output
    pdf(output)
    evidence = {}
    for request in cell_requests:
        evidence[request["id"]] = {**request, "text": "1,234", "value2": 1234, "number_format": "#,##0", "row_hidden": False, "column_hidden": False, "sheet_hidden": False, "print_area": "", "within_print_area": None, "has_formula": False, "formula": "1234", "merged_range": "$A$1"}
    return {"status": "rendered", "formula_recalculated": copied.suffix == ".xlsx", "cell_evidence": evidence}


def rewrite(path, mutate):
    with zipfile.ZipFile(path) as package:
        parts = {name: package.read(name) for name in package.namelist()}
    mutate(parts)
    with zipfile.ZipFile(path, "w") as package:
        for name, value in parts.items():
            package.writestr(name, value)


@pytest.mark.parametrize("suffix", [".docx", ".xlsx", ".pptx"])
def test_render_atomic_pdf_from_private_copy_and_bind_source(tmp_path, suffix):
    source = source_file(tmp_path, suffix)
    original = source.read_bytes()
    output = tmp_path / "결과" / "미리보기.pdf"
    result = native.render_native_pdf(source, output, engine=fake_engine)
    assert result["status"] == "rendered" and result["engine"] == "callable"
    assert result["source_unchanged"] and result["source_sha256"] == hashlib.sha256(original).hexdigest()
    assert result["formula_recalculated"] == (suffix == ".xlsx")
    assert result["page_count"] == len(PdfReader(output).pages) == 2
    assert result["output_sha256"] == hashlib.sha256(output.read_bytes()).hexdigest()
    assert source.read_bytes() == original and not list(output.parent.glob(".native-*"))


def test_xlsx_requests_exact_sheet_and_private_display_metadata(tmp_path):
    source = source_file(tmp_path, ".xlsx")
    result = native.render_native_pdf(source, tmp_path / "preview.pdf", engine=fake_engine, cell_requests=[{"id": "amount", "part": "xl/worksheets/sheet1.xml", "coordinate": "A1"}])
    evidence = result["cell_evidence"]["amount"]
    assert evidence["sheet_name"] == "업무 자료" and evidence["text"] == "1,234"
    assert evidence["value2"] == 1234 and evidence["within_print_area"] is None


@pytest.mark.parametrize("requests", [[{"id": "x", "part": "bad", "coordinate": "A1"}], [{"id": "x", "part": "xl/worksheets/sheet1.xml", "coordinate": "XFE1"}], [{"id": "x", "part": "xl/worksheets/sheet1.xml", "coordinate": "A1048577"}], [{"id": "x", "part": "xl/worksheets/sheet1.xml", "coordinate": "A1", "sheet_name": "injected"}], [{"id": "x", "part": "xl/worksheets/sheet1.xml", "coordinate": "A1"}] * 2])
def test_invalid_cell_request_never_calls_engine(tmp_path, requests):
    called = []
    with pytest.raises(ValueError):
        native.render_native_pdf(source_file(tmp_path, ".xlsx"), tmp_path / "preview.pdf", engine=lambda *args, **kwargs: called.append(1), cell_requests=requests)
    assert not called


@pytest.mark.parametrize("mode", ["missing", "invalid", "failure", "timeout", "unavailable", "blocked", "mutated_copy", "mutated_source"])
def test_failure_preserves_existing_output_and_never_claims_success(tmp_path, mode):
    source = source_file(tmp_path)
    output = tmp_path / "preview.pdf"
    output.write_bytes(b"existing output")
    def engine(copied, candidate, **kwargs):
        if mode == "failure":
            raise RuntimeError("private secret must not appear in result")
        if mode == "timeout":
            raise TimeoutError("private secret")
        if mode == "unavailable":
            return {"status": "unavailable", "reason": "private secret"}
        if mode == "blocked":
            return {"status": "blocked"}
        if mode == "mutated_copy":
            copied.write_bytes(b"changed copy")
        if mode == "mutated_source":
            source.write_bytes(b"changed original")
        if mode != "missing":
            candidate.write_bytes(b"not PDF")
    if mode in {"failure", "timeout", "unavailable"}:
        result = native.render_native_pdf(source, output, engine=engine)
        assert result["status"] == "unavailable" and "private secret" not in json.dumps(result)
    else:
        with pytest.raises(ValueError):
            native.render_native_pdf(source, output, engine=engine)
    assert output.read_bytes() == b"existing output"


@pytest.mark.parametrize("mutation", ["macro", "external", "signature", "word_protection", "xml_entity", "embedded", "dde"])
def test_unsafe_docx_rejected_before_any_engine(tmp_path, mutation):
    source = source_file(tmp_path)
    def mutate(parts):
        if mutation == "macro": parts["word/vbaProject.bin"] = b"macro"
        elif mutation == "external": parts["word/_rels/document.xml.rels"] = b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="r99" Type="https://schemas/image" TargetMode="External" Target="https://example.test/private"/></Relationships>'
        elif mutation == "signature": parts["_xmlsignatures/sig1.xml"] = b"<signature/>"
        elif mutation == "word_protection": parts["word/settings.xml"] = b'<w:settings xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:documentProtection w:enforcement="1"/></w:settings>'
        elif mutation == "xml_entity": parts["word/document.xml"] = b'<!DOCTYPE document [<!ENTITY data SYSTEM "file:///private">]><document/>'
        elif mutation == "embedded": parts["word/embeddings/oleObject1.bin"] = b"OLE"
        else: parts["word/document.xml"] = b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:instrText>DDEAUTO server topic</w:instrText></w:document>'
    rewrite(source, mutate)
    called = []
    with pytest.raises(ValueError):
        native.render_native_pdf(source, tmp_path / "preview.pdf", engine=lambda *args, **kwargs: called.append(1))
    assert not called


def test_pdf_copy_is_identical_and_signed_encrypted_or_active_pdf_blocked(tmp_path):
    source = tmp_path / "source.pdf"
    pdf(source)
    original = source.read_bytes()
    result = native.render_native_pdf(source, tmp_path / "copy.pdf")
    assert result["engine"] == "pdf_copy" and (tmp_path / "copy.pdf").read_bytes() == original
    for mode in ("encrypted", "signed", "active"):
        writer = PdfWriter()
        writer.add_blank_page(300, 400)
        if mode == "encrypted": writer.encrypt("password")
        if mode == "signed": writer.root_object[NameObject("/Perms")] = DictionaryObject({NameObject("/DocMDP"): DictionaryObject()})
        if mode == "active": writer.root_object[NameObject("/OpenAction")] = DictionaryObject({NameObject("/S"): NameObject("/JavaScript"), NameObject("/JS"): TextStringObject("app.alert('private')")})
        writer.write(source)
        with pytest.raises(ValueError): native.render_native_pdf(source, tmp_path / "copy.pdf")
        assert (tmp_path / "copy.pdf").read_bytes() == original


@pytest.mark.parametrize("suffix", [".hwp", ".hwpx", ".doc", ".unknown"])
def test_unsupported_formats_never_start_native_app(tmp_path, suffix):
    source = tmp_path / ("source" + suffix)
    source.write_bytes(b"original")
    result = native.render_native_pdf(source, tmp_path / "preview.pdf", engine=lambda *args, **kwargs: pytest.fail("should not execute"))
    assert result["status"] == "unavailable" and not (tmp_path / "preview.pdf").exists()


def test_no_engine_is_unavailable_without_output_and_self_overwrite_blocked(tmp_path, monkeypatch):
    source = source_file(tmp_path)
    monkeypatch.setattr(native, "_powershell", lambda: None)
    assert native.render_native_pdf(source, tmp_path / "preview.pdf")["status"] == "unavailable"
    with pytest.raises(ValueError): native.render_native_pdf(source, source, engine=fake_engine)


def test_incomplete_or_wrong_cell_evidence_rejected(tmp_path):
    source = source_file(tmp_path, ".xlsx")
    request = [{"id": "amount", "part": "xl/worksheets/sheet1.xml", "coordinate": "A1"}]
    def engine(copied, output, **kwargs):
        pdf(output)
        return {"status": "rendered", "formula_recalculated": True, "cell_evidence": {"amount": {"id": "amount", "part": request[0]["part"], "coordinate": "A2", "text": "1234"}}}
    with pytest.raises(ValueError, match="위치"):
        native.render_native_pdf(source, tmp_path / "preview.pdf", engine=engine, cell_requests=request)


def test_owned_office_worker_checks_ownership_before_mutation_and_never_saves():
    worker = native.OFFICE_WORKER
    assert worker.index("$ownerPid -in $priorIds") < worker.index("$app.AutomationSecurity=3") < worker.index("$app.Documents.Open")
    assert "if($owned){" in worker and "if($count -ne 0){return}" in worker
    assert "Workbooks.Open($InputPath,0,$true" in worker and "Presentations.Open($InputPath,-1,0,0)" in worker
    assert "CalculateFullRebuild()" in worker and "$document.SaveAs($OutputPath,32)" in worker
    assert "while([int]$app.CalculationState -eq 1)" in worker
    assert "$result.formula_recalculated=($result.calculation_state -eq 0)" in worker
    assert "sheet_hidden=([int]$sheet.Visible -ne -1)" in worker
    assert "$fresh.Count -ne 1" in worker
    assert "PageSetup.PrintArea=" not in worker and "UsedRange.Columns.AutoFit" not in worker
    assert "process.StartTime.ToUniversalTime().Ticks -eq" in native.CLEANUP_WORKER
    assert worker.index("$created.CommandLine -notmatch") < worker.index("$owned=$true")


def test_excel_pending_keeps_preview_without_claiming_recalculation_complete(tmp_path):
    def engine(copied, target, **kwargs):
        pdf(target)
        return {'status': 'rendered', 'formula_recalculated': False, 'calculation_state': 2}
    source = source_file(tmp_path, '.xlsx')
    result = native.render_native_pdf(source, tmp_path / 'preview.pdf', engine=engine)
    assert result['status'] == 'rendered' and result['calculation_state'] == 2
    assert result['formula_recalculated'] is False and result['page_count'] == 2
    def wrong(copied, target, **kwargs):
        data = engine(copied, target, **kwargs)
        return {**data, 'formula_recalculated': True}
    with pytest.raises(ValueError, match='계산 상태'):
        native.render_native_pdf(source, tmp_path / 'wrong.pdf', engine=wrong)


def test_subprocess_timeout_kills_only_worker_then_verified_cleanup(tmp_path, monkeypatch):
    calls = []
    class Process:
        returncode = 0
        def wait(self, timeout):
            if timeout == 1: raise subprocess.TimeoutExpired("native", timeout)
        def kill(self): calls.append("worker killed")
    monkeypatch.setattr(native.subprocess, "Popen", lambda *args, **kwargs: Process())
    monkeypatch.setattr(native.subprocess, "run", lambda command, **kwargs: calls.append(command))
    monkeypatch.setattr(native, "_powershell", lambda: "powershell.exe")
    owner = tmp_path / "owner.json"
    owner.write_text('{"pid":123,"start_ticks":456,"process_name":"WINWORD"}')
    with pytest.raises(TimeoutError): native._run_worker(["worker"], 1, owner, tmp_path)
    assert calls[0] == "worker killed" and "-OwnerFile" in calls[1]
    assert "taskkill" not in " ".join(calls[1])


@pytest.mark.parametrize("mutation", ["sheet_protection", "workbook_protection", "webservice", "dangling_link"])
def test_protected_or_external_xlsx_never_reaches_renderer(tmp_path, mutation):
    source = source_file(tmp_path, ".xlsx")
    def mutate(parts):
        if mutation == "dangling_link":
            parts["xl/_rels/workbook.xml.rels"] = b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="r1" Type="http://schemas/worksheet" Target="../../private.xml"/></Relationships>'
            return
        part = "xl/workbook.xml" if mutation == "workbook_protection" else "xl/worksheets/sheet1.xml"
        root = ET.fromstring(parts[part])
        if mutation == "workbook_protection":
            root.find("{" + native.S + "}workbookProtection").set("lockStructure", "1")
        elif mutation == "sheet_protection": ET.SubElement(root, "{" + native.S + "}sheetProtection", sheet="1")
        else: root.find(".//{" + native.S + "}f").text = 'WEBSERVICE("https://private.test/")'
        parts[part] = ET.tostring(root)
    rewrite(source, mutate)
    with pytest.raises(ValueError):
        native.render_native_pdf(source, tmp_path / "preview.pdf", engine=lambda *args, **kwargs: pytest.fail("unsafe engine call"))


def test_missing_type_or_main_part_relationship_rejected(tmp_path):
    source = source_file(tmp_path)
    rewrite(source, lambda parts: parts.__setitem__("_rels/.rels", b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"/>'))
    with pytest.raises(ValueError, match="주 문서"):
        native.render_native_pdf(source, tmp_path / "preview.pdf", engine=fake_engine)


def test_copy_race_is_blocked_before_engine_and_empty_requests_are_noop(tmp_path, monkeypatch):
    source = source_file(tmp_path)
    assert native.render_native_pdf(source, tmp_path / "normal.pdf", engine=fake_engine, cell_requests=[])["status"] == "rendered"
    copy = native.shutil.copyfile
    def race(original, copied):
        result = copy(original, copied)
        if Path(copied).name == "input.docx": Path(copied).write_bytes(b"unvalidated changed input")
        return result
    monkeypatch.setattr(native.shutil, "copyfile", race)
    with pytest.raises(ValueError, match="실행 전에"):
        native.render_native_pdf(source, tmp_path / "bad.pdf", engine=lambda *args, **kwargs: pytest.fail("raced engine call"))
    assert not (tmp_path / "bad.pdf").exists()


@pytest.mark.parametrize("mode", ["word_protection", "signature", "macro", "encrypted_office"])
def test_safe_preview_unavailability_has_dedicated_exception_not_corruption(tmp_path, mode):
    source = source_file(tmp_path)
    def mutate(parts):
        if mode == "word_protection": parts["word/settings.xml"] = b'<w:settings xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:documentProtection w:enforcement="1" w:edit="forms"/></w:settings>'
        elif mode == "signature": parts["_xmlsignatures/sig1.xml"] = b"<signature/>"
        else: parts["word/vbaProject.bin"] = b"macro"
    if mode == "encrypted_office": source.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"EncryptedPackage")
    else: rewrite(source, mutate)
    original = source.read_bytes()
    with pytest.raises(native.NativeRenderingUnavailable):
        native.render_native_pdf(source, tmp_path / "preview.pdf", engine=lambda *args, **kwargs: pytest.fail("must not open protected source"))
    assert source.read_bytes() == original


@pytest.mark.parametrize("mode", ["bad_zip", "xml", "invalid_pdf_output", "bad_evidence"])
def test_corruption_or_renderer_evidence_error_is_not_safe_unavailability(tmp_path, mode):
    source = source_file(tmp_path)
    engine = fake_engine
    if mode == "bad_zip": source.write_bytes(b"damaged zip")
    elif mode == "xml": rewrite(source, lambda parts: parts.__setitem__("word/document.xml", b"<unclosed>"))
    elif mode == "invalid_pdf_output":
        def engine(copied, candidate, **kwargs):
            candidate.write_bytes(b"not a PDF")
    else:
        def engine(copied, candidate, **kwargs):
            pdf(candidate)
            return {"status": "rendered", "formula_recalculated": "yes"}
    with pytest.raises(ValueError) as exception:
        native.render_native_pdf(source, tmp_path / "preview.pdf", engine=engine)
    assert not isinstance(exception.value, native.NativeRenderingUnavailable)


def test_encrypted_renderer_output_is_error_but_encrypted_input_only_preview_unavailable(tmp_path):
    protected = tmp_path / "protected.pdf"
    writer = PdfWriter()
    writer.add_blank_page(300, 400)
    writer.encrypt("private")
    writer.write(protected)
    with pytest.raises(native.NativeRenderingUnavailable): native.render_native_pdf(protected, tmp_path / "copy.pdf")
    def engine(copied, candidate, **kwargs):
        candidate.write_bytes(protected.read_bytes())
    with pytest.raises(ValueError) as exception:
        native.render_native_pdf(source_file(tmp_path), tmp_path / "output.pdf", engine=engine)
    assert not isinstance(exception.value, native.NativeRenderingUnavailable)
