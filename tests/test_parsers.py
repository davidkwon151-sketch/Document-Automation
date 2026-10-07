from pathlib import Path
from zipfile import ZipFile

from docx import Document
from openpyxl import Workbook
import pytest
from reportlab.pdfgen import canvas

from parsers import ParseError, parse_file


ROOT = Path(__file__).resolve().parents[1]
SAMPLES = ROOT / "samples"


@pytest.mark.parametrize("suffix", ["pdf", "xlsx", "docx", "hwpx"])
def test_all_parsers_return_same_contract_and_korean_evidence(suffix):
    result = parse_file(SAMPLES / f"source.{suffix}")
    assert set(result) == {"파일명", "본문", "표 목록", "페이지/시트 정보"}
    assert result["파일명"] == f"source.{suffix}"
    assert "95" in result["본문"] and "4.5" in result["본문"]
    assert result["표 목록"]
    for block in result["페이지/시트 정보"]:
        assert set(block) == {"페이지", "시트", "위치", "본문", "표 목록"}
        assert block["위치"]
    if suffix in {"docx", "hwpx"}:
        assert all(block["페이지"] is None for block in result["페이지/시트 정보"])


def test_pdf_page_and_table_source_positions():
    result = parse_file(SAMPLES / "source.pdf")
    assert [block["페이지"] for block in result["페이지/시트 정보"]] == [1, 2]
    assert "4.5" in result["페이지/시트 정보"][1]["본문"]
    assert result["표 목록"][0]["행"] == [["항목", "실적"], ["교육 참석", "95명"]]
    assert result["표 목록"][0]["페이지"] == 1


def test_xlsx_retains_sheet_cell_reference_and_percentage():
    result = parse_file(SAMPLES / "source.xlsx")
    assert {block["시트"] for block in result["페이지/시트 정보"]} == {"교육 실적", "만족도"}
    percentage = next(block for block in result["페이지/시트 정보"] if "참석률" in block["본문"])
    assert "95%" in percentage["본문"]
    assert percentage["위치"] == "시트 교육 실적!A4:B4"


def test_docx_table_cells_are_addressable():
    result = parse_file(SAMPLES / "source.docx")
    location = next(block["위치"] for block in result["페이지/시트 정보"] if block["본문"] == "95명")
    assert location == "본문/표 1/행 2/열 2/문단 1"


def test_hwpx_does_not_duplicate_table_text_or_invent_page_number():
    result = parse_file(SAMPLES / "source.hwpx")
    assert result["본문"].count("95명") == 2  # 본문 한 문장과 표의 해당 셀 각각 1회
    assert result["표 목록"][0]["행"] == [["항목", "대상", "참석"], ["교육", "100명", "95명"]]
    assert all("Contents/section0.xml:" in block["위치"] for block in result["페이지/시트 정보"])


@pytest.mark.parametrize("suffix", ["pdf", "xlsx", "docx", "hwpx"])
def test_corrupted_file_is_reported_consistently(tmp_path, suffix):
    path = tmp_path / f"broken.{suffix}"
    path.write_bytes(b"not a document")
    with pytest.raises(ParseError) as error:
        parse_file(path)
    assert error.value.code == "invalid_file"


def test_missing_empty_and_unsupported_files_are_distinguished(tmp_path):
    with pytest.raises(ParseError, match="찾을 수 없음") as missing:
        parse_file(tmp_path / "missing.pdf")
    assert missing.value.code == "missing_file"
    empty = tmp_path / "empty.docx"
    empty.touch()
    with pytest.raises(ParseError) as error:
        parse_file(empty)
    assert error.value.code == "empty_file"
    unsupported = tmp_path / "data.txt"
    unsupported.write_text("자료", encoding="utf-8")
    with pytest.raises(ParseError) as error:
        parse_file(unsupported)
    assert error.value.code == "unsupported_format"


def test_scanned_or_blank_pdf_requires_ocr(tmp_path):
    path = tmp_path / "scan.pdf"
    pdf = canvas.Canvas(str(path))
    pdf.rect(20, 20, 200, 200)
    pdf.showPage()
    pdf.save()
    with pytest.warns(UserWarning, match="OCR"), pytest.raises(ParseError) as error:
        parse_file(path)
    assert error.value.code == "ocr_required"


@pytest.mark.parametrize("suffix", ["docx", "xlsx"])
def test_empty_documents_are_reported(tmp_path, suffix):
    path = tmp_path / f"empty.{suffix}"
    if suffix == "docx":
        Document().save(path)
    else:
        Workbook().save(path)
    with pytest.raises(ParseError) as error:
        parse_file(path)
    assert error.value.code == "empty_document"


def test_uncalculated_xlsx_formula_is_not_used_as_a_fact(tmp_path):
    path = tmp_path / "formula.xlsx"
    workbook = Workbook()
    workbook.active.append(["합계", "=1+2"])
    workbook.save(path)
    with pytest.raises(ParseError) as error:
        parse_file(path)
    assert error.value.code == "uncalculated_formula"


def test_xlsx_saved_formula_value_can_be_used(tmp_path):
    path = tmp_path / "formula.xlsx"
    workbook = Workbook()
    workbook.active.append(["합계", "=1+2"])
    workbook.save(path)
    with ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    parts["xl/worksheets/sheet1.xml"] = parts["xl/worksheets/sheet1.xml"].replace(b"<v></v>", b"<v>3</v>")
    with ZipFile(path, "w") as archive:
        for name, content in parts.items():
            archive.writestr(name, content)
    assert "3" in parse_file(path)["본문"]


def test_toy_xml_zip_is_not_accepted_as_hwpx(tmp_path):
    path = tmp_path / "fake.hwpx"
    with ZipFile(path, "w") as archive:
        archive.writestr("Contents/section0.xml", "<sec><p>자료</p></sec>")
    with pytest.raises(ParseError, match="정상적인 HWPX"):
        parse_file(path)


def test_nested_docx_table_and_merged_cell_are_extracted_once(tmp_path):
    document = Document()
    table = document.add_table(rows=1, cols=2)
    cell = table.cell(0, 0).merge(table.cell(0, 1))
    cell.text = "합계 100명"
    nested = cell.add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "세부 95명"
    path = tmp_path / "nested.docx"
    document.save(path)
    result = parse_file(path)
    assert result["본문"].count("합계 100명") == 1
    assert result["본문"].count("세부 95명") == 1
    assert len(result["표 목록"]) == 2
