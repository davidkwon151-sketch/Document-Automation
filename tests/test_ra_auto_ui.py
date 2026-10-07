"""Synthetic actual files: parser -> batch evidence -> fill -> checked download."""
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
import json

from docx import Document
import pytest
from streamlit.testing.v1 import AppTest

from app import ra_auto_ui as ui
from app import ra_workflow_ui as registered_ui
from templates.compatibility import analyze_template


class Upload:
    def __init__(self, name, data):
        self.name, self.data = name, data

    def getvalue(self):
        return self.data


def template_bytes():
    return ui.example_uploads()[0]


@pytest.fixture
def setup(tmp_path, monkeypatch):
    data = template_bytes()
    path = tmp_path / '합성시험.docx'
    path.write_bytes(data)
    profile = analyze_template(path)
    profile.update(domain='pharmaceutical_ra', ra_workflow='product_approval', citation_mode='sidecar',
                   configured=True, auto_mapping_confirmed=True)
    for field in profile['fields']:
        field['narrative_style_required'] = False
    profile_path = tmp_path / 'profile.json'
    profile_path.write_text(json.dumps(profile, ensure_ascii=False), encoding='utf-8')
    entry = {'id': 'synthetic', 'title': '합성 시험 양식', 'source_path': path.name,
             'source_sha256': sha256(data).hexdigest(), 'profile_path': profile_path.name,
             'profile_sha256': sha256(profile_path.read_bytes()).hexdigest(),
             'coverage_note': '프로젝트 제작 합성 시험이며 공식 기관 양식 아님'}
    catalog = tmp_path / 'catalog.json'
    catalog.write_text(json.dumps({'workflows': [entry]}, ensure_ascii=False), encoding='utf-8')
    monkeypatch.setattr(registered_ui, 'ROOT', tmp_path)
    monkeypatch.setattr(registered_ui, 'CATALOG', catalog)
    upload = Upload('합성원자료.csv', ('제품명,항목,값\nSyntheticDrugA,제품명,SyntheticDrugA\n'
                                   'SyntheticDrugA,성상,흰색 분말\nSyntheticDrugA,포장단위,합성 시험 포장\n').encode('utf-8-sig'))
    files = {'ra_auto_sources_mm_uploads': [upload], 'ra_auto_template_upload': Upload('내양식.docx', data)}
    monkeypatch.setattr(ui.st, 'file_uploader', lambda *a, **k: files.get(k.get('key')))
    return files, path, profile


def app():
    return AppTest.from_string('from app.ra_auto_ui import render_ra_auto\nrender_ra_auto()', default_timeout=30).run()


def clean(at):
    assert not at.exception, [item.message for item in at.exception]
    assert not at.error, [item.value for item in at.error]


def data_ready(at):
    at.text_input(key='ra_auto_product').set_value('SyntheticDrugA')
    at.text_input(key='ra_auto_variant').set_value('')
    at.run()
    clean(at)
    next(item for item in at.text_input if item.label == '성명').set_value('합성 담당자').run()
    clean(at)
    return at


def copy_generated(at):
    at.button(key='ra_auto_propose').click().run()
    clean(at)
    next(item for item in at.checkbox if '일괄 후보' in item.label).check().run()
    at.button(key='ra_auto_generate').click().run()
    clean(at)
    return at


def exported(at):
    next(item for item in at.checkbox if '현재 기입값' in item.label).check().run()
    at.button(key='ra_auto_export').click().run()
    clean(at)
    assert any(item.proto.label == '작성 문서 다운로드' for item in at.get('download_button'))
    return at


def test_actual_registered_template_csv_batch_to_real_docx_and_evidence(setup, monkeypatch):
    def forbidden(*a, **k):
        raise AssertionError('Source copy must not call a model')
    monkeypatch.setattr(ui.LLMClient, 'generate_json', forbidden)
    files, original, _ = setup
    original_data = original.read_bytes()
    at = exported(copy_generated(data_ready(app())))
    result = at.session_state['ra_auto_exports']
    document = Document(BytesIO(result['document']))
    values = [row.cells[1].text for row in document.tables[0].rows]
    assert values == ['SyntheticDrugA', '흰색 분말', '합성 시험 포장', '합성 담당자']
    assert original.read_bytes() == original_data
    evidence = json.loads(result['evidence'])
    assert not evidence['submission_ready'] and evidence['actual_model_requests'] == 0
    assert evidence['target_coverage']['filled_count'] == evidence['target_coverage']['target_count'] == 4
    assert evidence['output_verification']['docx']['status'] == 'passed'
    accepted = evidence['auto']['source_records']
    assert accepted and all(source['document_sha256'] == sha256(files['ra_auto_sources_mm_uploads'][0].data).hexdigest() for source in accepted)
    assert all(source['filename'] == '합성원자료.csv' and source['sheet'] for source in accepted)
    assert all(item.disabled for item in at.text_area if item.label.endswith('· 초안'))


def test_first_seen_uploaded_docx_requires_mapping_confirmation_and_real_fill(setup):
    at = app()
    at.radio(key='ra_auto_template_kind').set_value('내 양식 업로드').run()
    clean(at)
    assert not at.button and not any(item.proto.label == '작성 문서 다운로드' for item in at.get('download_button'))
    next(item for item in at.checkbox if '원본 입력 위치' in item.label).check().run()
    at = exported(copy_generated(data_ready(at)))
    assert Document(BytesIO(at.session_state['ra_auto_exports']['document'])).tables[0].rows[1].cells[1].text == '흰색 분말'
    next(item for item in at.checkbox if '원본 입력 위치' in item.label).uncheck().run()
    assert 'ra_auto_exports' not in at.session_state
    assert not any(item.proto.label == '작성 문서 다운로드' for item in at.get('download_button'))


def test_new_template_document_kind_change_invalidates_mapping_confirmation_and_download(setup):
    at = app()
    at.radio(key='ra_auto_template_kind').set_value('내 양식 업로드').run()
    next(item for item in at.checkbox if '원본 입력 위치' in item.label).check().run()
    at = exported(copy_generated(data_ready(at)))
    assert json.loads(at.session_state['ra_auto_exports']['evidence'])['template_profile']['document_kind'] == 'application'
    next(item for item in at.selectbox if item.label == '내 양식의 문서 종류').set_value('report').run()
    clean(at)
    assert 'ra_auto_exports' not in at.session_state and 'ra_auto_result' not in at.session_state
    assert not next(item for item in at.checkbox if '원본 입력 위치' in item.label).value
    assert not any(item.proto.label == '작성 문서 다운로드' for item in at.get('download_button'))


@pytest.mark.parametrize('change', ['product', 'mode', 'direct', 'source', 'candidate_uncheck', 'final_uncheck'])
def test_every_material_input_or_confirmation_change_removes_download(setup, change):
    files, *_ = setup
    at = exported(copy_generated(data_ready(app())))
    if change == 'product':
        at.text_input(key='ra_auto_product').set_value('OtherDrug')
    elif change == 'mode':
        at.radio(key='ra_auto_mode').set_value('실제 AI 작성')
    elif change == 'direct':
        next(item for item in at.text_input if item.label == '성명').set_value('다른 합성 담당자')
    elif change == 'source':
        files['ra_auto_sources_mm_uploads'][0].data = files['ra_auto_sources_mm_uploads'][0].data.replace('흰색'.encode(), '노란색'.encode())
    elif change == 'candidate_uncheck':
        next(item for item in at.checkbox if '일괄 후보' in item.label).uncheck()
    else:
        next(item for item in at.checkbox if '현재 기입값' in item.label).uncheck()
    at.run()
    clean(at)
    assert 'ra_auto_exports' not in at.session_state
    assert not any(item.proto.label == '작성 문서 다운로드' for item in at.get('download_button'))


def test_selected_required_field_missing_data_is_not_blank_success(setup):
    files, *_ = setup
    files['ra_auto_sources_mm_uploads'][0].data = '제품명,항목,값\nSyntheticDrugA,제품명,SyntheticDrugA\n'.encode('utf-8-sig')
    at = copy_generated(data_ready(app()))
    result = at.session_state['ra_auto_result']
    assert result['target_coverage']['target_count'] == 4
    assert set(result['target_coverage']['missing_keys']) == {'성상', '포장단위'}
    assert at.button(key='ra_auto_export').disabled
    assert not any(item.proto.label == '작성 문서 다운로드' for item in at.get('download_button'))


def test_actual_ai_quota_error_is_sanitized_without_copy_fallback(setup, monkeypatch):
    from llm.client import LLMError
    from app import ra_auto_service
    calls = []
    monkeypatch.setattr(ui, 'LLMClient', lambda: object())
    def fail(*a, **k):
        calls.append(k['mode'])
        raise LLMError('provider raw sk-secret-private', kind='quota')
    monkeypatch.setattr(ra_auto_service, 'generate_ra_auto', fail)
    at = data_ready(app())
    at.radio(key='ra_auto_mode').set_value('실제 AI 작성').run()
    at.button(key='ra_auto_generate').click().run()
    assert not at.exception
    assert calls == ['live']
    assert any('API 잔액' in item.value for item in at.error)
    assert all('sk-secret' not in item.value and 'provider raw' not in item.value for item in at.error)
    assert 'ra_auto_result' not in at.session_state and 'ra_auto_exports' not in at.session_state


def test_first_seen_template_filename_escape_rejected(setup):
    at = app()
    with pytest.raises(ValueError, match='경로'):
        ui.uploaded_template(Upload('../원본.docx', template_bytes()))


def test_model_is_never_constructed_before_explicit_generate(setup, monkeypatch):
    def forbidden(*a, **k):
        raise AssertionError('No model before user action')
    monkeypatch.setattr(ui, 'LLMClient', forbidden)
    at = data_ready(app())
    at.radio(key='ra_auto_mode').set_value('실제 AI 작성').run()
    clean(at)
    assert not at.button(key='ra_auto_generate').disabled


@pytest.mark.parametrize('change', ['confirm_uncheck', 'transcription_edit'])
def test_real_image_manual_transcription_confirmation_change_blocks_previous_output(setup, change):
    from PIL import Image
    files, *_ = setup
    image = BytesIO()
    Image.new('RGB', (100, 80), 'white').save(image, format='PNG')
    files['ra_auto_sources_mm_uploads'] = [Upload('합성사진.png', image.getvalue())]
    at = app()
    at.text_input(key='ra_auto_product').set_value('SyntheticDrugA').run()
    clean(at)
    assert not any(item.proto.label == '작성 문서 다운로드' for item in at.get('download_button'))
    next(item for item in at.text_area if '원문 전사' in item.label).set_value(
        '제품명: SyntheticDrugA\n성상: 흰색 분말\n포장단위: 합성 시험 포장').run()
    next(item for item in at.button if '전사문 등록' in item.label).click().run()
    clean(at)
    for checkbox in at.checkbox:
        if '이 원문 조각' in checkbox.label:
            checkbox.check()
    at.run()
    next(item for item in at.text_input if item.label == '성명').set_value('합성 담당자').run()
    at = exported(copy_generated(at))
    evidence = json.loads(at.session_state['ra_auto_exports']['evidence'])
    accepted = evidence['auto']['source_records']
    assert all(source['verification_receipt']['confirmed'] for source in accepted)
    if change == 'confirm_uncheck':
        next(item for item in at.checkbox if '이 원문 조각' in item.label).uncheck().run()
    else:
        next(item for item in at.text_area if '원문 전사' in item.label).set_value('제품명: SyntheticDrugA\n성상: 변경한 시험 문구').run()
    clean(at)
    assert 'ra_auto_exports' not in at.session_state and 'ra_auto_result' not in at.session_state
    assert not any(item.proto.label == '작성 문서 다운로드' for item in at.get('download_button'))


def test_mock_generation_preserves_first_ai_baseline_across_source_and_mode_changes_until_new_document(setup, monkeypatch):
    """State contract only; the mocked result is not an actual model/KPI measurement."""
    from app import ra_auto_service
    files, *_ = setup
    baseline = {'baseline_draft': {'성상': '최초 합성 초안'}, 'draft_started_at': 'synthetic-fixed-time'}
    previous = []
    def generate(*a, **kwargs):
        previous.append(deepcopy(kwargs['previous_metrics']))
        return {'draft': {}, 'mode': 'live', 'status': 'needs_information', 'actual_model_requests': 0,
                'questions': [], 'review': {'blocking': True, 'issues': []},
                'metrics': deepcopy(kwargs['previous_metrics'] or baseline)}
    monkeypatch.setattr(ra_auto_service, 'generate_ra_auto', generate)
    monkeypatch.setattr(ui, 'LLMClient', lambda: object())
    at = data_ready(app())
    at.radio(key='ra_auto_mode').set_value('실제 AI 작성').run()
    at.button(key='ra_auto_generate').click().run()
    clean(at)
    assert previous == [None]
    files['ra_auto_sources_mm_uploads'][0].data += b'\n'
    at.run()
    assert 'ra_auto_result' not in at.session_state
    assert at.session_state['ra_auto_active_metrics'] == baseline
    at.radio(key='ra_auto_mode').set_value('원문 일괄 기입').run()
    at.radio(key='ra_auto_mode').set_value('실제 AI 작성').run()
    at.button(key='ra_auto_generate').click().run()
    clean(at)
    assert previous == [None, baseline]
    at.button(key='ra_auto_new_document').click().run()
    assert 'ra_auto_active_metrics' not in at.session_state
    at.button(key='ra_auto_generate').click().run()
    clean(at)
    assert previous == [None, baseline, None]


def test_preview_only_exposes_actual_template_fields_and_preserves_internal_pipeline_values():
    at = AppTest.from_string('''
import streamlit as st
from app.ra_auto_ui import _show_result
result = {'draft': {'제목': '내부 제목', '요약': '내부 요약', '본문': '내부 본문', '제품명': '합성 제품'},
          'mode': 'live', 'review': {'blocking': False, 'issues': []}, 'actual_model_requests': 0}
profile = {'fields': [{'value_key': '제품명', 'required': True}]}
st.session_state['edited'] = _show_result(result, profile)
''').run()
    clean(at)
    assert [item.label for item in at.text_area] == ['제품명 · 초안']
    assert at.session_state['edited'] == {'제목': '내부 제목', '요약': '내부 요약', '본문': '내부 본문', '제품명': '합성 제품'}
