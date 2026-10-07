from hashlib import sha256
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.shared import Pt
from lxml import etree
import pytest

from parsers import parse_file
from samples.generate import VALUES
from templates import TemplateError, fill_template


ROOT = Path(__file__).resolve().parents[1]
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
HP = "http://www.hancom.co.kr/hwpml/2011/paragraph"


def parts(path):
    with ZipFile(path) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


@pytest.mark.parametrize("suffix", ["docx", "hwpx"])
def test_fill_preserves_original_and_replaces_split_and_table_placeholders(tmp_path, suffix):
    source = ROOT / f"templates/result_report.{suffix}"
    original_hash = sha256(source.read_bytes()).hexdigest()
    output = fill_template(source, VALUES, tmp_path / f"result.{suffix}")
    assert output.is_file()
    assert sha256(source.read_bytes()).hexdigest() == original_hash
    extracted = parse_file(output)
    assert VALUES["제목"] in extracted["본문"]
    assert VALUES["요약"] in extracted["본문"]
    assert VALUES["본문"] in extracted["본문"]
    assert "{{" not in extracted["본문"]
    assert any(VALUES["제목"] in cell for table in extracted["표 목록"] for row in table["행"] for cell in row)


def test_docx_fonts_paragraphs_tables_headers_and_zip_parts_stay_intact(tmp_path):
    source = ROOT / "templates/result_report.docx"
    output = fill_template(source, VALUES, tmp_path / "filled.docx")
    before, after = parts(source), parts(output)
    changed = {name for name in before if before[name] != after[name]}
    assert changed == {"word/document.xml", "word/header1.xml"}
    for name in changed:
        original, filled = etree.fromstring(before[name]), etree.fromstring(after[name])
        for tag in ("rPr", "pPr", "tblPr", "tblGrid", "tcPr", "sectPr"):
            old = original.xpath(f".//w:{tag}", namespaces={"w": W})
            new = filled.xpath(f".//w:{tag}", namespaces={"w": W})
            assert [etree.tostring(e) for e in old] == [etree.tostring(e) for e in new]
    document = Document(output)
    assert document.paragraphs[0].runs[0].bold
    assert document.paragraphs[0].runs[0].font.name == "맑은 고딕"
    assert VALUES["제목"] in document.sections[0].header.paragraphs[0].text
    assert document.sections[0].footer.paragraphs[0].text == "검토용 문서"
    assert after["word/document.xml"].count(b"<w:br") == 2


def test_hwpx_style_references_tables_and_untouched_package_members_stay_intact(tmp_path):
    source = ROOT / "templates/result_report.hwpx"
    output = fill_template(source, VALUES, tmp_path / "filled.hwpx")
    before, after = parts(source), parts(output)
    assert set(before) == set(after)
    assert {name for name in before if before[name] != after[name]} == {"Contents/section0.xml", "Preview/PrvText.txt"}
    old, new = etree.fromstring(before["Contents/section0.xml"]), etree.fromstring(after["Contents/section0.xml"])
    for tag in ("p", "run", "tbl", "tc", "cellAddr", "cellSpan", "cellSz", "cellMargin", "pagePr"):
        old_items = old.xpath(f".//*[local-name()='{tag}']")
        new_items = new.xpath(f".//*[local-name()='{tag}']")
        assert [dict(item.attrib) for item in old_items] == [dict(item.attrib) for item in new_items]
    assert len(new.findall(f".//{{{HP}}}lineBreak")) == 2
    assert "95명" in after["Preview/PrvText.txt"].decode("utf-8")
    with ZipFile(output) as archive:
        assert archive.infolist()[0].filename == "mimetype"
        assert archive.infolist()[0].compress_type == 0


@pytest.mark.parametrize("suffix", ["docx", "hwpx"])
def test_missing_value_fails_without_replacing_existing_output(tmp_path, suffix):
    output = tmp_path / f"result.{suffix}"
    output.write_bytes(b"existing output")
    with pytest.raises(TemplateError, match="필수 값"):
        fill_template(ROOT / f"templates/result_report.{suffix}", {"제목": "제목"}, output)
    assert output.read_bytes() == b"existing output"


@pytest.mark.parametrize("suffix", ["docx", "hwpx"])
def test_source_cannot_be_overwritten(suffix):
    source = ROOT / f"templates/result_report.{suffix}"
    with pytest.raises(TemplateError, match="덮어쓸"):
        fill_template(source, VALUES, source)


@pytest.mark.parametrize("suffix", ["docx", "hwpx"])
def test_user_xml_and_placeholder_like_text_is_literal(tmp_path, suffix):
    values = VALUES | {"제목": "<script>& {{요약}}", "요약": "□ <비용> & 절감함"}
    output = fill_template(ROOT / f"templates/result_report.{suffix}", values, tmp_path / f"filled.{suffix}")
    text = parse_file(output)["본문"]
    assert "<script>& {{요약}}" in text
    assert "□ <비용> & 절감함" in text


def test_multiple_split_placeholders_keep_surrounding_styles_and_line_order(tmp_path):
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("앞 ").italic = True
    paragraph.add_run("{{제").bold = True
    paragraph.add_run("목}} / {{요약}} / {{본문}}").font.size = Pt(14)
    paragraph.add_run(" 뒤").underline = True
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "{{본문}}"
    path = tmp_path / "source.docx"
    document.save(path)
    values = {"제목": "첫줄\n둘째줄", "요약": "요약", "본문": "□ 하나함\n○ 둘임\n- 셋함"}
    output = fill_template(path, values, tmp_path / "filled.docx")
    result = Document(output)
    assert result.paragraphs[0].text == "앞 첫줄\n둘째줄 / 요약 / □ 하나함\n○ 둘임\n- 셋함 뒤"
    assert result.paragraphs[0].runs[0].italic
    assert result.paragraphs[0].runs[-1].underline
    assert result.paragraphs[0].runs[1].bold
    assert result.tables[0].cell(0, 0).text == values["본문"]


@pytest.mark.parametrize("bad_values", [{"제목": 123}, {"제목": "\x00"}, {"제목": "", "요약": "x", "본문": "y"}])
def test_invalid_template_values_are_rejected(tmp_path, bad_values):
    with pytest.raises(TemplateError):
        fill_template(ROOT / "templates/result_report.docx", bad_values, tmp_path / "filled.docx")


@pytest.mark.parametrize("stem", ["weekly_report", "approval_request"])
@pytest.mark.parametrize("suffix", ["docx", "hwpx"])
def test_report_type_sample_templates_can_be_filled(tmp_path, stem, suffix):
    output = fill_template(ROOT / f"templates/{stem}.{suffix}", VALUES, tmp_path / f"result.{suffix}")
    assert VALUES["본문"] in parse_file(output)["본문"]
