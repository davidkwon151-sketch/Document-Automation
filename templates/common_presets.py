"""Five editable example forms for the general document workspace."""

from io import BytesIO
from pathlib import Path

from docx import Document

from llm.client import load_prompt
from templates.compatibility import analyze_template


FORMS = {
    'business_trip': ('출장 보고서', '결과보고서', 'report', '출장 목적 → 방문·협의 결과 → 후속 조치·담당자·기한'),
    'meeting_minutes': ('회의록', '결과보고서', 'minutes', '회의 목적 → 실제 논의 → 결정 사항 → 미결 사항·담당자·기한'),
    'weekly_report': ('주간보고서', '주간업무보고', 'report', '핵심 성과 → 금주 진행 → 이슈 → 차주 계획'),
    'monthly_report': ('월간보고서', '결과보고서', 'report', '월간 결론 → 주요 실적·목표 대비 → 이슈 → 다음 달 계획'),
    'approval': ('기안서', '품의서', 'report', '결재 요청 → 배경·필요성 → 실행안·예산·일정 → 기대 효과·확인 사항'),
}
FIELDS = {'제목', '요약', '본문'}
BASE = Path(__file__).resolve().parent / 'common'


def create_example(form_id: str) -> bytes:
    """Produce a project example, never an employer's official form."""
    if form_id not in FORMS:
        raise ValueError('지원하지 않는 공통 문서 종류임')
    title = FORMS[form_id][0]
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = 720000
    doc.add_heading(title, 0)
    doc.add_paragraph('프로젝트 예제 양식 · 회사 공식 양식 아님', style='Subtitle')
    table = doc.add_table(rows=2, cols=2)
    table.style = 'Table Grid'
    table.cell(0, 0).text, table.cell(0, 1).text = '제목', '{{제목}}'
    table.cell(1, 0).text, table.cell(1, 1).text = '핵심 요약', '{{요약}}'
    doc.add_heading('상세 내용', 1)
    doc.add_paragraph('{{본문}}')
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def template_path(form_id: str) -> Path:
    if form_id not in FORMS:
        raise ValueError('지원하지 않는 공통 문서 종류임')
    return BASE / f'{form_id}.docx'


def common_profile(path: str | Path, form_id: str) -> dict:
    if form_id not in FORMS:
        raise ValueError('지원하지 않는 공통 문서 종류임')
    profile = analyze_template(path)
    fields = profile['fields']
    if (len(fields) != 3 or {field['value_key'] for field in fields} != FIELDS
            or any(field['kind'] != 'placeholder' for field in fields)):
        raise ValueError('공통 작업실의 맞춤 양식은 {{제목}}, {{요약}}, {{본문}} 입력칸이 각각 한 개 필요함')
    title, report_type, document_kind, order = FORMS[form_id]
    profile.update(document_kind=document_kind, citation_mode='sidecar', common_form=form_id,
                   common_context={'title': title, 'report_type': report_type, 'section_order': order,
                                   'guidance': load_prompt('common_documents.md')})
    return profile
