"""반복 XLSX의 위치와 승인된 규칙 승계는 실제 ZIP/XML을 근거로 검사함."""

from copy import deepcopy
from hashlib import sha256
from pathlib import Path
from zipfile import ZipFile
import json

from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter, range_boundaries
import pytest

from agent.output_check import verify_output
from templates import analyze_template
from templates.compatibility import _selected_fields, _xlsx_fill
from templates.fill import TemplateError
from templates.repeat_xlsx_profile import xlsx_repeat_profile, inherit_xlsx_repeat_rules
from templates.value_rules import inspect_form_values
from templates.xlsx_validation import NS


ROOT = Path(__file__).resolve().parents[1]
PART = 'xl/worksheets/sheet1.xml'


def _prepare_fixture(source, prepared, count=3):
    """Production transformer-independent row fixture; no Excel calculation."""
    with ZipFile(source) as archive:
        entries = [(info, archive.read(info.filename)) for info in archive.infolist()]
    root = etree.fromstring(dict((info.filename, raw) for info, raw in entries)[PART])
    data = root.find(f'{{{NS}}}sheetData')
    prototype = next(row for row in data if row.get('r') == '2')
    delta = count - 1
    for row in data:
        if int(row.get('r')) > 2:
            row.set('r', str(int(row.get('r')) + delta))
            for cell in row:
                column, number, _, _ = range_boundaries(cell.get('r'))
                cell.set('r', f'{get_column_letter(column)}{number + delta}')
    for offset in range(1, count):
        clone = deepcopy(prototype)
        clone.set('r', str(2 + offset))
        for cell in clone:
            column, _, _, _ = range_boundaries(cell.get('r'))
            cell.set('r', f'{get_column_letter(column)}{2 + offset}')
        data.insert(data.index(prototype) + offset, clone)
    validation = root.find(f'./{{{NS}}}dataValidations/{{{NS}}}dataValidation')
    if validation is not None:
        validation.set('sqref', f'E2:E{1 + count}')
    with ZipFile(prepared, 'w') as archive:
        for info, raw in entries:
            archive.writestr(info, etree.tostring(root) if info.filename == PART else raw)
    plan = {'table_id': 'xlsx:' + PART, 'row': 2, 'count': count}
    binding = {'original_sha256': sha256(source.read_bytes()).hexdigest(),
               'prepared_sha256': sha256(prepared.read_bytes()).hexdigest(), 'plan': plan}
    return plan, binding


def setup(tmp_path, *, count=3, missing_choice=False, formula='"A,B,0"'):
    book = Workbook()
    sheet = book.active
    sheet.title = '신청'
    for column, label in enumerate(['품목', '첫 금액\n(원)', '둘째 금액 (원)', '합계 (원)', '결재 상태', '계산식', '고정 안내', '시작일', '종료일', '확인 전 후보'], 1):
        sheet.cell(1, column, label)
    sheet['A2'] = '{{항목}}'
    for column in (2, 3, 4, 10):
        sheet.cell(2, column).font = Font(name='맑은 고딕', size=10)
    if not missing_choice:
        sheet['E2'].font = Font(name='맑은 고딕', size=10)
    sheet['F2'] = '=SUM(B2:D2)'
    sheet['G2'] = '수정하지 않는 고정 안내'
    sheet['H3'], sheet['I3'] = '{{시작일}}', '{{종료일}}'
    native = DataValidation(type='list', formula1=formula)
    native.add('E2')
    sheet.add_data_validation(native)
    source, prepared = tmp_path/'source.xlsx', tmp_path/'prepared.xlsx'
    book.save(source)
    original = analyze_template(source)
    approved = []
    for field in original['fields']:
        coordinate = field['id'].split(':')[2]
        if coordinate in {'A2', 'B2', 'C2', 'D2', 'E2', 'H3', 'I3'}:
            field = deepcopy(field)
            if coordinate in {'B2', 'C2', 'D2'}:
                field.update(value_key={'B2': '첫 금액', 'C2': '둘째 금액', 'D2': '계 금액'}[coordinate],
                             required=True, validation={'type': 'integer', 'unit': '원', 'unit_location': 'label', 'min': 0},
                             input_mode='source_grounded', input_required=False, max_chars=16, narrative_style_required=False)
            if coordinate in {'H3', 'I3'}:
                field.update(validation={'type': 'date'}, max_chars=10, input_required=True, input_mode='user_provided', narrative_style_required=False)
            approved.append(field)
    if not any(field['id'] == f'xlsx:{PART}:B2' for field in approved):
        # Registered exact locations may include a blank whose neighbour is a
        # token, which the generic adjacency-only analyser intentionally omits.
        approved.append({'id': f'xlsx:{PART}:B2', 'kind': 'xlsx_cell', 'label': '첫 금액\n(원)', 'value_key': '첫 금액',
                         'required': True, 'input_required': False, 'input_mode': 'source_grounded', 'max_chars': 16,
                         'validation': {'type': 'integer', 'unit': '원', 'unit_location': 'label', 'min': 0}, 'narrative_style_required': False})
    original.update(fields=approved, domain='business_support', business_workflow='government_grant', document_kind='application', citation_mode='sidecar')
    original['constraints'] = {'relations': [{'kind': 'sum', 'total': '계 금액', 'parts': ['첫 금액', '둘째 금액']},
                                           {'kind': 'date_order', 'start': '시작일', 'end': '종료일'}],
                               'groups': [{'kind': 'all_or_none', 'fields': ['첫 금액', '둘째 금액', '계 금액'], 'required_fields': ['첫 금액', '둘째 금액']}]}
    plan, binding = _prepare_fixture(source, prepared, count)
    raw = xlsx_repeat_profile(prepared, plan, binding)
    return source, prepared, original, raw, plan, binding


def test_each_position_has_unique_keys_exact_excel_columns_and_preserved_tokens(tmp_path):
    source, prepared, _, profile, plan, binding = setup(tmp_path)
    before = source.read_bytes(), prepared.read_bytes()
    repeated = [field for field in profile['fields'] if field.get('repeat_info')]
    assert len(repeated) == 18
    assert len({field['value_key'] for field in repeated}) == 18
    for relative in range(1, 4):
        row_fields = [field for field in repeated if field['repeat_info']['row'] == relative]
        assert {field['repeat_info']['column'] for field in row_fields} == {1, 2, 3, 4, 5, 10}
        token = next(field for field in row_fields if field['kind'] == 'xlsx_placeholder')
        assert token['id'] == f'xlsx:{PART}:A{relative + 1}:항목' and token['required'] is True
        choice = next(field for field in row_fields if field.get('control_type'))
        assert choice['options'] == ['A', 'B', '0'] and choice['input_required'] is True
        assert choice['narrative_style_required'] is False
    assert not any(field['id'].split(':')[2].startswith(('F', 'G')) for field in repeated)
    assert profile['repeat_expansion'] == binding
    assert (source.read_bytes(), prepared.read_bytes()) == before
    assert plan['row'] == 2


@pytest.mark.parametrize('source_present', [True, False])
@pytest.mark.parametrize('missing_choice', [False, True])
def test_rules_relations_and_native_lists_are_inherited_without_mutation(tmp_path, source_present, missing_choice):
    source, prepared, original, raw, _, _ = setup(tmp_path, missing_choice=missing_choice)
    before = source.read_bytes(), prepared.read_bytes(), deepcopy(original), deepcopy(raw)
    profile = inherit_xlsx_repeat_rules(source if source_present else None, raw, original, prepared_path=prepared)
    assert (source.read_bytes(), prepared.read_bytes(), original, raw) == before
    assert profile['domain'] == 'business_support' and profile['document_kind'] == 'application'
    assert len(profile['constraints']['relations']) == 4
    for row in range(1, 4):
        members = [field for field in profile['fields'] if field.get('repeat_info', {}).get('row') == row]
        assert len(members) == 6
        for field in members:
            if field['repeat_info']['column'] in {2, 3, 4}:
                assert field['validation'] == {'type': 'integer', 'unit': '원', 'unit_location': 'label', 'min': 0}
                assert field['max_chars'] == 16 and field['required'] is True
            elif field.get('control_type'):
                assert field['options'] == ['A', 'B', '0'] and field['xlsx_list']['sqref'] == 'E2:E4'
    assert next(field for field in profile['fields'] if field['value_key'] == '시작일')['id'] == f'xlsx:{PART}:H5:시작일'
    assert next(field for field in profile['fields'] if field['value_key'] == '종료일')['input_mode'] == 'user_provided'
    assert profile['repeat_rule_verification'] == {'original_file_checked': source_present, 'prepared_file_checked': True}
    assert len(profile['repeat_unconfirmed_fields']) == 3
    assert any('매핑 확인' in message for message in profile['warnings'])


def _values(profile):
    values = {'시작일': '2026-01-01', '종료일': '2026-02-01'}
    for field in profile['fields']:
        if field.get('repeat_info'):
            column = field['repeat_info']['column']
            values[field['value_key']] = {1: '품목A', 2: '2', 3: '3', 4: '5', 5: '0', 10: ''}[column]
    return values


def test_correct_values_roundtrip_and_inherited_sum_group_choice_errors(tmp_path):
    source, prepared, original, raw, _, _ = setup(tmp_path)
    profile = inherit_xlsx_repeat_rules(source, raw, original, prepared_path=prepared)
    values = _values(profile)
    assert not inspect_form_values(values, profile)
    total = next(field['value_key'] for field in profile['fields'] if field.get('repeat_info') == {'row': 2, 'column': 4, 'source_label': '합계 (원)'})
    bad = dict(values); bad[total] = '6'
    assert any('sum' in issue['code'] for issue in inspect_form_values(bad, profile))
    bad = dict(values); bad[next(field['value_key'] for field in profile['fields'] if field.get('repeat_info', {}).get('row') == 2 and field['repeat_info']['column'] == 3)] = ''
    assert any(issue['code'] == 'form_group_required' for issue in inspect_form_values(bad, profile))
    bad = dict(values); bad[next(field['value_key'] for field in profile['fields'] if field.get('control_type'))] = '허위선택'
    assert inspect_form_values(bad, profile)
    # Isolate this adapter from the separate source-to-prepared checker/dispatcher.
    # Saved-cell verification remains independent, with unchanged native DV XML.
    standalone = deepcopy(profile); standalone.pop('repeat_expansion')
    output = tmp_path/'filled.xlsx'
    _xlsx_fill(prepared, output, _selected_fields(standalone, values, None))
    assert verify_output(prepared, output, values, profile=standalone)['status'] == 'passed'
    book = load_workbook(output, data_only=False)
    assert book.active['B3'].value == 2 and book.active['E3'].value == '0'
    assert book.active['F3'].value == '=SUM(B2:D2)'


@pytest.mark.parametrize('change', ['original_sha', 'prepared_sha', 'plan_extra', 'boolean_row', 'boolean_count', 'overflow', 'missing_sheet', 'source_row_missing'])
def test_plan_and_file_bindings_fail_closed(tmp_path, change):
    source, prepared, _, _, plan, binding = setup(tmp_path)
    plan, binding = deepcopy(plan), deepcopy(binding)
    if change == 'original_sha':
        binding['original_sha256'] = 'bogus'
    elif change == 'prepared_sha':
        binding['prepared_sha256'] = '0' * 64
    elif change == 'plan_extra':
        plan['other'] = 1
    elif change == 'boolean_row':
        plan['row'] = True
    elif change == 'boolean_count':
        plan['count'] = True
    elif change == 'overflow':
        plan['row'] = 1048576
    elif change == 'missing_sheet':
        plan['table_id'] = 'xlsx:xl/worksheets/sheet9.xml'
    else:
        plan['row'] = 10
    binding['plan'] = deepcopy(plan)
    with pytest.raises(TemplateError):
        xlsx_repeat_profile(prepared, plan, binding)


@pytest.mark.parametrize('change', ['formula', 'fixed_text', 'missing_cell', 'wrong_kind', 'native_disguise', 'native_options', 'native_definition', 'native_generated', 'token_name', 'token_optional', 'wrong_repeat_row'])
def test_source_authority_never_allows_forged_locations_or_options(tmp_path, change):
    source, prepared, original, raw, _, _ = setup(tmp_path)
    numeric = next(field for field in original['fields'] if field['id'].endswith(':B2'))
    native = next(field for field in original['fields'] if field.get('control_type'))
    token = next(field for field in original['fields'] if field['id'].endswith(':A2:항목'))
    if change in {'formula', 'fixed_text', 'missing_cell'}:
        numeric['id'] = f'xlsx:{PART}:' + {'formula': 'F2', 'fixed_text': 'G2', 'missing_cell': 'Z2'}[change]
    elif change == 'wrong_kind':
        numeric['kind'] = 'xlsx_placeholder'
    elif change == 'native_disguise':
        for key in ('control_type', 'options', 'choice_items', 'xlsx_list'):
            native.pop(key)
    elif change == 'native_options':
        native.update(options=['FAKE'], choice_items=[{'value': 'FAKE', 'label': 'FAKE'}])
    elif change == 'native_definition':
        native['xlsx_list']['formula1'] = '"FAKE"'
    elif change == 'native_generated':
        native.update(input_required=False, input_mode='source_grounded')
    elif change == 'token_name':
        token['label'] = '다른 토큰'
    elif change == 'token_optional':
        token['required'] = False
    else:
        next(field for field in raw['fields'] if field.get('repeat_info'))['repeat_info']['row'] = 99
    with pytest.raises((TemplateError, ValueError)):
        inherit_xlsx_repeat_rules(source, raw, original, prepared_path=prepared)


@pytest.mark.parametrize('source_present', [True, False])
def test_relationships_crossing_repeated_and_outside_scopes_are_not_guessed(tmp_path, source_present):
    source, prepared, original, raw, _, _ = setup(tmp_path)
    outside = next(field for field in original['fields'] if field['value_key'] == '시작일')
    outside['validation'] = {'type': 'integer', 'unit': '원', 'unit_location': 'label'}
    original['constraints'] = {'relations': [{'kind': 'less_equal', 'left': '첫 금액', 'right': '시작일'}]}
    with pytest.raises(TemplateError, match='반복행·바깥'):
        inherit_xlsx_repeat_rules(source if source_present else None, raw, original, prepared_path=prepared)


def test_same_key_multiple_columns_is_ambiguous_only_for_dependencies(tmp_path):
    source, prepared, original, raw, _, _ = setup(tmp_path)
    first = next(field for field in original['fields'] if field['value_key'] == '첫 금액')
    second = next(field for field in original['fields'] if field['value_key'] == '둘째 금액')
    second['value_key'] = first['value_key']
    original['constraints'] = {'relations': [{'kind': 'less_equal', 'left': '첫 금액', 'right': '계 금액'}]}
    with pytest.raises(TemplateError, match='모호하게'):
        inherit_xlsx_repeat_rules(source, raw, original, prepared_path=prepared)
    original['constraints'] = {}
    profile = inherit_xlsx_repeat_rules(source, raw, original, prepared_path=prepared)
    assert len({field['value_key'] for field in profile['fields'] if field.get('repeat_info')}) == 18


def test_private_values_and_local_paths_removed_while_native_option_values_remain(tmp_path):
    source, prepared, original, raw, _, _ = setup(tmp_path)
    original.update(demo_values={'name': 'SECRET PRIVATE'}, answers={'name': 'SECRET PRIVATE'}, source_path='SECRET PRIVATE',
                    qa_values={'name': 'SECRET PRIVATE'}, sources=[{'text': 'SECRET PRIVATE'}])
    original['fields'][0].update(value='SECRET PRIVATE', user_value='SECRET PRIVATE', path='SECRET PRIVATE')
    raw.update(field_values={'name': 'SECRET PRIVATE'}, expected_values={'name': 'SECRET PRIVATE'})
    profile = inherit_xlsx_repeat_rules(source, raw, original, prepared_path=prepared)
    assert 'SECRET PRIVATE' not in json.dumps(profile, ensure_ascii=False)
    assert next(field for field in profile['fields'] if field.get('control_type'))['choice_items'][2]['value'] == '0'
    assert profile['repeat_source_profile']['source_sha256'] == original['source_sha256']


def test_count_one_works_without_creating_extra_positions(tmp_path):
    source, prepared, original, raw, _, _ = setup(tmp_path, count=1)
    profile = inherit_xlsx_repeat_rules(None, raw, original, prepared_path=prepared)
    assert len([field for field in profile['fields'] if field.get('repeat_info')]) == 6
    assert next(field for field in profile['fields'] if field['value_key'] == '시작일')['id'].endswith(':H3:시작일')


def test_unresolved_native_list_is_never_inherited_as_free_text(tmp_path):
    source, prepared, original, raw, _, _ = setup(tmp_path, formula='INDIRECT("A1:A2")')
    assert any(field.get('control_type') == 'choice_unresolved' for field in raw['fields'])
    with pytest.raises(TemplateError, match='목록 확인'):
        inherit_xlsx_repeat_rules(source, raw, original, prepared_path=prepared)


def test_unregistered_outside_required_tokens_and_source_warnings_are_not_dropped(tmp_path):
    source, prepared, original, raw, _, _ = setup(tmp_path)
    original['fields'] = [field for field in original['fields'] if field['value_key'] not in {'시작일', '종료일'}]
    original['constraints'] = {}
    original['warnings'].append('원본 기관 제출 사항 확인 필요')
    profile = inherit_xlsx_repeat_rules(source, raw, original, prepared_path=prepared)
    for name in ('시작일', '종료일'):
        actual = next(field for field in profile['fields'] if field['value_key'] == name)
        assert actual['required'] is True and not actual.get('repeat_info')
    assert '원본 기관 제출 사항 확인 필요' in profile['warnings']
    assert any('바깥 필수 자리표시자' in warning for warning in profile['warnings'])


@pytest.mark.parametrize('change', ['source_sha', 'profile_sha', 'missing_prepared_path', 'source_bytes_changed'])
def test_inheritance_checks_actual_file_and_profile_hashes(tmp_path, change):
    source, prepared, original, raw, _, _ = setup(tmp_path)
    if change == 'source_sha':
        original['source_sha256'] = '0' * 64
    elif change == 'profile_sha':
        raw['source_sha256'] = '0' * 64
    elif change == 'source_bytes_changed':
        source.write_bytes(source.read_bytes() + b'changed')
    with pytest.raises(TemplateError):
        inherit_xlsx_repeat_rules(source, raw, original, prepared_path=None if change == 'missing_prepared_path' else prepared)


def test_signed_xlsx_profile_cannot_grant_writing_permissions(tmp_path):
    source, prepared, _, _, plan, binding = setup(tmp_path)
    with ZipFile(prepared, 'a') as archive:
        archive.writestr('_xmlsignatures/sig1.xml', '<signature/>')
    binding['prepared_sha256'] = sha256(prepared.read_bytes()).hexdigest()
    with pytest.raises(TemplateError, match='전자서명'):
        xlsx_repeat_profile(prepared, plan, binding)


def test_actual_original_cell_cannot_claim_a_different_parent_row(tmp_path):
    source, prepared, original, raw, _, _ = setup(tmp_path)
    with ZipFile(source) as archive:
        entries = [(info, archive.read(info.filename)) for info in archive.infolist()]
    root = etree.fromstring(dict((info.filename, data) for info, data in entries)[PART])
    prototype = root.find(f'./{{{NS}}}sheetData/{{{NS}}}row[@r="2"]')
    prototype.set('r', '99')
    with ZipFile(source, 'w') as archive:
        for info, data in entries:
            archive.writestr(info, etree.tostring(root) if info.filename == PART else data)
    actual_sha = sha256(source.read_bytes()).hexdigest()
    original['source_sha256'] = actual_sha
    raw['repeat_expansion']['original_sha256'] = actual_sha
    with pytest.raises(TemplateError, match='부모 행'):
        inherit_xlsx_repeat_rules(source, raw, original, prepared_path=prepared)


def test_real_public_kotra_blank_native_row_is_positioned_by_actual_excel_row():
    source = ROOT/'data/public_templates/business/business_kotra_usedcars_2026.xlsx'
    if not source.exists():
        pytest.skip('공식 공개 KOTRA 원본이 이 체크아웃에 없음')
    before = source.read_bytes()
    original = analyze_template(source)
    native = next(field for field in original['fields'] if field.get('control_type') == 'choice')
    _, part, coordinate = native['id'].split(':')
    _, row, _, _ = range_boundaries(coordinate)
    plan = {'table_id': 'xlsx:' + part, 'row': row, 'count': 1}
    binding = {'original_sha256': sha256(before).hexdigest(), 'prepared_sha256': sha256(before).hexdigest(), 'plan': plan}
    raw = xlsx_repeat_profile(source, plan, binding)
    approved = deepcopy(original); approved['fields'] = [native]
    profile = inherit_xlsx_repeat_rules(source, raw, approved, prepared_path=source)
    actual = next(field for field in profile['fields'] if field['id'] == native['id'])
    assert actual['options'] == native['options'] and actual['input_required'] is True
    assert actual['repeat_info']['row'] == 1
    assert source.read_bytes() == before


def test_real_public_mss_application_blank_item_row_expands_and_inherits_exact_columns(tmp_path):
    from templates.repeat_rows import prepare_repeat_template, expansion_binding, verify_prepared_template
    from templates import fill_compatible_template
    source = ROOT/'data/public_templates/business/business_mss_india_application_2026.xlsx'
    if not source.exists():
        pytest.skip('공식 공개 중기부 인도 입점사업 신청서가 이 체크아웃에 없음')
    before = source.read_bytes()
    original = analyze_template(source)
    # The sheet explicitly says to add rows for multiple items. Row6 is an
    # existing fully blank input row below the printed examples, without inventing
    # applicant/approval values or treating the example text as user information.
    original['fields'] = [field for field in original['fields'] if field['id'].split(':')[1] == PART
                          and range_boundaries(field['id'].split(':')[2])[1] == 6]
    assert len(original['fields']) == 21
    for field in original['fields']:
        column, _, _, _ = range_boundaries(field['id'].split(':')[2])
        field.update(label=f'확인된 시험 입력 {column}열', value_key=f'시험 입력 {column}', max_chars=8)
    plan = {'table_id': 'xlsx:' + PART, 'row': 6, 'count': 3}
    prepared = tmp_path/'prepared-public.xlsx'
    expansion = prepare_repeat_template(source, prepared, plan)
    binding = expansion_binding(expansion)
    raw = xlsx_repeat_profile(prepared, plan, binding)
    profile = inherit_xlsx_repeat_rules(source, raw, original, prepared_path=prepared)
    assert len(profile['fields']) == 63
    assert len({field['value_key'] for field in profile['fields']}) == 63
    assert verify_prepared_template(prepared, profile, expansion)['status'] == 'passed'
    values = {field['value_key']: f"QA{field['repeat_info']['row']}" for field in profile['fields']}
    output = fill_compatible_template(prepared, values, tmp_path/'filled-public.xlsx', profile=profile)
    assert verify_output(prepared, output, values, profile=profile)['status'] == 'passed'
    assert source.read_bytes() == before
