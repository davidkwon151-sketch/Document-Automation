"""Synthetic RA23/RA32 transport tests; no actual model or clinical certification."""

from copy import deepcopy
from hashlib import sha256
import json
from pathlib import Path

from docx import Document
from pypdf import PdfReader
import pytest

from agent.output_check import verify_output
from agent.field_citations import profile_field, split_field_citations
from agent.pipeline import build_downloads, review_result
from agent.review import CITATION_PATTERN
from agent.retrieve import chunk_documents
from app import ra_mvp_service as service
from evals.run import MockEvaluationClient
from parsers import parse_file

ROOT = Path(__file__).resolve().parents[1]
ARTIFACTS = ROOT / 'outputs/ra-clinical-safety/pipeline'
FORMS = ('ra_law_form_23_pdf', 'ra_law_form_32_pdf')
LINES = {
    'dose': '합성약A 시험 함량 5 mg임',
    'title': '합성약A 임상시험 제목 합성 계획 검토임',
    'stage': '합성약A 임상시험 단계 제안 단계임',
    'protocol': '합성약A 계획서 식별번호 SYN-01 버전 A임',
    'approval': '합성약A 국내 허가 여부 미승인임',
    'period': '합성약A 보고기간 2026-01-01 ~ 2026-06-30임',
    'patients': '합성약A 누적 투약 임상시험 대상자 수 12명임',
    'safety': '합성약A 안전성 종합 평가 요약은 결론 미확정임',
    'conclusion': '합성약A 결론 추가 확인 필요함',
}


class ClinicalSafetyClient(MockEvaluationClient):
    """Replies exercise interfaces, not semantic-model quality."""

    def __init__(self, form_id):
        super().__init__({'mock_brief': {
            '목적': '합성약A 시험 함량 임상시험 제목 단계 계획서 버전 허가 여부 보고기간 누적 인원 안전성 결론 확인',
            '보고 대상': 'RA 담당자', '보고서 유형': '결과보고서',
            '마감': '', '분량': '양식의 선정 칸', '부족한 정보': [], '질문': [],
        }})
        self.form_id = form_id
        self.calls = []

    def generate_json(self, name, payload):
        self.calls.append((name, deepcopy(payload)))
        if name == 'draft':
            def quote(value, evidence):
                record = next(item for item in payload['sources'] if evidence in item['text'])
                return f"{value} [{record['source_id']}]"

            selected = ('dose', 'stage', 'protocol', 'approval') if self.form_id == FORMS[0] else ('period', 'patients', 'safety', 'conclusion')
            body = [quote('□ ' + LINES[key], LINES[key]) for key in selected]
            draft = {'제목': '합성 임상자료 내부 검토', '요약': body[-1], '본문': '\n'.join(body)}
            if self.form_id == FORMS[0]:
                draft.update({
                    '시험약 제품명': quote('합성약A', LINES['dose']),
                    '임상시험 제목': quote('합성 계획 검토', LINES['title']),
                    '임상시험 단계': quote('제안 단계', LINES['stage']),
                    '계획서 식별번호': quote('SYN-01', LINES['protocol']),
                    '국내 허가 여부': quote('미승인임', LINES['approval']),
                })
            else:
                draft.update({
                    '제품명 성분명': quote('합성약A', LINES['patients']),
                    '임상시험 제목': quote('합성 계획 검토', LINES['title']),
                    '누적 투약 임상시험 대상자 수': quote('12', LINES['patients']),
                    '안전성 종합 평가 요약': quote('결론 미확정임', LINES['safety']),
                    '결론': quote('추가 확인 필요함', LINES['conclusion']),
                    '비고': quote('보고기간 2026-01-01 ~ 2026-06-30임', LINES['period']),
                })
            return draft
        if name == 'grounding':
            by_id = {record['source_id']: record['text'] for record in payload['sources']}
            return {'claims': [{'field': item['field'], 'line': item['line'], 'status': 'supported',
                'evidence': [{'source_id': identifier, 'quote': by_id[identifier]}
                             for identifier in CITATION_PATTERN.findall(item['text']) if identifier in by_id]}
                for item in payload['claims']]}
        if name == 'completeness':
            return {'issues': [], 'checked_fields': list(payload['draft'])}
        return super().generate_json(name, payload)


@pytest.fixture
def evidence(tmp_path):
    path = tmp_path / 'synthetic-clinical-safety-evidence.docx'
    doc = Document()
    doc.add_paragraph('합성 시험 원자료임. 실제 임상시험·허가·안전성 결과가 아님.')
    for text in LINES.values():
        doc.add_paragraph('제품명: 합성약A\n' + text)
    doc.add_paragraph('제품명: 합성약B\n합성약B 시험 함량 5 mg임')
    doc.save(path)
    return path


def generate(form_id, evidence):
    template, _, record = service.resolve_mvp_form(form_id)
    client = ClinicalSafetyClient(form_id)
    result = service.generate_mvp('합성약A의 근거만으로 제안 상태 및 보고기간을 유지하여 선정 칸을 작성',
                                  form_id, paths=[evidence], client=client)
    assert result['status'] == 'ready', result.get('review')
    return template, record, client, result


def export(template, result):
    return build_downloads(result, confirmed=True, template_paths={'pdf': template},
                           template_profiles={'pdf': result['template_profile']}, native_review='off')


def printed_values(result):
    # Reopen against the actual sidecar rendering contract; JSON retains citations.
    profile = result['template_profile']
    return {key: split_field_citations(value, profile_field(profile, key))[0]
            for key, value in result['draft'].items()}


def record_case(name, data):
    directory = ARTIFACTS / name
    directory.mkdir(parents=True, exist_ok=True)
    (directory / 'check.json').write_text(json.dumps({
        'scope': '합성 시험 원자료와 주입 mock을 사용한 연결·검수 회귀임',
        'actual_api_requests': 0, 'actual_model_responses': 0, 'human_kpi_observations': 0,
        'native_visual_qa': 'not_performed', 'submission_ready': False, **data,
    }, ensure_ascii=False, indent=2), encoding='utf-8')
    return directory


@pytest.mark.parametrize('form_id', FORMS)
def test_selected_clinical_safety_m1_pipeline_fresh_pdf_and_profile_authority(form_id, evidence, tmp_path):
    source_before = sha256(evidence.read_bytes()).hexdigest()
    template, record, client, result = generate(form_id, evidence)
    original_before = sha256(template.read_bytes()).hexdigest()
    assert result['documents'][0]['파일명'] == evidence.name
    assert all(text in result['documents'][0]['본문'] for text in LINES.values())
    assert result['document_kind'] == record['document_kind']
    assert result['ra_checks']['workflow'] == record['ra_workflow']
    assert not result['ra_checks']['blocking'] and not result['grounding']['blocking']
    assert not result['completeness']['blocking']
    assert {'brief', 'draft', 'boss_review', 'grounding', 'completeness'} <= {name for name, _ in client.calls}
    assert all(field['input_mode'] in {'user_provided', 'source_grounded'}
               and field['kind'] == 'pdf_overlay' for field in result['template_profile']['fields'])
    baseline = deepcopy(result['metrics']['baseline_draft'])
    assert result['metrics']['mvp_form_id'] == form_id
    assert result['metrics']['finalized_at'] is None
    assert result['metrics']['rejection_count'] == result['metrics']['rework_count'] == 0
    for key in record['user_input_keys']:
        assert result['draft'][key] == ''  # Never invent an applicant, approval date or signature.
    source_ids = {item['source_id'] for item in result['sources']}
    selected = {key: value for key, value in result['draft'].items()
                if key in record['source_based_keys'] and value}
    for value in selected.values():
        assert set(CITATION_PATTERN.findall(value)) <= source_ids
        assert CITATION_PATTERN.findall(value)
    payload = export(template, result)['pdf']
    output = tmp_path / 'filled.pdf'
    output.write_bytes(payload)
    independent = verify_output(template, output, printed_values(result), profile=result['template_profile'])
    assert independent['status'] == result['output_verification']['pdf']['status'] == 'passed'
    assert len(PdfReader(output).pages) == record['source_page_count']
    assert result['metrics']['baseline_draft'] == baseline
    assert sha256(template.read_bytes()).hexdigest() == original_before == record['source_sha256']
    assert sha256(evidence.read_bytes()).hexdigest() == source_before
    directory = record_case(form_id, {'source_sha256': source_before, 'original_sha256': original_before,
        'original_unchanged': True, 'profile_sha256': record['profile_sha256'],
        'document_kind': record['document_kind'], 'ra_workflow': record['ra_workflow'],
        'selected_field_count': len(selected), 'selected_fields': list(selected),
        'source_ids': sorted(source_ids), 'baseline_draft': baseline, 'draft': result['draft'],
        'output_sha256': sha256(payload).hexdigest(), 'page_count': len(PdfReader(output).pages),
        'output_verification': independent, 'prompt_calls': [name for name, _ in client.calls]})
    (directory / 'filled.pdf').write_bytes(payload)
    (directory / 'synthetic-evidence.docx').write_bytes(evidence.read_bytes())
    # The passed proof cannot authorize changed input permissions or field kinds.
    for attribute, value in [('input_mode', 'user_provided'), ('kind', 'pdf_form')]:
        forged = deepcopy(result['template_profile'])
        field = next(item for item in forged['fields'] if item['value_key'] in selected)
        field[attribute] = value
        with pytest.raises(ValueError, match='다름|다시|검수|원본|동일'):
            build_downloads(result, confirmed=True, template_paths={'pdf': template},
                            template_profiles={'pdf': forged}, native_review='off')


@pytest.mark.parametrize('form_id,field,before,after', [
    (FORMS[0], '본문', '5 mg', '50 mg'),
    (FORMS[1], '본문', '12명', '120명'),
])
def test_changed_dose_or_cumulative_patients_blocks_without_automatic_correction(form_id, field, before, after, evidence):
    template, _, client, result = generate(form_id, evidence)
    baseline = deepcopy(result['metrics']['baseline_draft'])
    edited = {**result['draft'], field: result['draft'][field].replace(before, after)}
    if form_id == FORMS[1]:
        edited['누적 투약 임상시험 대상자 수'] = result['draft']['누적 투약 임상시험 대상자 수'].replace('12 ', '120 ', 1)
    checked = review_result(result, edited, client=client)
    assert checked['blocking'] and not checked['corrected']
    assert after in checked['draft'][field]
    assert any(issue['code'] in {'ra_quantity_mismatch', 'ra_number_mismatch', 'number_mismatch'}
               and issue['field'] == field for issue in checked['warnings'])
    if form_id == FORMS[1]:
        assert any(issue['code'] in {'ra_number_mismatch', 'number_mismatch'}
                   and issue['field'] == '누적 투약 임상시험 대상자 수' for issue in checked['warnings'])
    result['draft'] = edited
    with pytest.raises(ValueError):
        export(template, result)
    assert result['metrics']['baseline_draft'] == baseline
    record_case('dose-change' if form_id == FORMS[0] else 'patients-change', {
        'blocking': checked['blocking'], 'corrected': checked['corrected'], 'warnings': checked['warnings']})


def test_proposed_plan_and_pending_safety_cannot_become_positive_conclusions(evidence):
    rows = []
    for form_id, field, before, after in [
        (FORMS[0], '임상시험 단계', '제안 단계', '승인 완료임'),
        (FORMS[0], '국내 허가 여부', '미승인임', '승인 완료임'),
        (FORMS[1], '결론', '추가 확인 필요함', '안전성 확인됨'),
    ]:
        template, _, client, result = generate(form_id, evidence)
        edited = {**result['draft'], field: result['draft'][field].replace(before, after)}
        checked = review_result(result, edited, client=client)
        assert checked['blocking'] and any(issue['code'] == 'ra_claim_unverified' for issue in checked['warnings'])
        result['draft'] = edited
        with pytest.raises(ValueError):
            export(template, result)
        rows.append({'form_id': form_id, 'field': field, 'blocking': checked['blocking'], 'warnings': checked['warnings']})
    record_case('unsupported-approval-and-safety', {'rows': rows})


def test_changed_protocol_identifier_cannot_reuse_same_source(evidence):
    template, _, client, result = generate(FORMS[0], evidence)
    assert '버전 A임' in result['draft']['본문']
    edited = {**result['draft'], '계획서 식별번호': result['draft']['계획서 식별번호'].replace('SYN-01', 'SYN-02')}
    checked = review_result(result, edited, client=client)
    assert checked['blocking']
    assert any(issue['code'] in {'ra_identifier_mismatch', 'ra_number_mismatch', 'number_mismatch'}
               and issue['field'] == '계획서 식별번호' for issue in checked['warnings'])
    result['draft'] = edited
    with pytest.raises(ValueError):
        export(template, result)
    record_case('changed-protocol-identifier', {'blocking': checked['blocking'], 'warnings': checked['warnings']})


def test_known_other_product_source_with_same_dose_does_not_ground_selected_product(evidence):
    template, _, client, result = generate(FORMS[0], evidence)
    other = next(item for item in result['sources'] if '합성약B 시험 함량 5 mg임' in item['text'])
    old_line = next(line for line in result['draft']['본문'].splitlines() if '5 mg' in line)
    old_id = CITATION_PATTERN.findall(old_line)[0]
    changed_line = old_line.replace(f'[{old_id}]', f"[{other['source_id']}]")
    edited = {**result['draft'], '본문': result['draft']['본문'].replace(old_line, changed_line)}
    checked = review_result(result, edited, client=client)
    assert checked['blocking']
    assert any(issue['code'] == 'ra_quantity_mismatch' and issue['field'] == '본문' for issue in checked['warnings'])
    result['draft'] = edited
    with pytest.raises(ValueError):
        export(template, result)
    rows = [{'field': '본문', 'blocking': checked['blocking'], 'warnings': checked['warnings']}]
    # These literal registered aliases are product-name controls as well.
    for form_id, field in [(FORMS[0], '시험약 제품명'), (FORMS[1], '제품명 성분명')]:
        template, _, client, result = generate(form_id, evidence)
        other = next(item for item in result['sources'] if '합성약B 시험 함량 5 mg임' in item['text'])
        old_id = CITATION_PATTERN.findall(result['draft'][field])[0]
        edited = {**result['draft'], field: result['draft'][field].replace(f'[{old_id}]', f"[{other['source_id']}]")}
        checked = review_result(result, edited, client=client)
        assert checked['blocking']
        assert any(issue['code'] == 'ra_product_unverified' and issue['field'] == field
                   for issue in checked['warnings'])
        result['draft'] = edited
        with pytest.raises(ValueError):
            export(template, result)
        rows.append({'form_id': form_id, 'field': field, 'blocking': checked['blocking'], 'warnings': checked['warnings']})
    record_case('wrong-product-known-citation', {'rows': rows})


def test_other_reporting_period_citation_is_not_interchangeable(evidence, tmp_path):
    template, _, client, result = generate(FORMS[1], evidence)
    other_file = tmp_path / 'synthetic-other-reporting-period.docx'
    document = Document()
    document.add_paragraph('제품명: 합성약A\n합성약A 보고기간 2026-07-01 ~ 2026-12-31임')
    document.save(other_file)
    other = next(item for item in chunk_documents([parse_file(other_file)])
                 if '2026-07-01 ~ 2026-12-31' in item['text'])
    assert other['source_id'] not in {item['source_id'] for item in result['sources']}
    result['sources'].append(other)
    old_id = CITATION_PATTERN.findall(result['draft']['비고'])[0]
    edited = {**result['draft'], '비고': result['draft']['비고'].replace(f'[{old_id}]', f"[{other['source_id']}]")}
    checked = review_result(result, edited, client=client)
    assert checked['blocking']
    assert any(issue['code'] == 'ra_date_mismatch' and issue['field'] == '비고' for issue in checked['warnings'])
    result['draft'] = edited
    with pytest.raises(ValueError):
        export(template, result)
    record_case('wrong-reporting-period', {'blocking': checked['blocking'], 'warnings': checked['warnings']})


def test_two_explicit_reporting_periods_can_select_first_without_date_fragment_conflict(evidence, tmp_path):
    document = Document(evidence)
    document.add_paragraph('제품명: 합성약A\n합성약A 보고기간 2026-07-01 ~ 2026-12-31임')
    document.save(evidence)
    template, _, _, result = generate(FORMS[1], evidence)
    assert not result['conflicts']
    assert any('2026-07-01 ~ 2026-12-31' in source['text'] for source in result['sources'])
    assert '2026-01-01 ~ 2026-06-30' in result['draft']['비고']
    assert '2026-07-01' not in result['draft']['비고']
    payload = export(template, result)['pdf']
    output = tmp_path / 'two-period-selected-first.pdf'
    output.write_bytes(payload)
    proof = verify_output(template, output, printed_values(result), profile=result['template_profile'])
    assert proof['status'] == 'passed'
    directory = record_case('two-period-selected-first', {'blocking': False,
        'source_periods': ['2026-01-01 ~ 2026-06-30', '2026-07-01 ~ 2026-12-31'],
        'selected_period': '2026-01-01 ~ 2026-06-30', 'output_verification': proof})
    (directory / 'filled.pdf').write_bytes(payload)


def test_new_forms_do_not_promote_direct_values_answers_or_instruction_to_product_evidence(evidence):
    counts = {'field_values_rejected': 0, 'answers_rejected': 0, 'instruction_rejected': 0}
    for form_id, field in [(FORMS[0], '시험약 제품명'), (FORMS[1], '제품명 성분명')]:
        for argument in ('field_values', 'answers'):
            client = ClinicalSafetyClient(form_id)
            with pytest.raises(ValueError, match='직접 입력|원자료'):
                service.generate_mvp('합성 시험', form_id, paths=[evidence], client=client,
                                     **{argument: {field: '합성약A'}})
            assert not client.calls  # Reject before mock/model or parsing costs.
            counts[argument + '_rejected'] += 1
        client = ClinicalSafetyClient(form_id)
        client.case['mock_brief']['양식 항목'] = {field: '합성약A'}
        with pytest.raises(ValueError, match='원자료 기반 항목.*작성 지시'):
            service.generate_mvp('합성약A 제품명을 직접 입력함', form_id, paths=[evidence], client=client)
        counts['instruction_rejected'] += 1
    record_case('user-grounded-promotion', counts)
