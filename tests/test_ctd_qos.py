"""Synthetic DMF to 2.3.S working-draft safeguards."""

from hashlib import sha256
import json
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from lxml import etree
from openpyxl import Workbook
from PIL import Image
import pytest
from reportlab.pdfgen import canvas

from agent.ctd_qos import QOS_S_SECTIONS, prepare_qos_package, propose_dmf_options
from agent.ctd_template import fill_ctd_working_template
from agent.multimodal_intake import collect_multimodal, confirm_intake, generation_sources
from app.qos_ui import build_qos_template
from parsers import parse_file
from templates.fill import PLACEHOLDER, _segments, _set_segment


SAMPLE = Path(__file__).resolve().parents[1] / 'samples' / 'qos_dmf_synthetic.txt'
PRODUCT = '시험정 5 mg'
SUBSTANCE = '시험원료 A'
MANUFACTURER = 'Example API Manufacturing Ltd.'


def _inputs(path=SAMPLE):
    sources = generation_sources(collect_multimodal([path]))
    digest = sha256(path.read_bytes()).hexdigest()
    links = {digest: {'substance_name': SUBSTANCE,
                      'manufacturer_name': MANUFACTURER,
                      'product_name': PRODUCT, 'confirmed': True}}
    return sources, links


def _template(path, sections):
    document = Document()
    document.add_heading('CTD 2.3.S 원료의약품 품질요약 · 검토용', 0)
    document.add_paragraph('합성 예제 양식 · 제출용 완성본 아님')
    table = document.add_table(rows=0, cols=2)
    for identifier in sections:
        cells = table.add_row().cells
        cells[0].text = identifier
        run = cells[1].paragraphs[0].add_run('{{' + identifier + '}}')
        run.bold = True
    document.save(path)
    return path


def test_dmf_qos_all_seven_sections_from_exact_source_and_checked_identity(tmp_path):
    sources, links = _inputs()
    options = propose_dmf_options(sources)
    assert options[0]['substance_names'][0]['value'] == SUBSTANCE
    assert options[0]['manufacturer_names'][0]['value'] == MANUFACTURER
    assert options[0]['manufacturer_names'][0]['evidence'][0]['document_sha256'] == next(iter(links))
    package = prepare_qos_package(sources, product_name=PRODUCT, dmf_links=links)
    assert [item['section_id'] for item in package['sections']] == [
        item['section_id'] for item in QOS_S_SECTIONS]
    assert package['coverage']['needs_manual_check'] == 7
    assert not package['missing_sections'] and not package['submission_ready']
    assert package['actual_model_requests'] == 0
    for section in package['sections']:
        assert section['status'] == 'manual_check'
        assert section['source_section_id'] == section['section_id'].replace('2.3.', '3.2.')
        assert section['draft'] == '\n'.join(
            item['quote'] + ' [' + item['source_id'] + ']'
            for item in section['evidence'])
        assert all(item['document_sha256'] == next(iter(links)) for item in section['evidence'])
    template = _template(tmp_path / 'qos.docx', [item['section_id'] for item in QOS_S_SECTIONS])
    target = tmp_path / 'filled.docx'
    output = fill_ctd_working_template(template, target, package, sources=sources,
                                       product_name=PRODUCT, confirmed_dmf_links=links)
    assert output['output_check']['status'] == 'passed'
    content = parse_file(target)['본문']
    assert '99.4%' in content and '12개월' in content
    assert '{{2.3.S.4}}' not in content
    sidecar = json.loads(output['evidence_path'].read_text(encoding='utf-8'))
    assert sidecar['document_kind'] == package['document_kind']
    assert sidecar['identity_evidence'][next(iter(links))]['manufacturer'][0]['quote'] == MANUFACTURER
    assert sidecar['evidence']['2.3.S.4'][0]['document_sha256'] == next(iter(links))
    assert sidecar['submission_ready'] is False


@pytest.mark.parametrize('change', [
    {'confirmed': False}, {'substance_name': '다른 원료'},
    {'manufacturer_name': '다른 제조원'}, {'product_name': '다른 완제'},
])
def test_qos_rejects_unconfirmed_or_mismatched_dmf_identity(change):
    sources, links = _inputs()
    digest = next(iter(links))
    links[digest].update(change)
    with pytest.raises(ValueError):
        prepare_qos_package(sources, product_name=PRODUCT, dmf_links=links)


def test_qos_rejects_mixed_manufacturers_and_stale_output(tmp_path):
    first = tmp_path / 'first.txt'
    second = tmp_path / 'second.txt'
    first.write_text(SAMPLE.read_text(encoding='utf-8'), encoding='utf-8')
    second.write_text(SAMPLE.read_text(encoding='utf-8').replace(
        MANUFACTURER, 'Another API Ltd.'), encoding='utf-8')
    sources = generation_sources(collect_multimodal([first, second]))
    links = {
        sha256(first.read_bytes()).hexdigest(): {'substance_name': SUBSTANCE,
            'manufacturer_name': MANUFACTURER, 'product_name': PRODUCT, 'confirmed': True},
        sha256(second.read_bytes()).hexdigest(): {'substance_name': SUBSTANCE,
            'manufacturer_name': 'Another API Ltd.', 'product_name': PRODUCT, 'confirmed': True},
    }
    with pytest.raises(ValueError, match='서로 다른'):
        prepare_qos_package(sources, product_name=PRODUCT, dmf_links=links)
    one_sources, one_links = _inputs()
    package = prepare_qos_package(one_sources, product_name=PRODUCT,
                                  dmf_links=one_links, selected_sections=['2.3.S.4'])
    template = _template(tmp_path / 'qos.docx', ['2.3.S.4'])
    with pytest.raises(ValueError, match='바뀜|일치하지'):
        fill_ctd_working_template(template, tmp_path / 'out.docx', package,
                                  sources=one_sources, product_name=PRODUCT,
                                  confirmed_dmf_links={next(iter(one_links)): {
                                      **next(iter(one_links.values())),
                                      'manufacturer_name': 'Another API Ltd.'}})
    assert not (tmp_path / 'out.docx').exists()


def test_qos_missing_section_never_prints_other_section_or_product(tmp_path):
    path = tmp_path / 'partial.txt'
    path.write_text('원료의약품명: 시험원료 A\n'
                    '원료의약품 제조원: Example API Manufacturing Ltd.\n'
                    '3.2.S.4: 관리\n규격: 순도 99.0% 이상\n'
                    '3.2.P.5: 완제 관리\n시험결과: 95%\n', encoding='utf-8')
    sources, links = _inputs(path)
    package = prepare_qos_package(sources, product_name=PRODUCT, dmf_links=links)
    assert package['missing_sections'] == [
        item['section_id'] for item in QOS_S_SECTIONS if item['section_id'] != '2.3.S.4']
    assert '95%' not in package['sections'][3]['draft']
    form = _template(tmp_path / 'full.docx', [item['section_id'] for item in QOS_S_SECTIONS])
    with pytest.raises(ValueError, match='누락|상충'):
        fill_ctd_working_template(form, tmp_path / 'out.docx', package,
                                  sources=sources, product_name=PRODUCT,
                                  confirmed_dmf_links=links)


def test_qos_mixed_office_pdf_and_sheet_keep_sha_and_section_boundaries(tmp_path):
    docx = tmp_path / 'identity.docx'
    word = Document()
    for line in ('Drug substance name: API Alpha', 'Manufacturer: Example API Ltd.',
                 '3.2.S.1: General information', 'Molecular formula: C10H12N2O'):
        word.add_paragraph(line)
    word.save(docx)
    pdf = tmp_path / 'control.pdf'
    drawing = canvas.Canvas(str(pdf))
    for index, line in enumerate(('Drug substance name: API Alpha',
                                  'Manufacturer: Example API Ltd.',
                                  '3.2.S.4: Control of drug substance',
                                  'Assay: 99.4%')):
        drawing.drawString(48, 770 - index * 24, line)
    drawing.save()
    xlsx = tmp_path / 'stability.xlsx'
    book = Workbook()
    sheet = book.active
    for row in (('Drug substance name', 'API Alpha'),
                ('Manufacturer', 'Example API Ltd.'),
                ('3.2.S.7', 'Stability'), ('Time point', '12 months'),
                ('Assay', '99.2%')):
        sheet.append(row)
    book.save(xlsx)
    sources = generation_sources(collect_multimodal([docx, pdf, xlsx]))
    links = {sha256(path.read_bytes()).hexdigest(): {
        'substance_name': 'API Alpha', 'manufacturer_name': 'Example API Ltd.',
        'product_name': 'Product Alpha', 'confirmed': True}
        for path in (docx, pdf, xlsx)}
    package = prepare_qos_package(sources, product_name='Product Alpha', dmf_links=links,
                                  selected_sections=['2.3.S.1', '2.3.S.4', '2.3.S.7'])
    sections = package['sections']
    assert [item['status'] for item in sections] == ['manual_check'] * 3
    assert [next(iter({e['document_sha256'] for e in item['evidence']})) for item in sections] == [
        sha256(path.read_bytes()).hexdigest() for path in (docx, pdf, xlsx)]
    assert sections[1]['evidence'][0]['page'] == 1
    assert any(e['sheet'] for e in sections[2]['evidence'])
    assert '99.4%' not in sections[2]['draft'] and '99.2%' not in sections[1]['draft']


def test_qos_image_transcription_requires_original_page_receipt(tmp_path):
    image = tmp_path / 'dmf.png'
    Image.new('RGB', (120, 80), 'white').save(image)
    digest = sha256(image.read_bytes()).hexdigest()
    transcription = ('원료의약품명: 시험원료 A\n'
                     '원료의약품 제조원: Example API Manufacturing Ltd.\n'
                     '3.2.S.4: 관리\n규격: 순도 99.0% 이상')
    intake = collect_multimodal([image], transcriptions={digest: [
        {'page': 1, 'text': transcription}]})
    assert not generation_sources(intake)
    unconfirmed = intake['sources'][0]
    with pytest.raises(ValueError):
        prepare_qos_package([unconfirmed], product_name=PRODUCT, dmf_links={digest: {
            'substance_name': SUBSTANCE, 'manufacturer_name': MANUFACTURER,
            'product_name': PRODUCT, 'confirmed': True}})
    checked = confirm_intake(intake, [
        {'source_id': source['source_id'], 'fingerprint': source['verification_fingerprint']}
        for source in intake['sources']], confirmed=True)
    package = prepare_qos_package(generation_sources(checked), product_name=PRODUCT,
                                  dmf_links={digest: {'substance_name': SUBSTANCE,
                                      'manufacturer_name': MANUFACTURER,
                                      'product_name': PRODUCT, 'confirmed': True}},
                                  selected_sections=['2.3.S.4'])
    assert package['sections'][0]['status'] == 'manual_check'
    assert package['sections'][0]['evidence'][0]['document_sha256'] == digest


def test_qos_hwpx_template_keeps_original_layout_and_source_sidecar(tmp_path):
    sources, links = _inputs()
    package = prepare_qos_package(sources, product_name=PRODUCT,
                                  dmf_links=links, selected_sections=['2.3.S.4'])
    original = SAMPLE.parents[1] / 'templates' / 'result_report.hwpx'
    template = tmp_path / 'qos.hwpx'
    inserted = False
    with ZipFile(original) as before, ZipFile(template, 'w') as after:
        for info in before.infolist():
            data = before.read(info.filename)
            if info.filename == 'Contents/section0.xml':
                root = etree.fromstring(data)
                for paragraph in root.xpath(".//*[local-name()='p']"):
                    segments = _segments(paragraph, 'hwpx')
                    if not segments or not PLACEHOLDER.search(''.join(s.value for s in segments)):
                        continue
                    _set_segment(segments[0], '{{2.3.S.4}}' if not inserted else '검토용 양식', 'hwpx')
                    for segment in segments[1:]:
                        _set_segment(segment, '', 'hwpx')
                    inserted = True
                data = etree.tostring(root, encoding='UTF-8', xml_declaration=True)
            after.writestr(info, data)
    assert inserted
    result = fill_ctd_working_template(template, tmp_path / 'filled.hwpx', package,
                                       sources=sources, product_name=PRODUCT,
                                       confirmed_dmf_links=links)
    assert result['output_check']['status'] == 'passed'
    assert '99.4%' in parse_file(result['output_path'])['본문']


def test_default_qos_example_form_fills_only_current_evidenced_sections(tmp_path):
    sources, links = _inputs()
    package = prepare_qos_package(sources, product_name=PRODUCT,
                                  dmf_links=links, selected_sections=['2.3.S.4'])
    template = tmp_path / 'example.docx'
    template.write_bytes(build_qos_template(['2.3.S.4']))
    result = fill_ctd_working_template(template, tmp_path / 'out.docx', package,
                                       sources=sources, product_name=PRODUCT,
                                       confirmed_dmf_links=links)
    assert result['output_check']['status'] == 'passed'
    assert '99.4%' in parse_file(result['output_path'])['본문']


def test_qos_does_not_treat_finished_product_manufacturer_as_dmf_manufacturer(tmp_path):
    path = tmp_path / 'mixed-roles.txt'
    path.write_text('Drug substance name: API Alpha\n'
                    '3.2.S.4: Control of drug substance\nAssay: 99.4%\n'
                    '3.2.P.3: Manufacture of drug product\n'
                    'Manufacturer: Finished Product Plant\n', encoding='utf-8')
    sources = generation_sources(collect_multimodal([path]))
    assert propose_dmf_options(sources)[0]['manufacturer_names'] == []
    digest = sha256(path.read_bytes()).hexdigest()
    with pytest.raises(ValueError, match='manufacturer'):
        prepare_qos_package(sources, product_name='Product Alpha',
                            dmf_links={digest: {'substance_name': 'API Alpha',
                                'manufacturer_name': 'Finished Product Plant',
                                'product_name': 'Product Alpha', 'confirmed': True}},
                            selected_sections=['2.3.S.4'])
