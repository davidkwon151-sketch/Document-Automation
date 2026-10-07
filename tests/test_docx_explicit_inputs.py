from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
import pytest

from templates import analyze_template, fill_compatible_template, TemplateError


def _parts(path):
    with ZipFile(path) as archive:
        return {item: archive.read(item) for item in archive.namelist()}


def _legacy(paragraph, enabled=True):
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    data = OxmlElement('w:ffData')
    name = OxmlElement('w:name'); name.set(qn('w:val'), 'Text1')
    data.append(name)
    flag = OxmlElement('w:enabled'); flag.set(qn('w:val'), '1' if enabled else '0')
    data.append(flag); data.append(OxmlElement('w:textInput')); begin.append(data)
    paragraph.add_run()._r.append(begin)
    instruction = OxmlElement('w:instrText'); instruction.text = ' FORMTEXT '
    paragraph.add_run()._r.append(instruction)
    separate = OxmlElement('w:fldChar'); separate.set(qn('w:fldCharType'), 'separate')
    paragraph.add_run()._r.append(separate)
    paragraph.add_run('\u2002\u2002').italic = True
    paragraph.add_run('\u2002\u2002')
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    paragraph.add_run()._r.append(end)


def test_protected_formtext_fills_result_only_preserving_protection_and_instructions(tmp_path):
    doc = Document()
    doc.add_paragraph('Read these instructions; do not overwrite.')
    paragraph = doc.add_paragraph('Prepared by: ')
    _legacy(paragraph)
    protection = OxmlElement('w:documentProtection')
    protection.set(qn('w:enforcement'), '1'); protection.set(qn('w:edit'), 'forms')
    protection.set(qn('w:hash'), 'unchanged-password-hash')
    doc.settings.element.append(protection)
    source = tmp_path / 'form.docx'; doc.save(source)
    profile = analyze_template(source)
    assert len(profile['fields']) == 1
    field = profile['fields'][0]
    assert field['kind'] == 'docx_legacy_text' and field['input_required']
    target = fill_compatible_template(source, {'author': '사용자 제공 이름'}, tmp_path / 'out.docx', {field['id']: 'author'}, profile)
    before, after = _parts(source), _parts(target)
    assert {part for part in before if before[part] != after[part]} == {'word/document.xml'}
    assert Document(target).paragraphs[1].text == 'Prepared by: 사용자 제공 이름'
    assert Document(target).paragraphs[0].text == 'Read these instructions; do not overwrite.'
    ns = {'w': qn('w:p').split('}')[0][1:]}
    a, b = etree.fromstring(before['word/document.xml']), etree.fromstring(after['word/document.xml'])
    for xpath in ['//w:ffData', '//w:instrText', '//w:fldChar', '//w:rPr', '//w:pPr']:
        assert [etree.tostring(x) for x in a.xpath(xpath, namespaces=ns)] == [etree.tostring(x) for x in b.xpath(xpath, namespaces=ns)]
    with pytest.raises(TemplateError, match='한 줄'):
        fill_compatible_template(source, {'author': 'two\nlines'}, tmp_path / 'bad.docx', {field['id']: 'author'}, profile)


@pytest.mark.parametrize('protection,enabled', [('readOnly', True), ('forms', False)])
def test_readonly_and_disabled_fields_remain_blocked(tmp_path, protection, enabled):
    doc = Document(); _legacy(doc.add_paragraph('Name: '), enabled)
    prop = OxmlElement('w:documentProtection'); prop.set(qn('w:edit'), protection); prop.set(qn('w:enforcement'), '1')
    doc.settings.element.append(prop)
    source = tmp_path / 'blocked.docx'; doc.save(source)
    assert not analyze_template(source)['fields']


@pytest.mark.parametrize('blank', ['_____', '.....', '…………', '    '])
def test_explicit_inline_blanks_fill_without_removing_label_or_run_style(tmp_path, blank):
    doc = Document(); p = doc.add_paragraph(); p.add_run('Subject:').bold = True
    p.add_run(blank).underline = True
    doc.add_paragraph('Instructions: review this before submitting...')
    source = tmp_path / 'inline.docx'; doc.save(source)
    fields = analyze_template(source)['fields']
    assert len(fields) == 1 and fields[0]['kind'] == 'docx_inline'
    target = fill_compatible_template(source, {'title': 'User provided value'}, tmp_path / 'out.docx', {fields[0]['id']: 'title'})
    output = Document(target)
    assert output.paragraphs[0].text == 'Subject:User provided value'
    assert output.paragraphs[0].runs[0].bold and output.paragraphs[0].runs[1].underline
    assert output.paragraphs[1].text == 'Instructions: review this before submitting...'


def test_dotted_table_cell_and_blank_cell_only_fill_actual_input_cells(tmp_path):
    doc = Document(); table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'Name'; table.cell(0, 1).text = '______'
    table.cell(1, 0).text = 'Department'
    source = tmp_path / 'table.docx'; doc.save(source)
    fields = analyze_template(source)['fields']
    assert len(fields) == 2
    target = fill_compatible_template(source, {'name': 'Person', 'department': 'Finance'}, tmp_path / 'out.docx', {fields[0]['id']: 'name', fields[1]['id']: 'department'})
    result = Document(target).tables[0]
    assert result.cell(0, 0).text == 'Name' and result.cell(1, 0).text == 'Department'
    assert result.cell(0, 1).text == 'Person' and result.cell(1, 1).text == 'Finance'


def test_learning_and_user_configuration_keep_legacy_direct_input_and_context(tmp_path):
    from agent.template_learning import learn_template
    from app.form_config import mapping_rows, configure_profile
    class Client:
        def generate_json(self, prompt, payload):
            original = payload['fields'][0]
            assert original['kind'] == 'docx_legacy_text'
            assert original['input_required'] is True
            assert 'Name:' in original['context']
            return {'fields': [dict(id=original['id'], kind=original['kind'], value_key='Author', required=False,
                input_required=True, input_mode='user_provided', max_chars=100, confidence=.9)]}
    doc = Document(); _legacy(doc.add_paragraph('Name: '))
    source = tmp_path / 'legacy.docx'; doc.save(source)
    profile = learn_template(source, Client(), cache_dir=tmp_path / 'cache')
    assert profile['fields'][0]['input_mode'] == 'user_provided'
    rows = mapping_rows(profile); rows[0]['채울 값'] = 'Author'; rows[0]['직접 입력'] = False
    configured, _ = configure_profile(profile, rows)
    assert configured['fields'][0]['input_required'] is True
    assert configured['fields'][0]['input_mode'] == 'user_provided'


def test_forged_profile_cannot_write_outside_protected_form_input(tmp_path):
    doc = Document(); doc.add_paragraph('Name:____')
    source = tmp_path / 'source.docx'; doc.save(source)
    profile = analyze_template(source)
    prop = OxmlElement('w:documentProtection'); prop.set(qn('w:edit'), 'forms'); prop.set(qn('w:enforcement'), '1')
    doc.settings.element.append(prop); doc.save(source)
    profile.pop('source_sha256')
    with pytest.raises(TemplateError, match='보호 범위'):
        fill_compatible_template(source, {'value':'FORGED'}, tmp_path / 'bad.docx', {profile['fields'][0]['id']: 'value'}, profile)


@pytest.mark.parametrize('label,direct', [
    ('Name', True), ('Prepared by', True), ('Responsible', True), ('Department', True),
    ('Signature', True), ('Tax ID', True), ('Consent', True), ('Vote', True),
    ('Description of applicant name', True), ('Content approval', True),
    ('Product description', False), ('Technical specification', False), ('Content', False),
    ('Supplier Material No.', False), ('Material Safety Data Sheet is necessary.', True),
    ('제품 설명', False), ('기술사양', False), ('작성자 이름', True),
    ('Produktmerkmale', False), ('Unterschrift', True), ('', True), ('Unknown', True),
    ('Spezifische Produktmerkmale', False), ('Product description FullName', True),
    ('Specific product characteristics (if approved for US-product)', False),
])
def test_legacy_content_classifier_preserves_sensitive_and_ambiguous_fields(tmp_path, label, direct):
    doc = Document(); _legacy(doc.add_paragraph(label + ': ' if label else ''))
    source = tmp_path / 'classified.docx'; doc.save(source)
    field = analyze_template(source)['fields'][0]
    assert field['input_required'] is direct
    assert field['input_mode'] == ('user_provided' if direct else 'source_grounded')
    assert field['narrative_style_required'] is False


@pytest.mark.parametrize('layout', ['previous_cell', 'previous_paragraph'])
def test_observed_business_label_context_grounding_and_ai_mapping_roundtrip(tmp_path, layout):
    from agent.template_learning import learn_template
    from app.form_config import mapping_rows, configure_profile
    from agent.output_check import verify_output
    doc = Document(); table = doc.add_table(rows=1, cols=2)
    if layout == 'previous_cell':
        table.cell(0, 0).text = 'Product description'
    else:
        table.cell(0, 1).text = 'Technical specification:'
    paragraph = table.cell(0, 1).add_paragraph()
    _legacy(paragraph)
    source = tmp_path / 'grounded.docx'; doc.save(source)
    class Client:
        def generate_json(self, prompt, payload):
            original = next(f for f in payload['fields'] if f['kind'] == 'docx_legacy_text')
            assert original['input_required'] is False
            assert 'specification' in original['context'] or 'description' in original['context']
            return {'fields': [dict(id=f['id'], kind=f['kind'], value_key=f['label'], required=False,
                input_required=f['input_required'], input_mode='user_provided' if f['input_required'] else 'source_grounded',
                max_chars=100, confidence=.9) for f in payload['fields']]}
    profile = learn_template(source, Client(), cache_dir=tmp_path / 'cache')
    rows = mapping_rows(profile)
    field = next(f for f in profile['fields'] if f['kind']=='docx_legacy_text')
    assert field['narrative_style_required'] is False
    for row in rows:
        row['채울 값'] = 'specification' if row['입력칸 ID']==field['id'] else ''
    configured, mapping = configure_profile(profile, rows)
    assert configured['fields'][0]['input_mode'] == 'source_grounded'
    values = {'specification': 'Verified source material: stainless steel [S1]'}
    output = fill_compatible_template(source, values, tmp_path / 'out.docx', mapping, configured)
    assert verify_output(source, output, values, profile=configured, mapping=mapping)['status']=='passed'
