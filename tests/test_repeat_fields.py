"""반복 행 매핑은 원문을 바꾸지 않고 위치·원본 제약을 보존해야 함."""

from copy import deepcopy
from hashlib import sha256
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
import pytest

from templates.fill import TemplateError
from templates.repeat_fields import repeat_profile
from templates.repeat_docx import inventory as docx_inventory, transform as docx_transform
from templates.repeat_hwpx import inventory as hwpx_inventory, transform as hwpx_transform, HP
from templates.value_rules import inspect_form_values, validate_rule_profile


ROOT = Path(__file__).resolve().parents[1]


def read_parts(path):
    with ZipFile(path) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def prepare(source, tmp_path, row=2, count=3):
    parts = read_parts(source)
    inventory, transform = (docx_inventory, docx_transform) if source.suffix == '.docx' else (hwpx_inventory, hwpx_transform)
    command = {'table_id':inventory(parts)[0]['id'],'row':row,'count':count}
    changed = transform(parts, command)
    target = tmp_path / ('prepared' + source.suffix)
    with ZipFile(target, 'w') as archive:
        for name, data in parts.items():
            archive.writestr(name, changed.get(name, data))
    binding = {'original_sha256':sha256(source.read_bytes()).hexdigest(),'prepared_sha256':sha256(target.read_bytes()).hexdigest(),'plan':command}
    return target, command, binding


def docx_fixture(tmp_path, outside=False, placeholders=True):
    doc = Document()
    doc.add_paragraph('고정 안내 {{제목}}')
    table = doc.add_table(rows=3, cols=4)
    for cell, value in zip(table.rows[0].cells, ['제품/함량','수량','상태','동의']):
        cell.text = value
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    prototype = table.rows[1]
    if placeholders:
        paragraph = prototype.cells[0].paragraphs[0]
        paragraph.add_run('{{제품').bold = True
        paragraph.add_run('명}} / {{함량}}')
    for column, checkbox in [(2,False),(3,True)]:
        sdt, props, content = OxmlElement('w:sdt'), OxmlElement('w:sdtPr'), OxmlElement('w:sdtContent')
        alias, identifier = OxmlElement('w:alias'), OxmlElement('w:id')
        alias.set(qn('w:val'), '동의' if checkbox else '처리 상태')
        identifier.set(qn('w:val'), str(column + 20))
        props.extend([alias,identifier])
        if checkbox:
            control, state = OxmlElement('w14:checkbox'), OxmlElement('w14:checked')
            state.set(qn('w14:val'), '0')
            control.append(state)
            text = '☐'
        else:
            control = OxmlElement('w:dropDownList')
            props.append(OxmlElement('w:showingPlcHdr'))
            for code, label in [('A','검토 전'),('B','검토 완료')]:
                item = OxmlElement('w:listItem')
                item.set(qn('w:value'),code)
                item.set(qn('w:displayText'),label)
                control.append(item)
            text = '{{선택 안내}}'
        props.append(control)
        content.append(prototype.cells[column].add_paragraph(text)._p)
        sdt.extend([props,content])
        prototype.cells[column]._tc.append(sdt)
    for cell in table.rows[2].cells:
        cell.text = '합계'
    if outside:
        doc.add_paragraph('기존 외부 입력 {{제품명}}')
    source = tmp_path / 'source.docx'
    doc.save(source)
    return source


@pytest.mark.parametrize('outside', [False,True])
def test_docx_slots_unique_native_priority_and_outside_global_tokens(tmp_path, outside):
    source = docx_fixture(tmp_path, outside)
    prepared, command, binding = prepare(source,tmp_path)
    before = source.read_bytes(), prepared.read_bytes()
    profile = repeat_profile(prepared,command,binding)
    assert (source.read_bytes(), prepared.read_bytes()) == before
    fields = profile['fields']
    repeats = [field for field in fields if 'repeat_info' in field]
    assert len(repeats) == 15
    assert len({field['value_key'] for field in fields}) == len(fields)
    assert len({field['id'] for field in fields}) == len(fields)
    assert any(field['id']=='placeholder:제목' for field in fields)
    assert any(field['id']=='placeholder:제품명' for field in fields) is outside
    assert not any(field['id']=='placeholder:함량' for field in fields)
    assert not any(field.get('placeholder_key')=='선택 안내' for field in fields)
    for row in range(1,4):
        members = [field for field in repeats if field['repeat_info']['row']==row]
        assert len(members)==5
        dropdown = next(field for field in members if field['kind']=='docx_choice')
        assert dropdown['options']==['A','B'] and dropdown['input_required'] is True
        assert dropdown['required'] is False
        assert dropdown['choice_items']==[{'value':'A','label':'검토 전'},{'value':'B','label':'검토 완료'}]
        checkbox = next(field for field in members if field['kind']=='docx_checkbox')
        assert checkbox['input_required'] is True and checkbox['required'] is False
        quantity = next(field for field in members if field['kind']=='docx_cell')
        assert quantity['repeat_info']['source_label']=='수량'
        assert len([field for field in members if field['kind']=='docx_repeat_placeholder'])==2
    assert profile['repeat_expansion']==binding
    assert set(profile['repeat_expansion'])=={'original_sha256','prepared_sha256','plan'}
    assert profile['source_sha256']==binding['prepared_sha256']
    validate_rule_profile(profile)


def test_original_unicode_offsets_split_runs_and_source_labels_are_exact(tmp_path):
    source = docx_fixture(tmp_path)
    prepared, command, binding = prepare(source,tmp_path)
    profile = repeat_profile(prepared,command,binding)
    slots = [field for field in profile['fields'] if field['kind']=='docx_repeat_placeholder']
    assert len(slots)==6
    for slot in slots:
        start,end = slot['placeholder_start'],slot['placeholder_end']
        assert slot['anchor_text'][start:end]=='{{'+slot['placeholder_key']+'}}'
        assert slot['id']=='repeat_docx:word/document.xml:'+slot['xml_path']+'#slot:'+str(start)
        assert slot['repeat_info']['column']==1
        assert slot['repeat_info']['source_label']==slot['label']
        assert slot['required'] is True


def test_row_group_checks_only_declared_required_without_promoting_native_controls(tmp_path):
    prepared,command,binding = prepare(docx_fixture(tmp_path),tmp_path)
    profile = repeat_profile(prepared,command,binding)
    assert len(profile['constraints']['groups'])==3
    group = profile['constraints']['groups'][0]
    assert len(group['fields'])==5 and len(group['required_fields'])==2
    optional = next(key for key in group['fields'] if key not in group['required_fields'])
    issues = inspect_form_values({optional:'입력함'},profile)
    assert sum(issue['code']=='form_group_required' for issue in issues)==2
    assert not inspect_form_values({},profile)


def test_optional_blank_rows_do_not_invent_required_fields(tmp_path):
    prepared,command,binding = prepare(docx_fixture(tmp_path,placeholders=False),tmp_path)
    profile = repeat_profile(prepared,command,binding)
    assert not profile['constraints'].get('groups')
    assert all(row['required_definition_pending'] for row in profile['repeat_rows'])
    assert all(not field['required'] for field in profile['fields'] if 'repeat_info' in field)
    assert any('원본 필수 항목 정의 없음' in warning for warning in profile['warnings'])


@pytest.mark.parametrize('change', ['prepared_sha','original_sha','plan','path','count','row','table'])
def test_binding_and_range_mismatches_are_blocked(tmp_path, change):
    prepared,command,binding = prepare(docx_fixture(tmp_path),tmp_path)
    if change=='prepared_sha': binding['prepared_sha256']='0'*64
    elif change=='original_sha': binding['original_sha256']='invalid'
    elif change=='plan': binding['plan']={**command,'count':2}
    elif change=='path': binding['source_path']='secret/file.docx'
    else:
        command = {**command, {'count':'count','row':'row','table':'table_id'}[change]: {'count':200,'row':90,'table':'docx:word/document.xml:/unknown'}[change]}
        binding['plan']=command
    with pytest.raises(TemplateError):
        repeat_profile(prepared,command,binding)


@pytest.mark.parametrize('kind', ['relation','group'])
def test_existing_row_dependencies_cannot_be_silently_dropped_or_guessed(tmp_path, monkeypatch, kind):
    from templates import repeat_fields
    prepared,command,binding = prepare(docx_fixture(tmp_path),tmp_path)
    analyzed = repeat_fields.analyze_template(prepared)
    if kind=='relation':
        analyzed['constraints']={'relations':[{'kind':'less_equal','left':'제품명','right':'함량'}]}
    else:
        analyzed['constraints']={'groups':[{'kind':'all_or_none','fields':['제품명','함량']}]}
    monkeypatch.setattr(repeat_fields,'analyze_template',lambda path: deepcopy(analyzed))
    with pytest.raises(TemplateError,match='관계·그룹'):
        repeat_profile(prepared,command,binding)


def test_outside_required_rules_and_dependencies_are_unchanged(tmp_path, monkeypatch):
    from templates import repeat_fields
    source = docx_fixture(tmp_path)
    doc = Document(source)
    doc.add_paragraph('{{시작일}} / {{종료일}}')
    doc.save(source)
    prepared,command,binding = prepare(source,tmp_path)
    analyzed = repeat_fields.analyze_template(prepared)
    for field in analyzed['fields']:
        if field['value_key'] in {'시작일','종료일'}:
            field['validation']={'type':'date'}
    analyzed['constraints']={'relations':[{'kind':'date_order','start':'시작일','end':'종료일'}]}
    monkeypatch.setattr(repeat_fields,'analyze_template',lambda path: deepcopy(analyzed))
    profile = repeat_profile(prepared,command,binding)
    assert profile['constraints']['relations']==analyzed['constraints']['relations']
    for key in ['시작일','종료일']:
        assert next(field for field in profile['fields'] if field['value_key']==key)==next(field for field in analyzed['fields'] if field['value_key']==key)


def test_global_alias_key_is_preserved_outside_and_source_label_stays_literal(tmp_path):
    source = docx_fixture(tmp_path)
    doc = Document(source)
    doc.add_paragraph('{{건명}}')
    doc.tables[0].rows[1].cells[0].text='{{건명}}'
    doc.save(source)
    prepared,command,binding = prepare(source,tmp_path)
    profile = repeat_profile(prepared,command,binding)
    assert next(field for field in profile['fields'] if field['id']=='placeholder:건명')['value_key']=='제목'
    aliases = [field for field in profile['fields'] if field.get('placeholder_key')=='건명']
    assert len(aliases)==3
    assert all(field['value_key'].startswith('제목 (반복 ') for field in aliases)


def test_hwpx_slots_use_section_paragraph_positions_and_preserve_styles(tmp_path):
    parts = read_parts(ROOT/'samples/sample_company_form.hwpx')
    root = etree.fromstring(parts['Contents/section0.xml'])
    table = root.find('.//{'+HP+'}tbl')
    row = table.findall('{'+HP+'}tr')[1]
    for column, cell in enumerate(row.findall('{'+HP+'}tc'),1):
        p = cell.find('{'+HP+'}subList/{'+HP+'}p')
        for run in list(p): p.remove(run)
        run = etree.SubElement(p,'{'+HP+'}run',charPrIDRef='0')
        etree.SubElement(run,'{'+HP+'}t').text='{{항목'+str(column)+'}}'
    parts['Contents/section0.xml']=etree.tostring(root)
    source = tmp_path/'source.hwpx'
    with ZipFile(source,'w') as archive:
        for name,data in parts.items(): archive.writestr(name,data)
    prepared,command,binding = prepare(source,tmp_path)
    before=prepared.read_bytes()
    profile = repeat_profile(prepared,command,binding)
    slots = [field for field in profile['fields'] if field['kind']=='hwpx_repeat_placeholder']
    assert len(slots)==9
    assert len(profile['constraints']['groups'])==3
    assert not any(field['kind']=='placeholder' for field in profile['fields'])
    assert prepared.read_bytes()==before
    for slot in slots:
        assert slot['id'].startswith('repeat_hwpx:Contents/section0.xml:')
        assert slot['anchor_text'][slot['placeholder_start']:slot['placeholder_end']]=='{{'+slot['placeholder_key']+'}}'
        assert slot['required'] is True


def test_same_token_twice_in_one_cell_gets_distinct_positions_and_keys(tmp_path):
    source = docx_fixture(tmp_path)
    doc = Document(source)
    doc.tables[0].rows[1].cells[0].text='{{제품명}} / {{제품명}}'
    doc.save(source)
    prepared,command,binding = prepare(source,tmp_path)
    profile = repeat_profile(prepared,command,binding)
    slots = [field for field in profile['fields'] if field['kind']=='docx_repeat_placeholder']
    assert len(slots)==6 and len({slot['id'] for slot in slots})==6
    assert len({slot['value_key'] for slot in slots})==6


def test_long_repeated_source_labels_keep_original_but_keys_fit_model_limit(tmp_path):
    source = docx_fixture(tmp_path)
    doc = Document(source)
    token = '원문긴항목' * 30
    doc.tables[0].rows[1].cells[0].text='{{'+token+'}} / {{'+token+'}}'
    doc.save(source)
    prepared,command,binding = prepare(source,tmp_path)
    profile = repeat_profile(prepared,command,binding)
    repeated = [field for field in profile['fields'] if field['kind']=='docx_repeat_placeholder']
    assert len(repeated)==6 and len({field['value_key'] for field in repeated})==6
    assert all(len(field['value_key'])<=100 for field in repeated)
    assert all(field['label']==token and field['repeat_info']['source_label']==token for field in repeated)
    assert all('(반복 ' in field['value_key'] for field in repeated)
    from agent.brief import model_profile
    assert model_profile(profile)
