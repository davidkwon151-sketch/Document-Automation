"""CTD Module 1/3 working draft from confirmed, source-backed uploads."""

from collections import Counter
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import tempfile
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
import streamlit as st

from app.multimodal_ui import render_multimodal_upload
from agent.ctd import (CTD_SECTIONS, prepare_ctd_package,
                       propose_ctd_product_options, propose_ctd_substance_options)
from agent.ctd_mapping import suggest_ctd_section_map
from agent.ctd_template import ctd_template_section_ids, fill_ctd_working_template
from llm.client import LLMClient, LLMError
from templates.compatibility import analyze_template


def _digest(value):
    return sha256(json.dumps(value, ensure_ascii=False, sort_keys=True,
                             allow_nan=False).encode('utf-8')).hexdigest()


def _use_product_option(product, variant):
    st.session_state['ctd_product'] = product
    if variant:
        st.session_state['ctd_variant'] = variant


def _use_template_sections(section_ids):
    st.session_state['ctd_sections'] = section_ids


def _invalidate_ctd(*, links=False):
    for key in ('ctd_result', 'ctd_engine_result', 'ctd_template_export',
                'ctd_result_signature'):
        st.session_state.pop(key, None)
    if links:
        for key in ('ctd_mapping_context', 'ctd_mapping', 'ctd_substance_links',
                    'ctd_ai_proposal', 'ctd_ai_proposal_context', 'ctd_mapping_ai_meta'):
            st.session_state.pop(key, None)


def export_ctd_working_draft(package):
    """Export a reviewable draft and full provenance, never an official dossier."""
    if package.get('submission_ready') is not False or not isinstance(package.get('sections'), list):
        raise ValueError('제출용으로 표시된 결과나 CTD 절 목록이 없는 결과는 내보낼 수 없음')
    document = Document()
    document.styles['Normal'].font.name = '맑은 고딕'
    document.add_heading('CTD Module 1·3 작업 초안', 0)
    document.add_paragraph('제출용 완성본 아님 · 선택한 원자료의 발췌와 빈 절을 담당자가 확인해야 함')
    expected = ['CTD Module 1·3 작업 초안',
                '제출용 완성본 아님 · 선택한 원자료의 발췌와 빈 절을 담당자가 확인해야 함']
    for section in package['sections']:
        identifier = section['section_id']
        heading = f"{identifier} {section['title']}"
        document.add_heading(heading, level=1)
        expected.append(heading)
        status = section['status']
        status_text = {'proposed': '원문 발췌 후보 · 담당자 확인 필요',
                       'missing': '근거 자료 없음',
                       'deferred': '근거 후보 있음 · 자동 작성 보류',
                       'ambiguous': '복수 후보 · 선택 필요',
                       'manual_check': '담당자 직접 확인 필요'}.get(status, '검토 필요')
        document.add_paragraph(status_text)
        expected.append(status_text)
        draft = section.get('draft', '')
        if draft:
            document.add_paragraph(draft)
            expected.append(draft)
        for evidence in section.get('evidence', []):
            caption = (f"출처 {evidence['source_id']} · {evidence['filename']} · "
                       f"{evidence.get('page') or evidence.get('sheet') or evidence.get('location') or '위치 확인 필요'}")
            if evidence.get('source_scope') == 'confirmed_substance_link':
                caption += ' · 담당자 확인 원료–완제 연결'
            document.add_paragraph(caption, style='Caption')
            expected.append(caption)
    docx = BytesIO()
    document.save(docx)
    reopened = Counter(paragraph.text for paragraph in Document(BytesIO(docx.getvalue())).paragraphs)
    if any(reopened[text] < count for text, count in Counter(expected).items()):
        raise ValueError('CTD 초안 저장 후 절·출처·상태를 대조하지 못함')
    evidence = json.dumps(package, ensure_ascii=False, indent=2, allow_nan=False).encode('utf-8')
    archive = BytesIO()
    with ZipFile(archive, 'w', ZIP_DEFLATED) as bundle:
        bundle.writestr('CTD_Module1_3_작업초안.docx', docx.getvalue())
        bundle.writestr('CTD_출처와_누락_검토.json', evidence)
    with ZipFile(BytesIO(archive.getvalue())) as reopened_zip:
        if (reopened_zip.read('CTD_Module1_3_작업초안.docx') != docx.getvalue()
                or reopened_zip.read('CTD_출처와_누락_검토.json') != evidence
                or json.loads(evidence) != package):
            raise ValueError('CTD 작업 묶음 저장 후 문서·근거를 대조하지 못함')
    return {'document': docx.getvalue(), 'evidence': evidence, 'zip': archive.getvalue()}


def render_ctd():
    """Streamlit pilot. Automated candidates still need product and source review."""
    st.subheader('CTD 신규 품목허가·신고 · Module 1/3')
    st.caption('제네릭·개량신약 RA 작업 초안. 이 화면의 절 이름은 작업 분류이며 최신 법정 제출 요건을 대신하지 않습니다.')
    product = st.text_input('원자료의 정확한 제품명', key='ctd_product').strip()
    variant = st.text_input('선택 제형·함량 (해당 시)', key='ctd_variant').strip()
    catalog = list(CTD_SECTIONS)
    ids = [item['section_id'] for item in catalog]
    titles = {item['section_id']: item['title'] for item in catalog}
    selected = st.multiselect('작성할 Module 1·3 절', ids, default=ids,
                              format_func=lambda identifier: f'{identifier} · {titles[identifier]}',
                              key='ctd_sections')
    template_upload = st.file_uploader('CTD 절 자리표시자 작업 양식 DOCX/HWPX (선택)',
                                       type=['docx', 'hwpx'], key='ctd_template_upload')
    template_bytes, template_sha = None, None
    if template_upload is not None:
        template_name = template_upload.name
        template_bytes = template_upload.getvalue()
        suffix = Path(template_name).suffix.lower()
        if (Path(template_name).name != template_name or '/' in template_name or '\\' in template_name
                or suffix not in {'.docx', '.hwpx'} or not isinstance(template_bytes, bytes)
                or not 0 < len(template_bytes) <= 30 * 1024 * 1024):
            _invalidate_ctd(links=True)
            st.error('CTD 작업 양식은 경로 없는 30 MiB 이하 DOCX/HWPX 파일이어야 합니다.')
            return
        template_sha = sha256(template_bytes).hexdigest()
        if st.session_state.get('ctd_template_profile_sha') != template_sha:
            try:
                with tempfile.TemporaryDirectory(prefix='ctd-form-analysis-') as directory:
                    path = Path(directory) / ('원본' + suffix)
                    path.write_bytes(template_bytes)
                    st.session_state['ctd_template_profile'] = analyze_template(path)
                    st.session_state['ctd_template_profile_sha'] = template_sha
            except Exception:
                st.session_state.pop('ctd_template_profile', None)
                st.warning('양식의 자리표시자를 미리 분석하지 못했습니다. 출력할 때 다시 확인합니다.')
        profile = st.session_state.get('ctd_template_profile') or {}
        try:
            form_ids = ctd_template_section_ids(profile)
        except ValueError:
            form_ids = []
        if form_ids and len(form_ids) == len(set(form_ids)) and all(identifier in ids for identifier in form_ids):
            st.caption('양식에서 찾은 CTD 절: ' + ', '.join(form_ids))
            st.button('양식 자리표시자 절만 선택', on_click=_use_template_sections,
                      args=(form_ids,), key='ctd_template_select_' + template_sha[:16])
        else:
            st.warning('양식에 CTD 절 이외의 칸이나 중복/불명확한 자리표시자가 있습니다. 정확한 절만 있는 양식을 사용해 주세요.')
        st.caption('이 양식은 검토용 작업 파일입니다. 양식에 있는 절이 선택되어 있고 각 절의 근거가 있어야 기입할 수 있습니다. 다른 선택 절은 근거 기록에 남습니다.')
    st.caption('PDF·DOCX·HWPX·XLSX·TXT와 사진·스캔을 함께 첨부할 수 있습니다. 사진·스캔의 원문은 직접 대조·확인해야 합니다. 이미지 모델 호출은 화면에서 명시적으로 선택할 때만 발생합니다.')
    try:
        intake = render_multimodal_upload(key_prefix='ctd', context={
            'product': product, 'variant': variant, 'sections': selected,
            'template_sha256': template_sha})
    except Exception as exc:
        _invalidate_ctd(links=True)
        st.error(str(exc) if isinstance(exc, ValueError) else '원자료를 읽지 못했습니다. 첨부 파일을 확인해 주세요.')
        return
    if intake['changed']:
        _invalidate_ctd(links=True)
    if not intake['ready']:
        _invalidate_ctd(links=True)
        st.warning('미확인 사진·스캔 원문 또는 읽지 못한 파일이 있습니다. 원본 대조·전사 후 작성할 수 있습니다.')
        return
    if intake['sources']:
        try:
            options = propose_ctd_product_options(intake['sources'])
        except ValueError:
            options = {'product_names': [], 'variant_parts': [], 'proposed_product_variant': None}
        if options['product_names']:
            with st.expander('원자료에서 제품명·제형 고르기'):
                names = {item['value']: item for item in options['product_names']}
                choices = ['', *names]
                chosen = st.selectbox('원문 제품명 후보', choices,
                                      index=1 if len(names) == 1 else 0,
                                      format_func=lambda value: value or '제품명 선택',
                                      key='ctd_product_candidate_' + intake['signature'][:16])
                if chosen:
                    for evidence in names[chosen]['evidence']:
                        st.caption(f"{evidence['filename']} · "
                                   f"{evidence.get('page') or evidence.get('sheet') or evidence.get('location')} · "
                                   f"SHA {evidence['document_sha256']}")
                        st.code(evidence['quote'], language=None)
                    variant_option = options.get('proposed_product_variant')
                    include_variant = False
                    if variant_option:
                        st.caption('원문 제형·함량 후보: ' + variant_option)
                        for item in options['variant_parts']:
                            for evidence in item['evidence']:
                                st.caption(f"{evidence['filename']} · SHA {evidence['document_sha256']}")
                                st.code(evidence['quote'], language=None)
                        include_variant = st.checkbox('원문 제형·함량도 함께 사용', value=True,
                                                      key='ctd_variant_candidate_' + intake['signature'][:16])
                    st.button('확인한 제품명 입력칸에 적용',
                              on_click=_use_product_option,
                              args=(chosen, variant_option if include_variant else ''),
                              key='ctd_product_apply_' + intake['signature'][:16])
    if not product or not selected:
        _invalidate_ctd(links=True)
        st.info('제품명과 작성할 절을 선택해 주세요.')
        return
    current = _digest({'intake': intake['signature'], 'product': product,
                       'variant': variant, 'sections': selected, 'template_sha256': template_sha})
    if st.session_state.get('ctd_mapping_context') != current:
        st.session_state['ctd_mapping_context'] = current
        st.session_state['ctd_mapping'] = {}
        st.session_state['ctd_substance_links'] = {}
        st.session_state.pop('ctd_mapping_ai_meta', None)
        st.session_state.pop('ctd_ai_proposal', None)
        _invalidate_ctd()
    mapping = {key: list(value) for key, value in st.session_state['ctd_mapping'].items()}
    substance_links = dict(st.session_state['ctd_substance_links'])
    remap = False
    if any(identifier.startswith('3.2.S.') for identifier in selected):
        try:
            substances = propose_ctd_substance_options(intake['sources'])
        except ValueError:
            substances = []
        if substances:
            with st.expander('원료의약품 자료와 선택 완제 제품 연결'):
                st.caption('원료명과 완제 제품명의 관계는 파일명이나 AI 추정으로 결정하지 않습니다. 담당자가 실제 원문과 제품을 확인한 경우에만 3.2.S 절 후보로 연결합니다.')
                by_sha = {item['document_sha256']: item for item in substances}
                digest = st.selectbox('원료 원본 파일', ['', *by_sha],
                                      format_func=lambda value: '원본 선택' if not value else (
                                          f"{by_sha[value]['filename']} · SHA {value}"),
                                      key='ctd_substance_file_' + current[:16])
                if digest:
                    item = by_sha[digest]
                    names = {row['value']: row for row in item['substance_names']}
                    substance = st.selectbox('원문 원료명', ['', *names],
                                             index=1 if len(names) == 1 else 0,
                                             format_func=lambda value: value or '원료명 선택',
                                             key='ctd_substance_name_' + current[:16] + '_' + digest[:12])
                    if substance:
                        if substance_links.get(digest, {}).get('substance_name') != substance:
                            _invalidate_ctd()
                        for evidence in names[substance]['evidence']:
                            st.caption(f"{evidence['source_id']} · {evidence['filename']} · "
                                       f"{evidence.get('page') or evidence.get('sheet') or evidence.get('location')} · "
                                       f"SHA {evidence['document_sha256']}")
                            st.code(evidence['quote'], language=None)
                        st.write(f'선택 완제 제품: {product}')
                        confirmed = st.checkbox('이 원료가 선택 완제 제품에 사용되는 관계를 원자료와 대조했음',
                                                key='ctd_substance_confirm_' + _digest((current, digest, substance))[:16])
                        if st.button('확인한 원료–완제 관계로 S 절 다시 작성',
                                     disabled=not confirmed, key='ctd_substance_apply'):
                            substance_links[digest] = {'substance_name': substance,
                                                        'product_name': product, 'confirmed': True}
                            st.session_state['ctd_substance_links'] = substance_links
                            st.session_state.pop('ctd_ai_proposal', None)
                            st.session_state.pop('ctd_mapping_ai_meta', None)
                            remap = True
                    if digest in substance_links and st.button('이 원료 연결 해제', key='ctd_substance_remove'):
                        substance_links.pop(digest)
                        st.session_state['ctd_substance_links'] = substance_links
                        st.session_state.pop('ctd_ai_proposal', None)
                        st.session_state.pop('ctd_mapping_ai_meta', None)
                        remap = True
                if substance_links:
                    st.caption(f'담당자가 확인한 원료 원본 {len(substance_links)}개 · S 절 후보에만 사용')
    with st.expander('절별 원자료 위치 선택·확인'):
        st.caption('자동 분류가 모호하거나 빠진 경우 원문 조각을 고릅니다. 선택은 현재 제품·첨부·절에만 적용됩니다.')
        st.caption('절 연결을 확인해도 본문 내용·법정 적합성 검토는 별도로 남습니다.')
        if mapping:
            st.dataframe([{'절': section_id, '확인한 출처 ID': ', '.join(source_ids)}
                          for section_id, source_ids in mapping.items()], hide_index=True)
        target = st.selectbox('연결할 CTD 절', ['', *selected],
                              format_func=lambda value: '절 선택' if not value else f'{value} · {titles[value]}',
                              key='ctd_mapping_target_' + current[:16])
        if target:
            by_id = {source['source_id']: source for source in intake['sources']}
            source_ids = ['', *by_id]
            selected_id = mapping.get(target, [''])[0]
            source_key = 'ctd_mapping_source_' + current[:16] + '_' + target
            source_id = st.selectbox('원문 조각 선택', source_ids,
                                     index=source_ids.index(selected_id) if selected_id in source_ids else 0,
                                     format_func=lambda value: '연결하지 않음' if not value else (
                                         f"{value} · {by_id[value]['filename']} · "
                                         f"{by_id[value].get('page') or by_id[value].get('sheet') or by_id[value].get('location')} · "
                                         f"{by_id[value]['text'][:70]}"),
                                     key=source_key)
            if source_id and source_id not in mapping.get(target, []):
                _invalidate_ctd()
            if source_id:
                source = by_id[source_id]
                st.caption(f"{source['filename']} · SHA {source['document_sha256']} · "
                           f"{source.get('page') or source.get('sheet') or source.get('location')} · 출처 ID {source_id}")
                st.code(source['text'], language=None)
                confirmed = st.checkbox('원본의 전체 문구·제품 범위·이 절과의 연결을 대조함',
                                        key='ctd_mapping_confirm_' + _digest((current, target, source_id,
                                                                              source['verification_fingerprint']))[:16])
                duplicate = any(source_id in values for section_id, values in mapping.items()
                                if section_id != target)
                if duplicate:
                    st.error('이 출처 ID는 이미 다른 CTD 절에 연결됨. 절별 고유 원문을 선택해야 함')
                if st.button('확인한 원문으로 절 다시 작성',
                             disabled=not confirmed or duplicate, key='ctd_mapping_apply'):
                    mapping[target] = [source_id]
                    st.session_state['ctd_mapping'] = mapping
                    st.session_state.pop('ctd_mapping_ai_meta', None)
                    st.session_state.pop('ctd_ai_proposal', None)
                    remap = True
            if target in mapping and st.button('이 절의 수동 연결 해제', key='ctd_mapping_remove'):
                mapping.pop(target)
                st.session_state['ctd_mapping'] = mapping
                st.session_state.pop('ctd_mapping_ai_meta', None)
                st.session_state.pop('ctd_ai_proposal', None)
                st.session_state[source_key] = ''
                remap = True
    ai_context = _digest({'current': current, 'substance_links': substance_links})
    if st.session_state.get('ctd_ai_proposal_context') != ai_context:
        st.session_state['ctd_ai_proposal_context'] = ai_context
        st.session_state.pop('ctd_ai_proposal', None)
    with st.expander('AI 절 연결 제안 · 선택적으로 사용'):
        st.caption('설정된 모델에 현재 제품의 확인된 원문 조각을 보내 절 ID만 제안받습니다. 제안을 원본과 대조한 뒤에만 적용합니다. 모델 연결이 안 되면 위 직접 선택과 기본 원문 분류를 사용할 수 있습니다.')
        if st.button('AI 절 연결 제안', key='ctd_ai_suggest'):
            st.session_state.pop('ctd_ai_proposal', None)
            try:
                st.session_state['ctd_ai_proposal'] = suggest_ctd_section_map(
                    intake['intake'], product_name=product, product_variant=variant,
                    selected_sections=selected, confirmed_substance_links=substance_links or None,
                    client=LLMClient())
            except LLMError as exc:
                st.error({'quota': 'API 잔액·할당량 부족', 'billing_limit': 'API 사용 한도 도달',
                          'rate_limit': 'API 일시 요청 제한', 'configuration': '모델 연결 설정 필요',
                          'connection': '모델 서버 연결 실패', 'timeout': '모델 요청 시간 초과'}.get(
                              exc.kind, '모델 응답 오류') + ' · 기존 원문 선택 기능은 계속 사용할 수 있습니다.')
            except ValueError as exc:
                st.error(str(exc))
            except Exception:
                st.error('AI 절 연결 제안을 완료하지 못했습니다. 원문을 직접 연결해 주세요.')
        proposal = st.session_state.get('ctd_ai_proposal')
        if proposal:
            proposed_map = proposal.get('section_map', {})
            by_id = {source['source_id']: source for source in intake['sources']}
            st.caption(f"AI 연결 후보 {sum(len(value) for value in proposed_map.values())}개 · "
                       f"연결하지 못한 원문 {len(proposal.get('unmapped_source_ids', []))}개 · 아직 적용되지 않음")
            for section_id, source_ids in proposed_map.items():
                st.write(section_id + ' · ' + titles[section_id])
                for source_id in source_ids:
                    source = by_id[source_id]
                    st.caption(f"{source_id} · {source['filename']} · "
                               f"{source.get('page') or source.get('sheet') or source.get('location')} · "
                               f"SHA {source['document_sha256']}")
                    st.code(source['text'], language=None)
            combined = {**mapping, **proposed_map}
            all_ids = [identifier for values in combined.values() for identifier in values]
            conflict = (any(section_id in mapping and mapping[section_id] != source_ids
                            for section_id, source_ids in proposed_map.items())
                        or len(all_ids) != len(set(all_ids)))
            if conflict:
                st.warning('이미 직접 확인한 절 연결과 다릅니다. 해당 연결을 해제하고 다시 제안받아야 합니다.')
            ai_confirmed = st.checkbox('제안된 원문 전체·제품·제형·절 연결을 원본과 대조했음',
                                       key='ctd_ai_confirm_' + _digest((ai_context, proposed_map))[:16])
            if st.button('확인한 AI 연결을 적용해 다시 작성',
                         disabled=not proposed_map or not ai_confirmed or conflict,
                         key='ctd_ai_apply'):
                mapping.update(proposed_map)
                st.session_state['ctd_mapping'] = mapping
                st.session_state['ctd_mapping_ai_meta'] = {
                    'mapping_origin': 'model_proposal_confirmed_by_user',
                    'mapping_model_requests': proposal.get('actual_model_requests', 0),
                    'mapping_model_attempts': proposal.get('model_request_attempts', 0)}
                st.session_state.pop('ctd_ai_proposal', None)
                remap = True
    result_signature = _digest({'current': current, 'section_map': mapping,
                                'substance_links': substance_links})
    create = st.button('원자료로 CTD 작업 초안 만들기', disabled=not intake['sources'], key='ctd_prepare')
    if create or remap:
        try:
            package = prepare_ctd_package(intake['sources'], product_name=product,
                                          product_variant=variant, selected_sections=selected,
                                          section_map=mapping or None,
                                          confirmed_substance_links=substance_links or None)
        except Exception as exc:
            _invalidate_ctd()
            st.error(str(exc) if isinstance(exc, ValueError) else 'CTD 절별 원자료 대조를 완료하지 못했습니다.')
            return
        st.session_state['ctd_engine_result'] = package
        displayed = deepcopy(package)
        displayed.update(st.session_state.get('ctd_mapping_ai_meta', {}))
        st.session_state['ctd_result'] = displayed
        st.session_state['ctd_result_signature'] = result_signature
    package = st.session_state.get('ctd_result')
    if not package or st.session_state.get('ctd_result_signature') != result_signature:
        return
    proposed = sum(section['status'] == 'proposed' for section in package['sections'])
    st.metric('원문 발췌 후보', f"{proposed}/{len(package['sections'])}절",
              help='선택한 절 가운데 근거 후보를 찾은 비율이며 작성 완료율·법정 제출 가능 판정이 아닙니다.')
    st.caption(f"원문 ID를 선택·확인해 연결한 절 {len(mapping)}/{len(selected)} · "
               f"복수 후보로 남은 절 {len(package.get('ambiguous_sections', []))}/{len(selected)}")
    st.caption(f"본문 작성 모델 요청 {package.get('actual_model_requests', 0)}회 · "
               f"절 연결 모델 요청 {package.get('mapping_model_requests', 0)}회 · 제출용 완성본 아님")
    status_labels = {'proposed': '원문 발췌·확인 대기', 'missing': '근거 없음',
                     'deferred': '근거 후보 있음·검토 보류',
                     'ambiguous': '복수 후보·선택 필요', 'manual_check': '담당자 확인 필요'}
    for section in package['sections']:
        identifier = section['section_id']
        with st.expander(f"{identifier} {section['title']} · {status_labels.get(section['status'], '검토 필요')}"):
            if section.get('draft'):
                st.text(section['draft'])
            if not section.get('evidence'):
                st.caption('연결된 원자료 없음')
            for evidence in section.get('evidence', []):
                st.caption(f"{evidence['source_id']} · {evidence['filename']} · "
                           f"{evidence.get('page') or evidence.get('sheet') or evidence.get('location') or '위치 확인 필요'} "
                           f"· SHA {evidence['document_sha256']}")
                if evidence.get('source_scope') == 'confirmed_substance_link':
                    st.caption('담당자가 원료명과 선택 완제 제품의 관계를 확인한 S 절 자료')
                st.code(evidence['quote'], language=None)
    for key, label in (('missing_sections', '근거를 찾지 못한 절'),
                       ('deferred_sections', '근거 후보가 있으나 자동 작성 보류한 절'),
                       ('ambiguous_sections', '복수 후보로 보류한 절'),
                       ('manual_checks', '담당자가 직접 확인할 사항')):
        entries = package.get(key) or []
        if entries:
            st.warning(label + ': ' + ', '.join(str(item) for item in entries))
    deferred = package.get('deferred_sources') or []
    if deferred:
        st.warning(f'원문 위치·제품 범위·절 연결을 확인하지 못한 자료 {len(deferred)}건이 있습니다.')
        with st.expander('보류된 원문 조각과 사유'):
            st.dataframe(deferred, hide_index=True)
    st.caption('담당자는 원자료의 제품·제형·절·전체 문구, 신청인·서명·동의 및 최신 제출 요건을 확인해야 합니다.')
    try:
        output = export_ctd_working_draft(package)
    except ValueError as exc:
        st.error(str(exc))
        return
    st.download_button('CTD 작업 초안 DOCX', output['document'],
                       file_name='CTD_Module1_3_작업초안.docx', key='ctd_docx')
    st.download_button('출처·누락 검토 JSON', output['evidence'],
                       file_name='CTD_출처와_누락_검토.json', key='ctd_evidence')
    st.download_button('초안과 근거 ZIP', output['zip'],
                       file_name='CTD_Module1_3_작업묶음.zip', key='ctd_zip')
    if template_bytes is not None:
        rendered_sections = {section['section_id']: section for section in package['sections']}
        fillable = (bool(form_ids) and len(form_ids) == len(set(form_ids))
                    and all(identifier in rendered_sections
                            and rendered_sections[identifier]['status'] in {'proposed', 'manual_check'}
                            and rendered_sections[identifier].get('draft')
                            and rendered_sections[identifier].get('evidence')
                            for identifier in form_ids))
        if not fillable:
            st.warning('양식에 있는 절의 근거 누락·상충·보류를 해결해야 기입할 수 있습니다.')
        else:
            st.caption('업로드 양식의 절만 기입합니다. 다른 선택 절의 누락·모호함은 근거 기록에 남습니다.')
        export_signature = _digest({'result': result_signature,
                                    'package_fingerprint': package['fingerprint'],
                                    'template_sha256': template_sha})
        if st.button('업로드한 CTD 양식에 초안 기입', disabled=not fillable,
                     key='ctd_template_fill'):
            st.session_state.pop('ctd_template_export', None)
            try:
                suffix = Path(template_upload.name).suffix.lower()
                with tempfile.TemporaryDirectory(prefix='ctd-template-export-') as directory:
                    source_path = Path(directory) / ('원본' + suffix)
                    target_path = Path(directory) / ('CTD_기입_작업초안' + suffix)
                    source_path.write_bytes(template_bytes)
                    produced = fill_ctd_working_template(
                        source_path, target_path, st.session_state['ctd_engine_result'],
                        sources=intake['sources'], product_name=product,
                        product_variant=variant, section_map=mapping or None,
                        confirmed_substance_links=substance_links or None)
                    filled_bytes = produced['output_path'].read_bytes()
                    sidecar_bytes = produced['evidence_path'].read_bytes()
                sidecar = json.loads(sidecar_bytes)
                if (produced.get('submission_ready') is not False
                        or produced['template_sha256'] != template_sha
                        or produced['output_sha256'] != sha256(filled_bytes).hexdigest()
                        or sidecar.get('package_fingerprint') != package['fingerprint']
                        or sidecar.get('output_sha256') != produced['output_sha256']):
                    raise ValueError('양식 기입 후 파일·출처 지문이 현재 작업과 일치하지 않음')
                st.session_state['ctd_template_export'] = {
                    'signature': export_signature, 'document': filled_bytes,
                    'evidence': sidecar_bytes, 'suffix': suffix}
            except ValueError as exc:
                st.error(str(exc))
            except Exception:
                st.error('CTD 양식 기입·저장 후 검수를 완료하지 못했습니다.')
        filled = st.session_state.get('ctd_template_export')
        if filled and filled['signature'] == export_signature:
            st.download_button('양식에 기입한 CTD 작업 초안', filled['document'],
                               file_name='CTD_기입_작업초안' + filled['suffix'], key='ctd_filled_template')
            st.download_button('기입 문서의 절별 출처 기록', filled['evidence'],
                               file_name='CTD_기입_작업초안' + filled['suffix'] + '.sources.json',
                               key='ctd_filled_evidence')


render = render_ctd
