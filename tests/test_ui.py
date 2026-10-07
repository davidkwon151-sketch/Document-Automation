from copy import deepcopy
from pathlib import Path

import pytest
from streamlit.testing.v1 import AppTest

from agent.pipeline import run_pipeline
from evals.run import MockEvaluationClient

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture
def ui(monkeypatch, tmp_path, native_unavailable):
    import agent.pipeline as pipeline

    case = {"mock_brief": {"목적": "매출", "보고 대상": "팀장", "보고서 유형": "결과보고서", "마감": "", "분량": "1쪽", "부족한 정보": [], "질문": []}}
    text = "매출 120만원을 달성함"
    docs = [{"파일명": "자료.xlsx", "본문": text, "표 목록": [], "페이지/시트 정보": [{"본문": text, "표 목록": [], "페이지": None, "시트": "실적", "위치": "실적!A1"}]}]
    ready = run_pipeline("매출 결과보고서", documents=docs, client=MockEvaluationClient(case))
    monkeypatch.setattr(pipeline, "run_pipeline", lambda *args, **kwargs: deepcopy(ready))
    monkeypatch.setenv("REPORT_AGENT_DATA_DIR", str(tmp_path))
    return AppTest.from_file(str(ROOT / "app" / "ui.py"), default_timeout=15)


def test_ui_initial_state_and_empty_instruction(ui):
    ui.run()
    assert not ui.exception
    assert ui.title[0].value == "문서 표준화 AI AGENT"
    ui.button(key="generate").click().run()
    assert ui.error and "지시" in ui.error[0].value


def test_ui_generation_confirmation_and_download(ui):
    ui.run()
    ui.text_area(key="instruction").set_value("팀장에게 매출 결과보고서를 작성해줘")
    ui.button(key="generate").click().run()
    assert not ui.exception
    assert "draft" in ui.session_state["result"]
    assert ui.button(key="prepare").disabled
    ui.checkbox(key="confirmed").check().run()
    ui.button(key="prepare").click().run()
    assert not ui.exception
    assert set(ui.session_state["exports"]) == {"docx", "hwpx"}
    ui.text_area(key="instruction").set_value("다른 보고서 작성해줘").run()
    assert "result" not in ui.session_state
    assert "exports" not in ui.session_state


def test_ui_manual_edits_revalidate_and_invalidate_old_exports(ui):
    ui.run()
    ui.text_area(key="instruction").set_value("매출 결과보고서")
    ui.button(key="generate").click().run()
    ui.checkbox(key="confirmed").check().run()
    ui.button(key="prepare").click().run()
    ui.text_area(key="draft_body").set_value("□ 신규 매출 999만원임").run()
    assert ui.button(key="prepare").disabled
    submit = next(button for button in ui.button if button.label == "수정 내용 저장·재검수")
    submit.click().run()
    assert not ui.exception
    assert ui.session_state["result"]["review"]["blocking"]
    assert "exports" not in ui.session_state
    assert 'output_verification' not in ui.session_state['result']
    assert ui.warning


def test_ui_invalid_summary_shows_error_instead_of_crashing(ui):
    ui.run()
    ui.text_area(key="instruction").set_value("매출 결과보고서")
    ui.button(key="generate").click().run()
    ui.text_area(key="draft_summary").set_value("한 줄\n두 줄\n세 줄\n네 줄")
    next(button for button in ui.button if button.label == "수정 내용 저장·재검수").click().run()
    assert not ui.exception
    assert ui.error and "3줄" in ui.error[0].value


def test_ui_follow_up_questions_accept_answers(ui, monkeypatch):
    import agent.pipeline as pipeline

    calls = []

    def missing(*args, **kwargs):
        calls.append(kwargs.get("answers"))
        return {"status": "needs_information", "brief": {}, "questions": ["누구에게 보고하나요?", "어떤 유형인가요?"]}

    monkeypatch.setattr(pipeline, "run_pipeline", missing)
    ui.run()
    ui.text_area(key="instruction").set_value("보고서 만들어줘")
    ui.button(key="generate").click().run()
    assert len([item for item in ui.text_input if item.key in {'answer_0', 'answer_1'}]) == 2
    ui.text_input(key="answer_0").set_value("팀장")
    ui.text_input(key="answer_1").set_value("결과보고서")
    ui.button(key="generate").click().run()
    assert calls[-1] == {"누구에게 보고하나요?": "팀장", "어떤 유형인가요?": "결과보고서"}


def test_ui_rejection_regeneration_resume_preserve_first_draft_and_new_report_resets(ui):
    ui.run()
    ui.text_area(key='instruction').set_value('매출 결과보고서')
    ui.button(key='generate').click().run()
    original = deepcopy(ui.session_state['result'])
    ui.text_input(key='rejection_reason').set_value('설명을 보완해 주세요')
    ui.button(key='record_rejection').click().run()
    ui.button(key='record_rework').click().run()
    ui.text_area(key='instruction').set_value('매출 결과보고서 설명을 보완해줘').run()
    assert 'result' not in ui.session_state
    ui.button(key='generate').click().run()
    assert not ui.exception
    result = ui.session_state['result']
    assert result['run_id'] == original['run_id']
    assert result['metrics']['baseline_draft'] == original['metrics']['baseline_draft']
    assert result['metrics']['draft_started_at'] == original['metrics']['draft_started_at']
    assert result['metrics']['rejection_count'] == result['metrics']['rework_count'] == 1
    ui.button(key='resume_report').click().run()
    assert not ui.exception
    assert ui.session_state['result']['metrics']['rejection_count'] == 1
    ui.button(key='new_report').click().run()
    assert 'active_run' not in ui.session_state and 'result' not in ui.session_state
    ui.text_area(key='instruction').set_value('다음 보고서')
    ui.button(key='generate').click().run()
    assert ui.session_state['result']['run_id'] != original['run_id']
    assert ui.session_state['result']['metrics']['rejection_count'] == 0


def test_ui_unknown_form_ai_mapping_must_be_confirmed_then_reuses_without_model(ui, monkeypatch):
    import agent.template_learning as learning
    import hashlib
    from templates import analyze_template
    path = ROOT / 'samples/sample_company_form.docx'
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    profile = analyze_template(path)
    profile['learning'] = {'origin': 'ai_proposal', 'needs_confirmation': True}
    monkeypatch.setattr(learning, 'learn_template', lambda *args, **kwargs: deepcopy(profile))
    ui.session_state['public_template_path'] = str(path)
    ui.run()
    assert not ui.exception
    ui.button(key=f'learn_template_{digest}').click().run()
    assert not ui.exception
    assert ui.button(key='generate').disabled
    assert any('확인하고' in error.value for error in ui.error)
    ui.button(key=f'save_mapping_{digest}').click().run()
    assert not ui.exception
    assert not ui.button(key='generate').disabled
    assert any('재사용' in caption.value for caption in ui.caption)


def test_ui_typed_direct_input_shows_field_error_and_disables_generation(ui, tmp_path, monkeypatch):
    import hashlib
    import os
    from docx import Document
    from templates import analyze_template
    from agent.template_learning import save_learned_profile
    path = tmp_path / 'birthday.docx'
    doc = Document()
    doc.add_paragraph('{{생년월일}}')
    doc.save(path)
    profile = analyze_template(path)
    profile['fields'][0].update(required=False, input_required=True, input_mode='user_provided',
                               max_chars=40, confidence=1.0, validation={'type': 'date'},
                               narrative_style_required=False)
    cache = Path(os.environ['REPORT_AGENT_DATA_DIR']) / 'template_mappings'
    save_learned_profile(path, profile, cache_dir=cache, user_confirmed=True)
    ui.session_state['public_template_path'] = str(path)
    ui.run()
    assert not ui.exception
    ui.text_input(key='form_input_생년월일').set_value('2026-02-30').run()
    assert not ui.exception
    assert ui.button(key='generate').disabled
    assert any('달력' in item.value for item in ui.error)
    ui.text_input(key='form_input_생년월일').set_value('2024-02-29').run()
    assert not ui.exception
    assert not ui.button(key='generate').disabled
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    ui.checkbox(key=f'edit_value_rules_{digest}').check().run()
    assert not ui.exception
    assert any('검사 규칙' in item.label for item in ui.checkbox)


def test_ui_native_choice_is_blank_until_selected_and_partial_row_blocks(ui, tmp_path):
    import hashlib
    import os
    from openpyxl import Workbook
    from openpyxl.worksheet.datavalidation import DataValidation
    from templates import analyze_template
    from agent.template_learning import save_learned_profile
    path = tmp_path / 'choice-row.xlsx'
    workbook = Workbook()
    workbook.active['A1'] = '구분'
    workbook.active['B1'] = ''
    workbook.active['A2'] = '수량'
    workbook.active['B2'] = ''
    rule = DataValidation(type='list', formula1='"EXP,DOM"')
    workbook.active.add_data_validation(rule)
    rule.add('B1')
    workbook.save(path)
    profile = analyze_template(path)
    fields = [field for field in profile['fields'] if field['id'].endswith((':B1', ':B2'))]
    assert len(fields) == 2
    for field in fields:
        field.update(required=False, input_required=True, input_mode='user_provided',
                     max_chars=100, confidence=1.0, narrative_style_required=False)
        if field['id'].endswith(':B2'):
            field['validation'] = {'type': 'integer', 'min': 1}
    profile['fields'] = fields
    profile['constraints'] = {'groups': [{'kind': 'all_or_none', 'fields': ['구분', '수량']}]}
    cache = Path(os.environ['REPORT_AGENT_DATA_DIR']) / 'template_mappings'
    save_learned_profile(path, profile, cache_dir=cache, user_confirmed=True)
    ui.session_state['public_template_path'] = str(path)
    ui.run()
    assert not ui.exception and not ui.button(key='generate').disabled
    assert ui.selectbox(key='form_input_구분').value == ''
    ui.selectbox(key='form_input_구분').select('EXP').run()
    assert not ui.exception and ui.button(key='generate').disabled
    assert any('수량' in item.value for item in ui.error)
    ui.text_input(key='form_input_수량').set_value('2').run()
    assert not ui.exception and not ui.button(key='generate').disabled
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    ui.checkbox(key=f'edit_value_rules_{digest}').check().run()
    assert not ui.exception and not ui.button(key='generate').disabled


def test_ui_invalid_restored_choice_requires_explicit_reset(ui, tmp_path):
    import os
    from docx import Document
    from templates import analyze_template
    from agent.template_learning import save_learned_profile
    path = tmp_path / 'declared-choice.docx'
    doc = Document()
    doc.add_paragraph('{{구분}}')
    doc.save(path)
    profile = analyze_template(path)
    profile['fields'][0].update(required=False, input_required=True, input_mode='user_provided',
        max_chars=100, confidence=1.0, validation={'type': 'choice', 'options': ['국내', '해외']})
    save_learned_profile(path, profile, cache_dir=Path(os.environ['REPORT_AGENT_DATA_DIR']) / 'template_mappings', user_confirmed=True)
    ui.session_state['public_template_path'] = str(path)
    ui.session_state['form_input_구분'] = '이전 잘못된 선택값'
    ui.run()
    assert not ui.exception and ui.button(key='generate').disabled
    assert ui.session_state['form_input_구분'] == '이전 잘못된 선택값'
    ui.button(key='reset_choice_구분').click().run()
    assert not ui.exception and not ui.button(key='generate').disabled
    assert ui.selectbox(key='form_input_구분').value == ''


def test_ui_follow_up_answers_survive_repeated_regeneration_and_merge(ui, monkeypatch):
    import agent.pipeline as pipeline
    ui.run()
    ui.text_area(key='instruction').set_value('매출 결과보고서')
    ui.button(key='generate').click().run()
    ready = deepcopy(ui.session_state['result'])
    calls = []
    def capture(*args, **kwargs):
        calls.append(kwargs['answers'])
        return deepcopy(ready)
    monkeypatch.setattr(pipeline, 'run_pipeline', capture)
    ui.session_state['answers'] = {'보고 대상': '팀장'}
    ui.session_state['result'] = {'status': 'needs_information', 'brief': {}, 'questions': ['마감은 언제인가요?']}
    ui.run()
    ui.text_input(key='answer_0').set_value('금요일')
    ui.button(key='generate').click().run()
    ui.button(key='generate').click().run()
    ui.button(key='generate').click().run()
    assert not ui.exception
    assert calls == [{'보고 대상': '팀장', '마감은 언제인가요?': '금요일'}]*3
    ui.text_area(key='instruction').set_value('매출 결과보고서 설명 보완').run()
    ui.button(key='generate').click().run()
    assert calls[-1] == calls[0]


def test_ui_resumed_report_restores_saved_evidence_and_answers_without_other_template(ui, monkeypatch):
    import os
    import agent.pipeline as pipeline
    from app.storage import save_record
    ui.run()
    ui.text_area(key='instruction').set_value('매출 결과보고서')
    ui.button(key='generate').click().run()
    ready = deepcopy(ui.session_state['result'])
    data_root = Path(os.environ['REPORT_AGENT_DATA_DIR'])
    source = data_root / 'runs' / ready['run_id'] / 'inputs' / '실적.txt'
    source.parent.mkdir(parents=True)
    source.write_text('매출 120만원을 달성함', encoding='utf-8')
    ready['input_paths'] = ['inputs/실적.txt']
    ready['answers'] = {'보고 대상': '팀장', '마감은 언제인가요?': '금요일'}
    save_record(ready, data_root / 'runs' / ready['run_id'])
    ui.session_state['answers'] = {'다른 보고서 질문': '다른 답변'}
    ui.session_state['public_template_path'] = str(ROOT / 'samples/sample_company_form.docx')
    ui.session_state['public_template_provenance'] = {'id': 'unrelated'}
    ui.run()
    ui.button(key='resume_report').click().run()
    assert not ui.exception
    assert 'public_template_path' not in ui.session_state
    assert 'public_template_provenance' not in ui.session_state
    assert ui.session_state['answers'] == ready['answers']
    calls = []
    def capture(instruction, paths, **kwargs):
        calls.append((paths, kwargs['answers']))
        return deepcopy(ready)
    monkeypatch.setattr(pipeline, 'run_pipeline', capture)
    ui.button(key='generate').click().run()
    assert not ui.exception
    assert calls == [([source], ready['answers'])]
    assert ui.session_state['result']['metrics']['baseline_draft'] == ready['metrics']['baseline_draft']
    import streamlit as st
    from types import SimpleNamespace
    original_uploader = st.file_uploader
    extra = SimpleNamespace(name='추가.txt', getvalue=lambda: '매출 설명 자료임'.encode('utf-8'))
    def uploader(*args, **kwargs):
        return [extra] if kwargs.get('key', '').startswith('attachments') else original_uploader(*args, **kwargs)
    monkeypatch.setattr(st, 'file_uploader', uploader)
    ui.run()
    ui.button(key='generate').click().run()
    assert not ui.exception
    assert len(calls[-1][0]) == 2 and calls[-1][0][0] == source
    assert calls[-1][0][1].read_bytes() == extra.getvalue()
    assert calls[-1][1] == ready['answers']


def test_ui_new_report_clears_previously_imported_template(ui):
    ui.session_state['public_template_path'] = str(ROOT / 'samples/sample_company_form.docx')
    ui.session_state['public_template_provenance'] = {'id': 'previous'}
    ui.run()
    ui.text_area(key='instruction').set_value('이전 보고서 내용').run()
    ui.button(key='new_report').click().run()
    assert not ui.exception
    assert 'public_template_path' not in ui.session_state and 'public_template_provenance' not in ui.session_state
    assert ui.text_area(key='instruction').value == ''
    assert ui.session_state['attachment_epoch'] == ui.session_state['custom_template_epoch'] == 1


def test_ui_next_question_batch_never_reuses_answer_to_different_question(ui, monkeypatch):
    import agent.pipeline as pipeline
    calls = []
    def missing(*args, **kwargs):
        calls.append(kwargs['answers'])
        question = '누구에게 보고하나요?' if len(calls) == 1 else '마감은 언제인가요?'
        return {'status': 'needs_information', 'brief': {}, 'questions': [question]}
    monkeypatch.setattr(pipeline, 'run_pipeline', missing)
    ui.run()
    ui.text_area(key='instruction').set_value('보고서 작성해줘')
    ui.button(key='generate').click().run()
    ui.text_input(key='answer_0').set_value('팀장')
    ui.button(key='generate').click().run()
    assert not ui.exception
    assert ui.text_input(key='answer_0').label == '마감은 언제인가요?'
    assert ui.text_input(key='answer_0').value == ''
    ui.text_input(key='answer_0').set_value('금요일')
    ui.button(key='generate').click().run()
    assert calls[-1] == {'누구에게 보고하나요?': '팀장', '마감은 언제인가요?': '금요일'}
