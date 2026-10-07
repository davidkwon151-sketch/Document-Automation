"""DOCX 원본의 export/display·직접 입력·보호·독립 검수 계약."""

from copy import deepcopy
from hashlib import sha256
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
import pytest

from agent.output_check import verify_output
from templates import TemplateError, analyze_template, fill_compatible_template


NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}


def fixture(tmp_path, *, combo=False, inline=False, placeholder=True, locked=False,
            protected=False, bound=False, items=None, unsupported=None):
    doc = Document()
    control = OxmlElement('w:sdt')
    properties = OxmlElement('w:sdtPr')
    alias = OxmlElement('w:alias')
    alias.set(qn('w:val'), '결재 상태')
    properties.append(alias)
    if placeholder:
        properties.append(OxmlElement('w:showingPlcHdr'))
    if locked:
        lock = OxmlElement('w:lock')
        lock.set(qn('w:val'), 'sdtContentLocked')
        properties.append(lock)
    if bound:
        binding = OxmlElement('w:dataBinding')
        binding.set(qn('w:xpath'), '/root/status')
        properties.append(binding)
    choice = OxmlElement('w:' + (unsupported or ('comboBox' if combo else 'dropDownList')))
    if unsupported is None:
        for value, label in (items if items is not None else [('A', '승인 대기'), ('B', '검토 완료')]):
            item = OxmlElement('w:listItem')
            if value is not None:
                item.set(qn('w:value'), value)
            if label is not None:
                item.set(qn('w:displayText'), label)
            choice.append(item)
    properties.append(choice)
    content = OxmlElement('w:sdtContent')
    paragraph = doc.add_paragraph()
    run = paragraph.add_run('항목 선택' if placeholder else '기존 입력')
    run.bold = True
    run.font.name = '맑은 고딕'
    if inline:
        content.append(run._r)
        paragraph.add_run(' / 변경 금지 안내')
        paragraph._p.insert(0, control)
    else:
        content.append(paragraph._p)
        doc._element.body.insert(0, control)
        doc.add_paragraph('변경 금지 안내')
    control.append(properties)
    control.append(content)
    if protected:
        protection = OxmlElement('w:documentProtection')
        protection.set(qn('w:edit'), 'forms')
        protection.set(qn('w:enforcement'), '1')
        doc.settings._element.append(protection)
    source = tmp_path / 'choices.docx'
    doc.save(source)
    return source


def parts(path):
    with ZipFile(path) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def rewrite(path, mutate):
    data = parts(path)
    xml = etree.fromstring(data['word/document.xml'])
    mutate(xml)
    data['word/document.xml'] = etree.tostring(xml, xml_declaration=True, encoding='UTF-8', standalone=True)
    with ZipFile(path, 'w') as result:
        for name, raw in data.items():
            result.writestr(name, raw)


@pytest.mark.parametrize('combo', [False, True])
@pytest.mark.parametrize('inline', [False, True])
def test_native_options_are_export_values_and_declared_direct_input(tmp_path, combo, inline):
    source = fixture(tmp_path, combo=combo, inline=inline)
    field = analyze_template(source)['fields'][0]
    assert field['kind'] == ('docx_combobox' if combo else 'docx_choice')
    assert field['control_type'] == ('combobox' if combo else 'choice')
    assert field['options'] == ['A', 'B']
    assert field['choice_items'] == [{'value':'A','label':'승인 대기'}, {'value':'B','label':'검토 완료'}]
    assert field['allow_custom'] is combo and field['editable_native'] is combo
    assert field['input_required'] is True and field['input_mode'] == 'user_provided'
    assert field['narrative_style_required'] is False


@pytest.mark.parametrize('combo', [False, True])
@pytest.mark.parametrize('inline', [False, True])
def test_export_value_writes_original_display_and_last_value_preserving_parts(tmp_path, combo, inline):
    source = fixture(tmp_path, combo=combo, inline=inline)
    before = parts(source)
    profile = analyze_template(source)
    values = {'결재 상태':'B'}
    output = fill_compatible_template(source, values, tmp_path/'filled.docx', profile=profile)
    assert verify_output(source, output, values, profile=profile)['status'] == 'passed'
    actual = etree.fromstring(parts(output)['word/document.xml'])
    assert actual.xpath('.//w:sdtContent//w:t/text()', namespaces=NS) == ['검토 완료']
    assert actual.xpath('.//w:sdtPr/w:dropDownList/@w:lastValue | .//w:sdtPr/w:comboBox/@w:lastValue', namespaces=NS) == ['B']
    assert not actual.xpath('.//w:showingPlcHdr', namespaces=NS)
    assert actual.xpath('.//w:sdtContent//w:b', namespaces=NS)
    for name, raw in before.items():
        if name != 'word/document.xml':
            assert parts(output)[name] == raw
    assert parts(source) == before


@pytest.mark.parametrize('custom', ['사용자 확인 완료', 'X & <사용자>'])
def test_combo_accepts_custom_one_line_without_adding_native_options(tmp_path, custom):
    source = fixture(tmp_path, combo=True)
    profile = analyze_template(source)
    values = {'결재 상태':custom}
    output = fill_compatible_template(source, values, tmp_path/'custom.docx', profile=profile)
    assert verify_output(source, output, values, profile=profile)['status'] == 'passed'
    actual = etree.fromstring(parts(output)['word/document.xml'])
    assert actual.xpath('.//w:comboBox/w:listItem/@w:value', namespaces=NS) == ['A', 'B']
    assert actual.xpath('.//w:comboBox/@w:lastValue', namespaces=NS) == [custom]


@pytest.mark.parametrize('bad', ['검토 완료', 'C', 'B\n추가'])
def test_closed_dropdown_rejects_display_alias_unknown_or_multiline_before_output(tmp_path, bad):
    source = fixture(tmp_path)
    profile = analyze_template(source)
    output = tmp_path/'previous.docx'
    output.write_bytes(b'PREVIOUS')
    with pytest.raises(TemplateError):
        fill_compatible_template(source, {'결재 상태':bad}, output, profile=profile)
    assert output.read_bytes() == b'PREVIOUS'


@pytest.mark.parametrize('kwargs', [{'locked':True}, {'protected':True}, {'bound':True}, {'placeholder':False}])
def test_locked_protected_bound_or_existing_controls_are_not_fill_targets(tmp_path, kwargs):
    source = fixture(tmp_path, **kwargs)
    before = sha256(source.read_bytes()).hexdigest()
    profile = analyze_template(source)
    assert not profile['fields']
    assert profile['warnings']
    assert sha256(source.read_bytes()).hexdigest() == before


@pytest.mark.parametrize('items', [[('A','동일'), ('A','다름')], [('A','동일'), ('B','동일')], [(None,'값')], [('', '값')], [(' A', '값')], []])
def test_invalid_or_ambiguous_native_options_warn_and_are_excluded(tmp_path, items):
    source = fixture(tmp_path, items=items)
    profile = analyze_template(source)
    assert not profile['fields'] and any('선택' in warning for warning in profile['warnings'])


@pytest.mark.parametrize('unsupported', ['date','picture','group'])
def test_unsupported_structured_controls_are_not_generic_text(tmp_path, unsupported):
    profile = analyze_template(fixture(tmp_path, unsupported=unsupported))
    assert not profile['fields']
    assert any('자동 입력에서 제외' in warning for warning in profile['warnings'])


@pytest.mark.parametrize('mutation', ['display','lastValue','option','alias','guide'])
def test_independent_verifier_rejects_selection_definition_or_guide_corruption(tmp_path, mutation):
    source = fixture(tmp_path, inline=True)
    profile = analyze_template(source)
    values = {'결재 상태':'B'}
    output = fill_compatible_template(source, values, tmp_path/'tampered.docx', profile=profile)
    def corrupt(xml):
        if mutation == 'display':
            xml.xpath('.//w:sdtContent//w:t', namespaces=NS)[0].text = 'B'
        elif mutation == 'lastValue':
            xml.xpath('.//w:dropDownList', namespaces=NS)[0].set(qn('w:lastValue'), 'A')
        elif mutation == 'option':
            xml.xpath('.//w:listItem', namespaces=NS)[1].set(qn('w:displayText'), '변조')
        elif mutation == 'alias':
            xml.xpath('.//w:alias', namespaces=NS)[0].set(qn('w:val'), '변조')
        else:
            xml.xpath('.//w:t[not(ancestor::w:sdt)]', namespaces=NS)[0].text = '안내 변조'
    rewrite(output, corrupt)
    with pytest.raises(ValueError):
        verify_output(source, output, values, profile=profile)


def test_forged_generic_kind_cannot_bypass_native_closed_choice(tmp_path):
    source = fixture(tmp_path)
    profile = analyze_template(source)
    forged = deepcopy(profile)
    forged['fields'][0]['kind'] = 'docx_sdt'
    forged['fields'][0].pop('control_type')
    with pytest.raises(TemplateError):
        fill_compatible_template(source, {'결재 상태':'임의 입력'}, tmp_path/'forged.docx', profile=forged)


@pytest.mark.parametrize('combo', [False, True])
@pytest.mark.parametrize('wrong_kind', ['docx_sdt', 'docx_cell'])
def test_independent_verifier_reads_native_source_kind_even_when_profile_claims_generic(tmp_path, combo, wrong_kind):
    source = fixture(tmp_path, combo=combo)
    before = source.read_bytes()
    forged = deepcopy(analyze_template(source))
    forged['fields'][0]['kind'] = wrong_kind
    forged['fields'][0].pop('control_type')
    forged['fields'][0].pop('validation', None)
    output = tmp_path/'manually_forged.docx'
    output.write_bytes(before)
    def manually_write(xml):
        xml.xpath('.//w:sdtContent//w:t', namespaces=NS)[0].text = 'C'
        for header in xml.xpath('.//w:sdtPr/w:showingPlcHdr', namespaces=NS):
            header.getparent().remove(header)
    rewrite(output, manually_write)  # 채우기 함수의 차단에 의존하지 않고 출력 파일을 직접 구성함.
    with pytest.raises(ValueError, match='원본 선택 컨트롤'):
        verify_output(source, output, {'결재 상태':'C'}, profile=forged)
    assert source.read_bytes() == before


def test_independent_verifier_rejects_a_generic_parent_paragraph_covering_native_control(tmp_path):
    source = fixture(tmp_path, inline=True)
    profile = analyze_template(source)
    original = etree.fromstring(parts(source)['word/document.xml'])
    content = original.xpath('.//w:sdtContent', namespaces=NS)[0]
    parent = content.xpath('ancestor::w:p', namespaces=NS)[0]
    forged = deepcopy(profile)
    forged['fields'][0].update(kind='docx_paragraph', id='docx:word/document.xml:'+original.getroottree().getpath(parent))
    forged['fields'][0].pop('control_type')
    output = tmp_path/'parent_forged.docx'
    output.write_bytes(source.read_bytes())
    rewrite(output, lambda xml: setattr(xml.xpath('.//w:sdtContent//w:t', namespaces=NS)[0], 'text', 'C'))
    with pytest.raises(ValueError, match='원본 선택 컨트롤 밖'):
        verify_output(source, output, {'결재 상태':'C'}, profile=forged)


def test_combo_without_suggestions_is_editable_and_not_a_closed_empty_list(tmp_path):
    source = fixture(tmp_path, combo=True, items=[])
    profile = analyze_template(source)
    assert profile['fields'][0]['allow_custom'] is True
    assert profile['fields'][0]['options'] == []
    values = {'결재 상태':'직접 확인'}
    output = fill_compatible_template(source, values, tmp_path/'editable.docx', profile=profile)
    assert verify_output(source, output, values, profile=profile)['status'] == 'passed'


@pytest.mark.parametrize('combo,locked', [(False,False), (True,False), (False,True)])
def test_mixed_placeholders_do_not_reinterpret_native_or_locked_control_text(tmp_path, combo, locked):
    source = fixture(tmp_path, combo=combo, locked=locked,
                     items=[('A','{{제목}}'), ('B','그대로')])
    rewrite(source, lambda xml: setattr(xml.xpath('.//w:sdtContent//w:t', namespaces=NS)[0], 'text', '{{선택 안내}}'))
    doc = Document(source)
    doc.add_paragraph('{{제목}}')
    doc.save(source)
    profile = analyze_template(source)
    assert not any(field['label']=='선택 안내' for field in profile['fields'])
    values = {'제목':'사용자 제목'}
    if not locked:
        values['결재 상태'] = 'A' if not combo else '{{직접입력}}'
    output = fill_compatible_template(source, values, tmp_path/'mixed.docx', profile=profile)
    assert verify_output(source, output, values, profile=profile)['status'] == 'passed'
    actual = etree.fromstring(parts(output)['word/document.xml'])
    expected = '{{선택 안내}}' if locked else ('{{직접입력}}' if combo else '{{제목}}')
    assert actual.xpath('.//w:sdtContent//w:t/text()', namespaces=NS) == [expected]
    assert '사용자 제목' in ''.join(actual.xpath('.//w:t/text()', namespaces=NS))
