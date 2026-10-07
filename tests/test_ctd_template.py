"""Synthetic CTD working forms: original layout and evidence stay separate."""

from hashlib import sha256
import json
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from lxml import etree
import pytest

from agent.ctd import prepare_ctd_package
from agent.ctd_template import fill_ctd_working_template
from agent.multimodal_intake import collect_multimodal, generation_sources
from parsers import parse_file
from templates.fill import PLACEHOLDER, _segments, _set_segment


ROOT = Path(__file__).resolve().parents[1]


def _inputs(tmp_path, *, value='98.7%', token='[S1]'):
    path = tmp_path / 'coa.txt'
    path.write_text('제품명: ALPHA정\n3.2.P.5.4: 배치 분석\n'
                    '배치번호: B-001\n'
                    f'시험결과: {value}\n시험항목 코드: {token}\n', encoding='utf-8')
    sources = generation_sources(collect_multimodal([path]))
    package = prepare_ctd_package(sources, product_name='ALPHA정',
                                  selected_sections=['3.2.P.5.4'])
    assert package['sections'][0]['status'] == 'proposed'
    return sources, package


def _docx(path, token='{{3.2.P.5.4}}'):
    document = Document()
    document.add_paragraph('검토용 품질 자료')
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = '배치 분석'
    paragraph = table.cell(0, 1).paragraphs[0]
    run = paragraph.add_run(token)
    run.bold = True
    run.font.name = '맑은 고딕'
    document.save(path)
    return path


def _hwpx(path):
    """Adapt the repository's valid HWPX sample to one unique CTD token."""
    original = ROOT / 'templates/result_report.hwpx'
    inserted = False
    with ZipFile(original) as before, ZipFile(path, 'w') as after:
        for info in before.infolist():
            data = before.read(info.filename)
            if info.filename == 'Contents/section0.xml':
                root = etree.fromstring(data)
                for paragraph in root.xpath(".//*[local-name()='p']"):
                    segments = _segments(paragraph, 'hwpx')
                    if not segments:
                        continue
                    text = ''.join(segment.value for segment in segments)
                    if not PLACEHOLDER.search(text):
                        continue
                    replacement = '{{3.2.P.5.4}}' if not inserted else '검토용 양식'
                    inserted = True
                    _set_segment(segments[0], replacement, 'hwpx')
                    for segment in segments[1:]:
                        _set_segment(segment, '', 'hwpx')
                data = etree.tostring(root, encoding='UTF-8', xml_declaration=True)
            after.writestr(info, data)
    assert inserted
    return path


@pytest.mark.parametrize('suffix', ['docx', 'hwpx'])
def test_ctd_template_prints_exact_quotes_and_keeps_citations_in_sidecar(tmp_path, suffix):
    sources, package = _inputs(tmp_path)
    source = (_docx(tmp_path / 'working.docx') if suffix == 'docx'
              else _hwpx(tmp_path / 'working.hwpx'))
    source_sha = sha256(source.read_bytes()).hexdigest()
    result = fill_ctd_working_template(source, tmp_path / f'filled.{suffix}', package,
                                       sources=sources, product_name='ALPHA정')
    assert result['output_check']['status'] == 'passed'
    assert result['template_sha256'] == source_sha
    assert sha256(source.read_bytes()).hexdigest() == source_sha
    assert sha256(result['output_path'].read_bytes()).hexdigest() == result['output_sha256']
    extracted = parse_file(result['output_path'])
    assert '98.7%' in extracted['본문'] and '[S1]' in extracted['본문']
    assert '{{3.2.P.5.4}}' not in extracted['본문']
    assert '[S' + package['sections'][0]['evidence'][0]['source_id'][1:] + ']' not in extracted['본문']
    sidecar = json.loads(result['evidence_path'].read_text(encoding='utf-8'))
    assert sidecar['template_sha256'] == source_sha
    assert sidecar['output_sha256'] == result['output_sha256']
    assert sidecar['package_fingerprint'] == package['fingerprint']
    assert sidecar['evidence']['3.2.P.5.4'][0]['document_sha256'] == sources[0]['document_sha256']
    assert not sidecar['submission_ready'] and not result['submission_ready']
    if suffix == 'docx':
        assert Document(result['output_path']).tables[0].cell(0, 1).paragraphs[0].runs[0].bold


def test_ctd_export_blocks_missing_stale_extra_and_duplicate_sections(tmp_path):
    sources, package = _inputs(tmp_path)
    source = _docx(tmp_path / 'template.docx')
    target = tmp_path / 'out.docx'
    changed_sources, _ = _inputs(tmp_path, value='99.1%')
    with pytest.raises(ValueError, match='바뀜'):
        fill_ctd_working_template(source, target, package,
                                  sources=changed_sources, product_name='ALPHA정')
    assert not target.exists()
    missing = prepare_ctd_package(sources, product_name='ALPHA정',
                                  selected_sections=['3.2.P.1'])
    with pytest.raises(ValueError, match='누락|상충'):
        fill_ctd_working_template(_docx(tmp_path / 'missing.docx', '{{3.2.P.1}}'),
                                  target, missing, sources=sources, product_name='ALPHA정')
    assert not target.exists()
    with pytest.raises(ValueError, match='자리표시자'):
        fill_ctd_working_template(_docx(tmp_path / 'wrong.docx', '{{3.2.P.1}}'),
                                  target, package, sources=sources, product_name='ALPHA정')
    duplicate = Document()
    duplicate.add_paragraph('{{3.2.P.5.4}}')
    duplicate.add_paragraph('{{3.2.P.5.4}}')
    duplicate_path = tmp_path / 'duplicate.docx'
    duplicate.save(duplicate_path)
    with pytest.raises(ValueError, match='자리표시자'):
        fill_ctd_working_template(duplicate_path, target, package,
                                  sources=sources, product_name='ALPHA정')
    assert not target.exists()


def test_ctd_template_renders_its_selected_subset_and_records_other_gaps(tmp_path):
    sources, _ = _inputs(tmp_path)
    package = prepare_ctd_package(sources, product_name='ALPHA정',
                                  selected_sections=['3.2.P.5.4', '3.2.P.1'])
    assert package['missing_sections'] == ['3.2.P.1']
    template = _docx(tmp_path / 'subset.docx')
    result = fill_ctd_working_template(template, tmp_path / 'subset_filled.docx', package,
                                       sources=sources, product_name='ALPHA정')
    assert result['output_check']['status'] == 'passed'
    sidecar = json.loads(result['evidence_path'].read_text(encoding='utf-8'))
    assert sidecar['selected_sections'] == ['3.2.P.5.4', '3.2.P.1']
    assert sidecar['rendered_sections'] == ['3.2.P.5.4']
    assert sidecar['missing_sections'] == ['3.2.P.1']
    assert not sidecar['submission_ready']


def test_ctd_export_blocks_unowned_form_inputs_and_preserves_existing_output(tmp_path):
    sources, package = _inputs(tmp_path)
    document = Document()
    document.add_paragraph('{{3.2.P.5.4}}')
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = '신청인'
    path = tmp_path / 'direct.docx'
    document.save(path)
    target = tmp_path / 'out.docx'
    with pytest.raises(ValueError, match='절 자리표시자 외'):
        fill_ctd_working_template(path, target, package,
                                  sources=sources, product_name='ALPHA정')
    assert not target.exists()
    normal = _docx(tmp_path / 'normal.docx')
    target.write_bytes(b'KEEP')
    with pytest.raises(ValueError, match='덮어쓰지'):
        fill_ctd_working_template(normal, target, package,
                                  sources=sources, product_name='ALPHA정')
    assert target.read_bytes() == b'KEEP'


def test_ctd_export_rejects_unverified_mapping_origin_metadata(tmp_path):
    sources, package = _inputs(tmp_path)
    source = _docx(tmp_path / 'template.docx')
    with_metadata = dict(package, mapping_model_requests=1,
                         mapping_origin='claude_mcp_client_supplied')
    with pytest.raises(ValueError, match='바뀜'):
        fill_ctd_working_template(source, tmp_path / 'out.docx', with_metadata,
                                  sources=sources, product_name='ALPHA정')
    with pytest.raises(ValueError, match='바뀜'):
        fill_ctd_working_template(source, tmp_path / 'other.docx',
                                  dict(package, unverified_claim='approved'),
                                  sources=sources, product_name='ALPHA정')


def test_ctd_export_does_not_remove_another_writer_sidecar(tmp_path, monkeypatch):
    sources, package = _inputs(tmp_path)
    source = _docx(tmp_path / 'template.docx')
    target = tmp_path / 'filled.docx'
    sidecar = Path(str(target) + '.sources.json')
    original_open = Path.open

    def simultaneous_create(path, mode='r', *args, **kwargs):
        if path == sidecar and mode == 'x':
            sidecar.write_text('OTHER WRITER', encoding='utf-8')
        return original_open(path, mode, *args, **kwargs)

    monkeypatch.setattr(Path, 'open', simultaneous_create)
    with pytest.raises(FileExistsError):
        fill_ctd_working_template(source, target, package,
                                  sources=sources, product_name='ALPHA정')
    assert sidecar.read_text(encoding='utf-8') == 'OTHER WRITER'
    assert not target.exists()
