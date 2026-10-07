"""Position-bound repeat profiles, staged-file binding and exact neighbours."""
from copy import deepcopy
from hashlib import sha256
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from lxml import etree
import pytest

from agent.output_check import verify_output
from app.form_config import configure_profile, mapping_rows
from templates import fill_compatible_template
from templates.repeat_fields import repeat_profile
from templates.repeat_rows import (inspect_repeat_tables, prepare_repeat_template,
                                   expansion_binding, verify_prepared_template)


W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
NS = {'w': W}


def fixture(tmp_path):
    original = tmp_path / 'original.docx'; prepared = tmp_path / 'prepared.docx'
    document = Document(); document.add_paragraph('문서: {{제목}}')
    table = document.add_table(rows=3, cols=3); table.style = 'Table Grid'
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for cell, text in zip(table.rows[0].cells, ['품목', '수량', '금액']): cell.text = text
    for cell, text in zip(table.rows[1].cells, ['품목 {{품목}} / 보충 {{보충}}', '{{수량}} kg', '{{금액}} 원']): cell.text = text
    for cell, text in zip(table.rows[2].cells, ['합계', '원본 안내 유지', '{{합계금액}}']): cell.text = text
    document.save(original)
    table_record = inspect_repeat_tables(original)[0]
    plan = {'table_id': table_record['id'], 'row': 2, 'count': 2}
    expansion = prepare_repeat_template(original, prepared, plan)
    profile = repeat_profile(prepared, plan, expansion_binding(expansion))
    values = {}
    for field in profile['fields']:
        row = field.get('repeat_info', {}).get('row')
        key = field.get('placeholder_key')
        values[field['value_key']] = ('2' if key == '수량' else '1000' if key == '금액'
                                    else f'ITEM-{row}' if key == '품목' else f'DETAIL-{row}' if key == '보충'
                                    else '2000' if field['value_key'] == '합계금액' else 'TEST REPORT')
    return original, prepared, expansion, profile, values


def repeated(profile):
    return [field for field in profile['fields'] if field['kind'] == 'docx_repeat_placeholder']


def raw_change(path, callback):
    with ZipFile(path) as archive: parts = {name: archive.read(name) for name in archive.namelist()}
    root = etree.fromstring(parts['word/document.xml']); callback(root)
    parts['word/document.xml'] = etree.tostring(root)
    with ZipFile(path, 'w') as archive:
        for name, data in parts.items(): archive.writestr(name, data)


def test_exact_multislot_values_and_printed_units_preserve_all_inputs(tmp_path):
    original, prepared, expansion, profile, values = fixture(tmp_path)
    before = {path: path.read_bytes() for path in [original, prepared]}
    output = tmp_path / 'result.docx'
    assert verify_prepared_template(prepared, profile, expansion)['status'] == 'passed'
    fill_compatible_template(prepared, values, output, profile=profile)
    assert verify_output(prepared, output, values, profile=profile)['status'] == 'passed'
    table = Document(output).tables[0]
    assert table.cell(1, 0).text == '품목 ITEM-1 / 보충 DETAIL-1'
    assert table.cell(2, 0).text == '품목 ITEM-2 / 보충 DETAIL-2'
    assert table.cell(1, 1).text == table.cell(2, 1).text == '2 kg'
    assert table.cell(1, 2).text == table.cell(2, 2).text == '1000 원'
    assert all(path.read_bytes() == data for path, data in before.items())


@pytest.mark.parametrize('attack', ['drop_physical_fields', 'global_override', 'retarget_outside', 'duplicate_slot', 'wrong_offset', 'weaken_required', 'erase_groups', 'collapse_row_keys'])
def test_repeat_profile_cannot_weaken_authoritative_physical_scope(tmp_path, attack):
    _, prepared, expansion, profile, values = fixture(tmp_path)
    bad = deepcopy(profile); fields = repeated(bad)
    if attack == 'drop_physical_fields':
        bad['fields'].remove(fields[0]); bad['constraints']['groups'] = []
    elif attack == 'global_override':
        bad['fields'] = [field for field in bad['fields'] if field['kind'] != 'docx_repeat_placeholder']
        bad['fields'].extend({'id': 'placeholder:' + key, 'kind': 'placeholder', 'label': key,
                              'value_key': key, 'required': True, 'input_required': False} for key in ['품목', '보충', '수량', '금액'])
        bad['constraints']['groups'] = []
        values.update({key: 'GLOBAL' for key in ['품목', '보충', '수량', '금액']})
    elif attack == 'retarget_outside':
        field = fields[0]
        with ZipFile(prepared) as archive: root = etree.fromstring(archive.read('word/document.xml'))
        paragraph = root.find('w:body/w:p', NS)
        text = ''.join(paragraph.xpath('.//w:t/text()', namespaces=NS))
        path = root.getroottree().getpath(paragraph); start = text.index('{{제목}}')
        field.update(id=f'repeat_docx:word/document.xml:{path}#slot:{start}', xml_path=path,
                     anchor_text=text, placeholder_start=start, placeholder_end=start + len('{{제목}}'), placeholder_key='제목')
    elif attack == 'duplicate_slot':
        duplicate = deepcopy(fields[0]); duplicate['value_key'] = 'duplicate'
        bad['fields'].append(duplicate); values['duplicate'] = 'OTHER'
    elif attack == 'wrong_offset':
        field = fields[0]; field['placeholder_start'] += 1
        field['id'] = f"repeat_docx:word/document.xml:{field['xml_path']}#slot:{field['placeholder_start']}"
    elif attack == 'weaken_required': fields[0]['required'] = False
    elif attack == 'erase_groups': bad['constraints']['groups'] = []
    else: fields[1]['value_key'] = fields[0]['value_key']
    with pytest.raises(ValueError): verify_prepared_template(prepared, bad, expansion)
    output = tmp_path / 'blocked.docx'; output.write_bytes(b'prior output remains')
    with pytest.raises(ValueError): fill_compatible_template(prepared, values, output, profile=bad)
    assert output.read_bytes() == b'prior output remains'


def test_confirmed_logical_key_rename_rebinds_groups_without_changing_slots(tmp_path):
    _, prepared, expansion, profile, values = fixture(tmp_path)
    rows = mapping_rows(profile)
    for index, row in enumerate(rows): row['채울 값'] = f'사용자 항목 {index + 1}'
    configured, mapping = configure_profile(profile, rows)
    renamed = {mapping[field['id']]: values[field['value_key']] for field in profile['fields']}
    assert verify_prepared_template(prepared, configured, expansion)['status'] == 'passed'
    output = tmp_path / 'renamed.docx'
    fill_compatible_template(prepared, renamed, output, mapping=mapping, profile=configured)
    assert verify_output(prepared, output, renamed, mapping=mapping, profile=configured)['status'] == 'passed'


@pytest.mark.parametrize('target', ['original', 'prepared'])
def test_changed_stage_files_are_not_reused(tmp_path, target):
    original, prepared, expansion, profile, _ = fixture(tmp_path)
    changed = original if target == 'original' else prepared
    raw_change(changed, lambda root: root.find('.//w:t', NS).__setattr__('text', '문서 변조'))
    if target == 'prepared':
        profile['source_sha256'] = sha256(prepared.read_bytes()).hexdigest()
    with pytest.raises(ValueError): verify_prepared_template(prepared, profile, expansion)


@pytest.mark.parametrize('attack', ['unit', 'row_value', 'outside'])
def test_independent_final_check_rejects_changed_neighbours_and_row_values(tmp_path, attack):
    _, prepared, _, profile, values = fixture(tmp_path)
    output = tmp_path / 'result.docx'; fill_compatible_template(prepared, values, output, profile=profile)
    def change(root):
        if attack == 'unit':
            text = root.findall('.//w:tbl/w:tr', NS)[1].findall('w:tc', NS)[1].find('.//w:t', NS)
            text.text = text.text.replace('kg', 'mL')
        elif attack == 'row_value':
            text = root.findall('.//w:tbl/w:tr', NS)[1].find('.//w:t', NS)
            text.text = text.text.replace('ITEM-1', 'ITEM-2')
        else: root.find('.//w:t', NS).text = '다른 문서'
    raw_change(output, change)
    with pytest.raises(ValueError): verify_output(prepared, output, values, profile=profile)
