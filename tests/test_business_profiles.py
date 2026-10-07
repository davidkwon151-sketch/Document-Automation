"""실제 공개 사업 양식의 지정 빈칸과 원본 안내 보존을 확인함."""

from copy import deepcopy
from hashlib import sha256
import json
from pathlib import Path
from zipfile import ZipFile

from lxml import etree
from openpyxl import load_workbook
import pytest

from agent.output_check import _same_position, verify_output
from templates import TemplateError, fill_compatible_template, load_form_profile

ROOT = Path(__file__).resolve().parents[1]
IDS = ('business_startup_2026', 'business_kotra_digital_2026', 'business_kotra_salesforce_2026')


def _profile(entry_id):
    return json.loads((ROOT / f'templates/profiles/{entry_id}.json').read_text(encoding='utf-8'))


@pytest.mark.parametrize('entry_id', IDS)
def test_business_profiles_define_version_scope_and_safe_input_modes(entry_id):
    profile = _profile(entry_id)
    assert profile['domain'] == 'business_support'
    assert profile['business_workflow'] in {'trade_sales', 'overseas_business', 'government_grant', 'rd_project'}
    assert profile['resource_kind'] == 'blank_form'
    assert profile['full_form_filled'] is False
    assert profile['submission_ready'] is False
    assert profile['attachments_review_required'] is True
    assert '실제 제출용' in profile['demo_notice']
    assert len(profile['source_sha256']) == 64
    assert profile['source_url'].startswith('https://')
    assert profile['form_version']
    for field in profile['fields']:
        assert field['input_required'] == (field['input_mode'] == 'user_provided')
        assert field['narrative_style_required'] is False
        assert field['value_key'] in profile['demo_values']
    if 'kotra' in entry_id:
        assert profile['deadline_status'] == 'expired'
        assert profile['application_deadline']
        assert any('종료' in warning for warning in profile['warnings'])


@pytest.mark.parametrize('entry_id', IDS)
def test_business_official_blank_fill_preserves_original_examples_assets_and_values(entry_id, tmp_path):
    profile = _profile(entry_id)
    source = ROOT / 'data/public_templates/business' / profile['source_filename']
    if not source.is_file():
        pytest.skip('공식 공개 원본은 별도 corpus로 내려받음')
    assert load_form_profile(source, entry_id) == profile
    before = source.read_bytes()
    assert sha256(before).hexdigest() == profile['source_sha256']
    output = fill_compatible_template(source, profile['demo_values'], tmp_path / source.name, profile=profile)
    assert verify_output(source, output, profile['demo_values'], profile=profile)['status'] == 'passed'
    assert source.read_bytes() == before
    with ZipFile(source) as original, ZipFile(output) as result:
        for part in original.namelist():
            if part == 'word/document.xml' or part.startswith('xl/worksheets/'):
                continue
            assert original.read(part) == result.read(part), part
        if profile['format'] == 'docx':
            original_xml = etree.fromstring(original.read('word/document.xml'))
            filled_xml = etree.fromstring(result.read('word/document.xml'))
            ns = {key:value for key,value in original_xml.nsmap.items() if key}
            # 한 줄이라도 원본 안내·예시가 있으면 동일 위치에서 정확히 보존함.
            for paragraph in original_xml.xpath('//w:p', namespaces=ns):
                text = ''.join(paragraph.xpath('.//w:t/text()', namespaces=ns))
                if text:
                    same = _same_position(filled_xml, paragraph)
                    assert same is not None
                    assert ''.join(same.xpath('.//w:t/text()', namespaces=ns)) == text
    if profile['format'] == 'xlsx':
        original, filled = load_workbook(source).active, load_workbook(output).active
        assert [c.value for c in original[3]] == [c.value for c in filled[3]]
        assert [c.value for c in original[2]] == [c.value for c in filled[2]]
        assert filled['I4'].value == '0000'  # 식별 문자열의 선행 0을 지움 없이 저장함.
        assert filled['R4'].value == '0000-00'
        assert list(original.merged_cells.ranges) == list(filled.merged_cells.ranges)


def test_business_profile_blocks_oversized_value_before_output(tmp_path):
    profile = _profile('business_kotra_salesforce_2026')
    source = ROOT / 'data/public_templates/business' / profile['source_filename']
    if not source.is_file():
        pytest.skip('공식 원본 별도 corpus')
    values = deepcopy(profile['demo_values'])
    field = profile['fields'][0]
    values[field['value_key']] = '긴' * (field['max_chars'] + 1)
    target = tmp_path / 'bad.xlsx'
    with pytest.raises(TemplateError, match='길이|글자|초과'):
        fill_compatible_template(source, values, target, profile=profile)
    assert not target.exists()
