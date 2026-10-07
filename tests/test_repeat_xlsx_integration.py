"""End-to-end repeated cost inputs, source policy, native controls and export proof."""
from copy import deepcopy
from hashlib import sha256
import json
from pathlib import Path
from zipfile import ZipFile

from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
import pytest
from streamlit.testing.v1 import AppTest

from agent.output_check import verify_output
from agent.template_learning import learn_template, load_learned_profile, save_learned_profile
from templates import analyze_template, fill_compatible_template
from templates.repeat_fields import repeat_profile
from templates.repeat_rows import (expansion_binding, inspect_repeat_tables, prepare_repeat_template,
                                   validate_repeat_fields, verify_prepared_template)
from templates.repeat_rules import inherit_repeat_rules
from templates.value_rules import mapped_rule_profile


NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
PART = 'xl/worksheets/sheet1.xml'


def cost_form(path):
    """Synthetic RA cost schedule; values are test inputs, not drug evidence."""
    book = Workbook()
    sheet = book.active
    sheet.title = '시험내역'
    sheet.merge_cells('A1:E1')
    sheet['A1'] = 'RA 시험 비용 산출서 (합성 검증 양식)'
    sheet['A1'].font = Font(name='맑은 고딕', size=15, bold=True)
    sheet.row_dimensions[1].height = 30
    headers = ['시험명', '수량(건)', '단가(원)', '금액(원)', '검토 상태']
    for column, text in enumerate(headers, 1):
        cell = sheet.cell(3, column, text)
        cell.font = Font(name='맑은 고딕', size=11, bold=True)
        cell.fill = PatternFill('solid', fgColor='D9E2F3')
        cell.alignment = Alignment(horizontal='center')
    for column in range(1, 6):
        cell = sheet.cell(4, column)
        cell.font = Font(name='맑은 고딕', size=11)
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        cell.border = Border(bottom=Side(style='thin', color='B7B7B7'))
    sheet['A4'] = '{{시험명}}'
    sheet['D4'] = '=B4*C4'
    for coordinate in ['B4', 'C4', 'D4', 'D5']:
        sheet[coordinate].number_format = '#,##0'
    choice = DataValidation(type='list', formula1='"확인,대기"', allow_blank=True)
    choice.errorTitle, choice.error, choice.showErrorMessage = '입력 확인', '목록의 값만 선택함', True
    sheet.add_data_validation(choice)
    choice.add('E4')
    sheet['A5'], sheet['D5'] = '합계', '=SUM(D4:D4)'
    sheet['D5'].font = Font(name='맑은 고딕', size=11, bold=True)
    for letter, width in [('A', 32), ('B', 12), ('C', 17), ('D', 19), ('E', 17)]:
        sheet.column_dimensions[letter].width = width
    sheet.row_dimensions[4].height = 30
    sheet.print_area = 'A1:E5'
    sheet.print_title_rows = '1:3'
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
    summary = book.create_sheet('요약')
    summary['A1'], summary['B1'] = '총 시험 비용(원)', "='시험내역'!$D$5"
    summary['A3'], summary['B3'] = '내역 재집계(원)', "=SUM('시험내역'!D4:D4)"
    summary.column_dimensions['A'].width, summary.column_dimensions['B'].width = 24, 24
    for row in [1, 3]:
        summary.cell(row, 1).font = Font(name='맑은 고딕', size=11)
        summary.cell(row, 2).number_format = '#,##0'
    summary.print_area = 'A1:B3'
    book.save(path)
    book.close()
    return path


def cost_policy(source):
    original = analyze_template(source)
    # This confirmed quantity slot is physically blank; the generic adjacent-label
    # heuristic intentionally does not treat the preceding token as its label.
    original['fields'].append({'id': 'xlsx:' + PART + ':B4', 'kind': 'xlsx_cell',
        'label': '수량(건)', 'value_key': '수량(건)', 'required': True,
        'input_required': False, 'confidence': 1.0})
    for field in original['fields']:
        field.update(max_chars=1000, input_mode='user_provided' if field['input_required'] else 'source_grounded')
        coordinate = field['id'].split(':')[2]
        if coordinate in {'B4', 'C4'}:
            field['validation'] = {'type': 'integer', 'min': 0}
            field.update(required=True, input_required=False, input_mode='source_grounded', max_chars=15)
    return original


def prepared_costs(tmp_path, count=3):
    source = cost_form(tmp_path / 'source.xlsx')
    original = cost_policy(source)
    plan = {'table_id': 'xlsx:' + PART, 'row': 4, 'count': count}
    prepared = tmp_path / 'prepared.xlsx'
    expansion = prepare_repeat_template(source, prepared, plan)
    raw = repeat_profile(prepared, plan, expansion_binding(expansion))
    profile = inherit_repeat_rules(source, raw, original, prepared_path=prepared)
    return source, prepared, expansion, original, profile


def cost_values(profile):
    result = {}
    for field in profile['fields']:
        if not field.get('repeat_info'):
            continue
        row, column = field['repeat_info']['row'], field['repeat_info']['column']
        result[field['value_key']] = {1: f'시험 {row}', 2: str(row + 1),
                                      3: str(row * 1000), 5: '확인'}[column]
    return result


def test_repeated_numeric_inputs_keep_native_formulas_lists_styles_and_source(tmp_path):
    source, prepared, expansion, _, profile = prepared_costs(tmp_path)
    original_raw = source.read_bytes()
    fields = [field for field in profile['fields'] if field.get('repeat_info')]
    assert len(fields) == 12 and len({field['value_key'] for field in fields}) == 12
    assert all(field['repeat_info']['column'] != 4 for field in fields)
    values = cost_values(profile)
    target = fill_compatible_template(prepared, values, tmp_path / 'filled.xlsx', profile=profile)
    assert verify_output(prepared, target, values, profile=profile)['status'] == 'passed'
    assert verify_prepared_template(prepared, profile, expansion)['status'] == 'passed'
    with ZipFile(target) as archive:
        root = etree.fromstring(archive.read(PART))
        for row in range(4, 7):
            quantity = root.xpath(f'//s:c[@r="B{row}"]', namespaces=NS)[0]
            assert quantity.get('t') == 'n'
            assert quantity.find('s:v', NS).text == str(row - 2)
    book = load_workbook(target, data_only=False)
    original = load_workbook(source, data_only=False)
    try:
        for row in range(4, 7):
            assert book['시험내역'][f'D{row}'].value == f'=B{row}*C{row}'
            assert book['시험내역'][f'B{row}']._style == original['시험내역']['B4']._style
        assert book['시험내역']['D7'].value == '=SUM(D4:D6)'
        assert book['요약']['B1'].value == "='시험내역'!$D$7"
        assert book['요약']['B3'].value == "=SUM('시험내역'!D4:D6)"
        assert 'E4' in book['시험내역'].data_validations.dataValidation[0].sqref
        assert 'E6' in book['시험내역'].data_validations.dataValidation[0].sqref
    finally:
        book.close()
        original.close()
    assert source.read_bytes() == original_raw


@pytest.mark.parametrize('change', ['invalid_choice', 'fractional_quantity', 'relaxed_native', 'removed_required'])
def test_fill_rejects_weakened_or_invalid_repeated_inputs(tmp_path, change):
    _, prepared, _, _, profile = prepared_costs(tmp_path)
    values = cost_values(profile)
    repeated = [field for field in profile['fields'] if field.get('repeat_info')]
    if change == 'invalid_choice':
        values[next(field['value_key'] for field in repeated if field['repeat_info']['column'] == 5)] = '임의 승인'
    elif change == 'fractional_quantity':
        values[next(field['value_key'] for field in repeated if field['repeat_info']['column'] == 2)] = '2.5'
    elif change == 'relaxed_native':
        field = next(field for field in repeated if field['repeat_info']['column'] == 5)
        field['xlsx_list']['sqref'] = 'E1:E100'
    else:
        profile['fields'].remove(next(field for field in repeated if field['repeat_info']['column'] == 2))
    with pytest.raises(ValueError):
        fill_compatible_template(prepared, values, tmp_path / 'bad.xlsx', profile=profile)
    assert not (tmp_path / 'bad.xlsx').exists()


def test_saved_repeated_formula_damage_fails_second_independent_check(tmp_path):
    _, prepared, _, _, profile = prepared_costs(tmp_path)
    values = cost_values(profile)
    target = fill_compatible_template(prepared, values, tmp_path / 'filled.xlsx', profile=profile)
    with ZipFile(target) as archive:
        parts = [(item, archive.read(item)) for item in archive.infolist()]
    with ZipFile(target, 'w') as archive:
        for item, raw in parts:
            if item.filename == PART:
                root = etree.fromstring(raw)
                root.xpath('//s:c[@r="D7"]/s:f', namespaces=NS)[0].text = 'SUM(D4:D5)'
                raw = etree.tostring(root)
            archive.writestr(item, raw)
    with pytest.raises(ValueError, match='수식'):
        verify_output(prepared, target, values, profile=profile)


def test_count_one_cache_is_distinct_and_remapping_preserves_input_policy(tmp_path):
    source, prepared, _, original, profile = prepared_costs(tmp_path, count=1)
    class Client:
        calls = []
        def generate_json(self, prompt, payload):
            self.calls.append(deepcopy(payload))
            by_id = {field['id']: field for field in profile['fields']}
            return {'fields': [{**{key: by_id[field['id']][key] for key in
                ['id', 'kind', 'required', 'input_required', 'max_chars']},
                'value_key': '확인_' + by_id[field['id']]['value_key'],
                'input_mode': by_id[field['id']].get('input_mode', 'source_grounded'), 'confidence': 0.99,
                **({'validation': by_id[field['id']]['validation']} if by_id[field['id']].get('validation') else {})}
                for field in payload['fields']]}
    cache = tmp_path / 'cache'
    ordinary = save_learned_profile(source, original, cache_dir=cache, user_confirmed=True)
    old_bytes = ordinary.read_bytes()
    learned = learn_template(prepared, client=Client(), force=True, base_profile=profile)
    validate_repeat_fields(prepared, learned, original_path=source)
    repeated = save_learned_profile(prepared, learned, cache_dir=cache, user_confirmed=True)
    assert ordinary != repeated and ordinary.read_bytes() == old_bytes
    assert load_learned_profile(source, cache).get('repeat_expansion') is None
    cached = load_learned_profile(prepared, cache, repeat_expansion=profile['repeat_expansion'])
    normalized = mapped_rule_profile(learned, {field['id']: field['value_key'] for field in learned['fields']})
    assert cached['fields'] == normalized['fields']
    validate_repeat_fields(prepared, cached, original_path=source)


@pytest.mark.parametrize('example_marker', ['(예시)', '예 시', 'Example'])
def test_new_mapping_uses_headers_instead_of_example_or_numbered_data(tmp_path, example_marker):
    book = Workbook()
    sheet = book.active
    for column, label in enumerate(['순번', '품목명', '단가(원)'], 1):
        sheet.cell(2, column, label)
    for row, first, name, price in [(3, example_marker, '예시상품QA', 1117),
                                    (4, 1, '이전기입상품QA', 2229)]:
        sheet.cell(row, 1, first)
        sheet.cell(row, 2, name)
        sheet.cell(row, 3, price)
    for column in [2, 3]:
        sheet.cell(5, column).font = Font(name='맑은 고딕', size=11)
    source, prepared = tmp_path / 'new.xlsx', tmp_path / 'prepared.xlsx'
    book.save(source)
    book.close()
    original_bytes = source.read_bytes()
    plan = {'table_id': 'xlsx:' + PART, 'row': 5, 'count': 2}
    expansion = prepare_repeat_template(source, prepared, plan)
    profile = repeat_profile(prepared, plan, expansion_binding(expansion))
    repeated = [field for field in profile['fields'] if field.get('repeat_info')]
    assert {field['label'] for field in repeated} == {'품목명', '단가(원)'}

    class Client:
        payloads = []
        def generate_json(self, prompt, payload):
            self.payloads.append(deepcopy(payload))
            by_id = {field['id']: field for field in profile['fields']}
            return {'fields': [{key: by_id[field['id']].get(key, default)
                for key, default in [('id', ''), ('value_key', ''), ('required', False),
                    ('input_required', False), ('input_mode', 'source_grounded'),
                    ('max_chars', 100), ('confidence', 0.5)]}
                for field in payload['fields']]}
    client = Client()
    learned = learn_template(prepared, client=client, force=True, base_profile=profile)
    contexts = [field['context'] for payload in client.payloads for field in payload['fields']
                if field['id'] in {item['id'] for item in repeated}]
    assert len(contexts) == 4
    assert all('B2: 품목명' in context and 'C2: 단가(원)' in context for context in contexts)
    assert all('예시상품QA' not in context and '이전기입상품QA' not in context
               and '1117' not in context and '2229' not in context for context in contexts)
    assert learned['learning']['needs_confirmation'] is True
    assert source.read_bytes() == original_bytes


def test_excel_repeat_ui_is_opt_in_and_resumes_from_original(tmp_path):
    source = cost_form(tmp_path / 'source.xlsx')
    source_bytes = source.read_bytes()
    script = 'from pathlib import Path\nimport streamlit as st\nfrom app.repeat_ui import repeat_options, restore_repeat_selection\n'
    script += f'path, profile, expansion = repeat_options(Path({str(source)!r}), Path({str(tmp_path / "data")!r}))\n'
    script += "st.json({'count': expansion['plan']['count'] if expansion else None, 'path': str(path), 'expansion': expansion})\n"
    ui = AppTest.from_string(script, default_timeout=30).run()
    assert not ui.exception and json.loads(ui.json[0].value)['count'] is None
    digest = sha256(source_bytes).hexdigest()
    ui.checkbox(key=f'repeat_enabled_{digest}').check().run()
    assert not ui.exception
    ui.selectbox(key=f'repeat_row_{digest}').set_value(4).run()
    ui.number_input(key=f'repeat_count_{digest}').set_value(3).run()
    assert not ui.exception
    record = json.loads(ui.json[0].value)
    assert record['count'] == 3
    resume = AppTest.from_string('import streamlit as st\nfrom app.repeat_ui import restore_repeat_selection\n'
        + f'expansion = {record["expansion"]!r}\n'
        + "st.json({'restored': str(restore_repeat_selection(expansion))})\n", default_timeout=30).run()
    assert not resume.exception and Path(json.loads(resume.json[0].value)['restored']) == source
    assert source.read_bytes() == source_bytes


def test_preparation_failure_keeps_prior_output(tmp_path, monkeypatch):
    source = cost_form(tmp_path / 'source.xlsx')
    original = source.read_bytes()
    target = tmp_path / 'previous.xlsx'
    target.write_bytes(b'prior output')
    def rejected(*args):
        raise ValueError('independent XLSX rejection')
    monkeypatch.setattr('agent.repeat_check.verify_repeat_expansion', rejected)
    with pytest.raises(ValueError, match='rejection'):
        prepare_repeat_template(source, target, {'table_id': 'xlsx:' + PART, 'row': 4, 'count': 3})
    assert target.read_bytes() == b'prior output' and source.read_bytes() == original
    assert set(tmp_path.iterdir()) == {source, target}


def test_pipeline_export_requires_current_repeat_proof_before_download(tmp_path, monkeypatch):
    import agent.pipeline as pipeline
    from evals.run import MockEvaluationClient
    _, prepared, expansion, _, profile = prepared_costs(tmp_path)
    text = '시험 수량 2건, 단가 1000원임'
    document = {'파일명': '합성 시험자료.txt', '본문': text, '표 목록': [],
                '페이지/시트 정보': [{'본문': text, '표 목록': [], '페이지': 1}]}
    case = {'mock_brief': {'목적': '시험 비용', '보고 대상': '팀장', '보고서 유형': '결과보고서',
                          '마감': '', '분량': '1쪽', '부족한 정보': [], '질문': []}}
    result = pipeline.run_pipeline('시험 비용 결과보고서', documents=[document], client=MockEvaluationClient(case))
    values = cost_values(profile)
    result.update(draft=values, template_profile=profile, template_expansion=expansion)
    # This test isolates final export permissions from the separately-tested content review.
    monkeypatch.setattr(pipeline, 'review_result', lambda value: {'draft': values, 'blocking': False})
    output = pipeline.build_downloads(result, confirmed=True, template_paths={'xlsx': prepared})
    assert output['xlsx'] and result['output_verification']['xlsx']['repeat_expansion']['status'] == 'passed'
    result['template_expansion']['plan']['count'] = 2
    with pytest.raises(ValueError):
        pipeline.build_downloads(result, confirmed=True, template_paths={'xlsx': prepared})
